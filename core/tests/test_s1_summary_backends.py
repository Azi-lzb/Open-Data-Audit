from __future__ import annotations

from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from base_audit.name_config import USED_RANGE_SUMMARY_FUNCTION, FeatureMapping
from base_audit.native import summary_source_reader as source_reader_module
from base_audit.native.summary_source_reader import open_summary_source


def _make_source(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["B2"] = "日期"
    sheet["C2"] = "金额"
    sheet["D2"] = "文字"
    sheet["E2"] = "标记"
    sheet["B3"] = datetime(2025, 6, 30)
    sheet["C3"] = 0
    sheet["D3"] = "正常"
    sheet["E3"] = True
    sheet["B4"] = datetime(2025, 7, 31)
    sheet["C4"] = 2.25
    sheet["D4"] = ""
    sheet["E4"] = False
    sheet["B5"] = None
    sheet["C5"] = 3
    sheet["D5"] = "尾行"
    book.save(path)
    book.close()


def _make_template(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet.append(("", "", "", "", ""))
    sheet.append((None, "日期", "金额", "文字", "标记"))
    sheet.append((None, None, None, None, None))
    book.defined_names.add(DefinedName("任意行汇总区域", attr_text="'数据'!$B$3:$E$10"))
    book.defined_names.add(DefinedName("表头区域", attr_text="'数据'!$B$2:$E$2"))
    book.save(path)
    book.close()


def _snapshot(path: Path):
    book = load_workbook(path, read_only=False, data_only=True)
    try:
        return {
            "sheetnames": book.sheetnames,
            "rows": {
                sheet.title: list(sheet.iter_rows(values_only=True))
                for sheet in book.worksheets
            },
            "freeze": {sheet.title: sheet.freeze_panes for sheet in book.worksheets},
            "filter": {sheet.title: sheet.auto_filter.ref for sheet in book.worksheets},
            "widths": _effective_widths(path, book.sheetnames, [sheet.max_column for sheet in book.worksheets]),
        }
    finally:
        book.close()


def _effective_widths(path: Path, sheetnames, max_columns):
    namespaces = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    result = {}
    with ZipFile(path) as archive:
        for index, (name, count) in enumerate(zip(sheetnames, max_columns), start=1):
            root = ET.fromstring(archive.read(f"xl/worksheets/sheet{index}.xml"))
            widths = [13.0] * count
            cols = root.find("main:cols", namespaces)
            for column in cols if cols is not None else ():
                width = float(column.get("width", 13.0))
                for number in range(int(column.get("min", 1)), int(column.get("max", 1)) + 1):
                    if number <= count:
                        widths[number - 1] = width
            result[name] = widths
    return result


def test_calamine_summary_reader_preserves_absolute_positions_and_values(tmp_path, monkeypatch):
    if source_reader_module._calamine_load is None:
        pytest.skip("python-calamine is not installed")
    source = tmp_path / "source.xlsx"
    _make_source(source)
    readings = {}
    for engine in ("OPENPYXL", "CALAMINE"):
        with open_summary_source(source, engine=engine) as book:
            sheet = book["数据"]
            readings[engine] = list(sheet.iter_rows(
                min_row=2, max_row=5, min_col=2, max_col=5, values_only=True,
            ))
            assert sheet.cell(3, 2).value == datetime(2025, 6, 30)
            assert sheet.max_row >= 5
    assert readings["CALAMINE"] == readings["OPENPYXL"]


def test_summary_reader_auto_falls_back_but_manual_calamine_fails(tmp_path, monkeypatch):
    source = tmp_path / "source.xlsx"
    _make_source(source)
    monkeypatch.setattr(source_reader_module, "_calamine_load", None)
    engines, reasons = [], []
    with open_summary_source(
        source, engine="AUTO", on_engine=engines.append, on_fallback=reasons.append,
    ) as book:
        assert book["数据"]["C3"].value == 0
    assert engines == ["OpenPyXL"]
    assert reasons and "python-calamine" in reasons[0]
    with pytest.raises(RuntimeError, match="Calamine 无法读取"):
        with open_summary_source(source, engine="CALAMINE"):
            pass


def test_s1_f02_calamine_and_rust_writer_match_openpyxl(tmp_path):
    if source_reader_module._calamine_load is None:
        pytest.skip("python-calamine is not installed")
    from base_audit.native.s1_summary_writer import writer_available
    from base_audit.native.summary import run_region_summaries

    if not writer_available():
        pytest.skip("S1-F02 Rust writer binary was not built for this platform")
    template = tmp_path / "template.xlsx"
    _make_template(template)
    source_dir = tmp_path / "sources"
    source_dir.mkdir()
    _make_source(source_dir / "机构甲_202606.xlsx")
    feature = FeatureMapping(
        "本地校验汇总", USED_RANGE_SUMMARY_FUNCTION,
        ("任意行汇总区域",), False, "", "",
    )
    outputs = {}
    logs = {}
    for label, reader, writer in (
        ("reference", "OPENPYXL", "OPENPYXL"),
        ("optimized", "CALAMINE", "RUST_XLSXWRITER"),
    ):
        log = []
        outputs[label] = run_region_summaries(
            template_path=template,
            input_dir=source_dir,
            output_dir=tmp_path / label,
            features=[feature],
            recursive=False,
            output_name="S1F02",
            source_reader=reader,
            result_writer=writer,
            on_step=log.append,
        )
        logs[label] = "\n".join(log)
    reference, optimized = _snapshot(outputs["reference"]), _snapshot(outputs["optimized"])
    assert optimized["sheetnames"] == reference["sheetnames"]
    assert optimized["rows"] == reference["rows"]
    assert optimized["freeze"] == reference["freeze"]
    assert optimized["filter"] == reference["filter"]
    assert optimized["widths"] == reference["widths"]
    assert "实际引擎=Calamine" in logs["optimized"]
    assert "实际写入引擎=s1-summary-writer" in logs["optimized"]


def test_summary_writer_auto_falls_back_when_binary_is_missing(tmp_path, monkeypatch):
    from base_audit.native import s1_summary_writer

    output = tmp_path / "result.xlsx"
    monkeypatch.setattr(s1_summary_writer, "writer_binary", lambda: None)
    calls, logs = [], []

    def openpyxl_writer():
        calls.append(True)
        output.touch()
        return output

    result = s1_summary_writer.write_summary_payload(
        {"sheets": []}, output, mode="AUTO", openpyxl_writer=openpyxl_writer,
        on_engine=logs.append,
    )
    assert result == output and output.is_file() and calls == [True]
    assert "自动回退" in logs[0]
    with pytest.raises(RuntimeError, match="Rust xlsxwriter"):
        s1_summary_writer.write_summary_payload(
            {"sheets": []}, output, mode="RUST_XLSXWRITER",
            openpyxl_writer=openpyxl_writer,
        )
