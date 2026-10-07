"""跨期/数值核对引擎测试：合成两期 CSV 验证 VBA 语义（0.01 哨兵、差异幅度、
取反、无误抑制、一边有一边无、规则错误日志）。"""

from __future__ import annotations

import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "core" / "src"))

from base_audit.systems.s3_central_statistics.cross_period_engine import (
    _build_data_store,
    _rule_frequency_ok,
    _substitute_tokens,
    run_cross_period,
)
from base_audit.systems.s3_central_statistics.models import CentralRecord
from base_audit.systems.s3_central_statistics.csv_importer import read_central_csv

CSV_HEADERS = "业务类,数据日期,机构类代码,机构类名称,地区代码,地区名称,指标顺序码,指标代码,指标名称,数据属性,币种,频度,批次,数据值"


def _write_csv(path: Path, rows: list[tuple[str, str, str]]) -> None:
    lines = [CSV_HEADERS]
    for indicator, attr, value in rows:
        lines.append(f"人民币,2026-08-31,6k0i,示例甲银行,4400000,广东省,1,{indicator},{indicator}名称,{attr},人民币,月,2,{value}")
    path.write_text("\n".join(lines), encoding="gbk")


def _record(indicator: str, attr: str, value) -> CentralRecord:
    return CentralRecord(
        biz_class="人民币", record_date="2026-08-31", org_code="6k0i", org_name="示例甲银行",
        region_code="4400000", region_name="广东省", order_code="1", indicator=indicator,
        indicator_name=indicator, data_attr=attr, currency="人民币", frequency="月",
        batch="2", value=value,
    )


def _rule(**overrides) -> dict:
    row = {
        "校验编码": "K001", "校验名称": "占比核对", "机构类代码": "", "地区代码": "",
        "前提条件": "", "left": "[33370,余额,人民币,月,2]", "right": "{33370,余额,人民币,月,2}",
        "校验公式": "0.8 * left > right", "校验提示": "信贷资产占比小于80%",
        "取反标识": "", "无误是否提示": "是", "是否逻辑校验": "", "禁用": "", "备注": "", "__行号__": "2",
    }
    row.update(overrides)
    return row


def _config(rules: list[dict]):
    class _Config:
        cross_rules = rules
    return _Config()


class TokenTests(unittest.TestCase):
    def test_five_segment_token_joins_uni_code(self):
        entry = {("33370" + "余额" + "人民币" + "月" + "2"): {"cur": 800.0, "pre": 1000.0}}
        self.assertEqual(_substitute_tokens("[33370,余额,人民币,月,2]", entry), "800.0")
        self.assertEqual(_substitute_tokens("{33370,余额,人民币,月,2}", entry), "1000.0")

    def test_missing_token_becomes_zero(self):
        self.assertEqual(_substitute_tokens("[99999,余额,人民币,月,2]", {}), "0")

    def test_frequency_check(self):
        fres = {"月2"}
        self.assertTrue(_rule_frequency_ok("[33370,余额,人民币,月,2]", fres))
        self.assertFalse(_rule_frequency_ok("[33370,余额,人民币,月,1]", fres))
        self.assertFalse(_rule_frequency_ok("[33370,余额,人民币,月]", fres))


class DataStoreTests(unittest.TestCase):
    def test_zero_becomes_sentinel(self):
        store, _fres, _o, _r = _build_data_store(
            _dataset([_record("33370", "余额", 0.0)]), None, target_unit="亿元"
        )
        value = store["6k0i4400000"]["33370余额人民币月2"]["cur"]
        self.assertAlmostEqual(value, 0.01 / 1e8)  # 0.01 元 → 亿元


def _dataset(records) -> object:
    from base_audit.systems.s3_central_statistics.models import CentralDataset

    dataset = CentralDataset(records=list(records), record_date="2026-08-31")
    dataset.build_index()
    return dataset


class RunCrossPeriodTests(unittest.TestCase):
    def test_rule_remark_carried_into_output(self):
        """配置“备注”（日月报核对/月报12批核对等比对口径）逐行带出到结果。"""
        current = _dataset([_record("33370", "余额", 600.0)])
        previous = _dataset([_record("33370", "余额", 100.0)])
        rule = _rule(备注="月报12批核对")
        rows, _logs = run_cross_period(current=current, previous=previous, config=_config([rule]))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["备注"], "月报12批核对")

    def test_output_title_places_remark_before_values(self):
        from base_audit.systems.s3_central_statistics.cross_period_engine import DEFAULT_TITLE

        self.assertIn("备注", DEFAULT_TITLE)
        self.assertLess(DEFAULT_TITLE.index("备注"), DEFAULT_TITLE.index("左值"))

    def test_ratio_rule_hits_and_suppresses(self):
        current = _dataset([_record("33370", "余额", 600.0)])
        previous = _dataset([_record("33370", "余额", 100.0)])
        # left=600, right=100：0.8*600=480 > 100 为真 → 命中并输出提示
        rows, logs = run_cross_period(current=current, previous=previous, config=_config([_rule()]))
        self.assertEqual(len(rows), 1)
        self.assertIn("信贷资产占比小于80%", rows[0]["是否说明"])
        self.assertAlmostEqual(rows[0]["差异"], 500.0)
        self.assertAlmostEqual(rows[0]["差异幅度(%)"], 83.3333333, places=4)
        # left=600, right=1000：0.8*600=480 > 1000 为假 → 无误抑制（无误是否提示=是）
        current2 = _dataset([_record("33370", "余额", 600.0)])
        previous2 = _dataset([_record("33370", "余额", 1000.0)])
        rows2, _logs = run_cross_period(current=current2, previous=previous2, config=_config([_rule()]))
        self.assertEqual(rows2, [])

    def test_difference_below_e14_treated_as_equal(self):
        # 浮点舍入残留按绝对+相对双阈值判等：5.68E-14（2^-44，300 量级典型残差，
        # 实测曾漏过 1E-14 固定阈值）必须归零，真实差异（0.001）必须保留。
        previous = _dataset([_record("33370", "余额", 300.0)])
        current = _dataset([_record("33370", "余额", 300.00000000000006)])
        rule = _rule(校验公式="", 无误是否提示="")
        rows, _logs = run_cross_period(current=current, previous=previous, config=_config([rule]))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["差异"], 0.0)
        self.assertEqual(rows[0]["差异绝对值"], 0.0)
        self.assertEqual(rows[0]["差异幅度(%)"], 0.0)
        # 阈值之上的真实差异不受相对阈值影响。
        current2 = _dataset([_record("33370", "余额", 300.001)])
        rows2, _logs2 = run_cross_period(current=current2, previous=previous, config=_config([rule]))
        self.assertAlmostEqual(rows2[0]["差异"], 0.001, places=10)

    def test_check_workbook_shows_two_decimals_without_changing_values(self):
        from base_audit.systems.s3_central_statistics.service import _write_check_workbook

        with TemporaryDirectory() as folder:
            path = Path(folder) / "核对.xlsx"
            rows = [{
                "数据日期": "2026-08-31", "机构类代码": "6k0i", "机构类名称": "示例甲银行",
                "地区代码": "4400000", "地区名称": "广东省", "校验编码": "K001",
                "校验名称": "核对", "校验类型": "数值核对", "备注": "",
                "左值": 1234.5678, "右值": 1234.5656, "差异": 0.0022,
                "差异绝对值": 0.0022, "差异幅度(%)": 0.000178, "是否说明": "", "说明内容": "", "计算公式": "",
            }]
            rows.append({
                "数据日期": "2026-08-31", "机构类代码": "6k0u", "机构类名称": "示例乙银行",
                "地区代码": "4400000", "地区名称": "广东省", "校验编码": "K002",
                "校验名称": "核对", "校验类型": "数值核对", "备注": "",
                "左值": 0.0, "右值": 100.0, "差异": 0.0,
                "差异绝对值": 0.0, "差异幅度(%)": 0.0, "是否说明": "", "说明内容": "", "计算公式": "",
            })
            _write_check_workbook(rows, path, target_unit="亿元")
            book = load_workbook(path)
            sheet = book["对比结果"]
            # 左值(10)/右值(11)/差异(12)/差异绝对值(13)/差异幅度(14) 显示两位小数，
            # 0 值不显示（三段式格式零段留空），单元格数值不四舍五入、0 仍是 0。
            for column in (10, 11, 12, 13, 14):
                self.assertEqual(sheet.cell(2, column).number_format, "0.00;-0.00;")
            self.assertEqual(sheet.cell(2, 10).value, 1234.5678)
            self.assertEqual(sheet.cell(2, 14).value, 0.000178)
            # 0 值单元格：数值保留为 0，由三段式格式负责显示为空。
            self.assertEqual(sheet.cell(3, 12).value, 0.0)
            self.assertEqual(sheet.cell(3, 13).value, 0.0)
            self.assertEqual(sheet.cell(3, 14).value, 0.0)
            book.close()

    def test_inverted_rule(self):
        current = _dataset([_record("33370", "余额", 9000.0)])
        previous = _dataset([_record("33370", "余额", 1000.0)])
        rule = _rule(取反标识="是")
        rows, _logs = run_cross_period(current=current, previous=previous, config=_config([rule]))
        # 取反：公式为真视为无误 → 不输出
        self.assertEqual(rows, [])

    def test_one_side_missing_message(self):
        current = _dataset([_record("33370", "余额", 500.0)])
        previous = _dataset([])
        rows, _logs = run_cross_period(current=current, previous=previous, config=_config([_rule()]))
        self.assertEqual(len(rows), 1)
        self.assertIn("一边有，一边无;", rows[0]["是否说明"])

    def test_formula_error_still_outputs_and_logs(self):
        current = _dataset([_record("33370", "余额", 500.0)])
        previous = _dataset([_record("33370", "余额", 1000.0)])
        rule = _rule(校验公式="0.8 * left > [33370,余额,人民币,月,2]")  # 公式内 token 不替换 → 错误
        rows, logs = run_cross_period(current=current, previous=previous, config=_config([rule]))
        self.assertEqual(len(rows), 1)
        self.assertIn("校验公式计算有误", rows[0]["是否说明"])
        self.assertTrue(any("校验公式 计算有误" in str(item) for item in logs))

    def test_premise_false_skips(self):
        current = _dataset([_record("33370", "余额", 500.0), _record("33371", "余额", 1.0)])
        previous = _dataset([_record("33370", "余额", 1000.0), _record("33371", "余额", 1.0)])
        rule = _rule(前提条件="[33371,余额,人民币,月,2] > 100")
        rows, _logs = run_cross_period(current=current, previous=previous, config=_config([rule]))
        self.assertEqual(rows, [])

    def test_both_zero_row_not_output(self):
        """输出门槛（VBA 302-310）：left=0 且 right=0 → 整条不输出。

        真实数据对拍（20260831 日月报核对查询 vs VBA《日月报比较.xls》）曾因
        缺此门槛多输出 303 行双零记录；修复后 167 行与 VBA 逐行逐值一致。
        """
        # 33399 在数据中不存在：缺失 token 求值为 0 → left=0 且 right=0。
        # （数据值为 0 的格会被零哨兵替换成 0.01，不会形成双零——与 VBA 一致。）
        current = _dataset([_record("33371", "余额", 1.0)])
        previous = _dataset([_record("33371", "余额", 1.0)])
        zero_rule = _rule(
            校验编码="K002", 校验名称="双零不输出",
            left="[33399,余额,人民币,月,2]", right="{33399,余额,人民币,月,2}",
            校验公式="", 无误是否提示="")
        hit_rule = _rule(
            校验编码="K001", 校验名称="占比核对",
            left="[33371,余额,人民币,月,2]", right="{33371,余额,人民币,月,2}",
            校验公式="left > 0", 无误是否提示="")
        rows, _logs = run_cross_period(
            current=current, previous=previous,
            config=_config([zero_rule, hit_rule]))
        codes = [row["校验编码"] for row in rows]
        self.assertNotIn("K002", codes)      # 双零整条不输出
        self.assertIn("K001", codes)         # 非双零正常输出


class CsvRoundTripTests(unittest.TestCase):
    def test_comparison_business_output_is_independent_of_log_export(self):
        from base_audit.systems.s3_central_statistics.config import ensure_split_central_configs
        from base_audit.systems.s3_central_statistics.service import run_comparison

        with TemporaryDirectory() as folder:
            root = Path(folder)
            current_csv, previous_csv = root / "本期.csv", root / "上期.csv"
            _write_csv(current_csv, [("33370", "余额", "600")])
            _write_csv(previous_csv, [("33370", "余额", "1000")])
            configs = ensure_split_central_configs(root / "config")
            config_paths = (configs["common"], configs["comparison"])
            simple = run_comparison(
                current_csv=current_csv, previous_csv=previous_csv,
                output_path=root / "比较结果_简洁.xlsx", config_path=config_paths,
                write_flow_logs=False,
            )
            detailed = run_comparison(
                current_csv=current_csv, previous_csv=previous_csv,
                output_path=root / "比较结果_导出日志.xlsx", config_path=config_paths,
                write_flow_logs=True,
            )
            self.assertEqual(simple.rows, detailed.rows)

            def output_rows(path):
                workbook = load_workbook(path, read_only=True, data_only=False)
                try:
                    sheet = workbook[workbook.sheetnames[0]]
                    return tuple(tuple(row) for row in sheet.iter_rows(values_only=True))
                finally:
                    workbook.close()

            self.assertEqual(output_rows(simple.output_path), output_rows(detailed.output_path))
            self.assertEqual(len(list(root.glob("执行比较_运行日志_*.xlsx"))), 1)

    def test_service_interface(self):
        from base_audit.systems.s3_central_statistics.service import run_cross_period_check

        with TemporaryDirectory() as folder:
            root = Path(folder)
            current_csv = root / "本期.csv"
            previous_csv = root / "上期.csv"
            _write_csv(current_csv, [("33370", "余额", "600")])
            _write_csv(previous_csv, [("33370", "余额", "1000")])
            from base_audit.systems.s3_central_statistics.config import ensure_split_central_configs
            outputs = ensure_split_central_configs(root / "config")
            config_path = outputs["cross"]
            from openpyxl import load_workbook

            book = load_workbook(config_path)
            sheet = book["跨期数值核对"]
            sheet.append(["K009", "占比核对", "", "", "", "[33370,余额,人民币,月,2]",
                          "{33370,余额,人民币,月,2}", "0.8 * left > right", "信贷资产占比小于80%",
                          "", "", "", "", "日月报核对"])
            book.save(config_path)
            book.close()

            result = run_cross_period_check(
                current_csv=current_csv, previous_csv=previous_csv,
                output_path=root / "核对结果.xlsx", config_path=(outputs["common"], config_path),
                write_flow_logs=False,
            )
            result_with_logs = run_cross_period_check(
                current_csv=current_csv, previous_csv=previous_csv,
                output_path=root / "核对结果_启用日志.xlsx", config_path=(outputs["common"], config_path),
                write_flow_logs=True,
            )
            self.assertEqual(result.rows, result_with_logs.rows)
            def output_rows(path):
                workbook = load_workbook(path, read_only=True, data_only=False)
                try:
                    return tuple(tuple(row) for row in workbook["对比结果"].iter_rows(values_only=True))
                finally:
                    workbook.close()
            self.assertEqual(output_rows(result.output_path), output_rows(result_with_logs.output_path))
            self.assertEqual(len(list(root.glob("*运行日志_*.xlsx"))), 1)
            by_code = {row["校验编码"]: row for row in result.rows}
            self.assertEqual(by_code["K009"]["备注"], "日月报核对")
            from openpyxl import load_workbook as _load

            out = _load(result.output_path, read_only=True)
            header = next(out["对比结果"].iter_rows(values_only=True))
            out.close()
            self.assertIn("备注", header)
            self.assertTrue(result.output_path.is_file())
            self.assertEqual(len(list(root.glob("*运行日志_*.xlsx"))), 1)


if __name__ == "__main__":
    unittest.main()


def test_check_cross_formulas_detection_matrix(tmp_path):
    """3.2 检查补强回归（用户报障：加逗号/坏值查不出来）：

    token 结构检查覆盖全部四列（校验公式列的 7 段此前静默通过）；
    单段=问题（恒取 0 的静默错误）；多缺陷逐项报出；作用域列正则
    编译与左中/右中检查；词表提示（只提示不算错误）经 stats 输出。
    """
    from base_audit.systems.s3_central_statistics.config import check_cross_formulas, load_central_config

    from openpyxl import Workbook
    path = tmp_path / "3.2大集中本期数值核对_配置.xlsx"
    book = Workbook()
    sheet = book.active; sheet.title = "跨期数值核对"
    sheet.append(["校验编码", "校验名称", "机构类代码", "地区代码", "前提条件",
                  "left", "right", "校验公式", "校验提示", "取反标识",
                  "无误是否提示", "是否逻辑校验", "禁用", "备注"])
    rows = [
        ("K-1", "正常五段", "", "", "", "[33370,余额,人民币,月,2]",
         "{33370,余额,人民币,月,2}", "0.8 * left > right"),
        ("K-2", "单段", "", "", "", "[33370]", "[33370]", "left > right"),
        ("K-3", "空频度被跳过", "", "", "", "[33370,余额,人民币,,]",
         "[33370,余额,人民币,,]", "left > right"),
        ("K-4", "公式语法错", "", "", "", "[33370,余额,人民币,月,2]",
         "[33370,余额,人民币,月,2]", "0.8 * left >"),
        ("K-5", "校验公式列7段", "", "", "", "[33370,余额,人民币,月,2]",
         "[33370,余额,人民币,月,2]", "[3710f,,余额,,人民币,日,1] > 5"),
        ("K-6", "6段", "", "", "", "[3710f,余额,,人民币,日,1]",
         "[33370,余额,人民币,月,2]", "left > right"),
        ("K-7", "坏正则", "^[79", "", "", "[33370,余额,人民币,月,2]",
         "[33370,余额,人民币,月,2]", "left > right"),
        ("K-8", "左中右中", "^左中^f右中", "", "", "[33370,余额,人民币,月,2]",
         "[33370,余额,人民币,月,2]", "left > right"),
        ("K-9", "双缺陷", "", "", "", "[,余额,人民币,月,]",
         "[33370,余额,人民币,月,2]", "left > right"),
        ("K-10", "词表外值", "", "", "", "[33370,净值,美元,旬,9]",
         "[33370,余额,人民币,月,2]", "left > right"),
    ]
    for rule_id, name, org, region, premise, left, right, formula in rows:
        sheet.append([rule_id, name, org, region, premise, left, right, formula,
                      "", "", "", "", "否", ""])
    book.save(path); book.close()
    stats: dict = {}
    issues = {i["规则编号"]: i
              for i in check_cross_formulas(load_central_config([path]), stats_out=stats)}
    assert not issues["K-1"]["问题"]
    assert stats["total"] == 10 and stats["enabled"] == 10
    # 单段=问题（不再是建议）：恒取 0 的静默错误。
    assert any("段数不对" in q and "大概率取不到数据" in q for q in issues["K-2"]["问题"])
    assert any("频度/批次为空" in q for q in issues["K-3"]["问题"])
    assert any("语法/求值" in q for q in issues["K-4"]["问题"])
    # 校验公式列的 token 结构检查（此前静默通过的用户场景）。
    assert any("7 段" in q and "校验公式" in q for q in issues["K-5"]["问题"])
    assert any("6 段" in q for q in issues["K-6"]["问题"])
    # 作用域列：坏正则与左中/右中转写。
    assert any("正则无法编译" in q for q in issues["K-7"]["问题"])
    assert any("左中/右中" in q for q in issues["K-8"]["问题"])
    # 同一 token 的多个缺陷逐项报出。
    assert sum(1 for q in issues["K-9"]["问题"] if "缺指标代码" in q or "缺批次" in q) == 2
    # 词表外值：只提示不算错误。
    assert not issues["K-10"]["问题"]
    assert any("净值" in v for v in stats["vocabulary"])
    assert any("旬" in v for v in stats["vocabulary"])
