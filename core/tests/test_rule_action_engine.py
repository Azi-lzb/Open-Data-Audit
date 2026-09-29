"""V3 规则动作引擎：启动过滤、指标索引、双轨等价、k表迁移格式与 2×2 接口。"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from base_audit.systems.s3_central_statistics.config import (
    ACTION_RULE_SHEET,
    CentralConfig,
    _normalize_action_rule,
    load_central_config,
)
from base_audit.systems.s3_central_statistics.indicator_rule_engine import (
    build_prepared_from_config,
    enrich_indicator_rules,
)
from base_audit.systems.s3_central_statistics.rule_action_engine import (
    filter_applicable,
    prepare_action_rules,
)
from base_audit.systems.s3_central_statistics.models import CentralDataset, CentralRecord, ComparisonRow

FORMAL_CONFIG = ROOT.parent / "config" / "3.1大集中执行比较_配置.xlsx"


def _record(indicator: str, value, *, data_attr: str = "余额", frequency: str = "月",
            batch: str = "1", currency: str = "人民币") -> CentralRecord:
    return CentralRecord(
        biz_class="人民币", record_date="2026-08-31", org_code="6k0i", org_name="测试机构",
        region_code="4400000", region_name="广东省", order_code="1", indicator=indicator,
        indicator_name=indicator, data_attr=data_attr, currency=currency,
        frequency=frequency, batch=batch, value=value,
    )


def _action_row(rule_id: str, code: str, action: str, *, source: str = "自定义",
                frequency: str = "", batch: str = "", scene: str = "",
                data_attr: str = "", currency: str = "", enabled: str = "是",
                rule_note: str = "") -> dict:
    """「规则动作」表原始行（迁移脚本列序）。"""
    return {
        "规则编号": rule_id, "规则说明": f"指标{rule_id}", "规则动作": action,
        "指标代码": code, "数据属性": data_attr, "币种": currency, "比较指标代码": "",
        "人民币阈值(亿元)": "", "美元阈值(亿美元)": "", "变幅阈值(%)": "",
        "来源分组": source, "适用表单": "", "适用频度": frequency, "适用批次": batch,
        "适用场景": scene, "详细说明": "", "启用": enabled, "停用原因": "",
        "备注": rule_note, "__行号__": "2",
    }


def test_action_output_includes_id_action_effective_threshold_and_config_note() -> None:
    actions = [
        _action_row("SP0051", "12M9N", "指标比上期增加需说明。",
                    rule_note="从境内证券业金融机构买入"),
        _action_row("SP0495", "12M9N", "指标有变动需说明。",
                    frequency="月", batch="1", rule_note="境内证券业金融机构"),
    ]
    actions[1]["人民币阈值(亿元)"] = "2"
    actions[1]["美元阈值(亿美元)"] = "0.2"
    config = CentralConfig(
        path=Path("v3.xlsx"),
        action_rules=[_normalize_action_rule(action) for action in actions],
    )
    assert [rule["配置备注"] for rule in config.action_rules] == [
        "从境内证券业金融机构买入", "境内证券业金融机构",
    ]
    record = _record("12M9N", 6.48)
    row = ComparisonRow(record=record, change=6.48)
    prepared = prepare_action_rules(config, current_date="2026-08-31", fres=FRES)
    enrich_indicator_rules(
        [row], config, current_index={record.key(): record},
        current_date="2026-08-31", previous_date="2026-07-31",
        prepared=prepared,
    )
    assert row.need_explain == (
        "SP0051-指标比上期增加需说明。-阈值0亿元-从境内证券业金融机构买入"
        "    ||    "
        "SP0495-指标有变动需说明。-阈值2亿元-境内证券业金融机构"
    )


def test_action_output_omits_empty_config_note_without_trailing_separator() -> None:
    action = _action_row("SP0476", "A", "指标有变动需说明。")
    action["人民币阈值(亿元)"] = "2"
    config = CentralConfig(path=Path("v3.xlsx"), action_rules=[_normalize_action_rule(action)])
    record = _record("A", 3.0)
    row = ComparisonRow(record=record, prev_value=0.0, change=3.0)
    enrich_indicator_rules(
        [row], config, current_index={record.key(): record},
        current_date="2026-08-31", previous_date="2026-07-31",
        prepared=prepare_action_rules(config, current_date="2026-08-31", fres=FRES),
    )
    assert row.need_explain == "SP0476-指标有变动需说明。-阈值2亿元"


def test_action_output_distinguishes_usd_and_movement_amplitude() -> None:
    amount = _action_row("SP0495", "A", "指标有变动需说明。", rule_note="美元测试")
    amount["美元阈值(亿美元)"] = "0.2"
    amount["人民币阈值(亿元)"] = "2"
    amplitude = _action_row("SP-AMP", "B", "指标有变动需说明。", rule_note="变幅测试")
    amplitude["人民币阈值(亿元)"] = "2"
    amplitude["变幅阈值(%)"] = "0.3"
    config = CentralConfig(path=Path("v3.xlsx"), action_rules=[
        _normalize_action_rule(amount), _normalize_action_rule(amplitude),
    ])
    usd = _record("A", 0.2, currency="美元合计")
    percent = _record("B", 2.5)
    rows = [
        ComparisonRow(record=usd, change=0.2),
        ComparisonRow(record=percent, change=2.5, ratio=0.5),
    ]
    enrich_indicator_rules(
        rows, config, current_index={usd.key(): usd, percent.key(): percent},
        current_date="2026-08-31", previous_date="2026-07-31",
        prepared=prepare_action_rules(config, current_date="2026-08-31", fres=FRES),
    )
    assert rows[0].need_explain == "SP0495-指标有变动需说明。-阈值0.2亿美元-美元测试"
    assert rows[1].need_explain == (
        "SP-AMP-指标有变动需说明。-阈值2亿元、变动幅度30%-变幅测试"
    )


def test_blank_usd_threshold_for_increase_displays_zero_usd() -> None:
    rule = _normalize_action_rule(_action_row(
        "SP0051", "A", "指标比上期增加需说明。", rule_note="美元零阈值",
    ))
    config = CentralConfig(path=Path("v3.xlsx"), action_rules=[rule])
    record = _record("A", 1.0, currency="美元合计")
    row = ComparisonRow(record=record, change=1.0)
    enrich_indicator_rules(
        [row], config, current_index={record.key(): record},
        current_date="2026-08-31", previous_date="2026-07-31",
        prepared=prepare_action_rules(config, current_date="2026-08-31", fres=FRES),
    )
    assert row.need_explain == "SP0051-指标比上期增加需说明。-阈值0亿美元-美元零阈值"


def test_new_indicator_shows_threshold_column_actually_used() -> None:
    action = _action_row("SP-NEW", "A", "指标新增需说明。", rule_note="新增核查")
    action["人民币阈值(亿元)"] = "1"
    action["美元阈值(亿美元)"] = "0.2"
    config = CentralConfig(path=Path("v3.xlsx"), action_rules=[_normalize_action_rule(action)])
    record = _record("A", 2.0, currency="美元合计")
    row = ComparisonRow(record=record, prev_value=0.0, change=2.0)
    enrich_indicator_rules(
        [row], config, current_index={record.key(): record},
        current_date="2026-08-31", previous_date="2026-07-31",
        prepared=prepare_action_rules(config, current_date="2026-08-31", fres=FRES),
    )
    assert row.need_explain == "SP-NEW-指标新增需说明。-阈值1亿元-新增核查"


def test_cumulative_actions_keep_each_rule_id_when_keys_overlap() -> None:
    config = CentralConfig(path=Path("v3.xlsx"), action_rules=[
        _normalize_action_rule(_action_row(
            "AC173", "A", "当年累计指标比上期不应减少。", source="当年累计",
            frequency="月", data_attr="发生额", rule_note="甲",
        )),
        _normalize_action_rule(_action_row(
            "AC467", "A", "当年累计指标比上期不应减少。", source="当年累计",
            frequency="月", data_attr="发生额", rule_note="乙",
        )),
    ])
    record = _record("A", 9.0, data_attr="发生额")
    row = ComparisonRow(record=record, prev_value=10.0, change=-1.0)
    prepared = prepare_action_rules(config, current_date="2026-08-31", fres=FRES)
    enrich_indicator_rules(
        [row], config, current_index={record.key(): record},
        current_date="2026-08-31", previous_date="2026-07-31", prepared=prepared,
    )
    assert row.need_explain.count("当年累计指标比上期不应减少。-阈值不适用") == 2
    assert row.need_explain.startswith("AC173-")
    assert "    ||    AC467-" in row.need_explain


def _legacy_rows_from_action(action_rows: list[dict]) -> list[dict]:
    """同一批规则写回旧三表形状（累计不应下降 + 特殊阈值），用于双轨等价。"""
    legacy: list[dict] = []
    for row in action_rows:
        normalized = _normalize_action_rule(dict(row))
        if normalized["类型"] == "累计不应下降":
            normalized["指标名称"] = normalized.get("指标名称", "")
        legacy.append(normalized)
    return legacy


def _dataset_with(records: list[CentralRecord]) -> CentralDataset:
    dataset = CentralDataset(records=records, record_date="2026-08-31")
    dataset.build_index()
    return dataset


# ---------------------------------------------------------------------------
# 启动过滤（VBA 1121-1143 语义）
# ---------------------------------------------------------------------------

FRES = {"", "月", "月1", "月2", "季", "季1"}


def test_filter_drops_jiezhuan_rules_outside_jan1() -> None:
    rules = _legacy_rows_from_action([
        _action_row("SP1", "A", "指标不应有数，需说明。", scene="结转"),
        _action_row("SP2", "B", "指标不应有数，需说明。"),
    ])
    kept = filter_applicable(rules, current_date="2026-08-31", fres=FRES)
    assert [rule["规则编号"] for rule in kept] == ["SP2"]


@pytest.mark.parametrize("action", [
    "指标不应有数，需说明。",
    "指标有变动需说明。",
    "指标需要对应存在。",
])
def test_all_rollover_scene_actions_run_only_on_january_first(action: str) -> None:
    rule = _normalize_action_rule(_action_row("SP1", "A", action, scene="结转"))
    assert rule["场景"] == "结转"
    assert filter_applicable([rule], current_date="2026-08-31", fres=FRES) == []
    assert filter_applicable([rule], current_date="2027-01-01", fres=FRES) == [rule]


def test_rollover_scene_also_gates_cumulative_action() -> None:
    rule = _normalize_action_rule(_action_row(
        "AC1", "A", "当年累计指标比上期不应减少。",
        source="当年累计", scene="结转", frequency="月", data_attr="发生额",
    ))
    assert rule["场景"] == "结转"
    config = CentralConfig(path=Path("v3.xlsx"), action_rules=[rule])
    august = prepare_action_rules(config, current_date="2026-08-31", fres=FRES)
    january = prepare_action_rules(config, current_date="2027-01-01", fres=FRES)
    assert august.accu_keys == set()
    assert august.stats["filtered_out"] == 1
    assert january.accu_keys == {("A", "发生额", "月")}


def test_filter_keeps_only_jiezhuan_rules_on_jan1() -> None:
    rules = _legacy_rows_from_action([
        _action_row("SP1", "A", "指标不应有数，需说明。", scene="结转"),
        _action_row("SP2", "B", "指标不应有数，需说明。"),
        _action_row("AC1", "C", "当年累计指标比上期不应减少。", source="当年累计", frequency="月"),
    ])
    kept = filter_applicable(rules, current_date="2026-01-01", fres=FRES)
    # 1 月 1 日仅保留结转规则；当年累计不受门控。
    assert [rule["规则编号"] for rule in kept] == ["SP1", "AC1"]


def test_filter_drops_rules_whose_frequency_absent_from_data() -> None:
    rules = _legacy_rows_from_action([
        _action_row("SP1", "A", "指标不应有数，需说明。", frequency="月", batch="1"),
        _action_row("SP2", "B", "指标不应有数，需说明。", frequency="日", batch="1"),
        _action_row("SP3", "C", "指标不应有数，需说明。"),
    ])
    kept = filter_applicable(rules, current_date="2026-08-31", fres=FRES)
    assert [rule["规则编号"] for rule in kept] == ["SP1", "SP3"]


def test_prepare_reports_filter_stats() -> None:
    action_rows = [
        _action_row("AC1", "A", "当年累计指标比上期不应减少。", source="当年累计", frequency="月"),
        _action_row("SP1", "B", "指标不应有数，需说明。", frequency="月", batch="1"),
        _action_row("SP2", "C", "指标不应有数，需说明。", frequency="日", batch="1"),
    ]
    config = CentralConfig(path=Path("v3.xlsx"), action_rules=[_normalize_action_rule(dict(r)) for r in action_rows])
    prepared = prepare_action_rules(config, current_date="2026-08-31", fres=FRES)
    assert prepared.stats["loaded"] == 3
    assert prepared.stats["accu"] == 1
    assert prepared.stats["special"] == 1
    assert prepared.stats["filtered_out"] == 1
    assert prepared.index is not None and set(prepared.index) == {"B"}


def test_prepare_empty_without_action_sheet() -> None:
    config = CentralConfig(path=Path("legacy.xlsx"))
    prepared = prepare_action_rules(config, current_date="2026-08-31", fres=FRES)
    assert not prepared.accu_keys and not prepared.special_rules


def test_real_config_all_rollover_scene_actions_are_date_gated() -> None:
    config = load_central_config(ROOT.parent / "config" / "3.1大集中执行比较_配置.xlsx")
    rollover = [
        row for row in config.action_rules
        if str(row.get("场景") or "").strip() == "结转"
    ]
    assert rollover  # 检查真实工作簿的整组规则，而非只检查 12MRK。
    assert filter_applicable(rollover, current_date="2026-08-31", fres=None) == []
    assert filter_applicable(rollover, current_date="2027-01-01", fres=None) == rollover
    august = prepare_action_rules(config, current_date="2026-08-31", fres=None)
    assert not any(rule.get("场景") == "结转" for rule in august.special_rules)


# ---------------------------------------------------------------------------
# 双轨等价：legacy 三表装载 vs V3 规则动作装载，动作结果必须一致
# ---------------------------------------------------------------------------

def _run_both(action_rows: list[dict], records: list[CentralRecord],
              *, current_date="2026-08-31", previous_date="2026-07-31",
              change: float = 0.0, ratio: float = 0.0) -> tuple[str, str]:
    dataset = _dataset_with(records)
    row = ComparisonRow(
        record=records[0],
        prev_value=records[1].value if len(records) > 1 else None,
        change=change, ratio=ratio,
    )
    legacy_config = CentralConfig(path=Path("legacy.xlsx"), rules=_legacy_rows_from_action(action_rows))
    enrich_indicator_rules(
        [row], legacy_config, current_index=dataset.key_index,
        current_date=current_date, previous_date=previous_date,
    )
    legacy_explain = row.need_explain

    row2 = ComparisonRow(
        record=records[0],
        prev_value=records[1].value if len(records) > 1 else None,
        change=change, ratio=ratio,
    )
    v3_config = CentralConfig(
        path=Path("v3.xlsx"),
        action_rules=[_normalize_action_rule(dict(r)) for r in action_rows],
    )
    prepared = prepare_action_rules(v3_config, current_date=current_date, fres=FRES)
    enrich_indicator_rules(
        [row2], v3_config, current_index=dataset.key_index,
        current_date=current_date, previous_date=previous_date, prepared=prepared,
    )
    return legacy_explain, row2.need_explain


def test_dual_track_negative_balance_action() -> None:
    action_rows = [_action_row("SP1", "A", "指标为负数需核实。")]
    records = [_record("A", -5.0), _record("A_prev", 1.0)]
    legacy, v3 = _run_both(action_rows, records)
    assert legacy == v3 == "SP1-指标为负数需核实。-阈值不适用"


def test_dual_track_accumulated_rule_january_message() -> None:
    action_rows = [
        _action_row("AC1", "A", "当年累计指标比上期不应减少。", source="当年累计",
                    frequency="月", data_attr="发生额"),
    ]
    records = [
        _record("A", 10.0, data_attr="发生额"),
        _record("A", 20.0, data_attr="发生额"),
    ]
    legacy, v3 = _run_both(action_rows, records, current_date="2026-08-31", change=-10.0)
    assert legacy == v3
    assert "当年累计指标比上期不应减少" in legacy


def test_dual_track_corresponding_exists_action() -> None:
    action_rows = [_action_row("SP1", "A", "指标需要对应存在。")]
    records = [_record("A", 10.0), _record("B", 0.0)]
    legacy, v3 = _run_both(action_rows, records)
    assert legacy == v3
    assert "对应指标无数" in legacy


def test_dual_track_disabled_rule_never_fires() -> None:
    action_rows = [_action_row("SP1", "A", "指标为负数需核实。", enabled="否")]
    records = [_record("A", -5.0)]
    legacy, v3 = _run_both(action_rows, records)
    assert legacy == "" == v3


# ---------------------------------------------------------------------------
# 正式配置装载
# ---------------------------------------------------------------------------

def test_formal_config_loads_action_rules() -> None:
    assert FORMAL_CONFIG.is_file(), f"缺少正式 3.1 配置：{FORMAL_CONFIG}"
    config = load_central_config(FORMAL_CONFIG)
    assert len(config.action_rules) == 5082
    accum = [row for row in config.action_rules if row.get("类型") == "累计不应下降"]
    assert len(accum) == 521
    assert all(row.get("类型") in ("累计不应下降", "特殊阈值") for row in config.action_rules)
    # 正式配置中的动作文本贯通到归一化字段（特殊阈值行动作文案存于“备注”键）。
    sample = next(row for row in config.action_rules
                  if row.get("类型") == "特殊阈值" and row.get("备注"))
    assert str(sample["备注"]).endswith("。")


# ---------------------------------------------------------------------------
# k表迁移格式：8 段 token、机构正则折入、表达式可编译
# ---------------------------------------------------------------------------

def test_expression_sheet_codes_match_groups() -> None:
    """正式配置契约：跨期核对行编号编入 CX跨 段；V2 归档与三列已移除。"""
    import re
    from openpyxl import load_workbook

    config_path = ROOT.parent / "config" / "3.1大集中执行比较_配置.xlsx"
    book = load_workbook(config_path, read_only=True, data_only=True)
    try:
        assert "表达式校验V2" not in book.sheetnames
        # 正式配置已双轨：旧 8 段表在「表达式校验（兼容）」（三列移除断言针对它）。
        sheet = book["表达式校验（兼容）"]
        headers = [str(c.value or "").strip() for c in next(sheet.iter_rows(max_row=1))]
        for name in ("适用频度", "适用批次", "适用场景"):
            assert name not in headers, f"{name} 信息列应已移除"
        group_column = headers.index("来源分组") + 1
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            group = str(row[group_column - 1] or "").strip()
            code = str(row[0] or "").strip()
            if group == "跨期核对":
                assert re.fullmatch(r"CX跨\d{4}", code), f"跨期核对编号应为 CX跨XXXX：{code}"
    finally:
        book.close()


# ---------------------------------------------------------------------------
# 表达式 2×2 接口
# ---------------------------------------------------------------------------

def test_expression_settings_default_sbe_python() -> None:
    from base_audit.systems.s3_central_statistics.expression_backend import (
        BACKEND_PYTHON, MODE_SBE, resolve_expression_settings,
    )

    assert resolve_expression_settings({}) == (MODE_SBE, BACKEND_PYTHON)


def test_expression_settings_reject_unknown() -> None:
    from base_audit.systems.s3_central_statistics.expression_backend import resolve_expression_settings

    with pytest.raises(Exception):
        resolve_expression_settings({"表达式处理方式": "MAGIC"})


def test_expression_office_backend_fails_loudly() -> None:
    from base_audit.systems.s3_central_statistics.expression_backend import (
        BACKEND_OFFICE, ExpressionBackendError, MODE_SBE, evaluate_expression_rule,
    )

    with pytest.raises(ExpressionBackendError):
        evaluate_expression_rule("1>0", trigger_inverted=False,
                                 mode=MODE_SBE, backend=BACKEND_OFFICE)


def test_expression_lae_python_short_circuit_no_division() -> None:
    from base_audit.systems.s3_central_statistics.expression_backend import (
        BACKEND_PYTHON, MODE_LAE, evaluate_expression_rule,
    )
    from base_audit.systems.s3_central_statistics.expression_lae import ExpressionContext

    # 上期索引为空：{...} 缺失按 0 处理 → IF 命中 TRUE 分支；
    # 惰性求值下除法分支不得求值（无除零/求值错误）。
    context = ExpressionContext(
        previous_index={},
        current_index={},
        row=ComparisonRow(record=_record("A", 5.0)),
    )
    result = evaluate_expression_rule(
        "IF({,,,A,余额,人民币,月,1}=0, TRUE, "
        "[,,,A,余额,人民币,月,1]/{,,,A,余额,人民币,月,1}<1.3)",
        trigger_inverted=False, mode=MODE_LAE, backend=BACKEND_PYTHON,
        context=context,
    )
    assert result.hit is True
    assert result.status == "OK"


def test_expression_ast_cache_reuses_parse() -> None:
    from base_audit.systems.s3_central_statistics import expression_lae

    expression_lae._AST_CACHE.clear()
    first = expression_lae.parse_expression("ABS(1-2)>0")
    assert expression_lae._AST_CACHE.get("ABS(1-2)>0") is first
    assert expression_lae.parse_expression("ABS(1-2)>0") is first
