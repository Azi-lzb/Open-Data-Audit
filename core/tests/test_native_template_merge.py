"""UOS/麒麟 联合模板制作（openpyxl 版）测试。

覆盖 COM 版同等的业务契约：基准模板权威、同名工作表跳过、命名区域改写为
``<原名>_<工作表名>``、跨簿引用本地化、产物可供后续审核使用。
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from base_audit.native.template_merge import run_template_merge_native  # noqa: E402


def _make_base(path: Path) -> None:
    """基准模板：含外部依赖表（集中系统数据/参照表）与自己的业务表。"""
    book = Workbook()
    sheet = book.active
    sheet.title = "单位贷款信息"
    sheet["A1"], sheet["B1"] = "指标", "数值"
    sheet["A2"], sheet["B2"] = "余额", 100
    # 依赖外部工作表：
    dep = book.create_sheet("集中系统数据")
    dep["A1"] = "机构"
    dep["B1"] = 1
    ref = book.create_sheet("参照表")
    ref["A1"], ref["B1"] = "代码", "名称"
    # 跨簿公式（引用其它模板的工作表，联合后应本地化）
    sheet["C2"] = "='[其他模板.xlsx]存量债券投资'!B2*2"
    book.defined_names.add(DefinedName("校验区域_001", attr_text="'单位贷款信息'!$A$2:$B$2"))
    book.defined_names.add(DefinedName("表结构区域_001", attr_text="'单位贷款信息'!$A$1:$B$1"))
    book.save(path)
    book.close()


def _make_source(path: Path) -> None:
    """来源模板：一张新表 + 一张与基准同名表（应跳过）+ 同名命名区域。"""
    book = Workbook()
    other = book.active
    other.title = "存量债券投资"
    other["A1"], other["B1"] = "债券", "面值"
    other["A2"], other["B2"] = "国债", 500
    other["C2"] = "=B2*1.05"
    other["A3"] = "带批注"
    from openpyxl.comments import Comment

    other["A3"].comment = Comment("这是说明", "报送系统")
    other.conditional_formatting.add(
        "B2:B2", CellIsRule(operator="greaterThan", formula=["100"],
                            fill=PatternFill("solid", fgColor="FF0000"))
    )
    # 与基准同名的命名区域 → 迁移后须带本表后缀，不能覆盖基准
    book.defined_names.add(DefinedName("校验区域_001", attr_text="'存量债券投资'!$A$2:$B$2"))
    book.defined_names.add(DefinedName("表结构区域_001", attr_text="'存量债券投资'!$A$1:$B$1"))
    # 同名工作表：应被跳过，基准版本保留
    dup = book.create_sheet("单位贷款信息")
    dup["A1"] = "这是来源里的同名表，不应覆盖基准"
    book.save(path)
    book.close()


@pytest.fixture()
def templates(tmp_path: Path):
    base = tmp_path / "！10.基础数据-单位贷款202609.xlsx"
    source = tmp_path / "！70.债券业务202608.xlsx"
    _make_base(base)
    _make_source(source)
    return base, source


def test_base_template_is_authoritative(templates) -> None:
    """基准模板的工作表、依赖表与跨表引用必须保持权威。"""
    base, source = templates
    result = run_template_merge_native(base_template=base, source_templates=[source])
    book = load_workbook(result.output_path)
    try:
        # 基准三张表（含外部依赖表）全部保留
        for title in ("单位贷款信息", "集中系统数据", "参照表"):
            assert title in book.sheetnames
        # 同名表跳过：基准内容未被来源覆盖
        assert book["单位贷款信息"]["A2"].value == "余额"
        assert book["单位贷款信息"]["A1"].value == "指标"
        # 来源的新表已复制，含批注与条件格式
        assert "存量债券投资" in book.sheetnames
        copied = book["存量债券投资"]
        assert copied["B2"].value == 500
        assert copied["A3"].comment is not None
        assert sum(1 for _ in copied.conditional_formatting) >= 1
        # 跳过的同名表进入报告
        assert (source.name, "单位贷款信息") in result.skipped_sheets
    finally:
        book.close()


def test_named_ranges_get_sheet_suffix_and_no_collision(templates) -> None:
    """同名命名区域不得互相覆盖：一律改写为 <原名>_<工作表名>。"""
    base, source = templates
    result = run_template_merge_native(base_template=base, source_templates=[source])
    book = load_workbook(result.output_path)
    try:
        names = set(str(key) for key in book.defined_names)
        # 两个模板都有 校验区域_001 / 表结构区域_001 → 各自带工作表后缀
        assert "校验区域_001_单位贷款信息" in names
        assert "校验区域_001_存量债券投资" in names
        assert "表结构区域_001_单位贷款信息" in names
        assert "表结构区域_001_存量债券投资" in names
        # 原名必须清除，避免歧义
        assert "校验区域_001" not in names
        assert "表结构区域_001" not in names
        # 引用位置指向正确的工作表
        assert "单位贷款信息" in book.defined_names["校验区域_001_单位贷款信息"].attr_text
        assert "存量债券投资" in book.defined_names["校验区域_001_存量债券投资"].attr_text
        # 报告记录迁移条数（基准 2 + 来源 2）
        assert sum(item[5] == "已复制" for item in result.copied_named_ranges) >= 4
    finally:
        book.close()


def test_cross_workbook_reference_is_localized(templates) -> None:
    """跨簿引用改写为本地工作表引用；被引用的表不在产物中则暴露为 #REF!。"""
    base, source = templates
    result = run_template_merge_native(base_template=base, source_templates=[source])
    book = load_workbook(result.output_path)
    try:
        formula = book["单位贷款信息"]["C2"].value
        # 原来的 [其他模板.xlsx] 前缀必须消失
        assert "[" not in formula
        assert "存量债券投资" in formula
    finally:
        book.close()


def test_output_never_modifies_originals(templates) -> None:
    """原始基准与来源模板均不被修改；产物在基准同级的“联合模板”目录。"""
    import hashlib

    def sha(path: Path) -> str:
        return hashlib.sha256(path.read_bytes()).hexdigest()

    base, source = templates
    before = (sha(base), sha(source))
    result = run_template_merge_native(base_template=base, source_templates=[source])
    assert (sha(base), sha(source)) == before
    assert result.output_path.parent.name == "联合模板"
    assert result.output_path.is_file()
    assert result.report_path.is_file()


def test_report_lists_sheets_and_names(templates) -> None:
    """检查报告须含合并说明、工作表处理清单、命名区域处理清单、公式检查。"""
    base, source = templates
    result = run_template_merge_native(base_template=base, source_templates=[source])
    report = load_workbook(result.report_path)
    try:
        assert {"合并说明", "工作表处理清单", "命名区域处理清单", "公式检查"} <= set(report.sheetnames)
        copied_rows = [
            row for row in report["工作表处理清单"].iter_rows(values_only=True)
        ]
        assert any(row[2] == "已复制" for row in copied_rows[1:])
        assert any(row[2] == "保留为底稿" for row in copied_rows[1:])
    finally:
        report.close()


def test_rejects_xls_and_empty_sources(tmp_path: Path) -> None:
    """openpyxl 版不支持 .xls；来源为空要明确报错。"""
    base = tmp_path / "基准.xlsx"
    _make_base(base)
    legacy = tmp_path / "旧版.xls"
    legacy.write_bytes(b"not a real xls")
    with pytest.raises(ValueError):
        run_template_merge_native(base_template=base, source_templates=[legacy])
    with pytest.raises(ValueError):
        run_template_merge_native(base_template=base, source_templates=[])
