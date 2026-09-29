from __future__ import annotations

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from base_audit.name_config import (
    FORMULA_COPY_FUNCTION,
    NAMED_RANGE_CHECK_FUNCTION,
    STRUCTURE_COMPARE_FUNCTION,
    FeatureMapping,
)


def test_template_snapshot_loads_workbook_once(monkeypatch, tmp_path):
    from base_audit import template_snapshot

    path = tmp_path / "模板.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["A1"] = "指标"
    sheet["B2"] = '=IF(A1="指标","错|指标|说明|1|0|1","")'
    book.defined_names.add(DefinedName("表结构区域", attr_text="'数据'!$A$1"))
    book.defined_names.add(DefinedName("校验区域", attr_text="'数据'!$B$2"))
    book.save(path)
    book.close()

    mappings = (
        FeatureMapping("检查校验区域", NAMED_RANGE_CHECK_FUNCTION, ("校验区域",), False, "", ""),
        FeatureMapping("表结构比对", STRUCTURE_COMPARE_FUNCTION, ("表结构区域",), False, "", ""),
        FeatureMapping("公式校验复制", FORMULA_COPY_FUNCTION, ("校验区域",), False, "", ""),
    )
    real_load = template_snapshot.load_workbook
    calls = []

    def counted_load(*args, **kwargs):
        calls.append(args[0])
        return real_load(*args, **kwargs)

    monkeypatch.setattr(template_snapshot, "load_workbook", counted_load)
    snapshot = template_snapshot.load_template_snapshot(path, mappings)

    assert calls == [path]
    assert snapshot.definition.copy_ranges[0].address == "B2"
    assert snapshot.require_feature_ranges("检查校验区域")[0].address == "B2"
    assert snapshot.structure_values[0][1] == (("指标",),)
    assert snapshot.formula_overrides[("数据", "B2")].startswith("=IF(")
    assert snapshot.formula_cells[("数据", "B2")].startswith("=IF(")
    assert snapshot.formulas == (sheet["B2"].value,)
