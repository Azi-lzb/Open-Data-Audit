from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from base_audit.history import HISTORY_HEADERS
from base_audit.models import Issue
from base_audit.native.result_writer import write_audit_summary


def _issue(source: Path, audit: Path) -> Issue:
    return Issue(
        issue_id="甲银行｜数据表｜C6｜单位贷款", period="2026-09", batch_id="batch",
        audit_time="2026-09-11 12:00:00", triggered=True, status="新增",
        first_seen_period="2026-09", previous_seen_period="", consecutive_count=1,
        org_code="001", org_name="甲银行", report_code="单位贷款", sheet_name="数据表",
        rule_id="R001", severity="错误", formula_cell="C6", target_cell="C6",
        target_value=0, formula_result="错|单位贷款|不应该为0|0|0|0",
        message="不应该为0", source_file=str(source), audit_file=str(audit),
        check_field="单位贷款", check_type="逾期核对", comparison_value=0, difference_value=0,
        detail="不应该为0", institution_feedback="历史说明", auditor_opinion="复核",
    )


def test_openpyxl_result_writer_preserves_summary_contract(tmp_path):
    source = tmp_path / "甲银行.xlsx"
    audit = tmp_path / "甲银行_审核版.xlsx"
    for path in (source, audit):
        book = Workbook()
        book.active.title = "数据表"
        book.save(path)
        book.close()
    result = tmp_path / "结果.xlsx"

    write_audit_summary(result, [_issue(source, audit)], mode="OPENPYXL")

    book = load_workbook(result)
    try:
        sheet = book["本期审核结果"]
        assert tuple(cell.value for cell in sheet[1]) == HISTORY_HEADERS
        assert tuple(cell.value for cell in sheet[2]) == (
            "甲银行.xlsx", "数据表", "C6", "错误", "逾期核对", "单位贷款", "不应该为0",
            "0", "0", "0", "甲银行｜数据表｜C6｜单位贷款", "历史说明", "复核",
        )
        assert sheet["C2"].hyperlink.target == str(audit.resolve())
        assert sheet["C2"].hyperlink.location == "'数据表'!C6"
        assert sheet.freeze_panes == "A2"
    finally:
        book.close()


def test_rust_result_writer_preserves_summary_values_and_hyperlinks(tmp_path):
    from base_audit.native.s1_summary_writer import writer_available

    if not writer_available():
        pytest.skip("S1 Rust result writer binary is not built for this platform")
    source = tmp_path / "甲银行.xlsx"
    audit = tmp_path / "甲银行 审核版.xlsx"
    for path in (source, audit):
        book = Workbook()
        book.active.title = "数据表"
        book.save(path)
        book.close()
    issues = [_issue(source, audit)]
    reference = tmp_path / "结果-openpyxl.xlsx"
    optimized = tmp_path / "结果-rust.xlsx"

    from base_audit.native.result_writer import write_audit_summary
    write_audit_summary(reference, issues, mode="OPENPYXL")
    write_audit_summary(optimized, issues, mode="RUST_XLSXWRITER")

    ref_book = load_workbook(reference)
    rust_book = load_workbook(optimized)
    try:
        ref_sheet = ref_book["本期审核结果"]
        rust_sheet = rust_book["本期审核结果"]
        assert list(ref_sheet.values) == list(rust_sheet.values)
        assert rust_sheet.freeze_panes == ref_sheet.freeze_panes == "A2"
        assert rust_sheet.auto_filter.ref == ref_sheet.auto_filter.ref == "A1:M2"
        def effective_widths(sheet):
            widths = [13.0] * 13
            for dimension in sheet.column_dimensions.values():
                if dimension.min is not None and dimension.max is not None:
                    for index in range(dimension.min, min(dimension.max, 13) + 1):
                        widths[index - 1] = dimension.width
            return widths
        assert effective_widths(rust_sheet) == effective_widths(ref_sheet)
        for column in range(1, 14):
            assert rust_sheet.cell(1, column).font.bold is True
            assert rust_sheet.cell(1, column).fill.fgColor.rgb[-6:] == ref_sheet.cell(1, column).fill.fgColor.rgb[-6:]
            assert rust_sheet.cell(2, column).number_format == ref_sheet.cell(2, column).number_format
        ref_link = ref_sheet["C2"].hyperlink
        rust_link = rust_sheet["C2"].hyperlink
        assert ref_link is not None and rust_link is not None
        assert rust_link.location == ref_link.location == "'数据表'!C6"
        assert rust_sheet["C2"].value == ref_sheet["C2"].value == "C6"
        assert rust_link.target == ref_link.target == str(audit.resolve())
        assert rust_sheet["C2"].font.underline == ref_sheet["C2"].font.underline
        assert ref_sheet["C2"].font.color.theme == 10
        assert rust_sheet["C2"].font.color.rgb == "FF0000FF"
    finally:
        ref_book.close()
        rust_book.close()


def test_audit_summary_writer_auto_falls_back_but_manual_rust_fails(tmp_path, monkeypatch):
    from base_audit.native import s1_summary_writer

    output = tmp_path / "audit-summary.xlsx"
    monkeypatch.setattr(s1_summary_writer, "writer_binary", lambda: None)
    calls, logs = [], []

    def openpyxl_writer():
        calls.append(True)
        output.touch()
        return output

    result = s1_summary_writer.write_audit_summary_payload(
        {"sheet_name": "结果", "headers": [], "rows": [], "hyperlinks": []},
        output,
        mode="AUTO",
        openpyxl_writer=openpyxl_writer,
        on_engine=logs.append,
    )
    assert result == output and output.is_file() and calls == [True]
    assert "自动回退" in logs[0]
    with pytest.raises(RuntimeError, match="Rust xlsxwriter"):
        s1_summary_writer.write_audit_summary_payload(
            {"sheet_name": "结果", "headers": [], "rows": [], "hyperlinks": []},
            output,
            mode="RUST_XLSXWRITER",
            openpyxl_writer=openpyxl_writer,
        )
