import unittest

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from base_audit.systems.s3_central_statistics.expression_parser import (
    ExpressionError,
    evaluate_boolean,
    evaluate_expression,
)


class ArithmeticTests(unittest.TestCase):
    def test_basic_arithmetic_and_precedence(self):
        self.assertEqual(evaluate_expression("1+2*3"), 7.0)
        self.assertEqual(evaluate_expression("(1+2)*3"), 9.0)
        self.assertEqual(evaluate_expression("10/4"), 2.5)

    def test_excel_power_semantics(self):
        # Excel：负号先于乘方作用于底数；乘方左结合；指数可带符号。
        self.assertEqual(evaluate_expression("-2^2"), 4.0)
        self.assertEqual(evaluate_expression("2^3^2"), 64.0)
        self.assertEqual(evaluate_expression("2^-2"), 0.25)
        self.assertEqual(evaluate_expression("(-2)^2"), 4.0)

    def test_scientific_literal(self):
        self.assertEqual(evaluate_expression("1.5E+3"), 1500.0)
        self.assertEqual(evaluate_expression("2e-1"), 0.2)


class ComparisonTests(unittest.TestCase):
    def test_comparison_returns_bool(self):
        self.assertIs(evaluate_expression("1<2"), True)
        self.assertIs(evaluate_expression("1<>2"), True)
        self.assertIs(evaluate_expression("2<=2"), True)
        self.assertIs(evaluate_expression("1=2"), False)

    def test_chained_comparison_is_left_assoc(self):
        # Excel: (1<2)<1 → TRUE(1) < 1 → FALSE
        self.assertIs(evaluate_expression("1<2<1"), False)

    def test_bool_in_arithmetic(self):
        self.assertEqual(evaluate_expression("(1<2)+1"), 2.0)


class FunctionTests(unittest.TestCase):
    def test_and_or_not_case_insensitive(self):
        self.assertIs(evaluate_boolean("AND(1,2)"), True)
        self.assertIs(evaluate_boolean("and(1,0)"), False)
        self.assertIs(evaluate_boolean("OR(0,3)"), True)
        self.assertIs(evaluate_boolean("NOT(0)"), True)
        self.assertIs(evaluate_boolean("not(5)"), False)

    def test_abs_max_min_sum(self):
        self.assertEqual(evaluate_expression("ABS(0-3.5)"), 3.5)
        self.assertEqual(evaluate_expression("MAX(1,9,4)"), 9.0)
        self.assertEqual(evaluate_expression("MIN(1,9,4)"), 1.0)
        self.assertEqual(evaluate_expression("SUM(1,2,3)"), 6.0)
        self.assertEqual(evaluate_expression("SUM(1+1,MAX(2,5))"), 7.0)

    def test_variables_left_right_thd(self):
        variables = {"left": 8.0, "right": 2.0, "Thd": 0.5}
        self.assertEqual(evaluate_expression("left-right", variables), 6.0)
        self.assertIs(evaluate_boolean("0.5*left>right", variables), True)
        self.assertIs(evaluate_boolean("ABS(left-right)<Thd", variables), False)

    def test_true_false_literals(self):
        self.assertEqual(evaluate_expression("TRUE+1"), 2.0)
        self.assertEqual(evaluate_expression("FALSE"), 0.0)
        self.assertIs(evaluate_boolean("FALSE"), False)


class ErrorTests(unittest.TestCase):
    def test_division_by_zero_raises(self):
        with self.assertRaises(ExpressionError) as ctx:
            evaluate_expression("1/0")
        self.assertIn("除零", str(ctx.exception))

    def test_unknown_name_raises(self):
        with self.assertRaises(ExpressionError) as ctx:
            evaluate_expression("foo+1")
        self.assertIn("未知名称", str(ctx.exception))

    def test_unknown_char_raises(self):
        with self.assertRaises(ExpressionError):
            evaluate_expression("1 & 2")

    def test_unbalanced_paren_raises(self):
        with self.assertRaises(ExpressionError):
            evaluate_expression("(1+2")

    def test_trailing_garbage_raises(self):
        with self.assertRaises(ExpressionError):
            evaluate_expression("1+2 3")

    def test_empty_raises(self):
        with self.assertRaises(ExpressionError):
            evaluate_expression("   ")

    def test_abs_wrong_arity_raises(self):
        with self.assertRaises(ExpressionError):
            evaluate_expression("ABS(1,2)")


if __name__ == "__main__":
    unittest.main()


def test_expression_completion_suggestion_fills_to_eight_segments() -> None:
    """3.1 补全建议：不足 8 段的 token 给出补全形态（空段=继承，语义不变）。

    3.1 编译存活检查要求恰好 8 段：尾缺 5 段（[,,,20203045,]）与单段
    （[20203045]，code 会落「业务类」位）都编译失败且执行时无法挂接；
    补全为 8 段后语义与原意图一致。完整 8 段（含指定值/^）不做建议。
    """
    from base_audit.systems.s3_central_statistics.config import _expression_completion_suggestion
    assert _expression_completion_suggestion("[,,,20203045,]") == "[,,,20203045,,,,]"
    assert _expression_completion_suggestion("{,,,20203045,}") == "{,,,20203045,,,,}"
    assert _expression_completion_suggestion("[20203045]") == "[,,,20203045,,,,]"
    assert _expression_completion_suggestion("[人民币,余额,人民币,20203045,月]") == (
        "[人民币,余额,人民币,20203045,月,,,]")
    assert _expression_completion_suggestion("[本外币,,,20203045,余额,人民币,日,1]") is None  # 完整 8 段
    assert _expression_completion_suggestion("[,,,,,,]") is None            # 补不出指标段
    assert _expression_completion_suggestion("[余额,人民币,日,1]") is None   # 无指标样段


def test_check_expression_rules_reports_suggestions(tmp_path):
    from base_audit.systems.s3_central_statistics.config import check_expression_rules, load_central_config

    from openpyxl import Workbook
    path = tmp_path / "3.1大集中执行比较_配置.xlsx"
    book = Workbook()
    sheet = book.active; sheet.title = "表达式校验（兼容）"
    sheet.append(["规则编号", "来源分组", "来源表单", "规则说明", "校验表达式", "触发方式", "容差值(万元)", "启用", "停用原因", "备注"])
    sheet.append(["R-OLD", "自定义", "A0000", "八段规则", "[,,,20203045,] < 0.055 * [,,,20203048,]", "表达式成立", "", "是", "", ""])
    sheet.append(["R-NEW", "自定义", "A0000", "单段规则", "[20203045] > 0", "表达式成立", "", "是", "", ""])
    sheet.append(["R-FIX", "自定义", "A0000", "指定币种", "[人民币,,^44,20203045,,,,]<>0", "表达式成立", "", "是", "", ""])
    book.save(path); book.close()
    config = load_central_config([path])
    issues = check_expression_rules(config)
    by_id = {i["规则编号"]: i for i in issues}
    # 尾缺 5 段：建议补全 8 段（同时会有编译问题——建议供用户修正后通过）。
    assert any("[,,,20203045,] → [,,,20203045,,,,]" in s for s in by_id["R-OLD"]["建议"])
    assert not by_id["R-FIX"]["建议"]      # 完整 8 段（含指定值）：正式形态，无建议
    # 单段：编译失败（code 落业务类位），建议补全且指标落第 4 段。
    assert any("问题" in " ".join(by_id["R-NEW"]["问题"]) or by_id["R-NEW"]["问题"])
    assert any("[20203045] → [,,,20203045,,,,]" in s for s in by_id["R-NEW"]["建议"])


def test_expression_overlong_token_suggests_removing_empty_segment() -> None:
    """超长 9 段 token（真实 CX自0002 形态）：建议删除多余空段（唯一解才给）。"""
    from base_audit.systems.s3_central_statistics.config import (
        _expression_overlong_suggestion, check_expression_rules, load_central_config)

    assert _expression_overlong_suggestion("[本外币,,,,37103,余额,人民币,日,1]") == (
        "[本外币,,,37103,余额,人民币,日,1]")
    # 多个不同合法解 → 不猜，转人工。
    assert _expression_overlong_suggestion("[,,,37103,余额,人民币,日,1,]") is None

    from openpyxl import Workbook
    path = (Path(__file__).resolve().parent.parent / "..")
    with tempfile.TemporaryDirectory() as folder:
        cfg = Path(folder) / "31.xlsx"
        book = Workbook(); sheet = book.active; sheet.title = "表达式校验（兼容）"
        sheet.append(["规则编号", "来源分组", "来源表单", "规则说明", "校验表达式",
                      "触发方式", "容差值(万元)", "启用", "停用原因", "备注"])
        sheet.append(["CX-T", "自定义", "A3702", "九段",
                      "[本外币,,,,37103,余额,人民币,日,1] > 0", "表达式成立", "", "是", "", ""])
        book.save(cfg); book.close()
        issues = {i["规则编号"]: i for i in check_expression_rules(load_central_config([cfg]))}
        assert any("[本外币,,,,37103,余额,人民币,日,1] → [本外币,,,37103,余额,人民币,日,1]" in s
                   for s in issues["CX-T"]["建议"])


def test_check_expression_rules_five_mode_stats_and_vocabulary(tmp_path):
    """5 段式体检：stats 输出总数/启用；词表提示（只提示不算错误）独立于问题。

    指标段含中文逗号 → 10 段 → TOKEN_SEGMENT_COUNT（既有）；数据属性写
    「净值」（不在 余额/发生额 词表）→ 词表提示而非错误。
    """
    from base_audit.systems.s3_central_statistics.config import (
        check_expression_rules, load_central_config,
    )
    from openpyxl import Workbook

    path = tmp_path / "3.1大集中执行比较_配置.xlsx"
    book = Workbook()
    compat = book.active; compat.title = "表达式校验（兼容）"
    compat.append(["规则编号", "来源分组", "来源表单", "规则说明", "校验表达式",
                   "触发方式", "容差值(万元)", "启用", "停用原因", "备注"])
    compat.append(["R5-A", "自定义", "A0000", "属性不在词表",
                   "[人民币,,,12N0D,净值,人民币,月,1] > 0", "表达式成立", "", "是", "", ""])
    compat.append(["R5-B", "自定义", "A0000", "正常",
                   "[人民币,,,12N0D,余额,人民币,月,1] > 0", "表达式成立", "", "是", "", ""])
    five = book.create_sheet("表达式校验")
    five.append(["规则编号", "来源分组", "来源表单", "机构类代码", "地区代码", "规则说明",
                 "校验表达式", "触发方式", "容差值(万元)", "启用", "停用原因", "备注"])
    five.append(["R5-A", "自定义", "A0000", "", "", "属性不在词表",
                 "[12N0D,净值,人民币,月,1] > 0", "表达式成立", "", "是", "", ""])
    five.append(["R5-B", "自定义", "A0000", "", "", "正常",
                 "[12N0D,余额,人民币,月,1] > 0", "表达式成立", "", "是", "", ""])
    book.save(path); book.close()
    stats: dict = {}
    issues = check_expression_rules(load_central_config([path]), schema="FIVE_SEGMENT_V1",
                                    stats_out=stats)
    assert stats["total"] == 2 and stats["enabled"] == 2 and stats["failed"] == 0
    assert stats["vocabulary"] and "净值" in stats["vocabulary"][0]
    # 词表提示不算错误：问题清单为空。
    assert not any(i["问题"] for i in issues)
