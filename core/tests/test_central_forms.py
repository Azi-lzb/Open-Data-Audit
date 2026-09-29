"""系统导数转 Excel 测试：合成模板 + 合成比较结果 → 每机构 xlsx。"""

from __future__ import annotations

import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "core" / "src"))

from base_audit.systems.s3_central_statistics.config import ensure_split_central_configs
from base_audit.systems.s3_central_statistics.form_template import (
    embed_financial_template_in_config,
    has_embedded_form_template,
)
from base_audit.systems.s3_central_statistics.service import render_financial_forms

COMPARE_HEADERS = [
    "业务类", "数据日期", "机构类代码", "机构类名称", "地区代码", "地区名称",
    "指标顺序码", "指标代码", "指标名称", "数据属性", "币种", "频度", "批次",
    "数据值", "上期值", "增减额(亿元)", "环比", "备注", "是否说明", "说明内容", "计算过程",
]


def _write_template(path: Path) -> None:
    book = Workbook()
    listing = book.active
    listing.title = "报表清单"
    listing.append(("报表代码", "报表名称", "频度", "批次"))
    listing.append(("A1411", "资产负债项目月报表", "月", "1"))
    sheet = book.create_sheet("A1411")
    sheet.append(("指标代码", "指标名称", "行序号", "余额"))
    sheet.append(("12A09", "一、房地产开发贷款", 1, None))
    sheet.append(("12M01", "一、资产类总计", 2, None))
    book.save(path)
    book.close()


def _write_comparison(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "2026年8月对比结果"
    sheet.append(COMPARE_HEADERS)
    # 人民币 + 12A09 + 余额：命中 A1411 余额列
    sheet.append(("人民币", "2026-08-31", "6k0i", "示例甲银行", "4400000", "广东省",
                  "12A09", "12A09", "一、房地产开发贷款", "余额", "人民币", "月", "1",
                  120.5, 100.0, 20.5, 0.205, "", "", "", ""))
    # 人民币 + 12M01 + 余额
    sheet.append(("人民币", "2026-08-31", "6k0i", "示例甲银行", "4400000", "广东省",
                  "12M01", "12M01", "一、资产类总计", "余额", "人民币", "月", "1",
                  30.0, None, 30.0, None, "本期有，上期无", "", "", ""))
    book.save(path)
    book.close()


class RenderFormsTests(unittest.TestCase):
    def test_renders_one_file_per_org_with_values_and_colors(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            template = root / "金融表单.xlsx"
            comparison = root / "比较结果.xlsx"
            _write_template(template)
            _write_comparison(comparison)
            outputs = ensure_split_central_configs(root / "config")

            result = render_financial_forms(
                comparison_file=comparison, template_file=template,
                output_dir=root / "输出", config_path=(outputs["common"], outputs["forms"]),
            )
            self.assertEqual([p.name for p in result.files], ["2026.8 示例甲银行 广东省.xlsx"])
            book = load_workbook(result.files[0], data_only=False)
            try:
                self.assertIn("A1411", book.sheetnames)
                sheet = book["A1411"]
                self.assertEqual(sheet.max_column, 7)   # 3 固定 + 4 数据块
                self.assertEqual(sheet.cell(1, 5).value, "上期-余额")
                self.assertEqual(sheet.cell(1, 6).value, "增减额-余额")
                self.assertEqual(sheet.cell(1, 7).value, "环比-余额")
                self.assertEqual(sheet.cell(2, 4).value, 120.5)
                self.assertEqual(sheet.cell(2, 5).value, 100.0)
                self.assertEqual(sheet.cell(2, 6).value, 20.5)
                self.assertAlmostEqual(sheet.cell(2, 7).value, 0.205)
                # 显示格式：本期块 0.00、环比块 0.00%（底层值保持完整精度）
                self.assertEqual(sheet.cell(2, 4).number_format, "0.00")
                self.assertEqual(sheet.cell(2, 7).number_format, "0.00%")
                self.assertEqual(sheet.cell(3, 4).number_format, "0.00")
                # 12M01 本期有上期无 → 红色 3（FFFF0000）
                fill = sheet.cell(3, 4).fill
                self.assertEqual(fill.fgColor.rgb, "FFFF0000")
                self.assertNotIn("月报1批", book.sheetnames)  # 不在清单的表不受影响
            finally:
                book.close()

    def test_embedded_config_is_used_as_template_and_not_emitted(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            template = root / "金融表单.xlsx"
            comparison = root / "比较结果.xlsx"
            outputs = ensure_split_central_configs(root / "config")
            config = outputs["forms"]
            _write_template(template)
            _write_comparison(comparison)

            embed_financial_template_in_config(config, template)
            self.assertTrue(has_embedded_form_template(config))
            embedded = load_workbook(config, read_only=True, data_only=False)
            try:
                self.assertIn("报表清单", embedded.sheetnames)
                self.assertIn("A1411", embedded.sheetnames)
                self.assertIn("转表设置", embedded.sheetnames)
            finally:
                embedded.close()

            result = render_financial_forms(
                comparison_file=comparison, output_dir=root / "输出", config_path=(outputs["common"], config),
                template_file=None, render_mode="OPENPYXL",
            )
            rendered = load_workbook(result.files[0], data_only=False)
            try:
                self.assertIn("A1411", rendered.sheetnames)
                self.assertNotIn("使用说明", rendered.sheetnames)
                self.assertNotIn("金融表单映射", rendered.sheetnames)
                self.assertNotIn("转表设置", rendered.sheetnames)
            finally:
                rendered.close()


if __name__ == "__main__":
    unittest.main()
