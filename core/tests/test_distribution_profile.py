"""交付范围只控制显示；完整版和 S1/S2 版共用同一业务实现。"""

from __future__ import annotations

import json
import importlib.util
from pathlib import Path
import re
import shutil
import subprocess

import pytest


ROOT = Path(__file__).resolve().parents[2]
NODE = shutil.which("node")
FRONTEND = ROOT / "core/frontend/web/index.html"


@pytest.mark.parametrize("shell", ["shell-flask", "shell-pywebview"])
@pytest.mark.parametrize("profile_id", ["full", "s1s2"])
def test_external_reference_and_its_default_are_shipped_without_changing_source(tmp_path, monkeypatch, shell, profile_id):
    manifest = json.loads((ROOT / "版本管理/VERSION_MANIFEST.json").read_text(encoding="utf-8"))
    filename = Path(manifest["configs"]["S2-CFG-02"]["file"]).name
    private = tmp_path / "config-real"
    (private / "默认配置").mkdir(parents=True)
    (private / filename).write_bytes(b"private-current-template")
    (private / "默认配置" / filename).write_bytes(b"independent-default-template")
    (private / "3.0-test.xlsx").write_bytes(b"s3-only")
    (private / "S5_模板文件").mkdir()
    (private / "S5_模板文件/合成布局.xlsx").write_bytes(b"s5-layout")
    (private / "模板文件_命名区域").mkdir()
    (private / "模板文件_命名区域/共享命名区域模板.xlsx").write_bytes(b"shared-named-range-template")
    (private / "S5_审核配置.xlsx").write_bytes(b"s5-config")
    (private / "5.逐笔统计系统_审核配置.xlsx").write_bytes(b"renamed-s5-config")
    before = {p.relative_to(private): p.read_bytes() for p in private.rglob("*.xlsx")}
    monkeypatch.syspath_prepend(str(ROOT / "shell-flask"))
    spec = importlib.util.spec_from_file_location(f"external_pack_{shell.replace('-', '_')}", ROOT / shell / "build_exe.py")
    builder = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(builder)
    monkeypatch.setattr(builder, "CONFIG_TEMPLATE", private)
    release = tmp_path / "release"
    if shell == "shell-flask":
        monkeypatch.setattr(builder, "DIST", release)
        builder.assemble_config_template()
    else:
        builder.assemble_config_template(release)
    from distribution_profile import DEFAULT_CONFIG, filter_release_config, load_config
    filter_release_config(release / "config", load_config(DEFAULT_CONFIG), profile_id)
    assert (release / "config" / filename).read_bytes() == before[Path(filename)]
    assert (release / "config/默认配置" / filename).read_bytes() == before[Path("默认配置") / filename]
    assert (release / "config/3.0-test.xlsx").exists() is (profile_id == "full")
    assert (release / "config/S5_模板文件/合成布局.xlsx").exists() is (profile_id == "full")
    assert (release / "config/S5_审核配置.xlsx").exists() is (profile_id == "full")
    assert (release / "config/5.逐笔统计系统_审核配置.xlsx").exists() is (profile_id == "full")
    assert (release / "config/模板文件_命名区域/共享命名区域模板.xlsx").exists()
    assert {p.relative_to(private): p.read_bytes() for p in private.rglob("*.xlsx")} == before


@pytest.mark.parametrize("shell", ["shell-flask", "shell-pywebview"])
def test_windows_pe_product_metadata_uses_audittools_brand(tmp_path, monkeypatch, shell):
    """两种 Windows 外壳的 PE 属性均显示指定中英文名称和公司名。"""
    monkeypatch.syspath_prepend(str(ROOT / "shell-flask"))
    spec = importlib.util.spec_from_file_location(
        f"version_resource_{shell.replace('-', '_')}", ROOT / shell / "build_exe.py"
    )
    builder = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(builder)
    monkeypatch.setattr(builder, "ROOT", tmp_path / shell)

    resource = builder.write_version_file("AuditTools", {"version": "26.1.7.1", "name": "审核工具"})
    contents = resource.read_text(encoding="utf-8")
    assert "StringStruct('CompanyName', 'HZPBC')" in contents
    assert "StringStruct('FileDescription', '审核工具')" in contents
    assert "StringStruct('ProductName', '审核工具')" in contents
    assert "StringStruct('InternalName', 'AuditTools')" in contents
    assert "StringStruct('ProductName', 'AuditTools')" in contents
    assert "StringTable('080404B0'" in contents
    assert "StringTable('040904B0'" in contents
    assert "[2052, 1200, 1033, 1200]" in contents
    compile(contents, str(resource), "exec")


def test_linux_desktop_application_uses_audittools_english_name():
    script = (ROOT / "shell-flask/Linux-2-打包DEB与TARGZ.sh").read_text(encoding="utf-8")
    assert "Name=AuditTools V$APP_VERSION" in script
    assert "Name[zh_CN]=审核工具 V$APP_VERSION" in script
    assert "Maintainer: HZPBC" in script


def test_linux_packager_includes_shared_named_range_templates_in_s1s2_and_full_profiles():
    script = (ROOT / "shell-flask/Linux-2-打包DEB与TARGZ.sh").read_text(encoding="utf-8")
    assert 'if [ "$CONFIG_PROFILE" = internal-config ] && [ -d "$CONFIG_SOURCE/模板文件_命名区域" ]; then' in script
    assert 'cp -R "$CONFIG_SOURCE/模板文件_命名区域" "$PAYLOAD/config/"' in script


def test_linux_public_payload_keeps_external_headers_and_removes_private_rows(tmp_path):
    import sys
    from openpyxl import load_workbook
    script = (ROOT / "shell-flask/Linux-2-打包DEB与TARGZ.sh").read_text(encoding="utf-8")
    block = re.search(r"<<'PY'\n(import shutil\nimport sys\nimport json\n.*?\n)PY\n", script, re.S)
    assert block is not None
    manifest_path = ROOT / "版本管理/VERSION_MANIFEST.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    filename = Path(manifest["configs"]["S2-CFG-02"]["file"]).name
    source = tmp_path / "source"
    shutil.copytree(ROOT / "config", source)
    book = load_workbook(source / filename)
    headers = [c.value for c in book["集中系统数据"][2]]
    book["集中系统数据"]["B3"] = "private-bank"
    book["集中系统数据"]["C3"] = 17
    book["参照表"].append(["private-code", "private-bank", "private-region"])
    book.save(source / filename)
    book.close()
    before = (source / filename).read_bytes()
    target = tmp_path / "release/config"
    (target / "默认配置").mkdir(parents=True)
    result = subprocess.run([sys.executable, "-X", "utf8", "-", str(source), str(target), str(manifest_path)], input=block[1], text=True, encoding="utf-8", capture_output=True)
    assert result.returncode == 0, result.stderr
    for path in (target / filename, target / "默认配置" / filename):
        exported = load_workbook(path)
        assert [c.value for c in exported["集中系统数据"][2]] == headers
        assert exported["集中系统数据"].max_row == 2
        assert exported["参照表"].max_row == 1
        exported.close()
    assert (source / filename).read_bytes() == before


@pytest.mark.parametrize("profile_id", ["full", "s1s2"])
def test_windows_build_finalizes_with_same_profile_as_linux(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, profile_id: str
) -> None:
    """两个 Windows 构建均覆盖前端后，必须仍写入与 Linux 相同的交付方案。"""
    monkeypatch.syspath_prepend(str(ROOT / "shell-flask"))
    spec = importlib.util.spec_from_file_location(
        "windows_delivery_build_test", ROOT / "shell-flask/build_windows_all.py"
    )
    assert spec is not None and spec.loader is not None
    builder = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(builder)
    monkeypatch.setattr(builder, "ROOT", tmp_path)
    monkeypatch.setattr(builder, "DIST_ROOT", tmp_path / "dist")
    commands = []
    source = FRONTEND.read_bytes()

    def run_builder(command, *, cwd, env, stdin):
        commands.append(command)
        output = Path(env["FLASK_DIST_DIR"])
        win7 = Path(command[-1]).name == "build-win7.bat"
        (output / (builder.executable_stem("flask", win7=win7) + ".exe")).touch()
        # Win7 跟现代版一样会重新复制完整前端，方案必须在两者完成后才注入。
        page = output / "core/frontend/web/index.html"
        page.parent.mkdir(parents=True, exist_ok=True)
        page.write_bytes(source)
        (output / "config").mkdir(exist_ok=True)
        if win7:
            (output / "运行库修复").mkdir()
            (output / "运行环境说明-Win7.txt").touch()
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(builder.subprocess, "run", run_builder)
    assert builder.main(["--profile", profile_id]) == 0
    outputs = list((tmp_path / "dist").iterdir())
    assert len(outputs) == 1
    output = outputs[0]
    assert len(commands) == 2
    assert [Path(command[-1]).name for command in commands] == [
        "build-modern.bat", "build-win7.bat"
    ]
    assert not (output / "BUILD_INCOMPLETE.txt").exists()
    assert (output / "distribution-profile.txt").read_text().strip() == profile_id
    assert output.name.endswith("_S1-S2") is (profile_id == "s1s2")

    from distribution_profile import DEFAULT_CONFIG, embed_profile, load_config

    linux_page = tmp_path / "linux-page.html"
    linux_page.write_bytes(source)
    embed_profile(linux_page, load_config(DEFAULT_CONFIG), profile_id)
    windows_page = output / "core/frontend/web/index.html"
    assert windows_page.read_bytes() == linux_page.read_bytes()
    assert FRONTEND.read_bytes() == source
    payload = json.loads(re.search(
        r'<script id="audit-distribution-config" type="application/json">(.*?)</script>',
        windows_page.read_text(encoding="utf-8"), re.S,
    ).group(1))
    assert payload["enabledSystems"] == (
        ["S1", "S2"] if profile_id == "s1s2" else ["S1", "S2", "S3", "S4", "S5"]
    )


@pytest.mark.skipif(NODE is None, reason="未安装 Node.js，无法执行前端行为检查")
def test_delivery_profiles_render_only_their_visible_systems(tmp_path: Path) -> None:
    html = FRONTEND.read_text(encoding="utf-8")
    # 实际页面脚本须仍能被浏览器解析。
    inline = "\n".join(re.findall(r"<script(?:\s[^>]*)?>(.*?)</script>", html, re.S))
    script = tmp_path / "frontend.js"
    script.write_text(inline, encoding="utf-8")
    subprocess.run([NODE, "--check", str(script)], check=True, capture_output=True, text=True)

    def section(start: str, end: str) -> str:
        return html.split(start, 1)[1].split(end, 1)[0]

    functions = "function isS1S2Edition" + section(
        "function isS1S2Edition", "function updateAuditWorkbenchMode"
    )
    functions += "function renderGeneral" + section(
        "function renderGeneral", "/* ---- 高级设置页签"
    )
    functions += "function renderGuide" + section(
        "function renderGuide", "/* ---- 常见问题列表"
    )
    flows = "const GUIDE_FLOWS=" + section("const GUIDE_FLOWS=", ";") + ";"
    faq = json.loads((ROOT / "docs/常见问题.json").read_text(encoding="utf-8"))
    program = r"""
const assert=require('node:assert/strict');
let reduced=false;
const pane={innerHTML:''};
const state={};
const document={body:{classList:{contains:()=>reduced}},getElementById:()=>pane};
const esc=value=>String(value);
const configDefaultStatus=()=>'';
const centralConfigDescription=()=>'';
const syncGeneralControls=()=>{};
const auditExtrasUnlocked=()=>true;
""" + functions + flows + "\nconst faq=" + json.dumps(faq, ensure_ascii=False) + r""";
const entries=[{date:'2026-10-15',groups:[{label:'优化',items:['S1 逐笔审核','S3 大集中统计系统','S4 地金系统','报表采集外部核对']}]}];
assert.equal(editionFaqItems(faq),faq);
assert.equal(editionUpdateEntries(entries),entries);
renderGeneral();
assert.match(pane.innerHTML,/大集中通用配置/);
renderGuide();
assert.match(pane.innerHTML,/本期数值核对/);
const guideCodes=()=>[...pane.innerHTML.matchAll(/class="guide-flow-num">(S\dF\d+)/g)].map(match=>match[1]);
assert.deepEqual(guideCodes(),['S1F1','S1F2','S2F1','S3F1','S3F2','S3F3','S3F4','S3F5','S1F3','S1F4','S1F5','S5F1']);
reduced=true;
const visibleFaq=editionFaqItems(faq);
assert.ok(visibleFaq.length>0&&visibleFaq.length<faq.length);
assert.ok(visibleFaq.some(item=>item.question==='第一次使用，我需要在界面上做哪些准备？'));
assert.ok(visibleFaq.some(item=>item.answer.includes('大集中数据文件')));
for(const item of visibleFaq){
  assert.equal(includesUndeliveredSystemInfo(item.question+item.answer),false);
  assert.doesNotMatch(item.answer,/六册|三个系统/);
}
assert.deepEqual(editionUpdateEntries(entries)[0].groups[0].items,['S1 逐笔审核','报表采集外部核对']);
assert.equal(entries[0].groups[0].items.length,4);
const taggedItems=['[通用] 跨系统说明','[S1] 逐笔校验','[S2] 报表采集使用大集中系统数据','[S3] 大集中比较','[S5] 配置版规则'];
const taggedVisible=editionUpdateEntries([{groups:[{label:'新增',items:taggedItems}]}]);
assert.deepEqual(taggedVisible[0].groups[0].items,taggedItems.slice(0,3));
assert.deepEqual(updateItemInfo('[S1/S2] 共用更新').systems,['S1','S2']);
assert.equal(updateItemInfo('[S2] 报表采集使用大集中系统数据').text,'报表采集使用大集中系统数据');
state.distributionUnlocked=true;
assert.deepEqual(editionUpdateEntries([{groups:[{items:taggedItems}]}])[0].groups[0].items,taggedItems.slice(0,3));
state.distributionUnlocked=false;
assert.deepEqual(editionUpdateEntries([{groups:[{items:['S3 功能']}]}]),[]);
assert.equal(includesUndeliveredSystemInfo('升级后检查 3.3「报表清单」'),true);
assert.equal(includesUndeliveredSystemInfo('整包版本 V26.1.3.1'),false);
assert.deepEqual(editionUpdateEntries([{groups:[{items:['大集中指标查询本期/上期']}]}]),[]);
renderGeneral();
assert.match(pane.innerHTML,/逐笔统计系统配置/);
assert.match(pane.innerHTML,/报表采集系统配置/);
assert.doesNotMatch(pane.innerHTML,/大集中通用配置|大集中执行比较配置|大集中本期数值核对配置|大集中指标比较拆分配置/);
renderGuide();
assert.doesNotMatch(pane.innerHTML,/本期数值核对|两期数据 ⇨ 比较文件|机构说明 ⇨ 说明文件/);
// 交付版过滤后保留功能身份，不能按卡片位置重新编号。
assert.deepEqual(guideCodes(),['S1F1','S1F2','S2F1','S1F3','S1F4','S1F5']);
console.log('profile behavior passed');
"""
    result = subprocess.run(
        [NODE, "-e", program], check=True, capture_output=True, text=True, encoding="utf-8"
    )
    assert "profile behavior passed" in result.stdout


@pytest.mark.skipif(NODE is None, reason="未安装 Node.js，无法执行前端行为检查")
def test_period_pair_targets_show_credit_codes_without_org_name_mapping() -> None:
    html = FRONTEND.read_text(encoding="utf-8")
    render = "function pairListFilters" + html.split("function pairListFilters", 1)[1].split("function openPcPath", 1)[0]
    program = r"""
const assert=require('node:assert/strict');
const box={style:{},innerHTML:''};
const document={getElementById:id=>id==='pairList'?box:null};
const esc=value=>String(value).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;');
const shortPairBook=value=>value;
""" + render + r"""
renderPairList([
  {orgName:'示例银行',org:'001122334455667788',socialCreditCode:'001122334455667788',form:'20201',side:'both',curName:'current.xlsx',preName:'previous.xlsx'},
  {org:'9144BB',socialCreditCode:'9144BB',form:'20202',side:'pre',preName:'banks#20260331#01#20202.xlsx'},
  {orgName:'非标准机构文件',org:'非标准机构文件',form:'20201',side:'cur',curName:'ordinary.xlsx'},
  {orgName:'测试名称',org:'code',socialCreditCode:'"><svg onload=1>',form:'20201',side:'cur',curName:'safe.xlsx'}
]);
const codes=[...box.innerHTML.matchAll(/class="pair-target-code"[^>]*>(.*?)<\/span>/g)].map(match=>match[1]);
assert.deepEqual(codes,[
  '信用代码：001122334455667788','信用代码：9144BB','信用代码：未识别','信用代码：&quot;&gt;&lt;svg onload=1&gt;'
]);
assert.ok(box.innerHTML.includes('统一社会信用代码：001122334455667788'));
assert.ok(box.innerHTML.includes('示例银行'));
assert.ok(box.innerHTML.includes('1 对齐全，3 项单边'));
assert.ok(box.innerHTML.includes('20202'));
assert.equal(box.innerHTML.includes('<svg onload=1>'),false);
renderPairList([]);
assert.equal(box.style.display,'none');
"""
    result = subprocess.run([NODE, "-e", program], capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stderr


@pytest.mark.skipif(NODE is None, reason="未安装 Node.js，无法执行前端行为检查")
def test_period_pair_filters_combine_clear_and_preserve_full_audit_scope() -> None:
    html = FRONTEND.read_text(encoding="utf-8")
    render = "function pairListFilters" + html.split("function pairListFilters", 1)[1].split("function openPcPath", 1)[0]
    confirm = "function runPeriodAudit" + html.split("function runPeriodAudit", 1)[1].split("async function doCreatePeriodCompare", 1)[0]
    program = r"""
const assert=require('node:assert/strict');
const elements={
  pairList:{innerHTML:'',style:{}},
  pairFilters:{hidden:true},
  pairFilterToggle:{textContent:'筛选',attributes:{},setAttribute(key,value){this.attributes[key]=value;}},
  pairOrgFilter:{value:''},pairCreditFilter:{value:''},
  pairFormFilter:{value:''},pairSideFilter:{value:''}
};
let optionWrites=0,optionHtml='';
Object.defineProperty(elements.pairFormFilter,'innerHTML',{
  get:()=>optionHtml,set:value=>{optionWrites++;optionHtml=value;}
});
let listWrites=0,listHtml='';
elements.pairList.scrollTop=23;
Object.defineProperty(elements.pairList,'innerHTML',{
  get:()=>listHtml,set:value=>{listWrites++;listHtml=value;elements.pairList.scrollTop=0;}
});
const document={getElementById:id=>elements[id]||null};
const esc=value=>String(value).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;');
const shortPairBook=value=>value;
const state={pcCurDir:'current',pcPreDir:'previous',periodPairs:[
  {orgName:'甲银行',socialCreditCode:'0011AA',form:'20201',side:'both',curName:'a.xlsx',preName:'b.xlsx'},
  {orgName:'甲银行',socialCreditCode:'0011AA',form:'20202',side:'pre',preName:'b.xlsx'},
  {orgName:'乙银行',socialCreditCode:'0022BB',form:'20201',side:'cur',curName:'c.xlsx'},
  {orgName:'Gamma Bank',socialCreditCode:'0033CC',form:'20203',side:'both',curName:'d.xlsx',preName:'e.xlsx'},
  {orgName:'无代码机构',org:'无代码机构',form:'20201',side:'cur',curName:'f.xlsx'}
]};
const before=JSON.stringify(state);
let confirmation='';
const showRunConfirm=(label,run,summary)=>{confirmation=summary;};
const doCreatePeriodAudit=()=>{};
const alert=message=>{throw new Error(message);};
const names=()=>[...elements.pairList.innerHTML.matchAll(/class="pair-target-org">(.*?)<\/span>/g)].map(match=>match[1]);
""" + render + confirm + r"""
renderPairList(state.periodPairs);
assert.equal(names().length,5);
assert.ok(elements.pairList.innerHTML.includes('class="pair-row pair-header"'));
assert.equal(elements.pairList.scrollTop,23,'重建列表后保留滚动位置');
const initialListWrites=listWrites;
renderPairList(state.periodPairs);
assert.equal(listWrites,initialListWrites,'相同配对数据和筛选条件复用现有DOM');
assert.equal(elements.pairList.scrollTop,23);
assert.ok(optionHtml.includes('value="20201"'));
assert.ok(optionHtml.indexOf('value="20201"')<optionHtml.indexOf('value="20202"'));
togglePairFilters();
assert.equal(elements.pairFilters.hidden,false);
assert.equal(elements.pairFilterToggle.attributes['aria-expanded'],'true');
elements.pairOrgFilter.value=' 甲 ';
refreshPairFilters();
assert.deepEqual(names(),['甲银行','甲银行']);
elements.pairCreditFilter.value=' aa ';
elements.pairFormFilter.value='20202';
refreshPairFilters();
assert.deepEqual(names(),['甲银行']);
assert.ok(elements.pairList.innerHTML.includes('筛选 1/5 项，0 对齐全，1 项单边'));
assert.equal(elements.pairFilterToggle.textContent,'筛选（3）');
elements.pairSideFilter.value='pre';
refreshPairFilters();
assert.deepEqual(names(),['甲银行']);
assert.equal(elements.pairFilterToggle.textContent,'筛选（4）');
runPeriodAudit();
assert.ok(confirmation.includes('对比对象：5 项'));
const writes=optionWrites;
renderPairList(state.periodPairs);
assert.equal(elements.pairOrgFilter.value,' 甲 ');
assert.equal(elements.pairCreditFilter.value,' aa ');
assert.equal(elements.pairFormFilter.value,'20202');
assert.equal(elements.pairSideFilter.value,'pre');
assert.equal(optionWrites,writes);
togglePairFilters();
assert.equal(elements.pairFilters.hidden,true);
assert.equal(elements.pairFilterToggle.attributes['aria-expanded'],'false');
assert.deepEqual(names(),['甲银行']);
elements.pairCreditFilter.value='BB';
refreshPairFilters();
assert.equal(names().length,0);
assert.ok(elements.pairList.innerHTML.includes('没有符合筛选条件的对比对象'));
assert.ok(elements.pairList.innerHTML.includes('筛选 0/5 项'));
resetPairFilters();
assert.equal(elements.pairOrgFilter.value,'');
assert.equal(elements.pairCreditFilter.value,'');
assert.equal(elements.pairFormFilter.value,'');
assert.equal(elements.pairSideFilter.value,'');
assert.equal(names().length,5);
assert.equal(elements.pairFilterToggle.textContent,'筛选');
elements.pairCreditFilter.value='无代码机构';
refreshPairFilters();
assert.equal(names().length,0,'机构名称不可充作信用代码进行筛选');
resetPairFilters();
elements.pairFormFilter.value='20201';
refreshPairFilters();
assert.deepEqual(names(),['甲银行','乙银行','无代码机构']);
resetPairFilters();
elements.pairOrgFilter.value='gAMMA';
refreshPairFilters();
assert.deepEqual(names(),['Gamma Bank']);
resetPairFilters();
assert.equal(JSON.stringify(state),before,'筛选和清空不得修改正式审核数据或范围');
elements.pairFormFilter.value='20203';
renderPairList([state.periodPairs[0]]);
assert.equal(elements.pairFormFilter.value,'','目录表单集合变化时清除失效选项');
assert.equal(names().length,1);
renderPairList([{orgName:'测试',socialCreditCode:'0011',form:'"><svg onload=1>',side:'cur',curName:'safe.xlsx'}]);
assert.equal(optionHtml.includes('<svg'),false);
assert.equal(elements.pairList.innerHTML.includes('<svg'),false);
renderPairList([]);
assert.equal(elements.pairList.style.display,'none');
assert.equal(elements.pairList.innerHTML,'');
assert.equal(optionHtml,'<option value="">全部表单</option>');
"""
    result = subprocess.run([NODE, "-e", program], capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stderr


@pytest.mark.skipif(NODE is None, reason="未安装 Node.js，无法执行前端行为检查")
def test_s1_pending_file_filters_and_visible_select_all_preserve_hidden_selection() -> None:
    html = FRONTEND.read_text(encoding="utf-8")
    start = html.index("const S1_FILE_MATCH_LABELS=")
    end = html.index("function checkedPaths()", start)
    filter_helpers = html[start:end]
    selection_helpers = "function checkedPaths()" + html.split("function checkedPaths()", 1)[1].split("// 主按钮一次调用", 1)[0]
    assert 'id="s1-file-filter-toggle"' in html
    assert 'id="s1-file-filters"' in html
    assert 'id="s1-match-filter"' in html
    assert 'id="s1-template-filter"' in html
    assert "(s1FileMatchesFilters(f)?'':' hidden')" in html

    program = r"""
const assert=require('node:assert/strict');
const allBoxes=[
  {value:'matched.xlsx',checked:true,visible:true},
  {value:'hidden.xlsx',checked:true,visible:false},
  {value:'second.xlsx',checked:false,visible:true}
];
const document={querySelectorAll(selector){
  if(selector==='#fileList .file-check:checked')return allBoxes.filter(item=>item.checked);
  if(selector==='#fileList .file-row:not([hidden]) .file-check')return allBoxes.filter(item=>item.visible);
  return [];
}};
const files=[
  {path:'matched.xlsx',template:'模板甲.xlsx'},
  {path:'hidden.xlsx',template:'未识别'},
  {path:'second.xlsx',template:'模板乙.xlsx'}
];
const state={sourceFiles:files};
""" + filter_helpers + r"""
assert.deepEqual(s1VisibleFiles(files).map(item=>item.path),['matched.xlsx','hidden.xlsx','second.xlsx']);
s1FileFilters={match:'matched',template:''};
assert.deepEqual(s1VisibleFiles(files).map(item=>item.path),['matched.xlsx','second.xlsx']);
s1FileFilters={match:'',template:'模板乙.xlsx'};
assert.deepEqual(s1VisibleFiles(files).map(item=>item.path),['second.xlsx']);
""" + selection_helpers + r"""
let saved=[];
const bridge={api:{async set_selected_files(paths){saved=paths;}}};
const refresh=async()=>{};
(async()=>{
  await selectAll(true);
  assert.deepEqual(saved,['matched.xlsx','hidden.xlsx','second.xlsx'],'隐藏文件的原勾选保留');
  await selectAll(false);
  assert.deepEqual(saved,['hidden.xlsx'],'全选或取消只作用于可见文件');
})().catch(error=>{console.error(error);process.exitCode=1});
"""
    result = subprocess.run([NODE, "-e", program], capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stderr


@pytest.mark.skipif(NODE is None, reason="未安装 Node.js，无法执行前端行为检查")
def test_s2_phase_progress_reflects_actual_state_and_hides_stale_values() -> None:
    html = FRONTEND.read_text(encoding="utf-8")
    assert 'id="periodProgress"' not in html
    assert '.period-progress' not in html
    assert 'id="progressDetail"' in html
    progress_ui = "function periodElapsedLabel" + html.split("function periodElapsedLabel", 1)[1].split("function openPeriodResult", 1)[0]
    progress_bar = "function countNew" + html.split("function countNew", 1)[1].split("const ICON", 1)[0]
    program = r"""
const assert=require('node:assert/strict');
const elements={
  progressDetail:{hidden:true,textContent:'',title:''},
  progress:{hidden:true},progressFill:{style:{},className:''},progressPct:{textContent:''}
};
const document={getElementById:id=>elements[id]||null};
const PROGRESS={decay:.80,holdMs:1400,maxPct:99.5};
let run=null,runTimer=null,prevLogTail='';
let timerSequence=0;
const timers=new Map(),clearedTimers=[];
const setTimeout=(callback,delay)=>{const id=++timerSequence;timers.set(id,{callback,delay});return id;};
const clearTimeout=id=>{if(id!=null){clearedTimers.push(id);timers.delete(id);}};
function fireProgressTimer(id){const timer=timers.get(id);assert.ok(timer,'expected a pending hold timer');timers.delete(id);timer.callback();}
""" + progress_ui + progress_bar + r"""
renderProgressDetail(null);
assert.equal(elements.progressDetail.hidden,true);
applyProgress({busy:true,periodProgress:{status:'running',label:'读取本期报表',current:4,total:10,percent:37.6,message:'C:\\reports\\甲.xls',elapsedSeconds:8},log:[{text:'S2 event'}]});
assert.equal(elements.progressFill.style.width,'38%','底部进度使用后端实际百分比');
assert.equal(run.period,true);
assert.equal(elements.progressDetail.textContent,'读取本期报表 · 4 / 10 · 耗时 8 秒');
assert.equal(elements.progressDetail.title,'C:\\reports\\甲.xls');
assert.equal(elements.progressDetail.hidden,false);
applyProgress({busy:true,periodProgress:{status:'running',label:'检查配置',current:0,total:0,percent:0,elapsedSeconds:1},log:[{text:'S2 event'}]});
assert.equal(elements.progressDetail.textContent,'检查配置 · 耗时 1 秒','无计数阶段不显示0/0');
applyProgress({busy:false,periodProgress:{status:'failed',percent:42.7},log:[{text:'S2 event'}]});
assert.equal(elements.progressFill.style.width,'43%','底部失败进度保留当前百分比');
assert.equal(elements.progressFill.className,'progress-fill error');
const failedTimer=runTimer;
assert.equal(timers.get(failedTimer).delay,PROGRESS.holdMs,'失败进度按原 hold 时长收起');
applyProgress({busy:false,periodProgress:{status:'failed',percent:42.7},log:[{text:'S2 event'}]});
assert.equal(runTimer,failedTimer,'重复失败轮询不重设 hold 定时器');
assert.equal(clearedTimers.includes(failedTimer),false);
applyProgress({busy:true,periodProgress:{status:'failed',percent:42.7},log:[{text:'S2 event'},{text:'Other feature event'}]});
assert.ok(clearedTimers.includes(failedTimer),'新任务启动时清理失败进度的旧定时器');
assert.equal(elements.progressFill.style.width,'20%','其他功能运行时按原日志估算通用进度');
assert.equal(elements.progressFill.className,'progress-fill working');
assert.equal(elements.progressDetail.hidden,true,'其他功能运行时清除S2阶段详情');

run=null;runTimer=null;prevLogTail='';
applyProgress({busy:true,periodProgress:{status:'running',percent:51},log:[{text:'Independent S2 failed job'}]});
applyProgress({busy:false,periodProgress:{status:'failed',percent:42.7},log:[{text:'Independent S2 failed job'}]});
const failedHideTimer=runTimer;
assert.equal(elements.progressFill.style.width,'43%');
fireProgressTimer(failedHideTimer);
assert.equal(elements.progress.hidden,true,'失败进度在 hold 时长后隐藏');
assert.equal(elements.progressDetail.hidden,true);

run=null;runTimer=null;prevLogTail='';
applyProgress({busy:true,periodProgress:{status:'running',percent:68},log:[{text:'Independent S2 completed job'}]});
applyProgress({busy:false,periodProgress:{status:'completed',percent:68},log:[{text:'Independent S2 completed job'}]});
assert.equal(elements.progressFill.style.width,'100%');
assert.equal(elements.progressFill.className,'progress-fill success');
const completedTimer=runTimer;
assert.equal(timers.get(completedTimer).delay,PROGRESS.holdMs);
applyProgress({busy:false,periodProgress:{status:'completed',percent:68},log:[{text:'Independent S2 completed job'}]});
assert.equal(runTimer,completedTimer,'重复完成轮询不重设 hold 定时器');
assert.equal(clearedTimers.includes(completedTimer),false);
fireProgressTimer(completedTimer);
assert.equal(elements.progress.hidden,true,'完成进度在 hold 时长后隐藏');
assert.equal(elements.progressDetail.hidden,true);

run=null;runTimer=null;prevLogTail='';
applyProgress({busy:true,periodProgress:{status:'completed',percent:100},log:[{text:'Other feature event'}]});
assert.equal(elements.progressFill.style.width,'20%','其他功能运行时旧 S2 结果不覆盖通用进度');
assert.equal(elements.progressFill.className,'progress-fill working');
"""
    result = subprocess.run([NODE, "-e", program], capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stderr


@pytest.mark.skipif(NODE is None, reason="未安装 Node.js，无法执行前端行为检查")
@pytest.mark.parametrize(("selection", "expected_names", "expected_gaps"), [
    ("", ["齐全甲", "仅上期乙", "仅本期丙", "齐全丁"], ["本期缺少", "上期缺少"]),
    ("single", ["仅上期乙", "仅本期丙"], ["本期缺少", "上期缺少"]),
    ("pre", ["仅上期乙"], ["本期缺少"]),
    ("cur", ["仅本期丙"], ["上期缺少"]),
    ("both", ["齐全甲", "齐全丁"], []),
])
def test_period_pair_status_filter_identifies_the_missing_period(selection, expected_names, expected_gaps) -> None:
    html = FRONTEND.read_text(encoding="utf-8")
    assert '<option value="pre">本期缺少</option>' in html
    assert '<option value="cur">上期缺少</option>' in html
    render = "function pairListFilters" + html.split("function pairListFilters", 1)[1].split("function openPcPath", 1)[0]
    program = r"""
const assert=require('node:assert/strict');
const box={innerHTML:'',style:{}};
const side={value:""" + json.dumps(selection) + r"""};
const document={getElementById:id=>id==='pairList'?box:id==='pairSideFilter'?side:null};
const esc=value=>String(value).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;');
const shortPairBook=value=>value;
const state={periodPairs:[
  {orgName:'齐全甲',form:'20201',side:'both',curName:'a.xlsx',preName:'b.xlsx'},
  {orgName:'仅上期乙',form:'20202',side:'pre',preName:'b.xlsx'},
  {orgName:'仅本期丙',form:'20201',side:'cur',curName:'c.xlsx'},
  {orgName:'齐全丁',form:'20203',side:'both',curName:'d.xlsx',preName:'e.xlsx'}
]};
const before=JSON.stringify(state);
""" + render + r"""
refreshPairFilters();
const names=[...box.innerHTML.matchAll(/class="pair-target-org">(.*?)<\/span>/g)].map(match=>match[1]);
const gaps=[...box.innerHTML.matchAll(/class="pair-source-name">— (本期缺少|上期缺少)<\/span>/g)].map(match=>match[1]);
assert.deepEqual(names,""" + json.dumps(expected_names, ensure_ascii=False) + r""");
assert.deepEqual(gaps,""" + json.dumps(expected_gaps, ensure_ascii=False) + r""");
assert.equal(JSON.stringify(state),before);
resetPairFilters();
assert.equal(side.value,'');
assert.equal([...box.innerHTML.matchAll(/class="pair-target-org">/g)].length,4);
"""
    result = subprocess.run([NODE, "-e", program], capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stderr


@pytest.mark.skipif(NODE is None, reason="未安装 Node.js，无法执行前端行为检查")
def test_period_banks_warning_renders_and_reaches_execution_confirmation() -> None:
    html = FRONTEND.read_text(encoding="utf-8")
    render = "function renderPeriodSourceWarning" + html.split("function renderPeriodSourceWarning", 1)[1].split("function renderPairList", 1)[0]
    confirm = "function runPeriodAudit" + html.split("function runPeriodAudit", 1)[1].split("async function doCreatePeriodCompare", 1)[0]
    program = r"""
const assert=require('node:assert/strict');
const notice={textContent:'',style:{display:'none'}};
const document={getElementById:id=>id==='periodSourceWarning'?notice:null};
const warning='banks 名称需依赖机构参照映射，建议默认单sheet或按报表单sheet';
const state={pcCurDir:'current',pcPreDir:'previous',periodPairs:[],periodSourceWarning:warning};
let summary='';
const showRunConfirm=(label,run,message)=>{summary=message;};
const doCreatePeriodAudit=()=>{};
const alert=message=>{throw new Error(message);};
""" + render + confirm + r"""
renderPeriodSourceWarning(state.periodSourceWarning);
assert.equal(notice.style.display,'block');
assert.equal(notice.textContent,warning);
runPeriodAudit();
assert.ok(summary.includes(warning));
state.periodSourceWarning='';
renderPeriodSourceWarning(state.periodSourceWarning);
assert.equal(notice.style.display,'none');
assert.equal(notice.textContent,'');
runPeriodAudit();
assert.equal(summary.includes(warning),false);
"""
    result = subprocess.run([NODE, "-e", program], capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stderr


@pytest.mark.skipif(NODE is None, reason="未安装 Node.js，无法执行前端行为检查")
def test_xml_tool_notes_use_one_collapse_mechanism() -> None:
    """说明使用统一卡片折叠，避免 details 原生切换与 collapsed 样式相互抵消。"""
    html = FRONTEND.read_text(encoding="utf-8")
    render = "function renderAdvanced" + html.split(
        "function renderAdvanced", 1
    )[1].split("/* 高级设置双布局", 1)[0]
    listener = html.split("/* 设置中心卡片标题条", 1)[1].split(
        "/* ---- 更新与答疑", 1
    )[0].split("*/", 1)[1]
    program = r"""
const assert=require('node:assert/strict');
let click;
const pane={innerHTML:'',addEventListener:(event,fn)=>{if(event==='click')click=fn;}};
const document={getElementById:()=>pane};
const state={};
const esc=value=>String(value);
const advancedLayoutMode=()=> 'system';
const envGroupCollapsed=()=>true,envZoneCollapsed=()=>true;
const isSystemVisible=()=>true,syncGeneralControls=()=>{};
""" + render + listener + r"""
renderAdvanced();
const notes=pane.innerHTML.match(/<div class="gen-card"><p class="gen-card-title">处理工具与兼容性说明<\/p>([\s\S]*?)<\/div><\/div>/);
assert.ok(notes,'说明应默认展开，使用普通卡片而非原生 details');
assert.match(notes[1],/zipfile/);
assert.doesNotMatch(pane.innerHTML,/<details class="gen-card"/);
const classes=new Set(['gen-card']);
const card={classList:{contains:name=>classes.has(name),toggle:name=>{
  if(classes.has(name)){classes.delete(name);}else{classes.add(name);}
}}};
const title={parentElement:card};
const event={target:{closest:()=>title}};
assert.equal(classes.has('collapsed'),false);
click(event);assert.equal(classes.has('collapsed'),true);
click(event);assert.equal(classes.has('collapsed'),false);
"""
    result = subprocess.run([NODE, "-e", program], capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stderr


def test_distribution_unlock_password_and_process_lifetime(tmp_path: Path) -> None:
    from base_audit.web_app import WebApi

    api = WebApi(tmp_path)
    paths = [tmp_path / "data" / name for name in ("开发设置.json", "用户设置.json")]
    original = {path: path.read_bytes() for path in paths if path.exists()}
    assert api.get_state()["distributionUnlocked"] is False
    developer_access = api.get_state()["developerAccess"]
    assert all(developer_access.values())
    for password in ("", "HZPBC", "hzpbc790", None):
        assert api.unlock_distribution_ui(password)["ok"] is False
        assert api.get_state()["distributionUnlocked"] is False
        assert api.get_state()["developerAccess"] == developer_access
    result = api.unlock_distribution_ui("HZPBC790")
    assert result["ok"] is True
    assert result["state"]["distributionUnlocked"] is True
    assert all(result["state"]["developerAccess"].values())
    assert api.get_state()["distributionUnlocked"] is True
    assert {path: path.read_bytes() for path in original} == original
    restarted = WebApi(tmp_path)
    assert restarted.get_state()["distributionUnlocked"] is False
    assert restarted.get_state()["developerAccess"] == developer_access


@pytest.mark.skipif(NODE is None, reason="未安装 Node.js，无法执行前端行为检查")
def test_about_five_click_unlock_restores_embedded_profile_ui(tmp_path: Path) -> None:
    html = FRONTEND.read_text(encoding="utf-8")
    assert '<p class="about-credits">Instructor: cyc · Developer: lzb · Tester: yyy</p>' in html
    assert ".about-credits{margin:4px 0 0;font-size:11px;color:var(--text-3);text-align:right}" in html

    def section(start: str, end: str) -> str:
        return start + html.split(start, 1)[1].split(end, 1)[0]

    functions = section("function isS1S2Edition", "function updateAuditWorkbenchMode")
    functions += section("let distributionUnlockClicks=", "let developerUnlockTarget=")
    functions += section("let developerUnlockTarget=", "let pendingConfirm=")
    config = json.loads((ROOT / "shell-flask/distribution-profiles.json").read_text(encoding="utf-8"))
    payload = {"enabledSystems": config["profiles"]["s1s2"]["enabledSystems"], "systems": config["systems"]}
    program = r"""
const assert=require('node:assert/strict');
let state={}, now=0, settingsRenders=0, accept=false;
const Date={now:()=>now};
const currentPage='settings', CFG={data:{},tab:'guide'};
const setSettingsTab=()=>settingsRenders++;
const switchPage=()=>{throw new Error('must remain on current page')};
const nodes={};
const element=id=>nodes[id]||(nodes[id]={style:{display:'none'},value:'',textContent:'',focus:()=>{}});
const classes=new Set();
const nav=['S1','S2','S3','S4','S5'].map(system=>({dataset:{system},hidden:false}));
const document={activeElement:{focus:()=>{}},body:{dataset:{},classList:{
  contains:value=>classes.has(value),toggle:(value,on)=>on?classes.add(value):classes.delete(value)
}},getElementById:element,querySelectorAll:()=>nav};
element('audit-distribution-config').textContent=PAYLOAD;
const bridge={api:{unlock_distribution_ui:async password=>{
  assert.equal(password,'typed-password');
  return accept?{ok:true,state:{distributionUnlocked:true,developerAccess:{S1:true,S2:true,S3:true,S4:true,S5:true}}}:{ok:false,message:'wrong password'};
}}};
""".replace("PAYLOAD", json.dumps(json.dumps(payload, ensure_ascii=False), ensure_ascii=False))
    program += functions + r"""
(async()=>{
applyDistributionProfile();
assert.equal(nav[2].hidden,true);
openAbout();
for(let i=0;i<4;i++)registerDistributionUnlockClick();
assert.equal(element('developerUnlockMask').style.display,'none');
now=2100;
registerDistributionUnlockClick();
assert.equal(element('developerUnlockMask').style.display,'none');
for(let i=0;i<4;i++)registerDistributionUnlockClick();
assert.equal(element('developerUnlockMask').style.display,'flex');
assert.equal(element('developerPassword').value,'');
element('developerPassword').value='typed-password';
const pending=forceDeveloperEntry();
closeDeveloperUnlock();
assert.equal(element('developerUnlockMask').style.display,'flex');
await pending;
assert.equal(element('developerUnlockError').textContent,'wrong password');
assert.equal(nav[2].hidden,true);
closeDeveloperUnlock();
assert.equal(element('developerPassword').value,'');
assert.equal(isSystemVisible('S3'),false);
openAbout();
for(let i=0;i<5;i++)registerDistributionUnlockClick();
element('developerPassword').value='typed-password';
accept=true;
await forceDeveloperEntry();
assert.equal(element('developerUnlockMask').style.display,'none');
assert.ok(nav.every(item=>!item.hidden));
assert.ok(classes.has('distribution-ui-unlocked'));
assert.equal(settingsRenders,1);
assert.equal(isSystemVisible('S3'),true);
const faq=[{category:'大集中统计系统',question:'S3'}];
assert.equal(editionFaqItems(faq),faq);
assert.deepEqual(DISTRIBUTION_CONFIG.enabledSystems,['S1','S2']);
state={};
DISTRIBUTION_CONFIG.enabledSystems=['S1','S2','S3','S4','S5'];
openAbout();
for(let i=0;i<5;i++)registerDistributionUnlockClick();
assert.equal(element('developerUnlockMask').style.display,'none');
console.log('about unlock passed');
})().catch(error=>{console.error(error);process.exit(1)});
"""
    script = tmp_path / "about_unlock_test.js"
    script.write_text(program, encoding="utf-8")
    result = subprocess.run([NODE, str(script)], capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stderr
    assert "about unlock passed" in result.stdout
