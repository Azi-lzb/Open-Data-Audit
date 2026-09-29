from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "core" / "src"))

from base_audit.systems.s3_central_statistics.explanation_exporter import (
    export_selected_explanations,
    import_explanation_feedback,
)


class ExplanationExporterTests(unittest.TestCase):
    def test_exports_only_rows_marked_for_explanation(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "比较结果.xlsx"
            output = Path(folder) / "指标说明.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "2026年8月对比结果"
            sheet.append(("机构类代码", "机构类名称", "地区代码", "地区名称", "指标代码", "指标名称", "数据值", "上期值", "增减额(亿元)", "输出说明", "说明内容"))
            sheet.append(("6k0i", "测试机构", "4400000", "广东省", "A001", "指标一", 2, 1, 1, "是", "本期增加"))
            sheet.append(("6k0i", "测试机构", "4400000", "广东省", "A002", "指标二", 3, 2, 1, "否", "不导出"))
            book.save(source)
            book.close()

            self.assertEqual(export_selected_explanations(source, output), 1)
            result = load_workbook(output, data_only=True)
            explanation = result["说明文件"]
            self.assertEqual(explanation.max_row, 3)
            self.assertEqual(explanation["A1"].value, "xx说明")
            self.assertIn("A1:H1", {str(area) for area in explanation.merged_cells.ranges})
            self.assertEqual(explanation.max_column, 8)
            self.assertEqual(explanation.cell(3, 3).value, "A001")
            self.assertEqual(explanation.cell(3, 8).value, "本期增加")
            self.assertEqual(result["生成信息"].cell(4, 2).value, 1)
            result.close()

    def test_uses_precise_then_wildcard_org_reference_and_keeps_source_region(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "比较结果.xlsx"
            config = Path(folder) / "3.0大集中通用配置.xlsx"
            output = Path(folder) / "指标说明.xlsx"

            book = Workbook()
            sheet = book.active
            sheet.title = "比较结果"
            sheet.append(("机构类代码", "机构类名称", "地区代码", "地区名称", "指标代码", "指标名称", "数据值", "上期值", "增减额(亿元)", "输出说明", "说明内容"))
            sheet.append(("7020", "村镇银行", "33445", "广东省", "A001", "指标一", 2, 1, 1, "是", ""))
            sheet.append(("6k0i", "示例甲银行", "4400000", "广东省", "A002", "指标二", 3, 2, 1, "是", ""))
            book.save(source)
            book.close()

            book = Workbook()
            sheet = book.active
            sheet.title = "机构地区参照"
            sheet.append(("机构类代码", "地区代码", "机构名称"))
            sheet.append(("7020", "33445", "示例乙银行"))
            # 唯一名称会自动视为该机构类通用映射，无需人为复制全部地区代码。
            sheet.append(("6k0i", "*", "示例甲银行"))
            book.save(config)
            book.close()

            self.assertEqual(export_selected_explanations(source, output, config_path=config), 2)
            result = load_workbook(output, data_only=True)
            explanation = result["说明文件"]
            self.assertEqual(explanation.cell(3, 2).value, "示例乙银行（广东省）")
            self.assertEqual(explanation.cell(4, 2).value, "示例甲银行（广东省）")
            result.close()

    def test_imports_feedback_by_four_business_identity_columns(self):
        with tempfile.TemporaryDirectory() as folder:
            target_path = Path(folder) / "指标说明.xlsx"
            feedback_path = Path(folder) / "机构反馈.xlsx"
            for path, reason in ((target_path, "待机构填写"), (feedback_path, "本期业务正常增加")):
                book = Workbook()
                sheet = book.active
                sheet.title = "说明文件"
                sheet.merge_cells("A1:H1")
                sheet["A1"] = "xx说明"
                sheet.append(("机构代码", "机构名称", "指标代码", "指标名称", "数据值", "上期值", "增减额(亿元)", "变动原因"))
                sheet.append(("6k0i", "测试机构", "A001", "指标一", 2.0, 1.0, 1.0, reason))
                book.save(path)
                book.close()

            result = import_explanation_feedback(target_path, [feedback_path])
            self.assertEqual((result.imported, result.unmatched, result.conflicts), (1, 0, 0))
            self.assertTrue(result.backup_path.is_file())
            target = load_workbook(target_path, data_only=True)
            self.assertEqual(target["说明文件"].cell(3, 8).value, "本期业务正常增加")
            target.close()

    def test_keeps_same_named_indicators_separate_by_indicator_code(self):
        with tempfile.TemporaryDirectory() as folder:
            target_path = Path(folder) / "指标说明.xlsx"
            feedback_path = Path(folder) / "机构反馈.xlsx"
            for path, reasons in (
                (target_path, ("待填写一", "待填写二")),
                (feedback_path, ("原因一", "原因二")),
            ):
                book = Workbook()
                sheet = book.active
                sheet.title = "说明文件"
                sheet.merge_cells("A1:H1")
                sheet["A1"] = "xx说明"
                sheet.append(("机构代码", "机构名称", "指标代码", "指标名称", "数据值", "上期值", "增减额(亿元)", "变动原因"))
                sheet.append(("6k0p", "示例乙银行（广东省）", "12P1Z", "农产品加工贷款", 0.05024, 0.05046, -0.00022, reasons[0]))
                sheet.append(("6k0p", "示例乙银行（广东省）", "12P28", "农产品加工贷款", 0.05024, 0.05046, -0.00022, reasons[1]))
                book.save(path)
                book.close()

            result = import_explanation_feedback(target_path, [feedback_path])
            self.assertEqual((result.imported, result.unmatched, result.conflicts), (2, 0, 0))
            target = load_workbook(target_path, data_only=True)
            self.assertEqual(target["说明文件"].cell(3, 8).value, "原因一")
            self.assertEqual(target["说明文件"].cell(4, 8).value, "原因二")
            target.close()

    def test_writes_conflict_report_with_feedback_source_and_reason(self):
        with tempfile.TemporaryDirectory() as folder:
            target_path = Path(folder) / "指标说明.xlsx"
            feedback_a = Path(folder) / "机构反馈A.xlsx"
            feedback_b = Path(folder) / "机构反馈B.xlsx"
            for path, reason in (
                (target_path, "待机构填写"),
                (feedback_a, "业务原因甲"),
                (feedback_b, "业务原因乙"),
            ):
                book = Workbook()
                sheet = book.active
                sheet.title = "说明文件"
                sheet.merge_cells("A1:H1")
                sheet["A1"] = "xx说明"
                sheet.append(("机构代码", "机构名称", "指标代码", "指标名称", "数据值", "上期值", "增减额(亿元)", "变动原因"))
                sheet.append(("6k0i", "测试机构", "A001", "指标一", 2.0, 1.0, 1.0, reason))
                book.save(path)
                book.close()

            result = import_explanation_feedback(target_path, [feedback_a, feedback_b])
            self.assertEqual(result.conflicts, 1)
            self.assertIsNotNone(result.conflict_report_path)
            self.assertTrue(result.conflict_report_path.is_file())
            report = load_workbook(result.conflict_report_path, data_only=True)
            row = list(report["冲突明细"].iter_rows(min_row=2, values_only=True))[0]
            self.assertEqual(row[0], "冲突")
            self.assertIn("多个反馈", row[-1])
            self.assertIn("机构反馈A.xlsx", "\n".join(str(value or "") for value in row))
            self.assertIn("业务原因甲", "\n".join(str(value or "") for value in row))
            report.close()
