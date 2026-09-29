"""黄金基线：汇总校验结果说明（Windows COM 管线，端到端输出契约）。

冻结现状：结构比对通过文件才参与汇总、汇总工作簿的 sheet 命名/表头/行序/
固定字段、输出文件名、运行日志。阶段 0 只冻结行为，不改变任何执行路径。
"""

from __future__ import annotations

from pathlib import Path

import pytest

import golden_fixtures as gf

pytestmark = [
    pytest.mark.skipif(
        not gf.excel_available(), reason="需要 Windows + 可用的 Excel/WPS COM 引擎"
    ),
    pytest.mark.golden_com,
]

EXPECTED_HEADERS = ["来源文件", "来源工作表", "指标编号", "指标名称", "本期情况"]


@pytest.fixture()
def summary_env(tmp_path: Path):
    src = tmp_path / "源数据"
    src.mkdir()
    gf.build_source(src / "甲银行.xlsx", deposit=2_000_000, loan=3_000_000, reserve=30_000)
    gf.build_source(src / "乙银行.xlsx", deposit=5_000_000, loan=4_000_000, reserve=10_000)
    gf.build_bad_source(src / "丙银行.xlsx")
    template = tmp_path / "模板.xlsx"
    gf.build_audit_template(template)
    config = gf.make_config(tmp_path)
    hashes = {path.name: gf.sha256(path) for path in sorted(src.glob("*.xlsx"))}
    return tmp_path, src, template, config, hashes


def test_summary_flow_golden_contract(summary_env):
    tmp_path, src, template, config, hashes = summary_env
    from base_audit.service import AuditService
    from base_audit.workflow.runner import run_dag_flow

    output_dir = tmp_path / "输出"
    service = AuditService(config_path=config)
    result = run_dag_flow(
        "dag:汇总校验结果说明",
        service=service,
        template_path=template,
        input_dir=src,
        output_dir=output_dir,
        period="2026-06",
        history_path=config,
        recursive=False,
        write_flow_logs=True,
    )

    # ---- 只汇总结构比对通过的文件；丙进 skipped（元素为 (路径, SourceMatch) 元组） ----
    assert all(not match.matched for _path, match in result.skipped_files)
    skipped_names = {Path(path).name for path, _match in result.skipped_files}
    assert skipped_names == {"丙银行.xlsx"}

    # ---- 源文件只读 ----
    for name, before in hashes.items():
        assert gf.sha256(src / name) == before, name

    # ---- 汇总工作簿：文件名取步骤“输出文件名”，按来源工作表命名 sheet ----
    summary_path = gf.glob_one(output_dir, "汇总校验结果说明_[0-9]*.xlsx")
    summary = gf.snapshot_workbook(summary_path)
    assert list(summary) == [gf.DATA_SHEET]
    sheet = summary[gf.DATA_SHEET]
    assert sheet["headers"] == EXPECTED_HEADERS
    # 行序 = 文件名字典序（Unicode 码点：乙 U+4E59 < 甲 U+7532，故乙先处理），
    # 组内按命名区域行序（A4:C6 三行）。
    rows = sheet["rows"]
    assert len(rows) == 6
    assert [row[0] for row in rows] == ["乙银行"] * 3 + ["甲银行"] * 3
    assert [row[1] for row in rows] == [gf.DATA_SHEET] * 6
    assert [row[2] for row in rows] == [20202001, 20202002, 20202003] * 2
    assert rows[0][4] == 5_000_000 and rows[3][4] == 2_000_000
    # 无“历史触发条数”等三列：普通汇总表头不含 8 个身份字段
    assert "历史触发条数" not in sheet["headers"]

    # ---- 运行日志 ----
    log_path = gf.glob_one(output_dir, "汇总校验结果说明_运行日志_*.xlsx")
    log = gf.snapshot_workbook(log_path)
    assert {"检查表结构区域", "检查汇总区域", "表结构比对", "任意行汇总"} <= set(log)


def test_summary_flow_without_validation_region(tmp_path: Path):
    """纯汇总模板不含“校验区域/审核规则”时也能运行（回归：曾误用审核链体检）。

    真实汇总模板（如“！汇总校验结果及报送说明.xlsx”）只定义汇总区域与表头
    区域，没有可复制的校验公式；模板装载必须走 summary_template 模式。
    """
    from openpyxl import load_workbook
    from base_audit.service import AuditService
    from base_audit.workflow.runner import run_dag_flow

    src = tmp_path / "源数据"
    src.mkdir()
    gf.build_source(src / "甲银行.xlsx", deposit=2_000_000, loan=3_000_000, reserve=30_000)
    template = tmp_path / "模板.xlsx"
    gf.build_audit_template(template)
    # 删掉“校验区域”“审核规则”与公式单元格，模拟纯汇总模板。
    book = load_workbook(template)
    if "校验区域" in book.defined_names:
        del book.defined_names["校验区域"]
    if "审核规则" in book.sheetnames:
        del book["审核规则"]
    sheet = book[gf.DATA_SHEET]
    sheet["D4"] = None
    sheet["D5"] = None
    book.save(template)
    book.close()

    config = gf.make_config(tmp_path)
    result = run_dag_flow(
        "dag:汇总校验结果说明",
        service=AuditService(config_path=config),
        template_path=template,
        input_dir=src,
        output_dir=tmp_path / "输出",
        period="2026-06",
        history_path=config,
        recursive=False,
        write_flow_logs=True,
    )
    summary = gf.snapshot_workbook(result.output_path)
    assert summary[gf.DATA_SHEET]["headers"] == EXPECTED_HEADERS
    assert len(summary[gf.DATA_SHEET]["rows"]) == 3
