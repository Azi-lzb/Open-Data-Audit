"""Focused coverage for read-only S3 3.1 configuration preflight."""

from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from base_audit.systems.s3_central_statistics.config import (
    ACTION_RULE_SHEET,
    COMPARISON_CONFIG_NAME,
    COMMON_CONFIG_NAME,
    EXPRESSION_RULE_SHEET,
    FIVE_SEGMENT_RULE_SHEET,
    INDICATOR_SHEET,
    CentralConfig,
    check_action_rules,
    load_central_config,
)
from base_audit.systems.s3_central_statistics.indicator_rule_engine import (
    supports_special_action,
)
from base_audit.web_app import WebApi


REPO_ROOT = Path(__file__).resolve().parents[2]
CONFIG_DIR = REPO_ROOT / "config"
COMMON_CONFIG = CONFIG_DIR / COMMON_CONFIG_NAME
COMPARISON_CONFIG = CONFIG_DIR / COMPARISON_CONFIG_NAME
FORMAL_CONFIGS_READY = COMMON_CONFIG.is_file() and COMPARISON_CONFIG.is_file()


def _api_for_config(common_path: Path, comparison_path: Path, schema: str) -> WebApi:
    """Build only the state used by the check API, avoiding app startup writes."""
    api = WebApi.__new__(WebApi)
    api.project_root = REPO_ROOT
    api.state = {
        "busy": False,
        "centralExpressionSchema": schema,
        "centralCommonConfig": str(common_path),
        "centralComparisonConfig": str(comparison_path),
        "log": [],
    }
    api._log_detail = lambda _text: None
    api._refresh_config_issues = lambda: None
    return api


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


@pytest.mark.skipif(not FORMAL_CONFIGS_READY, reason="缺少正式 3.0/3.1 配置工作簿")
@pytest.mark.parametrize(
    ("schema", "expression_sheet", "mode_label", "opposite_label"),
    [
        (
            "FIVE_SEGMENT_V1",
            FIVE_SEGMENT_RULE_SHEET,
            "五段式",
            "八段式（兼容）",
        ),
        (
            "LEGACY_8",
            EXPRESSION_RULE_SHEET,
            "八段式（兼容）",
            "五段式",
        ),
    ],
)
def test_formal_config_report_names_only_the_selected_workbooks_and_sheets(
    schema: str, expression_sheet: str, mode_label: str, opposite_label: str
) -> None:
    before = (_sha256(COMMON_CONFIG), _sha256(COMPARISON_CONFIG))

    result = _api_for_config(COMMON_CONFIG, COMPARISON_CONFIG, schema).check_central_expression_config()

    assert result["schema"] == schema
    assert result["source_workbooks"] == [
        str(COMMON_CONFIG.resolve()),
        str(COMPARISON_CONFIG.resolve()),
    ]
    assert result["active_sheets"] == [INDICATOR_SHEET, ACTION_RULE_SHEET, expression_sheet]
    report = result["report_text"]
    assert f"执行比较 3.1（{mode_label}）" in report
    assert COMMON_CONFIG.name in report
    assert COMPARISON_CONFIG.name in report
    assert f"{INDICATOR_SHEET}、{ACTION_RULE_SHEET}、{expression_sheet}" in report
    assert opposite_label not in report
    assert (_sha256(COMMON_CONFIG), _sha256(COMPARISON_CONFIG)) == before


@pytest.mark.skipif(not FORMAL_CONFIGS_READY, reason="缺少正式 3.0/3.1 配置工作簿")
def test_formal_action_rule_count_and_enabled_actions_are_accounted_for() -> None:
    config = load_central_config((COMMON_CONFIG, COMPARISON_CONFIG))
    stats: dict = {}

    errors = check_action_rules(config, stats_out=stats)

    assert stats["total"] == 5082
    assert stats["enabled"] == 4246
    assert stats["disabled"] == 836
    assert stats["recognized"] == stats["enabled"]
    assert errors == []


def test_unsupported_enabled_action_reports_sheet_row_rule_and_original_action() -> None:
    action = "此动作不属于系统可识别动作"
    config = CentralConfig(
        path=Path("in-memory.xlsx"),
        action_rules=[
            {
                "规则编号": "TEST-ACTION-UNKNOWN",
                "类型": "特殊阈值",
                "指标代码": "TEST-001",
                "备注": action,
                "禁用": "",
                "__行号__": "47",
            }
        ],
    )
    stats: dict = {}

    errors = check_action_rules(config, stats_out=stats)

    assert len(errors) == 1
    assert ACTION_RULE_SHEET in errors[0]
    assert "第47行" in errors[0]
    assert "TEST-ACTION-UNKNOWN" in errors[0]
    assert action in errors[0]
    assert stats["enabled"] == 1
    assert stats["recognized"] == 0


def test_disabled_unsupported_action_is_ignored() -> None:
    config = CentralConfig(
        path=Path("in-memory.xlsx"),
        action_rules=[
            {
                "规则编号": "TEST-ACTION-DISABLED",
                "类型": "特殊阈值",
                "指标代码": "TEST-001",
                "备注": "此动作不属于系统可识别动作",
                "禁用": "是",
                "__行号__": "48",
            }
        ],
    )
    stats: dict = {}

    errors = check_action_rules(config, stats_out=stats)

    assert errors == []
    assert stats["total"] == 1
    assert stats["enabled"] == 0
    assert stats["disabled"] == 1
    assert stats["recognized"] == 0


def test_blank_cumulative_action_and_representative_runtime_actions_are_recognized() -> None:
    actions = [
        "指标不应有数，需说明。",
        "指标应相等核查。",
        "指标相除结果需核查。",
        "指标相减结果需核查。",
        "指标太小核查。",
    ]
    rules = [
        {
            "规则编号": "TEST-CUMULATIVE",
            "类型": "累计不应下降",
            "指标代码": "TEST-ACC-001",
            "备注": "",
            "禁用": "",
            "__行号__": "49",
        },
        *[
            {
                "规则编号": f"TEST-RUNTIME-{index}",
                "类型": "特殊阈值",
                "指标代码": f"TEST-{index:03}",
                "备注": action,
                "禁用": "",
                "__行号__": str(50 + index),
            }
            for index, action in enumerate(actions, start=1)
        ],
    ]
    stats: dict = {}

    errors = check_action_rules(
        CentralConfig(path=Path("in-memory.xlsx"), action_rules=rules),
        stats_out=stats,
    )

    assert errors == []
    assert stats["cumulative"] == 1
    assert stats["special"] == len(actions)
    assert stats["recognized"] == 1 + len(actions)
    assert all(supports_special_action(action) for action in actions)


@pytest.mark.skipif(not FORMAL_CONFIGS_READY, reason="缺少正式 3.0/3.1 配置工作簿")
def test_webapi_reports_an_unsupported_action_from_a_disposable_workbook_copy(
    tmp_path: Path,
) -> None:
    common_copy = tmp_path / COMMON_CONFIG.name
    comparison_copy = tmp_path / COMPARISON_CONFIG.name
    shutil.copy2(COMMON_CONFIG, common_copy)
    shutil.copy2(COMPARISON_CONFIG, comparison_copy)

    book = load_workbook(comparison_copy)
    try:
        sheet = book[ACTION_RULE_SHEET]
        header_row = None
        columns: dict[str, int] = {}
        for row_index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            candidate = {str(value).strip(): index + 1 for index, value in enumerate(values) if value is not None}
            if {"规则编号", "规则动作", "指标代码", "来源分组", "启用"} <= candidate.keys():
                header_row = row_index
                columns = candidate
                break
        assert header_row is not None

        target_row = None
        for row_index in range(header_row + 1, sheet.max_row + 1):
            rule_id = str(sheet.cell(row_index, columns["规则编号"]).value or "").strip()
            action = str(sheet.cell(row_index, columns["规则动作"]).value or "").strip()
            indicator = str(sheet.cell(row_index, columns["指标代码"]).value or "").strip()
            group = str(sheet.cell(row_index, columns["来源分组"]).value or "").strip()
            enabled = str(sheet.cell(row_index, columns["启用"]).value or "").strip()
            if rule_id and action and indicator and group != "当年累计" and enabled not in {"否", "0", "false", "False"}:
                target_row = row_index
                break
        assert target_row is not None
        rule_id = str(sheet.cell(target_row, columns["规则编号"]).value).strip()
        invalid_action = "集成测试不可识别动作"
        sheet.cell(target_row, columns["规则动作"], invalid_action)
        book.save(comparison_copy)
    finally:
        book.close()

    result = _api_for_config(common_copy, comparison_copy, "FIVE_SEGMENT_V1").check_central_expression_config()

    assert len(result["action_errors"]) == 1
    error = result["action_errors"][0]
    assert ACTION_RULE_SHEET in error
    assert f"第{target_row}行" in error
    assert rule_id in error
    assert invalid_action in error
    assert result["total_issues"] >= 1
    assert error in result["report_text"]
