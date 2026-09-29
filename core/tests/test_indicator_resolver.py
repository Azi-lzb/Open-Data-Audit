"""指标定位与统一结果构建测试。

覆盖：IndicatorResolver 的 COM 语义对齐（合并锚点、多级表头、数值过滤）、
MessageResolver 的「批注优先、公式兜底」、CommentReader 双后端
（openpyxl / Direct OOXML）一致性、统一 Builder 的业务字段。
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "core" / "src"))

from base_audit.indicator_resolver import IndicatorResolver
from base_audit.models import CopyRange


def _build_book() -> Workbook:
    """模拟“存量单位贷款信息”表头结构：

    - A1:A8 纵向合并 = 借款人证件类型（父级表头，值只在 A1，子行是
      MergedCell——合并区内的 A3/A5/A7 即使赋值也不可见）
    - B 列行子标签（统一社会信用代码/其他），C 列有数值格
    - D3:E3 横向合并 = 余额（分组表头），D4/E4 = 本期余额/环比增幅
    """
    book = Workbook()
    ws = book.active
    ws.title = "存量单位贷款信息"
    ws["A1"] = "借款人证件类型"
    ws["B5"] = "统一社会信用代码"
    ws["B7"] = "其他"
    ws["C5"] = 1234.5
    ws["C4"] = "1,234"
    ws["D3"] = "余额"
    ws["D4"] = "本期余额"
    ws["E4"] = "环比增幅"
    ws.merge_cells("A1:A8")
    ws.merge_cells("D3:E3")
    return book


def _structure() -> list[CopyRange]:
    return [CopyRange("存量单位贷款信息", "A2:E8")]


class IndicatorResolverTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.book = _build_book()
        cls.resolver = IndicatorResolver(cls.book, _structure())

    def test_merged_parent_header_visible_to_children(self):
        """A 列纵向合并的父级表头对子行可见（MergeArea 锚点语义）。"""
        # 触发格 D5：行向左扫 = A5(锚点 A1=借款人证件类型)、B5；
        # 列向上扫 = D3(锚点=余额)、D4(本期余额)，行序自上而下。
        self.assertEqual(
            self.resolver.resolve("存量单位贷款信息", 5, 4),
            "借款人证件类型_统一社会信用代码｜余额_本期余额",
        )

    def test_multi_level_left_labels_in_area_order(self):
        """行指标从结构区域左端向右收集去重，顺序与 COM 一致。"""
        self.assertEqual(
            self.resolver.resolve("存量单位贷款信息", 7, 5),
            "借款人证件类型_其他｜余额_环比增幅",
        )

    def test_numeric_cells_excluded(self):
        """数值与千分位数字串（C5/C4）不能进入指标。"""
        text = self.resolver.resolve("存量单位贷款信息", 5, 4)
        self.assertNotIn("1234", text)
        self.assertNotIn("1,234", text)

    def test_no_structure_areas_returns_empty(self):
        """无表结构区域时与 COM 一致返回空串。"""
        resolver = IndicatorResolver(self.book, [])
        self.assertEqual(resolver.resolve("存量单位贷款信息", 5, 4), "")

    def test_missing_sheet_returns_empty(self):
        self.assertEqual(self.resolver.resolve("不存在的表", 5, 4), "")

    def test_multiple_areas_use_union_start(self):
        """多个结构区域取 min_row/min_col 作为扫描起点（COM 语义）。"""
        resolver = IndicatorResolver(
            self.book, [CopyRange("存量单位贷款信息", "B2:E8"),
                        CopyRange("存量单位贷款信息", "A3:D8")])
        # min_col=1（第二个区域），行指标从 A 列开始收集
        self.assertEqual(
            resolver.resolve("存量单位贷款信息", 5, 4),
            "借款人证件类型_统一社会信用代码｜余额_本期余额",
        )

    def test_resolve_seconds_recorded(self):
        resolver = IndicatorResolver(self.book, _structure())
        resolver.resolve("存量单位贷款信息", 5, 4)
        self.assertGreaterEqual(resolver.resolve_seconds, 0.0)


class MessageResolverTests(unittest.TestCase):
    def test_comment_preferred(self):
        from base_audit.conditional_context import CellComment, MessageResolver

        comment = CellComment(cell="C26", author="l z b", text="这是一个测试",
                              formatted="l z b:\n这是一个测试")
        self.assertEqual(
            MessageResolver.resolve(comment, "条件格式规则：C26>0"),
            "l z b:\n这是一个测试",
        )

    def test_formula_fallback_when_no_comment(self):
        from base_audit.conditional_context import MessageResolver

        self.assertEqual(
            MessageResolver.resolve(None, "条件格式规则：D26>0"),
            "条件格式规则：D26>0",
        )

    def test_blank_comment_body_falls_back(self):
        from base_audit.conditional_context import CellComment, MessageResolver

        comment = CellComment(cell="A1", author="x", text="", formatted="x:")
        self.assertEqual(
            MessageResolver.resolve(comment, "条件格式规则：A1>0"),
            "条件格式规则：A1>0",
        )

    def test_standardize_author_line(self):
        """既有业务格式「作者:\n正文」拆分为 author/text，formatted 原样保留。"""
        from base_audit.conditional_context import _standardize_comment

        comment = _standardize_comment("C26", "l z b:\n这是一个测试", "l z b")
        self.assertEqual(comment.author, "l z b")
        self.assertEqual(comment.text, "这是一个测试")
        self.assertEqual(comment.formatted, "l z b:\n这是一个测试")

    def test_standardize_without_author_line(self):
        from base_audit.conditional_context import _standardize_comment

        comment = _standardize_comment("A1", "不应有数，请核实", "Administrator")
        self.assertEqual(comment.author, "Administrator")
        self.assertEqual(comment.text, "不应有数，请核实")


class CommentReaderTests(unittest.TestCase):
    """openpyxl 与 Direct OOXML 两后端对同一文件产出一致。"""

    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory(prefix="cctx-")
        cls.path = Path(cls._tmp.name) / "甲银行_审核版.xlsx"
        book = Workbook()
        ws = book.active
        ws.title = "存量单位贷款信息"
        ws["B2"] = 100
        ws["C26"] = 2831236.25
        from openpyxl.comments import Comment

        ws["C26"].comment = Comment("l z b:\n这是一个测试", "l z b")
        ws["D26"].comment = Comment("不应有数，请核实", "Administrator")
        book.save(cls.path)
        book.close()

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_direct_ooxml_reader(self):
        from base_audit.conditional_context import OoxmlCommentReader

        comments = OoxmlCommentReader(self.path).read()
        self.assertIn("存量单位贷款信息", comments)
        by_cell = comments["存量单位贷款信息"]
        c26 = by_cell["C26"]
        self.assertEqual(c26.author, "l z b")
        self.assertEqual(c26.text, "这是一个测试")
        self.assertEqual(c26.formatted, "l z b:\n这是一个测试")
        self.assertEqual(c26.comment_type, "note")
        self.assertEqual(by_cell["D26"].text, "不应有数，请核实")
        self.assertNotIn("B2", by_cell)

    def test_openpyxl_reader_matches_direct_ooxml(self):
        from base_audit.conditional_context import (
            OoxmlCommentReader, OpenpyxlCommentReader,
        )

        direct = OoxmlCommentReader(self.path).read()["存量单位贷款信息"]
        book = load_workbook(self.path, data_only=True)
        try:
            reader = OpenpyxlCommentReader(book)
            for cell in ("C26", "D26", "B2"):
                via_ooxml = direct.get(cell)
                via_openpyxl = reader.comment_for("存量单位贷款信息", cell)
                if via_ooxml is None:
                    self.assertIsNone(via_openpyxl)
                else:
                    self.assertIsNotNone(via_openpyxl)
                    self.assertEqual(via_ooxml.formatted, via_openpyxl.formatted)
                    self.assertEqual(via_ooxml.author, via_openpyxl.author)
        finally:
            book.close()

    def test_context_build_prefers_direct_ooxml_backend(self):
        from base_audit.conditional_context import ConditionalFormatContext

        book = load_workbook(self.path, data_only=True)
        try:
            context = ConditionalFormatContext.build(
                workbook_path=self.path, static_book=book,
                structure_ranges=[CopyRange("存量单位贷款信息", "A1:E8")])
            self.assertEqual(context.comment_backend, "direct_ooxml")
            comment = context.comment_for("存量单位贷款信息", "C26")
            self.assertIsNotNone(comment)
            self.assertEqual(comment.formatted, "l z b:\n这是一个测试")
        finally:
            book.close()


class ResultsToIssuesIntegrationTests(unittest.TestCase):
    """results_to_issues 经统一 Builder 产出完整 check_field 与批注优先 message。"""

    def test_issue_check_field_and_formula_fallback(self):
        from base_audit.conditional_engine import (
            ConditionalFormatResult, results_to_issues,
        )

        book = _build_book()
        hit = ConditionalFormatResult(
            sheet="存量单位贷款信息", cell="D5", triggered=True,
            message="条件格式规则：D5>0",
        )
        ranges = [CopyRange("存量单位贷款信息", "D5")]
        issues = results_to_issues(
            [hit], workbook_path=Path("甲银行.xlsx"), ranges=ranges,
            structure_ranges=_structure(), period="2026-07", batch_id="T",
            audit_time="t", source_file=Path("甲银行_在线核查表.xlsx"),
            static_book=book,
        )
        self.assertEqual(len(issues), 1)
        issue = issues[0]
        self.assertEqual(
            issue.check_field,
            "借款人证件类型_统一社会信用代码｜余额_本期余额")
        self.assertEqual(issue.message, "条件格式规则：D5>0")
        self.assertEqual(issue.detail, "条件格式规则：D5>0")
        # identity 保留“校验指标未识别”兜底语义（excel_com 同口径）
        self.assertIn("借款人证件类型", issue.issue_id)


if __name__ == "__main__":
    unittest.main()
