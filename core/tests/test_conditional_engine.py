"""条件格式规则触发检测引擎测试：四态求值、priority/stopIfTrue、统一输出。

背景：旧实现「颜色变化 = 触发」；新实现「规则是否成立 = 触发」，颜色只是
辅助信息。真实数据 A/B 验收见 test_conditional_compare.py。
"""

from __future__ import annotations

import sys
from pathlib import Path
from tempfile import TemporaryDirectory

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "core" / "src"))

import pytest
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule, Rule
from openpyxl.styles import PatternFill
from openpyxl.styles.differential import DifferentialStyle

from base_audit.conditional_engine import (
    CONFIDENCE_PARTIAL,
    CONDITIONAL_FORMAT_MODE_NATIVE,
    CONDITIONAL_FORMAT_MODE_OOXML,
    ConditionalFormatResult,
    OoxmlConditionalFormatEngine,
    create_conditional_format_engine,
)
from base_audit.conditional_format import (
    STATUS_ERROR,
    STATUS_FALSE,
    STATUS_TRUE,
    STATUS_UNSUPPORTED,
    evaluate_cellis_rule,
    evaluate_expression_condition,
)
from base_audit.models import CopyRange


RED_FILL = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")


def _rule(**kwargs):
    return Rule(**kwargs)


def _dxf_rule(**kwargs) -> Rule:
    kwargs.setdefault("dxf", DifferentialStyle(fill=RED_FILL))
    return Rule(**kwargs)


# ---------------------------------------------------------------------------
# 四态求值器
# ---------------------------------------------------------------------------


class _Sheet:
    """evaluate_cellis_rule / evaluate_expression_condition 的最小工作表桩。"""

    def __init__(self, values: dict[str, object]):
        self._values = values

    def cell(self, row: int, column: int, value=None):
        from openpyxl.utils.cell import get_column_letter

        return SimpleNamespaceCell(value=self._values.get(
            f"{get_column_letter(column)}{row}"))


class SimpleNamespaceCell:
    def __init__(self, value):
        self.value = value


class TestCellIsEvaluation:
    def test_numeric_comparison_true_and_false(self):
        result = evaluate_cellis_rule(
            "greaterThan", ["0"], 5, resolve_operand=lambda _f, _i: 0)
        assert result.status == STATUS_TRUE
        result = evaluate_cellis_rule(
            "greaterThan", ["0"], -1, resolve_operand=lambda _f, _i: 0)
        assert result.status == STATUS_FALSE

    def test_unsupported_operator_is_not_false(self):
        result = evaluate_cellis_rule("containsText", ["abc"], "abc",
                                      resolve_operand=lambda _f, _i: "abc")
        assert result.status == STATUS_UNSUPPORTED

    def test_expression_operand_is_unsupported_not_false(self):
        # 操作数是公式表达式（如 MAX(D:D)）时无法离线求值：必须 UNSUPPORTED。
        result = evaluate_cellis_rule(
            "greaterThan", ["MAX(D:D)"], 5,
            resolve_operand=lambda _f, _i: (_ for _ in ()).throw(
                __import__("base_audit.conditional_format", fromlist=["UnsupportedFormulaError"])
                .UnsupportedFormulaError("表达式操作数")),
        )
        assert result.status == STATUS_UNSUPPORTED

    def test_between(self):
        resolve = lambda _f, i: [1, 10][i]
        assert evaluate_cellis_rule("between", ["1", "10"], 5,
                                    resolve_operand=resolve).status == STATUS_TRUE
        assert evaluate_cellis_rule("between", ["1", "10"], 11,
                                    resolve_operand=resolve).status == STATUS_FALSE
        assert evaluate_cellis_rule("notBetween", ["1", "10"], 11,
                                    resolve_operand=resolve).status == STATUS_TRUE

    def test_empty_cell_does_not_trigger(self):
        assert evaluate_cellis_rule("greaterThan", ["0"], "",
                                    resolve_operand=lambda _f, _i: 0).status == STATUS_FALSE


class TestExpressionEvaluation:
    def test_true_false(self):
        sheet = _Sheet({"C2": 5, "D2": 3})
        assert evaluate_expression_condition(
            "AND($D$2>0,C2>$D$2)", sheet, 2, 3, 2, 3).status == STATUS_TRUE
        assert evaluate_expression_condition(
            "AND($D$2>0,C2>$D$2)", sheet, 2, 3, 2, 3).triggered is True

    def test_unsupported_function(self):
        result = evaluate_expression_condition("OFFSET(C1,1,0)>0", _Sheet({}), 1, 3, 2, 3)
        assert result.status == STATUS_UNSUPPORTED
        assert "OFFSET" in result.reason

    def test_garbage_formula_is_error_not_false(self):
        result = evaluate_expression_condition("(((", _Sheet({}), 1, 1, 1, 1)
        assert result.status == STATUS_ERROR


# ---------------------------------------------------------------------------
# OOXML 引擎：扫描、优先级、stopIfTrue、颜色独立性
# ---------------------------------------------------------------------------


def _write_workbook(path: Path, rules, values: dict[str, object],
                    sheet_name: str = "数据") -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    for ref, value in values.items():
        sheet[ref] = value
    for target, rule in rules:
        sheet.conditional_formatting.add(target, rule)
    book.save(path)
    book.close()


def _extract(path: Path, area="B2:B3", sheet="数据"):
    engine = OoxmlConditionalFormatEngine()
    return engine.extract(path, [CopyRange(sheet, area)], [])


class TestOoxmlEngine:
    def test_triggering_cell_reported_with_result_fields(self, tmp_path):
        path = tmp_path / "机构A.xlsx"
        _write_workbook(path, [("B2:B3", CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL))],
                        {"B2": 10, "B3": 0})
        extraction = _extract(path)
        assert [(r.cell, r.triggered, r.rule_type) for r in extraction.results] == [("B2", True, "cellis")]
        result = extraction.results[0]
        assert result.source == "OOXML"
        assert result.message == "条件格式规则：B2>0"
        assert result.color is not None    # 颜色作为辅助信息输出
        assert result.formula == "0"

    def test_color_missing_does_not_affect_triggered(self, tmp_path):
        # 规则没有 dxf 填充：颜色为 None，但规则成立仍必须 triggered=True。
        path = tmp_path / "机构B.xlsx"
        _write_workbook(path, [("B2:B3", CellIsRule(operator="greaterThan", formula=["0"]))],
                        {"B2": 10})
        extraction = _extract(path)
        assert len(extraction.results) == 1
        assert extraction.results[0].triggered is True
        assert extraction.results[0].color is None

    def test_higher_priority_wins_over_lower(self, tmp_path):
        # 规则1（priority 1，高优先级）与规则2（priority 2）同时命中：
        # 报告最高优先级规则的描述，而不是 priority 数值最大者。
        path = tmp_path / "机构C.xlsx"
        rule_high = _dxf_rule(type="cellIs", operator="greaterThan", formula=["0"],
                              priority=1, stopIfTrue=True)
        rule_low = _dxf_rule(type="cellIs", operator="lessThan", formula=["100"],
                             priority=2)
        _write_workbook(path, [("B2:B2", rule_high), ("B2:B2", rule_low)], {"B2": 5})
        extraction = _extract(path, area="B2:B2")
        assert len(extraction.results) == 1
        assert extraction.results[0].message == "条件格式规则：B2>0"

    def test_unmatched_stop_if_true_does_not_block_lower_rule(self, tmp_path):
        # 规则1 未命中时其 stopIfTrue 不生效：低优先级规则照常判定。
        path = tmp_path / "机构D.xlsx"
        rule_high = _dxf_rule(type="cellIs", operator="greaterThan", formula=["1000"],
                              priority=1, stopIfTrue=True)
        rule_low = _dxf_rule(type="cellIs", operator="greaterThan", formula=["10"],
                             priority=2)
        _write_workbook(path, [("B2:B2", rule_high), ("B2:B2", rule_low)], {"B2": 50})
        extraction = _extract(path, area="B2:B2")
        assert [r.message for r in extraction.results] == ["条件格式规则：B2>10"]

    def test_unsupported_higher_priority_marks_partial_confidence(self, tmp_path):
        # 更高优先级处存在不可离线求值的规则（colorScale）：低优先级命中仍报告，
        # 但置信度降为 PARTIAL，且 unsupported 清单可查。
        path = tmp_path / "机构E.xlsx"
        gradient = Rule(type="colorScale", priority=1, colorScale=None)
        rule_low = _dxf_rule(type="cellIs", operator="greaterThan", formula=["0"],
                             priority=2)
        _write_workbook(path, [("B2:B2", gradient), ("B2:B2", rule_low)], {"B2": 5})
        extraction = _extract(path, area="B2:B2")
        assert extraction.unsupported and extraction.unsupported[0].rule_type == "colorScale"
        assert len(extraction.results) == 1
        assert extraction.results[0].confidence == CONFIDENCE_PARTIAL

    def test_expression_rule_triggers_with_relative_shift(self, tmp_path):
        path = tmp_path / "机构F.xlsx"
        _write_workbook(path, [("C2:C2", FormulaRule(formula=["AND($D$2>0,C2>$D$2)"], fill=RED_FILL))],
                        {"C2": 5, "D2": 3})
        extraction = _extract(path, area="C2:C2")
        assert [r.triggered for r in extraction.results] == [True]
        assert extraction.results[0].rule_type == "expression"


# ---------------------------------------------------------------------------
# 工厂：平台与模式分派
# ---------------------------------------------------------------------------


class TestEngineFactory:
    def test_native_pipeline_forces_ooxml(self):
        engine, honoured = create_conditional_format_engine(
            CONDITIONAL_FORMAT_MODE_NATIVE, pipeline="native")
        assert isinstance(engine, OoxmlConditionalFormatEngine)
        assert honoured is False    # NATIVE 请求未被采纳（UOS 固定 OOXML）

    def test_com_pipeline_respects_ooxml_mode(self):
        engine, honoured = create_conditional_format_engine(
            CONDITIONAL_FORMAT_MODE_OOXML, pipeline="com")
        assert isinstance(engine, OoxmlConditionalFormatEngine)
        assert honoured is True

    def test_com_pipeline_default_is_native(self):
        engine, _honoured = create_conditional_format_engine(
            CONDITIONAL_FORMAT_MODE_NATIVE, pipeline="com", session=object())
        assert not isinstance(engine, OoxmlConditionalFormatEngine)

    def test_com_pipeline_native_requires_session(self):
        with pytest.raises(RuntimeError):
            create_conditional_format_engine(
                CONDITIONAL_FORMAT_MODE_NATIVE, pipeline="com", session=None)


# ---------------------------------------------------------------------------
# 统一业务输出：结果 → Issue（AuditResult）字段口径
# ---------------------------------------------------------------------------


class TestResultsToIssues:
    def test_issue_fields_match_legacy_contract(self, tmp_path):
        from base_audit.conditional_engine import results_to_issues
        from base_audit.models import Issue

        path = tmp_path / "机构G.xlsx"
        _write_workbook(path, [("B2:B2", CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL))],
                        {"B2": 10, "A2": "各项贷款"})
        engine = OoxmlConditionalFormatEngine()
        issues, extraction = engine.extract_issues(
            path, [CopyRange("数据", "B2:B2")], [CopyRange("数据", "A2:A2")],
            period="2026-08", batch_id="B1", audit_time="2026-08-31 10:00:00",
            source_file=path,
        )
        assert len(issues) == 1
        issue = issues[0]
        assert isinstance(issue, Issue)
        assert issue.sheet_name == "数据"
        assert issue.rule_id == "条件格式填充"
        assert issue.severity == "条件格式触发"
        assert issue.target_cell == "B2"
        assert issue.detail == "条件格式规则：B2>0"
        assert issue.check_field == "各项贷款"
