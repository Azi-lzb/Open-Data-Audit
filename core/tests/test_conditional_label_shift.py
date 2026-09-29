"""条件格式规则描述的「平移到对应单元格」渲染测试（Office 口径）。

对应真机问题：J36/J45 等命中多区域规则时，描述曾显示锚点原式
``AND(J8<>"",OR(J8<-30,J8>30))``，应显示平移到当前格的公式。
"""

from __future__ import annotations

import sys
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402

from base_audit.conditional_engine import (  # noqa: E402
    ConditionalFormatExtraction,
    OoxmlConditionalFormatEngine,
    rule_label,
)
from base_audit.conditional_format import translate_formula_refs  # noqa: E402
from base_audit.cf_reader import OoxmlConditionalRuleReader  # noqa: E402
from base_audit.conditional_evaluators import (  # noqa: E402
    ComFormulaEvaluator,
    EVALUATOR_PYTHON,
    ConditionalFormatService,
    _coerce_com_truth,
    _shift_refs,
    rule_label_normalized,
)


class _Rule:
    """最小规则桩：expression 规则。"""

    def __init__(self, formula: str, priority: int = 1, stop_if: bool = False):
        self.type = "expression"
        self.formula = [formula]
        self.priority = priority
        self.stopIfTrue = stop_if
        self.dxfId = None
        self.operator = None


class _Cell:
    def __init__(self, value):
        self.value = value
        self.comment = None


class _Sheet:
    """最小工作表桩：坐标 → 取值。"""

    def __init__(self, cells: dict[tuple[int, int], object]):
        self._cells = cells
        self.title = "汇总表"

    def cell(self, row: int, column: int) -> _Cell:
        return _Cell(self._cells.get((row, column)))


def test_translate_shifts_row_relative_ref():
    assert translate_formula_refs('AND($J8<>"",OR($J8<-30,$J8>30))', 28, 0) == \
        'AND($J36<>"",OR($J36<-30,$J36>30))'


def test_translate_shifts_plain_refs_both_dims():
    assert translate_formula_refs("B2>C2", 0, 1) == "C2>D2"
    assert translate_formula_refs("A1:B2", 1, 0) == "A2:B3"


def test_translate_keeps_locked_dims():
    assert translate_formula_refs("$F$8+F8", 2, 0) == "$F$8+F10"


def test_translate_keeps_string_literals_and_function_names():
    assert translate_formula_refs('"A1 文本"', 1, 1) == '"A1 文本"'
    assert translate_formula_refs("LOG10(A1)", 1, 0) == "LOG10(A2)"


def test_translate_out_of_sheet_becomes_ref_error():
    assert translate_formula_refs("A1", -5, 0) == "#REF!"


def test_translate_zero_shift_returns_original():
    text = 'AND($J8<>"",OR($J8<-30,$J8>30))'
    assert translate_formula_refs(text, 0, 0) == text


def test_rule_label_renders_shifted_formula():
    rule = _Rule('AND($J8<>"",OR($J8<-30,$J8>30))')
    label = rule_label(rule, address="J36", dr=28, dc=0)
    assert label == '条件格式规则：AND(J36<>"",OR(J36<-30,J36>30))'


def test_rule_label_zero_shift_keeps_anchored_text():
    rule = _Rule('AND($J8<>"",OR($J8<-30,$J8>30))')
    assert rule_label(rule, address="J8") == \
        '条件格式规则：AND(J8<>"",OR(J8<-30,J8>30))'


def test_scan_area_description_shifts_to_hit_cell():
    """多区域规则：J9（第一区域 J8:J19 内）命中时描述应平移为 $J9 口径。"""
    ws = _Sheet({
        (8, 10): 0.0,      # J8：锚点自身（0 不触发）
        (9, 10): -40.0,    # J9：触发，描述应为 J9 口径
    })
    engine = OoxmlConditionalFormatEngine()
    extraction = ConditionalFormatExtraction()
    rule = _Rule('AND($J8<>"",OR($J8<-30,$J8>30))', priority=1)
    spec = {
        "rule": rule,
        # 多区域 sqref（真机模板的保存顺序）：锚点取所有区域的最小行列 = J8
        "bounds": [(10, 35, 10, 47), (10, 8, 10, 19), (10, 21, 10, 33)],
        "anchor_row": 8, "anchor_col": 10, "colour": None,
        "priority": 1, "order": 0, "stop_if": False,
        "sqref": "J35:J47 J8:J19 J21:J33", "unsupported": None,
    }
    results = engine._scan_area(ws, "J8:J19", [spec], extraction)
    by_cell = {item.cell: item for item in results}
    assert "J9" in by_cell and "J8" not in by_cell
    assert by_cell["J9"].message == '条件格式规则：AND(J9<>"",OR(J9<-30,J9>30))'
    # J36/J45 在扫描区域（J8:J19）之外，不入结果：边界由扫描区域决定。
    assert all(item.cell in ("J8", "J9") for item in results)


def test_applies_to_fallback_expands_every_area():
    """无名区域回退必须展开全部区域，不得只取第一段。"""
    from openpyxl.worksheet.cell_range import MultiCellRange

    engine = OoxmlConditionalFormatEngine()

    class _CF:
        def __init__(self, sqref):
            self.sqref = MultiCellRange(sqref)
            self.rules = [_Rule("A1>0")]

    class _WS:
        title = "汇总表"

        def __init__(self):
            self.conditional_formatting = [_CF("J35:J47 J8:J19 J21:J33")]

    class _Book:
        worksheets = [_WS()]

    ranges = engine._applies_to_ranges(_Book())
    addresses = sorted(area.address for area in ranges)
    assert addresses == ["J21:J33", "J35:J47", "J8:J19"]


def test_buluo_multi_area_only_j10_triggers():
    """多区域条件格式案例（任务书十二）：J 列多区域规则，仅 J10 应触发。

    真实值（Office 重算后缓存）：J8=7.78、J9=14.78、J10=-50、J11=-3.17、
    J21=7.78、J26=7.16、J27=7.16、J37=8.89；阈值 |偏差|>30。
    按 OOXML 语义（锚点 = 全区域最小行列 = J8，逐格平移自引用），
    触发集合必须恰好是 {J10}——J12–J19/J22–J25 空格与阈内格都不得触发。
    """
    values = {
        (8, 10): 7.77777777777778,
        (9, 10): 14.7826086956522,
        (10, 10): -50.0,
        (11, 10): -3.17460317460317,
        (21, 10): 7.77777777777778,
        (26, 10): 7.16071568091257,
        (27, 10): 7.16435411373852,
        (37, 10): 8.88517745302714,
    }
    ws = _Sheet(values)
    engine = OoxmlConditionalFormatEngine()
    extraction = ConditionalFormatExtraction()
    spec = {
        "rule": _Rule('AND($J8<>"",OR($J8<-30,$J8>30))'),
        "bounds": [(10, 8, 10, 19), (10, 21, 10, 33), (10, 35, 10, 47)],
        "anchor_row": 8, "anchor_col": 10, "colour": None,
        "priority": 1, "order": 0, "stop_if": False,
        "sqref": "J35:J47 J8:J19 J21:J33", "unsupported": None,
    }
    results = engine._scan_area(ws, "J8:J47", [spec], extraction)
    flagged = sorted(item.cell for item in results)
    assert flagged == ["J10"], flagged
    # 描述按当前格平移（J10 命中时锚点 J8 → 平移 +2 → $J10）。
    assert results[0].message == '条件格式规则：AND(J10<>"",OR(J10<-30,J10>30))'


def _build_raw_order_multi_area_book(path: Path) -> None:
    """构造源文件常见的 J35 在前 sqref，验证规范锚点不依赖 XML 顺序。"""
    book = Workbook()
    ws = book.active
    ws.title = "汇总表"
    for row in range(8, 48):
        ws.cell(row=row, column=10, value=None)
    for row, value in {8: 7.0, 10: -50.0, 36: -40.0, 37: 8.0, 45: 50.0}.items():
        ws.cell(row=row, column=10, value=value)
    ws.conditional_formatting.add(
        "J8:J19 J21:J33 J35:J47",
        FormulaRule(formula=['AND($J8<>"",OR($J8<-30,$J8>30))']),
    )
    book.save(path)
    book.close()

    # openpyxl canonicalizes MultiCellRange order on save. Rewrite only the
    # worksheet sqref in a temporary copy to reproduce the raw source order.
    temp = path.with_name(path.stem + ".raw-order.xlsx")
    with ZipFile(path, "r") as source, ZipFile(temp, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(
                    b'sqref="J8:J19 J21:J33 J35:J47"',
                    b'sqref="J35:J47 J8:J19 J21:J33"',
                )
            target.writestr(item, data)
    os.replace(temp, path)


def test_direct_reader_and_python_use_order_independent_canonical_anchor():
    with TemporaryDirectory(prefix="cf-anchor-") as folder:
        workbook = Path(folder) / "raw-order.xlsx"
        _build_raw_order_multi_area_book(workbook)
        reader = OoxmlConditionalRuleReader().read(workbook)
        rule = next(rule for rule in reader.rules if rule.sheet == "汇总表")
        assert rule.anchor_cell == "J8"

        issues, _stats = ConditionalFormatService(EVALUATOR_PYTHON).extract_issues(
            workbook_path=workbook, ranges=[], structure_ranges=[],
            period="2026-08", batch_id="anchor", audit_time="",
            source_file=workbook,
        )
        assert {issue.target_cell for issue in issues} == {"J10", "J36", "J45"}


def test_normalized_rule_label_is_relative_to_target_cell():
    from base_audit.cf_reader import NormalizedConditionalRule

    rule = NormalizedConditionalRule(
        rule_id="t", sheet="汇总表", applies_to="J8:J19 J21:J33 J35:J47",
        rule_type="expression", operator="",
        formula1='AND($J8<>"",OR($J8<-30,$J8>30))', formula2="",
        priority=1, stop_if_true=True, anchor_cell="J8", dxf_id=None,
        bounds=((10, 8, 10, 19), (10, 21, 10, 33), (10, 35, 10, 47)),
    )
    assert rule_label_normalized(rule, target_cell="J36") == \
        '条件格式规则：AND(J36<>"",OR(J36<-30,J36>30))'


def test_com_formula_shift_does_not_skip_comma_argument():
    formula = 'AND($J8<>"",OR($J8<-30,$J8>30))'
    assert _shift_refs(formula, 28, 0) == \
        'AND($J36<>"",OR($J36<-30,$J36>30))'


def test_com_truth_rejects_arbitrary_objects_and_accepts_boolean_scalars():
    assert _coerce_com_truth(True) is True
    assert _coerce_com_truth(False) is False
    assert _coerce_com_truth(1) is True
    assert _coerce_com_truth(-1) is True
    assert _coerce_com_truth(2) is False
    assert _coerce_com_truth("TRUE") is False
    assert _coerce_com_truth(object()) is False


def test_com_evaluate_does_not_broadcast_one_trigger_to_rule_area():
    """COM_EVALUATE 必须按目标单元格求值，不能复用规则级真假。"""

    from base_audit.cf_reader import CellValueProvider, CfRuleReadResult, NormalizedConditionalRule

    class _ComSheet:
        def __init__(self):
            self.calls = []

        def Evaluate(self, expression):
            self.calls.append(expression)
            # 仅 J10 的平移公式满足条件；J8/J9 必须保持 FALSE。
            return "J10" in str(expression)

    class _ComBook:
        def __init__(self, sheet):
            self.sheet = sheet

        def Worksheets(self, name):
            assert name == "汇总表"
            return self.sheet

    class _Session:
        engine_name = "EXCEL"

    sheet = _ComSheet()
    rule = NormalizedConditionalRule(
        rule_id="r1", sheet="汇总表", applies_to="J8:J10",
        rule_type="expression", operator="",
        formula1='AND($J8<>"",OR($J8<-30,$J8>30))', formula2="",
        priority=1, stop_if_true=False, anchor_cell="J8", dxf_id=None,
        bounds=((10, 8, 10, 10),),
    )
    reader_result = CfRuleReadResult(
        rules=[rule], sheet_parts={}, cell_values=CellValueProvider({}, []),
        read_seconds=0.0, normalize_seconds=0.0, sheets_with_rules=["汇总表"],
    )

    results, stats = ComFormulaEvaluator(_Session(), _ComBook(sheet)).evaluate_sheet(
        "汇总表", [rule], [(10, 8, 10, 10)], reader_result)

    assert [item.cell for item in results] == ["J10"]
    assert stats["com_call_count"] == 3
    assert all(str(call).startswith("=") for call in sheet.calls)
