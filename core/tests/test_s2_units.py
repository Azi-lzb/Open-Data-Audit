from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from config_paths import resolve_test_config

from base_audit.systems.s2_report_collection.config_check import check_s2_config
from base_audit.systems.s2_report_collection.config import load_config, load_unit_settings, write_default_config
from base_audit.systems.s2_report_collection.external_engine import ExternalComparisonEngine
from base_audit.systems.s2_report_collection.importer import (
    ExternalValue,
    import_period_directory,
)
from base_audit.systems.s2_report_collection.models import (
    ExternalCheckRule,
    IndicatorConfig,
    IndicatorRecord,
    InstitutionProfile,
    PeriodBandConfig,
    ReportAuditConfig,
    RuleConfig,
    SourceRef,
    UnitSettings,
)
from base_audit.systems.s2_report_collection.period_engine import PeriodComparisonEngine
from base_audit.systems.s2_report_collection.rule_engine import RuleEngine
from base_audit.systems.s2_report_collection.unit_conversion import convert_amount, to_yuan


def _write_general_settings(
    book: Workbook,
    *,
    source: str = "元",
    external: str = "元",
    output: str = "元",
    hide_unchanged: str = "否",
    hide_matching: str = "否",
) -> None:
    sheet = book.create_sheet("通用设置")
    sheet.append(["设置项", "设置值", "说明"])
    sheet.append(["源数据单位", source, ""])
    sheet.append(["外部文件单位", external, ""])
    sheet.append(["输出文件单位", output, ""])
    sheet.append(["换算范围", "", "仅余额和累发"])
    sheet.append(["是否隐藏无变动指标", hide_unchanged, ""])
    sheet.append(["是否隐藏与大集中一致指标", hide_matching, ""])


def _make_config(path: Path, *, source: str = "元", external: str = "元", output: str = "元") -> None:
    book = Workbook()
    indicators = book.active
    indicators.title = "指标参照"
    indicators.append([
        "指标代码", "指标名称", "表单代码", "数据属性", "环比启用",
        "环比策略组", "最小变动值", "禁用", "备注",
    ])
    indicators.append(["20202001", "存款余额", "20202", "余额", "是", "BAL01", 0, "否", ""])
    indicators.append(["20202002", "累计金额", "20202", "累发", "是", "ACC01", 0, "否", ""])
    indicators.append(["20201028", "持股比例", "20201", "百分数", "是", "PCT01", 0, "否", ""])
    indicators.append(["20202003", "机构数量", "20202", "个数", "是", "CNT01", 0, "否", ""])

    orgs = book.create_sheet("机构参照")
    orgs.append(["机构名称", "社会信用代码", "机构类别", "承接行", "地区", "报表项目", "禁用"])
    orgs.append(["甲银行", "9144", "农商行", "一部", "示例地区", "row-a", "否"])

    strategies = book.create_sheet("环比策略")
    strategies.append(["策略组", "规则编号", "数据属性", "下限", "上限", "审核级别", "分类说明", "禁用"])
    for group, data_type in (("BAL01", "余额"), ("ACC01", "累发"), ("PCT01", "百分数"), ("CNT01", "个数")):
        strategies.append([group, "B001", data_type, None, 0, "软性", "下降", "否"])
        strategies.append([group, "B002", data_type, 0, 30, "", "", "否"])
        strategies.append([group, "B003", data_type, 30, None, "软性", "增长", "否"])

    rules = book.create_sheet("校验规则")
    rules.append(["规则编号", "类型", "描述", "规则内容", "Thd值(万元)", "取反", "级别", "禁用", "备注"])

    external_rules = book.create_sheet("外部核对规则")
    external_rules.append(["规则编号", "指标代码", "比较方式", "外部指标名称", "容差", "级别", "启用"])
    external_rules.append(["D001", "20202001", "等于", "存款余额", 0, "硬性", "是"])
    _write_general_settings(book, source=source, external=external, output=output)
    book.save(path)
    book.close()


def test_shipped_s2_configs_keep_global_unit_settings_in_main_workbook() -> None:
    for path in (
        resolve_test_config("2.报表采集系统_配置.xlsx"),
        resolve_test_config("默认配置/2.报表采集系统_配置.xlsx"),
    ):
        report = check_s2_config(path)
        assert report["passed"], report["issues"]
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            headers = next(book["指标参照"].iter_rows(values_only=True))
            assert not {"源数据单位", "大集中数据单位"}.intersection(headers)
        finally:
            book.close()
        general = load_workbook(path, read_only=True, data_only=True)
        try:
            configured = {row[0]: row[1] for row in general["通用设置"].iter_rows(min_row=2, values_only=True) if row[0] in {"源数据单位", "外部文件单位", "输出文件单位"}}
            loaded = load_config(path).unit_settings
            assert configured == {
                "源数据单位": loaded.source_unit,
                "外部文件单位": loaded.external_file_unit,
                "输出文件单位": loaded.output_file_unit,
            }
            assert set(configured) == {"源数据单位", "外部文件单位", "输出文件单位"}
            assert set(configured.values()) <= {"元", "万元", "亿元"}
            assert configured["源数据单位"] == "元"
            # 输出单位由配置维护；上方已核对合法值及加载一致性，不限定为“元”。
            # 三项默认级别随配置工作簿维护；正式配置应写明“硬性”。
            loaded_config = load_config(path)
            assert loaded_config.text_change_severity == "硬性"
            assert loaded_config.current_missing_severity == "硬性"
            assert loaded_config.previous_missing_severity == "硬性"
        finally:
            general.close()


def test_default_config_writer_embeds_unit_settings_in_same_workbook(tmp_path: Path) -> None:
    path = tmp_path / "2.报表采集系统_配置.xlsx"
    write_default_config(path)

    book = load_workbook(path, read_only=False, data_only=False)
    try:
        assert "通用设置" in book.sheetnames
        assert [cell.value for cell in book["外部核对规则"][1]] == [
            "规则编号", "指标代码", "比较方式", "外部指标名称", "容差", "级别", "启用",
        ]
        settings = {
            row[0]: row[1]
            for row in book["通用设置"].iter_rows(min_row=2, values_only=True)
            if row[0] in {
                "源数据单位", "外部文件单位", "输出文件单位",
                "是否隐藏无变动指标", "是否隐藏与大集中一致指标",
            }
        }
        assert settings == {
            "源数据单位": "元", "外部文件单位": "元", "输出文件单位": "元",
            "是否隐藏无变动指标": "否", "是否隐藏与大集中一致指标": "否",
        }
        general = book["通用设置"]
        assert [general.cell(row=1, column=column).value for column in range(1, 4)] == [
            "设置项", "设置值", "说明",
        ]
        assert [general.cell(row=row, column=1).value for row in range(2, 11)] == [
            "源数据单位", "外部文件单位", "输出文件单位", "换算范围",
            "是否隐藏无变动指标", "是否隐藏与大集中一致指标",
            "文字变动默认级别", "本期缺失默认级别", "上期缺失默认级别",
        ]
        assert [general.cell(row=row, column=2).value for row in range(8, 11)] == ["硬性", "硬性", "硬性"]
        assert all(cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb.endswith("1F4E78") for cell in general[1])
        assert all(cell.font.bold and cell.font.color.rgb.endswith("FFFFFF") for cell in general[1])
        body_cells = [cell for row in general.iter_rows(min_row=2) for cell in row if cell.value not in (None, "")]
        assert all(cell.fill.fill_type is None for cell in body_cells)
        validations = general.data_validations.dataValidation
        assert [(item.type, item.formula1, str(item.sqref)) for item in validations] == [
            ("list", '"元,万元,亿元"', "B2:B4"),
            ("list", '"是,否"', "B6:B7"),
            ("list", '"硬性,软性"', "B8:B10"),
        ]
    finally:
        book.close()
    assert not path.with_name("2.报表采集系统_单位配置.xlsx").exists()
    config = load_config(path)
    assert config.unit_settings.source_unit == "元"
    assert config.hide_unchanged_period_rows is False
    assert config.hide_matching_external_rows is False


def test_config_checker_rejects_unsupported_global_unit(tmp_path: Path) -> None:
    path = tmp_path / "s2-units.xlsx"
    _make_config(path)
    book = load_workbook(path)
    book["通用设置"]["B2"] = "千元"
    book.save(path)
    book.close()

    report = check_s2_config(path)
    assert not report["passed"]
    assert any("必须为元、万元或亿元" in issue["message"] for issue in report["issues"])


def test_config_checker_rejects_missing_global_unit_setting(tmp_path: Path) -> None:
    path = tmp_path / "s2-units.xlsx"
    _make_config(path)
    book = load_workbook(path)
    book["通用设置"]["B4"] = None
    book.save(path)
    book.close()

    report = check_s2_config(path)
    assert not report["passed"]
    assert any("必须为元、万元或亿元" in issue["message"] for issue in report["issues"])


def test_config_checker_rejects_missing_required_unit_settings_sheet(tmp_path: Path) -> None:
    path = tmp_path / "s2-units.xlsx"
    _make_config(path)
    book = load_workbook(path)
    book["通用设置"].title = "单位设置"
    book.save(path)
    book.close()
    report = check_s2_config(path)
    assert not report["passed"]
    assert any("缺少工作表：通用设置" in issue["message"] for issue in report["issues"])


def test_runtime_loads_all_three_units_from_main_workbook(tmp_path: Path) -> None:
    path = tmp_path / "s2-units.xlsx"
    _make_config(path, source="元", external="亿元", output="万元")

    config = load_config(path)

    assert config.unit_settings == UnitSettings("元", "亿元", "万元")
    assert load_unit_settings(path) == config.unit_settings


def test_inline_unit_columns_are_rejected(tmp_path: Path) -> None:
    path = tmp_path / "old-s2-config.xlsx"
    _make_config(path)
    book = load_workbook(path)
    book["指标参照"]["J1"] = "源数据单位"
    book.save(path)
    book.close()
    report = check_s2_config(path)
    assert not report["passed"]
    assert any("通用设置" in issue["message"] for issue in report["issues"])


@pytest.mark.parametrize(
    ("hide_unchanged", "hide_matching"),
    [(False, False), (False, True), (True, False), (True, True)],
)
def test_runtime_loads_all_general_visibility_setting_combinations(
    tmp_path: Path, hide_unchanged: bool, hide_matching: bool,
) -> None:
    path = tmp_path / "s2-general.xlsx"
    _make_config(path)
    book = load_workbook(path)
    book["通用设置"]["B6"] = "是" if hide_unchanged else "否"
    book["通用设置"]["B7"] = "是" if hide_matching else "否"
    book.save(path)
    book.close()

    config = load_config(path)

    assert config.hide_unchanged_period_rows is hide_unchanged
    assert config.hide_matching_external_rows is hide_matching


def test_report_config_visibility_defaults_are_false() -> None:
    config = ReportAuditConfig()

    assert config.hide_unchanged_period_rows is False
    assert config.hide_matching_external_rows is False


def test_report_config_severity_defaults_are_hard() -> None:
    config = ReportAuditConfig()

    assert config.text_change_severity == "硬性"
    assert config.current_missing_severity == "硬性"
    assert config.previous_missing_severity == "硬性"


@pytest.mark.parametrize(
    ("setting", "field"),
    [
        ("文字变动默认级别", "text_change_severity"),
        ("本期缺失默认级别", "current_missing_severity"),
        ("上期缺失默认级别", "previous_missing_severity"),
    ],
)
def test_severity_default_settings_are_optional_enums(tmp_path: Path, setting: str, field: str) -> None:
    path = tmp_path / "s2-general.xlsx"
    _make_config(path)
    # 旧配置不写该行：缺省按硬性执行，不阻断。
    assert getattr(load_config(path), field) == "硬性"

    book = load_workbook(path)
    book["通用设置"].append([setting, "软性", ""])
    book.save(path)
    book.close()
    assert getattr(load_config(path), field) == "软性"

    book = load_workbook(path)
    book["通用设置"]["B8"] = "提示"
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match=f"{setting}.*必须为硬性或软性"):
        load_config(path)
    report = check_s2_config(path)
    assert not report["passed"]
    assert any(setting in issue["message"] and "必须为硬性或软性" in issue["message"] for issue in report["issues"])


@pytest.mark.parametrize(
    ("cell", "setting", "bad_value"),
    [
        ("B6", "是否隐藏无变动指标", "true"),
        ("B7", "是否隐藏与大集中一致指标", "不隐藏"),
    ],
)
def test_runtime_and_checker_reject_non_yes_no_visibility_values(
    tmp_path: Path, cell: str, setting: str, bad_value: str,
) -> None:
    path = tmp_path / "s2-general.xlsx"
    _make_config(path)
    book = load_workbook(path)
    book["通用设置"][cell] = bad_value
    book.save(path)
    book.close()

    with pytest.raises(ValueError, match=f"{setting}.*必须为“是”或“否”"):
        load_config(path)
    report = check_s2_config(path)
    assert not report["passed"]
    assert any(setting in issue["message"] and "必须为“是”或“否”" in issue["message"] for issue in report["issues"])


@pytest.mark.parametrize(
    ("setting", "cell", "default_text"),
    [
        ("是否隐藏无变动指标", "B6", "否"),
        ("是否隐藏与大集中一致指标", "B7", "否"),
    ],
)
def test_runtime_uses_default_when_general_setting_row_is_missing(
    tmp_path: Path, setting: str, cell: str, default_text: str,
) -> None:
    path = tmp_path / "s2-general.xlsx"
    _make_config(path)
    book = load_workbook(path)
    book["通用设置"][cell.replace("B", "A")] = None
    book["通用设置"][cell] = None
    book.save(path)
    book.close()

    config = load_config(path)
    assert not config.hide_unchanged_period_rows
    assert not config.hide_matching_external_rows
    assert any(
        setting in warning and f"已按默认“{default_text}”执行" in warning
        for warning in config.warnings
    )
    report = check_s2_config(path)
    assert report["passed"]
    assert any(
        issue["level"] == "warning" and setting in issue["message"] and "已按默认" in issue["message"]
        for issue in report["issues"]
    )


def test_general_settings_allows_empty_sheet_and_runs_on_defaults(tmp_path: Path) -> None:
    """设置行全部丢失时按默认执行：单位=元、隐藏=否、级别=硬性；逐项提示补齐。"""
    path = tmp_path / "s2-general.xlsx"
    _make_config(path)
    book = load_workbook(path)
    general = book["通用设置"]
    for row in range(general.max_row, 1, -1):
        general.delete_rows(row)
    book.save(path)
    book.close()

    config = load_config(path)
    assert config.unit_settings == UnitSettings("元", "元", "元")
    assert config.hide_unchanged_period_rows is False
    assert config.hide_matching_external_rows is False
    assert config.text_change_severity == "硬性"
    assert config.current_missing_severity == "硬性"
    assert config.previous_missing_severity == "硬性"
    # 换算范围是纯说明行，缺行不提示；其余 8 个取值设置项各提示一条。
    assert len(config.warnings) == 8
    assert all("已按默认" in warning for warning in config.warnings)
    assert not any("换算范围" in warning for warning in config.warnings)


def test_runtime_blocks_duplicate_general_setting(tmp_path: Path) -> None:
    path = tmp_path / "s2-general.xlsx"
    _make_config(path)
    book = load_workbook(path)
    book["通用设置"].append(["是否隐藏无变动指标", "是", "重复项"])
    book.save(path)
    book.close()

    with pytest.raises(ValueError, match="设置项重复：是否隐藏无变动指标"):
        load_config(path)
    report = check_s2_config(path)
    assert not report["passed"]
    assert any("设置项重复：是否隐藏无变动指标" in issue["message"] for issue in report["issues"])


def test_runtime_rejects_general_settings_header_mismatch(tmp_path: Path) -> None:
    path = tmp_path / "s2-general.xlsx"
    _make_config(path)
    book = load_workbook(path)
    book["通用设置"]["B1"] = "单位"
    book.save(path)
    book.close()

    with pytest.raises(ValueError, match="表头必须为：A1=设置项、B1=设置值、C1=说明"):
        load_config(path)


def test_amount_conversion_is_exact_without_decimal_context_rounding() -> None:
    value = Decimal("123456789012345678901234567890.12345678")
    converted = convert_amount(value, "元", "万元")
    assert converted == Decimal("12345678901234567890123456.789012345678")
    assert convert_amount(Decimal("1234567.89"), "元", "万元") == Decimal("123.456789")
    assert convert_amount(Decimal("0"), "亿元", "元") == Decimal("0")


def test_source_amounts_convert_to_output_unit_before_engines_and_keep_other_types(tmp_path: Path) -> None:
    source_dir = tmp_path / "current"
    source_dir.mkdir()
    path = source_dir / "9144#2026-06-30#x#20202#甲银行.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "20202"
    sheet.append(["报表", None, None])
    sheet.append([None, None, None])
    sheet.append(["指标代码", "指标名称", "本期情况"])
    sheet.append(["20202001", "存款余额", 1_234_567.89])
    sheet.append(["20202002", "累计金额", 12_345_678.9])
    sheet.append(["20201028", "持股比例", 9.92])
    sheet.append(["20202003", "机构数量", 42])
    sheet.append(["20202004", "真实零值", 0])
    book.save(path)
    book.close()
    config = ReportAuditConfig(
        unit_settings=UnitSettings("元", "亿元", "万元"),
        indicators={
            "20202001": IndicatorConfig("20202001", data_type="余额"),
            "20202002": IndicatorConfig("20202002", data_type="累发"),
            "20201028": IndicatorConfig("20201028", data_type="百分数"),
            "20202003": IndicatorConfig("20202003", data_type="个数"),
            "20202004": IndicatorConfig("20202004", data_type="余额"),
        },
    )

    result = import_period_directory(source_dir, config=config, label="本期")
    records = {record.indicator_code: record for record in result.records}
    assert records["20202001"].normalized_value == Decimal("123.456789")
    assert records["20202001"].normalized_unit == "万元"
    assert records["20202002"].normalized_value == Decimal("1234.56789")
    assert records["20201028"].normalized_value == 9.92
    assert records["20201028"].normalized_unit == ""
    assert records["20202003"].normalized_value == 42
    assert records["20202003"].normalized_unit == ""
    assert records["20202004"].normalized_value == Decimal("0")


def test_unit_conversion_rejects_unsupported_unit() -> None:
    with pytest.raises(ValueError, match="不支持的金额单位"):
        to_yuan(1, "千元")


def _record(code: str, value: Decimal | float, *, data_type: str = "余额", unit: str = "万元") -> IndicatorRecord:
    source = SourceRef("source.xlsx", "20202", 4)
    return IndicatorRecord(
        "9144", "甲银行", "9144", "农商行", "一部", "示例地区", "20202", code,
        code, "2026-06", "20260630", value, data_type, unit, value, unit, source,
    )


def test_external_values_convert_to_output_unit_and_compare_there() -> None:
    record = _record("20202001", Decimal("1.234567"))
    config = ReportAuditConfig(
        unit_settings=UnitSettings("元", "亿元", "万元"),
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        external_rules=[ExternalCheckRule("D001", "20202001", "大集中", "存款余额", "等于", 0.000001, "硬性")],
    )
    external = {("row-a", "存款余额"): ExternalValue("row-a", "存款余额", Decimal("0.0001234567"), record.source)}
    findings = ExternalComparisonEngine().run(
        {record.logical_key: record}, external, {"9144": "row-a"}, config,
    )
    assert len(findings) == 1
    assert findings[0].external_value == Decimal("1.234567")
    assert findings[0].difference_value == Decimal("0.000000")
    assert findings[0].status == "正常"
    assert findings[0].calculation_trace == "差异绝对值= abs(1.23 - 1.23) = 0.00 万元"


def test_external_mapping_precedes_cf2_and_cf2_is_used_as_fallback() -> None:
    record = _record("20202001", Decimal("10"))
    config = ReportAuditConfig(
        institutions_by_id={
            "9144": InstitutionProfile("9144", "甲银行", "9144", report_item="配置项目"),
        },
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        external_rules=[ExternalCheckRule("D001", "20202001", "大集中", "存款余额", "等于", 0, "硬性")],
    )
    external = {
        ("外部项目", "存款余额"): ExternalValue("外部项目", "存款余额", Decimal("10"), record.source),
        ("配置项目", "存款余额"): ExternalValue("配置项目", "存款余额", Decimal("30"), record.source),
    }

    from_external = ExternalComparisonEngine().run(
        {record.logical_key: record}, external, {"9144": "外部项目"}, config,
    )[0]
    from_config = ExternalComparisonEngine().run(
        {record.logical_key: record}, external, {}, config,
    )[0]

    assert from_external.external_value == Decimal("10")
    assert from_external.report_item == "外部项目"
    assert from_external.status == "正常"
    assert from_config.external_value == Decimal("30")
    assert from_config.report_item == "配置项目"
    assert from_config.status == "异常"


def test_external_calculation_trace_formats_absolute_difference_to_two_decimals() -> None:
    record = _record("20202001", Decimal("5145253.70000"))
    config = ReportAuditConfig(
        unit_settings=UnitSettings("元", "万元", "万元"),
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        external_rules=[ExternalCheckRule("D001", "20202001", "大集中", "存款余额", "等于", 0, "硬性")],
    )
    external = {
        ("row-a", "存款余额"): ExternalValue("row-a", "存款余额", Decimal("5124804.24"), record.source),
    }

    finding = ExternalComparisonEngine().run(
        {record.logical_key: record}, external, {"9144": "row-a"}, config,
    )[0]

    assert finding.calculation_trace == (
        "差异绝对值= abs(5145253.70 - 5124804.24) = 20449.46 万元"
    )
    assert finding.difference_value == Decimal("20449.46000")
    assert finding.status == "异常"


@pytest.mark.parametrize(
    ("method", "left", "right", "tolerance", "expected_status"),
    [
        ("等于", "10", "10", 0, "正常"),
        ("等于", "11", "10", 1, "正常"),  # 容差边界包含等于。
        ("等于", "11.0001", "10", 1, "异常"),
        ("不等于", "10", "11", 0, "正常"),
        ("不等于", "10", "10", 0, "异常"),
        ("大于", "11", "10", 0, "正常"),
        ("大于", "10", "10", 0, "异常"),
        ("大于等于", "10", "10", 0, "正常"),
        ("大于等于", "9", "10", 0, "异常"),
        ("小于", "9", "10", 0, "正常"),
        ("小于", "10", "10", 0, "异常"),
        ("小于等于", "10", "10", 0, "正常"),
        ("小于等于", "11", "10", 0, "异常"),
    ],
)
def test_external_comparison_methods_use_report_value_as_left_operand_and_test_boundaries(
    method: str, left: str, right: str, tolerance: float, expected_status: str,
) -> None:
    report_value, external_value = Decimal(left), Decimal(right)
    record = _record("20202001", report_value)
    config = ReportAuditConfig(
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        external_rules=[ExternalCheckRule(
            "D001", "20202001", "大集中", "存款余额", method, tolerance, "硬性",
        )],
    )
    external = {
        ("row-a", "存款余额"): ExternalValue("row-a", "存款余额", external_value, record.source),
    }

    finding = ExternalComparisonEngine().run(
        {record.logical_key: record}, external, {"9144": "row-a"}, config,
    )[0]

    assert finding.current_value == report_value
    assert finding.external_value == external_value
    assert finding.status == expected_status
    assert finding.calculation_trace == (
        f"差异绝对值= abs({report_value:.2f} - {external_value:.2f}) = "
        f"{abs(report_value - external_value):.2f} 元"
    )
    assert "成立" not in finding.calculation_trace
    assert "不成立" not in finding.calculation_trace
    assert "容差" not in finding.calculation_trace
    if method == "等于":
        assert finding.description == (
            "与大集中系统数据一致" if expected_status == "正常" else "与大集中系统数据不一致"
        )
    else:
        assert finding.description == (
            "满足外部核对规则" if expected_status == "正常" else "不满足外部核对规则"
        )


def test_external_missing_value_is_not_treated_as_real_zero() -> None:
    record = _record("20202001", Decimal("0"))
    config = ReportAuditConfig(
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        external_rules=[ExternalCheckRule("D001", "20202001", "大集中", "存款余额", "等于", 0, "硬性")],
    )

    zero = ExternalValue("row-a", "存款余额", Decimal("0"), record.source)
    zero_finding = ExternalComparisonEngine().run(
        {record.logical_key: record}, {("row-a", "存款余额"): zero}, {"9144": "row-a"}, config,
    )[0]
    missing_finding = ExternalComparisonEngine().run(
        {record.logical_key: record}, {}, {"9144": "row-a"}, config,
    )[0]

    assert zero_finding.status == "正常"
    assert zero_finding.external_value == Decimal("0")
    assert zero_finding.report_item == "row-a"
    assert zero_finding.difference_value == Decimal("0")
    assert missing_finding.status == "异常"
    assert missing_finding.external_value is None
    assert missing_finding.report_item == "row-a"
    assert missing_finding.retrieval_note == "外部未匹配"
    assert missing_finding.description == "大集中系统数据未找到该报表项目，请先完善大集中数据的报表定义"
    assert missing_finding.rule_description == missing_finding.description
    assert missing_finding.calculation_trace == "未比较：报表项目“row-a”、指标“存款余额”无可读取数值"


def test_external_uses_explicit_zero_fallback_when_institution_label_is_missing() -> None:
    from base_audit.systems.s2_report_collection.importer import (
        MISSING_INSTITUTION_LABEL_FALLBACK_KEY,
        MISSING_INSTITUTION_LABEL_FALLBACK_REASON,
    )

    record = _record("20202001", Decimal("5"))
    config = ReportAuditConfig(
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        external_rules=[ExternalCheckRule("D001", "20202001", "大集中", "存款余额", "等于", 0, "硬性")],
    )
    zero_fallback = ExternalValue(
        "", "", Decimal("0"), record.source, MISSING_INSTITUTION_LABEL_FALLBACK_REASON,
    )

    finding = ExternalComparisonEngine().run(
        {record.logical_key: record},
        {MISSING_INSTITUTION_LABEL_FALLBACK_KEY: zero_fallback},
        {},
        config,
    )[0]

    assert finding.external_value == Decimal("0")
    assert finding.report_item == ""
    assert finding.difference_value == Decimal("5")
    assert finding.status == "异常"
    assert MISSING_INSTITUTION_LABEL_FALLBACK_REASON in finding.description
    assert "差异绝对值= abs(5.00 - 0.00) = 5.00 元" == finding.calculation_trace


@pytest.mark.parametrize(
    ("data_type", "indicator_code", "report_value", "external_value", "tolerance"),
    [
        ("百分数", "20201028", "9.92", "9.9205", 0.001),
        ("个数", "20202003", "42", "42.0005", 0.001),
    ],
)
def test_non_amount_external_values_and_tolerance_keep_native_scale(
    data_type: str, indicator_code: str, report_value: str, external_value: str, tolerance: float,
) -> None:
    record = _record(indicator_code, Decimal(report_value), data_type=data_type, unit="")
    config = ReportAuditConfig(
        unit_settings=UnitSettings("元", "亿元", "万元"),
        indicators={indicator_code: IndicatorConfig(indicator_code, data_type=data_type)},
        external_rules=[ExternalCheckRule(
            "D001", indicator_code, "大集中", "指标", "等于", tolerance, "硬性",
        )],
    )
    external = ExternalValue("row-a", "指标", Decimal(external_value), record.source)

    finding = ExternalComparisonEngine().run(
        {record.logical_key: record}, {("row-a", "指标"): external}, {"9144": "row-a"}, config,
    )[0]

    assert finding.current_value == Decimal(report_value)
    assert finding.external_value == Decimal(external_value)
    assert finding.difference_value == Decimal(external_value) - Decimal(report_value)
    assert finding.status == "正常"


@pytest.mark.parametrize(
    ("external_value", "expected_status"),
    [("0.01234568", "正常"), ("0.01234569", "异常")],
)
def test_amount_comparison_converts_both_sides_to_output_unit_and_keeps_tolerance_in_that_unit(
    external_value: str, expected_status: str,
) -> None:
    source = SourceRef("source.xlsx", "20202", 4)
    record = IndicatorRecord(
        "9144", "甲银行", "9144", "农商行", "一部", "示例地区", "20202", "20202001",
        "存款余额", "2026-06", "20260630", Decimal("1234567"), "余额", "元",
        Decimal("123.4567"), "万元", source,
    )
    config = ReportAuditConfig(
        unit_settings=UnitSettings("元", "亿元", "万元"),
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        external_rules=[ExternalCheckRule(
            "D001", "20202001", "大集中", "存款余额", "等于", 0.0001, "硬性",
        )],
    )
    external = ExternalValue("row-a", "存款余额", Decimal(external_value), source)

    finding = ExternalComparisonEngine().run(
        {record.logical_key: record}, {("row-a", "存款余额"): external}, {"9144": "row-a"}, config,
    )[0]

    assert finding.current_value == Decimal("123.4567")
    assert finding.external_value == convert_amount(Decimal(external_value), "亿元", "万元")
    assert finding.difference_value == abs(finding.current_value - finding.external_value)
    assert finding.status == expected_status


def test_period_comparison_runs_after_source_values_are_normalized(tmp_path: Path) -> None:
    config_path = tmp_path / "s2.xlsx"
    _make_config(config_path, source="元", output="万元")
    config = ReportAuditConfig(
        unit_settings=UnitSettings("元", "元", "万元"),
        indicators={
            "20202001": IndicatorConfig(
                "20202001", data_type="余额", period_strategy_group="BAL01", min_change_value=0,
            ),
        },
        bands=[
            PeriodBandConfig("B001", None, 0, "下降", "软性", False, "BAL01", "余额"),
            PeriodBandConfig("B002", 0, 30, "增长", "软性", False, "BAL01", "余额"),
            PeriodBandConfig("B003", 30, None, "大幅增长", "软性", False, "BAL01", "余额"),
        ],
    )
    current, previous = tmp_path / "current", tmp_path / "previous"
    current.mkdir()
    previous.mkdir()
    for directory, date, value in ((current, "2026-06-30", 1_234_567.0), (previous, "2026-03-31", 1_000_000.0)):
        path = directory / f"9144#{date}#x#20202#甲银行.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "20202"
        sheet.append(["标题", None, None])
        sheet.append([None, None, None])
        sheet.append(["指标代码", "指标名称", "本期情况"])
        sheet.append(["20202001", "存款余额", value])
        book.save(path)
        book.close()
    cur_records = import_period_directory(current, config=config, label="本期").by_logical_key()
    pre_records = import_period_directory(previous, config=config, label="上期").by_logical_key()

    finding = PeriodComparisonEngine().run(cur_records, pre_records, config)[0]
    assert finding.current_value == Decimal("123.4567")
    assert finding.previous_value == Decimal("100")
    assert finding.difference_value == Decimal("23.4567")
    assert finding.band == "BAL01-B002"


def test_rule_threshold_is_converted_to_output_unit_before_evaluation() -> None:
    record = _record("20202001", Decimal("0.00005"), unit="亿元")
    config = ReportAuditConfig(
        unit_settings=UnitSettings("元", "亿元", "亿元"),
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        rules=[RuleConfig("R001", "表达式", "阈值检查", "[20202001] > Thd", "软性", 1)],
    )
    findings, _quality = RuleEngine().run({record.logical_key: record}, {}, config)
    assert findings == []  # Thd=1万元 → 0.0001亿元；本期 0.00005亿元不超过阈值。


def test_rule_expression_preserves_decimal_digits_beyond_excel_float_precision() -> None:
    record = _record("20202001", Decimal("1234567890123456.7"), unit="元")
    config = ReportAuditConfig(
        unit_settings=UnitSettings("元", "元", "元"),
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        rules=[RuleConfig("R002", "表达式", "精度检查", "[20202001] > 1234567890123456.6", "硬性")],
    )

    findings, quality = RuleEngine().run({record.logical_key: record}, {}, config)

    assert not quality
    assert len(findings) == 1
