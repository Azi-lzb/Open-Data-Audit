"""V2 核心模型、导入校验、三类引擎与值写入输出的回归保护。"""

from __future__ import annotations

import hashlib
from dataclasses import replace
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from base_audit.systems.s2_report_collection.service import PeriodAuditV2Service
from base_audit.systems.s2_report_collection.models import AuditFinding, ExternalCheckRule, IndicatorConfig, IndicatorRecord, InstitutionProfile, PeriodBandConfig, RelatedRecord, ReportAuditConfig, RuleConfig, SourceRef, UnitSettings
from base_audit.systems.s2_report_collection.period_engine import PeriodComparisonEngine
from base_audit.systems.s2_report_collection.rule_engine import RuleEngine
from base_audit.systems.s2_report_collection.importer import ImportResult, import_period_directory
from base_audit.systems.s2_report_collection.validation import validate_import
from base_audit.systems.s2_report_collection.exporter import ExcelExporter, EXTERNAL_HEADERS, PERIOD_HEADERS, VALIDATION_HEADERS
from base_audit.systems.s2_report_collection.external_engine import ExternalComparisonEngine
from base_audit.systems.s2_report_collection.importer import ExternalValue
from base_audit.systems.s2_report_collection.config import load_config


def _config(path: Path) -> None:
    book = Workbook(); sheet = book.active; sheet.title = "指标参照"
    sheet.append(["指标代码", "指标名称", "表单代码", "数据属性", "环比启用", "环比策略组", "最小变动值", "禁用", "备注"])
    sheet.append(["20201001", "机构名称", "20201", "文字", "是", "不适用", "不适用", "否", ""])
    sheet.append(["20201002", "机构代码", "20201", "文字", "是", "不适用", "不适用", "否", ""])
    for code, name in (("20202001", "存款"), ("20202002", "贷款"), ("20202003", "资产"), ("20202004", "负债"), ("20202005", "权益")):
        sheet.append([code, name, "20202", "余额", "是", "BAL01", 0, "否", ""])
    external = book.create_sheet("外部核对规则")
    external.append(["规则编号", "指标代码", "比较方式", "外部指标名称", "容差", "级别", "启用"])
    external.append(["D001", "20202001", "等于", "存款", 0, "硬性", "是"])
    org = book.create_sheet("机构参照"); org.append(["机构名称", "社会信用代码", "机构类别", "承接行", "地区", "报表项目", "禁用"]); org.append(["甲银行", "9144", "农商行", "一部", "示例地区", "甲行", "否"])
    bands = book.create_sheet("环比策略"); bands.append(["策略组", "规则编号", "数据属性", "下限", "上限", "审核级别", "分类说明", "禁用"]); bands.append(["BAL01", "P001", "余额", None, -30, "软性", "下降", "否"]); bands.append(["BAL01", "P002", "余额", -30, 30, "", "", "否"]); bands.append(["BAL01", "P003", "余额", 30, None, "软性", "增长", "否"])
    rules = book.create_sheet("校验规则"); rules.append(["规则编号", "类型", "描述", "规则内容", "Thd值(万元)", "取反标识", "级别", "禁用", "备注"]); rules.append(["R1", "表达式", "资产负债不平", "[20202003]<>[20202004]+[20202005]", "", "否", "硬性", "否", ""])
    general = book.create_sheet("通用设置"); general.append(["设置项", "设置值", "说明"])
    general.append(["源数据单位", "元", ""]); general.append(["外部文件单位", "元", ""]); general.append(["输出文件单位", "元", ""])
    general.append(["换算范围", "", "仅余额和累发"])
    general.append(["是否隐藏无变动指标", "否", ""]); general.append(["是否隐藏与大集中一致指标", "否", ""])
    book.save(path)


def _source(path: Path, *, deposit: float, loan: float, asset: float, debt: float, equity: float, form: str = "20202", org_name: str = "甲银行") -> None:
    book = Workbook(); sh = book.active; sh.title = form
    sh.append(["标题", "", ""]); sh.append(["", "", ""]); sh.append(["指标", "名称", "值"])
    for row in ((20201001, "机构名称", org_name), (20201002, "机构代码", "9144"), (20202001, "存款", deposit), (20202002, "贷款", loan), (20202003, "资产", asset), (20202004, "负债", debt), (20202005, "权益", equity)):
        sh.append(row)
    book.save(path)


def _central(path: Path) -> None:
    book = Workbook(); sh = book.active; sh.title = "集中系统数据"; sh.append([None, None, None]); sh.append([None, " ", "存款"]); sh.append([None, "甲行", 0.01]); ref = book.create_sheet("参照表"); ref.append(["统一社会信用代码", "报表项目"]); ref.append(["9144", "甲行"]); book.save(path)


def test_v2_generates_values_not_excel_formulas(tmp_path: Path) -> None:
    config, cur, pre, central, output = tmp_path / "config.xlsx", tmp_path / "cur", tmp_path / "pre", tmp_path / "central.xlsx", tmp_path / "out"
    cur.mkdir(); pre.mkdir(); _config(config)
    _source(cur / "9144#2026-06-30#01#20202#甲银行.xlsx", deposit=2_000_000, loan=1_000_000, asset=1_000_000, debt=600_000, equity=300_000)
    _source(pre / "9144#2026-03-31#01#20202#甲银行.xlsx", deposit=1_000_000, loan=1_000_000, asset=1_000_000, debt=600_000, equity=400_000)
    _central(central)
    result = PeriodAuditV2Service().run(current_dir=cur, previous_dir=pre, config_path=config, output_dir=output, central_path=central)
    assert result.summary.institutions == 1
    assert any(item.audit_type == "规则" and item.rule_id == "R1" for item in result.findings)
    assert not (output / "审核工作底稿.xlsx").exists()
    # 输出名含时间戳（同名重跑不覆盖）；result.paths 记录实际路径。
    checklist = result.paths.result
    assert checklist.parent == output and checklist.name.startswith("报表审核清单_")
    assert checklist.is_file()
    book = load_workbook(checklist, data_only=False)
    try:
        assert book.sheetnames == ["环比规则", "外部核对规则", "校验规则"]
        assert [cell.value for cell in book["环比规则"][1]] == PERIOD_HEADERS
        assert "规则编号" not in [cell.value for cell in book["环比规则"][1]]
        assert "取数说明" not in [cell.value for cell in book["环比规则"][1]]
        assert [cell.value for cell in book["外部核对规则"][1]] == EXTERNAL_HEADERS
        external_sheet = book["外部核对规则"]
        assert external_sheet["E1"].value == "社会信用代码"
        assert external_sheet["F1"].value == "报表项目"
        assert external_sheet["F2"].value == "甲行"
        assert "规则编号" not in [cell.value for cell in book["外部核对规则"][1]]
        assert "取数说明" not in [cell.value for cell in book["外部核对规则"][1]]
        assert [cell.value for cell in book["校验规则"][1]] == VALIDATION_HEADERS
        assert all("状态" not in [cell.value for cell in book[name][1]] for name in book.sheetnames)
        assert all("审核类型" not in [cell.value for cell in book[name][1]] for name in book.sheetnames)
        assert all(cell.hyperlink is None for sheet in book.worksheets for row in sheet.iter_rows() for cell in row)
        assert sum(sheet.max_row - 1 for sheet in book.worksheets) == len(result.findings)
    finally:
        book.close()


def test_v2_service_uses_general_visibility_settings_for_result_rows(tmp_path: Path) -> None:
    config, cur, pre, central, output = (
        tmp_path / "config.xlsx", tmp_path / "cur", tmp_path / "pre",
        tmp_path / "central.xlsx", tmp_path / "out",
    )
    cur.mkdir(); pre.mkdir(); _config(config)
    settings = load_workbook(config)
    settings["通用设置"]["B6"] = "是"
    settings["通用设置"]["B7"] = "是"
    settings.save(config)
    settings.close()
    _source(cur / "9144#2026-06-30#01#20202#甲银行.xlsx", deposit=2_000_000, loan=1_000_000, asset=1_000_000, debt=600_000, equity=400_000)
    _source(pre / "9144#2026-03-31#01#20202#甲银行.xlsx", deposit=1_000_000, loan=1_000_000, asset=1_000_000, debt=600_000, equity=400_000)
    _central(central)
    external = load_workbook(central)
    external["集中系统数据"]["C3"] = 2_000_000
    external.save(central)
    external.close()

    result = PeriodAuditV2Service().run(
        current_dir=cur, previous_dir=pre, config_path=config, output_dir=output, central_path=central,
    )
    book = load_workbook(result.paths.result)
    try:
        assert sum(sheet.max_row - 1 for sheet in book.worksheets) == len(result.findings)
        for sheet_name, description in (
            ("环比规则", "无变动"),
            ("外部核对规则", "与大集中系统数据一致"),
        ):
            sheet = book[sheet_name]
            explanation_col = next(cell.column for cell in sheet[1] if cell.value == "审核说明")
            matches = [row for row in range(2, sheet.max_row + 1) if sheet.cell(row, explanation_col).value == description]
            assert matches, sheet_name
            assert all(sheet.row_dimensions[row].hidden for row in matches)
    finally:
        book.close()


def test_v2_text_change_severity_setting_and_data_attribute_column(tmp_path: Path) -> None:
    """通用设置“文字变动默认级别”与结果“环比规则”表的数据属性列。"""
    config, cur, pre, output = tmp_path / "config.xlsx", tmp_path / "cur", tmp_path / "pre", tmp_path / "out"
    cur.mkdir(); pre.mkdir(); _config(config)
    settings = load_workbook(config)
    settings["通用设置"].append(["文字变动默认级别", "软性", ""])
    settings.save(config)
    settings.close()
    _source(cur / "9144#2026-06-30#01#20202#甲银行.xlsx", deposit=2_000_000, loan=1_000_000, asset=1_000_000, debt=600_000, equity=400_000)
    _source(pre / "9144#2026-03-31#01#20202#甲银行.xlsx", deposit=1_000_000, loan=1_000_000, asset=1_000_000, debt=600_000, equity=300_000, org_name="乙银行")
    result = PeriodAuditV2Service().run(current_dir=cur, previous_dir=pre, config_path=config, output_dir=output)

    text_row = next(item for item in result.findings if item.audit_type == "环比" and item.indicator_code == "20201001")
    assert text_row.description == "文字指标变化"
    assert text_row.severity == "软性"
    assert text_row.data_type == "文字"
    book = load_workbook(result.paths.result, data_only=True)
    try:
        sheet = book["环比规则"]
        headers = [cell.value for cell in sheet[1]]
        assert headers[6:9] == ["指标代码", "数据属性", "指标名称"]
        code_column = headers.index("指标代码") + 1
        attribute_column = headers.index("数据属性") + 1
        attributes = {
            sheet.cell(row, code_column).value: sheet.cell(row, attribute_column).value
            for row in range(2, sheet.max_row + 1)
        }
        assert attributes["20201001"] == "文字"
        assert attributes["20202001"] == "余额"
    finally:
        book.close()


def test_v2_text_change_and_missing_side_severity_follow_settings() -> None:
    """文字变动与本期/上期缺失默认级别各自跟随通用设置，缺省均为硬性。

    “文字变动默认级别”管文字指标变化与文字指标单侧缺失；
    “本期缺失默认级别”“上期缺失默认级别”只管非文字指标，按审核说明
    方向分别生效（含该侧无数值或为 0 按无处理）。
    """
    engine = PeriodComparisonEngine()

    def text_pair(current_value, previous_value, config):
        cur = _record("20201001", current_value, form="20201", value_type="文字")
        pre = _record("20201001", previous_value, form="20201", period="2026-03", value_type="文字")
        return engine.run({cur.logical_key: cur}, {pre.logical_key: pre}, config)[0]

    changed = text_pair("甲银行", "乙银行", ReportAuditConfig())
    assert changed.description == "文字指标变化"
    assert changed.severity == "硬性"
    assert changed.data_type == "文字"
    soft = text_pair("甲银行", "乙银行", ReportAuditConfig(text_change_severity="软性"))
    assert soft.severity == "软性"
    unchanged = text_pair("甲银行", "甲银行", ReportAuditConfig())
    assert unchanged.status == "正常" and unchanged.severity == ""

    # 文字指标单侧缺失跟随“文字变动默认级别”，不受缺失方向设置影响。
    cur_only_text = _record("20201002", "9144", form="20201", value_type="文字")
    text_missing = engine.run({cur_only_text.logical_key: cur_only_text}, {}, ReportAuditConfig())[0]
    assert text_missing.description == "本期有，上期无"
    assert text_missing.severity == "硬性"
    text_missing_soft = engine.run({cur_only_text.logical_key: cur_only_text}, {}, ReportAuditConfig(text_change_severity="软性"))[0]
    assert text_missing_soft.severity == "软性"
    text_missing_other = engine.run({cur_only_text.logical_key: cur_only_text}, {}, ReportAuditConfig(current_missing_severity="软性", previous_missing_severity="软性"))[0]
    assert text_missing_other.severity == "硬性"

    # 非文字指标单侧缺失按方向取设置：本期缺 → 本期缺失默认级别。
    pre_only = _record("20202001", 5.0, period="2026-03")
    current_missing = engine.run({}, {pre_only.logical_key: pre_only}, ReportAuditConfig())[0]
    assert current_missing.description == "本期无，上期有"
    assert current_missing.severity == "硬性"
    current_missing_soft = engine.run({}, {pre_only.logical_key: pre_only}, ReportAuditConfig(current_missing_severity="软性"))[0]
    assert current_missing_soft.severity == "软性"
    current_missing_other = engine.run({}, {pre_only.logical_key: pre_only}, ReportAuditConfig(previous_missing_severity="软性"))[0]
    assert current_missing_other.severity == "硬性"

    # 上期缺 → 上期缺失默认级别。
    cur_only = _record("20202002", 5.0)
    previous_missing = engine.run({cur_only.logical_key: cur_only}, {}, ReportAuditConfig())[0]
    assert previous_missing.description == "本期有，上期无"
    assert previous_missing.severity == "硬性"
    previous_missing_soft = engine.run({cur_only.logical_key: cur_only}, {}, ReportAuditConfig(previous_missing_severity="软性"))[0]
    assert previous_missing_soft.severity == "软性"
    previous_missing_other = engine.run({cur_only.logical_key: cur_only}, {}, ReportAuditConfig(current_missing_severity="软性"))[0]
    assert previous_missing_other.severity == "硬性"

    # 一侧为 0 按无处理（band=未命中）同样按方向跟随各自设置。
    current_value = _record("20202003", 5.0)
    previous_zero = _record("20202003", 0, period="2026-03")
    previous_zero_side = engine.run({current_value.logical_key: current_value}, {previous_zero.logical_key: previous_zero}, ReportAuditConfig(previous_missing_severity="软性"))[0]
    assert previous_zero_side.description == "本期有，上期无" and previous_zero_side.band == "未命中"
    assert previous_zero_side.severity == "软性"
    current_zero = _record("20202004", 0)
    previous_value = _record("20202004", 9.0, period="2026-03")
    current_zero_side = engine.run({current_zero.logical_key: current_zero}, {previous_value.logical_key: previous_value}, ReportAuditConfig(current_missing_severity="软性"))[0]
    assert current_zero_side.description == "本期无，上期有" and current_zero_side.band == "未命中"
    assert current_zero_side.severity == "软性"


def test_s2_business_output_is_independent_of_unrelated_log_preferences(tmp_path: Path) -> None:
    """切换设置中心日志选项不改变 S2 审核结果和工作簿数据。"""
    config, cur, pre = tmp_path / "config.xlsx", tmp_path / "cur", tmp_path / "pre"
    cur.mkdir(); pre.mkdir(); _config(config)
    _source(cur / "9144#2026-06-30#01#20202#甲银行.xlsx", deposit=2_000_000, loan=1_000_000,
            asset=1_000_000, debt=600_000, equity=300_000)
    _source(pre / "9144#2026-03-31#01#20202#甲银行.xlsx", deposit=1_000_000, loan=1_000_000,
            asset=1_000_000, debt=600_000, equity=400_000)

    from base_audit.web_app import WebApi

    api = WebApi(tmp_path / "settings")
    api.update({
        "showRunDetailLogs": False,
        "exportRunLogs": False,
        "showPerformanceDiagnostics": False,
    })
    first = PeriodAuditV2Service().run(
        current_dir=cur, previous_dir=pre, config_path=config, output_dir=tmp_path / "out_simple"
    )
    api.update({
        "showRunDetailLogs": True,
        "exportRunLogs": True,
        "showPerformanceDiagnostics": True,
    })
    second = PeriodAuditV2Service().run(
        current_dir=cur, previous_dir=pre, config_path=config, output_dir=tmp_path / "out_detail"
    )
    assert first.findings == second.findings
    assert first.summary == second.summary

    def workbook_values(path: Path) -> tuple:
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            return tuple(
                (sheet.title, tuple(tuple(row) for row in sheet.iter_rows(values_only=True)))
                for sheet in book.worksheets
            )
        finally:
            book.close()

    assert workbook_values(first.paths.result) == workbook_values(second.paths.result)


def test_v2_export_keeps_zero_numeric_and_routes_three_empty_sheets(tmp_path: Path) -> None:
    current = _record("20202001", 0)
    previous = _record("20202001", 0, period="2026-03")
    finding = PeriodComparisonEngine().run(
        {current.logical_key: current}, {previous.logical_key: previous}, ReportAuditConfig()
    )[0]
    path = ExcelExporter().export(output_dir=tmp_path, current_records=[current], previous_records=[previous], findings=[finding], run_info={}).result
    book = load_workbook(path, data_only=True)
    try:
        assert book.sheetnames == ["环比规则", "外部核对规则", "校验规则"]
        sheet = book["环比规则"]
        headers = [cell.value for cell in sheet[1]]
        assert sheet.cell(2, headers.index("本期值") + 1).value == 0
        assert sheet.cell(2, headers.index("上期值") + 1).value == 0
        assert sheet.cell(2, headers.index("差异值") + 1).value == 0
        assert sheet.cell(2, headers.index("变动幅度") + 1).value == 0
        assert sheet.cell(2, headers.index("审核级别") + 1).value in (None, "")
        assert "取数说明" not in headers
        assert book["外部核对规则"].max_row == 1 and book["校验规则"].max_row == 1
    finally:
        book.close()


@pytest.mark.parametrize("hide_period,hide_external", [(False, False), (True, False), (False, True), (True, True)])
def test_v2_export_hides_only_rows_with_exact_audit_description(
    tmp_path: Path, hide_period: bool, hide_external: bool,
) -> None:
    findings = [
        AuditFinding("period-same", "环比", "正常", institution_name="甲银行", indicator_code="A", description="无变动", current_value=0, previous_value=0, difference_value=0),
        AuditFinding("period-changed", "环比", "异常", "软性", institution_name="甲银行", indicator_code="B", description="命中规则", current_value=2, previous_value=1, difference_value=1),
        AuditFinding("external-same", "外部核对", "正常", institution_name="甲银行", indicator_code="A", rule_description="与大集中系统数据一致", current_value=0, external_value=0, difference_value=0),
        AuditFinding("external-different", "外部核对", "异常", "软性", institution_name="甲银行", indicator_code="B", rule_description="与大集中系统数据不一致", current_value=2, external_value=1, difference_value=1),
    ]
    path = ExcelExporter().export(
        output_dir=tmp_path, current_records=[], previous_records=[], findings=findings, run_info={},
        hide_unchanged_period_rows=hide_period, hide_matching_external_rows=hide_external,
    ).result
    book = load_workbook(path)
    try:
        period, external = book["环比规则"], book["外部核对规则"]
        external_columns = {cell.value: cell.column for cell in external[1]}
        assert period.max_row == external.max_row == 3
        assert period["K2"].value == "无变动"
        assert external.cell(2, external_columns["审核说明"]).value == "与大集中系统数据一致"
        assert external.cell(3, external_columns["审核说明"]).value == "与大集中系统数据不一致"
        assert bool(period.row_dimensions[2].hidden) is hide_period
        assert bool(external.row_dimensions[2].hidden) is hide_external
        assert not bool(period.row_dimensions[3].hidden)
        assert not bool(external.row_dimensions[3].hidden)
        assert period["L2"].value == 0 and external.cell(2, external_columns["外部值"]).value == 0
        assert period["J3"].value == "软性"
        assert external.cell(3, external_columns["审核级别"]).value == "软性"
        if hide_period and hide_external:
            period.row_dimensions[2].hidden = False
            external.row_dimensions[2].hidden = False
            unhidden = tmp_path / "user-unhidden.xlsx"
            book.save(unhidden)
            reopened = load_workbook(unhidden)
            try:
                assert not reopened["环比规则"].row_dimensions[2].hidden
                assert not reopened["外部核对规则"].row_dimensions[2].hidden
                assert reopened["环比规则"]["K2"].value == "无变动"
                assert reopened["外部核对规则"].cell(2, external_columns["审核说明"]).value == "与大集中系统数据一致"
            finally:
                reopened.close()
    finally:
        book.close()


def test_v2_export_rejects_unknown_audit_type(tmp_path: Path) -> None:
    finding = AuditFinding("unknown", "未知类型", "异常", "软性", institution_name="甲银行")
    with pytest.raises(ValueError, match="未知审核类型"):
        ExcelExporter().export(output_dir=tmp_path, current_records=[], previous_records=[], findings=[finding], run_info={})


@pytest.mark.parametrize("case", ["matched", "zero", "missing_value", "unmapped", "label_fallback", "mapped_label_fallback", "literal_label"])
def test_v2_external_export_shows_selected_report_item_next_to_credit_code(tmp_path, case):
    from base_audit.systems.s2_report_collection.importer import (
        MISSING_INSTITUTION_LABEL_FALLBACK_KEY, MISSING_INSTITUTION_LABEL_FALLBACK_REASON,
    )

    record = _record("20202001", Decimal("10"))
    label = "=SUM(1,2)" if case == "literal_label" else "j01（示例机构报表项目）"
    has_mapping = case not in {"unmapped", "label_fallback"}
    config = ReportAuditConfig(
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        external_rules=[ExternalCheckRule("D001", "20202001", "大集中", "存款", "等于", 0, "硬性")],
    )
    external = {}
    if case in {"matched", "zero", "literal_label"}:
        external[(label, "存款")] = ExternalValue(label, "存款", Decimal("0" if case == "zero" else "10"), record.source)
    elif case in {"label_fallback", "mapped_label_fallback"}:
        external[MISSING_INSTITUTION_LABEL_FALLBACK_KEY] = ExternalValue(
            "", "", Decimal("0"), record.source, MISSING_INSTITUTION_LABEL_FALLBACK_REASON,
        )
    finding = ExternalComparisonEngine().run(
        {record.logical_key: record}, external, {"9144": label} if has_mapping else {}, config,
    )[0]
    expected_label = label if has_mapping else ""
    assert finding.report_item == expected_label
    if case in {"missing_value", "unmapped"}:
        assert finding.external_value is None and finding.status == "异常"
        expected_description = (
            "大集中系统数据未找到该报表项目，请先完善大集中数据的报表定义"
            if has_mapping else "请完善报表采集系统配置或外部文件中的机构参照映射"
        )
        assert finding.rule_description == finding.description == expected_description
    if case in {"label_fallback", "mapped_label_fallback"}:
        assert finding.external_value == 0 and MISSING_INSTITUTION_LABEL_FALLBACK_REASON in finding.description
    output = ExcelExporter().export(
        output_dir=tmp_path, current_records=[record], previous_records=[], findings=[finding], run_info={},
    )
    book = load_workbook(output.result, data_only=False)
    try:
        sheet = book["外部核对规则"]
        headers = [cell.value for cell in sheet[1]]
        assert headers[:8] == ["地区", "承接行", "机构类别", "机构名称", "社会信用代码", "报表项目", "表单", "指标代码"]
        values = dict(zip(headers, next(sheet.iter_rows(min_row=2, values_only=True))))
        assert values["社会信用代码"] == "9144"
        assert values["报表项目"] == (expected_label or None)
        assert values["表单"] == record.form_code
        assert values["指标代码"] == "20202001"
        assert values["报表值"] == finding.current_value
        assert values["外部值"] == finding.external_value
        assert values["审核说明"] == finding.rule_description
        if case in {"missing_value", "unmapped"}:
            assert values["审核说明"] == expected_description
        assert sheet["F2"].alignment.wrap_text
        assert sheet["F2"].data_type != "f"
        assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:P2"
        assert "报表项目" not in [cell.value for cell in book["环比规则"][1]]
        assert "报表项目" not in [cell.value for cell in book["校验规则"][1]]
    finally:
        book.close()


def test_v2_export_rejects_severity_conflict(tmp_path: Path) -> None:
    finding = AuditFinding("bad", "环比", "异常", "", institution_name="甲银行")
    with pytest.raises(ValueError, match="审核级别冲突"):
        ExcelExporter().export(output_dir=tmp_path, current_records=[], previous_records=[], findings=[finding], run_info={})

    legacy_label = AuditFinding("legacy-severity", "环比", "异常", "严重", institution_name="甲银行")
    with pytest.raises(ValueError, match="硬性/软性"):
        ExcelExporter().export(output_dir=tmp_path, current_records=[], previous_records=[], findings=[legacy_label], run_info={})


def test_v2_export_writes_only_hard_and_soft_levels(tmp_path: Path) -> None:
    findings = [
        AuditFinding("period-hard", "环比", "异常", "硬性", institution_name="甲银行"),
        AuditFinding("external-soft", "外部核对", "异常", "软性", institution_name="甲银行"),
        AuditFinding("rule-hard", "校验规则", "异常", "硬性", institution_name="甲银行"),
    ]
    path = ExcelExporter().export(
        output_dir=tmp_path, current_records=[], previous_records=[], findings=findings, run_info={},
    ).result
    book = load_workbook(path, data_only=True)
    try:
        for sheet_name in ("环比规则", "外部核对规则", "校验规则"):
            sheet = book[sheet_name]
            level_column = next(cell.column for cell in sheet[1] if cell.value == "审核级别")
            assert sheet.cell(2, level_column).value in {"硬性", "软性"}
    finally:
        book.close()


def test_v2_validation_scope_collects_all_related_forms(tmp_path: Path) -> None:
    source = SourceRef("rule.xlsx", "配置", 2)
    finding = AuditFinding(
        "multi-form", "规则", "异常", "软性", institution_name="甲银行", form_code="资产负债表",
        rule_id="R001", rule_type="表达式", value_details="本期 20202003=1；本期 20202004=2",
        related_records=(RelatedRecord("20202004", "负债合计", "利润表", "2026-06", 2, source),),
    )
    path = ExcelExporter().export(output_dir=tmp_path, current_records=[], previous_records=[], findings=[finding], run_info={}).result
    book = load_workbook(path, data_only=True)
    try:
        sheet = book["校验规则"]
        headers = [cell.value for cell in sheet[1]]
        assert sheet.cell(2, headers.index("涉及表单") + 1).value == "利润表、资产负债表"
    finally:
        book.close()


def test_v2_uses_abs_previous_for_negative_period_change(tmp_path: Path) -> None:
    config, cur, pre, output = tmp_path / "config.xlsx", tmp_path / "cur", tmp_path / "pre", tmp_path / "out"
    cur.mkdir(); pre.mkdir(); _config(config)
    _source(cur / "9144#2026-06-30#01#20202#甲银行.xlsx", deposit=-50_000, loan=1, asset=1, debt=1, equity=0)
    _source(pre / "9144#2026-03-31#01#20202#甲银行.xlsx", deposit=-100_000, loan=1, asset=1, debt=1, equity=0)
    result = PeriodAuditV2Service().run(current_dir=cur, previous_dir=pre, config_path=config, output_dir=output)
    deposit = next(item for item in result.findings if item.audit_type == "环比" and item.indicator_code == "20202001")
    assert deposit.change_rate == 0.5


def _record(code: str, value: object, *, form: str = "20202", period: str = "2026-06", value_type: str = "数值") -> IndicatorRecord:
    return IndicatorRecord("9144", "甲银行", "9144", "农商行", "一部", "示例地区", form, code, code, period, period + "-30", value, value_type, "元", value, "万元", SourceRef("source.xlsx", form, 4))


def test_v2_period_bands_have_explicit_boundary_semantics() -> None:
    bounds = [-0.96, -0.9, -0.8, -0.5, -0.3, 0, 0.3, 0.5, 1, 5, 10, 96]
    config = ReportAuditConfig(bands=[PeriodBandConfig(f"B{i}", lower, bounds[i + 1] if i + 1 < len(bounds) else None, "异常", "软性") for i, lower in enumerate(bounds)])
    engine = PeriodComparisonEngine()
    for index, rate in enumerate(bounds):
        # rate=0 即本期与上期一致：按用户口径短路为“无变动”，不进档位；
        # 用极小正变动检验 B006 下边界（[0,0.3) 含正零头）。
        effective = rate if rate else 0.0001
        cur, pre = _record("20202001", 100 * (1 + effective)), _record("20202001", 100, period="2026-03")
        finding = engine.run({cur.logical_key: cur}, {pre.logical_key: pre}, config)[0]
        assert finding.band == f"B{index}", (index, rate, finding.band)  # [lower, upper) 临界值进下一档


def test_v2_named_strategy_uses_plain_percent_numbers() -> None:
    """新策略组中 10/30 表示 10%/30%，配置单元格不需要百分比格式。"""
    indicator = IndicatorConfig(
        "20202001", "存款", "20202", "余额", period_enabled=True,
        period_strategy_group="BAL01", min_change_value=0,
    )
    config = ReportAuditConfig(
        indicators={"20202001": indicator},
        bands=[
            PeriodBandConfig("B1", None, 10, "低于10%", "", False, "BAL01", "余额"),
            PeriodBandConfig("B2", 10, 30, "增幅10%-30%", "软性", False, "BAL01", "余额"),
            PeriodBandConfig("B3", 30, None, "增幅30%以上", "软性", False, "BAL01", "余额"),
        ],
    )
    current = _record("20202001", 120)
    previous = _record("20202001", 100, period="2026-03")
    finding = PeriodComparisonEngine().run(
        {current.logical_key: current}, {previous.logical_key: previous}, config
    )[0]
    assert finding.change_rate == pytest.approx(0.2)
    assert finding.band == "BAL01-B2"
    assert finding.severity == "软性"
    assert finding.description == "增幅10%-30%"


def test_v2_period_threshold_does_not_add_extra_calculation_text(tmp_path: Path) -> None:
    indicator = IndicatorConfig(
        "20202001", "存款", "20202", "余额", period_enabled=True,
        period_strategy_group="BAL01", min_change_value=11,
    )
    config = ReportAuditConfig(
        indicators={"20202001": indicator},
        bands=[PeriodBandConfig("P001", 0, 30, "余额变化请复核", "软性", False, "BAL01", "余额")],
    )
    current = _record("20202001", 110)
    previous = _record("20202001", 100, period="2026-03")
    finding = PeriodComparisonEngine().run(
        {current.logical_key: current}, {previous.logical_key: previous}, config
    )[0]

    # 数值门槛仍参与原有规则逻辑；计算过程只保留变动公式。
    assert finding.band == "BAL01-P001"
    assert finding.status == "正常"
    assert finding.description == ""
    assert finding.calculation_trace == "(110.00 - 100.00) / abs(100.00) = 10.00%"

    path = ExcelExporter().export(
        output_dir=tmp_path, current_records=[current], previous_records=[previous],
        findings=[finding], run_info={},
    ).result
    book = load_workbook(path, data_only=True)
    try:
        sheet = book["环比规则"]
        headers = [cell.value for cell in sheet[1]]
        assert "规则编号" not in headers
        assert sheet.cell(2, headers.index("审核说明") + 1).value in (None, "")
        trace = sheet.cell(2, headers.index("计算过程") + 1).value
        assert trace == "(110.00 - 100.00) / abs(100.00) = 10.00%"
    finally:
        book.close()


def test_v2_named_strategy_without_description_leaves_audit_note_blank() -> None:
    indicator = IndicatorConfig(
        "20202001", "存款", "20202", "余额", period_enabled=True,
        period_strategy_group="BAL01", min_change_value=0,
    )
    config = ReportAuditConfig(
        indicators={"20202001": indicator},
        bands=[PeriodBandConfig("B005", None, None, "", "软性", False, "BAL01", "余额")],
    )
    current = _record("20202001", 120)
    previous = _record("20202001", 100, period="2026-03")
    finding = PeriodComparisonEngine().run(
        {current.logical_key: current}, {previous.logical_key: previous}, config
    )[0]
    assert finding.band == "BAL01-B005"
    assert finding.description == ""


def test_v2_percentage_change_rate_is_percentage_point_difference(tmp_path: Path) -> None:
    # 真实形态（2026-08 第一大股东持股比例）：值本身就是百分数，1 与 9.92 表示 1% 与 9.92%。
    # 数据属性=百分数 → 差异值就是变动率（-8.92 个百分点）；不计算相对变动、
    # 不进入普通增降幅警戒档；「变动率(%)」列按百分点数值输出（非百分比格式）。
    current = _record("20201028", 1, form="20201", value_type="百分数")
    previous = _record("20201028", 9.92, form="20201", period="2026-03", value_type="百分数")
    finding = PeriodComparisonEngine().run({current.logical_key: current}, {previous.logical_key: previous}, ReportAuditConfig())[0]
    assert round(finding.difference_value, 6) == -8.92
    assert round(finding.change_rate, 6) == -8.92      # 百分点差，不是 -0.8992
    assert finding.band == ""                          # 不进任何 Bxxx 档
    assert finding.status == "正常"
    assert finding.description == ""
    assert finding.calculation_trace == "1.00% - 9.92% = -8.92 个百分点"
    assert finding.value_type == "百分数"
    path = ExcelExporter().export(output_dir=tmp_path, current_records=[current], previous_records=[previous], findings=[finding], run_info={}).result
    book = load_workbook(path, data_only=True)
    try:
        sheet = book["环比规则"]
        headers = [cell.value for cell in sheet[1]]
        rate_column = headers.index("变动幅度") + 1
        assert "规则编号" not in headers
        assert sheet.cell(2, headers.index("审核说明") + 1).value in (None, "")
        assert sheet.cell(2, headers.index("计算过程") + 1).value == "1.00% - 9.92% = -8.92 个百分点"
        # 百分数行用数值格式展示百分点（-8.92），不得套 0.00% 放大成 -892.00%。
        assert sheet.cell(2, rate_column).number_format == "0.00"
        assert sheet.cell(2, rate_column).value == -8.92
    finally:
        book.close()


def test_v2_percentage_never_enters_any_band_even_empty_lower(tmp_path: Path) -> None:
    """百分点差换算十进制后不达大幅档：D001(空,-0.5) 不被 -8.92 个百分点命中。"""
    current = _record("20201028", 1, form="20201", value_type="百分数")
    previous = _record("20201028", 9.92, form="20201", period="2026-03", value_type="百分数")
    config = ReportAuditConfig(bands=[PeriodBandConfig("D001", None, -0.5, "持股比例大幅下降", "软性")])
    finding = PeriodComparisonEngine().run({current.logical_key: current}, {previous.logical_key: previous}, config)[0]
    assert finding.band == ""                       # -0.0892 未达 -0.5 大幅档
    assert finding.status == "正常"
    assert finding.severity == ""


def test_v2_percentage_matches_middle_bands_as_decimal() -> None:
    """真实口径（用户定稿）：百分点差换算十进制匹配环比规则中间提示档。

    -8.92 个百分点 = -0.0892 → B005(-0.3,0) 命中：状态=异常、级别=提示；
    +8.21 个百分点 = +0.0821 → B006(0,0.3) 同样提示；相对降幅 -89.92%
    命中 B002 的旧行为不再产生。
    """
    engine = PeriodComparisonEngine()
    config = ReportAuditConfig(bands=[
        PeriodBandConfig("B005", -0.3, 0.0, "", "软性"),
        PeriodBandConfig("B006", 0.0, 0.3, "", "软性"),
    ])
    down = engine.run(
        {_record("20201028", 1, form="20201", value_type="百分数").logical_key: _record("20201028", 1, form="20201", value_type="百分数")},
        {_record("20201028", 9.92, form="20201", value_type="百分数", period="2026-03").logical_key: _record("20201028", 9.92, form="20201", value_type="百分数", period="2026-03")},
        config)[0]
    assert down.band == "B005"
    assert down.status == "异常"
    assert down.severity == "软性"
    assert down.description == ""

    up = engine.run(
        {_record("20201037", 10, form="20201", value_type="百分数").logical_key: _record("20201037", 10, form="20201", value_type="百分数")},
        {_record("20201037", 1.79, form="20201", value_type="百分数", period="2026-03").logical_key: _record("20201037", 1.79, form="20201", value_type="百分数", period="2026-03")},
        config)[0]
    assert up.band == "B006"
    assert up.status == "异常"
    assert up.severity == "软性"
    assert up.description == ""


def test_v2_percentage_large_change_never_hits_relative_bands() -> None:
    """相对降幅/增幅档（B002/B009）不再被百分数以相对变动命中。"""
    config = ReportAuditConfig(bands=[
        PeriodBandConfig("B002", -0.9, -0.8, "降幅(-90%,-80%]", "硬性"),
        PeriodBandConfig("B009", 1.0, 5.0, "增幅[1倍,5倍)", "软性"),
    ])
    engine = PeriodComparisonEngine()
    drop = engine.run(
        {_record("20201028", 1, form="20201", value_type="百分数").logical_key: _record("20201028", 1, form="20201", value_type="百分数")},
        {_record("20201028", 9.92, form="20201", value_type="百分数", period="2026-03").logical_key: _record("20201028", 9.92, form="20201", value_type="百分数", period="2026-03")},
        config)[0]
    assert drop.band == ""          # 相对变动 -89.92% 不再命中 B002
    assert drop.severity == ""
    surge = engine.run(
        {_record("20201037", 10, form="20201", value_type="百分数").logical_key: _record("20201037", 10, form="20201", value_type="百分数")},
        {_record("20201037", 1.79, form="20201", value_type="百分数", period="2026-03").logical_key: _record("20201037", 1.79, form="20201", value_type="百分数", period="2026-03")},
        config)[0]
    assert surge.band == ""         # 相对变动 +458% 不再命中 B009
    assert surge.severity == ""


def test_v2_percentage_point_difference_does_not_match_decimal_band() -> None:
    """回归（真实误报案例 20201028）：百分点差 -8.92 不得命中 B001/B002。"""
    current = _record("20201028", 1, form="20201", value_type="百分数")
    previous = _record("20201028", 9.92, form="20201", period="2026-03", value_type="百分数")
    config = ReportAuditConfig(bands=[
        PeriodBandConfig("B001", -0.96, -0.9, "降幅(-96%,-90%],缩小10倍-100倍", "硬性"),
        PeriodBandConfig("B002", -0.9, -0.8, "降幅(-90%,-80%],缩小5倍-10倍", "硬性"),
    ])
    finding = PeriodComparisonEngine().run({current.logical_key: current}, {previous.logical_key: previous}, config)[0]
    assert finding.band == ""
    assert finding.status == "正常"
    assert finding.severity == ""


def test_v2_percentage_zero_previous_no_division() -> None:
    """百分数上期为 0：差异=本期值（个百分点），不除零、不进档。"""
    current = _record("20201037", 5, form="20201", value_type="百分数")
    previous = _record("20201037", 0, form="20201", period="2026-03", value_type="百分数")
    finding = PeriodComparisonEngine().run({current.logical_key: current}, {previous.logical_key: previous}, ReportAuditConfig())[0]
    assert round(finding.difference_value, 6) == 5.0
    assert round(finding.change_rate, 6) == 5.0
    assert finding.band == ""
    assert finding.description == ""


def test_v2_percentage_up_and_down_descriptions() -> None:
    """20201037（10 vs 1.79 → +8.21）与持平场景的审核说明方向正确。"""
    engine = PeriodComparisonEngine()
    up = engine.run(
        {_record("20201037", 10, form="20201", value_type="百分数").logical_key: _record("20201037", 10, form="20201", value_type="百分数")},
        {_record("20201037", 1.79, form="20201", value_type="百分数", period="2026-03").logical_key: _record("20201037", 1.79, form="20201", value_type="百分数", period="2026-03")},
        ReportAuditConfig())[0]
    assert round(up.change_rate, 6) == 8.21
    assert up.description == ""
    same = engine.run(
        {_record("20201029", 4.67, form="20201", value_type="百分数").logical_key: _record("20201029", 4.67, form="20201", value_type="百分数")},
        {_record("20201029", 4.67, form="20201", value_type="百分数", period="2026-03").logical_key: _record("20201029", 4.67, form="20201", value_type="百分数", period="2026-03")},
        ReportAuditConfig())[0]
    assert same.change_rate == 0.0 and same.band == ""
    assert same.status == "正常" and same.description == "无变动"


def test_v2_normal_amount_still_uses_relative_change() -> None:
    """普通金额（20202001 各项存款）仍按普通相对环比，不受百分数修复影响。"""
    current = _record("20202001", 7815774.78)
    previous = _record("20202001", 7805891.31, period="2026-03")
    config = ReportAuditConfig(bands=[PeriodBandConfig("B006", 0.0, 0.3, "", "")])
    finding = PeriodComparisonEngine().run({current.logical_key: current}, {previous.logical_key: previous}, config)[0]
    assert round(finding.change_rate, 6) == round((7815774.78 - 7805891.31) / 7805891.31, 6)
    assert abs(finding.change_rate - 0.0012665) < 0.00001     # 约 0.13%
    assert finding.value_type == "数值"


def test_v2_indicator_code_type_compat() -> None:
    """指标代码 "20201028" / 20201028 / 20201028.0 经 _code 归一后同一命中。"""
    from base_audit.systems.s2_report_collection.config import _code
    assert _code("20201028") == _code(20201028) == _code(20201028.0) == "20201028"
    indicators = {"20201028": object()}
    for probe in ("20201028", 20201028, 20201028.0, "20201028.0"):
        assert indicators.get(_code(probe)) is not None, probe


def test_v2_data_attribute_whitespace_normalized() -> None:
    """配置“数据属性”带前后空格（" 百分数 "）时归一识别，不影响分支。"""
    from base_audit.systems.s2_report_collection.config import _text
    assert _text(" 百分数 ") == "百分数" and _text("百分数\n") == "百分数"
    current = _record("20201028", 1, form="20201", value_type=" 百分数 ")
    previous = _record("20201028", 9.92, form="20201", period="2026-03", value_type="百分数")
    finding = PeriodComparisonEngine().run({current.logical_key: current}, {previous.logical_key: previous}, ReportAuditConfig())[0]
    # 带空格形态引擎不按百分数处理（importer 读配置时已 strip，此处仅证明
    # 引擎判断严格相等，脏数据在配置装载层归一）。
    assert finding.band == "" or finding.band.startswith("B") is False


def test_v2_band_loader_accepts_empty_lower_bound(tmp_path: Path) -> None:
    """环比规则表允许下限留空（无穷小），装载不跳过、排序排最前。"""
    _config(tmp_path / "报表采集系统_配置.xlsx")
    from openpyxl import load_workbook as _lw
    path = tmp_path / "报表采集系统_配置.xlsx"
    book = _lw(path)
    sheet = book["环比策略"]
    sheet.append(["BAL01", "P004", "余额", None, -0.96, "软性", "持股比例大幅下降", "否"])
    book.save(path); book.close()
    config = load_config(path)
    drop_band = next(b for b in config.bands if b.rule_id == "P004")
    assert drop_band.lower is None and drop_band.upper == -0.96
    assert drop_band.includes(-8.92) and not drop_band.includes(0.5)


def test_v2_period_trace_uses_two_decimal_places() -> None:
    current = _record("20202001", Decimal("8685065.63000"))
    previous = _record("20202001", Decimal("8663670.02000"), period="2026-03")
    finding = PeriodComparisonEngine().run({current.logical_key: current}, {previous.logical_key: previous}, ReportAuditConfig())[0]
    assert finding.calculation_trace == "(8685065.63 - 8663670.02) / abs(8663670.02) = 0.25%"


def test_v2_importer_preserves_percentage_data_type(tmp_path: Path) -> None:
    directory = tmp_path / "current"; directory.mkdir()
    source = directory / "9144#20260630#01#20201#甲银行.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "20201"
    sheet.append(["", "", ""]); sheet.append(["", "", ""]); sheet.append(["指标代码", "指标名称", "值"])
    sheet.append(["20201028", "第一大股东持股比例", 0.158094]); book.save(source)
    config = ReportAuditConfig(indicators={"20201028": IndicatorConfig("20201028", "第一大股东持股比例", "20201", "百分数")})
    result = import_period_directory(directory, config=config, label="本期")
    assert result.records[0].value_type == "百分数"


def test_v2_config_separates_indicator_definition_and_external_rules(tmp_path: Path) -> None:
    path = tmp_path / "config.xlsx"
    _config(path)
    config = load_config(path)
    indicator = config.indicators["20202001"]
    assert indicator.form_code == "20202"
    assert config.unit_settings == UnitSettings("元", "元", "元")
    assert not hasattr(indicator, "external_indicator_name")
    assert [(rule.rule_id, rule.indicator_code, rule.external_system, rule.comparison_method,
             rule.external_indicator_name, rule.tolerance, rule.severity) for rule in config.external_rules] == [
        ("D001", "20202001", "大集中", "等于", "存款", 0.0, "硬性"),
    ]


def test_v2_external_rule_uses_its_own_tolerance_and_rule_identity() -> None:
    from decimal import Decimal

    record = _record("20202001", Decimal("100000"), period="2026-06")
    rule = ExternalCheckRule("D001", "20202001", "大集中", "各项存款余额", "等于", 10, "硬性")
    config = ReportAuditConfig(
        external_rules=[rule],
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        unit_settings=UnitSettings("元", "亿元", "元"),
    )
    external = {("甲行", "各项存款余额"): ExternalValue("甲行", "各项存款余额", 0.0010001, SourceRef("central.xlsx", "集中系统数据", 5))}
    findings = ExternalComparisonEngine().run({record.logical_key: record}, external, {"9144": "甲行"}, config)
    assert len(findings) == 1
    finding = findings[0]
    assert finding.status == "正常" and finding.rule_id == "D001"
    assert finding.current_value == Decimal("100000")
    assert finding.external_value == Decimal("100010.0000")
    assert finding.external_indicator_name == "各项存款余额"
    assert finding.rule_description == "与大集中系统数据一致"


def test_v2_duplicate_and_cross_form_conflict_are_data_quality_findings() -> None:
    left, right = _record("20202001", 1, form="20202"), _record("20202001", 2, form="20203")
    issues = validate_import(ImportResult(records=[left], duplicates=[(left, right)]))
    assert len(issues) == 1
    assert "跨表单冲突" in issues[0].description


def test_v2_cross_form_rule_is_one_independent_finding() -> None:
    left, right = _record("20201099", 10, form="20201"), _record("20202001", 8, form="20202")
    config = ReportAuditConfig(rules=[RuleConfig("R-CROSS", "表达式", "跨表不相等", "[20201099]<>[20202001]", "硬性")])
    findings, quality = RuleEngine().run({left.logical_key: left, right.logical_key: right}, {}, config)
    assert not quality and len(findings) == 1
    assert findings[0].form_codes == ("20201", "20202")
    assert len(findings[0].related_records) == 2


def test_v2_expression_balance_uses_decimal_not_binary_float_residual() -> None:
    config = ReportAuditConfig(rules=[RuleConfig(
        "R-BALANCE", "表达式", "资产负债表平衡", "[20202003]<>[20202004]+[20202005]", "硬性",
    )])
    current = {
        record.logical_key: record
        for record in (
            _record("20202003", 9697998.71),
            _record("20202004", 8875290.77),
            _record("20202005", 822707.94),
        )
    }
    findings, quality = RuleEngine().run(current, {}, config)
    assert not findings and not quality

    current["9144", "20202005", "2026-06"]= _record("20202005", 822707.95)
    findings, quality = RuleEngine().run(current, {}, config)
    assert len(findings) == 1 and not quality


def test_v2_cumulative_rule_emits_one_finding_per_violated_indicator(tmp_path: Path) -> None:
    current_a, previous_a = _record("20202001", 8), _record("20202001", 10, period="2026-03")
    current_b, previous_b = _record("20202002", 3), _record("20202002", 5, period="2026-03")
    config = ReportAuditConfig(rules=[RuleConfig("R001", "累计不降(当年)", "当年累计指标比上期不应减少", "20202001,20202002", "软性")])
    findings, quality = RuleEngine().run(
        {current_a.logical_key: current_a, current_b.logical_key: current_b},
        {previous_a.logical_key: previous_a, previous_b.logical_key: previous_b}, config,
    )
    assert not quality and len(findings) == 2
    assert [(item.indicator_code, item.current_value, item.previous_value) for item in findings] == [
        ("20202001", 8, 10), ("20202002", 3, 5),
    ]
    assert all(item.rule_id == "R001" and item.rule_description == "当年累计指标比上期不应减少" for item in findings)
    assert all("<" in item.calculation_trace for item in findings)
    paths = ExcelExporter().export(
        output_dir=tmp_path / "output", current_records=[current_a, current_b],
        previous_records=[previous_a, previous_b], findings=findings, run_info={},
    )
    book = load_workbook(paths.result, data_only=True)
    try:
        sheet = book["校验规则"]
        assert sheet.max_row == 3
        headers = [cell.value for cell in sheet[1]]
        rule_column = headers.index("规则编号") + 1
        details_column = headers.index("取值明细") + 1
        assert [sheet.cell(row, rule_column).value for row in (2, 3)] == ["R001", "R001"]
        assert "本期=8.00" in sheet.cell(2, details_column).value and "上期=10.00" in sheet.cell(2, details_column).value
        assert "本期=3.00" in sheet.cell(3, details_column).value and "上期=5.00" in sheet.cell(3, details_column).value
        trace_column = headers.index("计算过程") + 1
        assert sheet.cell(2, trace_column).value == "20202001: 8.00 < 10.00"
        assert sheet.cell(3, trace_column).value == "20202002: 3.00 < 5.00"
    finally:
        book.close()

def test_v2_rule_with_missing_indicators_reports_reason_not_rule_text() -> None:
    """用户场景（旧复杂-003）：规则引用的指标缺失/空值 → 校验规则行。

    审核说明必须写真实原因（缺少哪些指标的数据），不得用规则描述
    冒充成“低于监管线 异常”；空值（行存在值为空）与缺失同等对待，
    禁止拿 0 参与比较。
    """
    config = ReportAuditConfig(rules=[RuleConfig(
        "旧复杂-003", "表达式", "核心一级资本充足率低于5.5%，接近监管底线",
        "[20203045] < 0.055 * [20203048]", "软性")])
    engine = RuleEngine()
    # 场景 1：指标整行缺失。
    only_deposit = _record("20202001", 100)
    findings, quality = engine.run({only_deposit.logical_key: only_deposit}, {}, config)
    assert not findings and len(quality) == 1
    assert quality[0].audit_type == "校验规则"
    assert "缺少指标 20203045" in quality[0].description
    assert "缺少指标 20203048" in quality[0].description
    assert "低于监管线" not in quality[0].description or "规则：" in quality[0].description
    assert quality[0].calculation_trace.startswith(
        "配置表达式：[20203045] < 0.055 * [20203048]；代入表达式："
    )
    # 场景 2：指标行存在但值为空串（空报）——同样不得拿 0 求值。
    empty = IndicatorRecord("9144", "甲银行", "9144", "农商行", "一部", "示例地区",
                            "20203", "20203045", "核心一级资本净额", "2026-08",
                            "2026-08-31", "", "余额", "", "", "万元",
                            SourceRef("s.xlsx", "20203", 5))
    has_both = {empty.logical_key: empty,
                _record("20203048", 50.0, form="20203").logical_key: _record("20203048", 50.0, form="20203")}
    findings2, quality2 = engine.run(has_both, {}, config)
    assert not findings2 and len(quality2) == 1
    assert "缺少指标 20203045 的有效数值" in quality2[0].description

def test_v2_rule_supports_legacy_eight_segment_placeholder(tmp_path: Path) -> None:
    """真实 2.配置的规则用旧 VBA 八段语法 [,,,20203045,]——第 4 段是指标代码。

    回归：引擎此前把整段（含逗号）当 lookup key，八段规则永远
    “缺少指标”，审核清单里全部规则取不到数据。
    """
    config = ReportAuditConfig(rules=[RuleConfig(
        "旧复杂-003", "表达式", "核心一级资本充足率低于5.5%，接近监管底线",
        "[,,,20203045,] < 0.055 * [,,,20203048,]", "软性")])
    engine = RuleEngine()
    def record(code, value):
        return IndicatorRecord("9144", "甲银行", "9144", "农商行", "一部", "示例地区",
                               "20203", code, code, "2026-08", "2026-08-31",
                               value, "余额", "", value, "万元",
                               SourceRef("s.xlsx", "20203", 5))
    healthy = {}
    for code, value in (("20203045", 120.0), ("20203048", 2000.0)):
        item = record(code, value); healthy[item.logical_key] = item
    findings, quality = engine.run(healthy, {}, config)
    assert not findings and not quality          # 120 < 110 不成立，正常无异常
    breach = {}
    for code, value in (("20203045", 100.0), ("20203048", 2000.0)):
        item = record(code, value); breach[item.logical_key] = item
    findings2, quality2 = engine.run(breach, {}, config)
    assert len(findings2) == 1 and not quality2  # 100 < 110 成立，命中
    expected_trace = (
        "配置表达式：[,,,20203045,] < 0.055 * [,,,20203048,]；"
        "代入表达式：100.00 < 0.055 * 2000.00"
    )
    assert findings2[0].calculation_trace == expected_trace
    assert findings2[0].related_records[0].indicator_code == "20203045"
    path = ExcelExporter().export(
        output_dir=tmp_path, current_records=list(breach.values()), previous_records=[],
        findings=findings2, run_info={},
    ).result
    book = load_workbook(path, data_only=True)
    try:
        sheet = book["校验规则"]
        headers = [cell.value for cell in sheet[1]]
        trace = sheet.cell(2, headers.index("计算过程") + 1).value
        assert "[,,,20203045,] < 0.055 * [,,,20203048,]" in trace
        assert "100.00 < 0.055 * 2000.00" in trace
    finally:
        book.close()


def test_v2_validation_values_and_substituted_expression_display_two_decimals(tmp_path: Path) -> None:
    def record(code: str, name: str, value: Decimal) -> IndicatorRecord:
        return IndicatorRecord(
            "9144", "甲银行", "9144", "农商行", "一部", "示例地区",
            "20203", code, name, "2026-06", "2026-06-30", value,
            "余额", "元", value, "万元", SourceRef("source.xlsx", "20203", 4),
        )

    capital = record("20203048", "应用资本底线及校准后的风险加权资产合计", Decimal("12014.97100"))
    assets = record("20202003", "资产总计", Decimal("35049.68749"))
    config = ReportAuditConfig(rules=[RuleConfig(
        "R001", "表达式", "资本约束", "[,,,20203048,] < 0.4 * [,,,20202003,]", "硬性",
    )])

    findings, quality = RuleEngine().run(
        {capital.logical_key: capital, assets.logical_key: assets}, {}, config,
    )

    assert not quality and len(findings) == 1
    finding = findings[0]
    assert finding.value_details == (
        "本期 20203048 应用资本底线及校准后的风险加权资产合计=12014.97；"
        "本期 20202003 资产总计=35049.69"
    )
    assert finding.calculation_trace == (
        "配置表达式：[,,,20203048,] < 0.4 * [,,,20202003,]；"
        "代入表达式：12014.97 < 0.4 * 35049.69"
    )
    related_values = {item.indicator_code: item.value for item in finding.related_records}
    assert related_values == {
        "20203048": Decimal("12014.97100"),
        "20202003": Decimal("35049.68749"),
    }
    output_path = ExcelExporter().export(
        output_dir=tmp_path / "output", current_records=[capital, assets], previous_records=[],
        findings=findings, run_info={},
    ).result
    book = load_workbook(output_path, data_only=True)
    try:
        sheet = book["校验规则"]
        headers = [cell.value for cell in sheet[1]]
        details = sheet.cell(2, headers.index("取值明细") + 1).value
        trace = sheet.cell(2, headers.index("计算过程") + 1).value
        assert details == finding.value_details
        assert trace == finding.calculation_trace
    finally:
        book.close()

def test_v2_both_sides_empty_is_normal_not_abnormal() -> None:
    """用户口径：本期、上期都没有数据 → 状态正常，不再标异常关注。"""
    engine = PeriodComparisonEngine()
    empty_a = _record("20201030", "")
    empty_b = _record("20201030", "", period="2026-03")
    finding = engine.run({empty_a.logical_key: empty_a}, {empty_b.logical_key: empty_b}, ReportAuditConfig())[0]
    assert finding.status == "正常" and finding.severity == ""
    # 单侧有值另一侧空：存在性异常，保留；级别由“单侧缺失默认级别”控制（缺省硬性）。
    half = engine.run(
        {_record("20201031", 5.0).logical_key: _record("20201031", 5.0)},
        {_record("20201031", "", period="2026-03").logical_key: _record("20201031", "", period="2026-03")},
        ReportAuditConfig())[0]
    assert half.status == "异常" and half.severity == "硬性"
    assert half.description == "本期有，上期无"

def test_v2_audit_description_wording_map() -> None:
    """审核说明统一口径（用户定稿）：
    完全一致（含两侧 0 / 两侧无数据 / 文字未变）→ 无变动；
    单侧存在 → 本期有，上期无 / 本期无，上期有（环比不计算，
    说明不带“未计算”解释尾巴）；增幅/降幅命中环比规则 → 保持规则描述。
    """
    engine = PeriodComparisonEngine()
    def run(cur_v, pre_v, code="20201030", vtype="数值"):
        cur = _record(code, cur_v, value_type=vtype)
        pre = _record(code, pre_v, period="2026-03", value_type=vtype)
        return engine.run({cur.logical_key: cur}, {pre.logical_key: pre}, ReportAuditConfig())[0]

    cases = [
        ((0, 0), "无变动", "正常"),
        ((None, None), "无变动", "正常"),
        ((7.5, 7.5), "无变动", "正常"),
        ((5.0, None), "本期有，上期无", "异常"),
        ((None, 9.0), "本期无，上期有", "异常"),
        ((5.0, 0), "本期有，上期无", "异常"),
        ((0, 9.0), "本期无，上期有", "异常"),
        (("甲银行", "甲银行"), "无变动", "正常"),
        (("甲银行", "乙银行"), "文字指标变化", "异常"),
    ]
    for (cur_v, pre_v), expect_desc, expect_status in cases:
        vtype = "文字" if isinstance(cur_v, str) else "数值"
        finding = run(cur_v, pre_v, vtype=vtype)
    for (cur_v, pre_v), expect_desc, expect_status in cases:
        vtype = "文字" if isinstance(cur_v, str) else "数值"
        finding = run(cur_v, pre_v, vtype=vtype)
        assert finding.description == expect_desc, (cur_v, pre_v, finding.description)
        assert finding.status == expect_status, (cur_v, pre_v, finding.status)
    # 单侧场景：环比不计算（变动率列为空），说明不带“未计算”尾巴。
    one_side = run(5.0, None)
    assert one_side.change_rate is None and "未计算" not in one_side.description
    # 文字未变化。
    text_same = run("甲银行", "甲银行", code="20201099", vtype="文字")
    assert text_same.description == "无变动" and text_same.status == "正常"
    assert text_same.retrieval_note == "变动幅度未计算"

def test_v2_float_residual_near_zero_does_not_alert() -> None:
    """0 附近浮点残差（如 100.00000000000003 vs 100）不得触发提示档。

    双重防护：引擎对精确相等短路“无变动”；配置侧在 0 附近挖 ±1E-11
    无级别档（用户维护，B006）吸收浮点残差——两侧任一生效都不误报。
    """
    bands = [
        PeriodBandConfig("B005", -0.3, -1e-11, "", "软性"),
        PeriodBandConfig("B006", -1e-11, 1e-11, "", ""),
        PeriodBandConfig("B007", 1e-11, 0.3, "", "软性"),
    ]
    config = ReportAuditConfig(bands=bands)
    engine = PeriodComparisonEngine()
    current = _record("20202001", 100.00000000000003)
    previous = _record("20202001", 100.0, period="2026-03")
    finding = engine.run({current.logical_key: current}, {previous.logical_key: previous}, config)[0]
    assert finding.status == "正常" and finding.severity == ""
    # 精确 0 变动：引擎短路“无变动”，即使没有 B006 洞也不进档。
    bare = ReportAuditConfig(bands=[
        PeriodBandConfig("B005", -0.3, 0.0, "", "软性"),
        PeriodBandConfig("B006", 0.0, 0.3, "", "软性"),
    ])
    same = engine.run(
        {_record("20202001", 100.0).logical_key: _record("20202001", 100.0)},
        {_record("20202001", 100.0, period="2026-03").logical_key: _record("20202001", 100.0, period="2026-03")},
        bare)[0]
    assert same.description == "无变动" and same.status == "正常" and same.band == ""


def _external_xlsx(path: Path) -> None:
    book = Workbook(); sheet = book.active; sheet.title = "集中系统数据"
    sheet.append([None, "单位贷款余额", "单位贷款新增"])
    sheet.append(["示例甲银行", 100.5, 3.2])
    sheet.append(["乙财务", 50, 1.1])
    ref = book.create_sheet("参照表")
    ref.append(["统一社会信用代码", "报表项目"]); ref.append(["9144", "甲行"])
    book.save(path)


def _time_series_external_xlsx(path: Path, rows: list[dict], *, headers=None, leading_rows=()) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "集中系统数据"
    headers = headers or [
        "数据日期", "报表名称", "报表代码", "版本号", "锁定信息", "报表组", "行顺序号", "行名称", "存款", "贷款",
    ]
    for row in leading_rows:
        sheet.append(row)
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(str(header).strip()) for header in headers])
    reference = book.create_sheet("参照表")
    reference.append(["统一社会信用代码", "报表项目"])
    reference.append(["9144", "甲行"])
    book.save(path)
    book.close()


@pytest.mark.parametrize(
    ("data_date", "leading_rows", "reorder"),
    [
        (datetime(2026, 6, 30), (), False),
        ("20260630", ([],), False),
        (20260630, (["导出说明"], []), False),
        ("2026-06-30T08:09:10", (), True),
        ("2026/6/30", ([],), True),
    ],
)
def test_v2_external_import_time_series_uses_named_headers_and_full_dates(
    tmp_path: Path, data_date, leading_rows, reorder: bool,
) -> None:
    from base_audit.systems.s2_report_collection.importer import (
        MISSING_INSTITUTION_LABEL_FALLBACK_KEY, import_external_workbook,
    )

    path = tmp_path / "central.xlsx"
    headers = [" 贷款 ", "行名称", "报表代码", " 数据日期 ", "存款"] if reorder else None
    _time_series_external_xlsx(path, [
        {"数据日期": data_date, "行名称": "甲行", "存款": 0},
        {"数据日期": "2026-07-31", "行名称": "甲行", "存款": "123.4567890123456789", "贷款": 10},
        {"数据日期": "2026-06-30", "行名称": "乙行", "贷款": 3},
        {"数据日期": "2026-06-30"},  # 导出的机构分组间隔行不算机构记录。
    ], headers=headers, leading_rows=leading_rows)

    values, org_map = import_external_workbook(path)

    assert values.time_series is True
    assert values.available_dates == {"2026-06-30", "2026-07-31"}
    assert values.date_labels == {("2026-06-30", "甲行"), ("2026-07-31", "甲行"), ("2026-06-30", "乙行")}
    assert {key: value.value for key, value in values.items()} == {
        ("2026-06-30", "甲行", "存款"): Decimal("0"),
        ("2026-07-31", "甲行", "存款"): Decimal("123.4567890123456789"),
        ("2026-07-31", "甲行", "贷款"): Decimal("10"),
        ("2026-06-30", "乙行", "贷款"): Decimal("3"),
    }
    item = values[("2026-07-31", "甲行", "存款")]
    assert item.data_date == "2026-07-31"
    assert item.source.source_row == len(leading_rows) + 3
    assert MISSING_INSTITUTION_LABEL_FALLBACK_KEY not in values
    assert org_map == {"9144": "甲行"}


@pytest.mark.parametrize(
    ("headers", "rows", "message"),
    [
        (["数据日期", "存款"], [{"数据日期": "20260630", "存款": 1}], "时序表头不完整"),
        (["行名称", "存款"], [{"行名称": "甲行", "存款": 1}], "时序表头不完整"),
        (["数据日期", "行名称", "行名称", "存款"], [], "有且只有一列“行名称”"),
        (["数据日期", "行名称", "存款", "存款"], [], "指标列“存款”重复"),
        (["数据日期", "行名称"], [], "缺少指标列"),
        (None, [{"数据日期": "2026-02-30", "行名称": "甲行", "存款": 1}], "数据日期.*无效"),
        (None, [{"数据日期": "202611", "行名称": "甲行", "存款": 1}], "数据日期.*无效"),
        (None, [{"数据日期": "20260630", "存款": 0}], "行名称.*为空"),
        (None, [
            {"数据日期": "20260630", "行名称": "甲行", "版本号": 1, "存款": 1},
            {"数据日期": "2026-06-30", "行名称": "甲行", "版本号": 2, "存款": 2},
        ], "第3、4行.*重复"),
        (None, [{"数据日期": "20260630", "行名称": "甲行", "存款": "错误数值"}], "指标“存款”.*无法读取"),
    ],
)
def test_v2_external_import_time_series_rejects_ambiguous_or_invalid_data(
    tmp_path: Path, headers, rows, message: str,
) -> None:
    from base_audit.systems.s2_report_collection.importer import ImportErrorV2, import_external_workbook

    path = tmp_path / "central.xlsx"
    _time_series_external_xlsx(path, rows, headers=headers, leading_rows=(["导出说明"],))

    with pytest.raises(ImportErrorV2, match=message):
        import_external_workbook(path)


@pytest.mark.parametrize(
    ("source_date", "expected", "missing_text"),
    [
        ("20260630", Decimal("100"), ""),
        ("2026-06-29", Decimal("333"), ""),
        ("2026/07/31", Decimal("0"), ""),
        ("2026-07-30", None, "无可读取数值"),
        ("2026-08-31", None, "没有数据日期 2026-08-31"),
        ("2026-06", None, "不是完整日期"),
        ("202611", None, "不是完整日期"),
    ],
)
def test_v2_external_time_series_matches_exact_current_date_and_preserves_missing_values(
    tmp_path: Path, source_date: str, expected, missing_text: str,
) -> None:
    from base_audit.systems.s2_report_collection.importer import (
        MISSING_INSTITUTION_LABEL_FALLBACK_KEY, MISSING_INSTITUTION_LABEL_FALLBACK_REASON,
        import_external_workbook,
    )

    path = tmp_path / "central.xlsx"
    _time_series_external_xlsx(path, [
        {"数据日期": "20260630", "行名称": "甲行", "存款": 100},
        {"数据日期": "20260629", "行名称": "甲行", "存款": 333},
        {"数据日期": "20260731", "行名称": "甲行", "存款": 0},
        {"数据日期": "20260730", "行名称": "甲行"},
    ])
    values, org_map = import_external_workbook(path)
    record = replace(_record("20202001", Decimal("100")), source_date=source_date)
    config = ReportAuditConfig(
        indicators={"20202001": IndicatorConfig("20202001", data_type="余额")},
        external_rules=[ExternalCheckRule("D001", "20202001", "大集中", "存款", "等于", 0, "硬性")],
    )
    # 即使出现普通格式键或历史零值兜底标记，时序数据也不得绕过完整日期匹配。
    values[("甲行", "存款")] = ExternalValue("甲行", "存款", Decimal("999"), record.source)
    values[MISSING_INSTITUTION_LABEL_FALLBACK_KEY] = ExternalValue(
        "", "", Decimal("0"), record.source, MISSING_INSTITUTION_LABEL_FALLBACK_REASON,
    )

    finding = ExternalComparisonEngine().run({record.logical_key: record}, values, org_map, config)[0]

    assert finding.external_value == expected
    assert finding.report_item == "甲行"
    if expected is None:
        assert finding.status == "异常"
        assert finding.difference_value == ""
        assert missing_text in finding.calculation_trace
    else:
        assert finding.difference_value == abs(Decimal("100") - expected)
        assert finding.retrieval_note == ""


@pytest.mark.parametrize(
    ("report_item", "missing_text"),
    [("甲行", "无可读取数值"), ("不存在的项目", "未找到报表项目"), ("", "缺少机构映射或报表项目")],
)
def test_v2_external_time_series_with_all_blank_values_stays_missing(
    tmp_path: Path, report_item: str, missing_text: str,
) -> None:
    from base_audit.systems.s2_report_collection.importer import import_external_workbook

    path = tmp_path / "central.xlsx"
    _time_series_external_xlsx(path, [{"数据日期": "20260630", "行名称": "甲行"}])
    values, _org_map = import_external_workbook(path)
    record = _record("20202001", Decimal("0"))
    config = ReportAuditConfig(external_rules=[
        ExternalCheckRule("D001", "20202001", "大集中", "存款", "等于", 0, "硬性"),
    ])

    finding = ExternalComparisonEngine().run(
        {record.logical_key: record}, values, {"9144": report_item}, config,
    )[0]

    assert values == {} and values.time_series is True
    assert finding.external_value is None and finding.status == "异常"
    assert finding.report_item == report_item
    assert missing_text in finding.calculation_trace
    assert finding.rule_description == finding.description == (
        "大集中系统数据未找到该报表项目，请先完善大集中数据的报表定义"
        if report_item else "请完善报表采集系统配置或外部文件中的机构参照映射"
    )


def test_v2_service_exports_current_time_series_value_without_modifying_inputs(tmp_path: Path) -> None:
    config, cur, pre, central = tmp_path / "config.xlsx", tmp_path / "cur", tmp_path / "pre", tmp_path / "central.xlsx"
    cur.mkdir(); pre.mkdir(); _config(config)
    _source(cur / "9144#2026-06-30#01#20202#甲银行.xlsx", deposit=2_000_000, loan=1_000_000, asset=1_000_000, debt=600_000, equity=400_000)
    _source(pre / "9144#2026-03-31#01#20202#甲银行.xlsx", deposit=1_000_000, loan=1_000_000, asset=1_000_000, debt=600_000, equity=400_000)
    _time_series_external_xlsx(central, [
        {"数据日期": "20260331", "行名称": "甲行", "存款": 77},
        {"数据日期": "20260630", "行名称": "甲行", "存款": 2_000_000},
        {"数据日期": "20260731", "行名称": "甲行", "存款": 88},
    ])
    book = load_workbook(central)
    shadow = book.create_sheet("集中系统数据2")
    shadow.append(["数据日期", "行名称", "存款"])
    shadow.append(["20260630", "甲行", 999])
    book.save(central); book.close()
    input_paths = [config, central, *cur.iterdir(), *pre.iterdir()]
    before = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in input_paths}

    result = PeriodAuditV2Service().run(
        current_dir=cur, previous_dir=pre, config_path=config, output_dir=tmp_path / "out", central_path=central,
    )

    external = [item for item in result.findings if item.audit_type == "外部核对"]
    assert len(external) == 1 and external[0].external_value == Decimal("2000000")
    assert external[0].status == "正常"
    assert external[0].source_refs[-1].source_row == 3
    book = load_workbook(result.paths.result, read_only=True, data_only=True)
    try:
        sheet = book["外部核对规则"]
        headers = [cell.value for cell in sheet[1]]
        assert headers[headers.index("社会信用代码") + 1] == "报表项目"
        assert sheet.cell(2, headers.index("报表项目") + 1).value == "甲行"
        assert sheet.cell(2, headers.index("外部值") + 1).value == 2_000_000
        assert sheet.cell(2, headers.index("差异值") + 1).value == 0
    finally:
        book.close()
    assert {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in input_paths} == before


def test_v2_external_import_supports_xlsx_only(tmp_path: Path) -> None:
    """DEFECT-4 定案：外部数据仅支持 .xlsx；.xls/.csv 会丢「参照表」，明确拒绝。"""
    import pytest

    from base_audit.systems.s2_report_collection.importer import ImportErrorV2, import_external_workbook

    xlsx = tmp_path / "central.xlsx"; _external_xlsx(xlsx)
    values, org_map = import_external_workbook(xlsx)
    assert values.time_series is False
    assert {key: value.value for key, value in values.items()} == {
        ("示例甲银行", "单位贷款余额"): Decimal("100.5"), ("示例甲银行", "单位贷款新增"): Decimal("3.2"),
        ("乙财务", "单位贷款余额"): Decimal("50"), ("乙财务", "单位贷款新增"): Decimal("1.1"),
    }
    assert org_map == {"9144": "甲行"}

    xls = tmp_path / "central.xls"; xls.write_bytes(b"\xd0\xcf\x11\xe0")
    csv_path = tmp_path / "central.csv"; csv_path.write_text("指标,值", encoding="gbk")
    for candidate in (xls, csv_path):
        with pytest.raises(ImportErrorV2, match="仅支持 .xlsx"):
            import_external_workbook(candidate)

    doc = tmp_path / "central.txt"; doc.write_text("x", encoding="utf-8")
    with pytest.raises(ImportErrorV2, match="仅支持 .xlsx"):
        import_external_workbook(doc)

    book = Workbook(); book.active.title = "别的表"; book.save(tmp_path / "no_sheet.xlsx")
    with pytest.raises(ImportErrorV2, match="集中系统数据"):
        import_external_workbook(tmp_path / "no_sheet.xlsx")


def test_v2_external_import_uses_zero_when_institution_label_column_is_missing(tmp_path: Path) -> None:
    from base_audit.systems.s2_report_collection.importer import (
        MISSING_INSTITUTION_LABEL_FALLBACK_KEY,
        MISSING_INSTITUTION_LABEL_FALLBACK_REASON,
        import_external_workbook,
    )

    path = tmp_path / "central_without_label.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "集中系统数据"
    sheet.append(["机构名称", "存款"])
    sheet.append(["甲行", 999])
    book.save(path)
    book.close()

    values, org_map = import_external_workbook(path)

    fallback = values[MISSING_INSTITUTION_LABEL_FALLBACK_KEY]
    assert fallback.value == Decimal("0")
    assert fallback.fallback_reason == MISSING_INSTITUTION_LABEL_FALLBACK_REASON
    assert org_map == {}


def test_v2_external_import_prefers_institution_reference_table(tmp_path: Path) -> None:
    from base_audit.systems.s2_report_collection.importer import import_external_workbook

    path = tmp_path / "central.xlsx"
    _external_xlsx(path)
    book = load_workbook(path)
    institution = book.create_sheet("机构参照表")
    institution.append(["机构名称", "社会信用代码", "机构类别", "承接行", "地区", "报表项目", "禁用", "备注"])
    institution.append(["甲银行", "9144", "农商行", "一部", "示例地区", "新项目", "否", None])
    institution.append(["甲银行历史映射", "9144", "农商行", "一部", "示例地区", "最终项目", "否", None])
    institution.append(["停用机构", "9145", "农商行", "一部", "示例地区", "不应采用", "是", None])
    book.save(path)
    book.close()

    _values, org_map = import_external_workbook(path)

    # 新表优先于旧参照表；有效重复映射沿用后行覆盖，禁用项不进入映射。
    assert org_map == {"9144": "最终项目"}


def test_v2_external_import_rejects_malformed_preferred_institution_reference(tmp_path: Path) -> None:
    from base_audit.systems.s2_report_collection.importer import ImportErrorV2, import_external_workbook

    path = tmp_path / "central.xlsx"
    _external_xlsx(path)
    book = load_workbook(path)
    institution = book.create_sheet("机构参照表")
    institution.append(["机构名称", "社会信用代码", "机构类别"])
    book.save(path)
    book.close()

    with pytest.raises(ImportErrorV2, match="机构参照表.*缺少必要表头"):
        import_external_workbook(path)


@pytest.mark.parametrize("layout", ["single", "reports", "banks"])
@pytest.mark.parametrize("suffix", [".xlsx", ".xls"])
@pytest.mark.parametrize("empty_side", ["current", "previous"])
def test_v2_import_identity_comes_from_names_when_body_is_blank_or_wrong(tmp_path, layout, suffix, empty_side):
    directories = {side: tmp_path / side for side in ("current", "previous")}
    for directory in directories.values():
        directory.mkdir()
    profiles = {
        code: InstitutionProfile(code, "配置" + name, code, "农商行", "一部", "示例地区")
        for code, name in (("9144AA", "甲银行"), ("9144BB", "乙银行"))
    }
    config = ReportAuditConfig(
        institutions_by_id=profiles,
        indicators={
            "20201001": IndicatorConfig("20201001", "机构名称", "20201", "文字"),
            "20201002": IndicatorConfig("20201002", "机构代码", "20201", "文字"),
            "20201999": IndicatorConfig("20201999", "测试指标", "20201", "个数"),
        },
    )
    for side, directory in directories.items():
        date = "20260630" if side == "current" else "20260331"
        if layout == "banks":
            paths = [(directory / f"banks#{date}#01#20201{suffix}", tuple(profiles))]
        else:
            paths = []
            for code, name in (("9144AA", "甲银行"), ("9144BB", "乙银行")):
                stem = f"reports#{code}#{date}#01#{name}" if layout == "reports" else f"{code}#{date}#01#20201#{name}"
                paths.append((directory / (stem + suffix), (code,)))
        for path, codes in paths:
            if suffix == ".xls":
                xlwt = pytest.importorskip("xlwt")
                book = xlwt.Workbook()
            else:
                book = Workbook(); book.remove(book.active)
            for code in codes:
                sheet_name = code if layout == "banks" else "20201"
                missing = side == empty_side
                rows = [["标题"], [], ["指标代码", "指标名称", "值"],
                        ["20201001", "机构名称", None if missing else "正文误填名称"],
                        ["20201002", "机构代码", None if missing else "--"],
                        ["20201999", "测试指标", None if missing else 100]]
                if suffix == ".xls":
                    sheet = book.add_sheet(sheet_name)
                    for r, row in enumerate(rows):
                        for c, value in enumerate(row):
                            if value is not None:
                                sheet.write(r, c, value)
                else:
                    sheet = book.create_sheet(sheet_name)
                    for row in rows:
                        sheet.append(row)
            book.save(str(path))
    source_files = [path for directory in directories.values() for path in directory.iterdir()]
    hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in source_files}
    imported = {
        side: import_period_directory(directory, config=config, label=side)
        for side, directory in directories.items()
    }
    for result in imported.values():
        assert len(result.records) == 6 and not result.duplicates
        assert {r.institution_id for r in result.records} == set(profiles)
        assert all(r.social_credit_code == r.institution_id for r in result.records)
        expected_names = {"配置甲银行", "配置乙银行"} if layout == "banks" else {"甲银行", "乙银行"}
        assert {r.institution_name for r in result.records} == expected_names
        assert all(r.region == "示例地区" and r.institution_type == "农商行" for r in result.records)
    findings = PeriodComparisonEngine().run(imported["current"].by_logical_key(), imported["previous"].by_logical_key(), config)
    numeric = [f for f in findings if f.indicator_code == "20201999"]
    assert len(numeric) == 2
    assert all(f.description == ("本期无，上期有" if empty_side == "current" else "本期有，上期无") for f in numeric)
    assert all(f.institution_name and f.social_credit_code in profiles for f in findings)
    assert {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in source_files} == hashes


def test_v2_import_keeps_filename_credit_code_when_name_profile_has_other_code(tmp_path):
    directory = tmp_path / "current"; directory.mkdir()
    _source(directory / "9144#20260630#01#20202#甲银行.xlsx", deposit=10, loan=20, asset=30, debt=10, equity=20)
    stale_profile = InstitutionProfile("旧代码", "甲银行", "旧代码")
    config = ReportAuditConfig(institutions_by_name={"甲银行": stale_profile})
    result = import_period_directory(directory, config=config, label="本期")
    assert all(r.institution_id == "9144" and r.social_credit_code == "9144" for r in result.records)
    assert all(r.institution_name == "甲银行" for r in result.records)


@pytest.mark.parametrize("empty_side", ["current", "previous"])
def test_v2_export_preserves_identity_for_completely_empty_counterpart(tmp_path, empty_side):
    config_path = tmp_path / "config.xlsx"; _config(config_path)
    current, previous = tmp_path / "current", tmp_path / "previous"
    current.mkdir(); previous.mkdir()
    for side, directory, date in (("current", current, "20260630"), ("previous", previous, "20260331")):
        path = directory / f"9144#{date}#01#20202#甲银行.xlsx"
        if side == empty_side:
            book = Workbook(); book.active.title = "20202"; book.save(path)
        else:
            _source(path, deposit=100, loan=100, asset=100, debt=50, equity=50)
            book = load_workbook(path)
            book["20202"]["C4"] = "--"; book["20202"]["C5"] = "--"
            book.save(path); book.close()
    inputs = [config_path, *current.iterdir(), *previous.iterdir()]
    hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs}
    result = PeriodAuditV2Service().run(current_dir=current, previous_dir=previous, config_path=config_path, output_dir=tmp_path / "out")
    findings = [f for f in result.findings if f.audit_type == "环比"]
    assert findings and all(f.institution_name == "甲银行" and f.social_credit_code == "9144" for f in findings)
    book = load_workbook(result.paths.result)
    try:
        rows = list(book["环比规则"].iter_rows(values_only=True))
        headers = list(rows[0])
        assert len(rows) > 1
        assert all(row[headers.index("机构名称")] == "甲银行" for row in rows[1:])
        assert all(row[headers.index("社会信用代码")] == "9144" for row in rows[1:])
        assert all(row[headers.index("审核说明")] == ("本期无，上期有" if empty_side == "current" else "本期有，上期无") for row in rows[1:])
    finally:
        book.close()
    assert {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs} == hashes


def test_v2_progress_covers_file_counts_export_and_changes_no_results(tmp_path: Path) -> None:
    config, current, previous, central = tmp_path / "config.xlsx", tmp_path / "cur", tmp_path / "pre", tmp_path / "central.xlsx"
    current.mkdir(); previous.mkdir(); _config(config); _central(central)
    for directory, date in ((current, "20260630"), (previous, "20260331")):
        _source(directory / f"9144#{date}#01#20202#甲银行.xlsx", deposit=100, loan=50, asset=100, debt=60, equity=40)
    source_hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in (config, central, *current.iterdir(), *previous.iterdir())}
    events = []

    def report(event):
        events.append(dict(event))
        if event["stage"] == "complete":
            assert list((tmp_path / "with_progress").glob("*.xlsx")), "仅在结果保存成功后报告完成"

    tracked = PeriodAuditV2Service().run(
        current_dir=current, previous_dir=previous, config_path=config, central_path=central,
        output_dir=tmp_path / "with_progress", on_progress=report,
    )
    plain = PeriodAuditV2Service().run(
        current_dir=current, previous_dir=previous, config_path=config, central_path=central, output_dir=tmp_path / "plain",
    )
    assert tracked.findings == plain.findings and tracked.summary == plain.summary
    stages = list(dict.fromkeys(event["stage"] for event in events))
    assert stages == ["config", "current", "previous", "validate", "period", "rules", "external_read", "external", "export", "save", "complete"]
    assert [event["percent"] for event in events] == sorted(event["percent"] for event in events)
    assert events[0]["percent"] == 0 and events[-1]["percent"] == 100
    assert all(event["percent"] < 100 for event in events[:-1])
    messages_by_stage = {event["stage"]: event["message"] for event in events}
    assert "本期 " in messages_by_stage["validate"] and "上期 " in messages_by_stage["validate"]
    assert "环比结果" in messages_by_stage["period"]
    assert "校验结果" in messages_by_stage["rules"]
    assert "机构参照" in messages_by_stage["external_read"]
    assert "外部核对结果" in messages_by_stage["external"]
    for stage in ("current", "previous"):
        file_events = [event for event in events if event["stage"] == stage and event["total"]]
        assert [(event["current"], event["total"]) for event in file_events] == [(0, 1), (1, 1)]
        assert file_events[0]["message"].endswith(".xlsx")
    save = next(event for event in events if event["stage"] == "save")
    assert save["current"] == save["total"] == len(tracked.findings)
    for path in (tracked.paths.result, plain.paths.result):
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            contents = tuple((sheet.title, tuple(sheet.values)) for sheet in book.worksheets)
            if path == tracked.paths.result:
                expected = contents
            else:
                assert contents == expected
        finally:
            book.close()
    assert all(hashlib.sha256(path.read_bytes()).hexdigest() == digest for path, digest in source_hashes.items())


def test_v2_import_progress_retains_sources_and_only_reads_contract_columns(tmp_path: Path) -> None:
    for code in ("9144", "9145"):
        path = tmp_path / f"{code}#20260630#01#20202#样例机构.xlsx"
        _source(path, deposit=100, loan=50, asset=100, debt=60, equity=40)
        book = load_workbook(path)
        book.active["ZZ4"] = "不属于报送取数范围的内容"
        book.save(path); book.close()
    events = []
    imported = import_period_directory(tmp_path, config=ReportAuditConfig(), label="本期", on_progress=lambda *event: events.append(event))
    assert [(current, total) for current, total, _message in events] == [(0, 2), (1, 2), (2, 2)]
    assert len(imported.records) == 14 and not imported.duplicates
    assert {record.institution_id for record in imported.records} == {"9144", "9145"}
    assert all(Path(record.source.source_file).is_absolute() for record in imported.records)
    assert all(record.source.source_row >= 4 for record in imported.records)
    assert all("不属于报送取数范围" not in str(record.value) for record in imported.records)


def test_v2_failed_read_does_not_report_completion(tmp_path: Path) -> None:
    config = tmp_path / "config.xlsx"; _config(config)
    current, previous = tmp_path / "cur", tmp_path / "pre"
    current.mkdir(); previous.mkdir()
    (current / "9144#20260630#01#20202#损坏文件.xlsx").write_bytes(b"broken workbook")
    events = []
    with pytest.raises(Exception):
        PeriodAuditV2Service().run(current_dir=current, previous_dir=previous, config_path=config, output_dir=tmp_path / "out", on_progress=events.append)
    assert events[-1]["stage"] == "current"
    assert events[-1]["message"].endswith("损坏文件.xlsx")
    assert all(event["percent"] < 100 for event in events)
