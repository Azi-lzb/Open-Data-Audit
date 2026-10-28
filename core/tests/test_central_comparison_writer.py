from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from base_audit.native.s3_comparison_writer import writer_available
from base_audit.systems.s3_central_statistics.comparison_exporter import (
    write_comparison_workbook,
)
from base_audit.systems.s3_central_statistics.models import CentralRecord, ComparisonRow


def sample_rows() -> list[ComparisonRow]:
    first = CentralRecord(
        "本外币", "2026-10-31", "001", "测试机构", "440000", "广东",
        "001", "0000123", "测试指标", "余额", "人民币", "月", "1", 12.3456,
    )
    second = CentralRecord(
        "人民币", "2026-10-31", "002", "另一机构", "440100", "广州",
        "002", "0000456", "文本指标", "发生额", "人民币", "月", "2", "异常文本",
    )
    return [
        ComparisonRow(
            record=first,
            order_code_out="001-1",
            indicator_name_out="别名指标",
            prev_value=10.0,
            change=2.35,
            ratio=0.23456,
            remark="触发提示",
            need_explain="是",
            explain_content="人工说明",
            process="本期减上期",
            fill_color=6,
            fill_scope="row",
        ),
        ComparisonRow(
            record=second,
            prev_value="旧文本",
            change=None,
            ratio="除数为0",
            remark="备注颜色",
            fill_color=3,
            fill_scope="remark",
        ),
    ]


def workbook_snapshot(path: Path) -> dict:
    book = load_workbook(path, data_only=True)
    try:
        sheet = book.active
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
        styles_book = load_workbook(path, data_only=False)
        try:
            styled = styles_book.active
            def width_for_column(index: int) -> float:
                for dimension in styled.column_dimensions.values():
                    if dimension.min <= index <= dimension.max:
                        return round(dimension.width, 3)
                return round(styled.sheet_format.defaultColWidth or 13.0, 3)

            return {
                "sheetnames": book.sheetnames,
                "values": values,
                "freeze": styled.freeze_panes,
                "filter": styled.auto_filter.ref,
                "widths": tuple(width_for_column(i) for i in range(1, 23)),
                "header_fill": tuple(
                    cell.fill.fgColor.rgb[-6:] for cell in styled[1]
                ),
                "sample_fills": {
                    "A2": styled["A2"].fill.fgColor.rgb[-6:],
                    "R3": styled["R3"].fill.fgColor.rgb[-6:],
                },
                "formats": tuple(
                    styled.cell(2, column).number_format
                    for column in (2, 14, 15, 16, 17)
                ),
            }
        finally:
            styles_book.close()
    finally:
        book.close()


class S3ComparisonWriterTests(unittest.TestCase):
    def test_auto_without_sidecar_uses_openpyxl_and_reports_route(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "result.xlsx"
            messages: list[str] = []
            with patch(
                "base_audit.native.s3_comparison_writer.writer_binary", return_value=None
            ):
                write_comparison_workbook(
                    sample_rows(), target, writer="AUTO", on_engine=messages.append
                )
                self.assertTrue(target.is_file())
                self.assertIn("实际引擎=OpenPyXL", messages[0])
                self.assertIn("自动回退原因", messages[0])

                with self.assertRaisesRegex(RuntimeError, "当前程序未附带"):
                    write_comparison_workbook(
                        sample_rows(), target, writer="RUST_XLSXWRITER"
                    )

    @unittest.skipUnless(writer_available(), "本平台未构建可选 Rust S3 writer")
    def test_rust_writer_matches_openpyxl_values_and_report_contract(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            openpyxl_path = root / "openpyxl.xlsx"
            rust_path = root / "rust.xlsx"
            messages: list[str] = []
            rows = sample_rows()
            write_comparison_workbook(rows, openpyxl_path, writer="OPENPYXL")
            write_comparison_workbook(
                rows, rust_path, writer="RUST_XLSXWRITER", on_engine=messages.append
            )
            self.assertEqual(workbook_snapshot(rust_path), workbook_snapshot(openpyxl_path))
            self.assertIn("0.100.0", messages[0])


if __name__ == "__main__":
    unittest.main()
