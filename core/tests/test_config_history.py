from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from src.base_audit.service import (
    AuditService,
    CONFIG_HISTORY_SHEET,
    _output_prefix,
)
from src.base_audit.name_config import (
    FLOW_HEADERS,
    FLOW_LEGACY_HEADERS,
    FLOW_PREVIOUS_HEADERS,
    FLOW_PREVIOUS_9_HEADERS,
    FIXED_ROW_SUMMARY_FUNCTION,
    FORMULA_COPY_FUNCTION,
    COMBINE_SHEETS_FUNCTION,
    NAMED_RANGE_CHECK_FUNCTION,
    MERGE_ORG_FILES_FUNCTION,
    STRUCTURE_COMPARE_FUNCTION,
    USED_RANGE_SUMMARY_FUNCTION,
    WORKBOOK_TABLE_MERGE_FUNCTION,
    _initialize_legacy_config,
    initialize_config,
    load_feature_mappings,
    load_config_editor_data,
    reset_default_configuration,
    save_config_editor_draft,
    validate_config_editor_draft,
    validate_config,
    load_combine_sheets_plan,
)
from src.base_audit.excel_com import _a1_address
from src.base_audit.history import HISTORY_AUDIT_SHEET, LOCAL_VALIDATION_HISTORY_SHEET


class ConfigHistoryTests(unittest.TestCase):
    def test_config_workbook_is_the_history_destination(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            config = root / "config.xlsx"
            workbook = Workbook()
            workbook.active.title = "区域映射"
            workbook.save(config)
            destination, sheet, migration = AuditService(
                config_path=config
            )._history_storage(root / "历史审核库.xlsx")
        self.assertEqual(config.resolve(), destination)
        self.assertEqual(CONFIG_HISTORY_SHEET, sheet)
        self.assertEqual((root / "历史审核库.xlsx").resolve(), migration)

    def test_initialize_config_restores_default_region_mappings(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            mappings = load_feature_mappings(config, Path("任意模板.xlsx"))
        names = {mapping.feature_type: mapping.range_names for mapping in mappings}
        named_checks = {mapping.name: mapping.range_names for mapping in mappings if mapping.feature_type == NAMED_RANGE_CHECK_FUNCTION}
        self.assertEqual(("校验区域",), names[FORMULA_COPY_FUNCTION])
        self.assertEqual(("校验区域",), named_checks["检查校验区域"])
        self.assertEqual(("表结构区域",), named_checks["检查表结构区域"])
        self.assertEqual(("任意行汇总区域",), named_checks["检查汇总区域"])
        self.assertEqual(("表头区域",), named_checks["检查表头区域"])
        self.assertEqual(("表结构区域",), names[STRUCTURE_COMPARE_FUNCTION])
        self.assertEqual(("任意行汇总区域",), names[USED_RANGE_SUMMARY_FUNCTION])
        self.assertEqual(("固定行汇总区域",), names[FIXED_ROW_SUMMARY_FUNCTION])
        self.assertEqual((), names[WORKBOOK_TABLE_MERGE_FUNCTION])
        self.assertEqual((), names[MERGE_ORG_FILES_FUNCTION])
        self.assertNotIn("区域汇总", names)

    def test_config_editor_exposes_pristine_defaults_for_page_reset(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            editor = load_config_editor_data(config)
        self.assertTrue(editor["initialModules"])
        self.assertTrue(editor["initialFlows"])
        self.assertEqual(
            "校验区域",
            next(row for row in editor["initialModules"] if row["功能名"] == "公式校验复制")["命名区域名"],
        )
        self.assertEqual(
            "汇总核查表校验",
            editor["initialFlows"][0]["流程名"],
        )

    def legacy_test_initialize_migrates_excel_settings_to_json_and_keeps_history_only(self):
        with TemporaryDirectory() as folder:
            config = Path(folder) / "config.xlsx"
            # 使用旧初始化器建立一次 Excel 格式配置，模拟升级前用户环境。
            _initialize_legacy_config(config)
            workbook = load_workbook(config)
            workbook["执行流程"].append(
                ("我的流程", 10, "任意行汇总", "是", "停止", "是", "我的结果", "", "")
            )
            workbook["自定义按钮"].append(("我的按钮", "我的流程", "是", ""))
            workbook["历史审核结果"].append(("文件.xlsx", "Sheet1", "A1", "错误", "指标", "说明", "1", "", "", "规则", "历史", "意见"))
            workbook.save(config)
            workbook.close()

            initialize_config(config)
            editor = load_config_editor_data(config)
            check = load_workbook(config, read_only=True, data_only=True)
            try:
                self.assertEqual([HISTORY_AUDIT_SHEET, LOCAL_VALIDATION_HISTORY_SHEET, "历史追加配置"], check.sheetnames)
                self.assertEqual("文件.xlsx", check[HISTORY_AUDIT_SHEET].cell(2, 1).value)
            finally:
                check.close()
            # 旧“执行流程”行随旧模型退役：不再迁移进任何配置，历史行原样保留。
            self.assertNotIn("执行流程", check.sheetnames)

    def legacy_test_reset_defaults_keeps_history_and_custom_buttons(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            from openpyxl import load_workbook
            workbook = load_workbook(config)
            workbook["模块化功能"].append(("临时模块", FORMULA_COPY_FUNCTION, "临时区域", "否", "", "待删除"))
            workbook["执行流程"].append(("临时流程", 10, "临时模块", "是", "停止", "是", "", "待删除"))
            workbook["自定义按钮"].append(("我的按钮", "审核前检查", "是", "保留"))
            workbook["历史审核结果"].append(("文件.xlsx", "Sheet1", "A1", "错误", "指标", "说明", "1", "", "", "规则｜1", "历史", "意见"))
            workbook.save(config)
            workbook.close()
            reset_default_configuration(config)
            check = load_workbook(config, data_only=True)
            try:
                self.assertNotIn("临时模块", [row[0] for row in check["模块化功能"].iter_rows(min_row=2, values_only=True)])
                self.assertNotIn("临时流程", [row[0] for row in check["执行流程"].iter_rows(min_row=2, values_only=True)])
                self.assertEqual("我的按钮", check["自定义按钮"].cell(2, 1).value)
                self.assertEqual("文件.xlsx", check["历史审核结果"].cell(2, 1).value)
            finally:
                check.close()

    def legacy_test_reset_defaults_recovers_from_broken_module_sheet_headers(self):
        with TemporaryDirectory() as folder:
            config = Path(folder) / "config.xlsx"
            workbook = Workbook()
            workbook.active.title = "模块化功能"
            workbook.active.append(("被修改的表头",))
            workbook.save(config)
            workbook.close()
            reset_default_configuration(config)
            check = load_workbook(config, data_only=True)
            try:
                self.assertEqual(
                    ("功能名", "执行模块", "命名区域名", "是否限定工作簿", "工作簿关键字", "备注"),
                    tuple(cell.value for cell in check["模块化功能"][1]),
                )
                self.assertIn("执行流程", check.sheetnames)
            finally:
                check.close()

    def legacy_test_default_audit_flow_keeps_audit_copies_and_final_result(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            steps = load_flow_steps(config, "汇总核查表校验")
        self.assertEqual(
            ("检查校验区域", "检查表结构区域", "表结构比对", "外部文件添加", "公式校验复制", "校验结果提取", "条件格式结果提取", "审核结果输出"),
            tuple(step.feature_name for step in steps),
        )
        self.assertEqual(
            (False, False, False, False, True, False, False, True),
            tuple(step.output_result for step in steps),
        )
        self.assertTrue(all(step.on_failure == "跳过" for step in steps))
        self.assertEqual("机构审核副本", steps[4].output_name)

    def test_default_modules_describe_typical_input_and_output(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            editor = load_config_editor_data(config)
        modules = {row["功能名"]: row for row in editor["initialModules"]}
        self.assertEqual("源数据目录的条件格式区域", modules["条件格式结果提取"]["输入"])
        self.assertEqual("审核结果集", modules["条件格式结果提取"]["输出"])
        self.assertEqual("模板文件、源文件（或已添加外部表的审核副本）", modules["公式校验复制"]["输入"])
        self.assertEqual("已复制公式的审核副本", modules["公式校验复制"]["输出"])
        self.assertEqual("", modules["外部文件添加"]["命名区域名"])

    def legacy_test_default_merge_org_flow_is_a_single_output_step(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            steps = load_flow_steps(config, "组合联合核查表")
        self.assertEqual(("组合联合核查表",), tuple(step.feature_name for step in steps))
        self.assertTrue(steps[0].output_result)
        self.assertEqual("组合联合核查表", steps[0].output_name)

    def test_default_combine_sheets_flow_has_an_independent_plan(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            plan = load_combine_sheets_plan(config)
            editor = load_config_editor_data(config)
        self.assertEqual("regex", plan["mode"])
        self.assertIn("(?P<组合>", str(plan["pattern"]))
        self.assertEqual(
            {"regex", "name"},
            {str(item["mode"]) for item in editor["combineSheetsPlans"]},
        )
        self.assertNotIn("组合工作表", editor["defaultFlowNames"])
        self.assertIn("组合工作表", [row["流程名"] for row in editor["customFlows"]])

    def legacy_test_custom_flow_with_combine_module_is_detected_as_template_free(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            draft = {
                "customModules": [],
                "customFlows": [{
                    "流程名": "我的组合流程", "顺序": "10", "功能名": "组合工作表",
                    "启用": "是", "失败后处理": "停止", "是否输出结果": "是",
                    "输出文件名": "", "处理对象": "", "备注": "",
                }],
                "buttons": [],
            }
            save_config_editor_draft(config, draft)
            service = AuditService(config_path=config)
            self.assertTrue(service.is_standalone_combine_flow("我的组合流程"))
            self.assertFalse(service.is_standalone_combine_flow("汇总核查表校验"))

    def test_config_editor_validates_combine_sheets_plan(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            draft = {
                "customModules": [], "customFlows": [], "buttons": [],
                "combineSheetsPlan": {"name": "测试", "mode": "regex", "pattern": "没有分组"},
            }
            errors = validate_config_editor_draft(config, draft)
        self.assertTrue(any("正则必须" in error for error in errors))

    def test_combine_sheets_can_save_and_switch_grouping_profiles(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            draft = {
                "customModules": [], "customFlows": [], "buttons": [],
                "combineSheetsPlans": [
                    {"id": "by-org", "name": "按机构", "mode": "regex",
                     "pattern": r"(?P<组合>.+?)_金融基础数据-", "groups": []},
                    {"id": "name-special", "name": "专项组合", "mode": "name",
                     "pattern": "", "groups": [{"name": "贷款联合", "keywords": "个人贷款、单位贷款"}]},
                ],
                "activeCombineSheetsPlanId": "name-special",
            }
            self.assertEqual([], validate_config_editor_draft(config, draft))
            save_config_editor_draft(config, draft)
            active = load_combine_sheets_plan(config)
            editor = load_config_editor_data(config)
        self.assertEqual("专项组合", active["name"])
        self.assertEqual("name", active["mode"])
        self.assertEqual("贷款联合", active["groups"][0]["name"])
        self.assertEqual("name-special", editor["activeCombineSheetsPlanId"])
        self.assertEqual(2, len(editor["combineSheetsPlans"]))

    def legacy_test_config_editor_keeps_default_flows_and_history(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            workbook = load_workbook(config)
            workbook[HISTORY_AUDIT_SHEET].append(("文件.xlsx", "Sheet1", "A1", "错误", "指标", "说明", "1", "", "", "规则", "历史", "意见"))
            workbook.save(config)
            workbook.close()
            editor = load_config_editor_data(config)
            self.assertIn("汇总核查表校验", editor["defaultFlowNames"])
            self.assertTrue(editor["defaultFlows"])
            self.assertEqual(
                ["组合工作表"],
                [row["流程名"] for row in editor["customFlows"]],
            )
            custom = {
                "customModules": [{
                    "功能名": "我的任意汇总模块", "执行模块": USED_RANGE_SUMMARY_FUNCTION,
                    "命名区域名": "我的汇总区域", "是否限定工作簿": "否",
                    "工作簿关键字": "", "备注": "测试模块",
                }],
                "customFlows": [{
                    "流程名": "我的汇总流程", "顺序": "10", "功能名": "我的任意汇总模块",
                    "启用": "是", "失败后处理": "停止", "是否输出结果": "是",
                    "输出文件名": "我的汇总", "处理对象": "", "备注": "测试",
                }],
                "flowDisplay": {"我的汇总流程": True},
            }
            self.assertEqual([], validate_config_editor_draft(config, custom))
            save_config_editor_draft(config, custom)
            saved = load_config_editor_data(config)
            self.assertEqual("我的汇总流程", saved["customFlows"][0]["流程名"])
            self.assertIn("我的任意汇总模块", [item["功能名"] for item in saved["modules"]])
            check = load_workbook(config, data_only=True)
            try:
                self.assertEqual("文件.xlsx", check[HISTORY_AUDIT_SHEET].cell(2, 1).value)
                self.assertEqual([HISTORY_AUDIT_SHEET, LOCAL_VALIDATION_HISTORY_SHEET, "历史追加配置"], check.sheetnames)
            finally:
                check.close()

    def test_config_editor_allows_default_flow_edits(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            editor = load_config_editor_data(config)
            errors = validate_config_editor_draft(config, {
                "customModules": [],
                "customFlows": [{
                    "流程名": "汇总核查表校验", "顺序": "10", "功能名": "校验结果提取",
                    "启用": "是", "失败后处理": "停止", "是否输出结果": "是",
                    "输出文件名": "", "处理对象": "", "备注": "",
                }],
                "buttons": [],
            })
        self.assertEqual([], errors)

    def legacy_test_config_editor_saves_default_flow_and_module_overrides(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            editor = load_config_editor_data(config)
            modules = [
                {header: row.get(header, "") for header in editor["mappingHeaders"]}
                for row in editor["modules"]
            ]
            flows = [
                *[dict(row) for row in editor["defaultFlows"]],
                *[dict(row) for row in editor["customFlows"]],
            ]
            next(row for row in modules if row["功能名"] == "公式校验复制")["命名区域名"] = "我的校验区域"
            next(row for row in flows if row["功能名"] == "公式校验复制")["是否输出结果"] = "否"
            draft = {
                "customModules": modules,
                "customFlows": flows,
                "flowDisplay": editor["flowDisplay"],
                "combineSheetsPlans": editor["combineSheetsPlans"],
                "activeCombineSheetsPlanId": editor["activeCombineSheetsPlanId"],
            }
            self.assertEqual([], validate_config_editor_draft(config, draft))
            save_config_editor_draft(config, draft)
            self.assertEqual(
                "我的校验区域",
                next(item for item in load_feature_mappings(config, Path("模板.xlsx")) if item.name == "公式校验复制").range_names[0],
            )
            saved = load_config_editor_data(config)
            formula_row = next(
                row for row in [*saved["defaultFlows"], *saved["customFlows"]]
                if row["流程名"] == "汇总核查表校验" and row["功能名"] == "公式校验复制"
            )
            self.assertEqual("否", formula_row["是否输出结果"])

    def test_config_editor_allows_check_or_verify_steps_to_output_logs(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            editor = load_config_editor_data(config)
            flows = [
                *[dict(row) for row in editor["defaultFlows"]],
                *[dict(row) for row in editor["customFlows"]],
            ]
            next(row for row in flows if row["功能名"] == "检查校验区域")["是否输出结果"] = "是"
            errors = validate_config_editor_draft(config, {
                "customModules": [dict(row) for row in editor["modules"]],
                "customFlows": flows,
                "flowDisplay": editor["flowDisplay"],
                "combineSheetsPlans": editor["combineSheetsPlans"],
                "activeCombineSheetsPlanId": editor["activeCombineSheetsPlanId"],
            })
            self.assertEqual([], errors)
            save_config_editor_draft(config, {
                "customModules": [dict(row) for row in editor["modules"]],
                "customFlows": flows,
                "flowDisplay": editor["flowDisplay"],
                "combineSheetsPlans": editor["combineSheetsPlans"],
                "activeCombineSheetsPlanId": editor["activeCombineSheetsPlanId"],
            })
            saved = load_config_editor_data(config)
            check_row = next(
                row for row in [*saved["defaultFlows"], *saved["customFlows"]]
                if row["流程名"] == "汇总核查表校验" and row["功能名"] == "检查校验区域"
            )
            self.assertEqual("是", check_row["是否输出结果"])

    def legacy_test_default_explanation_flow_checks_structure_before_summary(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            steps = load_flow_steps(config, "汇总校验结果说明")
        self.assertEqual(
            ("检查表结构区域", "检查汇总区域", "表结构比对", "任意行汇总"),
            tuple(step.feature_name for step in steps),
        )

    def legacy_test_config_requires_an_output_step_per_flow(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            workbook = load_workbook(config)
            sheet = workbook["执行流程"]
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 1).value == "汇总核查表校验":
                    sheet.cell(row, 6, "否")
            workbook.save(config)
            workbook.close()
            errors = validate_config(config)
        self.assertIn("流程“汇总核查表校验”没有启用且“是否输出结果=是”的模块", errors)

    def legacy_test_default_flow_prefills_flow_name_as_output_name(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            audit_steps = load_flow_steps(config, "汇总核查表校验")
            summary_steps = load_flow_steps(config, "汇总校验结果说明")
        self.assertEqual("", audit_steps[0].output_name)
        self.assertEqual("汇总核查表校验", audit_steps[-1].output_name)
        self.assertEqual("汇总校验结果说明", summary_steps[-1].output_name)

    def legacy_test_initialize_config_migrates_legacy_seven_column_flow_sheet(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            workbook = load_workbook(config)
            # 构造 7 列旧配置：移除“输出文件名”“处理对象”两列。
            workbook["执行流程"].delete_cols(len(FLOW_LEGACY_HEADERS), 2)
            workbook["执行流程"].append(("我的自定义流程", 10, "检查校验区域", "是", "停止", "是", "自定义备注"))
            workbook.save(config)
            workbook.close()
            initialize_config(config)
            migrated = load_workbook(config, data_only=True)
            try:
                sheet = migrated["执行流程"]
                self.assertEqual(FLOW_HEADERS, tuple(cell.value for cell in sheet[1]))
                rows = {(row[0], row[2]): row for row in sheet.iter_rows(min_row=2, values_only=True)}
                self.assertIn(("我的自定义流程", "检查校验区域"), rows)
                self.assertEqual("汇总核查表校验", rows[("汇总核查表校验", "校验结果提取")][6])
                self.assertEqual("公式校验复制", rows[("汇总核查表校验", "校验结果提取")][7])
                self.assertEqual("输出正式本期审核结果", rows[("汇总核查表校验", "校验结果提取")][8])
                self.assertEqual("汇总校验结果说明", rows[("汇总校验结果说明", "任意行汇总")][6])
                self.assertEqual([], validate_config(config))
            finally:
                migrated.close()

    def legacy_test_initialize_config_migrates_previous_eight_column_flow_sheet(self):
        # 上一版 8 列配置（有“输出文件名”列、无“处理对象”列）也应增量迁移为 9 列。
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            workbook = load_workbook(config)
            workbook["执行流程"].delete_cols(len(FLOW_PREVIOUS_HEADERS))  # 移除“处理对象”列
            workbook.save(config)
            workbook.close()
            initialize_config(config)
            migrated = load_workbook(config, data_only=True)
            try:
                sheet = migrated["执行流程"]
                self.assertEqual(FLOW_HEADERS, tuple(cell.value for cell in sheet[1]))
                rows = {(row[0], row[2]): row for row in sheet.iter_rows(min_row=2, values_only=True)}
                self.assertEqual("公式校验复制", rows[("汇总核查表校验", "校验结果提取")][7])
                self.assertEqual("汇总校验结果说明", rows[("汇总校验结果说明", "任意行汇总")][6])
                self.assertEqual([], validate_config(config))
            finally:
                migrated.close()

    def legacy_test_initialize_config_renames_previous_nine_column_process_source_header(self):
        # 上一版 9 列配置（“处理对象”还叫“待处理对象”）应就地改名，
        # 并回填公式校验复制的默认处理对象“外部文件添加”。
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            workbook = load_workbook(config)
            workbook["执行流程"].cell(1, 8, "待处理对象")
            workbook["执行流程"].append(("我的自定义流程", 10, "检查校验区域", "是", "停止", "是", "", "", "自定义备注"))
            workbook.save(config)
            workbook.close()
            initialize_config(config)
            migrated = load_workbook(config, data_only=True)
            try:
                sheet = migrated["执行流程"]
                self.assertEqual(FLOW_HEADERS, tuple(cell.value for cell in sheet[1]))
                rows = {(row[0], row[2]): row for row in sheet.iter_rows(min_row=2, values_only=True)}
                self.assertIn(("我的自定义流程", "检查校验区域"), rows)
                self.assertEqual("外部文件添加", rows[("汇总核查表校验", "公式校验复制")][7])
                self.assertEqual("公式校验复制", rows[("汇总核查表校验", "校验结果提取")][7])
                self.assertEqual([], validate_config(config))
            finally:
                migrated.close()

    def legacy_test_legacy_seven_column_flow_sheet_loads_with_empty_output_name(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            workbook = load_workbook(config)
            workbook["执行流程"].delete_cols(len(FLOW_LEGACY_HEADERS))
            workbook.save(config)
            workbook.close()
            steps = load_flow_steps(config, "汇总核查表校验")
        self.assertTrue(steps)
        self.assertEqual("", steps[-1].output_name)

    def legacy_test_default_flow_prefills_process_source(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            audit_steps = load_flow_steps(config, "汇总核查表校验")
            summary_steps = load_flow_steps(config, "汇总校验结果说明")
        self.assertEqual("源数据目录", audit_steps[0].process_source)
        self.assertEqual("外部文件添加", audit_steps[4].process_source)
        self.assertEqual("公式校验复制", audit_steps[5].process_source)
        self.assertEqual("本期审核结果", audit_steps[5].result_set)
        self.assertEqual("本期审核结果", audit_steps[-1].result_set)
        self.assertEqual("源数据目录", summary_steps[-1].process_source)

    def legacy_test_validate_config_rejects_missing_process_source(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            workbook = load_workbook(config)
            sheet = workbook["执行流程"]
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 1).value == "汇总校验结果说明" and sheet.cell(row, 3).value == "任意行汇总":
                    sheet.cell(row, 8, "校验结果提取")
            workbook.save(config)
            workbook.close()
            errors = validate_config(config)
        self.assertTrue(any("处理对象" in error and "不存在" in error for error in errors))

    def test_output_prefix_uses_output_name_when_filled(self):
        self.assertEqual("汇总核查表校验", _output_prefix(60, "校验结果提取", "汇总核查表校验"))
        self.assertEqual("60_校验结果提取", _output_prefix(60, "校验结果提取", ""))
        self.assertEqual("40_任意行汇总", _output_prefix(40, "任意行汇总"))

    def legacy_test_validate_config_rejects_illegal_output_name(self):
        with TemporaryDirectory() as folder:
            config = initialize_config(Path(folder) / "config.xlsx")
            workbook = load_workbook(config)
            sheet = workbook["执行流程"]
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 1).value == "汇总核查表校验" and sheet.cell(row, 3).value == "校验结果提取":
                    sheet.cell(row, 7, "非法/名")
            workbook.save(config)
            workbook.close()
            errors = validate_config(config)
        self.assertTrue(any("包含非法字符" in error for error in errors))

    def test_structure_check_address_is_computed_without_excel(self):
        self.assertEqual("B3", _a1_address("B3:D5", 0, 0))
        self.assertEqual("D5", _a1_address("B3:D5", 2, 2))
        self.assertEqual("AA10", _a1_address("Z8:AA10", 2, 1))
