from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from src.base_audit.models import CopyRange, TemplateDefinition
from src.base_audit.preflight_xlsx import (
    summary_source_is_usable,
    validate_source_xlsx,
    write_structure_report_xlsx,
)


class PreflightXlsxTests(unittest.TestCase):
    def _validate(self, values, ranges):
        folder = TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        source = Path(folder.name) / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "贷款明细"
        for address, value in values.items():
            sheet[address] = value
        workbook.save(source)
        workbook.close()
        definition = TemplateDefinition(
            [], [CopyRange("贷款明细", "A1")], False,
            [item for item, _ in ranges],
        )
        return validate_source_xlsx(source, definition, ranges)

    def test_blank_template_cells_are_skipped_but_numbers_and_codes_are_checked(self):
        area = CopyRange("贷款明细", "A1:D1")
        result = self._validate(
            {"A1": 1001, "B1": "001", "C1": "任意值", "D1": ""},
            [(area, [[1001.0, "001", "   ", None]])],
        )
        self.assertTrue(result.matched)
        self.assertEqual((2, 2), (result.matched_labels, result.checked_labels))

    def test_numeric_and_text_values_with_same_rendering_do_not_match(self):
        area = CopyRange("贷款明细", "A1")
        result = self._validate({"A1": "1001"}, [(area, [[1001]])])
        self.assertFalse(result.matched)
        self.assertEqual(0, result.matched_labels)

    def test_multiple_and_single_cell_areas_are_compared(self):
        first = CopyRange("贷款明细", "A1:B1")
        second = CopyRange("贷款明细", "D5")
        result = self._validate(
            {"A1": "甲", "B1": "乙", "D5": 9},
            [(first, [["甲", "乙"]]), (second, [[9]])],
        )
        self.assertTrue(result.matched)
        self.assertEqual(3, len(result.structure_checks))

    def test_merged_cell_only_compares_nonblank_template_anchor(self):
        folder = TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        source = Path(folder.name) / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "贷款明细"
        sheet.merge_cells("A1:B1")
        sheet["A1"] = "合并表头"
        workbook.save(source)
        workbook.close()
        area = CopyRange("贷款明细", "A1:B1")
        definition = TemplateDefinition([], [CopyRange("贷款明细", "A1")], False, [area])
        result = validate_source_xlsx(source, definition, [(area, [["合并表头", None]])])
        self.assertTrue(result.matched)
        self.assertEqual(1, result.checked_labels)

    def test_merged_cell_subordinate_uses_merged_anchor_value(self):
        folder = TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        source = Path(folder.name) / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "贷款明细"
        sheet.merge_cells("A1:B1")
        sheet["A1"] = "合并表头"
        workbook.save(source)
        workbook.close()
        area = CopyRange("贷款明细", "B1")
        definition = TemplateDefinition([], [CopyRange("贷款明细", "A1")], False, [area])
        result = validate_source_xlsx(source, definition, [(area, [["合并表头"]])])
        self.assertTrue(result.matched)
        self.assertEqual(1, result.matched_labels)

    def test_structure_details_keep_sheet_name_and_cell_address(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "机构甲.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "贷款明细"
            sheet["A1"], sheet["B1"] = "固定标题", "错误字段"
            workbook.save(source)
            workbook.close()

            structure = CopyRange("贷款明细", "A1:B1")
            definition = TemplateDefinition([], [CopyRange("贷款明细", "A1")], False, [structure])
            result = validate_source_xlsx(source, definition, [(structure, [["固定标题", "正确字段"]])])
            self.assertEqual(2, len(result.structure_checks))
            self.assertEqual("贷款明细", result.structure_checks[1].sheet_name)
            self.assertEqual("B1", result.structure_checks[1].cell_address)
            self.assertFalse(result.structure_checks[1].matched)

            report = root / "表结构比对.xlsx"
            write_structure_report_xlsx(report, [(source, result)])
            saved = load_workbook(report, read_only=True, data_only=True)
            detail = saved["结构比对明细"]
            self.assertEqual("工作表", detail["B1"].value)
            self.assertEqual("贷款明细", detail["B3"].value)
            self.assertEqual("B1", detail["C3"].value)
            saved.close()

    def test_summary_keeps_present_sheets_when_another_template_sheet_is_missing(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "机构甲.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "有数据的表"
            sheet.append(("编号", "名称", "金额"))
            workbook.save(source)
            workbook.close()

            present = CopyRange("有数据的表", "A1:C1")
            absent = CopyRange("缺少的表", "A1:C1")
            definition = TemplateDefinition(
                [],
                [CopyRange("有数据的表", "A2:C10"), CopyRange("缺少的表", "A2:C10")],
                False,
                [present, absent],
            )
            expected = [
                (present, [["编号", "名称", "金额"]]),
                (absent, [["编号", "名称", "金额"]]),
            ]

            result = validate_source_xlsx(
                source, definition, expected, allow_missing_sheets=True
            )

            self.assertTrue(result.matched)
            self.assertEqual(("缺少的表",), result.missing_sheets)
            self.assertEqual(3, result.checked_labels)
            self.assertTrue(summary_source_is_usable(result, allow_missing_sheets=True))
            self.assertIn("其余工作表继续汇总", result.details)

    def test_summary_accepts_a_fallback_only_template(self):
        with TemporaryDirectory() as folder:
            source = Path(folder) / "机构甲.xlsx"
            workbook = Workbook()
            workbook.active.title = "新增业务表"
            workbook.save(source)
            workbook.close()

            definition = TemplateDefinition([], [], False, [])
            result = validate_source_xlsx(
                source, definition, [], allow_fallback_only=True
            )

            self.assertTrue(result.matched)
            self.assertTrue(summary_source_is_usable(result))
            self.assertIn("模板仅配置“兜底模板”", result.details)

    def test_summary_keeps_header_only_result_when_all_target_sheets_are_missing(self):
        with TemporaryDirectory() as folder:
            source = Path(folder) / "机构甲.xlsx"
            workbook = Workbook()
            workbook.active.title = "无关工作表"
            workbook.save(source)
            workbook.close()

            first = CopyRange("表甲", "A1:B1")
            second = CopyRange("表乙", "A1:B1")
            definition = TemplateDefinition(
                [], [CopyRange("表甲", "A2:B10"), CopyRange("表乙", "A2:B10")],
                False, [first, second],
            )
            result = validate_source_xlsx(
                source,
                definition,
                [(first, [["字段", "数值"]]), (second, [["字段", "数值"]])],
                allow_missing_sheets=True,
            )

            self.assertFalse(result.matched)
            self.assertEqual(("表乙", "表甲"), result.missing_sheets)
            self.assertEqual(0, result.checked_labels)
            self.assertTrue(summary_source_is_usable(result, allow_missing_sheets=True))

    def test_summary_still_rejects_misaligned_existing_sheet_if_another_is_missing(self):
        with TemporaryDirectory() as folder:
            source = Path(folder) / "机构甲.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "有数据的表"
            sheet.append(("编号", "错误字段", "金额"))
            workbook.save(source)
            workbook.close()

            present = CopyRange("有数据的表", "A1:C1")
            absent = CopyRange("缺少的表", "A1:C1")
            definition = TemplateDefinition(
                [],
                [CopyRange("有数据的表", "A2:C10"), CopyRange("缺少的表", "A2:C10")],
                False,
                [present, absent],
            )
            result = validate_source_xlsx(
                source,
                definition,
                [
                    (present, [["编号", "名称", "金额"]]),
                    (absent, [["编号", "名称", "金额"]]),
                ],
                allow_missing_sheets=True,
            )

            self.assertFalse(result.matched)
            self.assertEqual(("缺少的表",), result.missing_sheets)
            self.assertGreater(result.checked_labels, 0)
            self.assertFalse(summary_source_is_usable(result, allow_missing_sheets=True))


if __name__ == "__main__":
    unittest.main()
