from __future__ import annotations

import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from base_audit.config_guide import CONFIG_VERSION, GUIDE_SHEET, guide_rows, read_config_version, write_guide_to_file
from base_audit.name_config import (
    check_summary_config,
    ensure_summary_config_guide,
    initialize_history_workbook,
)
from base_audit.node_flow_config import (
    NODE_HEADERS,
    check_dag_config,
    needs_dag_config_simplification,
    node_flow_config_path,
    reset_dag_config,
    simplify_dag_config,
)
from base_audit.period_compare import check_period_config, ensure_period_config_guide
from base_audit.systems.s2_report_collection.config import load_config, write_default_config

from openpyxl import load_workbook


class ConfigGuideTests(unittest.TestCase):
    def test_central_split_guides_explain_common_config_dependencies(self) -> None:
        common = dict(guide_rows("central_common"))
        comparison = dict(guide_rows("central_comparison"))
        cross = dict(guide_rows("central_cross"))
        forms = dict(guide_rows("central_forms"))

        self.assertIn("默认源数据单位", common["运行参数"])
        self.assertIn("99 表示 99%", common["环比警戒"])
        self.assertIn("程序自动生成区间文字", common["环比警戒"])
        self.assertIn("3.1、3.2", common["单位换算例外"])
        self.assertIn("排序", comparison["指标参照"])
        self.assertIn("单位换算例外", comparison["同时读取 3.0"])
        self.assertIn("单位换算例外", cross["同时读取 3.0"])
        self.assertIn("环比警戒", forms["同时读取 3.0"])
        self.assertIn("唯一维护", forms["报表清单"])

    def test_reset_config_has_guide_version_and_order_column(self) -> None:
        with TemporaryDirectory() as folder:
            history = Path(folder) / "config" / "逐笔统计系统_配置.xlsx"
            config = reset_dag_config(history)
            self.assertEqual(read_config_version(config), CONFIG_VERSION)
            book = load_workbook(config, read_only=True)
            try:
                self.assertEqual(tuple(c.value for c in book["节点配置"][1]), NODE_HEADERS)
                flows = [row[0] for row in book["节点配置"].iter_rows(min_row=2, values_only=True)]
                self.assertIn("汇总校验结果说明", flows)
            finally:
                book.close()
            self.assertEqual(check_dag_config(history), "")

    def test_summary_workflow_nodes_are_granular(self) -> None:
        with TemporaryDirectory() as folder:
            history = Path(folder) / "config" / "逐笔统计系统_配置.xlsx"
            config = reset_dag_config(history)
            book = load_workbook(config, read_only=True)
            try:
                names = [row[2] for row in book["节点配置"].iter_rows(min_row=2, values_only=True)
                         if row[0] == "汇总校验结果说明"]
            finally:
                book.close()
            # 汇总流程必须细粒度到独立节点，而不是两个黑盒。
            self.assertEqual(names, [
                "报送文件来源", "模板装载与体检", "检查表结构区域", "检查汇总区域",
                "检查表头区域", "表结构比对", "任意行汇总", "历史说明富化",
            ])

    def test_simplify_is_idempotent_and_keeps_guide(self) -> None:
        with TemporaryDirectory() as folder:
            history = Path(folder) / "config" / "逐笔统计系统_配置.xlsx"
            config = reset_dag_config(history)
            self.assertFalse(needs_dag_config_simplification(history))
            simplify_dag_config(history)  # 幂等：不报错
            self.assertEqual(read_config_version(config), CONFIG_VERSION)

    def test_summary_config_guide_and_check(self) -> None:
        with TemporaryDirectory() as folder:
            history = Path(folder) / "config" / "逐笔统计系统_配置.xlsx"
            initialize_history_workbook(history)  # 建簿时即补“使用说明”
            self.assertTrue(history.is_file())
            self.assertEqual(read_config_version(history), CONFIG_VERSION)
            self.assertEqual(check_summary_config(history), "")
            self.assertFalse(ensure_summary_config_guide(history))  # 已是最新，不再写

    def test_summary_config_missing_reports_issue(self) -> None:
        with TemporaryDirectory() as folder:
            missing = Path(folder) / "config" / "不存在.xlsx"
            self.assertIn("缺失", check_summary_config(missing))

    def test_period_default_config_loads_and_has_guide(self) -> None:
        with TemporaryDirectory() as folder:
            path = Path(folder) / "config" / "报表采集系统_配置.xlsx"
            write_default_config(path)
            self.assertEqual(read_config_version(path), CONFIG_VERSION)
            self.assertEqual(check_period_config(path.parent / "x.xlsx", path), "")
            config = load_config(path)
            self.assertEqual(config.warnings, [])
            self.assertTrue(config.indicators)
            self.assertTrue(config.bands)
            # 默认示例行都是禁用/提示性质，不应引入真实机构数据
            self.assertIn("20201001", config.indicators)

    def test_period_config_missing_table_reports_issue(self) -> None:
        with TemporaryDirectory() as folder:
            path = Path(folder) / "报表采集系统_配置.xlsx"
            # 造一个只有“指标参照”的坏配置
            from openpyxl import Workbook
            wb = Workbook()
            wb.active.title = "指标参照"
            wb.save(path)
            wb.close()
            self.assertIn("外部核对规则", check_period_config(Path(folder) / "x.xlsx", path))

    def test_write_guide_is_idempotent_by_version(self) -> None:
        with TemporaryDirectory() as folder:
            history = Path(folder) / "config" / "逐笔统计系统_配置.xlsx"
            config = reset_dag_config(history)
            self.assertFalse(write_guide_to_file(config, "config"))
            book = load_workbook(config, read_only=True)
            try:
                self.assertIn(GUIDE_SHEET, book.sheetnames)
            finally:
                book.close()


if __name__ == "__main__":
    unittest.main()
