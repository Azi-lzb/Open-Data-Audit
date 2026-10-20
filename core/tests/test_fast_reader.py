"""Calamine 快速读取引擎的值域对拍与回退保护。

快速引擎必须与 xlrd/openpyxl 传统引擎逐值一致，导入结果才不依赖引擎
选择；环境缺少 python-calamine 时对拍用例自动跳过（回退路径不依赖安装
状态，始终执行）。
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from base_audit.systems.s2_report_collection import importer
from base_audit.systems.s2_report_collection.importer import import_period_directory
from base_audit.systems.s2_report_collection.models import IndicatorConfig, ReportAuditConfig, UnitSettings


CALAMINE_INSTALLED = importer._calamine_load is not None


def _legacy_open(path: Path):
    """屏蔽快速引擎，得到传统引擎的读取结果。"""
    original = importer._calamine_load
    importer._calamine_load = None
    try:
        return importer._open(path)
    finally:
        importer._calamine_load = original


def _read_all(result) -> dict[str, list[list]]:
    names, read_rows, close = result
    try:
        return {name: read_rows(name) for name in names}
    finally:
        close()


def _same_cell(left, right) -> bool:
    # 空串与 None 在导入层语义等价（_number/_code/str 兜底一致），对拍按等价比较。
    return bool(left == right) or (left in (None, "") and right in (None, ""))


def _same_matrix(left: dict[str, list[list]], right: dict[str, list[list]]) -> bool:
    if list(left) != list(right):
        return False
    for name in left:
        rows_left, rows_right = left[name], right[name]
        if len(rows_left) != len(rows_right):
            return False
        for row_left, row_right in zip(rows_left, rows_right):
            if len(row_left) != len(row_right) or not all(
                _same_cell(a, b) for a, b in zip(row_left, row_right)
            ):
                return False
    return True


def _write_xlsx(path: Path) -> None:
    book = Workbook(); sheet = book.active; sheet.title = "20202"
    # 前导两行全空（含显式空串形态）：快速引擎必须按原行位补齐，不得前移。
    sheet.append(["", "", ""]); sheet.append(["", "", ""])
    sheet.append(["指标", "名称", "值"])
    sheet.append([20201001, "机构名称", "甲银行"])
    sheet.append([20202001, "整数", 7])
    sheet.append([20202002, "小数", 1234.5])
    sheet.append([20202003, "空值", None])
    sheet.append([20202004, "布尔", True])
    sheet.append([20202005, "日期", datetime(2026, 6, 30)])
    sheet.append([])
    sheet.append(["尾部", None, 3])
    extra = book.create_sheet("20201")
    extra.append(["指标", "名称", "值"])
    extra.append([20201002, "机构代码", "9144"])
    book.save(path)


def _write_xls(path: Path) -> None:
    import xlwt

    book = xlwt.Workbook(); sheet = book.add_sheet("20202")
    # 前导两行留空：xlwt 不写行元素，xlrd 按行位读出空行；快速引擎同样补齐。
    rows = [
        ["指标", "名称", "值"],
        [20201001, "机构名称", "甲银行"],
        [20202001, "整数", 7],
        [20202002, "小数", 1234.5],
        [20202003, "空值", None],
        [20202004, "布尔", True],
    ]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if value is not None:
                sheet.write(r + 2, c, value)
    sheet.write(8, 0, "尾部"); sheet.write(8, 2, 3)
    second = book.add_sheet("20201")
    second.write(0, 0, "指标"); second.write(0, 1, "名称"); second.write(0, 2, "值")
    second.write(1, 0, 20201002); second.write(1, 1, "机构代码"); second.write(1, 2, "9144")
    book.save(path)


@pytest.mark.skipif(not CALAMINE_INSTALLED, reason="未安装 python-calamine")
@pytest.mark.parametrize(
    ("writer", "suffix"),
    [(_write_xlsx, ".xlsx"), (_write_xls, ".xls")],
    ids=["xlsx", "xls"],
)
def test_fast_reader_values_match_legacy_engine(tmp_path: Path, writer, suffix: str) -> None:
    """两套引擎读同一文件：表名、行列数与逐单元格取值/类型一致。"""
    path = tmp_path / f"9144#2026-06-30#01#20202#甲银行{suffix}"
    writer(path)
    fast = _read_all(importer._open(path))
    legacy = _read_all(_legacy_open(path))
    assert _same_matrix(fast, legacy)
    # 除“空串≈None”等价视例外，类型也必须一致（int/float/str/datetime/bool）。
    for name in fast:
        for row_fast, row_legacy in zip(fast[name], legacy[name]):
            for cell_fast, cell_legacy in zip(row_fast, row_legacy):
                if _same_cell(cell_fast, cell_legacy) and cell_fast not in (None, "") and cell_legacy not in (None, ""):
                    assert type(cell_fast) is type(cell_legacy), (name, cell_fast, cell_legacy)


def test_fast_reader_absent_falls_back_silently(tmp_path: Path) -> None:
    """兼容构建未打包快速引擎：静默走传统引擎，不产生任何提示。"""
    path = tmp_path / "9144#2026-06-30#01#20202#甲银行.xlsx"
    _write_xlsx(path)
    original = importer._calamine_load
    importer._calamine_load = None
    try:
        warnings: list[str] = []
        result = importer._open(path, warnings)
        assert _read_all(result)
        assert warnings == []
    finally:
        importer._calamine_load = original


@pytest.mark.skipif(not CALAMINE_INSTALLED, reason="未安装 python-calamine")
def test_fast_reader_failure_falls_back_with_warning(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """引擎在具体文件上失败：逐文件提示并用传统引擎重读，结果一致。"""
    path = tmp_path / "9144#2026-06-30#01#20202#甲银行.xlsx"
    _write_xlsx(path)

    def broken(_path):
        raise RuntimeError("boom")

    monkeypatch.setattr(importer, "_calamine_load", broken)
    warnings: list[str] = []
    result = importer._open(path, warnings)
    assert _same_matrix(_read_all(result), _read_all(_legacy_open(path)))
    assert any("快速读取引擎处理失败" in warning and path.name in warning for warning in warnings)


@pytest.mark.skipif(not CALAMINE_INSTALLED, reason="未安装 python-calamine")
def test_import_results_are_identical_between_engines(tmp_path: Path) -> None:
    """同一目录走快速/传统引擎完成导入：全部 IndicatorRecord 逐字段一致。"""
    directory = tmp_path / "current"; directory.mkdir()
    for index, suffix in enumerate((".xlsx", ".xls")):
        target = directory / f"9144#2026-06-30-{index}#01#20202#甲银行{suffix}"
        (_write_xls if suffix == ".xls" else _write_xlsx)(target)
    config = ReportAuditConfig(
        indicators={
            "20201001": IndicatorConfig("20201001", "机构名称", "20202", "文字"),
            "20201002": IndicatorConfig("20201002", "机构代码", "20201", "文字"),
            "20202001": IndicatorConfig("20202001", "整数指标", "20202", "余额"),
            "20202002": IndicatorConfig("20202002", "小数指标", "20202", "余额"),
            "20202005": IndicatorConfig("20202005", "日期指标", "20202", "文字"),
        },
        unit_settings=UnitSettings("元", "元", "元"),
    )

    fast_result = import_period_directory(directory, config=config, label="本期")
    original = importer._calamine_load
    importer._calamine_load = None
    try:
        legacy_result = import_period_directory(directory, config=config, label="本期")
    finally:
        importer._calamine_load = original

    assert fast_result.records == legacy_result.records
    assert fast_result.warnings == legacy_result.warnings
    assert fast_result.duplicates == legacy_result.duplicates
    assert fast_result.records, "样例应产生导入记录"
