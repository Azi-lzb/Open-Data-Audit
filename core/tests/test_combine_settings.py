from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from base_audit.settings import SettingsStore


def test_legacy_combine_plans_migrate_once_to_user_settings(tmp_path: Path) -> None:
    """退役 config.xlsx 中可编辑组合方案应迁入用户设置。"""
    legacy_path = tmp_path / "config" / "config.xlsx"
    legacy_path.parent.mkdir()
    book = Workbook()
    plans = book.active
    plans.title = "组合方案"
    plans.append(["方案名称", "分组方式", "文件名正则", "方案状态", "输出标识"])
    plans.append(["核查表固定方案", "文件名正则", r"(?P<组合>.+)", "固定只读", "核查表合并"])
    plans.append(["华东机构", "文件名关键字", "", "当前使用", "华东组合"])
    groups = book.create_sheet("组合分组")
    groups.append(["方案名称", "组合名称", "匹配关键字", "启用"])
    groups.append(["华东机构", "杭州", "杭州分行、杭州", "是"])
    book.save(legacy_path)
    book.close()

    store = SettingsStore(tmp_path / "data" / "用户设置.json")
    settings = store.load()

    assert store.sanitized is True
    assert settings.active_combine_sheets_plan_id == "legacy-1"
    assert settings.combine_sheets_plans == [{
        "id": "legacy-1",
        "name": "华东机构",
        "mode": "name",
        "pattern": "",
        "output_marker": "华东组合",
        "groups": [{"name": "杭州", "keywords": ["杭州分行", "杭州"]}],
    }]

    store.save(settings)
    migrated = SettingsStore(store.path).load()
    assert migrated.combine_sheets_plans == settings.combine_sheets_plans
    assert migrated.active_combine_sheets_plan_id == "legacy-1"


def test_invalid_legacy_plan_falls_back_to_safe_defaults(tmp_path: Path) -> None:
    """损坏的历史方案不得阻止启动，回落到可执行的默认正则方案。"""
    legacy_path = tmp_path / "config" / "config.xlsx"
    legacy_path.parent.mkdir()
    book = Workbook()
    plans = book.active
    plans.title = "组合方案"
    plans.append(["方案名称", "分组方式", "文件名正则", "方案状态"])
    plans.append(["损坏正则", "文件名正则", "[", "当前使用"])
    book.save(legacy_path)
    book.close()

    settings = SettingsStore(tmp_path / "data" / "用户设置.json").load()

    assert settings.active_combine_sheets_plan_id == "default"
    assert settings.combine_sheets_plans[0]["name"] == "按文件名正则提取组合"
