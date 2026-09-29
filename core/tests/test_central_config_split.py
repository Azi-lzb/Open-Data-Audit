from __future__ import annotations

import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import pytest

from base_audit.systems.s3_central_statistics.config import (
    CENTRAL_CONFIG_ERROR,
    COMMON_SHEETS,
    COMPARISON_SHEETS,
    ACTION_RULE_SHEET,
    CUMULATIVE_RULE_SHEET,
    EXPRESSION_RULE_SHEET,
    INDICATOR_SHEET,
    SPECIAL_RULE_SHEET,
    UNIT_EXCEPTION_SHEET,
    check_central_config_bundle,
    ensure_split_central_configs,
    load_central_config,
    validate_central_config,
    write_default_central_profile,
)
from base_audit.systems.s3_central_statistics.comparison_engine import (
    _append_explanation,
    _band_explanation,
    _band_remark,
)


def test_current_split_books_are_created_without_legacy_source() -> None:
    with TemporaryDirectory() as folder:
        root = Path(folder)
        outputs = ensure_split_central_configs(root)
        assert set(outputs) == {"common", "comparison", "cross", "forms"}
        assert all(path.is_file() for path in outputs.values())
        common = load_workbook(outputs["common"], read_only=True)
        try:
            assert set(COMMON_SHEETS) <= set(common.sheetnames)
            assert EXPRESSION_RULE_SHEET not in common.sheetnames
            assert INDICATOR_SHEET not in common.sheetnames
        finally:
            common.close()


def test_each_feature_loads_common_plus_only_its_own_rules() -> None:
    with TemporaryDirectory() as folder:
        outputs = ensure_split_central_configs(Path(folder))

        comparison = load_central_config((outputs["common"], outputs["comparison"]))
        assert comparison.alerts
        assert comparison.rules
        assert comparison.complex_rules
        assert not comparison.cross_rules
        assert not comparison.toform_params

        cross = load_central_config((outputs["common"], outputs["cross"]))
        assert cross.rules
        assert cross.cross_rules
        assert not cross.complex_rules

        forms = load_central_config((outputs["common"], outputs["forms"]))
        assert forms.alerts
        assert forms.toform_params
        assert not forms.cross_rules

        assert check_central_config_bundle(outputs["common"], outputs["comparison"], "comparison") == ""
        assert check_central_config_bundle(outputs["common"], outputs["cross"], "cross") == ""
        assert check_central_config_bundle(outputs["common"], outputs["forms"], "forms") == ""


def test_alert_ranges_have_explicit_upper_bounds_and_keep_boundary_semantics() -> None:
    with TemporaryDirectory() as folder:
        outputs = ensure_split_central_configs(Path(folder))
        config = load_central_config((outputs["common"], outputs["comparison"]))
        first = config.alerts[0]
        assert set(("下限", "上限", "备注文字", "填充颜色", "金额变动阈值", "是否说明")) <= set(first)
        # 配置按百分数数值填写：-99 表示下降99%，运行环比仍为 -0.99。
        assert float(first["上限"]) == -99
        # 上限边界归入下一档：与旧版“最大下限命中”的行为一致。
        first_remark = _band_remark(-0.91, config.alerts)[0]
        second_remark = _band_remark(-0.9, config.alerts)[0]
        assert first_remark == "降幅[-99%,-90%) | 缩小10倍-100倍"
        assert second_remark == "降幅[-90%,-80%) | 缩小5倍-10倍"


def test_alert_band_can_append_explanation_without_overwriting_existing_rules() -> None:
    band = {
        "备注文字": "增幅100倍以上",
        "金额变动阈值": "0.5",
        "是否说明": "绝对值变动超过{金额变动阈值}{目标单位}，且{备注文字}。",
    }
    alert_text = _band_explanation(0.5001, band, "亿元")
    combined = _append_explanation("表达式规则触发", alert_text)

    assert alert_text == "绝对值变动超过0.5亿元，且增幅100倍以上。"
    assert combined == (
        "表达式规则触发    ||    "
        "绝对值变动超过0.5亿元，且增幅100倍以上。"
    )
    assert _band_explanation(0.5, band, "亿元") == ""
    assert _append_explanation(combined, alert_text) == combined
    assert _append_explanation(
        "表达式规则触发    ||    环比规则触发",
        "环比规则触发 || 特殊规则触发",
    ) == "表达式规则触发    ||    环比规则触发    ||    特殊规则触发"


def test_rules_are_split_into_maintenance_friendly_sheets() -> None:
    with TemporaryDirectory() as folder:
        outputs = ensure_split_central_configs(Path(folder))
        common = load_workbook(outputs["common"], read_only=True, data_only=True)
        comparison = load_workbook(outputs["comparison"], read_only=True, data_only=True)
        try:
            assert [cell.value for cell in common[UNIT_EXCEPTION_SHEET][1]][:3] == [
                "规则编号", "来源表单", "指标代码",
            ]
            assert "机构地区参照" in common.sheetnames
            assert "比较指标代码" in [cell.value for cell in comparison[ACTION_RULE_SHEET][1]]
            assert "是否取反" in [cell.value for cell in comparison[EXPRESSION_RULE_SHEET][1]]
            assert CUMULATIVE_RULE_SHEET not in comparison.sheetnames
            assert SPECIAL_RULE_SHEET not in comparison.sheetnames
        finally:
            common.close()
            comparison.close()


def test_new_rule_sheets_are_normalized_for_existing_engines() -> None:
    with TemporaryDirectory() as folder:
        outputs = ensure_split_central_configs(Path(folder))
        config = load_central_config((outputs["common"], outputs["comparison"]))
        assert {row["类型"] for row in config.rules} == {"单位不转换"}
        assert len(config.action_rules) == 2
        assert {row["类型"] for row in config.action_rules} == {"累计不应下降", "特殊阈值"}
        assert config.complex_rules[0]["校验描述"].startswith("示例")


def test_date_cannot_be_saved_as_amplitude_threshold() -> None:
    with TemporaryDirectory() as folder:
        outputs = ensure_split_central_configs(Path(folder))
        path = outputs["comparison"]
        book = load_workbook(path)
        try:
            sheet = book[ACTION_RULE_SHEET]
            headers = [cell.value for cell in sheet[1]]
            sheet.cell(3, headers.index("变幅阈值(%)") + 1).value = "2026-01-20 00:00:00"
            book.save(path)
        finally:
            book.close()
        errors = validate_central_config(load_central_config(path))
        assert any("变幅阈值不能填写日期" in error for error in errors)


def test_amplitude_action_requires_amplitude_threshold() -> None:
    with TemporaryDirectory() as folder:
        outputs = ensure_split_central_configs(Path(folder))
        path = outputs["comparison"]
        book = load_workbook(path)
        try:
            sheet = book[ACTION_RULE_SHEET]
            headers = [cell.value for cell in sheet[1]]
            # 第 3 行 = SP001（特殊阈值形状，动作文本存“规则动作”列）。
            sheet.cell(3, headers.index("规则动作") + 1).value = "指标变幅异常需说明。"
            sheet.cell(3, headers.index("变幅阈值(%)") + 1).value = ""
            sheet.cell(3, headers.index("启用") + 1).value = "是"
            book.save(path)
        finally:
            book.close()
        errors = validate_central_config(load_central_config(path))
        assert any("必须填写变幅阈值" in error for error in errors)
