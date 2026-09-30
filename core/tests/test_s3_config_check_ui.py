"""Regression coverage for S3 central configuration checks and UI dispatch."""

from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from config_paths import resolve_test_config

from base_audit.systems.s3_central_statistics.config import (
    COMMON_CONFIG_NAME,
    CROSS_CONFIG_NAME,
    CROSS_SHEET,
    SHEET_HEADERS,
)
from base_audit.web_app import WebApi


REPO_ROOT = Path(__file__).resolve().parents[2]
COMMON_CONFIG = resolve_test_config(COMMON_CONFIG_NAME)
CROSS_CONFIG = resolve_test_config(CROSS_CONFIG_NAME)
FORMAL_CONFIGS_READY = COMMON_CONFIG.is_file() and CROSS_CONFIG.is_file()
FRONTEND = REPO_ROOT / "core" / "frontend" / "web" / "index.html"


def _api_for_config(common_path: Path, cross_path: Path) -> WebApi:
    """Create only the state used by the configuration-check API."""
    api = WebApi.__new__(WebApi)
    api.project_root = REPO_ROOT
    api.state = {
        "busy": False,
        "centralCommonConfig": str(common_path),
        "centralCrossConfig": str(cross_path),
        "log": [],
    }
    api._central_config_paths = lambda profile: (
        (common_path, cross_path) if profile == "cross" else ()
    )
    api._log_detail = lambda _text: None
    api._refresh_config_issues = lambda: None
    return api


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _header_columns(sheet) -> tuple[int, dict[str, int]]:
    required = set(SHEET_HEADERS[CROSS_SHEET])
    for row_index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        columns = {
            str(value).strip(): col_index
            for col_index, value in enumerate(values, start=1)
            if value is not None and str(value).strip()
        }
        if required <= columns.keys():
            return row_index, columns
    raise AssertionError(f"工作表缺少 3.2 表头：{CROSS_SHEET}")


def test_central_check_report_matches_h2_prompt_headings() -> None:
    report = WebApi._central_check_report(
        "本期数值核对配置（3.2）",
        [{"规则编号": "TEST-01", "启用": "是", "问题": ["公式语法错误"]}],
        {"total": 2, "enabled": 1, "vocabulary": ["token 含词表外值"]},
    )

    assert "共 2 条规则（启用 1，停用 1）" in report
    assert "【问题】1 处：" in report
    assert "【取数词表提示】共 1 处：" in report
    assert "【错误】" not in report
    assert "【提示】" not in report


@pytest.mark.skipif(not FORMAL_CONFIGS_READY, reason="缺少正式 3.0/3.2 配置工作簿")
def test_formal_32_check_reports_selected_workbooks_sheet_and_syntax_scope() -> None:
    before = (_sha256(COMMON_CONFIG), _sha256(CROSS_CONFIG))

    result = _api_for_config(COMMON_CONFIG, CROSS_CONFIG).check_central_cross_formulas()

    assert result["source_workbooks"] == [
        str(COMMON_CONFIG.resolve()),
        str(CROSS_CONFIG.resolve()),
    ]
    assert result["active_sheets"] == [CROSS_SHEET]
    assert result["checked"] > 0
    report = result["report_text"]
    assert "【本次读取的配置工作簿】" in report
    assert COMMON_CONFIG.name in report
    assert CROSS_CONFIG.name in report
    assert f"【本模式参与执行的 3.2 工作表】{CROSS_SHEET}" in report
    assert "语法" in report
    assert "token" in report
    assert "校验公式" in report
    assert (_sha256(COMMON_CONFIG), _sha256(CROSS_CONFIG)) == before


def test_32_check_reports_missing_required_sheet_in_disposable_workbook(tmp_path: Path) -> None:
    common_path = tmp_path / COMMON_CONFIG_NAME
    cross_path = tmp_path / CROSS_CONFIG_NAME
    Workbook().save(common_path)
    Workbook().save(cross_path)  # default Sheet is valid XLSX but lacks CROSS_SHEET

    with pytest.raises(ValueError) as caught:
        _api_for_config(common_path, cross_path).check_central_cross_formulas()

    message = str(caught.value)
    assert cross_path.name in message
    assert CROSS_SHEET in message
    assert "缺少必需工作表" in message


@pytest.mark.skipif(not FORMAL_CONFIGS_READY, reason="缺少正式 3.0/3.2 配置工作簿")
def test_32_check_reports_a_malformed_enabled_rule_from_a_disposable_copy(
    tmp_path: Path,
) -> None:
    common_copy = tmp_path / COMMON_CONFIG.name
    cross_copy = tmp_path / CROSS_CONFIG.name
    shutil.copy2(COMMON_CONFIG, common_copy)
    shutil.copy2(CROSS_CONFIG, cross_copy)

    book = load_workbook(cross_copy)
    try:
        sheet = book[CROSS_SHEET]
        _header_row, columns = _header_columns(sheet)
        target_row = sheet.max_row + 1
        values = {
            "校验编码": "TEST-CROSS-SYNTAX",
            "校验名称": "配置检查语法回归",
            "left": "[33370,余额,人民币,月,1]",
            "校验公式": "left >",
            "禁用": "否",
        }
        for header, value in values.items():
            sheet.cell(target_row, columns[header], value)
        book.save(cross_copy)
    finally:
        book.close()

    before = (_sha256(common_copy), _sha256(cross_copy))
    result = _api_for_config(common_copy, cross_copy).check_central_cross_formulas()

    assert result["total_issues"] >= 1
    issue = next(item for item in result["issues"] if item["规则编号"] == "TEST-CROSS-SYNTAX")
    assert any("校验公式" in problem and "语法/求值" in problem for problem in issue["问题"])
    assert "TEST-CROSS-SYNTAX" in result["report_text"]
    assert (_sha256(common_copy), _sha256(cross_copy)) == before


def test_central_config_buttons_dispatch_and_show_matching_titles() -> None:
    source = FRONTEND.read_text(encoding="utf-8")

    expected = {
        "common": ("check_central_common_config", "大集中通用 3.0"),
        "cross": ("check_central_cross_formulas", "本期数值核对 3.2"),
        "forms": ("check_central_forms_config", "指标比较拆分 3.3"),
    }
    for kind, (bridge_method, title) in expected.items():
        assert f"onclick=\"checkCentralConfig('{kind}')\"" in source
        assert f"{kind}:()=>bridge.api.{bridge_method}()" in source
        assert title in source
