from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from config_paths import resolve_test_config

from base_audit.systems.s2_report_collection.config import load_config, write_default_config
from base_audit.systems.s2_report_collection.config_check import check_s2_config, format_config_report
from base_audit.systems.s2_report_collection.models import IndicatorConfig, IndicatorRecord, PeriodBandConfig, ReportAuditConfig, SourceRef
from base_audit.systems.s2_report_collection.period_engine import PeriodComparisonEngine, PeriodRuleExecutionError


ROOT = Path(__file__).resolve().parents[2]


def _record(code: str, value: float | Decimal, *, period: str, value_type: str = "余额") -> IndicatorRecord:
    source = SourceRef("source.xlsx", "20202", 4)
    normalized_value = value if isinstance(value, Decimal) else Decimal(str(value))
    return IndicatorRecord("9144", "甲银行", "9144", "农商行", "一部", "示例地区", "20202", code, code,
                           period, period + "-30", normalized_value, value_type, "元", normalized_value, "元", source)


def test_formal_and_default_s2_configs_pass_read_only_check() -> None:
    for path in (
        resolve_test_config("2.报表采集系统_配置.xlsx"),
        resolve_test_config("默认配置/2.报表采集系统_配置.xlsx"),
    ):
        report = check_s2_config(path)
        assert report["passed"], format_config_report(report)
        checks = {check["id"]: check for check in report["checks"]}
        assert checks["file_read"]["status"] == "passed"
        assert checks["required_sheets"]["status"] == "passed"
        assert checks["indicators"]["status"] in {"passed", "warning"}
        assert checks["period_strategies"]["status"] in {"passed", "warning"}
        assert checks["runtime_parse"]["status"] in {"passed", "warning"}
        formatted = format_config_report(report)
        assert "实际检查项目：" in formatted
        assert "检查范围说明：" in formatted
        assert "不读取本期、上期报表或大集中外部数据文件" in formatted
        assert report["indicator_count"] == 105
        assert report["strategy_count"] == 56
        config = load_config(path)
        band = next(item for item in config.bands if item.strategy_group == "BAL01" and item.rule_id == "B008")
        assert (band.lower, band.upper) == (30, 50)
        book = load_workbook(path, read_only=False, data_only=False)
        try:
            sheet = book["环比策略"]
            # 当前正式工作簿没有配置筛选区域；不要把历史测试中的过时筛选断言
            # 与通用设置表结构耦合起来。
            assert sheet["D10"].number_format == "General"
            assert "%" not in sheet["D10"].number_format
            for business_name in ("指标参照", "机构参照", "环比策略", "校验规则", "外部核对规则"):
                business = book[business_name]
                assert all(cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb.endswith("1F4E78") for cell in business[1])
                assert all(cell.font.bold and cell.font.color.rgb.endswith("FFFFFF") for cell in business[1])
                body_cells = [cell for row in business.iter_rows(min_row=2) for cell in row if cell.value not in (None, "")]
                assert body_cells
                assert all(cell.fill.fill_type is None for cell in body_cells)
                assert all(not cell.font.bold and cell.font.color.rgb.endswith("000000") for cell in body_cells)
            assert [cell.value for cell in book["外部核对规则"][1]] == [
                "规则编号", "指标代码", "比较方式", "外部指标名称", "容差", "级别", "启用",
            ]
            general = book["通用设置"]
            unit_values = {row[0]: row[1] for row in general.iter_rows(min_row=2, values_only=True) if row[0] in {"源数据单位", "外部文件单位", "输出文件单位"}}
            assert unit_values == {
                "源数据单位": config.unit_settings.source_unit,
                "外部文件单位": config.unit_settings.external_file_unit,
                "输出文件单位": config.unit_settings.output_file_unit,
            }
            assert set(unit_values) == {"源数据单位", "外部文件单位", "输出文件单位"}
            assert set(unit_values.values()) <= {"元", "万元", "亿元"}
            assert unit_values["源数据单位"] == "元"
            assert unit_values["输出文件单位"] == "元"
            assert general["B6"].value in {"是", "否"}
            assert general["B7"].value in {"是", "否"}
            assert (general["B6"].value == "是") is config.hide_unchanged_period_rows
            assert (general["B7"].value == "是") is config.hide_matching_external_rows
        finally:
            book.close()


def test_checker_blocks_gap_and_duplicate_strategy(tmp_path: Path) -> None:
    path = tmp_path / "config.xlsx"
    book = Workbook()
    indicator = book.active
    indicator.title = "指标参照"
    indicator.append(["指标代码", "指标名称", "表单代码", "数据属性", "环比启用", "环比策略组", "最小变动值", "禁用"])
    indicator.append(["20202001", "存款", "20202", "余额", "是", "BAL01", 0, "否"])
    org = book.create_sheet("机构参照")
    org.append(["机构名称", "社会信用代码", "机构类别", "承接行", "地区", "报表项目", "禁用"])
    strategy = book.create_sheet("环比策略")
    strategy.append(["策略组", "规则编号", "数据属性", "下限", "上限", "审核级别", "分类说明", "禁用"])
    strategy.append(["BAL01", "B1", "余额", None, -20, "提示", "下降", "否"])
    strategy.append(["BAL01", "B2", "余额", -10, None, "关注", "增长", "否"])
    rules = book.create_sheet("校验规则")
    rules.append(["规则编号", "类型", "描述", "规则内容", "级别", "禁用", "备注"])
    external = book.create_sheet("外部核对规则")
    external.append(["规则编号", "指标代码", "比较方式", "外部指标名称", "容差", "级别", "启用"])
    units = book.create_sheet("通用设置")
    units.append(["设置项", "设置值", "说明"])
    units.append(["源数据单位", "元", ""])
    units.append(["外部文件单位", "元", ""])
    units.append(["输出文件单位", "元", ""])
    units.append(["换算范围", "", ""])
    units.append(["是否隐藏无变动指标", "否", ""])
    units.append(["是否隐藏与大集中一致指标", "否", ""])
    book.save(path)
    book.close()
    report = check_s2_config(path)
    assert not report["passed"]
    assert any("空档" in issue["message"] for issue in report["issues"])
    checks = {check["id"]: check for check in report["checks"]}
    assert checks["period_strategies"]["status"] == "failed"
    assert "阻断问题" in checks["period_strategies"]["detail"]


def test_report_marks_checks_not_run_when_config_file_is_missing(tmp_path: Path) -> None:
    report = check_s2_config(tmp_path / "missing.xlsx")
    checks = {check["id"]: check for check in report["checks"]}
    assert checks["file_read"]["status"] == "failed"
    assert all(checks[check_id]["status"] == "not_run" for check_id in (
        "required_sheets", "indicators", "period_strategies", "runtime_parse",
    ))
    assert "配置文件不存在" in format_config_report(report)


def test_s2_check_ui_displays_check_status_scope_and_issue_details() -> None:
    source = (ROOT / "core/frontend/web/index.html").read_text(encoding="utf-8")
    function = source.split("async function checkPeriodConfig()", 1)[1].split("/* ============ 设置中心 ============ */", 1)[0]
    assert "实际检查项目：" in function
    assert "检查范围说明：" in function
    assert "问题明细：" in function
    assert "checkLabels[c.status]" in function


def test_checker_and_runtime_reject_legacy_external_system_column(tmp_path: Path) -> None:
    path = tmp_path / "legacy-external-rules.xlsx"
    write_default_config(path)
    book = load_workbook(path)
    external = book["外部核对规则"]
    for column, value in enumerate(
        ["规则编号", "指标代码", "外部系统", "外部指标名称", "比较方式", "容差", "级别", "启用"],
        start=1,
    ):
        external.cell(row=1, column=column, value=value)
    for column, value in enumerate(
        ["D001", "20201001", "大集中", "示例外部指标（请替换）", "等于", 0, "关注", "否"],
        start=1,
    ):
        external.cell(row=2, column=column, value=value)
    book.save(path)
    book.close()

    with pytest.raises(ValueError, match="外部核对规则"):
        load_config(path)
    report = check_s2_config(path)
    assert not report["passed"]
    assert any("外部核对规则" in issue["message"] for issue in report["issues"])


def test_checker_requires_external_rule_header_order(tmp_path: Path) -> None:
    path = tmp_path / "misordered-external-rules.xlsx"
    write_default_config(path)
    book = load_workbook(path)
    external = book["外部核对规则"]
    for column, value in enumerate(
        ["规则编号", "指标代码", "外部指标名称", "比较方式", "容差", "级别", "启用"],
        start=1,
    ):
        external.cell(row=1, column=column, value=value)
    book.save(path)
    book.close()

    with pytest.raises(ValueError, match="外部核对规则"):
        load_config(path)
    report = check_s2_config(path)
    assert not report["passed"]
    assert any("外部核对规则" in issue["message"] for issue in report["issues"])


@pytest.mark.parametrize(
    ("method", "tolerance", "expected_fragment"),
    [("近似等于", 0, "比较方式"), ("大于", 0.1, "容差")],
)
def test_checker_blocks_invalid_external_method_or_nonzero_ordering_tolerance(
    tmp_path: Path, method: str, tolerance: float, expected_fragment: str,
) -> None:
    path = tmp_path / "invalid-external-rule.xlsx"
    write_default_config(path)
    book = load_workbook(path)
    external = book["外部核对规则"]
    external["C2"] = method
    external["E2"] = tolerance
    external["G2"] = "是"
    book.save(path)
    book.close()

    report = check_s2_config(path)
    assert not report["passed"]
    assert any(expected_fragment in issue["message"] for issue in report["issues"])


def test_period_strategy_requires_single_match_and_honors_minimum_gate() -> None:
    indicator = IndicatorConfig(
        "20202001", indicator_name="存款", form_code="20202", data_type="余额",
        period_strategy_group="BAL01", min_change_value=100,
    )
    config = ReportAuditConfig(
        indicators={"20202001": indicator},
        bands=[PeriodBandConfig("B1", None, None, "异常", "提示", False, "BAL01", "余额")],
    )
    current = _record("20202001", 950, period="2026-06")
    previous = _record("20202001", 1000, period="2026-03")
    finding = PeriodComparisonEngine().run({current.logical_key: current}, {previous.logical_key: previous}, config)[0]
    assert finding.status == "正常"
    assert finding.severity == ""
    assert "未达到最小变动值" in finding.calculation_trace

    config.bands.extend([PeriodBandConfig("B2", -20, 20, "重复", "提示", False, "BAL01", "余额")])
    try:
        PeriodComparisonEngine().run({current.logical_key: current}, {previous.logical_key: previous}, config)
    except PeriodRuleExecutionError as exc:
        assert "多个规则" in str(exc)
    else:
        raise AssertionError("overlapping strategy must fail explicitly")


def test_percentage_strategy_uses_percentage_point_scale_directly() -> None:
    indicator = IndicatorConfig(
        "20201028", indicator_name="持股比例", form_code="20201", data_type="百分数",
        period_strategy_group="PCT01", min_change_value=0,
    )
    config = ReportAuditConfig(
        indicators={"20201028": indicator},
        bands=[
            PeriodBandConfig("B005", -30, 0, "下降30个百分点以内", "提示", False, "PCT01", "百分数"),
            PeriodBandConfig("B006", 0, 30, "上升30个百分点以内", "提示", False, "PCT01", "百分数"),
        ],
    )
    current = _record("20201028", 1, period="2026-06", value_type="百分数")
    previous = _record("20201028", 9.92, period="2026-03", value_type="百分数")
    finding = PeriodComparisonEngine().run(
        {current.logical_key: current}, {previous.logical_key: previous}, config
    )[0]
    assert finding.change_rate == Decimal("-8.92")
    assert finding.band == "PCT01-B005"
    assert finding.severity == "提示"
    assert finding.description == "下降30个百分点以内"
