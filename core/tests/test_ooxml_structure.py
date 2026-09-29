# -*- coding: utf-8 -*-
"""OOXML 结构解析与完整性：namespace 无关、路径解析、悬空检测、删表保护。"""

from __future__ import annotations

import sys
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from base_audit.systems.s3_central_statistics.ooxml_structure import (  # noqa: E402
    OoxmlStructureError,
    parse_relationships,
    parse_workbook_sheets,
    read_workbook_structure,
    resolve_relationship_target,
    validate_workbook_relationships,
)

WB_DEFAULT = (
    '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    '<sheets><sheet name="报表清单" sheetId="1" r:id="rId1"/>'
    '<sheet name="A1411" sheetId="2" r:id="rId2"/></sheets></workbook>'
)
WB_NS0 = (
    '<ns0:workbook xmlns:ns0="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    '<ns0:sheets><ns0:sheet name="报表清单" sheetId="1" r:id="rId1"/></ns0:sheets>'
    '</ns0:workbook>'
)
RELS_DEFAULT = (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://x/worksheet" Target="worksheets/sheet1.xml"/>'
    '<Relationship Id="rId2" Type="http://x/worksheet" Target="worksheets/sheet2.xml"/>'
    '</Relationships>'
)
RELS_NS0 = (
    '<ns0:Relationships xmlns:ns0="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<ns0:Relationship Type="http://x/worksheet" Target="/xl/worksheets/sheet1.xml" Id="rId1"/>'
    '</ns0:Relationships>'
)
RELS_PKG_NON_SELF_CLOSED = (
    '<pkg:Relationships xmlns:pkg="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<pkg:Relationship Type="http://x/worksheet" Target="worksheets/sheet1.xml" Id="rId1">'
    '</pkg:Relationship></pkg:Relationships>'
)

PARTS = {
    "xl/workbook.xml": WB_DEFAULT.encode(),
    "xl/_rels/workbook.xml.rels": RELS_DEFAULT.encode(),
    "xl/worksheets/sheet1.xml": b"<x/>",
    "xl/worksheets/sheet2.xml": b"<x/>",
}


class ParseTests(unittest.TestCase):
    """场景 1/2/3/4：默认 namespace、ns0 前缀、pkg 前缀、非自闭合。"""

    def test_default_namespace_parses(self) -> None:
        sheets = parse_workbook_sheets(WB_DEFAULT)
        self.assertEqual([s["name"] for s in sheets], ["报表清单", "A1411"])
        rels = parse_relationships(RELS_DEFAULT)
        self.assertEqual(rels["rId1"].target, "worksheets/sheet1.xml")

    def test_ns0_prefix_parses(self) -> None:
        """openpyxl 重存后产生 ns0: 前缀（用户报障形态）。"""
        structure = read_workbook_structure(WB_NS0, RELS_NS0)
        self.assertEqual(structure.sheets[0]["name"], "报表清单")
        resolved = resolve_relationship_target("xl/workbook.xml", "/xl/worksheets/sheet1.xml")
        self.assertEqual(resolved, "xl/worksheets/sheet1.xml")

    def test_pkg_prefix_and_non_self_closed_parse(self) -> None:
        rels = parse_relationships(RELS_PKG_NON_SELF_CLOSED)
        self.assertIn("rId1", rels)
        self.assertEqual(rels["rId1"].target, "worksheets/sheet1.xml")

    def test_external_relationship_marked(self) -> None:
        rels = parse_relationships(
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId9" Target="http://example.com/x.xlsx" TargetMode="External"/>'
            "</Relationships>")
        self.assertTrue(rels["rId9"].is_external)


class ResolveTests(unittest.TestCase):
    """路径解析：相对 / 包根绝对 / 上跳。"""

    def test_relative_target(self) -> None:
        self.assertEqual(
            resolve_relationship_target("xl/workbook.xml", "worksheets/sheet1.xml"),
            "xl/worksheets/sheet1.xml")

    def test_absolute_target(self) -> None:
        self.assertEqual(
            resolve_relationship_target("xl/workbook.xml", "/xl/worksheets/sheet1.xml"),
            "xl/worksheets/sheet1.xml")

    def test_parent_target(self) -> None:
        self.assertEqual(
            resolve_relationship_target(
                "xl/worksheets/_rels/sheet1.xml.rels", "../theme/theme1.xml"),
            "xl/theme/theme1.xml")


class StructureErrorTests(unittest.TestCase):
    """场景 7/8：悬空 rId 与悬空 Target 主动报可读结构错误（无 KeyError）。"""

    def test_missing_relationship_raises_readable_error(self) -> None:
        workbook = (
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheets><sheet name="报表清单" sheetId="1" r:id="rId17"/></sheets></workbook>')
        with self.assertRaises(OoxmlStructureError) as ctx:
            read_workbook_structure(workbook, RELS_DEFAULT)
        self.assertIn("报表清单", str(ctx.exception))
        self.assertEqual(ctx.exception.error_code, "OOXML_WORKSHEET_RELATION_MISSING")
        # 用户提示不含技术词 rId
        self.assertNotIn("rId", ctx.exception.user_message)
        # 技术日志保留 rid
        self.assertEqual(ctx.exception.technical.get("rid"), "rId17")

    def test_dangling_target_raises(self) -> None:
        rels = (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Target="worksheets/missing.xml"/></Relationships>')
        workbook = (
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheets><sheet name="报表清单" sheetId="1" r:id="rId1"/></sheets></workbook>')
        with self.assertRaises(OoxmlStructureError) as ctx:
            validate_workbook_relationships(workbook, rels, PARTS)
        self.assertEqual(ctx.exception.error_code, "OOXML_WORKSHEET_PART_MISSING")
        self.assertFalse(ctx.exception.technical.get("part_exists", True))


class FastOoxmlCompileTests(unittest.TestCase):
    """场景 5/6：openpyxl 重存/删表后可编译；3.3 无金融表单映射可编译。"""

    def _template(self) -> Path:
        template = ROOT.parent / "config" / "3.3大集中指标比较拆分_配置.xlsx"
        if not template.is_file():
            self.skipTest("缺少 3.3 正式配置")
        return template

    def test_compile_real_33_without_form_mapping_sheet(self) -> None:
        from base_audit.systems.s3_central_statistics.fast_ooxml import TemplateCompiler

        compiled = TemplateCompiler(self._template()).compile()
        self.assertTrue(any(e["name"] == "报表清单" for e in compiled.sheet_registry))
        self.assertFalse(any(e["name"] == "金融表单映射" for e in compiled.sheet_registry))
        self.assertTrue(len(compiled.forms) >= 40)

    def test_openpyxl_resaved_template_compiles(self) -> None:
        import tempfile

        from openpyxl import load_workbook

        template = self._template()
        with tempfile.TemporaryDirectory() as folder:
            resaved = Path(folder) / "resaved.xlsx"
            book = load_workbook(template)
            book.save(resaved)
            book.close()
            from base_audit.systems.s3_central_statistics.fast_ooxml import TemplateCompiler

            compiled = TemplateCompiler(resaved).compile()
            self.assertTrue(any(e["name"] == "报表清单" for e in compiled.sheet_registry))

    def test_openpyxl_delete_sheet_then_compile(self) -> None:
        import tempfile

        from openpyxl import load_workbook

        template = self._template()
        with tempfile.TemporaryDirectory() as folder:
            trimmed = Path(folder) / "trimmed.xlsx"
            book = load_workbook(template)
            for name in list(book.sheetnames):
                if name.startswith("A1"):
                    del book[name]
            book.save(trimmed)
            book.close()
            from base_audit.systems.s3_central_statistics.fast_ooxml import TemplateCompiler

            compiled = TemplateCompiler(trimmed).compile()
            self.assertTrue(compiled.forms)


class DeleteSheetIntegrityTests(unittest.TestCase):
    """场景 9/10：Direct OOXML 删表后结构完整、openpyxl 可重开。"""

    def test_render_with_deletes_passes_integrity(self) -> None:
        import tempfile

        from openpyxl import load_workbook

        from base_audit.systems.s3_central_statistics.fast_ooxml import FastOoxmlRenderer

        template = ROOT.parent / "config" / "3.3大集中指标比较拆分_配置.xlsx"
        if not template.is_file():
            self.skipTest("缺少 3.3 正式配置")
        with tempfile.TemporaryDirectory() as folder:
            renderer = FastOoxmlRenderer(
                template, alerts=[], hide_empty_rows=False, delete_empty_sheets=True)
            output = Path(folder) / "out.xlsx"
            from base_audit.systems.s3_central_statistics.renderers import UnsupportedFastRenderOperation

            try:
                renderer.render_org(
                    org_name="测试机构", region_name="广东省",
                    record_date="2026-08-31", data={"指标": ("测试机构", 1.0)},
                    output_path=output)
            except UnsupportedFastRenderOperation as exc:
                # 结构无法可靠解析时保守终止（不伪实现）——
                # 关键是：可解释错误 + 无损坏输出。正常 definedNames
                # （_FilterDatabase 等）已支持安全删除/重映射，不应再触发。
                self.assertNotIn("KeyError", str(exc))
                self.assertFalse(output.exists())   # 不落一个损坏文件
                return
            with zipfile.ZipFile(output) as archive:
                workbook = archive.read("xl/workbook.xml").decode("utf-8")
                rels = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
                names = set(archive.namelist())
            # 无悬空（不抛错即通过）
            validate_workbook_relationships(workbook, rels, names)
            # openpyxl 再次打开成功
            book = load_workbook(output)
            self.assertTrue(book.sheetnames)
            book.close()


if __name__ == "__main__":
    unittest.main()
