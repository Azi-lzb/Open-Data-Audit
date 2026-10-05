from __future__ import annotations

import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from base_audit.web_app import WebApi
from base_audit.name_config import HISTORY_WORKBOOK_NAME, initialize_config
from base_audit.history import HISTORY_AUDIT_SHEET
from base_audit.workflow.catalog import save_custom_workflow
from base_audit.workflow.graph import WorkflowDefinition, WorkflowNode


class _Window:
    def __init__(self, selected: str) -> None:
        self.selected = selected
        self.calls: list[tuple[object, dict[str, object]]] = []

    def create_file_dialog(self, dialog_type: object, **kwargs: object) -> tuple[str, ...]:
        self.calls.append((dialog_type, kwargs))
        return (self.selected,)


class WebAppCompatibilityTests(unittest.TestCase):
    def test_run_log_preferences_are_independent_and_legacy_api_is_compatible(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            api = WebApi(root)
            api.update({
                "showRunDetailLogs": True,
                "exportRunLogs": False,
                "showPerformanceDiagnostics": True,
            })
            self.assertTrue(api.state["showRunDetailLogs"])
            self.assertFalse(api.state["exportRunLogs"])
            self.assertTrue(api.state["showPerformanceDiagnostics"])
            self.assertFalse(api.state["writeFlowLogs"])

            reloaded = WebApi(root)
            self.assertTrue(reloaded.state["showRunDetailLogs"])
            self.assertFalse(reloaded.state["exportRunLogs"])
            self.assertTrue(reloaded.state["showPerformanceDiagnostics"])

            reloaded.update({"writeFlowLogs": True})
            self.assertTrue(reloaded.state["showRunDetailLogs"])
            self.assertTrue(reloaded.state["exportRunLogs"])
            self.assertTrue(reloaded.state["showPerformanceDiagnostics"])
            reloaded.update({"writeFlowLogs": False})
            self.assertFalse(reloaded.state["showRunDetailLogs"])
            self.assertFalse(reloaded.state["exportRunLogs"])
            self.assertTrue(reloaded.state["showPerformanceDiagnostics"])

    def test_fresh_environment_leaves_template_dir_empty_until_dir_exists(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            api = WebApi(root)
            # 全新环境（发行包/新解压目录）：模板目录不留死路径，界面显示“未选择”。
            self.assertEqual(api.state["templateDir"], "")
            # 六册配置默认绑定不受影响：S1 配置回填默认路径。
            self.assertEqual(api.state["historyConfig"], str(api.history_path))

        with TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "templates").mkdir()
            api = WebApi(root)
            # 本机存在模板目录时仍预填，保持开发布局的便利。
            self.assertEqual(api.state["templateDir"], str(root / "templates"))

    def test_settings_ui_exposes_independent_log_options_and_existing_perf_filter(self) -> None:
        source = (ROOT / "frontend" / "web" / "index.html").read_text(encoding="utf-8")
        self.assertIn("row('运行明细日志'", source)
        self.assertIn("row('导出运行日志 Excel'", source)
        self.assertIn("row('性能与运行诊断信息'", source)
        self.assertNotIn("row('配置摘要'", source)
        self.assertIn("showPerformanceDiagnostics?lines.filter(isPerformanceLine):[]", source)
        self.assertIn("公式校验复制方式：", source)
        self.assertIn("条件格式判定方式：", source)
        self.assertNotIn("let logFull=", source)

    def test_period_config_check_button_is_in_field_header(self) -> None:
        source = (ROOT / "frontend" / "web" / "index.html").read_text(encoding="utf-8")
        check_button = source.index('onclick="checkPeriodConfig()"')
        header_start = source.rfind('<div class="field-head tight">', 0, check_button)
        normal_header_start = source.rfind('<div class="field-head">', 0, check_button)
        label = source.rfind("<label>报表采集系统配置</label>", 0, check_button)
        header_end = source.index("</div>", check_button)
        input_group = source.index('<div class="input-group">', check_button)
        picker_button = source.index('onclick="pickPc(\'pcConfig\')"', input_group)

        self.assertGreaterEqual(header_start, 0)
        self.assertGreater(header_start, normal_header_start)
        self.assertGreater(label, header_start)
        self.assertLess(check_button, header_end)
        self.assertLess(header_end, input_group)
        self.assertLess(input_group, picker_button)

    def test_central_workers_use_export_preference_without_coupling_page_options(self) -> None:
        from types import SimpleNamespace

        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            api.update({
                "showRunDetailLogs": True,
                "exportRunLogs": False,
                "showPerformanceDiagnostics": True,
            })
            api._central_config_paths = lambda _kind: Path(folder) / "config.xlsx"
            compare_result = SimpleNamespace(output_path=Path(folder) / "compare.xlsx", rows=[])
            with patch(
                "base_audit.systems.s3_central_statistics.service.run_comparison",
                return_value=compare_result,
            ) as run_comparison:
                api._central_comparison_worker(
                    Path(folder) / "current.csv", Path(folder) / "previous.csv",
                    Path(folder) / "compare.xlsx",
                )
            self.assertFalse(run_comparison.call_args.kwargs["write_flow_logs"])

            api.update({
                "showRunDetailLogs": False,
                "exportRunLogs": True,
                "showPerformanceDiagnostics": False,
            })
            cross_result = SimpleNamespace(output_path=Path(folder) / "cross.xlsx", rows=[])
            with patch(
                "base_audit.systems.s3_central_statistics.service.run_cross_period_check",
                return_value=cross_result,
            ) as run_cross_period:
                api._central_cross_period_worker(
                    Path(folder) / "current.csv", "", Path(folder) / "cross.xlsx",
                )
            self.assertTrue(run_cross_period.call_args.kwargs["write_flow_logs"])

    def test_cross_period_paths_are_persisted_independently(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            api = WebApi(root)
            api.state["centralCur"] = str(root / "执行比较本期.csv")
            api.state["centralOutput"] = str(root / "执行比较输出")
            api.state["centralOutputAuto"] = False
            api.state["centralCrossCur"] = str(root / "跨期核对本期.csv")
            api.state["centralCrossOutput"] = str(root / "跨期核对输出")
            api.state["centralCrossOutputAuto"] = False
            api._save_settings()

            reloaded = WebApi(root)
            self.assertEqual(reloaded.state["centralCur"], str(root / "执行比较本期.csv"))
            self.assertEqual(reloaded.state["centralCrossCur"], str(root / "跨期核对本期.csv"))
            self.assertEqual(reloaded.state["centralOutput"], str(root / "执行比较输出"))
            self.assertEqual(reloaded.state["centralCrossOutput"], str(root / "跨期核对输出"))

    def test_pywebview_file_filters_match_pywebview_regex(self) -> None:
        """系统原生过滤器必须是 pywebview 可解析的 `label (*.a;*.b)` 形态。"""
        import re

        from base_audit.web_app import pywebview_file_types

        valid = re.compile(r'^([\w ]+)\((\*(?:\.(?:\w+|\*))*(?:;\*(?:\.(?:\w+|\*))*)*)\)$')
        cases = pywebview_file_types((
            ("Excel 文件", "*.xlsx *.xls *.xlsm"),
            ("数据文件", "*.csv *.xlsx"),
            ("Excel 文件", "*.xlsx"),
        ))
        for item in cases:
            self.assertRegex(item, valid)
        self.assertEqual(cases[0], "Excel 文件 (*.xlsx;*.xls;*.xlsm)")
        self.assertEqual(cases[1], "数据文件 (*.csv;*.xlsx)")

    def test_openpyxl_read_pickers_do_not_offer_xls(self) -> None:
        """openpyxl 读不了 .xls：喂给它的选择器过滤器不得放行 .xls。"""
        import re
        from pathlib import Path as _Path

        source = (_Path(__file__).resolve().parents[2]
                  / "core" / "src" / "base_audit" / "web_app.py").read_text(encoding="utf-8")
        self.assertNotIn('*.xlsx *.xls *.xlsm', source)
        self.assertNotIn('*.xlsx *.xlsm *.xls', source)
        # 大集中/报表采集的数据文件 = 查询导出（csv/xlsx/xls，纯读取入口）。
        self.assertIn('("数据文件", "*.csv *.xlsx *.xls")', source)

    def test_central_rule_engine_is_fixed_to_v3(self) -> None:
        """V3 转正：引擎不再暴露选择；历史存档的 legacy 一律按 v3 执行。"""
        with TemporaryDirectory() as folder:
            root = Path(folder)
            api = WebApi(root)
            api.settings.central_rule_engine = "legacy"   # 旧存档
            # 界面 update 不再接受该键；持久化与状态输出固定 v3。
            api.update({"centralRuleEngine": "legacy"})
            api._save_settings()
            self.assertEqual(api.settings.central_rule_engine, "v3")
            self.assertEqual(api.get_state()["centralRuleEngine"], "v3")

    def test_saving_settings_draft_logs_generic_settings_message(self) -> None:
        """底部“保存设置”不应误报为仅保存了组合工作表方案。"""
        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            studio = api.get_dag_editor_data()
            result = api.save_config_editor_draft({
                "combineSheetsPlans": studio["combineSheetsPlans"],
                "activeCombineSheetsPlanId": studio["activeCombineSheetsPlanId"],
            })
            self.assertTrue(result["ok"])
            self.assertIn("设置已保存", api.state["log"][-1]["text"])

    def test_business_workflow_api_compiles_and_lists_all_templates(self) -> None:
        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            studio = api.get_dag_editor_data()
            template_ids = {item["templateId"] for item in studio["flowTemplates"]}
            self.assertEqual(template_ids, {"audit", "summary", "combine_sheets"})
            function_ids = {item["functionId"] for item in studio["functions"]}
            self.assertIn("merge.sheets", function_ids)
            self.assertNotIn("merge.organization", function_ids)

            from base_audit.workflow.compiler import default_business_draft

            draft = default_business_draft(
                "combine_sheets", "custom:api-combine", "接口组合流程"
            ).to_dict()
            result = api.save_business_workflow(draft)
            self.assertFalse(result["ok"])
            self.assertIn("不保存自定义 DAG", result["errors"][0])

    def test_custom_dag_persistence_is_disabled(self) -> None:
        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            definition = WorkflowDefinition(
                workflow_id="custom:delete-test",
                name="待删除组合流程",
                nodes=[WorkflowNode("combine", "excel.combine_sheets")],
            )
            result = api.save_dag_workflow(definition.to_dict())
            self.assertFalse(result["ok"])
            self.assertIn("不保存自定义 DAG", result["errors"][0])

    def test_old_pywebview_dialog_constants_are_supported(self) -> None:
        window = _Window("C:/审核输出")
        old_webview = SimpleNamespace(
            FOLDER_DIALOG=20,
            OPEN_DIALOG=10,
            windows=[window],
        )
        with TemporaryDirectory() as folder:
            with patch.dict(sys.modules, {"webview": old_webview}):
                api = WebApi(Path(folder))
                self.assertEqual(api.choose_folder("output"), "C:/审核输出")
                self.assertEqual(window.calls[0][0], 20)
                self.assertEqual(api.choose_file("external"), "C:/审核输出")
                self.assertEqual(window.calls[1][0], 10)

    def test_new_pywebview_file_dialog_enum_is_supported(self) -> None:
        window = _Window("C:/模板.xlsx")
        file_dialog = SimpleNamespace(FOLDER=101, OPEN=102)
        new_webview = SimpleNamespace(
            FileDialog=file_dialog,
            FOLDER_DIALOG=20,
            OPEN_DIALOG=10,
            windows=[window],
        )
        with TemporaryDirectory() as folder:
            with patch.dict(sys.modules, {"webview": new_webview}):
                api = WebApi(Path(folder))
                self.assertEqual(api.choose_folder("output"), "C:/模板.xlsx")
                self.assertEqual(window.calls[0][0], 101)
                self.assertEqual(api.choose_file("external"), "C:/模板.xlsx")
                self.assertEqual(window.calls[1][0], 102)

    def test_explicit_empty_selection_survives_pre_run_refresh(self) -> None:
        """用户取消全选不是“首次加载”，不得在刷新时恢复为全选。"""
        with TemporaryDirectory() as folder:
            root = Path(folder)
            input_dir = root / "报送"
            template_dir = root / "模板"
            input_dir.mkdir()
            template_dir.mkdir()
            row = {"path": str(input_dir / "机构A.xlsx"), "name": "机构A.xlsx", "template": "模板.xlsx"}
            api = WebApi(root)
            api.state["input"] = str(input_dir)
            api.state["templateDir"] = str(template_dir)
            api.state["sourceFiles"] = [row]
            api.set_selected_files([])
            with patch("base_audit.web_app.classify_source_files", return_value=[row]):
                api._recognize(allow_template_auto=False)
            self.assertEqual(api.state["selectedFiles"], [])
            self.assertTrue(api.state["selectedFilesInitialized"])

    def test_unpinned_output_follows_input_with_execution_results_name(self) -> None:
        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            api.update({"input": "C:/报送目录", "outputPinned": False})
            self.assertEqual(api.state["output"], str(Path("C:/报送目录") / "执行结果"))

    def test_stale_config_path_falls_back_to_release_dir(self) -> None:
        """旧布局目录被删除后，保存的配置路径应回退到程序 config，而非重建旧目录。"""
        with TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "data").mkdir()
            (root / "config").mkdir()
            stale = root / "unified" / "config" / "逐笔统计系统_配置.xlsx"
            (root / "data" / "用户设置.json").write_text(
                '{"version":1,"history_config":%s,"period_config_file":%s}'
                % (repr(str(stale)).replace("'", '"'), repr(str(stale.with_name("报表采集系统_配置.xlsx"))).replace("'", '"')),
                encoding="utf-8",
            )
            api = WebApi(root)
            self.assertEqual(api.history_path, root / "config" / "1.逐笔统计系统_配置.xlsx")
            self.assertFalse((root / "unified").exists())  # 不在旧位置凭空重建
            saved = (root / "data" / "用户设置.json").read_text(encoding="utf-8")
            self.assertNotIn("unified", saved)  # 已回写清理

    def test_renamed_default_configs_rebind_even_when_old_copies_remain(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "config").mkdir()
            api = WebApi(root)  # 先生成当前名称的默认配置
            config_dir = root / "config"
            old_paths = {
                "history_config": config_dir / "逐笔统计系统_配置.xlsx",
                "period_config_file": config_dir / "报表采集系统_配置.xlsx",
                "central_common_config": config_dir / "大集中通用配置.xlsx",
                "central_comparison_config": config_dir / "大集中执行比较_配置.xlsx",
                "central_cross_config": config_dir / "大集中本期数值核对_配置.xlsx",
                "central_forms_config": config_dir / "大集中指标比较拆分_配置.xlsx",
            }
            for field, path in old_paths.items():
                path.touch()
                setattr(api.settings, field, str(path))
            api.settings_store.save(api.settings)

            reloaded = WebApi(root)
            self.assertEqual(reloaded.history_path.name, "1.逐笔统计系统_配置.xlsx")
            self.assertEqual(Path(reloaded.state["pcConfig"]).name, "2.报表采集系统_配置.xlsx")
            self.assertEqual(Path(reloaded.state["centralCommonConfig"]).name, "3.0大集中通用配置.xlsx")
            self.assertEqual(Path(reloaded.state["centralComparisonConfig"]).name, "3.1大集中执行比较_配置.xlsx")
            self.assertEqual(Path(reloaded.state["centralCrossConfig"]).name, "3.2大集中本期数值核对_配置.xlsx")
            self.assertEqual(Path(reloaded.state["centralFormsConfig"]).name, "3.3大集中指标比较拆分_配置.xlsx")

    def test_central_config_card_is_after_workflow_card(self) -> None:
        html = (ROOT / "frontend" / "web" / "index.html").read_text(encoding="utf-8")
        central = html[
            html.index('<section class="page" id="page-central">'):
            html.index('</section><!-- /page-central -->')
        ]
        self.assertLess(
            central.index("<label>比较文件</label>"),
            central.index("<label>配置文件</label>"),
        )

    def test_collect_page_does_not_use_central_card_layout(self) -> None:
        html = (ROOT / "frontend" / "web" / "index.html").read_text(encoding="utf-8")
        collect = html[
            html.index('<section class="page" id="page-collect">'):
            html.index('</section><!-- /page-collect -->')
        ]
        self.assertNotIn("central-card-grid", collect)

    def test_reset_central_profile_keeps_custom_binding_and_writes_backup(self) -> None:
        from shutil import copy2

        with TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "config").mkdir()
            api = WebApi(root)
            custom = root / "external-config" / "我的通用配置.xlsx"
            custom.parent.mkdir()
            copy2(api.state["centralCommonConfig"], custom)
            api.state["centralCommonConfig"] = str(custom)
            api._save_settings()

            result = api.reset_central_config_profile("common")

            self.assertEqual(Path(result["centralCommonConfig"]), custom.resolve())
            self.assertTrue(custom.is_file())
            backup_dir = custom.parent / "备份文件夹"
            backups = list(backup_dir.glob("我的通用配置_备份_*.xlsx"))
            self.assertEqual(len(backups), 1)
            api.reset_central_config_profile("common")
            self.assertEqual(len(list(backup_dir.glob("我的通用配置_备份_*.xlsx"))), 2)

    def test_saved_config_defaults_restore_complete_workbooks(self) -> None:
        """“设为默认”保存完整数据，重置不再只恢复空表头。"""
        from openpyxl import load_workbook

        def write_marker(path: Path, value: str, sheet_name: str | None = None) -> None:
            book = load_workbook(path)
            try:
                book[sheet_name or book.sheetnames[0]]["Z1"] = value
                book.save(path)
            finally:
                book.close()

        def marker(path: Path, sheet_name: str | None = None) -> str:
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                return book[sheet_name or book.sheetnames[0]]["Z1"].value
            finally:
                book.close()

        with TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "config").mkdir()
            api = WebApi(root)
            actions = [
                ("summary", api.history_path, api.set_summary_config_default, api.reset_summary_config),
                ("period", Path(api.state["pcConfig"]), api.set_period_config_default, api.reset_period_config),
                ("common", Path(api.state["centralCommonConfig"]), lambda: api.set_central_config_default("common"), lambda: api.reset_central_config_profile("common")),
                ("comparison", Path(api.state["centralComparisonConfig"]), lambda: api.set_central_config_default("comparison"), lambda: api.reset_central_config_profile("comparison")),
                ("cross", Path(api.state["centralCrossConfig"]), lambda: api.set_central_config_default("cross"), lambda: api.reset_central_config_profile("cross")),
                ("forms", Path(api.state["centralFormsConfig"]), lambda: api.set_central_config_default("forms"), lambda: api.reset_central_config_profile("forms")),
            ]
            for key, path, save_default, reset in actions:
                expected = f"{key}-default"
                marker_sheet = "通用设置" if key == "period" else None
                write_marker(path, expected, marker_sheet)
                save_default()
                self.assertTrue(Path(api.state["configDefaults"][key]).is_file())
                write_marker(path, f"{key}-changed", marker_sheet)
                reset()
                self.assertEqual(marker(path, marker_sheet), expected)

    def test_reset_config_is_noop_after_global_config_retired(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "config").mkdir()
            api = WebApi(root)
            self.assertNotIn("config", api.state["configIssues"])
            result = api.reset_config()
            self.assertEqual(result["status"], "无需重置通用配置")
            self.assertEqual(api.state["status"], "无需重置通用配置")
            self.assertFalse((root / "config" / "config.xlsx").exists())

    def test_reset_summary_and_period_configs_rebuild_structure(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "config").mkdir()
            api = WebApi(root)
            api.reset_summary_config()
            api.reset_period_config()
            issues = api.state["configIssues"]
            self.assertEqual(issues["summary"], "")
            self.assertEqual(issues["period"], "")
            self.assertTrue((root / "config" / "1.逐笔统计系统_配置.xlsx").is_file())
            period_config = root / "config" / "2.报表采集系统_配置.xlsx"
            self.assertTrue(period_config.is_file())
            period_book = load_workbook(period_config, read_only=True, data_only=True)
            try:
                self.assertIn("通用设置", period_book.sheetnames)
            finally:
                period_book.close()

    def test_period_reset_backs_up_existing_config(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "config").mkdir()
            api = WebApi(root)
            api.reset_period_config()
            api.reset_period_config()  # 第二次应先备份
            backups = list((root / "config" / "备份文件夹").glob("2.报表采集系统_配置_备份_*.xlsx"))
            self.assertEqual(len(backups), 2)
            self.assertFalse((root / "config" / "2.报表采集系统_单位配置.xlsx").exists())

    def test_period_compare_paths_and_output_are_persisted_separately(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            current = root / "本期"
            previous = root / "上期"
            output = root / "跨期结果"
            current.mkdir()
            previous.mkdir()
            central = root / "集中系统.xlsx"
            central.touch()
            window = _Window(str(current))
            webview = SimpleNamespace(FOLDER_DIALOG=20, OPEN_DIALOG=10, windows=[window])
            with patch.dict(sys.modules, {"webview": webview}):
                api = WebApi(root)
                api.choose_period_compare("pcCurDir")
                self.assertEqual(api.state["pcOutput"], str(current / "执行结果"))
                window.selected = str(previous)
                api.choose_period_compare("pcPreDir")
                window.selected = str(central)
                api.choose_period_compare("pcCentral")
                window.selected = str(output)
                api.choose_period_compare("pcOutput")

            reloaded = WebApi(root)
            self.assertEqual(reloaded.state["pcCurDir"], str(current.resolve()))
            self.assertEqual(reloaded.state["pcPreDir"], str(previous.resolve()))
            self.assertEqual(reloaded.state["pcCentral"], str(central.resolve()))
            self.assertEqual(reloaded.state["pcOutput"], str(output.resolve()))
            self.assertFalse(reloaded.state["pcOutputAuto"])

    def test_central_selections_are_persisted_separately(self) -> None:
        """大集中统计系统的选择须写入用户设置，第二次打开无需重新选择。"""
        with TemporaryDirectory() as folder:
            root = Path(folder)
            selections = {
                "centralCur": root / "本期.csv",
                "centralPre": root / "上期.csv",
                "centralOutput": root / "比较输出",
                "centralCompareFile": root / "比较结果.xlsx",
                "centralCrossCur": root / "日月报比较.csv",
                "centralTemplate": root / "金融表单模板.xlsx",
                "centralFormOutput": root / "表单输出",
                "centralCommonConfig": root / "通用配置.xlsx",
                "centralComparisonConfig": root / "执行比较配置.xlsx",
                "centralCrossConfig": root / "日月报配置.xlsx",
                "centralFormsConfig": root / "指标拆分配置.xlsx",
            }
            for kind, path in selections.items():
                if kind in {"centralOutput", "centralFormOutput"}:
                    path.mkdir()
                else:
                    path.touch()
            window = _Window("")
            webview = SimpleNamespace(FOLDER_DIALOG=20, OPEN_DIALOG=10, windows=[window])
            with patch.dict(sys.modules, {"webview": webview}):
                api = WebApi(root)
                for kind, path in selections.items():
                    window.selected = str(path)
                    api.choose_central(kind)

            reloaded = WebApi(root)
            for kind, path in selections.items():
                self.assertEqual(reloaded.state[kind], str(path.resolve()))

    def test_period_output_pin_switches_between_fixed_and_follow_current_dir(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            current = root / "本期"
            current.mkdir()
            api = WebApi(root)
            api.state["pcCurDir"] = str(current)
            api.state["pcOutput"] = str(current / "执行结果")
            api.state["pcOutputAuto"] = True
            self.assertTrue(api.toggle_period_output_pin())
            self.assertFalse(api.state["pcOutputAuto"])
            self.assertFalse(api.toggle_period_output_pin())
            self.assertTrue(api.state["pcOutputAuto"])
            self.assertEqual(api.state["pcOutput"], str(current / "执行结果"))

    def test_clear_path_forgets_persisted_selection_without_deleting_files(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            selected_dir = root / "报送目录"
            selected_file = root / "外部文件.xlsx"
            selected_dir.mkdir()
            selected_file.touch()
            api = WebApi(root)
            api.state["input"] = str(selected_dir)
            api.state["external"] = str(selected_file)
            api.state["pcCurDir"] = str(selected_dir)
            api._save_settings()

            api.clear_path("input")
            api.clear_path("external")
            api.clear_path("pcCurDir")

            self.assertEqual(api.state["input"], "")
            self.assertEqual(api.state["external"], "")
            self.assertEqual(api.state["pcCurDir"], "")
            self.assertTrue(selected_dir.is_dir())
            self.assertTrue(selected_file.is_file())
            reloaded = WebApi(root)
            self.assertEqual(reloaded.state["input"], "")
            self.assertEqual(reloaded.state["external"], "")
            self.assertEqual(reloaded.state["pcCurDir"], "")

    def test_clear_path_supports_central_selections(self) -> None:
        """大集中页的数据路径与四册配置绑定都可单独清空。"""
        with TemporaryDirectory() as folder:
            root = Path(folder)
            data_file = root / "本期.csv"
            data_file.touch()
            api = WebApi(root)
            api.state["centralCur"] = str(data_file)
            config_fields = (
                "centralCommonConfig", "centralComparisonConfig",
                "centralCrossConfig", "centralFormsConfig",
            )
            for field in config_fields:
                api.state[field] = str(root / f"{field}.xlsx")
            api._save_settings()

            api.clear_path("centralCur")
            for field in config_fields:
                api.clear_path(field)

            self.assertEqual(api.state["centralCur"], "")
            for field in config_fields:
                self.assertEqual(api.state[field], "")
            self.assertTrue(data_file.is_file())
            reloaded = WebApi(root)
            self.assertEqual(reloaded.state["centralCur"], "")
            self.assertEqual(reloaded.settings.central_common_config, "")
            self.assertEqual(reloaded.settings.central_comparison_config, "")
            self.assertEqual(reloaded.settings.central_cross_config, "")
            self.assertEqual(reloaded.settings.central_forms_config, "")

    def test_calculation_engine_setting_is_persisted_in_state(self) -> None:
        """按当前平台取一个合法引擎值，验证持久化与非法值被拒。"""
        from base_audit.engines import valid_engine_values

        engines = sorted(valid_engine_values())
        # “自动”始终可用；再取一个平台专属引擎（Windows: Excel/WPS；Linux: LibreOffice）。
        candidate = next((name for name in engines if name != "自动"), "自动")
        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            api.update({"calculationEngine": candidate})
            self.assertEqual(candidate, api.state["calculationEngine"])
            api.update({"calculationEngine": "不存在的引擎"})
            self.assertEqual(candidate, api.state["calculationEngine"])

    def test_confirm_before_run_defaults_on_and_toggles(self) -> None:
        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            # 默认开启：误点主按钮先看到文件清单。
            self.assertTrue(api.state["confirmBeforeRun"])
            api.update({"confirmBeforeRun": False})
            self.assertFalse(api.state["confirmBeforeRun"])
            # 重启后保持关闭（写入用户设置）。
            reloaded = WebApi(Path(folder))
            self.assertFalse(reloaded.state["confirmBeforeRun"])

    def test_ui_theme_defaults_and_persists(self) -> None:
        """界面模式默认日间；夜间写入用户设置重启保持；非法值被拒；旧值迁移；恢复默认回日间。"""
        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            self.assertEqual(api.state["uiTheme"], "日间")
            api.update({"uiTheme": "夜间"})
            self.assertEqual(api.state["uiTheme"], "夜间")
            api.update({"uiTheme": "不存在的颜色"})
            self.assertEqual(api.state["uiTheme"], "夜间")
            reloaded = WebApi(Path(folder))
            self.assertEqual(reloaded.state["uiTheme"], "夜间")
            api.reset_general_settings()
            self.assertEqual(api.state["uiTheme"], "日间")

    def test_ui_theme_legacy_values_are_migrated(self) -> None:
        """旧三主题值在加载时迁移：浅色→日间，深色/纯黑→夜间。"""
        import json

        with TemporaryDirectory() as folder:
            root = Path(folder)
            data_dir = root / "data"
            data_dir.mkdir()
            for legacy, expected in (("浅色", "日间"), ("深色", "夜间"), ("纯黑", "夜间")):
                (data_dir / "用户设置.json").write_text(
                    json.dumps({"ui_theme": legacy}, ensure_ascii=False), encoding="utf-8"
                )
                api = WebApi(root)
                self.assertEqual(api.state["uiTheme"], expected, legacy)

    def test_file_picker_mode_persists_and_rejects_invalid_value(self) -> None:
        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            self.assertEqual(api.state["filePickerMode"], "系统原生")
            api.update({"filePickerMode": "浏览器内置"})
            self.assertEqual(api.state["filePickerMode"], "浏览器内置")
            api.update({"filePickerMode": "不存在"})
            self.assertEqual(api.state["filePickerMode"], "浏览器内置")
            self.assertEqual(WebApi(Path(folder)).state["filePickerMode"], "浏览器内置")
            api.reset_general_settings()
            self.assertEqual(api.state["filePickerMode"], "系统原生")

    def test_conditional_format_evaluator_and_reader_persist(self) -> None:
        """判定方式默认 Python 规则求值；三值持久化；旧值映射；读取方式正交。"""
        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            self.assertEqual(api.state["conditionalFormatEvaluator"], "PYTHON")
            self.assertEqual(api.state["conditionalFormatRuleReader"], "DIRECT_OOXML")
            # 该标志 = pipeline_kind()=="com"，仅 Windows 为真（UOS 上 COM 选项按设计禁用）。
            self.assertEqual(api.state["conditionalFormatNativeAvailable"], sys.platform == "win32")
            api.update({"conditionalFormatEvaluator": "PYTHON"})
            self.assertEqual(api.state["conditionalFormatEvaluator"], "PYTHON")
            # 旧口径 OOXML 映射为 PYTHON
            api.update({"conditionalFormatEvaluator": "OOXML"})
            self.assertEqual(api.state["conditionalFormatEvaluator"], "PYTHON")
            api.update({"conditionalFormatEvaluator": "AUTO"})
            self.assertEqual(api.state["conditionalFormatEvaluator"], "PYTHON")
            # 读取方式与判定方式正交：可独立切换
            api.update({"conditionalFormatRuleReader": "OPENPYXL"})
            self.assertEqual(api.state["conditionalFormatRuleReader"], "OPENPYXL")
            self.assertEqual(api.state["conditionalFormatEvaluator"], "PYTHON")
            api.update({"conditionalFormatRuleReader": "BOGUS"})
            self.assertEqual(api.state["conditionalFormatRuleReader"], "OPENPYXL")
            reloaded = WebApi(Path(folder))
            self.assertEqual(reloaded.state["conditionalFormatEvaluator"], "PYTHON")
            self.assertEqual(reloaded.state["conditionalFormatRuleReader"], "OPENPYXL")
            api.reset_general_settings()
            self.assertEqual(api.state["conditionalFormatEvaluator"], "PYTHON")
            self.assertEqual(api.state["conditionalFormatRuleReader"], "DIRECT_OOXML")

    def test_reset_general_settings_keeps_other_system_settings(self) -> None:
        """General-page reset must not erase period/record-level preferences."""
        with TemporaryDirectory() as folder:
            root = Path(folder)
            api = WebApi(root)
            api.state.update({
                "calculationEngine": "WPS 表格",
                "recursiveDepth": 2,
                "writeFlowLogs": True,
                "showRunDetailLogs": True,
                "exportRunLogs": True,
                "showPerformanceDiagnostics": True,
                "confirmBeforeRun": False,
                "input": str(root / "逐笔统计输入"),
                "pcCurDir": str(root / "报表采集本期"),
                "pcPreDir": str(root / "报表采集上期"),
                "pcConfig": str(root / "config" / "报表采集系统_配置.xlsx"),
            })
            api._save_settings()

            api.reset_general_settings()

            self.assertEqual("自动", api.state["calculationEngine"])
            self.assertEqual(-1, api.state["recursiveDepth"])
            self.assertFalse(api.state["writeFlowLogs"])
            self.assertFalse(api.state["showRunDetailLogs"])
            self.assertFalse(api.state["exportRunLogs"])
            self.assertFalse(api.state["showPerformanceDiagnostics"])
            self.assertTrue(api.state["confirmBeforeRun"])
            self.assertEqual(str(root / "逐笔统计输入"), api.state["input"])
            self.assertEqual(str(root / "报表采集本期"), api.state["pcCurDir"])
            self.assertEqual(str(root / "报表采集上期"), api.state["pcPreDir"])
            self.assertEqual(str(root / "config" / "报表采集系统_配置.xlsx"), api.state["pcConfig"])

            reloaded = WebApi(root)
            self.assertEqual("自动", reloaded.state["calculationEngine"])
            self.assertEqual(-1, reloaded.state["recursiveDepth"])
            self.assertFalse(reloaded.state["writeFlowLogs"])
            self.assertFalse(reloaded.state["showRunDetailLogs"])
            self.assertFalse(reloaded.state["exportRunLogs"])
            self.assertFalse(reloaded.state["showPerformanceDiagnostics"])
            self.assertTrue(reloaded.state["confirmBeforeRun"])
            self.assertEqual(str(root / "报表采集本期"), reloaded.state["pcCurDir"])

    def test_history_page_is_paginated_and_filterable(self) -> None:
        with TemporaryDirectory() as folder:
            root = Path(folder)
            config = root / "config" / HISTORY_WORKBOOK_NAME
            initialize_config(config)
            from openpyxl import load_workbook
            workbook = load_workbook(config)
            sheet = workbook[HISTORY_AUDIT_SHEET]
            sheet.append(("机构A.xlsx", "贷款表", "A1", "错误", "余额", "余额异常", "10", "8", "2", "规则A", "说明A", "通过"))
            sheet.append(("机构B.xlsx", "贷款表", "B2", "软性", "利率", "利率核实", "", "", "", "规则B", "说明B", "待核实"))
            workbook.save(config)
            workbook.close()
            api = WebApi(root)
            result = api.get_history_page({"page": 1, "pageSize": 10, "keyword": "余额", "errorType": "错误"})
        self.assertEqual(1, result["total"])
        self.assertEqual("机构A.xlsx", result["items"][0]["工作簿名"])
        self.assertIn("软性", result["errorTypes"])
        self.assertIn("通过", result["opinions"])

    def test_flow_start_logs_before_background_worker_for_win7_compatibility(self) -> None:
        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            with patch("base_audit.web_app.threading.Thread") as thread:
                self.assertTrue(api.start("flow:汇总核查表校验"))
            last = api.state["log"][-1]
            self.assertIn("开始执行：汇总核查表校验", last["text"])
            self.assertFalse(last["detail"])
            thread.return_value.start.assert_called_once()

    def test_build_diagnostic_summary_covers_all_sections(self) -> None:
        with TemporaryDirectory() as folder:
            api = WebApi(Path(folder))
            api._log("诊断摘要冒烟")
            text = api.build_diagnostic_summary()["text"]
        # 报障定位五段齐全：产品/环境/配置/非默认设置/运行日志。
        for section in ("【产品】", "【环境】", "【配置】", "【非默认设置】", "【运行日志】"):
            self.assertIn(section, text)
        # 六册配置逐条报版本或缺失；新建目录中配置由启动逻辑生成，应报出版本。
        self.assertIn("1.逐笔统计系统：", text)
        self.assertIn("配置检查：", text)
        # 运行日志段带本次打入的条目；默认设置下非默认段为"无"。
        self.assertIn("诊断摘要冒烟", text)
        self.assertIn("  无", text)
        with TemporaryDirectory() as folder2:
            api2 = WebApi(Path(folder2))
            api2.update({"showRunDetailLogs": True})
            text2 = api2.build_diagnostic_summary()["text"]
        self.assertIn("show_run_detail_logs = True（默认 False）", text2)


if __name__ == "__main__":
    unittest.main()
