from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from s1audit_configured.formula_error_scan import scan_validation_formula_errors


def test_scan_only_reports_error_values_at_template_validation_formulas(tmp_path):
    template_path = tmp_path / "template.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "校验表"
    sheet["A1"] = "=1/0"
    sheet["A2"] = "=1+1"
    sheet["B1"] = "=1/0"  # Outside 校验区域.
    template.defined_names.add(DefinedName("校验区域_001", attr_text="'校验表'!$A$1:$A$2"))
    template.save(template_path)
    template.close()

    source_path = tmp_path / "source.xlsx"
    source = Workbook()
    sheet = source.active
    sheet.title = "校验表"
    sheet["A1"] = "#DIV/0!"
    sheet["A1"].data_type = "e"
    sheet["A2"] = 2
    sheet["B1"] = "#REF!"
    sheet["B1"].data_type = "e"
    source.save(source_path)
    source.close()

    result = scan_validation_formula_errors(source_path, [template_path])

    assert result["template_count"] == 1
    assert result["formula_cells_scanned"] == 2
    assert result["missing_sheets"] == []
    assert len(result["errors"]) == 1
    assert {key: result["errors"][0][key] for key in
            ("sheet", "cell", "value", "formula", "source", "origin")} == {
        "sheet": "校验表", "cell": "A1", "value": "#DIV/0!", "formula": "=1/0",
        "source": str(source_path.resolve()), "origin": "source_cached_value",
    }


def test_scan_reports_static_template_error_references_but_ignores_quoted_text(tmp_path):
    template_path = tmp_path / "template-static-error.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "校验表"
    sheet["A1"] = '=IF(B1="#N/A",#REF!,0)'
    sheet["A2"] = '="#DIV/0!"'
    template.defined_names.add(DefinedName("校验区域_001", attr_text="'校验表'!$A$1:$A$2"))
    template.save(template_path)
    template.close()

    source_path = tmp_path / "source-static-error.xlsx"
    source = Workbook()
    source.active.title = "校验表"
    source.save(source_path)
    source.close()

    result = scan_validation_formula_errors(source_path, [template_path])

    assert [(error["sheet"], error["cell"], error["value"], error["origin"])
            for error in result["errors"]] == [("校验表", "A1", "#REF!", "template_formula")]


def test_formula_error_token_scan_ignores_error_text_inside_string_literals():
    from s1audit_configured.formula_error_scan import formula_error_tokens

    assert formula_error_tokens('=IF(A1="#N/A", "#REF!", 0)') == []
    assert formula_error_tokens('=IF(A1="#N/A", #REF!, 0)') == ["#REF!"]


def test_scan_reports_skipped_when_no_layout_template_was_used(tmp_path):
    source_path = tmp_path / "source.xlsx"
    Workbook().save(source_path)

    result = scan_validation_formula_errors(source_path, [])

    assert result["status"] == "SKIPPED"
    assert result["reason"]
    assert result["formula_cells_scanned"] == 0
