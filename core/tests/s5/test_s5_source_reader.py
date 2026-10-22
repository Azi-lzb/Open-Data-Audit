from datetime import date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName

from s1audit_configured import source_reader as readers
from s1audit_configured.layout import named_regions


def fixture_workbook(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "O'Brien"
    # Leading empty rows and columns must not shift absolute coordinates.
    sheet["C5"] = "统一社会信用代码"
    sheet.merge_cells("C5:D5")
    sheet["D5"].number_format = "0.00"
    sheet["E5"] = "DEMO001"
    for row, value in enumerate([7, 7.5, date(2026, 1, 2), datetime(2026, 1, 2, 15, 30),
                                  True, "", " 空 文本 ", "#N/A", "=E6+1", 9007199254740993,
                                  time(12, 30), timedelta(days=2, hours=3)], 6):
        sheet.cell(row, 5).value = value
    sheet["E17"].number_format = "[h]:mm:ss"
    sheet["F6"] = "=E6"
    sheet["F7"] = "=E7"
    sheet["F8"] = "=1/0"
    sheet["G6"] = "=1"
    sheet["G7"] = 7
    sheet["G8"] = '=IF(1,"","")'
    sheet["H22"].comment = Comment("comment without stored cell", "test")
    sheet["H23"].hyperlink = "https://example.test/empty-cell"
    sheet.conditional_formatting.add("E6:E7 G6:G7", FormulaRule(formula=["E6>0"], fill=PatternFill("solid", fgColor="FFFF0000")))
    sheet.conditional_formatting.add("E6:E8", CellIsRule(operator="greaterThan", formula=["1"]))
    sheet.conditional_formatting.add("E9:E10", ColorScaleRule(start_type="min", start_color="FF0000", end_type="max", end_color="00FF00"))
    book.defined_names.add(DefinedName("表结构区域_001", attr_text="'O''Brien'!$C$5:$E$5"))
    book.defined_names.add(DefinedName("校验区域_001", attr_text="'O''Brien'!$F$6:$F$8"))
    sheet.defined_names.add(DefinedName("数字区域_001", attr_text="'O''Brien'!$E$6:$E$17", localSheetId=0))
    book.create_sheet("空表")
    book.save(path)
    book.close()
    # Add cached values, a shared-formula follower, an explicit empty string,
    # and a numeric spelling whose float type must survive normalization.
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
    ns = readers.NS
    cells = {node.get("r"): node for node in root.iter(ns + "c")}
    cells["E14"].find(ns + "v").text = "8"
    for address, text in [("F6", "7"), ("F7", "7.5")]:
        cells[address].find(ns + "v").text = text
        formula = cells[address].find(ns + "f")
        formula.set("t", "shared")
        formula.set("si", "0")
        if address == "F6":
            formula.set("ref", "F6:F7")
        else:
            formula.text = None
    cells["F8"].set("t", "e")
    cells["F8"].find(ns + "v").text = "#DIV/0!"
    cells["G6"].find(ns + "v").text = "1"
    cells["G7"].find(ns + "v").text = "7.0"
    cells["E15"].find(ns + "v").text = "9007199254740993"
    cells["G8"].set("t", "str")
    for row in root.find(ns + "sheetData"):
        for cell in list(row):
            if cell.get("r") in {"H22", "H23"}:
                row.remove(cell)
    inline = ET.SubElement(cells["E11"], ns + "is")
    ET.SubElement(inline, ns + "t").text = ""
    entries["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, value in entries.items():
            archive.writestr(name, value)
    return path


def test_mixed_reader_matches_rich_cells_types_and_structures(tmp_path):
    path = fixture_workbook(tmp_path / "fixture.xlsx")
    with readers.SourceReader(path) as reader:
        assert reader.engine == "mixed", reader.fallback_reason
        for cached in (False, True):
            expected = load_workbook(path, data_only=cached)
            actual = reader.book(data_only=cached)
            assert actual.sheetnames == expected.sheetnames
            assert named_regions(actual, ("表结构区域", "数字区域")) == named_regions(expected, ("表结构区域", "数字区域"))
            for sheet in expected:
                fast = actual[sheet.title]
                assert (fast.max_row, fast.max_column) == (sheet.max_row, sheet.max_column)
                assert sorted(str(r) for r in fast.merged_cells.ranges) == sorted(str(r) for r in sheet.merged_cells.ranges)
                for row in sheet:
                    for cell in row:
                        target = fast.cell(cell.row, cell.column)
                        assert (target.value, type(target.value), target.data_type, target.number_format) == (cell.value, type(cell.value), cell.data_type, cell.number_format), cell.coordinate
                expected_cfs = [(str(cf.sqref), [ET.tostring(rule.to_tree()) for rule in cf.rules]) for cf in sheet.conditional_formatting]
                actual_cfs = [(str(cf.sqref), [ET.tostring(rule.to_tree()) for rule in cf.rules]) for cf in fast.conditional_formatting]
                assert actual_cfs == expected_cfs
                for cf in fast.conditional_formatting:
                    for rule in cf.rules:
                        if rule.dxfId is not None:
                            assert rule.dxf.fill.fgColor.rgb == "FFFF0000"
            expected.close()
        meta = reader.metadata["O'Brien"]
        assert set(meta.formulas) == {(14, 5), (6, 6), (7, 6), (8, 6), (6, 7), (8, 7)}
        assert {cell for cell, kind in meta.types.items() if kind == "e"} == {(13, 5), (8, 6)}
        assert reader.book()["O'Brien"]["E15"].value == 9007199254740993


@pytest.mark.parametrize("failure", ["missing", "file_error"])
def test_fallback_is_per_file_and_matches_rich(tmp_path, monkeypatch, failure):
    path = fixture_workbook(tmp_path / "fallback.xlsx")
    if failure == "missing":
        monkeypatch.setattr(readers, "CalamineWorkbook", None)
    else:
        def bad_file(_stream):
            raise RuntimeError("injected single-file failure")
        monkeypatch.setattr(readers, "CalamineWorkbook", SimpleNamespace(from_filelike=bad_file))
    with pytest.warns(RuntimeWarning, match="fallback.xlsx.*openpyxl"):
        with readers.SourceReader(path) as reader:
            assert reader.engine == "openpyxl"
            assert reader.fallback_reason
            assert reader.book()["O'Brien"]["E14"].data_type == "f"
            assert reader.book(data_only=True)["O'Brien"]["E14"].value == 8


def test_scope_reuses_single_parse_and_invalidates_changed_files(tmp_path, monkeypatch):
    path = fixture_workbook(tmp_path / "scoped.xlsx")
    calls = []
    original = readers.SourceReader
    def traced(path):
        calls.append(path)
        return original(path)
    monkeypatch.setattr(readers, "SourceReader", traced)
    with readers.reader_session():
        first = readers.get_reader(path)
        assert readers.get_reader(path) is first
        readers.read_book(path).close()
        with readers.reader_session():
            assert readers.get_reader(path) is first
            assert readers.read_book(path, data_only=True)["O'Brien"]["E14"].value == 8
        assert len(readers.reader_diagnostics()) == 1
        assert len(calls) == 1
        path.write_bytes(path.read_bytes() + b" ")
        assert readers.get_reader(path) is not first
        assert len(calls) == 2
    assert readers.reader_diagnostics() == []


def test_source_value_rejects_formula_and_error_and_uses_merge_anchor(tmp_path):
    from s1audit_configured.sources import _Workbook
    from s1audit_configured.model import SourceBinding
    from s1audit_configured.errors import RetrievalError
    path = fixture_workbook(tmp_path / "semantics.xlsx")
    workbook = _Workbook(SourceBinding("CUR", path, "测试", "20260831", "DEMO001"))
    try:
        for coordinate, error in [("E14", "SOURCE_FORMULA_UNVERIFIED"), ("E13", "VALUE_TYPE_INVALID")]:
            row, col = readers.coordinate_to_tuple(coordinate)
            with pytest.raises(RetrievalError) as caught:
                workbook.value("O'Brien", row, col)
            assert caught.value.code == error
        assert workbook.value("O'Brien", 5, 4)[1] == "C5"
    finally:
        workbook.close()


def test_four_consumers_share_one_parse_and_legacy_parity(tmp_path, monkeypatch):
    from s1audit_configured.conditional_audit import extract_conditional_triggers
    from s1audit_configured.formula_error_scan import scan_validation_formula_errors
    from s1audit_configured.institutions import InstitutionCatalog, resolve_institution
    from s1audit_configured.model import SourceBinding
    from s1audit_configured.sources import _Workbook
    path = fixture_workbook(tmp_path / "shared.xlsx")
    outcomes = []
    for engine in ("mixed", "openpyxl"):
        monkeypatch.setenv("AUDIT3_READER_ENGINE", engine)
        with readers.reader_session():
            identity = resolve_institution(path, InstitutionCatalog())
            workbook = _Workbook(SourceBinding("CUR", path))
            assert workbook.value("O'Brien", 6, 5)[0] == 7
            workbook.close()
            conditional = extract_conditional_triggers(path)
            errors = scan_validation_formula_errors(path, [path])
            outcomes.append((identity, conditional, errors))
            diagnostics = readers.reader_diagnostics()
            assert len(diagnostics) == 1
            assert diagnostics[0]["file_reads"] == 1
            assert diagnostics[0]["engine"] == engine
    assert outcomes[0] == outcomes[1]
