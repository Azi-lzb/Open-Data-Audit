from __future__ import annotations

import hashlib
import json
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from s1audit_configured import AuditEngine, ConfigurationError, RuntimeContext, SourceBinding, load_configuration
from s1audit_configured.dsl import evaluate_expression, parse_dictionary, validate_dictionary_dependencies
from s1audit_configured.model import EvalValue
from s1audit_configured import sources as configured_sources


RULE_HEADERS = [
    "规则ID", "规则修订", "是否启用", "业务类型", "默认数据源", "默认Sheet", "规则名称", "审核级别",
    "校验类型", "校验指标", "审核表达式", "表达式布尔值取反", "规则描述", "当前值表达式", "对比值表达式", "差值表达式",
    "取值字典", "生效起始日", "生效截止日", "适用范围", "适用表达式",
]
METRIC_HEADERS = [
    "业务类型", "工作表名称", "行路径", "列路径", "行表头标题", "数据类型", "指标类型", "来源单位",
    "标准单位", "存储尺度", "允许空值", "精度", "口径依据", "确认状态", "布局版本",
]
EXT_HEADERS = [
    "数据源ID", "指标表头", "数据类型", "指标类型", "来源单位", "标准单位", "存储尺度",
    "允许空值", "精度", "口径依据", "确认状态",
]
SOURCE_HEADERS = [
    "数据源ID", "来源类型", "默认Sheet", "机构Sheet", "机构代码标签", "机构名称标签", "期间标签",
    "映射Sheet", "映射表头行", "映射数据起始行", "映射机构字段", "映射目标字段", "数据表头行",
    "数据起始行", "数据键列", "指标表头行", "指标起始列", "报表项目标签", "期望期间", "契约版本",
    "期间必须", "机构必须",
]
VERSION_HEADERS = ["配置版本号", "状态", "发布日期", "生效起始日", "生效截止日", "发布说明"]
SCOPE_HEADERS = ["适用范围ID", "业务类型", "工作表名称", "行路径前缀", "范围说明", "确认状态"]


def _write_sheet(book: Workbook, name: str, headers: list[str], rows: list[dict] = None):
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for row in rows or []:
        sheet.append([row.get(header) for header in headers])
    return sheet


def _source_workbook(path: Path, current_value: int | float, *, period="20260831", formula=False,
                     insertions=False, duplicate_row=False, missing_column=False, extra_row=False) -> Path:
    book = Workbook()
    info = book.active
    info.title = "基本信息"
    info.append(["机构代码", "I001"])
    info.append(["统计期间", period])
    sheet = book.create_sheet("贷款表")
    sheet.append(["模拟单位贷款数据"])
    sheet.append([None, None, "余额", None])
    sheet.append(["一级", "二级", "本期", "上期"])
    sheet.append(["贷款产品类别", "经营贷款", f"={current_value}" if formula else current_value, 100])
    if missing_column:
        sheet.cell(3, 3, "错误表头")
    if duplicate_row:
        sheet.append(["贷款产品类别", "经营贷款", current_value, 100])
    if extra_row:
        sheet.append(["贷款产品类别", "消费贷款", current_value + 5, 100])
    if insertions:
        # A harmless inserted row and data column prove the locator follows paths, not coordinates.
        sheet.insert_rows(4, 1)
        sheet.insert_cols(3, 1)
        sheet.cell(4, 1, "新增指标层")
        sheet.cell(4, 2, "新增指标")
        sheet.cell(3, 3, "新增指标")
        sheet.cell(5, 3, "插入列")
        sheet.merge_cells("D2:E2")
    else:
        sheet.merge_cells("C2:D2")
    book.save(path)
    book.close()
    return path


def _four_level_named_source(path: Path, current_value: int | float, *, period="20260831",
                             row_shift=0, column_shift=0, row_label="经营贷款") -> Path:
    book = Workbook()
    info = book.active
    info.title = "基本信息"
    info.append(["机构代码", "I001"])
    info.append(["统计期间", period])
    sheet = book.create_sheet("贷款表")
    sheet["A1"] = "四层表头示例"
    header_row = 5 + row_shift
    row_start_col = 1 + column_shift
    data_start_col = row_start_col + 2
    for row, label in zip(range(header_row - 3, header_row), ("余额", "贷款类别", "子类别")):
        sheet.cell(row, data_start_col, label)
        sheet.merge_cells(start_row=row, start_column=data_start_col,
                          end_row=row, end_column=data_start_col + 1)
    sheet.cell(header_row, row_start_col, "一级")
    sheet.cell(header_row, row_start_col + 1, "二级")
    sheet.cell(header_row, data_start_col, "本期")
    sheet.cell(header_row, data_start_col + 1, "上期")
    data_row = header_row + 1
    sheet.cell(data_row, row_start_col, "贷款产品类别")
    sheet.cell(data_row, row_start_col + 1, row_label)
    sheet.cell(data_row, data_start_col, current_value)
    sheet.cell(data_row, data_start_col + 1, 300)
    row_ref = f"'贷款表'!${get_column_letter(row_start_col)}${header_row}:${get_column_letter(row_start_col + 1)}${data_row}"
    col_ref = f"'贷款表'!${get_column_letter(data_start_col)}${header_row - 3}:${get_column_letter(data_start_col + 1)}${header_row}"
    book.defined_names.add(DefinedName("表结构区域_001", attr_text=row_ref))
    book.defined_names.add(DefinedName("表结构区域_002", attr_text=col_ref))
    book.save(path)
    book.close()
    return path


def _external_workbook(path: Path, value: float = 0.015, *, duplicate_mapping=False,
                       duplicate_metric=False, blank=False, duplicate_report_item=False) -> Path:
    book = Workbook()
    mapping = book.active
    mapping.title = "机构参照表"
    mapping.append(["机构代码", "报表项目"])
    mapping.append(["I001", "单位贷款"])
    if duplicate_mapping:
        mapping.append(["I001", "单位贷款"])
    data = book.create_sheet("集中系统数据")
    data.append(["外部示例"])
    data.append(["序号", "报表项目", "贷款余额(亿元)", "备用指标"])
    if duplicate_metric:
        data.cell(2, 4, "贷款余额(亿元)")
    data.append([1, "单位贷款", None if blank else value, 1])
    if duplicate_report_item:
        data.append([2, "单位贷款", value, 2])
    book.save(path)
    book.close()
    return path


def _rule(rule_id: str, expr: str, dictionary: str, *, current="", comparison="", difference="",
          description="差异={差值}万元", enabled="是", default_source="CUR", default_sheet="贷款表",
          applicable="") -> dict:
    return {
        "规则ID": rule_id, "规则修订": 1, "是否启用": enabled, "业务类型": "单位贷款",
        "默认数据源": default_source, "默认Sheet": default_sheet, "规则名称": rule_id,
        "审核级别": "错误", "校验类型": "表间核对", "校验指标": "经营贷款余额",
        "审核表达式": expr, "表达式布尔值取反": "否", "规则描述": description, "当前值表达式": current,
        "对比值表达式": comparison, "差值表达式": difference,
        "取值字典": dictionary, "生效起始日": None, "生效截止日": None,
        "适用范围": None, "适用表达式": applicable,
    }


def _config(path: Path, *, rules: list[dict] | None = None, extra_metrics: list[dict] | None = None,
            scopes: list[dict] | None = None,
            metric_type="amount", source_unit="万元", standard_unit="万元",
            storage_scale="native") -> Path:
    book = Workbook()
    book.remove(book.active)
    if rules is None:
        dictionary = '{"本期": V_CUR(("贷款产品类别", "经营贷款"), ("余额", "本期")), "上期": V_PRE(("贷款产品类别", "经营贷款"), ("余额", "本期")), "差值": {本期} - {上期}}'
        rules = [
            _rule("PERIOD.PASS", "ABS({差值}) <= 20", dictionary,
                  current="{本期}", comparison="{上期}", difference="{差值}"),
            _rule("PERIOD.TRIGGER", "ABS({差值}) <= 10", dictionary,
                  current="{本期}", comparison="{上期}", difference="{差值}"),
        ]
    _write_sheet(book, "审核规则", RULE_HEADERS, rules)
    metric_rows = [{
        "业务类型": "单位贷款", "工作表名称": "贷款表", "行路径": "贷款产品类别\n经营贷款",
        "列路径": "余额\n本期", "行表头标题": "一级\n二级", "数据类型": "number",
        "指标类型": metric_type, "来源单位": source_unit, "标准单位": standard_unit,
        "存储尺度": storage_scale,
        "允许空值": "否", "精度": 2, "口径依据": "合成测试口径", "确认状态": "是", "布局版本": "demo-1",
    }, *(extra_metrics or [])]
    _write_sheet(book, "模板指标目录", METRIC_HEADERS, metric_rows)
    if scopes is not None:
        _write_sheet(book, "适用范围目录", SCOPE_HEADERS, scopes)
    _write_sheet(book, "外部指标目录", EXT_HEADERS, [{
        "数据源ID": "EXT", "指标表头": "贷款余额(亿元)", "数据类型": "number", "指标类型": "amount",
        "来源单位": "亿元", "标准单位": "万元", "存储尺度": "native", "允许空值": "否",
        "精度": 2, "口径依据": "合成测试口径", "确认状态": "是",
    }])
    _write_sheet(book, "数据源契约", SOURCE_HEADERS, [
        {
            "数据源ID": "CUR", "来源类型": "current", "默认Sheet": "贷款表", "机构Sheet": "基本信息",
            "机构代码标签": "机构代码", "期间标签": "统计期间", "契约版本": "1",
            "期间必须": "是", "机构必须": "是",
        },
        {
            "数据源ID": "EXT", "来源类型": "external", "默认Sheet": "集中系统数据",
            "映射Sheet": "机构参照表", "映射表头行": 1, "映射数据起始行": 2,
            "映射机构字段": "机构代码", "映射目标字段": "报表项目", "数据起始行": 3,
            "数据键列": 2, "指标表头行": 2, "指标起始列": 3, "契约版本": "1",
            "期间必须": "否", "机构必须": "否",
        },
    ])
    _write_sheet(book, "配置版本", VERSION_HEADERS, [{
        "配置版本号": "demo-1", "状态": "测试", "发布日期": "2026-10-19", "发布说明": "纯合成伪配置",
    }])
    book.save(path)
    book.close()
    return path


def _context(current: Path, previous: Path | None = None, external: Path | None = None,
             *, period="20260831", allow_external_time_series=False) -> RuntimeContext:
    return RuntimeContext(
        period=period, institution_code="I001", institution_name="合成机构",
        current_sources={"CUR": SourceBinding("CUR", current, business="单位贷款", period=period, institution_code="I001")},
        previous_sources=(
            {"CUR": SourceBinding("CUR", previous, business="单位贷款", period="20250630", institution_code="I001")}
            if previous else {}
        ),
        external_sources=(
            {"EXT": SourceBinding("EXT", external, business="单位贷款", period=period, institution_code="I001")}
            if external else {}
        ),
        default_external_source="EXT" if external else "",
        allow_external_time_series=allow_external_time_series,
    )


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture(autouse=True)
def isolate_schema_snapshots(monkeypatch, tmp_path):
    monkeypatch.setattr(configured_sources, "DEFAULT_SCHEMA_SNAPSHOT_DIR", tmp_path / ".schema-snapshots")


def test_dsl_safe_parser_and_dependency_validation():
    parsed = parse_dictionary('{"B": {A} + 1, "A": ROUND(1.255, 2),}')
    validate_dictionary_dependencies(parsed)
    assert evaluate_expression("ABS({x}) <= 1", variables={"x": EvalValue.number(Decimal("-0.5"), unit="万元")}).value is True
    assert evaluate_expression("SUM(1, 2, 3) = 6").value is True
    assert evaluate_expression('ISBLANK({x})', variables={"x": EvalValue(None, "blank")}).value is True
    assert evaluate_expression("ISNUMBER({x})", variables={"x": EvalValue.number(Decimal("2"))}).value is True
    with pytest.raises(ConfigurationError):
        parse_dictionary('{"x": __import__("os")}')
    with pytest.raises(ConfigurationError):
        validate_dictionary_dependencies(parse_dictionary('{"a": {b}, "b": {a}}'))


def test_cached_dictionary_parse_returns_independent_mapping():
    source = '{"A": ROUND(1.255, 2)}'
    first = parse_dictionary(source)
    first.clear()

    second = parse_dictionary(source)
    assert list(second) == ["A"]
    validate_dictionary_dependencies(second)


def test_excel_binary64_mode_preserves_template_formula_rounding():
    # Excel converges each calculation to 15 significant digits, while the
    # default engine keeps exact Decimal arithmetic for hand-authored rules.
    assert evaluate_expression("0.1 + 0.2 = 0.3").value is True
    excel = evaluate_expression("0.1 + 0.2 = 0.3", numeric_mode="excel_binary64")
    assert excel.value is True


def test_excel_binary64_rounds_source_and_intermediate_values_but_keeps_real_differences():
    assert evaluate_expression(
        "ABS(683745.4 - 683745.3999999999) > 0",
        numeric_mode="excel_binary64",
    ).value is False
    assert evaluate_expression(
        "ABS(1500.01 - 1500) > 0",
        numeric_mode="excel_binary64",
    ).value is True
    assert evaluate_expression(
        "ABS(1501 - 1500) > 0",
        numeric_mode="excel_binary64",
    ).value is True
    assert evaluate_expression(
        "ABS(683745.4 - 683745.3999999999) > 0",
    ).value is True


def test_rule_xlsx_can_select_excel_binary64_mode(tmp_path):
    source = tmp_path / "source.xlsx"
    _source_workbook(source, 120)
    rules = [
        _rule("EXCEL.FLOAT.EQUAL", "0.1 + 0.2 = 0.3", "{}", description="Excel 尾差相等"),
        _rule("EXCEL.FLOAT.CENT", "ABS(1500.01 - 1500) = 0", "{}", description="保留真实分差"),
    ]
    path = _config(tmp_path / "excel_mode.xlsx", rules=rules)
    book = load_workbook(path)
    sheet = book["审核规则"]
    mode_column = sheet.max_column + 1
    sheet.cell(1, mode_column, "校验公式单元格")
    sheet.cell(2, mode_column, "W5")
    sheet.cell(3, mode_column, "W6")
    sheet.cell(1, mode_column + 1, "数值运算模式")
    sheet.cell(2, mode_column + 1, "Excel双精度")
    sheet.cell(3, mode_column + 1, "Excel双精度")
    book.save(path)
    book.close()

    config = load_configuration(path)
    assert all(rule.numeric_mode == "excel_binary64" for rule in config.rules)
    assert [rule.formula_cell for rule in config.rules] == ["W5", "W6"]
    results = AuditEngine(config).run(
        _context(source), schema_snapshot_dir=tmp_path / "snapshots",
    ).results
    by_rule = {outcome.rule_id: outcome for outcome in results}
    assert by_rule["EXCEL.FLOAT.EQUAL"].status == "PASSED"
    assert by_rule["EXCEL.FLOAT.CENT"].status == "TRIGGERED"
    from s1audit_configured.outputs import _location
    assert _location(by_rule["EXCEL.FLOAT.EQUAL"], set()) == ("贷款表", "W5", None)
    assert _location(by_rule["EXCEL.FLOAT.CENT"], set()) == ("贷款表", "W6", None)


def test_blank_numeric_source_uses_excel_zero_coercion_without_losing_isblank():
    blank = EvalValue(None, "blank", unit="万元", metric_type="amount")
    variables = {"x": blank}
    assert evaluate_expression("{x} / 10000 = 0", variables=variables).value is True
    assert evaluate_expression("{x} = 0", variables=variables).value is True
    assert evaluate_expression("{x} > 0", variables=variables).value is False
    assert evaluate_expression("ISBLANK({x})", variables=variables).value is True


def test_unitless_numeric_literals_in_addition_and_subtraction_inherit_measure_unit():
    amount = EvalValue.number(Decimal("12.5"), unit="万元", metric_type="amount")
    result = evaluate_expression("{amount} + 1 - 2", variables={"amount": amount})
    assert result.value == Decimal("11.5")
    assert result.unit == "万元"
    reverse = evaluate_expression("1 + {amount}", variables={"amount": amount})
    assert reverse.value == Decimal("13.5")
    assert reverse.unit == "万元"


def test_sum_of_blank_measures_preserves_unit_for_later_arithmetic():
    blank = EvalValue(None, "blank", unit="万元", metric_type="amount")
    amount = EvalValue.number(Decimal("0"), unit="万元", metric_type="amount")
    total = evaluate_expression("SUM({blank})", variables={"blank": blank})
    assert total.value == Decimal("0")
    assert total.unit == "万元"
    assert evaluate_expression("SUM({blank}) - {amount} = 0", variables={
        "blank": blank, "amount": amount,
    }).value is True


def test_config_loader_rejects_duplicate_rules_and_cycles(tmp_path):
    current = tmp_path / "current.xlsx"
    _source_workbook(current, 120)
    cycle = _config(tmp_path / "cycle.xlsx", rules=[_rule("CYCLE", "{a} = 1", '{"a": {b}, "b": {a}}')])
    with pytest.raises(ConfigurationError, match="循环引用"):
        load_configuration(cycle)

    rules = [_rule("DUP", "TRUE", "{}", description="重复项"), _rule("DUP", "TRUE", "{}", description="重复项")]
    duplicated = _config(tmp_path / "duplicate.xlsx", rules=rules)
    with pytest.raises(ConfigurationError, match="重复"):
        load_configuration(duplicated)

    undefined = _config(tmp_path / "undefined.xlsx", rules=[
        _rule("UNDEFINED", "{missing} = 1", "{}", description="未定义变量")
    ])
    with pytest.raises(ConfigurationError, match="未定义变量"):
        load_configuration(undefined)
    with pytest.raises(ConfigurationError):
        evaluate_expression("UNKNOWN(1)")


def test_engine_runs_current_previous_and_writes_auditable_outputs(tmp_path):
    current = _source_workbook(tmp_path / "current.xlsx", 120)
    previous = _source_workbook(tmp_path / "previous.xlsx", 100, period="20250630")
    config = _config(tmp_path / "rules.xlsx")
    original_hashes = {path: _sha(path) for path in (current, previous, config)}
    result = AuditEngine(config).run(_context(current, previous), output_dir=tmp_path / "outputs", run_id="synthetic-run")

    by_id = {item.rule_id: item for item in result.results}
    assert by_id["PERIOD.PASS"].status == "PASSED"
    assert by_id["PERIOD.TRIGGER"].status == "TRIGGERED"
    assert by_id["PERIOD.PASS"].current_value == Decimal("120")
    assert by_id["PERIOD.PASS"].comparison_value == Decimal("100")
    assert by_id["PERIOD.PASS"].difference_value == Decimal("20")
    assert by_id["PERIOD.PASS"].difference_precision == 2
    assert {item.source_id for item in by_id["PERIOD.PASS"].evidence} == {"CUR"}
    assert any(item.sheet == "贷款表" and item.cell == "C4" for item in by_id["PERIOD.PASS"].evidence)
    assert result.status == "COMPLETED"
    assert all(_sha(path) == digest for path, digest in original_hashes.items())

    output_file = Path(result.output_files["result_workbook"])
    ledger_file = Path(result.output_files["run_ledger"])
    assert output_file.is_file() and ledger_file.is_file()
    ledger = json.loads(ledger_file.read_text(encoding="utf-8"))
    assert ledger["input_integrity"] == "UNCHANGED"
    result_book = load_workbook(output_file, data_only=False)
    assert result_book.sheetnames == ["审核结果", "条件格式触发"]
    assert result_book["审核结果"].max_row == 2
    assert result_book["审核结果"]["G2"].value == "校验硬性错误"
    evidence_cell = result_book["审核结果"]["F2"]
    assert evidence_cell.hyperlink is not None
    assert "审核副本" in evidence_cell.hyperlink.target
    result_book.close()


def test_compact_rule_sheet_infers_single_business_and_contract_defaults(tmp_path):
    current = _source_workbook(tmp_path / "current.xlsx", 120)
    previous = _source_workbook(tmp_path / "previous.xlsx", 100, period="20250630")
    config = _config(tmp_path / "compact-rules.xlsx")
    book = load_workbook(config)
    sheet = book["审核规则"]
    old_headers = [cell.value for cell in sheet[1]]
    old_record = dict(zip(old_headers, [cell.value for cell in sheet[2]]))
    sheet.delete_rows(1, sheet.max_row)
    compact_headers = [
        "规则序号", "是否启用", "规则名称", "审核级别", "校验类型", "校验指标",
        "审核表达式", "规则描述", "当前值表达式", "对比值表达式", "差值表达式", "取值字典",
    ]
    sheet.append(compact_headers)
    sheet.append([
        old_record["规则ID"], old_record["是否启用"], old_record["规则名称"],
        old_record["审核级别"], old_record["校验类型"], old_record["校验指标"],
        old_record["审核表达式"], old_record["规则描述"], old_record["当前值表达式"],
        old_record["对比值表达式"], old_record["差值表达式"], old_record["取值字典"],
    ])
    book.save(config)
    book.close()

    loaded = load_configuration(config)
    assert loaded.rules[0].business == "单位贷款"
    assert loaded.rules[0].revision == 1
    assert loaded.rules[0].default_source == ""
    assert loaded.rules[0].default_sheet == ""
    result = AuditEngine(loaded).run(_context(current, previous), run_id="compact-config")
    assert result.results[0].status == "PASSED"
    assert result.results[0].current_value == Decimal("120")


def test_single_rule_sheet_derives_template_and_external_paths(tmp_path):
    current = _source_workbook(tmp_path / "current.xlsx", 120)
    previous = _source_workbook(tmp_path / "previous.xlsx", 100, period="20250630")
    external = _external_workbook(tmp_path / "external.xlsx", 0.015)
    current_previous = _rule(
        "SINGLE.PERIOD", "ABS({差值}) <= 20",
        '{"本期": V_CUR("贷款表", ("贷款产品类别", "经营贷款"), ("余额", "本期")), '
        '"上期": V_PRE("贷款表", ("贷款产品类别", "经营贷款"), ("余额", "本期")), '
        '"差值": {本期} - {上期}}',
        current="{本期}", comparison="{上期}", difference="{差值}",
    )
    current_external = _rule(
        "SINGLE.EXTERNAL", "ABS({差值}) <= 30",
        '{"本期": V_CUR("贷款表", ("贷款产品类别", "经营贷款"), ("余额", "本期")), '
        '"大集中": V_EXT("贷款余额(亿元)"), "差值": {本期} - {大集中}}',
        current="{本期}", comparison="{大集中}", difference="{差值}",
    )
    config = _config(tmp_path / "single-sheet.xlsx", rules=[current_previous, current_external])
    book = load_workbook(config)
    for name in list(book.sheetnames):
        if name != "审核规则":
            del book[name]
    sheet = book["审核规则"]
    sheet.delete_rows(1, sheet.max_row)
    headers = [
        "规则序号", "是否启用", "业务名称", "工作表名称", "规则名称", "审核级别",
        "校验类型", "校验指标", "审核表达式", "规则描述", "当前值表达式",
        "对比值表达式", "差值表达式", "取值字典",
    ]
    sheet.append(headers)
    for rule in (current_previous, current_external):
        sheet.append([
            rule["规则ID"], rule["是否启用"], rule["业务类型"], rule["默认Sheet"],
            rule["规则名称"], rule["审核级别"], rule["校验类型"], rule["校验指标"],
            rule["审核表达式"], rule["规则描述"], rule["当前值表达式"],
            rule["对比值表达式"], rule["差值表达式"], rule["取值字典"],
        ])
    book.save(config)
    book.close()

    loaded = load_configuration(config)
    assert len(loaded.rules) == 2
    assert len(loaded.metrics) == 1
    assert len(loaded.external_metrics) == 1
    assert loaded.source_contracts == ()
    assert loaded.version == ""

    result = AuditEngine(loaded).run(
        _context(current, previous, external), run_id="single-rule-sheet",
    )
    by_id = {item.rule_id: item for item in result.results}
    assert by_id["SINGLE.PERIOD"].status == "PASSED"
    assert by_id["SINGLE.PERIOD"].current_value == Decimal("120")
    assert by_id["SINGLE.PERIOD"].comparison_value == Decimal("100")
    assert by_id["SINGLE.EXTERNAL"].status == "PASSED"
    assert by_id["SINGLE.EXTERNAL"].comparison_value == Decimal("150")
    assert any(item.mapping_sheet == "机构参照表" and item.mapping_key == "I001"
               for item in by_id["SINGLE.EXTERNAL"].evidence)


def test_locator_survives_inserted_data_row_and_column(tmp_path):
    current = _source_workbook(tmp_path / "shifted.xlsx", 120, insertions=True)
    previous = _source_workbook(tmp_path / "previous.xlsx", 100, period="20250630")
    config = _config(tmp_path / "rules.xlsx")
    result = AuditEngine(config).run(_context(current, previous), run_id="shifted")
    for item in result.results:
        assert item.status in {"PASSED", "TRIGGERED"}
        assert item.current_value == Decimal("120")
        assert any(evidence.cell == "D5" for evidence in item.evidence)


def test_named_regions_locate_four_level_merged_columns_and_snapshot_layout(tmp_path):
    source = _four_level_named_source(tmp_path / "four-level.xlsx", 321)
    call = 'V_CUR(("贷款产品类别", "经营贷款"), ("余额", "贷款类别", "子类别", "本期"))'
    rule = _rule("FOUR.LEVEL", f"{call} >= 321", "{}", current=call, description="四层列路径")
    metric = {
        "业务类型": "单位贷款", "工作表名称": "贷款表", "行路径": "贷款产品类别\n经营贷款",
        "列路径": "余额\n贷款类别\n子类别\n本期", "行表头标题": "一级\n二级",
        "数据类型": "number", "指标类型": "amount", "来源单位": "万元", "标准单位": "万元",
        "存储尺度": "native", "允许空值": "否", "精度": 2, "口径依据": "合成测试口径",
        "确认状态": "是", "布局版本": "four-level",
    }
    config = _config(tmp_path / "four-level-config.xlsx", rules=[rule], extra_metrics=[metric])
    snapshot_dir = tmp_path / "schema-snapshots"

    first = AuditEngine(config).run(
        _context(source), run_id="four-level-baseline", schema_snapshot_dir=snapshot_dir,
    )
    assert first.results[0].status == "PASSED"
    assert first.ledger["structure_checks"][0]["status"] == "BASELINE_CREATED"
    snapshot_text = next(snapshot_dir.glob("*.json")).read_text(encoding="utf-8")
    snapshot = json.loads(snapshot_text)
    assert "经营贷款" not in snapshot_text
    def contains_source_value(value):
        if isinstance(value, dict):
            return any(contains_source_value(item) for item in value.values())
        if isinstance(value, list):
            return any(contains_source_value(item) for item in value)
        return value == 321 or value == "321"
    assert not contains_source_value(snapshot)

    shifted_source = _four_level_named_source(
        tmp_path / "four-level-shifted.xlsx", 322, row_shift=1, column_shift=1,
    )
    shifted = AuditEngine(config).run(
        _context(shifted_source), run_id="four-level-shifted", schema_snapshot_dir=snapshot_dir,
    )
    assert shifted.results[0].status == "PASSED"
    assert shifted.results[0].current_value == Decimal("322")
    assert shifted.ledger["structure_checks"][0]["status"] == "POSITION_CHANGED"

    changed_source = _four_level_named_source(
        tmp_path / "four-level-changed.xlsx", 323, row_shift=1, column_shift=1,
        row_label="经营性贷款",
    )
    changed = AuditEngine(config).run(
        _context(changed_source), run_id="four-level-changed", schema_snapshot_dir=snapshot_dir,
    )
    assert changed.ledger["structure_checks"][0]["status"] == "STRUCTURE_CHANGED"
    assert changed.results[0].error_code == "ROW_PATH_NOT_FOUND"


def test_row_path_compacts_blank_header_fields_without_losing_title_depth(tmp_path):
    current = tmp_path / "sparse-row-path.xlsx"
    book = Workbook()
    info = book.active
    info.title = "基本信息"
    info.append(["机构代码", "I001"])
    info.append(["统计期间", "20260831"])
    sheet = book.create_sheet("贷款表")
    sheet["A3"] = "字段一"
    sheet["B3"] = "字段二"
    sheet["C3"] = "字段三"
    sheet["D3"] = "字段四"
    sheet["E2"] = "余额"
    sheet["E3"] = "本期"
    sheet["A4"] = "客户类型"
    sheet["C4"] = "企业"
    sheet["D4"] = "客户"
    sheet["E4"] = 321
    book.save(current)
    book.close()

    call = 'V_CUR(("客户类型", "企业", "客户"), ("余额", "本期"))'
    rule = _rule("SPARSE.PATH", f"{call} = 321", "{}", current=call, description="稀疏行路径")
    metric = {
        "业务类型": "单位贷款", "工作表名称": "贷款表", "行路径": "客户类型\n企业\n客户",
        "列路径": "余额\n本期", "行表头标题": "字段一\n字段二\n字段三\n字段四",
        "数据类型": "number", "指标类型": "amount", "来源单位": "万元", "标准单位": "万元",
        "存储尺度": "native", "允许空值": "否", "精度": 2, "口径依据": "合成测试",
        "确认状态": "是", "布局版本": "sparse-demo",
    }
    config = _config(tmp_path / "sparse-config.xlsx", rules=[rule], extra_metrics=[metric])
    result = AuditEngine(config).run(_context(current), run_id="sparse-row-path")
    assert result.results[0].status == "PASSED"
    assert result.results[0].current_value == Decimal("321")


def test_direct_v_cur_can_override_source_and_sheet(tmp_path):
    current = _source_workbook(tmp_path / "current.xlsx", 120)
    book = load_workbook(current)
    alternate = book.create_sheet("跨业务表")
    alternate.append(["备用业务"])
    alternate.append([None, None, "余额", None])
    alternate.append(["一级", "二级", "本期", "上期"])
    alternate.append(["贷款产品类别", "经营贷款", 999, 0])
    alternate.merge_cells("C2:D2")
    book.save(current)
    book.close()

    extra_metric = {
        "业务类型": "其他业务", "工作表名称": "跨业务表", "行路径": "贷款产品类别\n经营贷款",
        "列路径": "余额\n本期", "行表头标题": "一级\n二级", "数据类型": "number",
        "指标类型": "amount", "来源单位": "万元", "标准单位": "万元", "存储尺度": "native",
        "允许空值": "否", "精度": 2, "口径依据": "合成测试口径", "确认状态": "是", "布局版本": "demo-2",
    }
    direct_call = 'V_CUR(("贷款产品类别", "经营贷款"), ("余额", "本期"), source="ALT", sheet="跨业务表")'
    rule = _rule("CROSS.SOURCE", f"{direct_call} = 999", "{}", current=direct_call,
                 description="跨业务直接取值")
    config = _config(tmp_path / "cross-source.xlsx", rules=[rule], extra_metrics=[extra_metric])
    context = _context(current)
    context.current_sources["ALT"] = SourceBinding(
        "ALT", current, business="其他业务", period="20260831", institution_code="I001"
    )
    item = AuditEngine(config).run(context, run_id="cross-source").results[0]
    assert item.status == "PASSED"
    assert item.current_value == Decimal("999")
    assert item.evidence[0].source_id == "ALT"
    assert item.evidence[0].sheet == "跨业务表"
    assert item.evidence[0].cell == "C4"


def test_v_cur_accepts_positional_sheet_before_row_and_column_paths(tmp_path):
    current = _source_workbook(tmp_path / "current.xlsx", 120)
    book = load_workbook(current)
    alternate = book.create_sheet("跨业务表")
    alternate.append(["备用业务"])
    alternate.append([None, None, "余额", None])
    alternate.append(["一级", "二级", "本期", "上期"])
    alternate.append(["贷款产品类别", "经营贷款", 999, 0])
    alternate.merge_cells("C2:D2")
    book.save(current)
    book.close()

    extra_metric = {
        "业务类型": "单位贷款", "工作表名称": "跨业务表", "行路径": "贷款产品类别\n经营贷款",
        "列路径": "余额\n本期", "行表头标题": "一级\n二级", "数据类型": "number",
        "指标类型": "amount", "来源单位": "万元", "标准单位": "万元", "存储尺度": "native",
        "允许空值": "否", "精度": 2, "口径依据": "合成测试口径", "确认状态": "是", "布局版本": "demo-2",
    }
    direct_call = 'V_CUR("跨业务表", ("贷款产品类别", "经营贷款"), ("余额", "本期"))'
    rule = _rule("CROSS.SHEET", f"{direct_call} = 999", "{}", current=direct_call,
                 description="按显式Sheet跨表取值")
    config = _config(tmp_path / "cross-sheet.xlsx", rules=[rule], extra_metrics=[extra_metric])
    item = AuditEngine(config).run(_context(current), run_id="cross-sheet").results[0]
    assert item.status == "PASSED"
    assert item.current_value == Decimal("999")
    assert item.evidence[0].sheet == "跨业务表"
    assert item.evidence[0].cell == "C4"


@pytest.mark.parametrize(
    ("source_options", "expected"),
    [({"duplicate_row": True}, "METRIC_AMBIGUOUS"),
     ({"missing_column": True}, "COLUMN_PATH_NOT_FOUND")],
)
def test_template_path_missing_or_duplicate_is_not_guessed(tmp_path, source_options, expected):
    current = _source_workbook(tmp_path / "current.xlsx", 120, **source_options)
    previous = _source_workbook(tmp_path / "previous.xlsx", 100, period="20250630")
    config = _config(tmp_path / "rules.xlsx")
    item = AuditEngine(config).run(_context(current, previous), run_id="layout-boundary").results[0]
    assert item.status == "STRUCTURE_ERROR"
    assert item.error_code == expected
    assert item.error_evidence["source_id"] == "CUR"


def test_missing_previous_is_retrieval_failed_not_pass(tmp_path):
    current = _source_workbook(tmp_path / "current.xlsx", 120)
    config = _config(tmp_path / "rules.xlsx")
    results = AuditEngine(config).run(_context(current), run_id="no-prev").results
    assert all(item.status == "RETRIEVAL_FAILED" for item in results)
    assert all(item.error_code == "SOURCE_NOT_SELECTED" for item in results)


def test_disabled_rule_is_not_executed(tmp_path):
    current = _source_workbook(tmp_path / "current.xlsx", 120)
    config = _config(tmp_path / "disabled.xlsx", rules=[
        _rule("DISABLED", "TRUE", "{}", enabled="否", description="已停用")
    ])
    item = AuditEngine(config).run(_context(current), run_id="disabled").results[0]
    assert item.status == "NOT_EXECUTED"
    assert item.triggered is None


@pytest.mark.parametrize(
    ("dictionary", "expression", "expected_code"),
    [
        ('{"x": 1 / 0}', "{x} = 1", "DIVISION_BY_ZERO"),
        ('{"x": IF(TRUE, "文本", "备用")}', "ABS({x}) >= 0", "TYPE_ERROR"),
    ],
)
def test_calculation_errors_are_not_treated_as_rule_results(tmp_path, dictionary, expression, expected_code):
    current = _source_workbook(tmp_path / "current.xlsx", 120)
    rule = _rule("CALC.ERROR", expression, dictionary, description="计算失败")
    config = _config(tmp_path / "calculation-error.xlsx", rules=[rule])
    item = AuditEngine(config).run(_context(current), run_id="calculation-error").results[0]
    assert item.status == "CALCULATION_FAILED"
    assert item.error_code == expected_code


def test_v_pre_reads_the_independent_previous_file_not_current_file_prior_column(tmp_path):
    current = _source_workbook(tmp_path / "current.xlsx", 120)
    previous = _source_workbook(tmp_path / "previous.xlsx", 90, period="20250630")
    config = _config(tmp_path / "rules.xlsx")
    result = AuditEngine(config).run(_context(current, previous), run_id="prior-source").results[0]
    assert result.current_value == Decimal("120")
    assert result.comparison_value == Decimal("90")
    assert result.difference_value == Decimal("30")
    previous_evidence = next(item for item in result.evidence if item.source_sha256 == _sha(previous))
    assert previous_evidence.source_period == "20250630"


def test_audit_expression_can_be_inverted_by_a_boolean_config_flag(tmp_path):
    current = _source_workbook(tmp_path / "current.xlsx", 120)
    rule = _rule("INVERTED", "TRUE", "{}", description="布尔反转")
    rule["表达式布尔值取反"] = "是"
    config = _config(tmp_path / "rules.xlsx", rules=[rule])
    result = AuditEngine(config).run(_context(current), run_id="inverted")
    assert result.results[0].status == "TRIGGERED"


def test_external_mapping_unit_conversion_and_zero_semantics(tmp_path):
    current = _source_workbook(tmp_path / "current.xlsx", 1500)
    external = _external_workbook(tmp_path / "external.xlsx", value=0.15)
    dictionary = '{"本期": V_CUR(("贷款产品类别", "经营贷款"), ("余额", "本期")), "大集中": V_EXT("贷款余额(亿元)"), "差值": {本期} - {大集中}}'
    rule = _rule("EXT.MATCH", "ABS({差值}) <= 0", dictionary,
                 current="{本期}", comparison="{大集中}", difference="{差值}", default_source="CUR")
    config = _config(tmp_path / "rules.xlsx", rules=[rule])
    result = AuditEngine(config).run(_context(current, external=external), run_id="external")
    item = result.results[0]
    assert item.status == "PASSED"
    assert item.current_value == Decimal("1500")
    assert item.comparison_value == Decimal("1500.00")
    assert item.difference_value == Decimal("0.00")
    ext_evidence = next(evidence for evidence in item.evidence if evidence.source_id == "EXT")
    assert ext_evidence.conversion.startswith("亿元→万元")
    assert ext_evidence.mapping_key == "I001"
    assert ext_evidence.report_item == "单位贷款"

    zero_external = _external_workbook(tmp_path / "external-zero.xlsx", value=0)
    zero_result = AuditEngine(config).run(_context(current, external=zero_external), run_id="external-zero").results[0]
    assert zero_result.status == "TRIGGERED"
    assert zero_result.comparison_value == Decimal("0")


@pytest.mark.parametrize(
    ("options", "expected"),
    [
        ({"duplicate_mapping": True}, "IDENTITY_MAPPING_AMBIGUOUS"),
        ({"duplicate_metric": True}, "METRIC_AMBIGUOUS"),
        ({"duplicate_report_item": True}, "REPORT_ITEM_DUPLICATE"),
    ],
)
def test_external_retrieval_boundaries_are_not_silently_accepted(tmp_path, options, expected):
    current = _source_workbook(tmp_path / "current.xlsx", 1500)
    external = _external_workbook(tmp_path / "external.xlsx", **options)
    dictionary = '{"大集中": V_EXT("贷款余额(亿元)")}'
    rule = _rule("EXT.BOUNDARY", "{大集中} >= 0", dictionary,
                 current="{大集中}", default_source="EXT", default_sheet="", description="边界测试")
    config = _config(tmp_path / "rules.xlsx", rules=[rule])
    item = AuditEngine(config).run(_context(current, external=external), run_id="external-failure").results[0]
    expected_status = "RETRIEVAL_FAILED" if expected == "IDENTITY_MAPPING_AMBIGUOUS" else "STRUCTURE_ERROR"
    assert item.status == expected_status
    assert item.error_code == expected


@pytest.mark.parametrize("value", [None, "误填文字"])
def test_external_numeric_blank_or_text_is_zero_with_a_quality_warning(tmp_path, value):
    current = _source_workbook(tmp_path / "current.xlsx", 1500)
    external = _external_workbook(tmp_path / "external.xlsx", blank=value is None)
    if value is not None:
        book = load_workbook(external)
        book["集中系统数据"]["C3"] = value
        book.save(external)
        book.close()
    dictionary = '{"大集中": V_EXT("贷款余额(亿元)")}'
    rule = _rule("EXT.QUALITY", "{大集中} >= 0", dictionary,
                 current="{大集中}", default_source="EXT", default_sheet="", description="异常值提示")
    config = _config(tmp_path / "rules.xlsx", rules=[rule])

    result = AuditEngine(config).run(
        _context(current, external=external, allow_external_time_series=True), run_id="external-quality",
    )

    assert result.results[0].status == "PASSED"
    assert result.results[0].current_value == Decimal("0.00")
    warnings = result.ledger["data_quality_warnings"]
    assert len(warnings) == 1
    assert warnings[0]["location"] == "C3"
    assert "已按 0 参与审核" in warnings[0]["message"]


def test_source_formula_is_unverified_and_applicability_short_circuits(tmp_path):
    formula_source = _source_workbook(tmp_path / "formula.xlsx", 120, formula=True)
    previous = _source_workbook(tmp_path / "previous.xlsx", 100, period="20250630")
    config = _config(tmp_path / "rules.xlsx")
    results = AuditEngine(config).run(_context(formula_source, previous), run_id="formula").results
    assert all(item.status == "RETRIEVAL_FAILED" for item in results)
    assert all(item.error_code == "SOURCE_FORMULA_UNVERIFIED" for item in results)
    assert all(item.error_evidence["cell"] == "C4" for item in results)

    rule = _rule("NA.SHORT", "TRUE", '{"unused": V_PRE("x", "y")}', applicable="FALSE", description="不适用")
    na_config = _config(tmp_path / "not-applicable.xlsx", rules=[rule])
    item = AuditEngine(na_config).run(_context(previous), run_id="na").results[0]
    assert item.status == "NOT_APPLICABLE"
    assert item.unused_variables == ["unused"]


def test_count_unit_stays_as_count_and_amount_only_is_converted():
    # DSL unit-aware arithmetic rejects comparing different configured units.
    with pytest.raises(Exception, match="UNIT_MISMATCH"):
        evaluate_expression("{a} = {b}", variables={
            "a": EvalValue.number(Decimal(1), unit="户"),
            "b": EvalValue.number(Decimal(1), unit="万元"),
        })
    with pytest.raises(Exception, match="UNIT_MISMATCH"):
        evaluate_expression("{a} + {b}", variables={
            "a": EvalValue.number(Decimal(1), unit="户"),
            "b": EvalValue.number(Decimal(1), unit="万元"),
        })


def test_rate_fraction_is_scaled_but_count_unit_is_not_converted(tmp_path):
    current = _source_workbook(tmp_path / "rate.xlsx", 0.035)
    rate_rule = _rule(
        "RATE.SCALE", "{比例} = 3.5",
        '{"比例": V_CUR(("贷款产品类别", "经营贷款"), ("余额", "本期"))}',
        current="{比例}", description="百分比={比例}",
    )
    rate_config = _config(tmp_path / "rate-config.xlsx", rules=[rate_rule], metric_type="rate",
                           source_unit="比例", standard_unit="百分比", storage_scale="fraction")
    rate_result = AuditEngine(rate_config).run(_context(current), run_id="rate-scale").results[0]
    assert rate_result.status == "PASSED"
    assert rate_result.current_value == Decimal("3.500")
    assert rate_result.current_unit == "百分比"
    assert rate_result.current_precision == 2
    rate_evidence = rate_result.evidence[0]
    assert rate_evidence.conversion == "比例存储→百分比:×100"

    count_source = _source_workbook(tmp_path / "count.xlsx", 57)
    count_rule = _rule(
        "COUNT.NATIVE", "{户数} = 57",
        '{"户数": V_CUR(("贷款产品类别", "经营贷款"), ("余额", "本期"))}',
        current="{户数}", description="户数={户数}户",
    )
    count_config = _config(tmp_path / "count-config.xlsx", rules=[count_rule], metric_type="count",
                            source_unit="户", standard_unit="户")
    count_result = AuditEngine(count_config).run(_context(count_source), run_id="count-native").results[0]
    assert count_result.status == "PASSED"
    assert count_result.current_value == Decimal("57")
    assert count_result.current_unit == "户"
    assert count_result.evidence[0].conversion == ""


def test_unknown_amount_unit_fails_and_row_scope_expands_over_matching_metrics(tmp_path):
    bad_unit = _config(tmp_path / "bad-unit.xlsx", source_unit="", standard_unit="")
    with pytest.raises(ConfigurationError, match="单位"):
        load_configuration(bad_unit)
    source = _source_workbook(tmp_path / "scoped-source.xlsx", 120, extra_row=True)
    scoped_rule = _rule(
        "SCOPED", 'V_CUR("@当前行", ("余额", "本期")) >= 0', "{}",
        current='V_CUR("@当前行", ("余额", "本期"))', description="按行执行",
    )
    scoped_rule["适用范围"] = "全部贷款"
    extra_metric = {
        "业务类型": "单位贷款", "工作表名称": "贷款表", "行路径": "贷款产品类别\n消费贷款",
        "列路径": "余额\n本期", "行表头标题": "一级\n二级", "数据类型": "number",
        "指标类型": "amount", "来源单位": "万元", "标准单位": "万元", "存储尺度": "native",
        "允许空值": "否", "精度": 2, "口径依据": "合成测试口径", "确认状态": "是", "布局版本": "demo-1",
    }
    scopes = [{
        "适用范围ID": "全部贷款", "业务类型": "单位贷款", "工作表名称": "贷款表",
        "行路径前缀": "贷款产品类别", "范围说明": "按产品逐行执行", "确认状态": "是",
    }]
    scoped = _config(tmp_path / "scoped.xlsx", rules=[scoped_rule], extra_metrics=[extra_metric], scopes=scopes)
    loaded = load_configuration(scoped)
    assert len(loaded.applicability_scopes) == 1
    result = AuditEngine(loaded).run(_context(source), run_id="scoped-row-expansion")
    assert len(result.results) == 2
    assert {item.status for item in result.results} == {"PASSED"}
    assert {item.scope_path for item in result.results} == {
        "贷款产品类别 / 经营贷款", "贷款产品类别 / 消费贷款",
    }


def test_current_source_period_and_institution_must_match_run_context(tmp_path):
    current = _source_workbook(tmp_path / "current.xlsx", 120)
    previous = _source_workbook(tmp_path / "previous.xlsx", 100, period="20250630")
    config = _config(tmp_path / "rules.xlsx")
    context = _context(current, previous)
    binding = context.current_sources["CUR"]
    context.current_sources = {
        "CUR": SourceBinding("CUR", current, business="单位贷款", period="20250630", institution_code="I001")
    }
    result = AuditEngine(config).run(context, run_id="period-mismatch")
    assert all(item.status == "RETRIEVAL_FAILED" and item.error_code == "PERIOD_CONTRACT_MISMATCH" for item in result.results)

    context = _context(current, previous)
    context.current_sources = {
        "CUR": SourceBinding("CUR", current, business="单位贷款", period="20260831", institution_code="I999")
    }
    result = AuditEngine(config).run(context, run_id="institution-mismatch")
    assert all(item.status == "RETRIEVAL_FAILED" and item.error_code == "INSTITUTION_MISMATCH" for item in result.results)
