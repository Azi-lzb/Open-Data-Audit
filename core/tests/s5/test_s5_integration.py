from __future__ import annotations

import copy
import hashlib
import json
import shutil
import subprocess
import threading
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName

from base_audit.systems.s5_configured_statistics.contracts import S5Request
from base_audit.systems.s5_configured_statistics.service import run_s5_f1, validate_s5_config
from base_audit.systems.s5_configured_statistics.web_api import S5WebApiMixin
from s1audit_configured import source_reader
from s1audit_configured.runtime import AuditCancelled

NODE = shutil.which("node")


@pytest.mark.skipif(NODE is None, reason="未安装 Node.js，无法执行 S5 前端行为检查")
def test_s5_maintenance_tools_are_hidden_by_default_and_toggle_with_five_nav_clicks(tmp_path):
    root = Path(__file__).resolve().parents[3]
    html = (root / "core/frontend/web/index.html").read_text(encoding="utf-8")
    for label in (
        "F2 模板 V_CUR 表达式",
        "F3 外部文件 V_EXT 表达式",
        "F4 模板校验公式转配置",
    ):
        assert label in html
    assert '<button type="button" class="sidenav-item" data-page="configured" data-system="S5" onclick="openS5Workbench()"' in html
    assert '#page-configured .card[hidden]{display:none!important}' in html
    assert '.s5-maintenance-actions[hidden]{display:none!important}' in html
    assert html.index('data-action="S5-F01"') < html.index('data-action="S5-F02"')
    assert 'id="s5-file-filter-toggle"' in html
    assert 'id="s5-match-filter"' in html and 'id="s5-business-filter"' in html and 'id="s5-period-filter"' in html
    assert "只作用于当前显示的文件" in html
    action_start = html.index('<div class="card actionbar s5-action-card">')
    maintenance_view_start = html.index('<div id="s5-maintenance-view"', action_start)
    action_card_markup = html[action_start:maintenance_view_start]
    assert all(f'data-action="S5-F0{number}"' in action_card_markup for number in range(1, 5))
    assert action_card_markup.count('class="card') == 1
    assert 'S5 配置维护功能' not in action_card_markup
    assert '连续点击左侧 S5 导航 5 次显示或隐藏' not in action_card_markup
    assert 'F1 按 S5 审核规则执行，结果与运行总账保存到独立输出目录。' not in action_card_markup
    assert 'id="s5-result"' not in action_card_markup

    start = html.index("let s5EggClickTimes=[];")
    end = html.index("function s5FormValues()", start)
    functions = html[start:end]
    assert "uiS5MaintenanceUnlocked" not in functions
    program = r"""
const assert=require('node:assert/strict');
let now=1000;
const Date={now:()=>now};
const ids=['s5-maintenance-entry','s5-maintenance-view','s5-panel-f1','s5-tool-vcur','s5-tool-vext','s5-tool-candidates'];
const nodes=Object.fromEntries(ids.map(id=>[id,{id,hidden:id!=='s5-panel-f1'}]));
const document={getElementById:id=>nodes[id],querySelectorAll:selector=>selector==='.s5-maintenance-panel'?ids.slice(3).map(id=>nodes[id]):[]};
let lastToast='';
const showToast=message=>{lastToast=message};
const requestSystemNavigation=()=>{};
const isSystemVisible=()=>true;
let state={s5:{config_path:'config.xlsx',configurationCheckTask:{}}};
const refresh=async()=>{};
let automaticChecks=0;
const checkS5Config=async()=>{automaticChecks++;state.s5.configurationCheckTask={task_id:'checked-once'}};
"""+functions+r"""
 (async()=>{
assert.equal(nodes['s5-maintenance-entry'].hidden,true);
assert.equal(nodes['s5-maintenance-view'].hidden,true);
assert.equal(nodes['s5-panel-f1'].hidden,false);
for(let i=0;i<4;i++)await openS5Workbench();
assert.equal(nodes['s5-maintenance-entry'].hidden,true);
await openS5Workbench();
assert.equal(nodes['s5-maintenance-entry'].hidden,false);
assert.equal(lastToast,'已显示 S5 配置维护工具');
assert.equal(automaticChecks,0);
openS5MaintenanceTool('vcur');
assert.equal(nodes['s5-panel-f1'].hidden,true);
assert.equal(nodes['s5-maintenance-view'].hidden,false);
assert.equal(nodes['s5-tool-vcur'].hidden,false);
assert.equal(nodes['s5-tool-vext'].hidden,true);
for(let i=0;i<5;i++)await openS5Workbench();
assert.equal(nodes['s5-maintenance-entry'].hidden,true);
assert.equal(nodes['s5-maintenance-view'].hidden,true);
assert.equal(nodes['s5-panel-f1'].hidden,false);
assert.equal(lastToast,'已隐藏 S5 配置维护工具');
console.log('S5 maintenance visibility passed');
})().catch(error=>{console.error(error);process.exit(1)});
"""
    script = tmp_path / "s5_maintenance_visibility_test.js"
    script.write_text(program, encoding="utf-8")
    result = subprocess.run([NODE, str(script)], capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stderr
    assert "S5 maintenance visibility passed" in result.stdout


def test_s5_path_labels_reuse_the_shared_clickable_path_style():
    root = Path(__file__).resolve().parents[3]
    html = (root / "core/frontend/web/index.html").read_text(encoding="utf-8")
    assert 'onclick="checkS5Config(\'current_business\')" title="只检查当前匹配模板对应业务的启用规则">检查当前业务' in html
    assert 'onclick="checkS5Config(\'all_enabled\')" title="按模板目录逐业务检查所有启用规则">检查全量' in html
    assert '<div class="field-head tight"><label>匹配模板</label><button class="btn soft micro" onclick="recognizeS5Template()"' in html
    assert '>识别模板</button>' in html
    helper_start = html.index("function renderS5PathName(")
    helper_end = html.index("function renderS5Maintenance(", helper_start)
    helper = html[helper_start:helper_end]
    assert "link.className='path-link'" in helper
    assert "link.onclick=()=>openPath(path)" in helper
    assert "renderS5PathName('s5-'+key+'-name',form[key])" in html
    assert "renderS5PathName('s5-vcur-template-path',form.template_path)" in html
    assert "renderS5PathName('s5-vext-external-path',form.external_path)" in html


def test_s5_f1_does_not_render_duplicate_completion_card():
    root = Path(__file__).resolve().parents[3]
    html = (root / "core/frontend/web/index.html").read_text(encoding="utf-8")
    start = html.index("function renderS5(s){")
    end = html.index("function s5ProgressState(", start)
    renderer = html[start:end]
    assert 'id="s5-result"' not in html
    assert "task.finding_counts" not in renderer
    assert "打开审核结果" not in renderer
    assert "打开输出目录" not in renderer
    assert "report.message" not in renderer
    assert "report.configuration_errors" not in renderer
    web_api = (root / "core/src/base_audit/systems/s5_configured_statistics/web_api.py").read_text(encoding="utf-8")
    assert 'if phase == "timing":' in web_api
    assert "耗时明细" in web_api


def test_s5_file_status_labels_do_not_treat_template_only_sheets_as_warnings():
    root = Path(__file__).resolve().parents[3]
    html = (root / "core/frontend/web/index.html").read_text(encoding="utf-8")
    assert "partial:'含模板外工作表'" in html
    assert "+(unmatched?' warn':'')" in html
    assert "源表匹配：" in html


def test_s5_configuration_check_has_progress_dialog_and_report_action():
    root = Path(__file__).resolve().parents[3]
    html = (root / "core/frontend/web/index.html").read_text(encoding="utf-8")
    assert 'id="s5-config-check-overlay"' in html
    assert 'aria-modal="true"' in html
    assert 'id="s5-config-check-fill"' in html
    assert "bridge.api.get_s5_config_check()" in html
    assert 'id="s5-config-check-open-report"' in html
    assert '打开检查报告' in html


def test_s5_configuration_check_dialog_keeps_only_compact_status_and_report_link():
    root = Path(__file__).resolve().parents[3]
    html = (root / "core/frontend/web/index.html").read_text(encoding="utf-8")
    start = html.index("function renderS5ConfigCheckTask(")
    end = html.index("function closeS5ConfigCheck(", start)
    renderer = html[start:end]
    assert "发现 '+issueCount+' 项配置问题，详细内容请打开检查报告查看。" in renderer
    assert "有 '+warningCount+' 条提示，详细内容请打开检查报告查看。" in renderer
    assert "report.business_checks" not in renderer
    assert "report.configuration_errors" in renderer  # 只读取问题数量，不展开明细
    assert "Object.entries(report.checks||{})" not in renderer
    assert "open.onclick=()=>openPath(task.report_path)" in renderer


def test_s5_configuration_report_uses_output_dir_and_preserves_pinned_path(tmp_path):
    legacy = tmp_path / "S5_执行结果"
    assert S5WebApiMixin._s5_config_check_output_dir({
        "output_dir": str(legacy), "output_pinned": False,
    }) == tmp_path / "执行结果"
    assert S5WebApiMixin._s5_config_check_output_dir({
        "output_dir": str(legacy), "output_pinned": True,
    }) == legacy


def test_s5_configuration_check_reports_index_and_reference_progress(request_data):
    events = []
    report = validate_s5_config(request_data["config_path"], on_progress=events.append)
    phases = {event["phase"] for event in events}
    assert {"rules", "locations", "references", "template", "external", "complete"} <= phases
    assert all(0 <= event["percent"] <= 100 for event in events)
    assert [event["percent"] for event in events] == sorted(event["percent"] for event in events)
    assert report["enabled_rules"] == 1


def make_source(path, value=100, *, named=False, code="91440000000000001A", cf=True):
    book = Workbook()
    info = book.active
    info.title = "基本信息"
    info["C3"] = "法人机构统一社会信用代码："
    info["D3"] = code
    sheet = book.create_sheet("贷款表")
    sheet.append(["项目", "本期余额"])
    sheet.append(["合计", value])
    if named:
        for name, address in [("表结构区域_001", "$A$1:$A$2"), ("表结构区域_002", "$B$1"),
                              ("数字区域_001", "$B$2"), ("条件格式区域_001", "$B$2")]:
            book.defined_names.add(DefinedName(name, attr_text=f"'贷款表'!{address}"))
    if cf:
        sheet.conditional_formatting.add("B2", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FFFF0000")))
    book.save(path)
    book.close()
    return path


@pytest.fixture
def request_data(tmp_path):
    source_dir = tmp_path / "本期输入"
    source_dir.mkdir()
    source = make_source(source_dir / "合成银行_核查表.xlsx")
    template = make_source(tmp_path / "布局模板.xlsx", named=True, cf=False)
    book = Workbook()
    sheet = book.active
    sheet.title = "审核规则"
    sheet.append(["规则编号", "是否启用", "业务名称", "工作表名称", "审核表达式", "当前值表达式", "规则名称", "审核级别", "校验类型", "校验指标", "规则描述"])
    expression = 'V_CUR("贷款表", "合计", "本期余额")'
    sheet.append(["R1", "是", "合成业务", "贷款表", expression + " < 50", expression, "余额提示", "硬性", "校验硬性错误", "余额", "合成余额提示"])
    config = tmp_path / "审核配置.xlsx"
    book.save(config)
    book.close()
    return {"config_path": str(config), "business": "合成业务", "template_path": str(template),
            "period": "2026-08-31", "input_dir": str(source_dir), "files": [str(source)],
            "output_dir": str(tmp_path / "任意中文输出")}


def hash_files(data):
    return {p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
            for p in [data["config_path"], data["template_path"], *data["files"]]}


def test_complete_f1_explicit_template_source_without_names_final_cf_and_read_only(request_data, tmp_path, monkeypatch):
    before = hash_files(request_data)
    events = []
    monkeypatch.setattr(source_reader, "_new_file_digest",
                        lambda: pytest.fail("S5-F1 不应计算输入文件内容哈希"))
    result = run_s5_f1(S5Request.from_dict(request_data), data_dir=tmp_path / "可写用户数据", on_progress=events.append)
    assert result["task_state"] == "succeeded"
    output_root = Path(request_data["output_dir"]).resolve()
    result_path = Path(result["result_path"]).resolve()
    assert result_path.parent == output_root
    assert result_path.name.startswith("汇总核查表校验_")
    assert "ledger_path" not in result
    assert result["input_integrity"] == "NOT_CHECKED"
    assert sorted(path.name for path in output_root.iterdir()) == [result_path.name]
    assert result_path.suffix == ".xlsx"
    assert result["finding_counts"]["rule_triggered"] == 1
    assert result["finding_counts"]["conditional_triggered"] == 1
    book = load_workbook(result_path, read_only=False)
    assert book["审核结果"].max_row == 2
    assert book["条件格式触发"].max_row == 2
    assert Path(book["审核结果"]["F2"].hyperlink.target.split("#", 1)[0]).name == Path(request_data["files"][0]).name
    book.close()
    assert "result_path" not in result["institutions"][0]
    assert {event["phase"] for event in events} >= {"identity", "audit", "export", "timing"}
    assert {
        "configuration_load_seconds", "institution_reference_load_seconds",
        "file_recognition_seconds", "rule_audit_seconds", "formula_scan_seconds",
        "conditional_format_scan_seconds", "result_export_seconds", "total_seconds",
    } <= set(result["stage_timings"])
    assert all(value >= 0 for value in result["stage_timings"].values())
    assert before == hash_files(request_data)


def test_vpre_without_previous_is_visible_and_matching_previous_is_retained(request_data, tmp_path):
    config = Path(request_data["config_path"])
    book = load_workbook(config)
    book["审核规则"]["E2"] = 'V_CUR("贷款表", "合计", "本期余额") < V_PRE("贷款表", "合计", "本期余额")'
    book.save(config)
    book.close()
    missing = run_s5_f1(S5Request.from_dict(request_data), data_dir=tmp_path)
    assert missing["task_state"] == "partial_failed"
    assert missing["counts"]["RETRIEVAL_FAILED"] == 1
    previous = make_source(tmp_path / "上期.xlsx", value=80, cf=False)
    request_data.update(previous_files=[str(previous)], previous_period="2026-07-31")
    matched = run_s5_f1(S5Request.from_dict(request_data), data_dir=tmp_path)
    assert matched["task_state"] == "succeeded"
    assert matched["counts"]["TRIGGERED"] == 1
    assert matched["institutions"][0]["previous"] == str(previous)


def test_explicit_template_is_not_silently_replaced(request_data, tmp_path):
    book = load_workbook(request_data["template_path"])
    book["贷款表"]["B1"] = "其他指标"
    book.save(request_data["template_path"])
    book.close()
    result = run_s5_f1(S5Request.from_dict(request_data), data_dir=tmp_path)
    assert result["task_state"] == "partial_failed"
    assert result["finding_counts"]["rule_execution_errors"] >= 1


def test_cancellation_stops_before_next_institution_and_never_returns_success(request_data, tmp_path):
    other = make_source(Path(request_data["input_dir"]) / "第二家.xlsx", code="91440000000000002B")
    request_data["files"].append(str(other))
    cancel = threading.Event()
    seen = []

    def progress(event):
        seen.append(event)
        if event["phase"] == "audit":
            cancel.set()

    with pytest.raises(AuditCancelled):
        run_s5_f1(S5Request.from_dict(request_data), data_dir=tmp_path, on_progress=progress, cancel_event=cancel)
    assert len([e for e in seen if e["phase"] == "audit"]) == 1


def test_configuration_error_reports_physical_sheet_cell_and_rule(request_data):
    book = load_workbook(request_data["config_path"])
    book["审核规则"]["E2"] = "BOGUS("
    book["审核规则"]["B2"] = "是"
    book.save(request_data["config_path"])
    book.close()
    report = validate_s5_config(request_data["config_path"])
    assert not report["valid"]
    assert report["checks"]["expressions"]["status"] == "failed"
    assert report["rules"] == report["enabled_rules"] == 1
    assert report["disabled_rules"] == 0
    location = report["configuration_errors"][0]["locations"][0]
    assert location["sheet"] == "审核规则" and location["cell"] == "E2" and location["rule_id"] == "R1"


def test_s5_configuration_check_validates_dictionary_keys_for_enabled_rule(request_data):
    book = load_workbook(request_data["config_path"])
    book["审核规则"]["B2"] = "是"
    book["审核规则"]["E2"] = "{取不到的值} > 0"
    book.save(request_data["config_path"])
    book.close()
    report = validate_s5_config(request_data["config_path"])
    assert not report["valid"]
    assert "引用未定义变量" in report["configuration_errors"][0]["message"]
    assert report["configuration_errors"][0]["locations"][0]["rule_id"] == "R1"


def test_s5_configuration_check_skips_disabled_rules_and_reports_incomplete(request_data):
    book = load_workbook(request_data["config_path"])
    book["审核规则"]["B2"] = "否"
    book["审核规则"]["E2"] = "BOGUS("
    book.save(request_data["config_path"])
    book.close()
    report = validate_s5_config(
        request_data["config_path"], template_path=request_data["template_path"], business="合成业务",
    )
    assert report["valid"]
    assert not report["complete"]
    assert report["enabled_rules"] == 0
    assert report["checks"]["expressions"]["status"] == "not_checked"
    assert report["checks"]["dictionary_keys"]["status"] == "not_checked"
    assert report["checks"]["template_addresses"]["status"] == "not_checked"
    assert "全部禁用" in report["message"]
    assert "未执行" in report["message"]


def test_s5_configuration_check_resolves_template_addresses_with_runtime_locator(request_data):
    report = validate_s5_config(
        request_data["config_path"], template_path=request_data["template_path"], business="合成业务",
    )
    assert report["valid"]
    address_check = report["checks"]["template_addresses"]
    assert address_check["status"] == "passed"
    assert address_check["checked"] == address_check["resolved"] == 1
    assert report["complete"]


def test_s5_configuration_check_reports_unresolvable_template_address(request_data):
    book = load_workbook(request_data["config_path"])
    book["审核规则"]["E2"] = 'V_CUR("贷款表", "不存在的行", "本期余额") < 50'
    book.save(request_data["config_path"])
    book.close()
    report = validate_s5_config(
        request_data["config_path"], template_path=request_data["template_path"], business="合成业务",
    )
    assert not report["valid"]
    issue = report["configuration_errors"][0]
    assert "未找到完整行路径" in issue["message"]
    assert issue["locations"][0]["sheet"] == "审核规则"
    assert issue["locations"][0]["cell"] == "E2"
    assert issue["locations"][0]["rule_id"] == "R1"


def test_s5_configuration_check_marks_missing_template_address_check_as_incomplete(request_data):
    report = validate_s5_config(request_data["config_path"], business="合成业务")
    assert report["valid"]
    assert not report["complete"]
    assert report["checks"]["template_addresses"]["status"] == "not_checked"
    assert any("无法检查模板取数地址" in warning for warning in report["warnings"])
    assert all("V_PRE" not in warning for warning in report["warnings"])


def test_s5_current_business_check_ignores_other_business_expression_errors(request_data):
    book = load_workbook(request_data["config_path"])
    sheet = book["审核规则"]
    other = [cell.value for cell in sheet[2]]
    other[0] = "R2"
    other[2] = "其他业务"
    other[3] = "其他表"
    other[4] = "BOGUS("
    other[5] = ""
    sheet.append(other)
    book.save(request_data["config_path"])
    book.close()

    current = validate_s5_config(
        request_data["config_path"], template_path=request_data["template_path"],
        business="合成业务", scope="current_business",
    )
    assert current["valid"] and current["complete"]
    assert current["business"] == "合成业务"
    assert current["enabled_rules"] == 1
    assert current["checks"]["template_addresses"]["resolved"] == 1

    inferred = validate_s5_config(
        request_data["config_path"], template_path=request_data["template_path"],
        scope="current_business",
    )
    assert inferred["valid"] and inferred["business"] == "合成业务"

    full = validate_s5_config(
        request_data["config_path"], template_dir=str(Path(request_data["template_path"]).parent),
        scope="all_enabled",
    )
    assert not full["valid"]
    assert full["enabled_rules"] == 2
    assert full["checks"]["expressions"]["status"] == "failed"


def test_s5_full_config_check_maps_each_enabled_business_to_template(request_data, tmp_path):
    template_dir = tmp_path / "all-templates"
    template_dir.mkdir()
    shutil.copy2(request_data["template_path"], template_dir / "合成业务模板.xlsx")
    other_template = template_dir / "其他业务模板.xlsx"
    other_book = Workbook()
    other_sheet = other_book.active
    other_sheet.title = "其他表"
    other_sheet.append(["项目", "本期余额"])
    other_sheet.append(["合计", 10])
    other_book.save(other_template)
    other_book.close()

    book = load_workbook(request_data["config_path"])
    sheet = book["审核规则"]
    other = [cell.value for cell in sheet[2]]
    other[0] = "R2"
    other[2] = "其他业务"
    other[3] = "其他表"
    other[4] = 'V_CUR("其他表", "合计", "本期余额") < 50'
    other[5] = 'V_CUR("其他表", "合计", "本期余额")'
    sheet.append(other)
    book.save(request_data["config_path"])
    book.close()

    events = []
    report = validate_s5_config(
        request_data["config_path"], template_dir=str(template_dir), scope="all_enabled",
        on_progress=events.append,
    )
    assert report["valid"] and report["complete"]
    assert report["enabled_rules"] == 2
    assert report["checks"]["template_addresses"]["checked"] == 2
    assert report["checks"]["template_addresses"]["resolved"] == 2
    checks = {item["name"]: item for item in report["business_checks"]}
    assert checks["合成业务"]["template"] == "合成业务模板.xlsx"
    assert checks["其他业务"]["template"] == "其他业务模板.xlsx"
    assert {event["phase"] for event in events} >= {"rules", "template", "external", "complete"}
    assert events[-1]["percent"] == 100
    summary = S5WebApiMixin._s5_config_check_report_summary(report)
    assert summary["scope"] == "all_enabled"
    assert all("errors" not in item["template_addresses"] for item in summary["business_checks"])


def test_s5_full_config_check_is_incomplete_when_template_is_missing(request_data, tmp_path):
    empty_templates = tmp_path / "empty-templates"
    empty_templates.mkdir()
    report = validate_s5_config(
        request_data["config_path"], template_dir=str(empty_templates), scope="all_enabled",
    )
    assert report["valid"]
    assert not report["complete"]
    assert report["checks"]["template_addresses"]["status"] == "not_checked"
    assert report["business_checks"][0]["template_addresses"]["status"] == "not_checked"
    assert "没有可读取的 XLSX 模板" in report["business_checks"][0]["template_addresses"]["message"]


def test_no_enabled_rules_is_not_reported_as_pass(request_data, tmp_path):
    book = load_workbook(request_data["config_path"])
    book["审核规则"]["B2"] = "否"
    book.save(request_data["config_path"])
    book.close()
    with pytest.raises(ValueError, match="没有启用规则"):
        run_s5_f1(S5Request.from_dict(request_data), data_dir=tmp_path)


@pytest.mark.parametrize("fallback", [False, True])
@pytest.mark.parametrize("time_series", [False, True])
def test_external_code_mapping_runs_complete_f1_with_both_layouts_and_readers(
        request_data, tmp_path, monkeypatch, fallback, time_series):
    if fallback:
        monkeypatch.setenv("AUDIT3_DISABLE_CALAMINE", "1")
    book = Workbook()
    ref = book.active
    ref.title = "机构参照表"
    ref.append(["机构名称", "社会信用代码", "报表项目", "禁用"])
    ref.append(["合成银行", "91440000000000001A", "合成报表项目", "否"])
    data = book.create_sheet("集中系统数据")
    if time_series:
        data.append(["数据日期", "行名称", "贷款余额"])
        data.append(["2026-07-31", "合成报表项目", 999])
        data.append(["2026-08-31", "合成报表项目", 100])
        data.append(["2026-09-30", "合成报表项目", 888])
    else:
        data.append([])
        data.append([None, None, "贷款余额"])
        data.append([None, "合成报表项目", 100])
    external = tmp_path / "外部映射.xlsx"
    book.save(external)
    book.close()
    config = load_workbook(request_data["config_path"])
    config["审核规则"]["E2"] = 'V_CUR("贷款表", "合计", "本期余额") = V_EXT("贷款余额")'
    config.save(request_data["config_path"])
    config.close()
    request_data["external_path"] = str(external)
    before = hashlib.sha256(external.read_bytes()).hexdigest()
    result = run_s5_f1(S5Request.from_dict(request_data), data_dir=tmp_path)
    assert result["task_state"] == "succeeded"
    assert result["counts"]["PASSED"] == 1
    assert result["institutions"][0]["identity"]["report_item"] == "合成报表项目"
    if fallback:
        assert result["warnings"]
    assert hashlib.sha256(external.read_bytes()).hexdigest() == before


@pytest.mark.parametrize("time_series", [False, True])
@pytest.mark.parametrize("bad_value", [None, "误填文字"])
@pytest.mark.parametrize("writer_engine", ["direct", "openpyxl"])
def test_external_layout_anomalies_warn_and_continue_for_both_layouts(
        request_data, tmp_path, monkeypatch, time_series, bad_value, writer_engine):
    monkeypatch.setenv("AUDIT3_WRITER_ENGINE", writer_engine)
    book = Workbook()
    ref = book.active
    ref.title = "机构参照表"
    ref.append(["机构名称", "社会信用代码", "报表项目", "禁用"])
    ref.append(["合成银行", "91440000000000001A", "合成报表项目", "否"])
    data = book.create_sheet("集中系统数据")
    if time_series:
        data.append(["数据日期", "报表名称", "报表代码", "版本号", "锁定信息", "报表组", "行顺序号", "行名称", "贷款余额"])
        data.append(["日期错误", "金融基础数据核对模板", "R104400000g0e", "20265", "-", "219", "84w0", "合成报表项目", 999])
        data.append(["2026-08-31", "金融基础数据核对模板", "R104400000g0e", "20265", "-", "219", "84w1", "合成报表项目", bad_value])
        data.append(["2026-08-31", "金融基础数据核对模板", "R104400000g0e", "20265", "-", "219", "84w2", None, None])
        data.append([None] * 9)
        data.append(["2026-08-31", "金融基础数据核对模板", "R104400000g0e", "20265", "-", "219", "84w3", None, 456])
    else:
        data.append(["行顺序号", "报表项目", "贷款余额"])
        data.append(["84w0", "合成报表项目", bad_value])
        data.append(["84w1", None, None])
        data.append([None, None, None])
        data.append(["84w2", None, 456])
    external = tmp_path / "异常数据.xlsx"
    book.save(external)
    book.close()

    config = load_workbook(request_data["config_path"])
    config["审核规则"]["E2"] = 'V_CUR("贷款表", "合计", "本期余额") = V_EXT("贷款余额")'
    config.save(request_data["config_path"])
    config.close()
    request_data["external_path"] = str(external)

    result = run_s5_f1(S5Request.from_dict(request_data), data_dir=tmp_path)

    assert result["task_state"] == "succeeded"
    assert result["finding_counts"]["data_quality_warnings"] >= (4 if time_series else 2)
    assert any("外部数据质量提示" in item for item in result["warnings"])
    assert any("已按 0 参与审核" in item for item in result["warnings"])
    assert any("未参与时序匹配" in item for item in result["warnings"]) is time_series

    output = load_workbook(result["result_path"], read_only=True, data_only=True)
    try:
        quality = output["数据质量提示"]
        quality_rows = list(quality.iter_rows(min_row=2, values_only=True))
        messages = [str(row[5]) for row in quality_rows]
        assert any("已按 0 参与审核" in message for message in messages)
        assert any("行名称为空" in message for message in messages)
        metadata_only_row = "第4行" if time_series else "第3行"
        fully_empty_row = "第5行" if time_series else "第4行"
        assert not any(
            str(row[4]) == metadata_only_row and "行名称及指标数据均为空" in str(row[5])
            for row in quality_rows
        )
        assert any(
            str(row[4]) == fully_empty_row and "行名称及指标数据均为空" in str(row[5])
            for row in quality_rows
        )
        if time_series:
            assert any("本行未参与时序匹配" in message for message in messages)
    finally:
        output.close()


def test_time_series_missing_current_date_only_fails_affected_rule_and_continues_batch(request_data, tmp_path):
    second = make_source(
        Path(request_data["input_dir"]) / "第二家_核查表.xlsx",
        code="91440000000000002B",
        cf=False,
    )
    request_data["files"].append(str(second))
    book = Workbook()
    ref = book.active
    ref.title = "机构参照表"
    ref.append(["机构名称", "社会信用代码", "报表项目", "禁用"])
    ref.append(["合成银行", "91440000000000001A", "合成报表项目", "否"])
    ref.append(["第二家", "91440000000000002B", "第二报表项目", "否"])
    data = book.create_sheet("集中系统数据")
    data.append(["数据日期", "行名称", "贷款余额"])
    data.append(["2026-08-31", "合成报表项目", 100])
    data.append(["日期错误", "第二报表项目", 200])
    external = tmp_path / "部分日期无效.xlsx"
    book.save(external)
    book.close()

    config = load_workbook(request_data["config_path"])
    config["审核规则"]["E2"] = 'V_CUR("贷款表", "合计", "本期余额") = V_EXT("贷款余额")'
    config.save(request_data["config_path"])
    config.close()
    request_data["external_path"] = str(external)

    result = run_s5_f1(S5Request.from_dict(request_data), data_dir=tmp_path)

    assert result["task_state"] == "partial_failed"
    assert result["files_processed"] == 2
    assert result["counts"]["RETRIEVAL_FAILED"] == 1
    assert {record["status"] for record in result["institutions"]} == {
        "COMPLETED", "COMPLETED_WITH_ISSUES",
    }
    assert any("本行未参与时序匹配" in item for item in result["warnings"])


def test_s5_maintenance_uses_the_current_panel_paths(request_data, tmp_path, monkeypatch):
    template = Path(request_data["template_path"])
    external = tmp_path / "maintenance-external.xlsx"
    book = Workbook()
    mapping = book.active
    mapping.title = "机构参照表"
    mapping.append(["机构代码", "报表项目"])
    mapping.append(["91440000000000001A", "合成报表项目"])
    data = book.create_sheet("集中系统数据")
    data.append(["数据日期", "行名称", "贷款余额"])
    data.append(["2026-07-31", "合成报表项目", 90])
    data.append(["2026-08-31", "合成报表项目", 100])
    book.save(external)
    book.close()

    captured = []

    class Thread:
        def __init__(self, **kwargs):
            captured.append(kwargs)

        def start(self):
            pass

    class Host(S5WebApiMixin):
        project_root = tmp_path
        _execution_start_lock = threading.RLock()
        _cancel_event = threading.Event()
        state = {"busy": False, "log": []}

        def _execution_busy(self):
            return self.state["busy"]

        def _log(self, *args, **kwargs):
            pass

    host = Host()
    host._init_s5()
    host.state["s5"].update(
        template_path=str(template), template_dir=str(template.parent),
        external_path=str(external), recognition={"errors": {}},
    )
    host._s5_recognize = lambda: host.state["s5"]["recognition"]
    monkeypatch.setattr("base_audit.systems.s5_configured_statistics.web_api.threading.Thread", Thread)

    assert host.start_s5_maintenance("vcur", {"outputDir": str(tmp_path / "vcur")})["accepted"]
    assert host.start_s5_maintenance("vext", {"outputDir": str(tmp_path / "vext")})["accepted"]
    assert host.start_s5_maintenance("candidates", {"outputDir": str(tmp_path / "candidates")})["accepted"]
    args = [call["args"][2] for call in captured]
    assert args[0]["files"] == [str(template.resolve())]
    assert args[0]["layoutTemplate"] == str(template.resolve())
    assert args[1]["files"] == [str(external.resolve())]
    assert args[2]["files"] == [str(template.resolve())]
    assert args[2]["externalFile"] == str(external.resolve())


@pytest.mark.parametrize("time_series", [False, True])
def test_s5_vext_reference_export_supports_both_external_layouts(tmp_path, time_series):
    from base_audit.systems.s5_configured_statistics.expression_export import export_expressions

    external = tmp_path / ("timeseries.xlsx" if time_series else "ordinary.xlsx")
    book = Workbook()
    mapping = book.active
    mapping.title = "机构参照表"
    mapping.append(["机构代码", "报表项目"])
    mapping.append(["91440000000000001A", "合成报表项目"])
    data = book.create_sheet("集中系统数据")
    if time_series:
        data.append(["数据日期", "行名称", "贷款余额"])
        data.append(["2026-07-31", "合成报表项目", 90])
        data.append(["2026-08-31", "合成报表项目", 100])
    else:
        data.append(["說明"])
        data.append([None, "报表项目", "贷款余额"])
        data.append([1, "合成报表项目", 100])
    book.save(external)
    book.close()

    result = export_expressions([external], mode="vext", output_root=tmp_path / "out")
    assert result["counts"]["ready"] == (2 if time_series else 1)
    assert result["counts"]["failed"] == 0
    reference = load_workbook(result["files"][0]["path"], data_only=False)
    try:
        values = [cell.value for row in reference["集中系统数据"].iter_rows()
                  for cell in row if isinstance(cell.value, str) and cell.value.startswith("V_EXT(")]
        assert len(values) == (2 if time_series else 1)
        assert all(value == 'V_EXT("贷款余额")' for value in values)
    finally:
        reference.close()




def test_s5_state_does_not_write_s1_fields_and_mutual_exclusion(request_data, tmp_path):
    class Host(S5WebApiMixin):
        project_root = tmp_path
        _execution_start_lock = threading.RLock()
        _cancel_event = threading.Event()
        state = {"input": "S1 原目录", "template": "S1 原模板", "selectedFiles": ["S1 文件"], "busy": True, "log": []}

        def _execution_busy(self):
            return self.state["busy"]

        def _log(self, *args, **kwargs):
            pass

    host = Host()
    before = copy.deepcopy(host.state)
    host._init_s5()
    host.update_s5({"input_dir": request_data["input_dir"], "period": request_data["period"]})
    assert {key: host.state[key] for key in before} == before
    assert not host.start_s5("S5-F01")["accepted"]


def test_s5_configuration_check_runs_in_background_and_saves_report_under_execution_results(
        request_data, tmp_path, monkeypatch):
    captured = {}

    class Thread:
        def __init__(self, **kwargs):
            captured.update(kwargs)

        def start(self):
            captured["started"] = True

    class Host(S5WebApiMixin):
        project_root = tmp_path
        _execution_start_lock = threading.RLock()
        state = {"busy": False, "log": []}

        def _execution_busy(self):
            return bool(self.state["busy"])

        def _log(self, message, **kwargs):
            self.state["log"].append(message)

    host = Host()
    host._init_s5()
    input_dir = Path(request_data["input_dir"])
    host.update_s5({
        "input_dir": str(input_dir), "template_dir": str(Path(request_data["template_path"]).parent),
        "config_path": request_data["config_path"], "template_path": request_data["template_path"],
        "business": "合成业务", "period": "2026-08-31", "selected_files": request_data["files"],
        "output_dir": str(input_dir / "执行结果"),
    })
    monkeypatch.setattr("base_audit.systems.s5_configured_statistics.web_api.threading.Thread", Thread)

    started = host.check_s5_config()
    assert started["accepted"]
    assert captured["started"]
    assert host.state["s5"]["configurationCheckTask"]["task_state"] == "queued"
    assert host.state["s5"]["configurationCheckTask"]["scope"] == "current_business"
    alternate = host.check_s5_config("all_enabled")
    assert not alternate["accepted"]
    assert "正在运行" in alternate["message"]

    captured["target"](*captured["args"])
    task = host.get_s5_config_check()
    report_path = Path(task["report_path"])
    assert task["task_state"] == "succeeded"
    assert report_path.is_file()
    assert report_path.parent == input_dir / "执行结果"
    assert report_path.parent.name != "S5_执行结果"
    assert "审核表达式" in report_path.read_text(encoding="utf-8")
    assert task["progress"]["percent"] == 100


def test_s5_migrates_saved_release_defaults_and_preserves_custom_paths(tmp_path):
    config_dir = tmp_path / "config-real"
    new_config = config_dir / "5.逐笔统计系统_审核配置.xlsx"
    old_config = config_dir / "S5_逐笔统计系统_审核配置.xlsx"
    named_templates = config_dir / "模板文件_命名区域"
    legacy_templates = config_dir / "S5_模板文件"
    for path in (new_config, old_config):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.touch()
    named_templates.mkdir()
    legacy_templates.mkdir()
    settings_path = tmp_path / "data" / "s5" / "界面设置.json"
    settings_path.parent.mkdir(parents=True)
    settings_path.write_text(json.dumps({
        "config_path": str(old_config),
        "template_dir": str(legacy_templates),
    }), encoding="utf-8")

    class Host(S5WebApiMixin):
        project_root = tmp_path
        _execution_start_lock = threading.RLock()
        _cancel_event = threading.Event()
        state = {"busy": False, "log": []}

        def _execution_busy(self):
            return self.state["busy"]

        def _log(self, *args, **kwargs):
            pass

    host = Host()
    host._init_s5()
    assert host.state["s5"]["config_path"] == str(new_config)
    assert host.state["s5"]["template_dir"] == str(named_templates)

    custom_config = tmp_path / "custom" / "my-s5-config.xlsx"
    custom_templates = tmp_path / "custom" / "my-templates"
    custom_config.parent.mkdir(parents=True)
    custom_config.touch()
    custom_templates.mkdir()
    settings_path.write_text(json.dumps({
        "config_path": str(custom_config),
        "template_dir": str(custom_templates),
    }), encoding="utf-8")
    host = Host()
    host._init_s5()
    assert host.state["s5"]["config_path"] == str(custom_config)
    assert host.state["s5"]["template_dir"] == str(custom_templates)


def test_s5_start_freezes_parameters_before_thread(request_data, tmp_path, monkeypatch):
    captured = {}

    class Thread:
        def __init__(self, **kwargs):
            captured.update(kwargs)

        def start(self):
            pass

    class Host(S5WebApiMixin):
        project_root = tmp_path
        _execution_start_lock = threading.RLock()
        _cancel_event = threading.Event()
        state = {"busy": False}

        def _execution_busy(self):
            return self.state["busy"]

        def _log(self, *args, **kwargs):
            pass

    host = Host()
    host._init_s5()
    data = {key: value for key, value in request_data.items() if key != "files"}
    host.update_s5(data)
    monkeypatch.setattr("base_audit.systems.s5_configured_statistics.web_api.threading.Thread", Thread)
    assert host.start_s5("S5-F01")["accepted"]
    request = captured["args"][2]
    host.state["s5"]["template_path"] = "changed.xlsx"
    host.state["s5"]["selected_files"].clear()
    assert request.template_path == Path(request_data["template_path"])
    assert request.business == request_data["business"]
    assert len(request.files) == 1
    assert host.state["busy"]
