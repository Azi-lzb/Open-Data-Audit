from dataclasses import replace
from datetime import datetime
from decimal import Decimal

from openpyxl import load_workbook

from s1audit_configured import outputs
from s1audit_configured.model import AuditRunResult, Evidence, RuleResult


def signature(path):
    book = load_workbook(path)
    result = []
    for sheet in book:
        result.append((sheet.title, sheet.freeze_panes, sheet.auto_filter.ref,
                       sheet.sheet_view.showGridLines, str(sheet.page_setup),
                       str(sheet.sheet_properties.pageSetUpPr), sheet.print_title_rows,
                       [(key, value.width) for key, value in sheet.column_dimensions.items()],
                       [(key, value.height) for key, value in sheet.row_dimensions.items()],
                       [[(cell.value, cell.data_type, cell.number_format, str(cell.font), str(cell.fill),
                          str(cell.alignment), cell.hyperlink.target if cell.hyperlink else None)
                         for cell in row] for row in sheet]))
    book.close()
    return result


def sample_result(tmp_path):
    source = tmp_path / "source & 空格.xlsx"
    evidence = Evidence("CUR", str(source), "hash", "O'Brien", "C4", 123, "int", report_item="REPORT")
    item = RuleResult("RULE", 1, "业务", "20260831", "000000000000000001", "机构名", "TRIGGERED",
                      True, "错误", "比对", "全路径 → 指标", description="=不是公式\r\n<&> 中文",
                      current_value=Decimal("123.45678"), comparison_value=Decimal("100"), difference_value=Decimal("23.45678"),
                      current_precision=5, comparison_precision=0, difference_precision=2,
                      evidence=[evidence], rule_name="规则", sheet_name="O'Brien", run_id="run")
    results = [item, replace(item, rule_id="SOFT", severity="软性错误", current_value=True, comparison_value=None, difference_value=float("nan")),
               replace(item, rule_id="ERR", status="RETRIEVAL_FAILED", evidence=[], error_message="缺失路径",
                       error_evidence={"sheet": "O'Brien", "cell": "D9", "source_path": str(source), "cells": ["D9", "E9"]}),
               *[replace(item, rule_id=status, status=status) for status in ("PASSED", "NOT_EXECUTED", "NOT_APPLICABLE")]]
    conditional = [{"机构名称": "机构", "统一社会信用代码": "0001", "工作表名": "表", "定位单元格": "A1",
                    "单元格值": "=SUM(1)", "判定状态": "已触发", "规则公式": "=A1>0", "优先级": 1},
                   {"机构名称": "机构", "判定状态": "无法判定", "判定说明": "不支持 colorScale"}]
    return AuditRunResult("run", "20260831", "1", "hash", "COMPLETED", results, {}, conditional_format_results=conditional), {str(source.resolve()): source}


def test_direct_writer_matches_rich_cells_styles_links_and_layout(tmp_path, monkeypatch):
    result, copies = sample_result(tmp_path)
    options = {"batch_records": [{"run_id": "run", "source": str(next(iter(copies))), "identity": {"report_item": "REPORT"}, "status": "COMPLETED"},
                                  {"source": "未识别_表.xlsx", "status": "IDENTITY_FAILED", "message": "无法识别"}],
               "previous_records": [{"source": "上期_表.xlsx", "status": "FAILED", "message": "无效上期"}]}
    old, new = tmp_path / "rich.xlsx", tmp_path / "direct.xlsx"
    outputs._write_result_workbook_openpyxl(result, old, copies, **options)
    def forbidden(*args, **kwargs):
        raise AssertionError("supported result tables must not use the fallback")
    monkeypatch.setattr(outputs, "_write_result_workbook_openpyxl", forbidden)
    outputs._write_result_workbook(result, new, copies, **options)
    assert signature(new) == signature(old)
    book = load_workbook(new)
    assert book.sheetnames == ["审核结果", "条件格式触发", "条件格式无法判定"]
    assert book["审核结果"].max_row == 6
    assert book["审核结果"]["I2"].data_type == "s"
    assert book["条件格式触发"]["F2"].data_type == "s"
    book.close()


def test_uncommon_type_falls_back_before_writing_zip(tmp_path, monkeypatch):
    result, copies = sample_result(tmp_path)
    result.results[0].current_value = datetime(2026, 10, 21)
    called = []
    original = outputs._write_result_workbook_openpyxl
    def fallback(*args, **kwargs):
        called.append(True)
        return original(*args, **kwargs)
    monkeypatch.setattr(outputs, "_write_result_workbook_openpyxl", fallback)
    path = tmp_path / "fallback.xlsx"
    outputs._write_result_workbook(result, path, copies)
    assert called == [True]
    assert load_workbook(path).sheetnames == ["审核结果", "条件格式触发", "条件格式无法判定"]
