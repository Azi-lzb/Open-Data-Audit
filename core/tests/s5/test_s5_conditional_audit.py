from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName

from s1audit_configured.conditional_audit import extract_conditional_triggers
from s1audit_configured.outputs import _write_result_workbook
from s1audit_configured.model import AuditRunResult


def test_python_conditional_format_detection_reports_only_triggered_cells(tmp_path):
    path = tmp_path / "source.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "贷款表"
    sheet["B2"] = 12
    sheet["B3"] = 0
    sheet.conditional_formatting.add(
        "B2:B3", CellIsRule(operator="greaterThan", formula=["0"],
                            fill=PatternFill("solid", fgColor="FFFF0000")),
    )
    sheet.conditional_formatting.add(
        "C2:C3", FormulaRule(formula=["C2>0"],
                             fill=PatternFill("solid", fgColor="FFFF0000")),
    )
    sheet.conditional_formatting.add(
        "B4", FormulaRule(formula=['B4=""'],
                           fill=PatternFill("solid", fgColor="FFFF0000")),
    )
    sheet["C2"] = 3
    book.save(path)
    book.close()

    rows = extract_conditional_triggers(path, institution_name="测试机构", report_item="单位贷款",
                                        institution_code="TEST001")
    assert [(row["工作表名"], row["定位单元格"]) for row in rows] == [
        ("贷款表", "B2"), ("贷款表", "C2"), ("贷款表", "B4"),
    ]
    assert all(row["机构名称"] == "测试机构" for row in rows)


def test_conditional_formula_uses_one_anchor_for_disjoint_applies_to_ranges(tmp_path):
    path = tmp_path / "disjoint-ranges.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["A1"] = 0
    sheet["A2"] = 0
    sheet["C1"] = 1
    sheet["C2"] = 0
    sheet.conditional_formatting.add(
        "A1:A2 C1:C2", FormulaRule(formula=["A1>0"],
                                    fill=PatternFill("solid", fgColor="FFFF0000")),
    )
    book.save(path)
    book.close()

    rows = extract_conditional_triggers(path)
    assert [(row["定位单元格"], row["规则公式"]) for row in rows] == [("C1", "C1>0")]


def test_cellis_reference_operand_moves_with_applies_to_cell(tmp_path):
    path = tmp_path / "cellis-reference.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["B1"] = 1
    sheet["B2"] = 0
    sheet["C1"] = 1
    sheet["C2"] = 1
    sheet.conditional_formatting.add(
        "C1:C2", CellIsRule(operator="equal", formula=["B1"],
                             fill=PatternFill("solid", fgColor="FFFF0000")),
    )
    book.save(path)
    book.close()

    rows = extract_conditional_triggers(path)
    assert [row["定位单元格"] for row in rows] == ["C1"]


def test_source_conditional_rules_use_source_values_and_template_scan_region(tmp_path):
    template = tmp_path / "template.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "贷款表"
    book.defined_names.add(DefinedName("条件格式区域_001", attr_text="'贷款表'!$D$1"))
    book.save(template)
    book.close()

    source = tmp_path / "source.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "贷款表"
    sheet["B2"] = 0
    sheet.conditional_formatting.add(
        "D1", FormulaRule(formula=["B1<>B2+B3"],
                          fill=PatternFill("solid", fgColor="FFFF0000")),
    )
    book.save(source)
    book.close()

    # Preserve S1F1's Python evaluator semantics for a blank-vs-zero comparison.
    assert [(row["工作表名"], row["定位单元格"])
            for row in extract_conditional_triggers(source, template_path=template)] == [("贷款表", "D1")]

    book = load_workbook(source)
    book["贷款表"]["B2"] = 2
    book.save(source)
    book.close()
    rows = extract_conditional_triggers(source, template_path=template)
    assert [(row["工作表名"], row["定位单元格"]) for row in rows] == [("贷款表", "D1")]


def test_named_conditional_region_limits_source_rule_scan(tmp_path):
    template = tmp_path / "named-cf-template.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "贷款表"
    book.defined_names.add(DefinedName("条件格式区域_001", attr_text="'贷款表'!$B$2"))
    book.save(template)
    book.close()

    source = tmp_path / "named-cf-source.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "贷款表"
    sheet["B2"] = 1
    sheet["B3"] = 2
    sheet.conditional_formatting.add(
        "B2:B3", CellIsRule(operator="greaterThan", formula=["0"],
                             fill=PatternFill("solid", fgColor="FFFF0000")),
    )
    book.save(source)
    book.close()

    rows = extract_conditional_triggers(source, template_path=template)
    assert [(row["工作表名"], row["定位单元格"]) for row in rows] == [("贷款表", "B2")]


def test_result_workbook_has_separate_conditional_format_sheet(tmp_path):
    result = AuditRunResult(
        run_id="test", period="20260831", configuration_version="", configuration_sha256="",
        status="COMPLETED", results=[], ledger={},
        conditional_format_results=[{
            "机构名称": "测试机构", "报表项目": "单位贷款", "统一社会信用代码": "TEST001",
            "工作表名": "贷款表", "定位单元格": "B2", "单元格值": 12,
            "条件格式类型": "cellis", "规则公式": "0", "应用范围": "B2:B3",
            "优先级": 1, "来源文件": "source.xlsx",
        }],
    )
    path = tmp_path / "result.xlsx"
    _write_result_workbook(result, path, {})
    book = load_workbook(path, data_only=True)
    assert book.sheetnames == ["审核结果", "条件格式触发"]
    assert book["条件格式触发"]["A2"].value == "测试机构"
    assert book["条件格式触发"]["E2"].value == "B2"
    book.close()
