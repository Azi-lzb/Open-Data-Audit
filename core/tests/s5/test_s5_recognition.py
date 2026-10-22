from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from base_audit.systems.s5_configured_statistics.recognition import (
    classify_s5_files, detect_s5_period, file_period_hint, match_s5_template,
    periods_in, recognize_s5,
)
from base_audit.systems.s5_configured_statistics.web_api import S5WebApiMixin
from test_s5_integration import make_source, request_data


@pytest.mark.parametrize("value,expected", [
    ("甲银行_20260831_核查表", "2026-08-31"),
    ("reports#2026-08-31#机构", "2026-08-31"),
    ("报告2026年7月31日", "2026-07-31"),
    ("2026/7/3", "2026-07-03"),
    ("202608", "2026-08-31"),
    ("2024年2月", "2024-02-29"),
])
def test_period_formats_and_month_end_are_explicit(value, expected):
    assert periods_in(value) == {expected}


def test_invalid_date_is_not_silently_replaced_with_month_end():
    with pytest.raises(ValueError):
        periods_in("2026-02-30")
    assert not periods_in("=DATE(2026,8,31)")


def test_period_prefers_source_fields_and_rejects_mixed_dates(request_data, tmp_path):
    source = Path(request_data["files"][0])
    book = load_workbook(source)
    book["基本信息"]["A1"], book["基本信息"]["B1"] = "数据日期：", "2026年8月31日"
    book.save(source)
    book.close()
    period, evidence = detect_s5_period([source], source.parent)
    assert period == "2026-08-31" and evidence == "源表报告期字段"
    named = source.with_name("合成银行_20260930.xlsx")
    source.rename(named)
    assert detect_s5_period([named], named.parent)[0] == "2026-09-30"
    other = make_source(tmp_path / "其他_20260831.xlsx")
    with pytest.raises(ValueError, match="多个报告期"):
        detect_s5_period([named, other], tmp_path)


def test_no_date_reports_action_and_folder_month_is_supported(request_data, tmp_path):
    files = list(map(Path, request_data["files"]))
    with pytest.raises(ValueError, match="标明"):
        detect_s5_period(files, Path(request_data["input_dir"]))
    assert detect_s5_period(files, tmp_path / "202608" / "本期") == ("2026-08-31", "文件夹“202608”（仅年月时按月末）")
    assert detect_s5_period(files, tmp_path / "2026-08-31" / "本期") == ("2026-08-31", "文件夹“2026-08-31”")


def test_automatic_inputs_resolve_one_business_and_unique_template(request_data, tmp_path):
    templates = tmp_path / "模板目录"
    templates.mkdir()
    Path(request_data["template_path"]).rename(templates / "该业务模板.xlsx")
    form = {"selected_files": request_data["files"], "config_path": request_data["config_path"],
            "template_dir": str(templates), "input_dir": str(tmp_path / "202608")}
    result = recognize_s5(form, tmp_path / "宿主数据")
    assert result["errors"] == {}
    assert form["business"] == "合成业务"
    assert form["period"] == "2026-08-31"
    assert form["template_path"] == str((templates / "该业务模板.xlsx").resolve())
    form["selected_files"] = []
    recognize_s5(form, tmp_path / "宿主数据")
    assert not form["period"] and not form["business"] and not form["template_path"]


def test_template_tie_is_not_selected_arbitrarily(request_data, tmp_path):
    templates = tmp_path / "templates"
    templates.mkdir()
    make_source(templates / "A.xlsx", named=True)
    make_source(templates / "B.xlsx", named=True)
    with pytest.raises(ValueError, match="多个相近模板"):
        match_s5_template(request_data["files"], templates, tmp_path / "index.json")


def test_multiple_businesses_require_evidence_instead_of_first_rule(request_data, tmp_path):
    book = load_workbook(request_data["config_path"])
    sheet = book["审核规则"]
    row = [cell.value for cell in sheet[2]]
    row[0], row[2] = "R2", "另一个业务"
    sheet.append(row)
    book.save(request_data["config_path"])
    book.close()
    form = {"selected_files": request_data["files"], "config_path": request_data["config_path"],
            "template_path": request_data["template_path"], "template_manual": True, "period": "2026-08-31", "period_manual": True}
    result = recognize_s5(form, tmp_path)
    assert "无法唯一识别业务" in result["errors"]["business"]




def test_s5_defaults_use_own_config_and_output_pin(request_data, tmp_path):
    class Host(S5WebApiMixin):
        project_root = tmp_path
        state = {}
        def _log(self, *args, **kwargs):
            pass
    host = Host()
    host._init_s5()
    host.update_s5({"input_dir": request_data["input_dir"]})
    assert host.state["s5"]["output_dir"] == str(Path(request_data["input_dir"]) / "执行结果")
    pinned = str(tmp_path / "固定输出")
    host.update_s5({"output_dir": pinned})
    assert host.state["s5"]["output_pinned"]
    second = tmp_path / "另一批"
    second.mkdir()
    host.update_s5({"input_dir": str(second)})
    assert host.state["s5"]["output_dir"] == pinned


def test_s5_initialization_and_path_updates_do_not_auto_recognize(request_data, tmp_path):
    class Host(S5WebApiMixin):
        project_root = tmp_path
        state = {}

        def _log(self, *args, **kwargs):
            pass

        def _s5_recognize(self):
            raise AssertionError("recognition must require an explicit user action")

    host = Host()
    host._init_s5()
    host.update_s5({"input_dir": request_data["input_dir"],
                    "template_dir": str(Path(request_data["template_path"]).parent),
                    "config_path": request_data["config_path"]})
    assert host.state["s5"]["template_path"] == ""
    assert host.state["s5"]["period"] == ""
    assert host.state["s5"]["business"] == ""


def test_s5_manual_template_can_be_preserved_while_identifying_period_and_business(request_data, tmp_path):
    source_dir = Path(request_data["input_dir"])
    dated = make_source(source_dir / "合成银行_20260831.xlsx")

    class Host(S5WebApiMixin):
        project_root = tmp_path
        state = {}

        def _log(self, *args, **kwargs):
            pass

    host = Host()
    host._init_s5()
    host.update_s5({"input_dir": str(source_dir), "template_dir": str(Path(request_data["template_path"]).parent),
                    "config_path": request_data["config_path"]})
    host.update_s5({"selected_files": [str(dated.resolve())],
                    "template_path": request_data["template_path"]})
    host.recognize_s5_inputs(preserve_template=True)
    assert host.state["s5"]["template_path"] == str(Path(request_data["template_path"]).resolve())
    assert host.state["s5"]["business"] == "合成业务"
    assert host.state["s5"]["period"] == "2026-08-31"


def test_s5_file_list_classifies_template_business_and_date_per_file(request_data, tmp_path):
    source = make_source(tmp_path / "合成银行_2026-08-31.xlsx")
    partial = make_source(tmp_path / "部分匹配_20260831.xlsx")
    book = load_workbook(partial)
    book.create_sheet("额外工作表")
    book.save(partial)
    book.close()
    expanded_template = tmp_path / "模板含其他工作表.xlsx"
    book = load_workbook(request_data["template_path"])
    book.create_sheet("本期未报送的模板表")
    book.save(expanded_template)
    book.close()
    unmatched = tmp_path / "未匹配_20260831.xlsx"
    book = Workbook()
    book.active.title = "无关工作表"
    book.save(unmatched)
    book.close()

    rows = classify_s5_files(
        [source, partial, unmatched],
        template_path=expanded_template,
        config_path=request_data["config_path"],
        input_dir=tmp_path,
        selected_files=[source],
        business="合成业务",
    )
    by_name = {row["name"]: row for row in rows}
    assert by_name[source.name]["match_status"] == "matched"
    assert by_name[partial.name]["match_status"] == "partial"
    assert by_name[unmatched.name]["match_status"] == "unmatched"
    assert "不影响匹配" in by_name[source.name]["match_detail"]
    assert "源工作表匹配" in by_name[source.name]["match_detail"]
    assert by_name[source.name]["matched_sheets"] == by_name[source.name]["source_sheet_count"]
    assert by_name[source.name]["business"] == "合成业务"
    assert by_name[source.name]["period"] == "2026-08-31"
    assert by_name[partial.name]["period"] == "2026-08-31"


def test_s5_template_recognition_accepts_source_sheet_subset(request_data, tmp_path):
    templates = tmp_path / "templates"
    templates.mkdir()
    source = Path(request_data["files"][0])
    template = templates / "综合模板.xlsx"
    book = load_workbook(request_data["template_path"])
    for index in range(20):
        book.create_sheet(f"模板保留的其他表{index + 1}")
    book.save(template)
    book.close()

    selected = match_s5_template([source], templates, tmp_path / "index.json")

    assert Path(selected) == template


def test_s5_file_period_filter_uses_filename_and_parent_folder_without_reading_cells(tmp_path):
    root = tmp_path / "2026-08-31" / "本期"
    root.mkdir(parents=True)
    source = root / "没有日期的文件.xlsx"
    source.touch()
    assert file_period_hint(source, root) == "2026-08-31"


def test_s5_migrates_saved_legacy_default_output_folder(tmp_path):
    source = tmp_path / "报送文件"
    source.mkdir()
    settings = tmp_path / "data" / "s5" / "界面设置.json"
    settings.parent.mkdir(parents=True)
    settings.write_text(json.dumps({
        "input_dir": str(source),
        "output_dir": str(source / "S5_执行结果"),
        "output_pinned": False,
    }), encoding="utf-8")

    class Host(S5WebApiMixin):
        project_root = tmp_path
        state = {}

        def _log(self, *args, **kwargs):
            pass

    host = Host()
    host._init_s5()
    assert host.state["s5"]["output_dir"] == str(source / "执行结果")
