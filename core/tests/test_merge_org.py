from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from src.base_audit.merge_org import (
    _format_merged_sheet_sources,
    _copy_sheet_named_ranges_as_local,
    _delete_workbook_scoped_name,
    _localize_workbook_named_ranges,
    _name_scope_token,
    _names_for_source_sheet,
    group_key_from_plan,
    _organisation_from_source,
    _merge_one_org_openpyxl,
    _output_folder,
    _rebase_formula_to_local_sheets,
    _safe_sheet_name,
    _selected_sources,
    _source_period,
    _template_merge_output,
    _write_template_merge_report,
)


def _workbook(path: Path) -> None:
    book = Workbook()
    book.save(path)
    book.close()


def test_organisation_key_is_first_filename_segment() -> None:
    path = Path("示例财务有限公司_金融基础数据-单位贷款_B00_2026-07-31_在线核查表.xlsx")
    assert _organisation_from_source(path) == "示例财务有限公司"
    assert _source_period(path) == "2026-07-31"


def test_configurable_group_key_can_use_regex_or_name_matching() -> None:
    path = Path("示例财务有限公司_金融基础数据-单位贷款_B00_2026-07-31_在线核查表.xlsx")
    assert group_key_from_plan(path, {
        "mode": "regex", "pattern": r"(?P<组合>.+?)_金融基础数据-",
    }) == "示例财务有限公司"
    assert group_key_from_plan(path, {
        "mode": "name", "groups": [{"name": "贷款组合", "keywords": "个人贷款、单位贷款"}],
    }) == "贷款组合"


def test_configurable_group_key_reports_unmatched_file_clearly() -> None:
    path = Path("未分类文件.xlsx")
    try:
        group_key_from_plan(path, {"mode": "regex", "pattern": r"(?P<组合>.+?)_金融基础数据-"})
    except ValueError as exc:
        assert "未匹配正则" in str(exc)
    else:
        raise AssertionError("没有命中分组规则的文件应明确失败")


def test_merged_sheet_log_groups_sheets_by_source_workbook() -> None:
    result = _format_merged_sheet_sources((
        "机构A_个人贷款.xlsx｜存量个人贷款信息",
        "机构A_个人贷款.xlsx｜个人客户基础信息",
        "机构A_单位贷款.xlsx｜存量单位贷款信息",
    ))
    assert result == (
        "机构A_个人贷款.xlsx（存量个人贷款信息、个人客户基础信息）\n"
        "机构A_单位贷款.xlsx（存量单位贷款信息）"
    )


def test_sheet_prefix_is_sanitized_and_deduplicated() -> None:
    existing: set[str] = set()
    assert _safe_sheet_name("单位贷款_明细/表", existing) == "单位贷款_明细_表"
    assert _safe_sheet_name("单位贷款_明细/表", existing) == "单位贷款_明细_表_2"
    assert len(_safe_sheet_name("A" * 40, existing)) == 31


def test_empty_selection_means_all_source_files() -> None:
    with TemporaryDirectory() as folder:
        root = Path(folder)
        first = root / "机构A_表一_B00_2026-07-31_在线核查表.xlsx"
        second = root / "机构B_表二_B00_2026-07-31_在线核查表.xlsx"
        _workbook(first)
        _workbook(second)
        assert _selected_sources(root, [], recursive=True) == [first, second]


def test_output_folder_uses_requested_prefix_and_rejects_source_root() -> None:
    with TemporaryDirectory() as folder:
        root = Path(folder)
        result = _output_folder(root, root / "执行结果", "组合联合核查表")
        assert result.parent == root / "执行结果"
        assert result.name.startswith("组合联合核查表_")
        try:
            _output_folder(root, root, "组合联合核查表")
        except ValueError as exc:
            assert "不能与源数据目录相同" in str(exc)
        else:
            raise AssertionError("源数据目录应被拒绝作为输出目录")


def test_template_merge_always_creates_a_separate_copy_next_to_base() -> None:
    with TemporaryDirectory() as folder:
        base = Path(folder) / "！金融基础数据-单位贷款202605.xlsx"
        personal = Path(folder) / "！金融基础数据-个人贷款202605.xlsx"
        _workbook(base)
        _workbook(personal)
        output, report = _template_merge_output(base, [personal])
        assert output.parent == base.parent / "联合模板"
        assert output != base
        assert output.name.startswith("！联合模板_单位贷款+个人贷款_")
        assert output.suffix == ".xlsx"
        assert report.name.startswith("！联合模板_单位贷款+个人贷款_检查报告_")


def test_template_merge_report_includes_base_template_sheets() -> None:
    with TemporaryDirectory() as folder:
        root = Path(folder)
        base = root / "！个人贷款模板.xlsx"
        report = root / "联合模板检查报告.xlsx"
        _workbook(base)
        _write_template_merge_report(
            report,
            base_template=base,
            base_sheets=["集中系统数据", "参照表"],
            sources=[],
            copied=[],
            skipped=[],
            findings=[],
        )
        book = load_workbook(report, read_only=True, data_only=True)
        try:
            rows = list(book["工作表处理清单"].iter_rows(min_row=2, values_only=True))
            assert (base.name, "集中系统数据", "保留为底稿") in rows
            assert (base.name, "参照表", "保留为底稿") in rows
            assert "命名区域处理清单" in book.sheetnames
        finally:
            book.close()


def test_template_merge_converts_same_workbook_names_to_sheet_local_names() -> None:
    class Area:
        def __init__(self, sheet_name: str, address: str) -> None:
            self.Worksheet = type("SheetRef", (), {"Name": sheet_name})()
            self.Address = address

    class Name:
        def __init__(self, raw_name: str, sheet_name: str, address: str) -> None:
            self.Name = raw_name
            self.RefersTo = f"='{sheet_name}'!{address}"
            self.RefersToRange = type("Range", (), {"Areas": [Area(sheet_name, address)]})()

    class Names:
        def __init__(self, values=()) -> None:
            self.values = list(values)
            self.added: list[tuple[str, str]] = []
            self.Count = len(self.values)

        def Item(self, index: int):
            return self.values[index - 1]

        def Add(self, name: str, refers_to: str) -> None:
            self.added.append((name, refers_to))

    source_names = Names((
        Name("来源模板.xlsx!_xlfn.IFERROR", "贷款明细", "$A$1"),
        Name("来源模板.xlsx!表结构区域", "贷款明细", "$A$1:$D$2"),
        Name("来源模板.xlsx!校验区域", "贷款明细", "$M$2:$M$20"),
        Name("来源模板.xlsx!表结构区域_担保", "担保明细", "$A$1:$C$2"),
        Name("'担保明细'!局部区域", "担保明细", "$F$1"),
        Name("来源模板.xlsx!贷款明细!局部校验区域", "贷款明细", "$N$2:$N$20"),
        Name("'贷款明细'!跨表别名", "担保明细", "$G$1"),
    ))
    source_book = type("Book", (), {"Name": "来源模板.xlsx", "Names": source_names})()
    source_sheet = type("Sheet", (), {"Name": "贷款明细"})()
    target_names = Names()
    target_book = type("Book", (), {"Name": "联合模板.xlsx", "Names": target_names})()
    target_sheet = type("Sheet", (), {"Name": "贷款明细"})()

    records = _copy_sheet_named_ranges_as_local(source_book, source_sheet, target_book, target_sheet)

    assert target_names.added == [
        ("表结构区域_贷款明细", "='贷款明细'!$A$1:$D$2"),
        ("校验区域_贷款明细", "='贷款明细'!$M$2:$M$20"),
        ("局部校验区域_贷款明细", "='贷款明细'!$N$2:$N$20"),
    ]
    assert all(item[5] == "已复制" for item in records)


def test_named_range_copy_rejects_error_address_instead_of_claiming_success() -> None:
    class Area:
        Worksheet = type("SheetRef", (), {"Name": "贷款明细"})()
        Address = "#VALUE!"

    class Name:
        Name = "来源模板.xlsx!表结构区域"
        RefersTo = "=错误引用"
        RefersToRange = type("Range", (), {"Areas": [Area()]})()

    class Names:
        Count = 1
        def Item(self, _index: int): return Name()
        def Add(self, _name: str, _refers_to: str): raise AssertionError("不得写入错误地址")

    source_book = type("Book", (), {"Name": "来源模板.xlsx", "Names": Names()})()
    source_sheet = type("Sheet", (), {"Name": "贷款明细"})()
    target_book = type("Book", (), {"Name": "联合模板.xlsx", "Names": Names()})()
    target_sheet = type("Sheet", (), {"Name": "贷款明细"})()
    records = _copy_sheet_named_ranges_as_local(source_book, source_sheet, target_book, target_sheet)
    assert records[0][5].startswith("复制失败：无法取得有效单元格地址：#VALUE!")


def test_workbook_scope_token_supports_full_path_excel_qualifier() -> None:
    assert _name_scope_token("'C:\\模板\\[来源模板.xlsx]'") == "来源模板.xlsx"
    assert _name_scope_token("[来源模板.xlsx]") == "来源模板.xlsx"
    assert _name_scope_token("贷款明细") == "贷款明细"


def test_localizing_uses_explicit_workbook_name_delete_before_fallback() -> None:
    class OriginalName:
        deleted = False
        def Delete(self): self.deleted = True

    class ExplicitName:
        deleted = False
        def Delete(self): self.deleted = True

    explicit = ExplicitName()
    names = type("Names", (), {"Item": lambda self, key: explicit if key == "模板.xlsx!区域" else None})()
    book = type("Book", (), {"Names": names})()
    fallback = OriginalName()
    _delete_workbook_scoped_name(book, "模板.xlsx!区域", fallback)
    assert explicit.deleted is True
    assert fallback.deleted is False


def test_copy_sheet_names_reads_worksheet_local_names_missing_from_workbook_collection() -> None:
    class Area:
        def __init__(self) -> None:
            self.Worksheet = type("SheetRef", (), {"Name": "个人贷款极端数据核查信息"})()
            self.Address = "$A$4:$C$15"

    class Name:
        Name = "个人贷款极端数据核查信息!表结构区域_007"
        RefersTo = "='个人贷款极端数据核查信息'!$A$4:$C$15"
        RefersToRange = type("Range", (), {"Areas": [Area()]})()

    class Names:
        def __init__(self, values): self.values, self.Count = list(values), len(values)
        def Item(self, index): return self.values[index - 1]
        def Add(self, name, refers_to): self.added.append((name, refers_to))

    source_book = type("Book", (), {"Name": "个人贷款.xlsx", "Names": Names([])})()
    source_sheet = type("Sheet", (), {"Name": "个人贷款极端数据核查信息", "Names": Names([Name()])})()
    target_names = Names([]); target_names.added = []
    target_book = type("Book", (), {"Name": "联合模板.xlsx", "Names": target_names})()
    target_sheet = type("Sheet", (), {"Name": "个人贷款极端数据核查信息"})()
    records = _copy_sheet_named_ranges_as_local(source_book, source_sheet, target_book, target_sheet)
    assert target_names.added == [("表结构区域_007_个人贷款极端数据核查信息", "='个人贷款极端数据核查信息'!$A$4:$C$15")]
    assert records[0][5] == "已复制"


def test_base_template_names_become_unique_workbook_names_with_sheet_suffix() -> None:
    class Names:
        def __init__(self) -> None:
            self.added: list[tuple[str, str]] = []
            self.values: list[Name] = []
            self.Count = 0

        def Add(self, name: str, refers_to: str) -> None:
            self.added.append((name, refers_to))

        def Item(self, index: int) -> "Name":
            return self.values[index - 1]

    class Sheet:
        def __init__(self, name: str) -> None:
            self.Name = name
            self.Names = Names()

    class Area:
        def __init__(self, sheet: Sheet, address: str) -> None:
            self.Worksheet = sheet
            self.Address = address

    class Name:
        def __init__(self, raw_name: str, areas: list[Area]) -> None:
            self.Name = raw_name
            self.RefersTo = "=测试"
            self.RefersToRange = type("Range", (), {"Areas": areas})()
            self.deleted = False

        def Delete(self) -> None:
            self.deleted = True

    loan = Sheet("贷款明细")
    guarantee = Sheet("担保明细")
    structure = Name(
        "联合底稿.xlsx!表结构区域",
        [Area(loan, "$A$1:$D$2"), Area(guarantee, "$A$1:$C$2")],
    )
    formula = Name("联合底稿.xlsx!校验区域", [Area(loan, "$M$2:$M$20")])
    existing_local = Name("联合底稿.xlsx!贷款明细!已有局部区域", [Area(loan, "$F$1")])
    book_names = Names()
    book_names.values = [structure, formula, existing_local]
    book_names.Count = len(book_names.values)
    worksheets = type("Worksheets", (), {
        "Count": 2, "__call__": lambda self, index: [loan, guarantee][index - 1]
    })()
    book = type("Book", (), {
        "Name": "联合底稿.xlsx", "Names": book_names,
        "Worksheets": worksheets,
    })()

    records = _localize_workbook_named_ranges(book)

    assert book_names.added == [
        ("表结构区域_贷款明细", "='贷款明细'!$A$1:$D$2"),
        ("表结构区域_担保明细", "='担保明细'!$A$1:$C$2"),
        ("校验区域_贷款明细", "='贷款明细'!$M$2:$M$20"),
        ("已有局部区域_贷款明细", "='贷款明细'!$F$1"),
    ]
    assert structure.deleted is True
    assert formula.deleted is True
    assert existing_local.deleted is True
    assert len(records) == 4
    assert all(item[5] == "已复制" for item in records)


def test_every_external_formula_is_rebased_to_a_local_sheet_reference() -> None:
    original = (
        "=SUMIFS('C:\\模板\\[个人贷款.xlsx]个人贷款发生额信息'!$E:$E,"
        "[1]个人贷款发生额信息!$D:$D,$B5)"
    )
    result = _rebase_formula_to_local_sheets(original)
    assert result == "=SUMIFS('个人贷款发生额信息'!$E:$E,'个人贷款发生额信息'!$D:$D,$B5)"


def test_openpyxl_merge_keeps_all_sheets_and_formulas() -> None:
    with TemporaryDirectory() as folder:
        root = Path(folder)
        first = root / "机构A_个人贷款_B00_2026-07-31_在线核查表.xlsx"
        second = root / "机构A_单位贷款_B00_2026-07-31_在线核查表.xlsx"
        for path, sheet_name, hidden in ((first, "个人明细", False), (second, "单位明细", True)):
            book = Workbook()
            sheet = book.active
            sheet.title = sheet_name
            sheet["A1"] = "指标"
            sheet["B1"] = "金额"
            sheet["B2"] = "=1+1"
            sheet["A1"].font = Font(bold=True, color="FF0000")
            sheet["A1"].fill = PatternFill("solid", fgColor="FFFF00")
            sheet.column_dimensions["A"].width = 22
            sheet.row_dimensions[1].height = 24
            sheet.freeze_panes = "A2"
            sheet.merge_cells("A4:B4")
            if hidden:
                book.create_sheet("封面")
                sheet.sheet_state = "hidden"
            book.save(path)
            book.close()
        output = root / "机构A_合并_2026-07-31.xlsx"
        merged_sheets: list[str] = []
        assert _merge_one_org_openpyxl([first, second], output, merged_sheets) == 3
        assert f"{first.name}｜个人明细" in merged_sheets
        assert f"{second.name}｜单位明细" in merged_sheets
        merged = load_workbook(output, read_only=False, data_only=False)
        try:
            assert len(merged.worksheets) == 3
            assert merged["个人明细"]["B2"].value == "=1+1"
            assert "A4:B4" in {str(item) for item in merged["个人明细"].merged_cells.ranges}
            assert merged["个人明细"]["A1"].font.bold is True
            assert merged["个人明细"]["A1"].fill.fgColor.rgb == "00FFFF00"
            assert merged["个人明细"].column_dimensions["A"].width == 22
            assert merged["个人明细"].row_dimensions[1].height == 24
            assert merged["个人明细"].freeze_panes == "A2"
            assert merged["单位明细"].sheet_state == "hidden"
        finally:
            merged.close()


def test_openpyxl_merge_skips_later_duplicate_sheet_names() -> None:
    with TemporaryDirectory() as folder:
        root = Path(folder)
        first = root / "机构A_个人贷款_B00_2026-07-31_在线核查表.xlsx"
        second = root / "机构A_单位贷款_B00_2026-07-31_在线核查表.xlsx"
        for path, value in ((first, "首份"), (second, "后份")):
            book = Workbook()
            book.active.title = "集中系统数据"
            book.active["A1"] = value
            book.save(path)
            book.close()
        output = root / "机构A_合并_2026-07-31.xlsx"
        assert _merge_one_org_openpyxl([first, second], output) == 1
        merged = load_workbook(output, read_only=True)
        try:
            assert merged.sheetnames == ["集中系统数据"]
            assert merged["集中系统数据"]["A1"].value == "首份"
        finally:
            merged.close()
