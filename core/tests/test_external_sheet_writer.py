from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from base_audit.external_sheet_writer import (
    DirectOoxmlExternalSheetWriter,
    UnsupportedExternalSheetOperation,
)


def _audit(path: Path) -> None:
    book = Workbook()
    book.active.title = "报表"
    book.save(path)
    book.close()


def test_direct_ooxml_external_writer_adds_plain_data_sheets(tmp_path: Path) -> None:
    external = tmp_path / "外部.xlsx"
    book = Workbook()
    data = book.active
    data.title = "集中系统数据"
    data["A1"] = "机构"
    data["B2"] = 123.45
    data["C2"] = True
    reference = book.create_sheet("参照表")
    reference["A1"] = "代码"
    reference["B2"] = "A001"
    book.save(external)
    book.close()
    audit = tmp_path / "审核版.xlsx"
    _audit(audit)

    writer = DirectOoxmlExternalSheetWriter(external, ("集中系统数据", "参照表"))
    prepared = writer.prepare()
    written = writer.write(audit)

    assert prepared.sheet_count == 2
    assert written.cell_count == 5
    restored = load_workbook(audit, data_only=False)
    try:
        assert restored.sheetnames == ["报表", "集中系统数据", "参照表"]
        assert restored["集中系统数据"]["A1"].value == "机构"
        assert restored["集中系统数据"]["B2"].value == 123.45
        assert restored["集中系统数据"]["C2"].value is True
        assert restored["参照表"]["B2"].value == "A001"
    finally:
        restored.close()


def test_direct_ooxml_external_writer_rejects_formula_sheet(tmp_path: Path) -> None:
    external = tmp_path / "外部.xlsx"
    book = Workbook()
    book.active.title = "参照表"
    book.active["A1"] = "=1+1"
    book.save(external)
    book.close()

    with pytest.raises(UnsupportedExternalSheetOperation, match="含公式"):
        DirectOoxmlExternalSheetWriter(external, ("参照表",)).prepare()


def test_direct_ooxml_external_writer_explicitly_reports_dropped_comments(tmp_path: Path) -> None:
    external = tmp_path / "外部.xlsx"
    book = Workbook()
    book.active.title = "参照表"
    book.active["A1"] = "数据"
    book.active["A1"].comment = Comment("仅说明", "审核")
    book.save(external)
    book.close()

    prepared = DirectOoxmlExternalSheetWriter(external, ("参照表",)).prepare()
    assert "未复制：批注" in prepared.messages[0]
