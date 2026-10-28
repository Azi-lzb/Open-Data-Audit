"""外部文件添加的降级链：Direct OOXML → openpyxl → 记录失败但链路继续。

S1-F01 的 ``excel.external_sheet_copy`` 节点在 UOS/native 与 Windows/COM 两条
适配器上都必须能在 Direct OOXML 拒绝复杂外部表时继续添加工作表；两个写入器
都失败时也只能标记该文件失败，不能中止整条审核链。
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from base_audit.external import ExternalSheetPlan
from base_audit.workflow.adapters.native_audit_handlers import handle_external_copy
from base_audit.workflow.artifacts import FileSetArtifact, FileRecord
from base_audit.workflow.types import ArtifactSubtype, NodeRunStatus


def _audit_workbook(path: Path) -> None:
    book = Workbook()
    book.active.title = "报表"
    book.active["A1"] = "数据"
    book.save(path)
    book.close()


def _plain_external(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "集中系统数据"
    sheet["A1"] = "机构"
    sheet["B2"] = 123.0
    sheet["C2"] = "文本"
    book.save(path)
    book.close()


def _complex_external(path: Path) -> None:
    """模拟真实业务外部表：合并单元格 + 公式，Direct OOXML 明确不支持。"""
    book = Workbook()
    sheet = book.active
    sheet.title = "参照表"
    sheet["A1"] = "机构代码"
    sheet["B1"] = "机构名称"
    sheet.merge_cells("A1:B1")
    sheet["A2"] = "A001"
    sheet["B2"] = '=CONCAT("机构",A2)'
    book.save(path)
    book.close()


def _mixed_external(path: Path) -> None:
    """纯数据表与公式表共存，贴近真实外部文件结构。"""
    book = Workbook()
    plain = book.active
    plain.title = "集中系统数据"
    plain["A1"] = "机构"
    plain["B2"] = 456.0
    formula = book.create_sheet("参照表")
    formula["A1"] = "代码"
    formula["A2"] = "A001"
    formula["D2"] = "=VLOOKUP(A2,参照表!A:B,2,0)"
    book.save(path)
    book.close()


def _context(tmp_path: Path) -> SimpleNamespace:
    return SimpleNamespace(
        log=lambda message: None,
        run=SimpleNamespace(services={"service": SimpleNamespace(external_sheet_writer="DIRECT_OOXML")}),
    )


def _params(tmp_path: Path, external: Path) -> dict:
    return {
        "stage_dir": str(tmp_path / "stage"),
        "external_path": str(external),
        "feature_name": "外部文件添加",
        "flow_name": "S1-F01",
    }


def _inputs(audit: Path, sheets: tuple[str, ...]) -> dict:
    record = FileRecord(
        source=str(audit),
        audit=str(audit),
        org_code="机构",
        org_name="机构",
        version=0,
        stage_paths=(str(audit),),
    )
    return {
        "template": SimpleNamespace(payload={"external_plan": ExternalSheetPlan(sheets, "公式识别")}),
        "workbooks": FileSetArtifact(records=(record,), subtype=ArtifactSubtype.EXCEL_WORKBOOK_SET),
    }


def test_plain_external_uses_direct_ooxml_fast_path(tmp_path: Path) -> None:
    external = tmp_path / "外部纯数据.xlsx"
    _plain_external(external)
    audit = tmp_path / "审核.xlsx"
    _audit_workbook(audit)

    result = handle_external_copy(
        _context(tmp_path), _params(tmp_path, external), _inputs(audit, ("集中系统数据",)),
    )

    assert result.status is NodeRunStatus.SUCCEEDED
    assert result.metrics["external_openpyxl_fallback_writes"] == 0
    assert result.metrics["external_failed_files"] == 0
    written = Path(result.outputs["workbooks"].records[0].audit)
    book = load_workbook(written, data_only=False)
    try:
        assert "集中系统数据" in book.sheetnames
    finally:
        book.close()


def test_complex_external_falls_back_to_openpyxl_instead_of_aborting(tmp_path: Path) -> None:
    """合并单元格/公式曾让 native 节点直接抛错，整条 S1-F01 链被中止。"""
    external = tmp_path / "外部复杂.xlsx"
    _complex_external(external)
    audit = tmp_path / "审核.xlsx"
    _audit_workbook(audit)

    result = handle_external_copy(
        _context(tmp_path), _params(tmp_path, external), _inputs(audit, ("参照表",)),
    )

    assert result.status is NodeRunStatus.SUCCEEDED
    assert result.metrics["external_failed_files"] == 0
    written = Path(result.outputs["workbooks"].records[0].audit)
    book = load_workbook(written, data_only=False)
    try:
        assert book.sheetnames == ["报表", "参照表"]
        sheet = book["参照表"]
        assert [str(item) for item in sheet.merged_cells.ranges] == ["A1:B1"]
        assert sheet["B2"].value == '=CONCAT("机构",A2)'
    finally:
        book.close()


def test_mixed_plan_falls_back_for_whole_batch_and_keeps_formula(tmp_path: Path) -> None:
    """计划同时含纯数据表与公式表时，快路径整批失败，转 openpyxl 后两张表都要在。"""
    external = tmp_path / "外部混合.xlsx"
    _mixed_external(external)
    audit = tmp_path / "审核.xlsx"
    _audit_workbook(audit)

    result = handle_external_copy(
        _context(tmp_path), _params(tmp_path, external),
        _inputs(audit, ("集中系统数据", "参照表")),
    )

    assert result.status is NodeRunStatus.SUCCEEDED
    assert result.metrics["external_failed_files"] == 0
    written = Path(result.outputs["workbooks"].records[0].audit)
    book = load_workbook(written, data_only=False)
    try:
        assert book.sheetnames == ["报表", "集中系统数据", "参照表"]
        assert book["集中系统数据"]["B2"].value == 456.0
        assert book["参照表"]["D2"].value == "=VLOOKUP(A2,参照表!A:B,2,0)"
    finally:
        book.close()


def test_external_copy_keeps_chain_when_both_writers_fail(tmp_path: Path) -> None:
    """openpyxl 也失败时只标记该文件失败，审核副本保留，节点降级为 PARTIAL。"""
    missing = tmp_path / "不存在的外部.xlsx"
    audit = tmp_path / "审核.xlsx"
    _audit_workbook(audit)

    result = handle_external_copy(
        _context(tmp_path), _params(tmp_path, missing), _inputs(audit, ("参照表",)),
    )

    assert result.status is NodeRunStatus.PARTIAL
    assert result.metrics["external_failed_files"] == 1
    assert len(result.partial_failures) == 1
    record = result.outputs["workbooks"].records[0]
    assert record.stage_errors, "失败必须进入文件级结果，不能静默吞掉"
    assert "外部文件添加" in record.stage_errors[0]
    assert Path(record.audit).exists(), "审核副本必须保留以便链路继续"
