"""Read-only preflight coverage for the S3 3.0 common configuration."""

from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from config_paths import resolve_test_config

from base_audit.systems.s3_central_statistics.common_config_check import check_common_config
from base_audit.systems.s3_central_statistics.config import (
    ALERT_SHEET,
    COMMON_CONFIG_NAME,
    ORG_REFERENCE_SHEET,
    RUN_SHEET,
    UNIT_EXCEPTION_SHEET,
    write_default_central_config,
)


FORMAL_CONFIG = resolve_test_config(COMMON_CONFIG_NAME)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _make_config(path: Path) -> Path:
    return write_default_central_config(path)


def test_formal_common_config_passes_and_is_read_only() -> None:
    if not FORMAL_CONFIG.is_file():
        pytest.skip("缺少正式 3.0 通用配置工作簿")
    before = _sha256(FORMAL_CONFIG)

    report = check_common_config(FORMAL_CONFIG)

    assert report["passed"] is True
    assert report["errors"] == []
    assert report["warnings"] == []
    assert report["path"] == str(FORMAL_CONFIG.resolve())
    assert [item["sheet"] for item in report["sheet_stats"]] == [
        RUN_SHEET, ALERT_SHEET, UNIT_EXCEPTION_SHEET, ORG_REFERENCE_SHEET,
    ]
    assert _sha256(FORMAL_CONFIG) == before


@pytest.mark.parametrize(
    "sheet_name",
    [RUN_SHEET, ALERT_SHEET, UNIT_EXCEPTION_SHEET, ORG_REFERENCE_SHEET],
)
def test_missing_required_sheet_is_reported(tmp_path: Path, sheet_name: str) -> None:
    path = _make_config(tmp_path / "common.xlsx")
    workbook = load_workbook(path)
    del workbook[sheet_name]
    workbook.save(path)
    workbook.close()

    report = check_common_config(path)

    assert report["passed"] is False
    assert any(sheet_name in message for message in report["errors"])
    item = next(stat for stat in report["sheet_stats"] if stat["sheet"] == sheet_name)
    assert item["present"] is False


@pytest.mark.parametrize("unit", ["公斤", "", "元/万元"])
def test_unsupported_run_parameter_unit_is_an_error(tmp_path: Path, unit: str) -> None:
    path = _make_config(tmp_path / "common.xlsx")
    workbook = load_workbook(path)
    sheet = workbook[RUN_SHEET]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "默认源数据单位":
            sheet.cell(row, 2).value = unit
            break
    else:
        raise AssertionError("default source unit row was not generated")
    workbook.save(path)
    workbook.close()

    report = check_common_config(path)

    assert report["passed"] is False
    assert any("默认源数据单位" in message and "只允许" in message for message in report["errors"])


def test_alert_interval_gap_is_a_warning_not_an_error(tmp_path: Path) -> None:
    path = _make_config(tmp_path / "common.xlsx")
    workbook = load_workbook(path)
    sheet = workbook[ALERT_SHEET]
    # First interval ends at -99; start the next interval at -98 to create a gap.
    sheet.cell(3, 2).value = -98
    workbook.save(path)
    workbook.close()

    report = check_common_config(path)

    assert report["passed"] is True
    assert report["errors"] == []
    assert any("存在缺口" in message and ALERT_SHEET in message for message in report["warnings"])


def test_invalid_soft_color_is_reported_before_runtime(tmp_path: Path) -> None:
    path = _make_config(tmp_path / "common.xlsx")
    workbook = load_workbook(path)
    sheet = workbook[RUN_SHEET]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "复杂校验软性颜色":
            sheet.cell(row, 2).value = "蓝色"
            break
    workbook.save(path)
    workbook.close()

    report = check_common_config(path)

    assert report["passed"] is False
    assert any("复杂校验软性颜色" in message and "整数" in message for message in report["errors"])


def test_flexible_sheets_are_checked_for_presence_only(tmp_path: Path) -> None:
    path = _make_config(tmp_path / "common.xlsx")
    workbook = load_workbook(path)
    for sheet_name in (UNIT_EXCEPTION_SHEET, ORG_REFERENCE_SHEET):
        sheet = workbook[sheet_name]
        sheet.cell(1, 1).value = "用户自定义表头"
        sheet.cell(2, 1).value = "任意结构内容"
    workbook.save(path)
    workbook.close()

    report = check_common_config(path)

    assert report["passed"] is True
    assert report["errors"] == []
    assert report["warnings"] == []
    flexible_stats = {
        item["sheet"]: item for item in report["sheet_stats"]
        if item["sheet"] in (UNIT_EXCEPTION_SHEET, ORG_REFERENCE_SHEET)
    }
    assert all(item["present"] is True for item in flexible_stats.values())
    assert all(item["scope"] == "仅检查工作表是否存在" for item in flexible_stats.values())
