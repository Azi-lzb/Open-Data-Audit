"""Read-only preflight coverage for the S3 3.3 forms configuration."""

from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from config_paths import resolve_test_config

from base_audit.systems.s3_central_statistics.forms_config_check import check_forms_config
from base_audit.web_app import WebApi


REPO_ROOT = Path(__file__).resolve().parents[2]
COMMON_CONFIG = resolve_test_config("3.0大集中通用配置.xlsx")
FORMAL_CONFIG = resolve_test_config("3.3大集中指标比较拆分_配置.xlsx")
FORMAL_CONFIGS_READY = COMMON_CONFIG.is_file() and FORMAL_CONFIG.is_file()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _make_config(path: Path, *, include_settings: bool = True, include_form: bool = True) -> Path:
    book = Workbook()
    listing = book.active
    listing.title = "报表清单"
    listing.append(("报表代码", "报表名称", "频度", "批次"))
    listing.append(("A1411", "测试报表", "月", 1))

    if include_form:
        form = book.create_sheet("A1411")
        form.append(("指标代码", "指标名称", "行序号", "本期数据"))
        form.append(("TEST001", "测试指标", 1, None))

    if include_settings:
        settings = book.create_sheet("转表设置")
        settings.append(("参数", "值", "说明"))
        settings.append(("隐藏空行", "是", ""))
        settings.append(("空表删除", "是", ""))
        settings.append(("分析文件数据单位", "亿元", ""))

    book.save(path)
    book.close()
    return path


def _api_for_configs(common_path: Path, forms_path: Path) -> WebApi:
    """Build only the WebApi state used by the 3.3 preflight endpoint."""
    api = WebApi.__new__(WebApi)
    api.project_root = REPO_ROOT
    api.state = {
        "busy": False,
        "centralCommonConfig": str(common_path),
        "centralFormsConfig": str(forms_path),
        "log": [],
    }
    api._log_detail = lambda _text: None
    api._refresh_config_issues = lambda: None
    return api


@pytest.mark.skipif(not FORMAL_CONFIG.is_file(), reason="缺少正式 3.3 配置工作簿")
def test_formal_config_passes_and_is_not_modified() -> None:
    before = _sha256(FORMAL_CONFIG)

    report = check_forms_config(FORMAL_CONFIG)

    assert report["passed"] is True
    assert report["errors"] == []
    assert {check["id"] for check in report["checks"]} == {
        "file_read", "required_sheets", "settings", "embedded_template", "form_references"
    }
    assert any(stat["name"] == "转表设置" for stat in report["sheet_stats"])
    json.dumps(report, ensure_ascii=False)
    assert _sha256(FORMAL_CONFIG) == before


@pytest.mark.skipif(not FORMAL_CONFIGS_READY, reason="缺少正式 3.0/3.3 配置工作簿")
def test_webapi_formal_two_book_report_names_sources_and_preserves_hashes() -> None:
    before = (_sha256(COMMON_CONFIG), _sha256(FORMAL_CONFIG))

    result = _api_for_configs(COMMON_CONFIG, FORMAL_CONFIG).check_central_forms_config()

    assert result["passed"] is True
    assert result["source_workbooks"] == [str(COMMON_CONFIG), str(FORMAL_CONFIG)]
    assert result["common_report"]["errors"] == []
    assert result["forms_report"]["errors"] == []
    report_text = result["report_text"]
    assert "指标比较拆分（3.3）" in report_text
    assert COMMON_CONFIG.name in report_text
    assert FORMAL_CONFIG.name in report_text
    assert "内嵌金融表单模板可识别" in report_text
    assert "不逐格检查灵活表单模板" in report_text
    assert (_sha256(COMMON_CONFIG), _sha256(FORMAL_CONFIG)) == before


@pytest.mark.skipif(not FORMAL_CONFIGS_READY, reason="缺少正式 3.0/3.3 配置工作簿")
def test_webapi_form_report_renders_setting_and_missing_template_errors(tmp_path: Path) -> None:
    common_copy = tmp_path / COMMON_CONFIG.name
    forms_copy = tmp_path / FORMAL_CONFIG.name
    shutil.copy2(COMMON_CONFIG, common_copy)
    shutil.copy2(FORMAL_CONFIG, forms_copy)

    book = load_workbook(forms_copy)
    try:
        listing = book["报表清单"]
        form_code = next(
            str(listing.cell(row, 1).value).strip()
            for row in range(2, listing.max_row + 1)
            if str(listing.cell(row, 1).value or "").strip().startswith("A")
            and str(listing.cell(row, 1).value).strip() in book.sheetnames
        )
        settings = book["转表设置"]
        setting_row = next(
            row for row in range(2, settings.max_row + 1)
            if settings.cell(row, 1).value == "分析文件数据单位"
        )
        settings.cell(setting_row, 2, "千元")
        del book[form_code]
        book.save(forms_copy)
    finally:
        book.close()

    result = _api_for_configs(common_copy, forms_copy).check_central_forms_config()

    assert result["passed"] is False
    assert result["total_issues"] >= 2
    assert "必须为元、万元或亿元" in result["report_text"]
    assert form_code in result["report_text"]
    assert "未对应到内嵌表单工作表" in result["report_text"]
    assert "第" in result["report_text"]


def test_missing_required_sheet_is_reported(tmp_path: Path) -> None:
    path = _make_config(tmp_path / "missing-settings.xlsx", include_settings=False)

    report = check_forms_config(path)

    assert report["passed"] is False
    assert any(
        issue.get("sheet") == "转表设置" and "缺少工作表" in issue["message"]
        for issue in report["errors"]
    )


@pytest.mark.parametrize(
    ("setting_key", "bad_value", "expected_fragment"),
    [
        ("分析文件数据单位", "千元", "元、万元或亿元"),
        ("隐藏空行", "true", "必须为“是”或“否”"),
        ("空表删除", "1", "必须为“是”或“否”"),
    ],
)
def test_invalid_front_checkable_settings_are_reported(
    tmp_path: Path, setting_key: str, bad_value: str, expected_fragment: str
) -> None:
    path = _make_config(tmp_path / "bad-setting.xlsx")
    book = load_workbook(path)
    try:
        sheet = book["转表设置"]
        setting_rows = {
            str(sheet.cell(row, 1).value): row for row in range(2, sheet.max_row + 1)
        }
        sheet.cell(setting_rows[setting_key], 2, bad_value)
        book.save(path)
    finally:
        book.close()

    report = check_forms_config(path)

    assert report["passed"] is False
    assert any(
        issue.get("field") == setting_key and expected_fragment in issue["message"]
        for issue in report["errors"]
    )


def test_missing_referenced_embedded_form_sheet_is_reported(tmp_path: Path) -> None:
    path = _make_config(tmp_path / "missing-form-sheet.xlsx", include_form=False)

    report = check_forms_config(path)

    assert report["passed"] is False
    assert any(
        issue.get("sheet") == "报表清单"
        and issue.get("row") == 2
        and "A1411" in issue["message"]
        and "未对应" in issue["message"]
        for issue in report["errors"]
    )
    assert any("内嵌金融表单模板" in issue["message"] for issue in report["errors"])


def test_missing_and_duplicate_required_settings_are_reported(tmp_path: Path) -> None:
    path = _make_config(tmp_path / "bad-settings.xlsx")
    book = load_workbook(path)
    try:
        sheet = book["转表设置"]
        sheet.cell(4, 1, "隐藏空行")
        sheet.cell(4, 2, "否")
        sheet.cell(5, 1, "自定义项")
        sheet.cell(5, 2, "x")
        book.save(path)
    finally:
        book.close()

    report = check_forms_config(path)

    assert any("转表设置项重复：隐藏空行" == issue["message"] for issue in report["errors"])
    assert any("缺少转表设置项：分析文件数据单位" == issue["message"] for issue in report["errors"])
