"""Read-only institution recognition for one source workbook at a time."""
from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string

from .errors import RetrievalError
from .external_layout import IDENTITY_LABELS, TARGET_LABELS, normalized
from .source_reader import file_hashing_enabled, file_sha256, read_book, shared_readers


CODE_LABELS = (*IDENTITY_LABELS, "社会信用代码", "社会统一信用代码")
NAME_LABELS = ("法人机构名称", "机构名称", "机构全称", "法人名称")
INSTITUTION_SHEETS = {"机构参照表", "机构参照"}


@dataclass(frozen=True)
class InstitutionEntry:
    code: str
    name: str
    report_item: str
    disabled: bool
    path: str
    sheet: str
    row: int
    code_column: int = 0
    name_column: int = 0
    target_column: int = 0


@dataclass(frozen=True)
class InstitutionIdentity:
    code: str
    name: str
    report_item: str
    method: str
    source_cells: tuple[str, ...]
    references: tuple[InstitutionEntry, ...]
    source_code: str = ""
    mapping_method: str = ""

    def to_dict(self) -> dict:
        return asdict(self)


class InstitutionCatalog:
    def __init__(self, paths=()):
        self.entries: list[InstitutionEntry] = []
        self.paths = list(dict.fromkeys(str(Path(p).resolve()) for p in paths if p))
        self.hashes: dict[str, str] = {}
        for path in self.paths:
            if file_hashing_enabled():
                self.hashes[path] = file_sha256(path)
            self._load(Path(path))
            if file_hashing_enabled() and file_sha256(path) != self.hashes[path]:
                raise RetrievalError("INPUT_CHANGED_DURING_RUN", "读取机构参照期间文件发生变化，请重新运行")

    def _load(self, path: Path) -> None:
        book = read_book(path, data_only=False)
        try:
            sheet_name = next((name for name in ("机构参照表", "机构参照") if name in book.sheetnames), "")
            for sheet in ([book[sheet_name]] if sheet_name else []):
                rows = sheet.iter_rows(values_only=True)
                for row_no, values in enumerate(rows, 1):
                    if row_no > 30:
                        break
                    labels = [normalized(value) for value in values]
                    codes = [i for i, label in enumerate(labels) if label in CODE_LABELS]
                    names = [i for i, label in enumerate(labels) if label in NAME_LABELS]
                    targets = [i for i, label in enumerate(labels) if label in TARGET_LABELS]
                    if not codes or not (names or targets):
                        continue
                    if len(codes) != 1 or len(names) > 1 or len(targets) > 1:
                        raise RetrievalError("REFERENCE_HEADERS_AMBIGUOUS", f"{path.name} / {sheet.title}：机构参照表头重复，请保留唯一代码、名称和报表项目字段")
                    disabled_col = next((i for i, label in enumerate(labels) if label == "禁用"), None)
                    for data_row, values in enumerate(rows, row_no + 1):
                        def text(index):
                            value = values[index] if index is not None and index < len(values) else None
                            if isinstance(value, str) and value.startswith("="):
                                raise RetrievalError("REFERENCE_FORMULA_UNVERIFIED", f"{path.name} / {sheet.title} 第 {data_row} 行：机构映射不能使用未验证的公式")
                            return str(value if value is not None else "").strip()
                        code = text(codes[0])
                        name = text(names[0]) if names else ""
                        target = text(targets[0]) if targets else ""
                        if not code and not name:
                            continue
                        flag = normalized(text(disabled_col))
                        if flag not in {"", "否", "0", "false", "no", "是", "1", "true", "yes", "禁用"}:
                            raise RetrievalError("REFERENCE_FLAG_INVALID", f"{path.name} / {sheet.title} 第 {data_row} 行：禁用字段应填写是或否")
                        self.entries.append(InstitutionEntry(code, name, target, flag in {"是", "1", "true", "yes", "禁用"}, str(path), sheet.title, data_row,
                                                             codes[0] + 1, names[0] + 1 if names else 0, targets[0] + 1 if targets else 0))
                    break
        finally:
            book.close()

    def by_code(self, code: str) -> list[InstitutionEntry]:
        return [entry for entry in self.entries if normalized(entry.code) == normalized(code)]

    def by_name(self, name: str) -> list[InstitutionEntry]:
        return [entry for entry in self.entries if normalized(entry.name) == normalized(name) and normalized(name)]

    def validate(self, entries: list[InstitutionEntry]) -> tuple[str, str, str]:
        codes = {normalized(entry.code): entry.code for entry in entries if entry.code}
        targets = {normalized(entry.report_item): entry.report_item for entry in entries if entry.report_item}
        if len(codes) > 1 or len(targets) > 1:
            raise RetrievalError("INSTITUTION_REFERENCE_CONFLICT", "机构参照匹配到多个信用代码或报表项目，请核对外部文件各参照表中的映射")
        if any(entry.disabled for entry in entries):
            raise RetrievalError("INSTITUTION_DISABLED", "当前机构在机构参照中标为禁用，请先核对禁用设置")
        code = next(iter(codes.values()), "")
        names = {normalized(entry.name): entry.name for entry in entries if entry.name}
        return code, next(iter(names.values()), ""), next(iter(targets.values()), "")

    def match(self, codes=(), names=()) -> tuple[list[InstitutionEntry], str]:
        """Code takes precedence; names are consulted only when no code maps.

        Consistent copies across reference sheets are allowed. Duplicate rows in
        one sheet, contradictory targets and disabled entries never fall back.
        """
        entries = list(dict.fromkeys(entry for code in codes if normalized(code)
                                     for entry in self.by_code(code)))
        if any(entry.disabled for entry in entries):
            self.validate(entries)
        if not any(entry.report_item for entry in entries):
            entries = []
        method = "统一社会信用代码"
        if not entries:
            entries = list(dict.fromkeys(entry for name in names if normalized(name)
                                         for entry in self.by_name(name)))
            if not any(entry.report_item for entry in entries):
                entries = []
            method = "机构名称"
        self.validate(entries)
        seen = set()
        for entry in entries:
            key = (entry.path, entry.sheet, normalized(entry.name) if method == "机构名称" else normalized(entry.code))
            if key in seen:
                raise RetrievalError("IDENTITY_MAPPING_AMBIGUOUS", "外部文件同一参照表中机构匹配多行，请保留唯一机构记录")
            seen.add(key)
        return entries, method if entries else ""


def _is_code(value, catalog: InstitutionCatalog) -> bool:
    if value is None or isinstance(value, bool):
        return False
    code = str(value).strip()
    if code.startswith("="):
        return False
    return bool(catalog.by_code(code) or re.fullmatch(r"[A-Za-z0-9]{18}", code)
                or (re.fullmatch(r"[A-Za-z0-9]{4,32}", code) and re.search("[A-Za-z]", code)))


@shared_readers
def resolve_institution(path: str | Path, catalog: InstitutionCatalog, *,
                        identity_sheet: str = "", identity_cell: str = "") -> InstitutionIdentity:
    """Prefer source identity; filename prefix is an exact reference lookup fallback.

    Match the external file by code first, then exact institution name. Ambiguous
    or disabled code matches cannot be hidden by retrying a different name.
    """
    path = Path(path).resolve()
    book = read_book(path, data_only=False)
    codes: list[tuple[int, str, str]] = []
    names: list[tuple[int, str]] = []
    try:
        if identity_sheet and identity_sheet not in book.sheetnames:
            raise RetrievalError("IDENTITY_SHEET_MISSING", f"{path.name}：缺少指定身份工作表 {identity_sheet}")
        sheets = [book[identity_sheet]] if identity_sheet else book.worksheets
        if identity_cell:
            try:
                coordinate_from_string(identity_cell.replace("$", ""))
            except Exception as exc:
                raise RetrievalError("IDENTITY_CELL_INVALID", "身份单元格应为 D4 这样的单个坐标") from exc
        for sheet in sheets:
            def cell(row, column):
                for merged in sheet.merged_cells.ranges:
                    if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
                        return sheet.cell(merged.min_row, merged.min_col)
                return sheet.cell(row, column)
            if identity_cell:
                selected = sheet[identity_cell.replace("$", "")]
                selected = cell(selected.row, selected.column)
                if selected.data_type != "f" and _is_code(selected.value, catalog):
                    codes.append((3, str(selected.value).strip(), f"{sheet.title}!{selected.coordinate}"))
            # Read the institution banner, never search numeric business data.
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12), max_col=min(sheet.max_column, 20)):
                for label in row:
                    text = normalized(label.value)
                    if text not in (*CODE_LABELS, *NAME_LABELS):
                        continue
                    right_col, below_row = label.column + 1, label.row + 1
                    for merged in sheet.merged_cells.ranges:
                        if merged.min_row == label.row and merged.min_col == label.column:
                            right_col, below_row = merged.max_col + 1, merged.max_row + 1
                            break
                    right = cell(label.row, right_col)
                    below = cell(below_row, label.column)
                    chosen = right if right.value not in (None, "") else below
                    if chosen.data_type in {"f", "e"}:
                        continue
                    priority = 2 if text.startswith("法人") else 1
                    if not identity_cell and text in CODE_LABELS and _is_code(chosen.value, catalog):
                        codes.append((priority, str(chosen.value).strip(), f"{sheet.title}!{chosen.coordinate}"))
                    if text in NAME_LABELS and isinstance(chosen.value, str) and chosen.value.strip():
                        names.append((priority, chosen.value.strip()))
    finally:
        book.close()
    if identity_cell and not codes:
        raise RetrievalError("IDENTITY_CELL_INVALID", f"{path.name}：指定单元格没有有效机构代码，请核对工作表与坐标")
    if codes:
        priority = max(value[0] for value in codes)
        codes = [value for value in codes if value[0] == priority]
    distinct = {normalized(code): code for _, code, _ in codes}
    if len(distinct) > 1:
        raise RetrievalError("SOURCE_IDENTITY_CONFLICT", f"{path.name}：不同工作表提供了不同机构代码，请核对源文件")
    source_code = next(iter(distinct.values()), "")
    name_priority = max((priority for priority, _ in names), default=0)
    source_names = list(dict.fromkeys(name for priority, name in names if priority == name_priority))
    filename_name = path.stem.split("_", 1)[0].strip() if "_" in path.stem else ""
    entries, mapping_method = catalog.match([source_code], [filename_name, *source_names])
    reference_code, canonical_name, report_item = catalog.validate(entries)
    code = reference_code or source_code
    if not code:
        raise RetrievalError("INSTITUTION_NOT_IDENTIFIED", f"{path.name}：未读到表内机构代码，也未通过名称匹配外部文件机构参照；请核对外部文件的机构名称和信用代码")
    if catalog.entries and not entries:
        raise RetrievalError("IDENTITY_NOT_MAPPED", f"{path.name}：信用代码及机构名称均未匹配到外部文件的报表项目，请补充或修正外部机构参照")
    return InstitutionIdentity(code, canonical_name or (source_names[0] if source_names else filename_name), report_item,
                               "机构名称匹配" if mapping_method == "机构名称" and source_code else "指定单元格" if identity_cell else "表内身份字段" if source_code else "文件名与机构参照",
                               tuple(dict.fromkeys(location for _, _, location in codes)), tuple(entries), source_code, mapping_method)
