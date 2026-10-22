"""Shared, read-only path discovery for audit and expression export."""
from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from typing import Any, Iterable

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from .errors import RetrievalError


@dataclass(frozen=True)
class NamedRegion:
    name: str
    sheet: str
    reference: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int


def region_number(name: str) -> int:
    match = re.search(r"_(\d+)(?:#\d+)?$", name)
    return int(match.group(1)) if match else 1_000_000_000


def named_regions(workbook, prefixes: tuple[str, ...]) -> list[NamedRegion]:
    definitions = list(workbook.defined_names.values())
    for sheet in workbook.worksheets:
        definitions.extend(sheet.defined_names.values())
    regions: dict[tuple, NamedRegion] = {}
    for definition in definitions:
        name = str(definition.name)
        if not name.startswith(prefixes):
            continue
        try:
            destinations = list(definition.destinations)
        except (ValueError, TypeError, AttributeError):
            continue
        for index, (sheet, coordinate) in enumerate(destinations, start=1):
            if sheet not in workbook.sheetnames and sheet.replace("''", "'") in workbook.sheetnames:
                sheet = sheet.replace("''", "'")
            if sheet not in workbook.sheetnames:
                continue
            try:
                bounds = range_boundaries(coordinate)
            except ValueError:
                continue
            if not all(bounds):
                continue
            min_col, min_row, max_col, max_row = bounds
            key = (name, sheet, min_row, min_col, max_row, max_col)
            regions[key] = NamedRegion(
                name if len(destinations) == 1 else f"{name}#{index}", sheet,
                f"{sheet}!{coordinate}", min_row, min_col, max_row, max_col,
            )
    return sorted(regions.values(), key=lambda r: (r.sheet, region_number(r.name), r.min_row, r.min_col))


def numeric_regions(workbook) -> tuple[list[NamedRegion], str]:
    primary = named_regions(workbook, ("数字区域", "数值区域"))
    compatible = named_regions(workbook, ("条件格式区域",))
    inferred, _ = inferred_numeric_regions(workbook)
    regions, origins = [], []
    for sheet in workbook.sheetnames:
        for candidates, origin in ((primary, "数字区域/数值区域"), (compatible, "条件格式区域（兼容现有模板）"),
                                   (inferred, "表结构区域交叉推导")):
            selected = [r for r in candidates if r.sheet == sheet]
            if selected:
                regions.extend(selected)
                if origin not in origins:
                    origins.append(origin)
                break
    return regions, "、".join(origins)


def split_path(value: Any) -> tuple[str, ...]:
    if isinstance(value, str):
        return tuple(part.strip() for part in value.split("@"))
    if isinstance(value, (tuple, list)):
        return tuple(str(part).strip() for part in value)
    return ()


def _sequence(book, sheet: str, coordinates: Iterable[tuple[int, int]]) -> list[tuple[str, int]]:
    result = []
    previous = None
    for offset, (row, col) in enumerate(coordinates):
        anchor = book._anchor(sheet, row, col)
        value = book.cell_raw(sheet, *anchor)
        label = "" if value is None else str(value).strip()
        if previous != anchor and label:
            result.append((label, offset))
        previous = anchor
    return result


def row_path_candidates(book, sheet: str, region: NamedRegion, *, candidate_rows=None):
    """Return horizontal row paths plus safe group-qualified aliases for repeats.

    Some templates repeat the same child labels below different group rows, but
    place both the group and child labels in one header column. When the repeated
    child rows form identical contiguous blocks, the unique row immediately above
    each block is a stable group discriminator. We expose that longer path only
    when it is unique in the region; otherwise the original path remains
    ambiguous and callers must not guess.
    """
    rows = sorted(set(candidate_rows if candidate_rows is not None else
                      range(region.min_row, region.max_row + 1)))
    entries = {}
    anchor_occurrences = defaultdict(set)
    path_counts = defaultdict(int)
    for row in rows:
        tokens = []
        previous_anchor = None
        for col in range(region.min_col, region.max_col + 1):
            anchor = book._anchor(sheet, row, col)
            raw = book.cell_raw(sheet, *anchor)
            label = "" if raw is None else str(raw).strip()
            if anchor != previous_anchor and label:
                tokens.append((label, col, anchor))
                anchor_occurrences[(col, label)].add(anchor)
            previous_anchor = anchor
        path = tuple(item[0] for item in tokens)
        entries[row] = {"tokens": tokens, "path": path}
        if path:
            path_counts[path] += 1

    duplicate_paths = {path for path, count in path_counts.items() if count > 1}
    duplicate_rows = {row for row, item in entries.items() if item["path"] in duplicate_paths}
    runs = []
    for row in sorted(duplicate_rows):
        if not runs or row != runs[-1][-1] + 1:
            runs.append([row])
        else:
            runs[-1].append(row)
    run_groups = defaultdict(list)
    for run in runs:
        signature = tuple(entries[row]["path"] for row in run)
        run_groups[signature].append(run)

    proposals = {}
    for signature, matching_runs in run_groups.items():
        if len(matching_runs) < 2 or not signature:
            continue
        for run in matching_runs:
            first_tokens = entries[run[0]]["tokens"]
            repeated_tokens = [item for item in first_tokens
                               if len(anchor_occurrences[(item[1], item[0])]) > 1]
            if not repeated_tokens:
                continue
            # The rightmost duplicated label is the narrowest useful group
            # discriminator when a row has several header columns.
            repeated = repeated_tokens[-1]
            label, col, _anchor = repeated
            parent = None
            for parent_row in range(run[0] - 1, region.min_row - 1, -1):
                if parent_row not in entries:
                    continue
                raw = book.cell_raw(sheet, parent_row, col)
                parent_label = "" if raw is None else str(raw).strip()
                if not parent_label:
                    continue
                parent_anchor = book._anchor(sheet, parent_row, col)
                if len(anchor_occurrences[(col, parent_label)]) != 1:
                    continue
                parent_path = entries[parent_row]["path"]
                if not parent_path or path_counts[parent_path] != 1:
                    continue
                parent = parent_label
                break
            if not parent:
                continue
            for row in run:
                tokens = entries[row]["tokens"]
                token_index = next((index for index, item in enumerate(tokens)
                                    if item[1] == col and len(anchor_occurrences[(col, item[0])]) > 1), None)
                if token_index is None:
                    continue
                path = entries[row]["path"]
                proposals[row] = path[:token_index] + (parent,) + path[token_index:]

    proposal_counts = defaultdict(int)
    for path in proposals.values():
        proposal_counts[path] += 1
    base_paths = set(entries[row]["path"] for row in entries)
    for row, item in entries.items():
        base = item["path"]
        proposed = proposals.get(row)
        unique = bool(proposed and proposal_counts[proposed] == 1 and proposed not in base_paths)
        item["aliases"] = (base, proposed) if unique else ((base,) if base else ())
        item["primary"] = proposed if unique else base
    return entries


def _contextual_row_hits(book, sheet, region, paths, candidate_rows):
    """Resolve context-qualified paths when source files have no named ranges.

    The source may be wider than the template's row-header rectangle. In that
    case, match the path with its inserted group label removed against the row,
    then verify the group label above the repeated child within the same column.
    """
    rows = sorted(set(candidate_rows))
    if not rows:
        return {}
    base_to_paths = defaultdict(list)
    for path in paths:
        # A context label is inserted before the repeated child it qualifies.
        # The last path item therefore cannot be the inserted parent; allowing
        # it would misread a regular terminal row label as context and can
        # attach a neighboring row to the wrong path.
        for parent_index in range(0, len(path) - 1):
            base = path[:parent_index] + path[parent_index + 1:]
            if base:
                base_to_paths[base].append((path, parent_index, path[parent_index]))
    if not base_to_paths:
        return {}

    sequences = {}
    anchors_by_column_label = defaultdict(set)
    for row in rows:
        sequence = _sequence(book, sheet, ((row, col) for col in range(region.min_col, region.max_col + 1)))
        sequences[row] = sequence
        for col in range(region.min_col, region.max_col + 1):
            anchor = book._anchor(sheet, row, col)
            raw = book.cell_raw(sheet, *anchor)
            label = "" if raw is None else str(raw).strip()
            if label:
                anchors_by_column_label[(col, label)].add(anchor)

    hits_by_base = defaultdict(list)
    for row, sequence in sequences.items():
        for base in base_to_paths:
            for start in _starts(sequence, base):
                hits_by_base[base].append((row, start))

    result = defaultdict(set)
    eligible = set(rows)
    for base, candidates in hits_by_base.items():
        # Context is only used to disambiguate an actually repeated horizontal path.
        if len({row for row, _ in candidates}) < 2:
            continue
        for path, parent_index, expected_parent in base_to_paths[base]:
            for row, start in candidates:
                child_index = start + parent_index
                sequence = sequences[row]
                if child_index >= len(sequence):
                    continue
                child_col = region.min_col + sequence[child_index][1]
                nearest_unique = ""
                for parent_row in range(row - 1, region.min_row - 1, -1):
                    if parent_row not in eligible:
                        continue
                    anchor = book._anchor(sheet, parent_row, child_col)
                    raw = book.cell_raw(sheet, *anchor)
                    label = "" if raw is None else str(raw).strip()
                    if label and len(anchors_by_column_label[(child_col, label)]) == 1:
                        nearest_unique = label
                        break
                if nearest_unique != expected_parent:
                    continue
                first_col = region.min_col + sequence[start][1]
                last_col = region.min_col + sequence[start + len(base) - 1][1]
                result[path].add((row, first_col, last_col))
    return {path: sorted(matches) for path, matches in result.items()}


def _starts(sequence, path):
    for i in range(len(sequence) - len(path) + 1):
        if tuple(value for value, _ in sequence[i:i + len(path)]) == path:
            yield i


def paired_column_region(region: NamedRegion, regions: list[NamedRegion]) -> NamedRegion | None:
    candidates = [item for item in regions if item.sheet == region.sheet
                  and item.max_col > region.max_col and item.max_row < region.max_row
                  and item.min_row <= region.min_row
                  and not item.name.startswith("行索引区域")]
    numbered = [item for item in candidates if region_number(item.name) == region_number(region.name) + 1]
    candidates = numbered or candidates
    return min(candidates, key=lambda item: (item.min_col, abs(item.min_row - region.min_row))) if candidates else None


def projected_column_region(book, region: NamedRegion, regions: list[NamedRegion]) -> NamedRegion | None:
    """Resolve a sequential column axis misplaced onto a sibling sheet by geometry."""
    paired = paired_column_region(region, regions)
    if paired:
        return paired
    workbook = book.book if hasattr(book, "book") else book
    candidates = [item for item in named_regions(workbook, ("表结构区域", "行索引区域", "列索引区域"))
                  if region_number(item.name) == region_number(region.name) + 1
                  and item.min_col > region.max_col]
    if not candidates:
        return None
    candidate = min(candidates, key=lambda item: (item.min_col, item.min_row))
    target_sheet = workbook[region.sheet]
    if not any(target_sheet.cell(row, col).value is not None
               for row in range(candidate.min_row, candidate.max_row + 1)
               for col in range(candidate.min_col, candidate.max_col + 1)):
        return None
    return NamedRegion(
        candidate.name, region.sheet,
        f"{region.sheet}!{get_column_letter(candidate.min_col)}{candidate.min_row}:"
        f"{get_column_letter(candidate.max_col)}{candidate.max_row}",
        candidate.min_row, candidate.min_col, candidate.max_row, candidate.max_col,
    )


def inferred_numeric_regions(workbook) -> tuple[list[NamedRegion], list[NamedRegion]]:
    """Use the geometric cross of existing row and column structure regions."""
    axes = named_regions(workbook, ("表结构区域", "行索引区域", "列索引区域"))
    results = {}
    used_bounds = set()
    bounds = lambda r: (r.sheet, r.min_row, r.min_col, r.max_row, r.max_col)
    for row_axis in axes:
        if row_axis.name.startswith("列索引区域"):
            continue
        column_axis = paired_column_region(row_axis, axes)
        if not column_axis:
            # Some templates have a sequential column-header name that points to
            # a sibling sheet by mistake. Reuse only its geometry when the target
            # sheet actually has header content at those coordinates.
            numbered = [item for item in axes
                        if region_number(item.name) == region_number(row_axis.name) + 1
                        and item.min_col > row_axis.max_col]
            if numbered:
                candidate = min(numbered, key=lambda item: (item.min_col, item.min_row))
                target_sheet = workbook[row_axis.sheet]
                has_target_headers = any(
                    target_sheet.cell(row, col).value is not None
                    for row in range(candidate.min_row, candidate.max_row + 1)
                    for col in range(candidate.min_col, candidate.max_col + 1)
                )
                if has_target_headers:
                    column_axis = NamedRegion(
                        candidate.name, row_axis.sheet,
                        f"{row_axis.sheet}!{get_column_letter(candidate.min_col)}{candidate.min_row}:"
                        f"{get_column_letter(candidate.max_col)}{candidate.max_row}",
                        candidate.min_row, candidate.min_col, candidate.max_row, candidate.max_col,
                    )
                    used_bounds.add(bounds(candidate))
        if not column_axis:
            continue
        min_row = max(row_axis.min_row, column_axis.max_row + 1)
        min_col = max(row_axis.max_col + 1, column_axis.min_col)
        if min_row > row_axis.max_row or min_col > column_axis.max_col:
            continue
        key = (row_axis.sheet, min_row, min_col, row_axis.max_row, column_axis.max_col)
        if key not in results:
            results[key] = NamedRegion(
                f"数字区域_推导_{len(results) + 1:03d}", row_axis.sheet,
                f"{row_axis.sheet}!{get_column_letter(min_col)}{min_row}:{get_column_letter(column_axis.max_col)}{row_axis.max_row}",
                min_row, min_col, row_axis.max_row, column_axis.max_col,
            )
        used_bounds.update((bounds(row_axis), bounds(column_axis)))
    missing = {bounds(r): r for r in axes if bounds(r) not in used_bounds}
    return list(results.values()), list(missing.values())


def scan_rule_paths(book, sheet_name: str, key_pairs: set[tuple[tuple[str, ...], tuple[str, ...]]]):
    """Return the same coordinate index to the evaluator and the exporters."""
    sheet = book.book[sheet_name]
    data_regions = book.data_regions(sheet_name)
    regions = book.structure_regions(sheet_name) or [NamedRegion(
        "全工作表回退", sheet_name, f"{sheet_name}!A1:{get_column_letter(max(sheet.max_column, 1))}{max(sheet.max_row, 1)}",
        1, 1, max(sheet.max_row, 1), max(sheet.max_column, 1),
    )]
    row_paths = {row for row, _ in key_pairs}
    locations = defaultdict(set)
    row_locations = defaultdict(set)
    column_locations = defaultdict(set)
    tables = []
    for region in regions:
        if region.name.startswith("列索引区域"):
            continue
        row_hits = defaultdict(list)
        eligible_rows = {row for data in data_regions if data.sheet == sheet_name
                         for row in range(max(region.min_row, data.min_row), min(region.max_row, data.max_row) + 1)}
        if not data_regions:
            eligible_rows = set(range(region.min_row, region.max_row + 1))
        path_index = row_path_candidates(book, sheet_name, region, candidate_rows=eligible_rows)
        contextual_hits = _contextual_row_hits(book, sheet_name, region, row_paths, eligible_rows)
        for row, item in path_index.items():
            sequence = _sequence(book, sheet_name, ((row, col) for col in range(region.min_col, region.max_col + 1)))
            if not sequence:
                continue
            for path in row_paths:
                if path in item["aliases"]:
                    first_col = region.min_col + sequence[0][1]
                    last_col = region.min_col + sequence[-1][1]
                    row_hits[path].append((row, first_col, last_col))
                    row_locations[path].add(row)
                    continue
                for start in _starts(sequence, path):
                    first_col = region.min_col + sequence[start][1]
                    last_col = region.min_col + sequence[start + len(path) - 1][1]
                    row_hits[path].append((row, first_col, last_col))
                    row_locations[path].add(row)
        for path, matches in contextual_hits.items():
            for row, first_col, last_col in matches:
                row_hits[path].append((row, first_col, last_col))
                row_locations[path].add(row)
        if not row_hits:
            continue
        column_region = projected_column_region(book, region, regions)
        search = column_region or region
        column_hits = defaultdict(set)
        cache = {}
        for row_path, column_path in key_pairs:
            for row, _, last_col in row_hits.get(row_path, ()):
                first_col = max(search.min_col, region.max_col + 1) if column_region else max(search.min_col, last_col + 1)
                last_row = min(search.max_row, row - 1)
                for col in range(first_col, search.max_col + 1):
                    if data_regions and not in_numeric_regions(book, sheet_name, row, col, data_regions):
                        continue
                    cache_key = (col, last_row)
                    if cache_key not in cache:
                        cache[cache_key] = _sequence(book, sheet_name, ((r, col) for r in range(search.min_row, last_row + 1)))
                    if any(True for _ in _starts(cache[cache_key], column_path)):
                        column_hits[column_path].add(col)
                        column_locations[column_path].add(col)
                        locations[(row_path, column_path)].add((row, col))
        tables.append({"region": region, "column_region": column_region, "search_region": search,
                       "row_hits": dict(row_hits), "column_hits": dict(column_hits)})
    return ({key: sorted(value) for key, value in locations.items()},
            {key: sorted(value) for key, value in row_locations.items()},
            {key: sorted(value) for key, value in column_locations.items()}, tables)


def in_numeric_regions(book, sheet, row, col, regions):
    anchor_row, anchor_col = book._anchor(sheet, row, col)
    return any(region.min_row <= row <= region.max_row and region.min_col <= col <= region.max_col
               and region.min_row <= anchor_row <= region.max_row
               and region.min_col <= anchor_col <= region.max_col for region in regions)


def numeric_header_regions(book, data: NamedRegion) -> tuple[NamedRegion, NamedRegion]:
    """Find the row/column axes geometrically, allowing optional explicit axis names."""
    regions = named_regions(book.book, ("表结构区域", "行索引区域", "列索引区域"))
    rows = [r for r in regions if r.sheet == data.sheet and r.max_col < data.min_col
            and r.min_row <= data.min_row and r.max_row >= data.max_row
            and not r.name.startswith("列索引区域")]
    cols = [r for r in regions if r.sheet == data.sheet and r.max_row < data.min_row
            and r.min_col <= data.min_col and r.max_col >= data.max_col
            and not r.name.startswith("行索引区域")]
    def choose(candidates, explicit, axis):
        selected = [r for r in candidates if r.name.startswith(explicit)] or candidates
        bounds = {(r.min_row, r.min_col, r.max_row, r.max_col): r for r in selected}
        if not bounds:
            raise RetrievalError("HEADER_REGION_MISSING", f"数字区域缺少可确定的{axis}表头区域", {"region": data.name})
        # The closest full-width/full-height named axis owns this data rectangle.
        distance = (lambda r: data.min_col - r.max_col) if axis == "行" else (lambda r: data.min_row - r.max_row)
        nearest = min(distance(r) for r in bounds.values())
        selected = [r for r in bounds.values() if distance(r) == nearest]
        if len(selected) != 1:
            raise RetrievalError("HEADER_REGION_AMBIGUOUS", f"数字区域匹配多组{axis}表头区域", {"region": data.name})
        return selected[0]
    row_axis = choose(rows, "行索引区域", "行")
    if not cols:
        # Match the same cross-sheet geometry fallback used when inferring the
        # numeric rectangle. The target sheet's cells provide the actual labels.
        projected = [r for r in regions
                     if region_number(r.name) == region_number(row_axis.name) + 1
                     and r.min_col <= data.min_col and r.max_col >= data.max_col
                     and r.max_row < data.min_row]
        if projected:
            candidate = min(projected, key=lambda r: (r.min_row, r.min_col))
            target = book.book[data.sheet]
            if any(target.cell(row, col).value is not None
                   for row in range(candidate.min_row, candidate.max_row + 1)
                   for col in range(candidate.min_col, candidate.max_col + 1)):
                cols = [NamedRegion(
                    candidate.name, data.sheet,
                    f"{data.sheet}!{get_column_letter(candidate.min_col)}{candidate.min_row}:"
                    f"{get_column_letter(candidate.max_col)}{candidate.max_row}",
                    candidate.min_row, candidate.min_col, candidate.max_row, candidate.max_col,
                )]
    return row_axis, choose(cols, "列索引区域", "列")
