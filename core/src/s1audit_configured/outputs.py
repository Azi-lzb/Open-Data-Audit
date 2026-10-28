from __future__ import annotations

import json
import os
import re
import shutil
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .model import AuditRunResult, Configuration
from .sources import SourceRepository, file_sha256
from .fast_ooxml import DataTable, supports_tables, write_tables


_RESULT_HEADERS = (
    "机构名称", "报表项目", "统一社会信用代码",
    "规则名称", "工作表名", "定位单元格", "错误类型", "校验指标", "描述", "本期值", "对比值", "差值",
)
_FAILURE_TYPES = {
    "RETRIEVAL_FAILED": "取数失败", "STRUCTURE_ERROR": "表结构错误",
    "CALCULATION_FAILED": "计算失败", "CONFIG_ERROR": "配置错误",
}


def write_run_outputs(result: AuditRunResult, configuration: Configuration,
                      repository: SourceRepository,
                      output_root: Path) -> dict[str, str]:
    output_root = output_root.resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    run_dir = output_root / _safe_name(result.run_id)
    if run_dir.exists():
        raise FileExistsError(f"运行输出目录已存在，不覆盖已有结果：{run_dir}")
    run_dir.mkdir()
    copies_dir = run_dir / "审核副本"
    copies_dir.mkdir()

    source_copy_by_path: dict[str, Path] = {}
    copy_manifest: list[dict[str, Any]] = []
    for role, binding in repository.selected_bindings():
        source_path = Path(binding.path).resolve()
        if not source_path.is_file():
            continue
        path_key = str(source_path)
        if path_key not in source_copy_by_path:
            target = copies_dir / f"{_safe_name(role)}_{_safe_name(binding.source_id)}_{_safe_name(source_path.stem)}{source_path.suffix.lower()}"
            shutil.copy2(source_path, target)
            source_copy_by_path[path_key] = target
            source_after = file_sha256(source_path)
            source_before = repository.start_hashes.get(f"{role}:{binding.source_id}", "")
            copy_hash = file_sha256(target)
            copy_manifest.append({
                "role": role, "source_id": binding.source_id, "source_path": str(source_path),
                "copy_path": str(target), "sha256_before": source_before,
                "sha256_after": source_after, "sha256_copy": copy_hash,
                "matches_snapshot": source_before == source_after == copy_hash,
            })
    config_copy: Path | None = None
    if configuration.path and configuration.path.is_file():
        config_copy = copies_dir / f"配置_{_safe_name(configuration.path.stem)}{configuration.path.suffix.lower()}"
        shutil.copy2(configuration.path, config_copy)
        config_after = file_sha256(configuration.path)
        config_copy_hash = file_sha256(config_copy)
        copy_manifest.append({
            "role": "configuration", "source_id": "configuration", "source_path": str(configuration.path),
            "copy_path": str(config_copy), "sha256_before": configuration.sha256,
            "sha256_after": config_after, "sha256_copy": config_copy_hash,
            "matches_snapshot": configuration.sha256 == config_after == config_copy_hash,
        })

    _refresh_integrity(result, configuration, repository)
    structure_events = result.ledger.get("structure_checks", [])
    if result.status == "COMPLETED":
        result.ledger["structure_checks"] = repository.persist_structure_snapshots()
    else:
        for event in structure_events:
            if event.get("status") in {"BASELINE_CREATED", "UNCHANGED", "POSITION_CHANGED"}:
                event["status"] = "SNAPSHOT_NOT_SAVED"
                event["detail"] = "输入或配置在生成审核副本时发生变化，未更新结构快照"
        result.ledger["structure_checks"] = structure_events

    result_path = run_dir / "审核结果.xlsx"
    ledger_path = run_dir / "运行总账.json"
    result.output_files.update({
        "result_workbook": str(result_path),
        "run_ledger": str(ledger_path),
        "audit_copies": str(copies_dir),
    })
    if config_copy:
        result.output_files["configuration_copy"] = str(config_copy)
    result.ledger["audit_copy_manifest"] = copy_manifest
    result.ledger["output_files"] = result.output_files
    _write_result_workbook(result, result_path, source_copy_by_path)
    ledger_path.write_text(json.dumps(result.ledger, ensure_ascii=False, indent=2, default=_json_default), encoding="utf-8")
    return result.output_files


def _refresh_integrity(result: AuditRunResult, configuration: Configuration,
                       repository: SourceRepository) -> None:
    inputs_unchanged = repository.hashes_unchanged()
    config_after = ""
    config_unchanged = True
    if configuration.path:
        config_after = file_sha256(configuration.path) if configuration.path.is_file() else "MISSING"
        config_unchanged = config_after == configuration.sha256
    unchanged = inputs_unchanged and config_unchanged
    if not unchanged:
        result.status = "INPUT_CHANGED_DURING_RUN"
        for item in result.results:
            if item.status in {"PASSED", "TRIGGERED"}:
                item.status = "CALCULATION_FAILED"
                item.triggered = None
                item.error_code = "INPUT_CHANGED_DURING_RUN"
                item.error_message = "审核结果文件生成期间输入或配置哈希发生变化，结论不可用"
    result.ledger["run_status"] = result.status
    result.ledger["input_integrity"] = "UNCHANGED" if inputs_unchanged else "CHANGED_DURING_RUN"
    result.ledger["configuration_integrity"] = "UNCHANGED" if config_unchanged else "CHANGED_DURING_RUN"
    hashes = result.ledger.get("input_hashes", {})
    for key, value in repository.start_hashes.items():
        hashes[key] = {
            "sha256_before": value,
            "sha256_after": repository.end_hashes.get(key, ""),
        }
    if configuration.path:
        hashes["configuration"] = {
            "sha256_before": configuration.sha256,
            "sha256_after": config_after,
        }
    result.ledger["input_hashes"] = hashes
    counts: dict[str, int] = {}
    for item in result.results:
        counts[item.status] = counts.get(item.status, 0) + 1
    result.ledger["status_counts"] = counts


def _error_type(item) -> str:
    if item.status in _FAILURE_TYPES:
        return _FAILURE_TYPES[item.status]
    level = str(item.severity).strip().casefold()
    if level in {"条件格式", "条件格式触发", "条件格式错误"}:
        return "条件格式触发"
    if level in {"错误", "硬性", "硬性错误", "校验硬性错误", "error", "hard"}:
        return "校验硬性错误"
    if level in {"核实", "警告", "软性", "软性错误", "校验软性错误", "warn", "warning", "soft"}:
        return "校验软性错误"
    # Preserve an explicitly configured custom level instead of guessing.
    return str(item.severity)


def _location(item, current_paths):
    if item.status in _FAILURE_TYPES and item.error_evidence:
        return (str(item.error_evidence.get("sheet") or item.sheet_name),
                str(item.error_evidence.get("cell") or ""), None)
    if item.formula_cell:
        return item.sheet_name, item.formula_cell, None
    evidence = [entry for entry in item.evidence
                if not current_paths or str(Path(entry.source_path).resolve()) in current_paths]
    preferred = [entry for entry in evidence if entry.sheet == item.sheet_name]
    evidence = preferred or evidence
    if evidence:
        sheets = list(dict.fromkeys(entry.sheet for entry in evidence))
        cells = list(dict.fromkeys(entry.cell if len(sheets) == 1 else f"{entry.sheet}!{entry.cell}"
                                   for entry in evidence))
        return "；".join(sheets), "；".join(cells), evidence[0]
    # Ambiguous candidate addresses are not a resolved cell location.
    return (str(item.error_evidence.get("sheet") or item.sheet_name),
            str(item.error_evidence.get("cell") or ""), None)


def _description(item) -> str:
    if item.status == "TRIGGERED":
        return item.description
    message = item.error_message or _FAILURE_TYPES.get(item.status, item.status)
    candidates = item.error_evidence.get("cells")
    if candidates:
        message += "；候选单元格：" + "、".join(str(cell) for cell in candidates)
    return message


def _style_table(sheet, widths):
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 54
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            # User-maintained text remains literal even when it starts with '='.
            if isinstance(cell.value, str):
                cell.data_type = "s"


def _write_result_workbook(result: AuditRunResult, path: Path,
                           source_copy_by_path: dict[str, Path], *,
                           batch_records: list[dict[str, Any]] | None = None,
                           previous_records: list[dict[str, Any]] | None = None) -> None:
    """Use direct OOXML for literal result tables; keep rich export as fallback."""
    options = dict(batch_records=batch_records, previous_records=previous_records)
    engine = os.environ.get("AUDIT3_WRITER_ENGINE", "direct")
    if engine not in {"direct", "openpyxl"}:
        raise ValueError("AUDIT3_WRITER_ENGINE must be direct or openpyxl")
    if engine == "direct":
        tables = _result_tables(result, path, source_copy_by_path, **options)
        if supports_tables(tables):
            write_tables(path, tables)
            return
    _write_result_workbook_openpyxl(result, path, source_copy_by_path, **options)


def _result_tables(result, path, source_copy_by_path, *, batch_records=None, previous_records=None):
    table = DataTable("审核结果", _RESULT_HEADERS,
                      (34, 34, 26, 28, 30, 20, 20, 30, 70, 18, 18, 18),
                      numeric_columns=(10, 11, 12), text_columns=(3,))
    current_paths = {str(Path(entry["source_path"]).resolve())
                     for entry in result.ledger.get("audit_copy_manifest", []) if entry.get("role") == "current"}
    source_by_run = {record.get("run_id"): str(Path(record["source"]).resolve()) for record in batch_records or []}
    identity_by_run = {record.get("run_id"): record.get("identity", {}) for record in batch_records or []}
    for item in result.results:
        if item.status != "TRIGGERED" and item.status not in _FAILURE_TYPES:
            continue
        paths = {source_by_run[item.run_id]} if item.run_id in source_by_run else current_paths
        sheet, cells, evidence = _location(item, paths)
        identity = identity_by_run.get(item.run_id, {})
        report_item = identity.get("report_item") or next((entry.report_item for entry in item.evidence if entry.report_item), "")
        table.rows.append([item.institution_name, report_item, item.institution_code,
                           item.rule_name, sheet, cells, _error_type(item), item.metric,
                           _description(item), item.current_value, item.comparison_value, item.difference_value])
        row_no = len(table.rows) + 1
        target = None
        if evidence:
            copy_path = source_copy_by_path.get(str(Path(evidence.source_path).resolve()))
            if copy_path:
                safe_sheet = evidence.sheet.replace("'", "''")
                target = _source_cell_link(copy_path, path, safe_sheet, evidence.cell)
        elif item.error_evidence.get("source_path") and item.error_evidence.get("cell"):
            copy_path = source_copy_by_path.get(str(Path(item.error_evidence["source_path"]).resolve()))
            if copy_path:
                safe_sheet = str(item.error_evidence.get("sheet", "")).replace("'", "''")
                target = _source_cell_link(copy_path, path, safe_sheet, item.error_evidence["cell"])
        if target:
            table.hyperlinks[f"F{row_no}"] = target
            if table.rows[-1][5] is None:
                table.rows[-1][5] = target
        for col, precision in zip((10, 11, 12), (item.current_precision, item.comparison_precision, item.difference_precision)):
            if precision is not None:
                table.formats[(row_no, col)] = "0" if precision == 0 else "0." + "0" * precision
    if batch_records is not None:
        failed_run_ids = {item.run_id for item in result.results if item.status in _FAILURE_TYPES}
        for records, label in ((batch_records or [], "机构处理失败"), (previous_records or [], "上期文件识别失败")):
            for record in records:
                if record.get("status") in {"COMPLETED", "COMPLETED_WITH_ISSUES", "IDENTIFIED"}:
                    continue
                if record.get("run_id") in failed_run_ids:
                    continue
                identity = record.get("identity", {})
                name = identity.get("name") or Path(record.get("source", "")).stem.split("_", 1)[0]
                table.rows.append([name, identity.get("report_item", ""), identity.get("code", ""),
                                   "", "", "", label, "", record.get("message") or "文件未完成审核", None, None, None])
    cf_headers = ("机构名称", "报表项目", "统一社会信用代码", "工作表名", "定位单元格",
                  "单元格值", "判定状态", "条件格式类型", "规则公式", "应用范围", "优先级", "来源文件", "判定说明")
    widths = (34, 26, 22, 32, 16, 18, 18, 20, 70, 32, 12, 54, 60)
    conditional = DataTable("条件格式触发", cf_headers, widths, text_columns=(3, 5))
    unresolved = DataTable("条件格式无法判定", cf_headers, widths, text_columns=(3, 5))
    for item in result.conditional_format_results:
        target = unresolved if item.get("判定状态") == "无法判定" else conditional
        target.rows.append([item.get(header, "") for header in cf_headers])
    tables = [table, conditional] + ([unresolved] if unresolved.rows else [])
    data_warnings = result.ledger.get("data_quality_warnings", [])
    if data_warnings:
        warning_headers = ("机构名称", "统一社会信用代码", "来源文件", "工作表", "单元格/行", "提示")
        warning_table = DataTable("数据质量提示", warning_headers,
                                  (34, 26, 45, 24, 18, 90), text_columns=(2,))
        for warning in data_warnings:
            warning_table.rows.append([
                warning.get("institution_name", ""), warning.get("institution_code", ""),
                warning.get("source_file", ""), warning.get("sheet", ""),
                warning.get("location", ""), warning.get("message", ""),
            ])
        tables.append(warning_table)
    return tables


def _source_cell_link(source_path: Path, result_path: Path, sheet: str, cell: str) -> str:
    try:
        target = Path(os.path.relpath(source_path, result_path.parent)).as_posix()
    except ValueError:  # Windows paths on different drives cannot be made relative.
        target = Path(source_path).resolve().as_uri()
    return f"{target}#'{sheet}'!{cell}"


def _write_result_workbook_openpyxl(result: AuditRunResult, path: Path,
                           source_copy_by_path: dict[str, Path], *,
                           batch_records: list[dict[str, Any]] | None = None,
                           previous_records: list[dict[str, Any]] | None = None) -> None:
    """Export issues for readers; keep detailed execution evidence in JSON only."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "审核结果"
    batch = batch_records is not None
    sheet.append(list(_RESULT_HEADERS))
    current_paths = {str(Path(entry["source_path"]).resolve())
                     for entry in result.ledger.get("audit_copy_manifest", []) if entry.get("role") == "current"}
    source_by_run = {record.get("run_id"): str(Path(record["source"]).resolve()) for record in batch_records or []}
    identity_by_run = {record.get("run_id"): record.get("identity", {}) for record in batch_records or []}
    offset = 3
    for item in result.results:
        if item.status != "TRIGGERED" and item.status not in _FAILURE_TYPES:
            continue
        paths = {source_by_run[item.run_id]} if item.run_id in source_by_run else current_paths
        sheet_name, cells, first_evidence = _location(item, paths)
        row = [item.rule_name, sheet_name, cells, _error_type(item), item.metric,
               _description(item), item.current_value, item.comparison_value, item.difference_value]
        identity = identity_by_run.get(item.run_id, {})
        report_item = identity.get("report_item") or next((entry.report_item for entry in item.evidence if entry.report_item), "")
        sheet.append([item.institution_name, report_item, item.institution_code] + row)
        row_num = sheet.max_row
        if first_evidence:
            copy_path = source_copy_by_path.get(str(Path(first_evidence.source_path).resolve()))
            if copy_path:
                safe_sheet = first_evidence.sheet.replace("'", "''")
                sheet.cell(row_num, 3 + offset).hyperlink = _source_cell_link(
                    copy_path, path, safe_sheet, first_evidence.cell,
                )
        elif item.error_evidence.get("source_path") and item.error_evidence.get("cell"):
            copy_path = source_copy_by_path.get(str(Path(item.error_evidence["source_path"]).resolve()))
            if copy_path:
                safe_sheet = str(item.error_evidence.get("sheet", "")).replace("'", "''")
                sheet.cell(row_num, 3 + offset).hyperlink = _source_cell_link(
                    copy_path, path, safe_sheet, item.error_evidence["cell"],
                )
        for column, precision in zip((7, 8, 9),
                                     (item.current_precision, item.comparison_precision, item.difference_precision)):
            cell = sheet.cell(row_num, column + offset)
            if precision is not None:
                cell.number_format = "0" if precision == 0 else "0." + "0" * precision
    if batch:
        # A file which could not execute must remain visible in the issue list.
        for records, label in ((batch_records or [], "机构处理失败"),
                               (previous_records or [], "上期文件识别失败")):
            for record in records:
                if record.get("status") in {"COMPLETED", "COMPLETED_WITH_ISSUES", "IDENTIFIED"}:
                    continue
                # An invalidated run already has explicit calculation-failure rows.
                if any(item.run_id == record.get("run_id") and item.status in _FAILURE_TYPES for item in result.results):
                    continue
                identity = record.get("identity", {})
                name = identity.get("name") or Path(record.get("source", "")).stem.split("_", 1)[0]
                sheet.append([name, identity.get("report_item", ""), identity.get("code", ""),
                              "", "", "", label, "", record.get("message") or "文件未完成审核", None, None, None])
    widths = [34, 34, 26, 28, 30, 20, 20, 30, 70, 18, 18, 18]
    _style_table(sheet, widths)
    for row in sheet.iter_rows(min_row=2):
        row[2].number_format = "@"
        for column in (7, 8, 9):
            row[column - 1 + offset].alignment = Alignment(horizontal="right", vertical="top")
    conditional = workbook.create_sheet("条件格式触发")
    cf_headers = ("机构名称", "报表项目", "统一社会信用代码", "工作表名", "定位单元格",
                  "单元格值", "判定状态", "条件格式类型", "规则公式", "应用范围", "优先级", "来源文件", "判定说明")
    conditional.append(cf_headers)
    unresolved = None
    for item in result.conditional_format_results:
        target = conditional
        if item.get("判定状态") == "无法判定":
            if unresolved is None:
                unresolved = workbook.create_sheet("条件格式无法判定")
                unresolved.append(cf_headers)
            target = unresolved
        target.append([item.get(header, "") for header in cf_headers])
    for target in [conditional] + ([unresolved] if unresolved is not None else []):
        _style_table(target, (34, 26, 22, 32, 16, 18, 18, 20, 70, 32, 12, 54, 60))
        for row in target.iter_rows(min_row=2):
            row[2].number_format = "@"
            row[4].number_format = "@"
    data_warnings = result.ledger.get("data_quality_warnings", [])
    if data_warnings:
        quality = workbook.create_sheet("数据质量提示")
        quality.append(("机构名称", "统一社会信用代码", "来源文件", "工作表", "单元格/行", "提示"))
        for warning in data_warnings:
            quality.append([
                warning.get("institution_name", ""), warning.get("institution_code", ""),
                warning.get("source_file", ""), warning.get("sheet", ""),
                warning.get("location", ""), warning.get("message", ""),
            ])
        _style_table(quality, (34, 26, 45, 24, 18, 90))
        for row in quality.iter_rows(min_row=2):
            row[1].number_format = "@"
    workbook.save(path)
    workbook.close()


def _safe_name(value: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._\-\u3400-\u9fff]+", "_", str(value)).strip("_.")
    return safe[:80] or "source"


def _json_default(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, Path):
        return str(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if hasattr(value, "__dict__"):
        return value.__dict__
    return str(value)
