"""Evaluate source-workbook conditional formatting without Excel/LibreOffice."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange

from .layout import named_regions
from .source_reader import read_book, shared_readers
from .conditional_format_eval import (
    STATUS_ERROR, STATUS_TRUE, STATUS_UNSUPPORTED, evaluate_cellis_rule,
    evaluate_expression_condition, translate_formula_refs,
)


def _bounds(sqref):
    items = sqref.ranges if isinstance(sqref, MultiCellRange) else [sqref]
    for item in items:
        try:
            yield item.bounds
        except (AttributeError, ValueError, TypeError):
            try:
                yield CellRange(str(item)).bounds
            except (ValueError, TypeError):
                continue


def _operand(ws, formulas, index, *, anchor_row, anchor_col, row, col):
    if not formulas or index >= len(formulas):
        return None
    raw = str(formulas[index]).strip().lstrip("=")
    if raw.casefold() in {"true", "false"}:
        return raw.casefold() == "true"
    try:
        return float(raw) if any(char in raw.casefold() for char in (".", "e")) else int(raw)
    except ValueError:
        pass
    match = re.fullmatch(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)", raw)
    if match:
        col_locked, letters, row_locked, digits = match.groups()
        from openpyxl.utils.cell import column_index_from_string
        target_col = column_index_from_string(letters)
        target_row = int(digits)
        if not col_locked:
            target_col += col - anchor_col
        if not row_locked:
            target_row += row - anchor_row
        if target_row < 1 or target_col < 1:
            raise ValueError(f"条件格式引用越界：{raw}")
        return ws.cell(target_row, target_col).value
    raise ValueError(f"不支持的 cellIs 操作数：{raw}")


@shared_readers
def extract_conditional_triggers(path: str | Path, *, template_path: str | Path | None = None,
                                institution_name="", report_item="", institution_code="") -> list[dict]:
    """Return triggered conditional rules and explicit unsupported diagnostics.

    Rule formulas and values come from the source workbook because reporting
    files may update conditional formatting each period. The selected template
    supplies ``条件格式区域`` scan bounds when present; otherwise scan the source
    rules' AppliesTo ranges, matching S1F1. Excel/LibreOffice is not launched.
    Unsupported conditional-format rule types are reported as unresolved.
    """
    source_book = read_book(path, data_only=True)
    source_rule_book = read_book(path, data_only=False)
    layout_source = Path(template_path) if template_path else Path(path)
    template_book = (read_book(layout_source, data_only=False)
                     if layout_source.resolve() != Path(path).resolve() else source_rule_book)
    named_cf_regions = named_regions(template_book, ("条件格式区域",))
    regions_by_sheet = {}
    for region in named_cf_regions:
        regions_by_sheet.setdefault(region.sheet, []).append(
            (region.min_col, region.min_row, region.max_col, region.max_row))
    results = []
    unresolved = []
    try:
        for ws in source_book.worksheets:
            if ws.title not in source_rule_book.sheetnames:
                continue
            rule_ws = source_rule_book[ws.title]
            rules = []
            for order, cf in enumerate(rule_ws.conditional_formatting):
                for rule in cf.rules:
                    regions = list(_bounds(cf.sqref))
                    if not regions:
                        continue
                    # S1F1 treats the complete AppliesTo list as one relative
                    # reference frame. Using a fresh anchor for each disjoint
                    # area changes the meaning of rules that span multiple
                    # ranges (and can create hundreds of false positives).
                    anchor_col = min(bounds[0] for bounds in regions)
                    anchor_row = min(bounds[1] for bounds in regions)
                    rules.append((rule, cf.sqref, regions, order, anchor_row, anchor_col))
            if named_cf_regions:
                scan_regions = regions_by_sheet.get(ws.title, [])
                if not scan_regions:
                    continue
            else:
                # S1F1 fallback: when the selected template has no
                # 条件格式区域, scan the source workbook's AppliesTo ranges.
                scan_regions = list(dict.fromkeys(
                    bounds for _rule, _sqref, regions, _order, _anchor_row, _anchor_col in rules
                    for bounds in regions
                ))
            for rule, sqref, rule_regions, _order, _anchor_row, _anchor_col in rules:
                if str(rule.type or "").casefold() in {"cellis", "expression"}:
                    continue
                intersects = any(
                    max(scan[0], target[0]) <= min(scan[2], target[2])
                    and max(scan[1], target[1]) <= min(scan[3], target[3])
                    for scan in scan_regions for target in rule_regions
                )
                if intersects:
                    unresolved.append({
                        "机构名称": institution_name, "报表项目": report_item,
                        "统一社会信用代码": institution_code, "工作表名": ws.title,
                        "定位单元格": "", "单元格值": None,
                        "判定状态": "无法判定", "条件格式类型": str(rule.type or "未知"),
                        "规则公式": str((rule.formula or [""])[0] or ""),
                        "应用范围": str(sqref), "优先级": getattr(rule, "priority", None),
                        "来源文件": Path(path).name,
                        "判定说明": f"Python 离线求值暂不支持该条件格式类型：{rule.type or '未知'}",
                    })
            rules.sort(key=lambda entry: (int(getattr(entry[0], "priority", 0) or 0), entry[3]))
            cells = set()
            for scan_min_col, scan_min_row, scan_max_col, scan_max_row in scan_regions:
                scan_max_row = min(scan_max_row, max(ws.max_row, scan_min_row))
                scan_max_col = min(scan_max_col, max(ws.max_column, scan_min_col))
                for _rule, _sqref, rule_regions, _order, _anchor_row, _anchor_col in rules:
                    for rule_min_col, rule_min_row, rule_max_col, rule_max_row in rule_regions:
                        min_col = max(scan_min_col, rule_min_col)
                        min_row = max(scan_min_row, rule_min_row)
                        max_col = min(scan_max_col, rule_max_col)
                        max_row = min(scan_max_row, rule_max_row)
                        if min_col <= max_col and min_row <= max_row:
                            cells.update((row, col) for row in range(min_row, max_row + 1)
                                         for col in range(min_col, max_col + 1))
            for row, col in sorted(cells):
                cell = ws.cell(row, col)
                for rule, sqref, regions, _order, anchor_row, anchor_col in rules:
                    matching_bounds = next((bounds for bounds in regions
                                            if bounds[0] <= col <= bounds[2]
                                            and bounds[1] <= row <= bounds[3]), None)
                    if matching_bounds is None:
                        continue
                    rule_type = str(rule.type or "").casefold()
                    try:
                        if rule_type == "cellis":
                            evaluation = evaluate_cellis_rule(
                                rule.operator, list(rule.formula or []), cell.value,
                                lambda formulas, index: _operand(
                                    ws, formulas, index, anchor_row=anchor_row,
                                    anchor_col=anchor_col, row=row, col=col,
                                ),
                            )
                        elif rule_type == "expression":
                            formula = str((rule.formula or [""])[0] or "")
                            evaluation = evaluate_expression_condition(
                                formula, ws, anchor_row, anchor_col, row, col)
                        else:
                            continue
                    except Exception:
                        unresolved.append({
                            "机构名称": institution_name, "报表项目": report_item,
                            "统一社会信用代码": institution_code, "工作表名": ws.title,
                            "定位单元格": cell.coordinate, "单元格值": cell.value,
                            "判定状态": "无法判定", "条件格式类型": rule_type,
                            "规则公式": str((rule.formula or [""])[0] or ""),
                            "应用范围": str(sqref), "优先级": getattr(rule, "priority", None),
                            "来源文件": Path(path).name,
                            "判定说明": "条件格式规则计算失败",
                        })
                        continue
                    if evaluation.status in {STATUS_UNSUPPORTED, STATUS_ERROR}:
                        unresolved.append({
                            "机构名称": institution_name, "报表项目": report_item,
                            "统一社会信用代码": institution_code, "工作表名": ws.title,
                            "定位单元格": cell.coordinate, "单元格值": cell.value,
                            "判定状态": "无法判定", "条件格式类型": rule_type,
                            "规则公式": str((rule.formula or [""])[0] or ""),
                            "应用范围": str(sqref), "优先级": getattr(rule, "priority", None),
                            "来源文件": Path(path).name,
                            "判定说明": evaluation.reason or "条件格式规则计算失败",
                        })
                    if evaluation.status != STATUS_TRUE:
                        continue
                    formula = str((rule.formula or [""])[0] or "")
                    if rule_type == "expression":
                        formula = translate_formula_refs(formula, row-anchor_row, col-anchor_col)
                    results.append({
                        "机构名称": institution_name,
                        "报表项目": report_item,
                        "统一社会信用代码": institution_code,
                        "工作表名": ws.title,
                        "定位单元格": cell.coordinate,
                        "单元格值": cell.value,
                        "判定状态": "已触发",
                        "条件格式类型": rule_type,
                        "规则公式": formula,
                        "应用范围": str(sqref),
                        "优先级": getattr(rule, "priority", None),
                        "来源文件": Path(path).name,
                    })
                    break  # Highest-priority triggered rule wins for a cell.
            # resolve rule priority per cell, and de-duplicate rows from overlapping areas.
        seen = set()
        unique = []
        for row in [*results, *unresolved]:
            if row.get("判定状态") == "无法判定":
                key = ("无法判定", row["工作表名"], row["条件格式类型"],
                       row["规则公式"], row.get("应用范围"))
            else:
                key = (row["判定状态"], row["工作表名"], row["定位单元格"])
            if key not in seen:
                seen.add(key)
                unique.append(row)
        return unique
    finally:
        source_book.close()
        source_rule_book.close()
        if template_book is not source_rule_book:
            template_book.close()
