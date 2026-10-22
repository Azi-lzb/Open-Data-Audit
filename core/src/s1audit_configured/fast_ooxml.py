"""Small, literal-data result tables written directly as an XLSX package.

Complex workbook editing is deliberately outside this writer. The caller keeps
the original openpyxl exporter for types/features not supported here.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape, quoteattr
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.compat import safe_string
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.fonts import DEFAULT_FONT
from openpyxl.styles.numbers import BUILTIN_FORMATS_REVERSE
from openpyxl.utils.cell import get_column_letter
from openpyxl.xml.functions import tostring


NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_HEADER = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'


@dataclass
class DataTable:
    title: str
    headers: tuple
    widths: tuple
    rows: list = field(default_factory=list)
    formats: dict = field(default_factory=dict)
    hyperlinks: dict = field(default_factory=dict)
    numeric_columns: tuple = ()
    text_columns: tuple = ()


def supports_tables(tables):
    """Preflight before opening output, so fallback never leaves a partial ZIP."""
    for table in tables:
        for row in [table.headers, *table.rows]:
            for value in row:
                if value is not None and type(value) not in (str, bool, int, float, Decimal):
                    return False
                if isinstance(value, str) and ILLEGAL_CHARACTERS_RE.search(value):
                    return False
    return True


def _xml(value):
    return escape(str(value)).replace("\r", "&#13;")


def _styles(tables):
    root = ET.Element("styleSheet", xmlns=NS)
    formats = sorted({fmt for table in tables for fmt in table.formats.values()
                      if fmt not in BUILTIN_FORMATS_REVERSE})
    custom_ids = {fmt: index + 164 for index, fmt in enumerate(formats)}
    if formats:
        nodes = ET.SubElement(root, "numFmts", count=str(len(formats)))
        for fmt, index in custom_ids.items():
            ET.SubElement(nodes, "numFmt", numFmtId=str(index), formatCode=fmt)
    fonts = ET.SubElement(root, "fonts", count="3")
    for font in (DEFAULT_FONT, Font(name="微软雅黑", size=10, bold=True, color="FFFFFF"),
                 Font(name="微软雅黑", size=10, color="222222")):
        fonts.append(ET.fromstring(tostring(font.to_tree())))
    fills = ET.SubElement(root, "fills", count="3")
    for fill in (PatternFill(), PatternFill(patternType="gray125"), PatternFill("solid", fgColor="1F4E78")):
        fills.append(ET.fromstring(tostring(fill.to_tree())))
    borders = ET.SubElement(root, "borders", count="1")
    border = ET.SubElement(borders, "border")
    for tag in ("left", "right", "top", "bottom", "diagonal"):
        ET.SubElement(border, tag)
    base = ET.SubElement(root, "cellStyleXfs", count="1")
    ET.SubElement(base, "xf", numFmtId="0", fontId="0", fillId="0", borderId="0")
    xfs = ET.SubElement(root, "cellXfs")
    ET.SubElement(xfs, "xf", numFmtId="0", fontId="0", fillId="0", borderId="0", xfId="0")
    header = ET.SubElement(xfs, "xf", numFmtId="0", fontId="1", fillId="2", borderId="0", xfId="0", applyAlignment="1")
    header.append(ET.fromstring(tostring(Alignment(horizontal="center", vertical="center", wrap_text=True).to_tree())))
    combinations = {("General", False), ("@", False)}
    for table in tables:
        for row_no, row in enumerate(table.rows, 2):
            for col in range(1, len(table.headers) + 1):
                fmt = "@" if col in table.text_columns else table.formats.get((row_no, col), "General")
                combinations.add((fmt, col in table.numeric_columns))
    style_ids = {}
    for fmt, right in sorted(combinations):
        style_ids[(fmt, right)] = len(xfs)
        index = BUILTIN_FORMATS_REVERSE.get(fmt, custom_ids.get(fmt))
        xf = ET.SubElement(xfs, "xf", numFmtId=str(index), fontId="2", fillId="0", borderId="0", xfId="0", applyAlignment="1")
        alignment = Alignment(horizontal="right", vertical="top") if right else Alignment(vertical="top", wrap_text=True)
        xf.append(ET.fromstring(tostring(alignment.to_tree())))
    xfs.set("count", str(len(xfs)))
    styles = ET.SubElement(root, "cellStyles", count="1")
    ET.SubElement(styles, "cellStyle", name="Normal", xfId="0", builtinId="0", hidden="0")
    ET.SubElement(root, "tableStyles", count="0", defaultTableStyle="TableStyleMedium9", defaultPivotStyle="PivotStyleLight16")
    return XML_HEADER + ET.tostring(root), style_ids


def _cell(address, value, style):
    prefix = f'<c r="{address}" s="{style}"'
    if isinstance(value, str):
        value = value[:32767]
        if not value:
            return prefix + ' t="inlineStr"></c>'
        return prefix + ' t="inlineStr"><is><t xml:space="preserve">' + _xml(value) + '</t></is></c>'
    if value is None:
        return prefix + ' t="n"></c>'
    if isinstance(value, bool):
        return prefix + f' t="b"><v>{int(value)}</v></c>'
    return prefix + ' t="n"><v>' + safe_string(value) + '</v></c>'


def _write_sheet(archive, index, table, styles):
    letters = [get_column_letter(col) for col in range(1, len(table.headers) + 1)]
    dimension = f"A1:{letters[-1]}{len(table.rows) + 1}"
    with archive.open(f"xl/worksheets/sheet{index}.xml", "w") as stream:
        def write(text):
            stream.write(text.encode("utf-8"))
        stream.write(XML_HEADER)
        write(f'<worksheet xmlns="{NS}" xmlns:r="{RELS}">'
              '<sheetPr><outlinePr summaryBelow="1" summaryRight="1"/><pageSetUpPr fitToPage="1"/></sheetPr>'
              f'<dimension ref="{dimension}"/><sheetViews><sheetView showGridLines="0" workbookViewId="0">'
              '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
              '<selection pane="bottomLeft" activeCell="A1" sqref="A1"/></sheetView></sheetViews>'
              '<sheetFormatPr baseColWidth="8" defaultRowHeight="15"/><cols>')
        for col, width in enumerate(table.widths, 1):
            write(f'<col width="{width}" customWidth="1" min="{col}" max="{col}"/>')
        write('</cols><sheetData>')
        for row_no, row in enumerate([table.headers, *table.rows], 1):
            write(f'<row r="{row_no}" ht="{30 if row_no == 1 else 54}" customHeight="1">')
            parts = []
            for col, letter in enumerate(letters, 1):
                value = row[col - 1] if col <= len(row) else None
                if row_no == 1:
                    style = 1
                else:
                    fmt = "@" if col in table.text_columns else table.formats.get((row_no, col), "General")
                    style = styles[(fmt, col in table.numeric_columns)]
                parts.append(_cell(f"{letter}{row_no}", value, style))
            write("".join(parts) + '</row>')
        write(f'</sheetData><autoFilter ref="{dimension}"/>')
        if table.hyperlinks:
            write('<hyperlinks>')
            for rel_id, coordinate in enumerate(table.hyperlinks, 1):
                write(f'<hyperlink ref="{coordinate}" r:id="rId{rel_id}"/>')
            write('</hyperlinks>')
        write('<pageMargins left="0.75" right="0.75" top="1" bottom="1" header="0.5" footer="0.5"/>'
              '<pageSetup orientation="landscape" paperSize="8" fitToHeight="0" fitToWidth="1"/></worksheet>')
    if table.hyperlinks:
        relationships = ET.Element("Relationships", xmlns=PKG_RELS)
        for rel_id, target in enumerate(table.hyperlinks.values(), 1):
            ET.SubElement(relationships, "Relationship", Id=f"rId{rel_id}", Type=RELS + "/hyperlink", Target=target, TargetMode="External")
        archive.writestr(f"xl/worksheets/_rels/sheet{index}.xml.rels", XML_HEADER + ET.tostring(relationships))


def write_tables(path: Path, tables):
    if not supports_tables(tables):
        raise TypeError("结果包含直写引擎不支持的类型")
    styles_xml, styles = _styles(tables)
    workbook = ET.Element("workbook", xmlns=NS, attrib={"xmlns:r": RELS})
    ET.SubElement(workbook, "workbookPr")
    sheets = ET.SubElement(workbook, "sheets")
    names = ET.SubElement(workbook, "definedNames")
    relationships = ET.Element("Relationships", xmlns=PKG_RELS)
    types = ET.Element("Types", xmlns="http://schemas.openxmlformats.org/package/2006/content-types")
    ET.SubElement(types, "Default", Extension="rels", ContentType="application/vnd.openxmlformats-package.relationships+xml")
    ET.SubElement(types, "Default", Extension="xml", ContentType="application/xml")
    ET.SubElement(types, "Override", PartName="/xl/workbook.xml", ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml")
    ET.SubElement(types, "Override", PartName="/xl/styles.xml", ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml")
    with ZipFile(path, "w", ZIP_DEFLATED, compresslevel=1) as archive:
        for index, table in enumerate(tables, 1):
            ET.SubElement(sheets, "sheet", name=table.title, sheetId=str(index), attrib={"r:id": f"rId{index}"})
            name = ET.SubElement(names, "definedName", name="_xlnm.Print_Titles", localSheetId=str(index - 1))
            name.text = "'" + table.title.replace("'", "''") + "'!$1:$1"
            ET.SubElement(relationships, "Relationship", Id=f"rId{index}", Type=RELS + "/worksheet", Target=f"worksheets/sheet{index}.xml")
            ET.SubElement(types, "Override", PartName=f"/xl/worksheets/sheet{index}.xml", ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml")
            _write_sheet(archive, index, table, styles)
        ET.SubElement(relationships, "Relationship", Id=f"rId{len(tables) + 1}", Type=RELS + "/styles", Target="styles.xml")
        root_rels = ET.Element("Relationships", xmlns=PKG_RELS)
        ET.SubElement(root_rels, "Relationship", Id="rId1", Type=RELS + "/officeDocument", Target="xl/workbook.xml")
        archive.writestr("[Content_Types].xml", XML_HEADER + ET.tostring(types))
        archive.writestr("_rels/.rels", XML_HEADER + ET.tostring(root_rels))
        archive.writestr("xl/workbook.xml", XML_HEADER + ET.tostring(workbook))
        archive.writestr("xl/_rels/workbook.xml.rels", XML_HEADER + ET.tostring(relationships))
        archive.writestr("xl/styles.xml", styles_xml)
