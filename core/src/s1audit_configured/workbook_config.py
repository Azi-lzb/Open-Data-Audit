from __future__ import annotations

import hashlib
import re
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .dsl import (
    Call, Literal, TupleLiteral, calls_in, parse_dictionary, parse_expression, static_kind,
    validate_dictionary_dependencies, validate_static_types, variables_in,
)
from .errors import ConfigurationError
from .layout import split_path
from .model import (
    ApplicabilityScope, Configuration, ExternalMetricSpec, MetricSpec, RuleDefinition, SourceContract,
)


RULE_HEADERS = (
    "规则ID", "规则修订", "是否启用", "业务类型", "默认数据源", "默认Sheet", "规则名称",
    "审核级别", "校验类型", "校验指标", "审核表达式", "表达式布尔值取反", "规则描述", "当前值表达式",
    "对比值表达式", "差值表达式", "取值字典", "生效起始日", "生效截止日", "适用范围",
    "适用表达式",
)
METRIC_HEADERS = (
    "业务类型", "工作表名称", "行路径", "列路径", "行表头标题", "数据类型", "指标类型",
    "来源单位", "标准单位", "存储尺度", "允许空值", "精度", "口径依据", "确认状态", "布局版本",
)
EXTERNAL_METRIC_HEADERS = (
    "数据源ID", "指标表头", "数据类型", "指标类型", "来源单位", "标准单位", "存储尺度",
    "允许空值", "精度", "口径依据", "确认状态",
)
SOURCE_HEADERS = (
    "数据源ID", "来源类型", "默认Sheet", "机构Sheet", "机构代码标签", "机构名称标签", "期间标签",
    "映射Sheet", "映射表头行", "映射数据起始行", "映射机构字段", "映射目标字段", "数据表头行",
    "数据起始行", "数据键列", "指标表头行", "指标起始列", "报表项目标签", "期望期间",
    "契约版本", "期间必须", "机构必须",
)
VERSION_HEADERS = ("配置版本号", "状态", "发布日期", "生效起始日", "生效截止日", "发布说明")
SCOPE_HEADERS = ("适用范围ID", "业务类型", "工作表名称", "行路径前缀", "范围说明", "确认状态")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _flag(value: Any, *, field: str, default: bool | None = None) -> bool:
    text = _text(value).lower()
    if not text and default is not None:
        return default
    if text in {"是", "是的", "yes", "y", "true", "1", "启用", "已确认"}:
        return True
    if text in {"否", "不是", "no", "n", "false", "0", "禁用", "未确认"}:
        return False
    raise ConfigurationError(f"{field} 必须填写 是/否")


def _integer(value: Any, *, field: str, default: int | None = None) -> int:
    if value is None or _text(value) == "":
        if default is not None:
            return default
        raise ConfigurationError(f"{field} 不能为空")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ConfigurationError(f"{field} 必须是整数") from exc
    if isinstance(value, float) and value != number:
        raise ConfigurationError(f"{field} 必须是整数")
    return number


def _path_parts(value: Any, *, field: str) -> tuple[str, ...]:
    text = _text(value)
    if not text:
        raise ConfigurationError(f"{field} 不能为空")
    parts = tuple(part.strip() for part in re.split(r"\r?\n", text) if part.strip())
    if not parts:
        raise ConfigurationError(f"{field} 不能为空")
    return parts


class _ConfigRow(dict):
    """Keep physical Excel coordinates even after skipping blanks or aliasing headers."""

    def __init__(self, values, sheet, row_num, columns):
        super().__init__(values)
        self.sheet, self.row_num, self.columns = sheet, row_num, dict(columns)
        self.original_headers = {column: header for header, column in columns.items()}

    def location(self, header=""):
        column = self.columns.get(header)
        return {"sheet": self.sheet, "row": self.row_num, "column": column,
                "cell": f"{get_column_letter(column)}{self.row_num}" if column else "",
                "field": self.original_headers.get(column, header),
                "rule_id": _text(self.get("规则ID") or self.get("规则编号") or self.get("规则序号"))}

    def error(self, headers, message, *, related=()):
        headers = [headers] if isinstance(headers, str) else headers
        return ConfigurationError(message, locations=[self.location(header) for header in headers] + list(related))

    def check(self, headers, function, *args, **kwargs):
        try:
            return function(*args, **kwargs)
        except ConfigurationError as exc:
            if exc.locations:
                raise
            raise self.error(headers, exc.message) from exc

    def value(self, header, function, **kwargs):
        return self.check(header, function, self.get(header), **kwargs)

    def dates(self, start_header, end_header, label):
        start, end = _text(self.get(start_header)), _text(self.get(end_header))
        for header, value in ((start_header, start), (end_header, end)):
            if value:
                self.check(header, _validate_date_range, value, value, header)
        self.check((start_header, end_header), _validate_date_range, start, end, label)

    @contextmanager
    def context(self):
        try:
            yield
        except ConfigurationError as exc:
            if exc.locations:
                raise
            # For cross-field constraints, identify all fields named by the validator.
            fields = sorted((header for header in self.columns if header and header in exc.message), key=len, reverse=True)
            raise self.error(fields or [""], exc.message) from exc


def _numbered_rows(rows):
    for row in rows:
        yield row.row_num, row


def _ensure_unique_rows(rows, fields, message, *, keys=None):
    seen = {}
    keys = keys if keys is not None else [tuple(_text(row.get(field)) for field in fields) for row in rows]
    for row, key in zip(rows, keys):
        if key in seen:
            raise row.error(fields, message, related=[seen[key].location(field) for field in fields])
        seen[key] = row


def _header_error(sheet, message, columns=()):
    return ConfigurationError(message, locations=[{"sheet": sheet, "row": 1, "column": column,
        "cell": f"{get_column_letter(column)}1" if column else "", "field": "表头", "rule_id": ""}
        for column in columns or [None]])


def _sheet_rows(workbook, sheet_name: str, headers: Iterable[str], *, required: bool = False) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        if required:
            raise ConfigurationError(f"配置工作簿缺少工作表：{sheet_name}")
        return []
    sheet = workbook[sheet_name]
    header_values = [_text(sheet.cell(1, col).value) for col in range(1, sheet.max_column + 1)]
    while header_values and not header_values[-1]:
        header_values.pop()
    if not header_values:
        raise _header_error(sheet_name, "缺少表头，请在第1行填写字段名称")
    if len(header_values) != len(set(header_values)):
        repeated = next(header for header in header_values if header_values.count(header) > 1)
        columns = [index for index, header in enumerate(header_values, 1) if header == repeated]
        raise _header_error(sheet_name, f"存在重复表头：{repeated or '空白表头'}，请保留唯一字段名称", columns)
    missing = [header for header in headers if header not in header_values]
    if missing:
        raise _header_error(sheet_name, f"第1行缺少字段：{', '.join(missing)}；缺失字段尚无列号，请补充表头")
    columns = {header: index + 1 for index, header in enumerate(header_values)}
    rows: list[dict[str, Any]] = []
    for row_num in range(2, sheet.max_row + 1):
        values = _ConfigRow({header: sheet.cell(row_num, col).value for header, col in columns.items()}, sheet_name, row_num, columns)
        if not any(value is not None and _text(value) for value in values.values()):
            continue
        for header, col in columns.items():
            cell = sheet.cell(row_num, col)
            if cell.data_type == "f":
                raise values.error(header, "不允许使用 Excel 公式，请将审核表达式按文本保存，开头不要加 =")
        rows.append(values)
    return rows


def _required_text(row: dict[str, Any], header: str, *, sheet: str, row_num: int) -> str:
    value = _text(row.get(header))
    if not value:
        if isinstance(row, _ConfigRow):
            raise row.error(header, f"{header} 不能为空")
        raise ConfigurationError(f"{sheet} 第 {row_num} 行 {header} 不能为空")
    return value


def _description_variables(text: str) -> set[str]:
    found: set[str] = set()
    index = 0
    while index < len(text):
        if text.startswith("{{", index) or text.startswith("}}", index):
            index += 2
            continue
        if text[index] == "{":
            end = text.find("}", index + 1)
            if end < 0:
                raise ConfigurationError("规则描述中的变量占位符缺少结束大括号")
            name = text[index + 1:end]
            if not name or not re.fullmatch(r"[\w\u3400-\u9fff]+", name):
                raise ConfigurationError("规则描述中的变量占位符格式错误")
            found.add(name)
            index = end + 1
            continue
        if text[index] == "}":
            raise ConfigurationError("规则描述中的右大括号需要写成 }}")
        index += 1
    return found


def _literal_path(node: Any) -> tuple[str, ...]:
    if isinstance(node, Literal) and node.kind == "string":
        return split_path(node.value)
    if isinstance(node, TupleLiteral):
        return tuple(_text(value) for value in node.values if _text(value))
    return ()


def _metric_profile(labels: Iterable[str], *, source_unit: str = "", standard_unit: str = "",
                    metric_type: str = "") -> tuple[str, str, str, str]:
    """Infer the small set of common metric units when no catalog sheet is used."""
    label = " ".join(str(item) for item in labels if item).lower()
    kind = _canonical_metric_type(metric_type) if metric_type else ""
    if not kind:
        if source_unit in {"户", "笔"}:
            kind = "count"
        elif source_unit in {"百分比", "比例", "百分点", "%"} or any(token in label for token in ("利率", "利息率", "百分比", "占比", "%")):
            kind = "rate"
        elif any(token in label for token in ("户数", "贷款户", "客户户数", "（户）", "(户)")):
            kind = "count"
            source_unit = source_unit or "户"
        elif any(token in label for token in ("笔数", "业务笔数", "发生笔数", "（笔）", "(笔)")):
            kind = "count"
            source_unit = source_unit or "笔"
        else:
            kind = "amount"

    if kind == "amount":
        source_unit = source_unit or "万元"
        standard_unit = standard_unit or "万元"
        scale = "native"
    elif kind == "rate":
        source_unit = _canonical_unit(source_unit or "百分比", "rate")
        standard_unit = _canonical_unit(standard_unit or source_unit, "rate")
        scale = "native"
    elif kind == "count":
        source_unit = source_unit or ("户" if "户数" in label or "贷款户" in label else "笔")
        standard_unit = standard_unit or source_unit
        scale = "native"
    else:
        source_unit = source_unit or "万元"
        standard_unit = standard_unit or source_unit
        scale = "native"
    return kind, source_unit, standard_unit, scale


def _template_metric_from_call(call: Call, rule: RuleDefinition) -> MetricSpec:
    if len(call.arguments) == 3:
        sheet_node, row_node, column_node = call.arguments
        sheet = _text(sheet_node.value) if isinstance(sheet_node, Literal) and sheet_node.kind == "string" else ""
    elif len(call.arguments) == 2:
        row_node, column_node = call.arguments
        sheet = dict(call.keywords).get("sheet", "")
    else:
        raise ConfigurationError(f"{call.name} 必须传入行路径、列路径，可选工作表名")
    sheet = sheet or rule.default_sheet
    if not sheet:
        raise ConfigurationError(f"规则 {rule.rule_id} 的 {call.name} 缺少工作表名，请使用 V_CUR(工作表, 行路径, 列路径)")
    row_path = _literal_path(row_node)
    column_path = _literal_path(column_node)
    if not row_path or not column_path or any(not part for part in (*row_path, *column_path)):
        raise ConfigurationError(f"规则 {rule.rule_id} 的 {call.name} 行路径和列路径必须是字符串或字符串元组")
    keywords = dict(call.keywords)
    kind, source_unit, standard_unit, storage_scale = _metric_profile(
        (*row_path, *column_path),
        source_unit=keywords.get("unit", keywords.get("source_unit", "")),
        standard_unit=keywords.get("standard_unit", ""),
        metric_type=keywords.get("metric_type", ""),
    )
    return MetricSpec(
        business=rule.business, sheet=sheet, row_path=row_path, column_path=column_path,
        row_header_titles=(), data_type="number", metric_type=kind,
        source_unit=source_unit, standard_unit=standard_unit, storage_scale=storage_scale,
        allow_blank=True, precision=2, confirmed=True,
    )


def _external_metric_from_call(call: Call, rule: RuleDefinition) -> ExternalMetricSpec:
    if not call.arguments or not isinstance(call.arguments[0], Literal) or call.arguments[0].kind != "string":
        raise ConfigurationError(f"规则 {rule.rule_id} 的 V_EXT 指标名必须是字符串")
    metric_header = _text(call.arguments[0].value)
    if not metric_header:
        raise ConfigurationError(f"规则 {rule.rule_id} 的 V_EXT 指标名不能为空")
    keywords = dict(call.keywords)
    unit = keywords.get("unit") or keywords.get("source_unit") or ""
    if not unit:
        unit_match = re.search(r"[（(]\s*(元|万元|亿元|百分比|百分点|户|笔)\s*[）)]", metric_header)
        unit = unit_match.group(1) if unit_match else ""
    kind, source_unit, standard_unit, storage_scale = _metric_profile(
        (metric_header,),
        source_unit=unit,
        standard_unit=keywords.get("standard_unit", ""),
        metric_type=keywords.get("metric_type", ""),
    )
    return ExternalMetricSpec(
        source_id=keywords.get("source", "EXT") or "EXT", metric_header=metric_header,
        data_type="number", metric_type=kind, source_unit=source_unit,
        standard_unit=standard_unit, storage_scale=storage_scale,
        allow_blank=True, precision=2, confirmed=True,
    )


def _infer_rule_business(rule_num: int, calls: list[Call], metric_catalog: list[tuple[str, str, tuple[str, ...], tuple[str, ...]]],
                         scopes: list[ApplicabilityScope], applicability_scope: str) -> str:
    if applicability_scope:
        scope = next((item for item in scopes if item.scope_id == applicability_scope), None)
        if scope:
            return scope.business

    matching_business_sets: list[set[str]] = []
    for call in calls:
        if call.name not in {"V_CUR", "V_PRE"}:
            continue
        if len(call.arguments) == 3:
            sheet_node, row_node, column_node = call.arguments
            sheet = _text(sheet_node.value) if isinstance(sheet_node, Literal) and sheet_node.kind == "string" else ""
        else:
            row_node, column_node = call.arguments
            sheet = dict(call.keywords).get("sheet", "")
        row_path = _literal_path(row_node)
        column_path = _literal_path(column_node)
        if row_path == ("@当前行",):
            continue
        if not row_path or not column_path:
            continue
        matches = {
            business for business, metric_sheet, metric_row_path, metric_column_path in metric_catalog
            if metric_row_path == row_path and metric_column_path == column_path
            and (not sheet or metric_sheet == sheet)
        }
        if matches:
            matching_business_sets.append(matches)

    if matching_business_sets:
        common = set.intersection(*matching_business_sets)
        if len(common) == 1:
            return next(iter(common))
        if len(common) > 1:
            raise ConfigurationError(f"审核规则第 {rule_num} 行的业务类型无法由指标路径唯一推断，请补充业务类型")
        raise ConfigurationError(f"审核规则第 {rule_num} 行引用的指标跨越多个业务，请补充业务类型")

    businesses = {business for business, _, _, _ in metric_catalog if business}
    if len(businesses) == 1:
        return next(iter(businesses))
    raise ConfigurationError(f"审核规则第 {rule_num} 行缺少业务类型，且无法从指标路径唯一推断")


def load_configuration(path: str | Path, *, include_disabled: bool = False, on_progress=None,
                       business: str = "") -> Configuration:
    """Load and validate an XLSX configuration workbook.

    The workbook is read with formulas visible so configuration cells containing
    Excel formulas are rejected instead of silently using a cached value.
    """
    config_path = Path(path).resolve()
    if not config_path.is_file():
        raise ConfigurationError(f"配置工作簿不存在：{config_path}")
    try:
        config_hash_before = _sha256(config_path)
        workbook = load_workbook(config_path, data_only=False, read_only=False)
    except Exception as exc:
        raise ConfigurationError(f"配置工作簿无法读取：{exc}") from exc
    try:
        raw_rules = _sheet_rows(workbook, "审核规则", (), required=True)
        if not raw_rules:
            raise ConfigurationError("审核规则工作表没有规则数据")
        rule_headers = set(raw_rules[0])
        id_header = next((header for header in ("规则ID", "规则编号", "规则序号") if header in rule_headers), "")
        if not id_header:
            raise _header_error("审核规则", "第1行缺少字段：规则ID（也可使用规则编号或规则序号）；请补充表头")
        required_rule_headers = {"是否启用", "规则名称", "审核级别", "校验类型", "校验指标", "审核表达式", "规则描述"}
        missing_rule_headers = sorted(required_rule_headers - rule_headers)
        if missing_rule_headers:
            raise _header_error("审核规则", f"第1行缺少字段：{', '.join(missing_rule_headers)}；缺失字段尚无列号，请补充表头")
        for row in raw_rules:
            row["规则ID"] = _text(row.get("规则ID")) or _text(row.get(id_header))
            row.columns.setdefault("规则ID", row.columns[id_header])
            for alias, preferred in (("业务类型", "业务名称"), ("默认Sheet", "工作表名称")):
                if preferred in row.columns:
                    row.columns.setdefault(alias, row.columns[preferred])
        requested_business = _text(business)
        if requested_business:
            raw_rules = [
                row for row in raw_rules
                if (_text(row.get("业务名称")) or _text(row.get("业务类型"))) == requested_business
            ]
        raw_metric_rows = _sheet_rows(workbook, "模板指标目录", METRIC_HEADERS)
        if requested_business:
            raw_metric_rows = [row for row in raw_metric_rows
                               if _text(row.get("业务类型")) == requested_business]
        metric_catalog = [
            (
                _required_text(row, "业务类型", sheet="模板指标目录", row_num=row_num),
                _required_text(row, "工作表名称", sheet="模板指标目录", row_num=row_num),
                row.value("行路径", _path_parts, field=f"模板指标目录第 {row_num} 行行路径"),
                row.value("列路径", _path_parts, field=f"模板指标目录第 {row_num} 行列路径"),
            )
            for row_num, row in _numbered_rows(raw_metric_rows)
        ]
        scopes: list[ApplicabilityScope] = []
        raw_scope_rows = _sheet_rows(workbook, "适用范围目录", SCOPE_HEADERS)
        if requested_business:
            raw_scope_rows = [row for row in raw_scope_rows
                              if _text(row.get("业务类型")) == requested_business]
        _ensure_unique_rows(raw_scope_rows, ("适用范围ID",), "适用范围目录存在重复适用范围ID")
        for row_num, row in _numbered_rows(raw_scope_rows):
            with row.context():
                scopes.append(ApplicabilityScope(
                    scope_id=_required_text(row, "适用范围ID", sheet="适用范围目录", row_num=row_num),
                    business=_required_text(row, "业务类型", sheet="适用范围目录", row_num=row_num),
                    sheet=_required_text(row, "工作表名称", sheet="适用范围目录", row_num=row_num),
                    row_path_prefix=row.value("行路径前缀", _path_parts, field=f"适用范围目录第 {row_num} 行行路径前缀"),
                    description=_text(row.get("范围说明")),
                    confirmed=row.value("确认状态", _flag, field=f"适用范围目录第 {row_num} 行确认状态"),
                ))
        scope_ids = [scope.scope_id for scope in scopes]
        rules: list[RuleDefinition] = []
        derived_metrics: dict[tuple[str, str, tuple[str, ...], tuple[str, ...]], MetricSpec] = {}
        derived_external: dict[tuple[str, str], ExternalMetricSpec] = {}
        has_metric_sheet = "模板指标目录" in workbook.sheetnames
        has_external_sheet = "外部指标目录" in workbook.sheetnames
        has_source_sheet = "数据源契约" in workbook.sheetnames
        seen_revisions: dict[tuple[str, int], _ConfigRow] = {}
        referenced_template = False
        template_references = []
        referenced_external = []
        external_locations = {}
        enabled_rules_seen = 0
        for position, (row_num, row) in enumerate(_numbered_rows(raw_rules), 1):
            with row.context():
                rule_id = _required_text(row, "规则ID", sheet="审核规则", row_num=row_num)
                revision = row.value("规则修订", _integer, field=f"审核规则第 {row_num} 行规则修订", default=1)
                if revision < 1:
                    raise ConfigurationError(f"审核规则第 {row_num} 行规则修订必须大于零")
                identity = (rule_id, revision)
                if identity in seen_revisions:
                    raise row.error("规则ID", f"规则ID/修订重复：{rule_id} / {revision}", related=[seen_revisions[identity].location("规则ID")])
                seen_revisions[identity] = row
                enabled = row.value("是否启用", _flag, field=f"审核规则第 {row_num} 行是否启用")
                enabled_rules_seen += int(enabled)
                if on_progress:
                    on_progress({
                        "phase": "rules", "current": position, "total": len(raw_rules),
                        "enabled_rules": enabled_rules_seen,
                        "message": f"逐行检查审核规则配置 {position}/{len(raw_rules)}；已启用 {enabled_rules_seen} 条（禁用规则跳过表达式和字典检查）",
                    })
                business = _text(row.get("业务名称")) or _text(row.get("业务类型"))
                default_sheet = _text(row.get("工作表名称")) or _text(row.get("默认Sheet"))
                name = _required_text(row, "规则名称", sheet="审核规则", row_num=row_num)
                severity = _required_text(row, "审核级别", sheet="审核规则", row_num=row_num)
                check_type = _required_text(row, "校验类型", sheet="审核规则", row_num=row_num)
                metric = _required_text(row, "校验指标", sheet="审核规则", row_num=row_num)
                audit_expression = _required_text(row, "审核表达式", sheet="审核规则", row_num=row_num)
                numeric_mode_label = _text(row.get("数值运算模式")) or "十进制"
                numeric_modes = {"十进制": "decimal", "Decimal": "decimal",
                                 "Excel双精度": "excel_binary64", "Excel binary64": "excel_binary64"}
                if numeric_mode_label not in numeric_modes:
                    raise row.error("数值运算模式", "数值运算模式仅支持 十进制 或 Excel双精度")
                numeric_mode = numeric_modes[numeric_mode_label]
                formula_cell = _text(row.get("校验公式单元格"))
                invert_audit_expression = row.value(
                    "表达式布尔值取反", _flag,
                    field=f"审核规则第 {row_num} 行表达式布尔值取反", default=False,
                )
                description = _required_text(row, "规则描述", sheet="审核规则", row_num=row_num)
                rule_effective_from = _text(row.get("生效起始日"))
                rule_effective_to = _text(row.get("生效截止日"))
                applicability_scope = _text(row.get("适用范围"))
                row.dates("生效起始日", "生效截止日", "生效期")
                if applicability_scope and applicability_scope not in set(scope_ids):
                    raise ConfigurationError(f"审核规则第 {row_num} 行引用了不存在的适用范围ID：{applicability_scope}")
                if enabled or include_disabled:
                    dictionary_nodes = row.check("取值字典", parse_dictionary, _text(row.get("取值字典")))
                    row.check("取值字典", validate_dictionary_dependencies, dictionary_nodes)
                    for node in dictionary_nodes.values():
                        row.check("取值字典", validate_static_types, node)
                    expression_nodes = {
                        "审核表达式": row.check("审核表达式", parse_expression, audit_expression),
                    }
                    optional_expressions = {
                        "当前值表达式": _text(row.get("当前值表达式")),
                        "对比值表达式": _text(row.get("对比值表达式")),
                        "差值表达式": _text(row.get("差值表达式")),
                        "适用表达式": _text(row.get("适用表达式")),
                    }
                    for field_name, expression in optional_expressions.items():
                        if expression:
                            expression_nodes[field_name] = row.check(field_name, parse_expression, expression)
                    for field_name, node in expression_nodes.items():
                        row.check(field_name, validate_static_types, node)
                        kind = static_kind(node)
                        if field_name in {"审核表达式", "适用表达式"} and kind not in {None, "boolean"}:
                            raise row.error(field_name, f"{field_name} 必须返回布尔值")
                        if field_name == "差值表达式" and kind not in {None, "number"}:
                            raise row.error("差值表达式", "差值表达式必须返回数值")
                    for field_name, node in expression_nodes.items():
                        undefined = variables_in(node) - dictionary_nodes.keys()
                        if undefined:
                            raise row.error(field_name, f"{field_name} 引用未定义变量：{', '.join(sorted(undefined))}")
                    rule_calls = [
                        call
                        for node in [*dictionary_nodes.values(), *expression_nodes.values()]
                        for call in calls_in(node)
                    ]
                    current_row_calls = [
                        call for call in rule_calls
                        if call.name in {"V_CUR", "V_PRE"}
                        and ((len(call.arguments) == 2 and getattr(call.arguments[0], "value", None) == "@当前行")
                             or (len(call.arguments) == 3 and getattr(call.arguments[1], "value", None) == "@当前行"))
                    ]
                    uses_current_row = bool(current_row_calls)
                    if uses_current_row and not applicability_scope:
                        owner = next(field for field, node in [("取值字典", node) for node in dictionary_nodes.values()] + list(expression_nodes.items()) if current_row_calls[0] in calls_in(node))
                        raise row.error(owner, "使用 @当前行，但没有配置适用范围ID，请补充适用范围")
                    if applicability_scope and not uses_current_row:
                        raise ConfigurationError(f"审核规则第 {row_num} 行配置了适用范围，但表达式没有使用 @当前行")
                    if not business:
                        field = next((name for name in ("业务名称", "业务类型") if name in row.columns), "审核表达式")
                        business = row.check(field, _infer_rule_business, row_num, rule_calls, metric_catalog, scopes, applicability_scope)
                    if applicability_scope:
                        scope = next(item for item in scopes if item.scope_id == applicability_scope)
                        if not scope.confirmed:
                            raise ConfigurationError(f"审核规则第 {row_num} 行引用的适用范围尚未确认：{applicability_scope}")
                        if scope.business != business:
                            raise ConfigurationError(f"审核规则第 {row_num} 行适用范围业务与规则业务不一致")
                        if default_sheet and scope.sheet != default_sheet:
                            raise ConfigurationError(f"审核规则第 {row_num} 行适用范围工作表与默认Sheet不一致")
                        for call in current_row_calls:
                            if len(call.arguments) == 3 and getattr(call.arguments[0], "value", None) != scope.sheet:
                                raise ConfigurationError(
                                    f"审核规则第 {row_num} 行 @当前行 的显式工作表与适用范围工作表不一致"
                                )
                    for field_name, node in [("取值字典", node) for node in dictionary_nodes.values()] + list(expression_nodes.items()):
                        for call in calls_in(node):
                            if call.name in {"V_CUR", "V_PRE"}:
                                referenced_template = True
                                template_references.append((row, field_name))
                            elif call.name == "V_EXT":
                                metric_header = call.arguments[0].value if getattr(call.arguments[0], "kind", "") == "string" else ""
                                source_id = dict(call.keywords).get("source", "")
                                referenced_external.append((source_id, metric_header, row, field_name))
                    description_vars = row.check("规则描述", _description_variables, description)
                    if description_vars - dictionary_nodes.keys():
                        raise row.error("规则描述", "规则描述引用未定义变量：" + ", ".join(sorted(description_vars - dictionary_nodes.keys())))
                elif not business and not include_disabled:
                    field = next((name for name in ("业务名称", "业务类型") if name in row.columns), "规则名称")
                    business = row.check(field, _infer_rule_business, row_num, [], metric_catalog, scopes, applicability_scope)
                rule_definition = RuleDefinition(
                    rule_id=rule_id,
                    revision=revision,
                    enabled=enabled,
                    business=business,
                    default_source=_text(row.get("默认数据源")),
                    default_sheet=default_sheet,
                    name=name,
                    severity=severity,
                    check_type=check_type,
                    metric=metric,
                    audit_expression=audit_expression,
                    description=description,
                    audit_expression_inverted=invert_audit_expression,
                    current_expression=_text(row.get("当前值表达式")),
                    comparison_expression=_text(row.get("对比值表达式")),
                    difference_expression=_text(row.get("差值表达式")),
                    value_dictionary=_text(row.get("取值字典")),
                    effective_from=rule_effective_from,
                    effective_to=rule_effective_to,
                    applicable_expression=_text(row.get("适用表达式")),
                    applicability_scope=applicability_scope,
                    numeric_mode=numeric_mode,
                    formula_cell=formula_cell,
                )
                rules.append(rule_definition)
                if (enabled or include_disabled) and not has_metric_sheet:
                    for field_name, node in [("取值字典", node) for node in dictionary_nodes.values()] + list(expression_nodes.items()):
                        for call in calls_in(node):
                            if call.name in {"V_CUR", "V_PRE"}:
                                spec = row.check(field_name, _template_metric_from_call, call, rule_definition)
                                key = (spec.business, spec.sheet, spec.row_path, spec.column_path)
                                previous_spec = derived_metrics.get(key)
                                if previous_spec is not None and previous_spec != spec:
                                    raise row.error(field_name,
                                        f"审核规则引用同一模板路径时单位定义不一致：{spec.sheet} / {spec.row_path} / {spec.column_path}"
                                    )
                                derived_metrics[key] = spec
                            elif call.name == "V_EXT" and not has_external_sheet:
                                spec = row.check(field_name, _external_metric_from_call, call, rule_definition)
                                key = (spec.source_id, spec.metric_header)
                                previous_spec = derived_external.get(key)
                                if previous_spec is not None and previous_spec != spec:
                                    raise row.error(field_name,
                                        f"审核规则引用同一外部指标时单位定义不一致：{spec.source_id} / {spec.metric_header}"
                                    )
                                derived_external[key] = spec
                                external_locations[key] = (row, field_name)
        if on_progress:
            on_progress({
                "phase": "rules", "current": len(raw_rules), "total": len(raw_rules),
                "enabled_rules": enabled_rules_seen,
                "message": f"审核规则配置检查完成，共 {len(raw_rules)} 行",
            })

        metric_specs: list[MetricSpec] = []
        for row_num, row in _numbered_rows(raw_metric_rows):
            with row.context():
                confirmed = row.value("确认状态", _flag, field=f"模板指标目录第 {row_num} 行确认状态")
                metric_type = _canonical_metric_type(row.get("指标类型"))
                spec = MetricSpec(
                    business=_required_text(row, "业务类型", sheet="模板指标目录", row_num=row_num),
                    sheet=_required_text(row, "工作表名称", sheet="模板指标目录", row_num=row_num),
                    row_path=row.value("行路径", _path_parts, field=f"模板指标目录第 {row_num} 行行路径"),
                    column_path=row.value("列路径", _path_parts, field=f"模板指标目录第 {row_num} 行列路径"),
                    row_header_titles=row.value("行表头标题", _path_parts, field=f"模板指标目录第 {row_num} 行行表头标题"),
                    data_type=_canonical_data_type(row.get("数据类型")),
                    metric_type=metric_type,
                    source_unit=_canonical_unit(_text(row.get("来源单位")), metric_type),
                    standard_unit=_canonical_unit(_text(row.get("标准单位")), metric_type),
                    storage_scale=_canonical_storage_scale(row.get("存储尺度")),
                    allow_blank=row.value("允许空值", _flag, field=f"模板指标目录第 {row_num} 行允许空值", default=False),
                    precision=row.value("精度", _integer, field=f"模板指标目录第 {row_num} 行精度", default=2),
                    evidence=_text(row.get("口径依据")),
                    confirmed=confirmed,
                    layout_version=_text(row.get("布局版本")),
                )
                if len(spec.row_path) > len(spec.row_header_titles):
                    raise ConfigurationError(f"模板指标目录第 {row_num} 行行路径层级数不能多于行表头标题数")
                row.check("数据类型", _validate_data_type, spec.data_type, spec.metric_type, row_num)
                row.check("精度", _validate_precision, spec.precision, row_num)
                _validate_metric_units(spec.metric_type, spec.source_unit, spec.standard_unit, spec.storage_scale, row_num, row=row)
                metric_specs.append(spec)
        if not has_metric_sheet:
            metric_specs = list(derived_metrics.values())

        external_specs: list[ExternalMetricSpec] = []
        raw_external_rows = _sheet_rows(workbook, "外部指标目录", EXTERNAL_METRIC_HEADERS)
        if requested_business:
            referenced_external_keys = [
                (source_id, metric_header) for source_id, metric_header, _, _ in referenced_external
            ]
            raw_external_rows = [
                row for row in raw_external_rows
                if any(
                    metric_header == _text(row.get("指标表头"))
                    and (not source_id or source_id == _text(row.get("数据源ID")))
                    for source_id, metric_header in referenced_external_keys
                )
            ]
        for row_num, row in _numbered_rows(raw_external_rows):
            with row.context():
                metric_type = _canonical_metric_type(row.get("指标类型"))
                spec = ExternalMetricSpec(
                    source_id=_required_text(row, "数据源ID", sheet="外部指标目录", row_num=row_num),
                    metric_header=_required_text(row, "指标表头", sheet="外部指标目录", row_num=row_num),
                    data_type=_canonical_data_type(row.get("数据类型")),
                    metric_type=metric_type,
                    source_unit=_canonical_unit(_text(row.get("来源单位")), metric_type),
                    standard_unit=_canonical_unit(_text(row.get("标准单位")), metric_type),
                    storage_scale=_canonical_storage_scale(row.get("存储尺度")),
                    allow_blank=row.value("允许空值", _flag, field=f"外部指标目录第 {row_num} 行允许空值", default=False),
                    precision=row.value("精度", _integer, field=f"外部指标目录第 {row_num} 行精度", default=2),
                    evidence=_text(row.get("口径依据")),
                    confirmed=row.value("确认状态", _flag, field=f"外部指标目录第 {row_num} 行确认状态"),
                )
                row.check("数据类型", _validate_data_type, spec.data_type, spec.metric_type, row_num)
                row.check("精度", _validate_precision, spec.precision, row_num)
                _validate_metric_units(spec.metric_type, spec.source_unit, spec.standard_unit, spec.storage_scale, row_num, row=row)
                external_specs.append(spec)
                external_locations[(spec.source_id, spec.metric_header)] = (row, "数据源ID")
        if not has_external_sheet:
            external_specs = list(derived_external.values())

        contracts: list[SourceContract] = []
        raw_contract_rows = _sheet_rows(workbook, "数据源契约", SOURCE_HEADERS)
        if requested_business:
            referenced_source_ids = {spec.source_id for spec in external_specs if spec.source_id}
            raw_contract_rows = [row for row in raw_contract_rows
                                 if _text(row.get("数据源ID")) in referenced_source_ids]
        for row_num, row in _numbered_rows(raw_contract_rows):
            with row.context():
                source_type = _required_text(row, "来源类型", sheet="数据源契约", row_num=row_num).lower()
                if source_type not in {"current", "previous", "external", "本期", "上期", "外部"}:
                    raise ConfigurationError(f"数据源契约第 {row_num} 行来源类型不支持：{source_type}")
                contract = SourceContract(
                    source_id=_required_text(row, "数据源ID", sheet="数据源契约", row_num=row_num),
                    source_type=source_type,
                    default_sheet=_text(row.get("默认Sheet")),
                    identity_sheet=_text(row.get("机构Sheet")),
                    institution_code_label=_text(row.get("机构代码标签")),
                    institution_name_label=_text(row.get("机构名称标签")),
                    period_label=_text(row.get("期间标签")),
                    mapping_sheet=_text(row.get("映射Sheet")),
                    mapping_header_row=row.value("映射表头行", _integer, field=f"数据源契约第 {row_num} 行映射表头行", default=1),
                    mapping_data_start_row=row.value("映射数据起始行", _integer, field=f"数据源契约第 {row_num} 行映射数据起始行", default=2),
                    mapping_identity_field=_text(row.get("映射机构字段")),
                    mapping_target_field=_text(row.get("映射目标字段")),
                    data_header_row=row.value("数据表头行", _integer, field=f"数据源契约第 {row_num} 行数据表头行", default=1),
                    data_start_row=row.value("数据起始行", _integer, field=f"数据源契约第 {row_num} 行数据起始行", default=2),
                    data_key_column=row.value("数据键列", _integer, field=f"数据源契约第 {row_num} 行数据键列", default=1),
                    metric_header_row=row.value("指标表头行", _integer, field=f"数据源契约第 {row_num} 行指标表头行", default=1),
                    metric_first_column=row.value("指标起始列", _integer, field=f"数据源契约第 {row_num} 行指标起始列", default=2),
                    report_item_label=_text(row.get("报表项目标签")),
                    expected_period=_text(row.get("期望期间")),
                    version=_text(row.get("契约版本")),
                    period_required=row.value("期间必须", _flag, field=f"数据源契约第 {row_num} 行期间必须", default=False),
                    institution_required=row.value("机构必须", _flag, field=f"数据源契约第 {row_num} 行机构必须", default=False),
                )
                if min(contract.mapping_header_row, contract.mapping_data_start_row, contract.data_header_row,
                       contract.data_start_row, contract.data_key_column, contract.metric_header_row,
                       contract.metric_first_column) < 1:
                    fields = ("映射表头行", "映射数据起始行", "数据表头行", "数据起始行", "数据键列", "指标表头行", "指标起始列")
                    invalid = [field for field in fields if _integer(row.get(field), field=field, default=1) < 1]
                    raise row.error(invalid, "表头/起始行/列号必须大于零")
                if source_type in {"external", "外部"} and not all((contract.default_sheet, contract.mapping_sheet,
                                                                      contract.mapping_identity_field,
                                                                      contract.mapping_target_field)):
                    missing = [field for field in ("默认Sheet", "映射Sheet", "映射机构字段", "映射目标字段") if not _text(row.get(field))]
                    raise row.error(missing, "外部数据源契约必须配置默认Sheet和映射字段")
                contracts.append(contract)

        version = version_status = release_date = effective_from = effective_to = release_note = ""
        version_rows = _sheet_rows(workbook, "配置版本", VERSION_HEADERS)
        if len(version_rows) > 1:
            raise version_rows[1].error("配置版本号", "配置版本工作表只能有一行当前版本信息", related=[version_rows[0].location("配置版本号")])
        if version_rows:
            row = version_rows[0]
            version = _required_text(row, "配置版本号", sheet="配置版本", row_num=row.row_num)
            version_status = _required_text(row, "状态", sheet="配置版本", row_num=row.row_num)
            release_date = _text(row.get("发布日期"))
            if release_date:
                row.check("发布日期", _validate_date_range, release_date, release_date, "发布日期")
            effective_from = _text(row.get("生效起始日"))
            effective_to = _text(row.get("生效截止日"))
            row.dates("生效起始日", "生效截止日", "生效期")
            release_note = _text(row.get("发布说明"))

        _ensure_unique_rows(raw_metric_rows, ("业务类型", "工作表名称", "行路径", "列路径"), "模板指标目录存在重复业务/Sheet/行路径/列路径",
                            keys=[(spec.business, spec.sheet, spec.row_path, spec.column_path) for spec in metric_specs])
        _ensure_unique_rows(raw_external_rows, ("数据源ID", "指标表头"), "外部指标目录存在重复数据源/指标表头",
                            keys=[(spec.source_id, spec.metric_header) for spec in external_specs])
        _ensure_unique_rows(raw_contract_rows, ("数据源ID",), "数据源契约存在重复数据源ID")
        for scope in scopes:
            if not any(
                metric.business == scope.business and metric.sheet == scope.sheet
                and metric.row_path[:len(scope.row_path_prefix)] == scope.row_path_prefix
                for metric in metric_specs
            ):
                scope_row = next(row for row in raw_scope_rows if _text(row.get("适用范围ID")) == scope.scope_id)
                raise scope_row.error("行路径前缀", f"适用范围 {scope.scope_id} 没有匹配到模板指标目录行路径")
        if referenced_template and not metric_specs:
            owner, field = template_references[0]
            raise owner.error(field, "规则使用 V_CUR/V_PRE，但无法从审核规则表达式生成模板取值路径，请检查模板指标目录是否为空")
        if referenced_external and not external_specs:
            _, _, owner, field = referenced_external[0]
            raise owner.error(field, "规则使用 V_EXT，但无法从审核规则表达式生成外部指标信息，请检查外部指标目录是否为空")
        contract_ids = {contract.source_id for contract in contracts}
        contract_by_id = {contract.source_id: contract for contract in contracts}
        for spec in external_specs:
            owner, field = external_locations[(spec.source_id, spec.metric_header)]
            if spec.source_id not in contract_ids and has_source_sheet:
                raise owner.error(field, f"外部指标 {spec.metric_header} 缺少对应数据源契约：{spec.source_id}")
            if spec.source_id in contract_by_id and contract_by_id[spec.source_id].source_type not in {"external", "外部"}:
                raise owner.error(field, f"外部指标 {spec.metric_header} 引用了非外部来源契约：{spec.source_id}")
        external_keys = {(spec.source_id, spec.metric_header) for spec in external_specs}
        external_headers = {spec.metric_header for spec in external_specs}
        for source_id, metric_header, row, field_name in referenced_external:
            if source_id and (source_id, metric_header) not in external_keys and has_external_sheet:
                raise row.error(field_name, f"V_EXT 引用的外部指标未登记：{source_id} / {metric_header}")
            if not source_id and metric_header not in external_headers and has_external_sheet:
                raise row.error(field_name, f"V_EXT 引用的外部指标未登记：{metric_header}")
        config_hash_after = _sha256(config_path)
        if config_hash_before != config_hash_after:
            raise ConfigurationError("配置工作簿在加载期间发生变化，请保存稳定版本后重试")
        return Configuration(
            rules=tuple(rules), metrics=tuple(metric_specs), external_metrics=tuple(external_specs),
            source_contracts=tuple(contracts), applicability_scopes=tuple(scopes),
            version=version, version_status=version_status,
            effective_from=effective_from, effective_to=effective_to, release_date=release_date,
            release_note=release_note,
            sha256=config_hash_before, path=config_path,
        )
    finally:
        workbook.close()


def _validate_metric_units(metric_type: str, source_unit: str, standard_unit: str,
                           storage_scale: str, row_num: int, *, row=None) -> None:
    def fail(fields, message):
        raise row.error(fields, message) if row is not None else ConfigurationError(message)

    if metric_type not in {"amount", "rate", "count", "number", "text"}:
        fail("指标类型", f"指标类型不支持：{metric_type}")
    if metric_type == "amount":
        if source_unit not in {"元", "万元", "亿元"} or standard_unit not in {"元", "万元", "亿元"}:
            invalid = [field for field, unit in (("来源单位", source_unit), ("标准单位", standard_unit)) if unit not in {"元", "万元", "亿元"}]
            fail(invalid, "金额指标单位必须是元/万元/亿元")
    elif metric_type == "rate":
        if source_unit not in {"百分比", "%", "百分点", "比例"} or standard_unit not in {"百分比", "%", "百分点", "比例"}:
            invalid = [field for field, unit in (("来源单位", source_unit), ("标准单位", standard_unit)) if unit not in {"百分比", "%", "百分点", "比例"}]
            fail(invalid, "利率指标必须明确比例单位/尺度")
    elif metric_type == "count":
        if source_unit not in {"户", "笔"} or standard_unit != source_unit:
            fail(("来源单位", "标准单位"), "计数指标单位必须是户或笔且保持一致")
    elif metric_type == "number" and source_unit != standard_unit:
        fail(("来源单位", "标准单位"), "普通数值指标来源单位与标准单位必须相同")
    if metric_type == "rate" and storage_scale not in {"native", "fraction", "percent"}:
        fail("存储尺度", f"比例存储尺度不支持：{storage_scale}")


def _canonical_data_type(value: Any) -> str:
    text = _text(value).lower()
    aliases = {"": "number", "数值": "number", "数值型": "number", "decimal": "number",
               "整数": "integer", "文本": "text", "字符串": "text", "string": "text"}
    return aliases.get(text, text)


def _canonical_metric_type(value: Any) -> str:
    text = _text(value).lower()
    aliases = {"": "amount", "金额": "amount", "余额": "amount", "发生额": "amount",
               "利率": "rate", "比例": "rate", "户数": "count", "笔数": "count",
               "计数": "count", "普通数值": "number", "文本": "text"}
    return aliases.get(text, text)


def _canonical_storage_scale(value: Any) -> str:
    text = _text(value).lower()
    aliases = {"": "native", "原值": "native", "比例存储": "fraction", "小数比例": "fraction",
               "百分比存储": "percent", "百分数存储": "percent"}
    return aliases.get(text, text)


def _canonical_unit(value: str, metric_type: str) -> str:
    if metric_type == "rate" and value in {"%", "百分点"}:
        return "百分比"
    return value


def _validate_data_type(data_type: str, metric_type: str, row_num: int) -> None:
    if data_type not in {"number", "numeric", "decimal", "integer", "text", "string"}:
        raise ConfigurationError(f"指标目录第 {row_num} 行数据类型不支持：{data_type}")
    if metric_type == "text" and data_type not in {"text", "string"}:
        raise ConfigurationError(f"文本指标目录第 {row_num} 行的数据类型必须是文本")
    if metric_type != "text" and data_type in {"text", "string"}:
        raise ConfigurationError(f"数值指标目录第 {row_num} 行的数据类型不能是文本")


def _validate_precision(precision: int, row_num: int) -> None:
    if not 0 <= precision <= 12:
        raise ConfigurationError(f"指标目录第 {row_num} 行精度必须介于 0 到 12")


def _validate_date_range(start: str, end: str, field: str) -> None:
    def parse(value: str) -> date:
        patterns = ((r"^\d{8}$", "%Y%m%d"), (r"^\d{4}-\d{2}-\d{2}$", "%Y-%m-%d"),
                    (r"^\d{6}$", "%Y%m"), (r"^\d{4}-\d{2}$", "%Y-%m"))
        for pattern, fmt in patterns:
            if re.fullmatch(pattern, value):
                try:
                    return datetime.strptime(value, fmt).date()
                except ValueError as exc:
                    raise ConfigurationError(f"{field}日期无效：{value}") from exc
        raise ConfigurationError(f"{field}日期格式必须为 YYYYMMDD、YYYY-MM-DD、YYYYMM 或 YYYY-MM")

    start_date = parse(start) if start else None
    end_date = parse(end) if end else None
    if start_date and end_date and start_date > end_date:
        raise ConfigurationError(f"{field}的开始日期晚于结束日期")


def _validate_catalog_uniqueness(metrics: list[MetricSpec], external: list[ExternalMetricSpec],
                                 contracts: list[SourceContract]) -> None:
    keys = [(m.business, m.sheet, m.row_path, m.column_path) for m in metrics]
    if len(keys) != len(set(keys)):
        raise ConfigurationError("模板指标目录存在重复业务/Sheet/行路径/列路径")
    ext_keys = [(m.source_id, m.metric_header) for m in external]
    if len(ext_keys) != len(set(ext_keys)):
        raise ConfigurationError("外部指标目录存在重复数据源/指标表头")
    contract_keys = [c.source_id for c in contracts]
    if len(contract_keys) != len(set(contract_keys)):
        raise ConfigurationError("数据源契约存在重复数据源ID")
