"""一次性解析 Excel 审核模板，向 DAG 节点提供不可变的轻量快照。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .models import CopyRange, TemplateDefinition
from .name_config import (
    FORMULA_COPY_FUNCTION,
    ISSUE_EXTRACT_FUNCTION,
    STRUCTURE_COMPARE_FUNCTION,
    FeatureMapping,
)
from .native.openpyxl_workbook import (
    named_ranges,
    read_template,
    template_formulas,
    template_structure_values,
)
from .template import TemplateError


@dataclass(frozen=True)
class TemplateSnapshot:
    """关闭工作簿后仍可安全跨节点传递的模板数据。"""

    template_path: Path
    definition: TemplateDefinition
    formulas: tuple[object, ...]
    structure_values: tuple[tuple[CopyRange, tuple[tuple[Any, ...], ...]], ...]
    feature_ranges: Mapping[str, tuple[CopyRange, ...]]
    missing_feature_ranges: Mapping[str, tuple[str, ...]]
    formula_overrides: Mapping[tuple[str, str], Any]
    formula_cells: Mapping[tuple[str, str], Any]

    def require_feature_ranges(self, feature_name: str) -> tuple[CopyRange, ...]:
        missing = self.missing_feature_ranges.get(feature_name, ())
        if missing:
            raise TemplateError(
                "模块“{}”缺少命名区域：{}".format(feature_name, "、".join(missing))
            )
        return self.feature_ranges.get(feature_name, ())


def _formula_overrides(workbook: Any, definition: TemplateDefinition) -> dict[tuple[str, str], Any]:
    result: dict[tuple[str, str], Any] = {}
    for item in definition.copy_ranges:
        identity = (item.sheet_name, item.address.upper())
        if identity in result:
            continue
        min_col, min_row, max_col, max_row = range_boundaries(item.address)
        source_sheet = workbook[item.sheet_name]
        matrix = tuple(
            tuple(source_sheet.cell(row, column).value for column in range(min_col, max_col + 1))
            for row in range(min_row, max_row + 1)
        )
        result[identity] = matrix[0][0] if len(matrix) == 1 and len(matrix[0]) == 1 else matrix
    return result


def _feature_range_index(
    workbook: Any, mappings: tuple[FeatureMapping, ...],
) -> tuple[dict[str, tuple[CopyRange, ...]], dict[str, tuple[str, ...]]]:
    ranges_by_feature: dict[str, tuple[CopyRange, ...]] = {}
    missing_by_feature: dict[str, tuple[str, ...]] = {}
    for mapping in mappings:
        resolved: list[CopyRange] = []
        missing: list[str] = []
        for range_name in mapping.range_names:
            one_name = FeatureMapping(
                mapping.name, mapping.feature_type, (range_name,), mapping.workbook_limited,
                mapping.workbook_keyword, mapping.remark, mapping.output, mapping.input_note,
            )
            areas = named_ranges(workbook, (one_name,))
            if areas:
                resolved.extend(areas)
            else:
                missing.append(range_name)
        ranges_by_feature[mapping.name] = tuple(dict.fromkeys(resolved))
        if missing:
            missing_by_feature[mapping.name] = tuple(missing)
    return ranges_by_feature, missing_by_feature


def load_template_snapshot(
    template_path: Path,
    mappings: Iterable[FeatureMapping],
) -> TemplateSnapshot:
    """打开模板一次并提取后续审核节点所需的全部静态数据。"""
    template_path = Path(template_path)
    mappings = tuple(mappings)
    workbook = load_workbook(template_path, read_only=False, data_only=False, keep_links=False)
    try:
        definition = read_template(
            template_path,
            formula_mappings=(m for m in mappings if m.feature_type == FORMULA_COPY_FUNCTION),
            structure_mappings=(m for m in mappings if m.feature_type == STRUCTURE_COMPARE_FUNCTION),
            extraction_mappings=(m for m in mappings if m.feature_type == ISSUE_EXTRACT_FUNCTION),
            workbook=workbook,
        )
        ranges_by_feature, missing_by_feature = _feature_range_index(workbook, mappings)
        structure = template_structure_values(template_path, definition, workbook=workbook)
        return TemplateSnapshot(
            template_path=template_path,
            definition=definition,
            formulas=tuple(template_formulas(template_path, workbook=workbook)),
            structure_values=tuple(
                (item, tuple(tuple(row) for row in values)) for item, values in structure
            ),
            feature_ranges=ranges_by_feature,
            missing_feature_ranges=missing_by_feature,
            formula_overrides=_formula_overrides(workbook, definition),
            formula_cells={
                (rule.sheet_name, rule.formula_cell): workbook[rule.sheet_name][rule.formula_cell].value
                for rule in definition.rules
                if rule.sheet_name in workbook.sheetnames
            },
        )
    finally:
        workbook.close()
