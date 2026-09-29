"""条件格式规则触发检测引擎（统一接口）。

设计要点（区别于旧「颜色变化 = 触发」的判定口径）：

- **规则触发是主判定**：读取报送文件自身的 OOXML 条件格式规则
  （conditionalFormatting / cfRule / formula / dxf），对目标单元格求值，
  输出 triggered 结果；颜色（dxf fill）只是随规则的辅助信息，颜色解析
  失败不影响 triggered。
- **两个后端、一个接口**：
  - :class:`OoxmlConditionalFormatEngine`：纯 openpyxl 规则求值，UOS 固定
    使用，也是 Windows 上可供对比验证的快速通道；
  - :class:`NativeConditionalFormatEngine`：Windows Excel/WPS COM 的真实
    渲染结果（DisplayFormat），用于验证 OOXML 结果、覆盖 OOXML 解析不了
    的复杂规则。不包含 LibreOffice UNO——UNO 的 ``CellBackColor`` 只反映
    基础填充色，读不到规则触发后的显示状态（已在真实文件上验证）。
- **四态求值**：TRUE / FALSE / UNSUPPORTED / ERROR，统一出自
  ``conditional_format`` 模块；UNSUPPORTED 不等于 FALSE，会进入
  extraction.unsupported 供差异报告与运行日志使用。
- **priority / stopIfTrue**：按 Excel 语义模拟——规则按 priority 升序
  （数值小者优先），同优先级按文件内出现顺序；命中的最高优先级规则决定
  该格的判定，其 stopIfTrue 生效时停止继续求值更低优先级规则。若更高
  优先级处存在 UNSUPPORTED 规则，结果标记 confidence=PARTIAL。
- **统一输出**：:class:`ConditionalFormatResult`；业务输出（AuditResult/
  Issue）由 :func:`results_to_issues` 统一构造，字段口径与既有实现一致。

DAG 节点（条件格式提取）只面向 :class:`ConditionalFormatEngine` 与工厂
:func:`create_conditional_format_engine`，不感知具体实现。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from .conditional_format import (
    RuleEvaluation,
    UnsupportedFormulaError,
    evaluate_cellis_rule,
    evaluate_expression_condition,
    translate_formula_refs,
)
from .models import CopyRange, Issue

# 条件格式结果来源（ConditionalFormatResult.source）。
SOURCE_OOXML = "OOXML"
SOURCE_EXCEL_NATIVE = "EXCEL_NATIVE"
SOURCE_WPS_NATIVE = "WPS_NATIVE"

# 置信度：FULL=判定无歧义；PARTIAL=更高优先级处存在 UNSUPPORTED 规则，
# 真实办公软件可能给出不同结果。
CONFIDENCE_FULL = "FULL"
CONFIDENCE_PARTIAL = "PARTIAL"

# 白、透明与自动/无填充永远不代表可见结果。
NO_FILL_COLOURS = {0, -1, 16777215}

_CELLREF_RE = re.compile(r"^([A-Za-z]{1,3})(\d+)$")

_OPERATOR_LABELS = {
    "equal": "=", "notequal": "<>", "greaterthan": ">",
    "greaterthanorequal": ">=", "lessthan": "<", "lessthanorequal": "<=",
}

# COM FormatCondition.Type → 规则类型（xlCellValue=1 / xlExpression=2）。
_COM_RULE_TYPES = {1: "cellIs", 2: "expression"}


# ---------------------------------------------------------------------------
# 统一输出模型
# ---------------------------------------------------------------------------


@dataclass
class ConditionalFormatResult:
    """单个单元格的条件格式判定结果（两个后端共用同一结构）。"""

    sheet: str
    cell: str                      # 如 "C26"
    range: str = ""                # 命中规则的 AppliesTo（sqref 片段）
    rule_type: str = ""            # cellIs / expression / …
    formula: str = ""              # 规则公式原文（cellIs 为操作数）
    triggered: bool = False
    source: str = SOURCE_OOXML     # OOXML / EXCEL_NATIVE / WPS_NATIVE
    color: int | None = None       # 规则填充色（辅助信息；None 不影响判定）
    message: str = ""              # 规则描述（条件格式规则：C26>0 / 批注）
    confidence: str = CONFIDENCE_FULL
    unsupported_reason: str = ""   # confidence=PARTIAL 时的原因


@dataclass
class ConditionalFormatUnsupported:
    """一条无法由 OOXML 求值器解析的规则（不是「未触发」）。"""

    sheet: str
    rule_type: str
    detail: str                    # 公式或规则描述
    reason: str
    priority: int | None = None


@dataclass
class ConditionalFormatExtraction:
    """一次条件格式提取的完整产物。"""

    results: list[ConditionalFormatResult] = field(default_factory=list)
    unsupported: list[ConditionalFormatUnsupported] = field(default_factory=list)
    unsupported_cells: dict[tuple[str, int, int], str] = field(default_factory=dict)
    evaluated_cells: int = 0
    rules_total: int = 0
    rules_supported: int = 0
    # 实际生效的扫描区域（无名区域回退后与传入 ranges 不同）。
    scanned_ranges: list[CopyRange] = field(default_factory=list)


# ---------------------------------------------------------------------------
# OOXML 规则求值后端
# ---------------------------------------------------------------------------


def _color_to_hex(color) -> str | None:
    """Return an 8-digit AARRGGBB hex string for an openpyxl Color, if real."""
    if color is None:
        return None
    try:
        color_type = (color.type or "rgb").casefold()
    except Exception:
        color_type = "rgb"
    if color_type == "rgb":
        try:
            raw = str(color.rgb)
        except Exception:
            return None
        # openpyxl keeps file 8-char "AARRGGBB" verbatim but pads 6-char input
        # to "00RRGGBB".  A conditional-format background is always opaque, so
        # normalise to an FF alpha over the trailing RGB.
        if len(raw) >= 6 and re.fullmatch(r"[0-9A-Fa-f]{6,8}", raw):
            return "FF" + raw[-6:].upper()
        return None
    if color_type == "indexed":
        idx = color.indexed
        if idx is None:
            return None
        from openpyxl.styles.colors import COLOR_INDEX
        try:
            named = COLOR_INDEX[idx]
        except (IndexError, TypeError):
            return None
        if named in (None, "window", "windowText"):
            return None
        # COLOR_INDEX stores 8-char "00RRGGBB" strings (alpha byte set to 00).
        if isinstance(named, str) and re.fullmatch(r"[0-9A-Fa-f]{8}", named):
            return "FF" + named[2:].upper()
        if isinstance(named, str) and named.startswith("#"):
            return "FF" + named[1:].upper()
        return None
    if color_type == "theme":
        # Standard Office theme colours for the first 12 palette slots.
        theme_hex = {
            0: "FFFFFF", 1: "000000", 2: "EEECE1", 3: "1F497D",
            4: "4F81BD", 5: "C0504D", 6: "9BBB59", 7: "8064A2",
            8: "4BACC6", 9: "F79646", 10: "0000FF", 11: "800080",
        }.get(color.theme)
        if theme_hex is None:
            return None
        tint = color.tint or 0.0
        if tint:
            # Apply Excel's luminance tint formula.
            def _channel(value):
                return round((value if tint > 0 else value * (1 + tint)) * 255)
            base = [int(theme_hex[i:i + 2], 16) for i in (0, 2, 4)]
            if tint > 0:
                rgb = [round(255 - (255 - ch) * (1 - tint)) for ch in base]
            else:
                rgb = [round(ch * (1 + tint)) for ch in base]
            return "FF" + "".join("{:02X}".format(ch) for ch in rgb)
        return "FF" + theme_hex
    return None


def _dxf_fill_rgb(dxfs, dxf_id: int | None) -> str | None:
    """Resolve an Excel conditional-format dxfId to its visible fill RGB.

    Excel writes the condition background under ``dxf/fill/patternFill``:
    ``bgColor`` paints the whole cell (the usual conditional-format idiom),
    ``fgColor`` is the base.  openpyxl keeps them on the differential style.
    """
    if dxf_id is None:
        return None
    try:
        dxf = dxfs[dxf_id]
    except (IndexError, TypeError):
        return None
    if dxf is None:
        return None
    fill = getattr(dxf, "fill", None)
    if fill is None:
        return None
    fg_hex = _color_to_hex(getattr(fill, "fgColor", None))
    bg_hex = _color_to_hex(getattr(fill, "bgColor", None))
    for candidate in (bg_hex, fg_hex):
        if candidate and candidate.upper() not in {"00000000", "FFFFFFFF"}:
            return candidate
    return None


def rule_label(rule, *, address: str, dr: int = 0, dc: int = 0) -> str:
    """规则的可读描述：条件格式规则：C26>0 / 条件格式规则：AND(C5>0,C5>1)。

    保留原大小写、去 $ 锚点。``dr``/``dc`` 非零时（OOXML 引擎逐格求值），
    公式里的相对引用按当前格平移后再渲染——即 Office 对单格的公式渲染口
    径：「规则改到对应的单元格」。如 J36 命中锚点为 J8 的多区域规则时，
    描述为 ``AND(J36<>"",OR(J36<-30,J36>30))`` 而非锚点原式。

    注意：COM 原生路径（excel_com）当前仍输出锚点原式——两条路径对同一
    触发格的描述文本会因此不同（OOXML=平移式 / COM=锚点式），差异报告
    如出现该类 MESSAGE 差异以此为准；COM 侧同步平移口径待后续轮次。
    """
    rule_type = (rule.type or "").casefold()
    raw = str((rule.formula or [""])[0] or "").lstrip("=")
    if dr or dc:
        raw = translate_formula_refs(raw, dr, dc)
    if rule_type == "cellis":
        op_label = _OPERATOR_LABELS.get((rule.operator or "").casefold(), (rule.operator or "?"))
        operand = raw.replace("$", "")
        return "条件格式规则：{}{}{}".format(address, op_label, operand)
    if rule_type == "expression":
        return "条件格式规则：{}".format(raw.replace("$", ""))
    return "条件格式规则：{}".format(rule_type or "未知")


def cellis_operand(book_ws, formulas, index: int):
    """解析 cellIs 操作数：数值/文本/布尔字面量或单格引用。

    表达式形态的操作数无法离线求值——抛 UnsupportedFormulaError，由统一
    求值器记为 UNSUPPORTED，绝不返回 None 伪装成「未触发」。
    """
    if not formulas or index >= len(formulas):
        raise UnsupportedFormulaError("cellIs 缺少第 {} 个操作数".format(index + 1))
    raw = str(formulas[index]).strip()
    if raw.startswith("="):
        raw = raw[1:]
    lowered = raw.casefold()
    if lowered == "true":
        return True
    if lowered == "false":
        return False
    try:
        return float(raw) if ("." in raw or "e" in lowered) else int(raw)
    except ValueError:
        pass
    match = _CELLREF_RE.fullmatch(raw.replace("$", ""))
    if match:
        from openpyxl.utils.cell import column_index_from_string
        try:
            col = column_index_from_string(match.group(1))
            row = int(match.group(2))
        except Exception:
            raise UnsupportedFormulaError("cellIs 操作数无法解析：{}".format(raw)) from None
        return book_ws.cell(row=row, column=col).value
    raise UnsupportedFormulaError("cellIs 操作数不支持离线求值：{}".format(raw))


def evaluate_ooxml_rule(rule, value, ws, *, anchor_row: int, anchor_col: int,
                        cur_row: int, cur_col: int) -> RuleEvaluation:
    """单条 OOXML 规则在单个单元格上的四态求值（唯一入口）。"""
    rule_type = (rule.type or "").casefold()
    if rule_type == "cellis":
        return evaluate_cellis_rule(
            rule.operator, list(rule.formula or []), value,
            resolve_operand=lambda formulas, index: cellis_operand(ws, formulas, index),
        )
    if rule_type == "expression":
        formula = (rule.formula or [""])[0]
        if not str(formula or "").strip():
            return RuleEvaluation("ERROR", reason="expression 规则公式为空")
        return evaluate_expression_condition(
            str(formula), ws, anchor_row, anchor_col, cur_row, cur_col
        )
    # colorScale / dataBar / iconSet 等需要真实渲染器，明确不支持。
    return RuleEvaluation(
        "UNSUPPORTED", reason="暂不支持的条件格式类型：{}".format(rule.type or "未知"),
    )


def _covering_rules(ws):
    """按 (priority 升序, 文件内顺序) 返回工作表的条件格式规则清单。"""
    try:
        entries = list(ws.conditional_formatting)
    except AttributeError:
        return []
    specs = []
    for order, (cf) in enumerate(entries):
        for rule in cf.rules:
            specs.append((rule, cf.sqref, order))
    return sorted(specs, key=lambda item: (priority_of(item[0]), item[2]))


def priority_of(rule) -> int:
    """Excel 优先级：数值小者优先；缺失/非法按 0 处理。"""
    try:
        return int(rule.priority) if rule.priority is not None else 0
    except (TypeError, ValueError):
        return 0


def _sqref_bounds(sqref) -> list[tuple[int, int, int, int]]:
    """Expand a conditional-format sqref into [(min_col,min_row,max_col,max_row)]."""
    from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
    if isinstance(sqref, MultiCellRange):
        raw = sqref.ranges
    else:
        raw = [sqref]
    result = []
    for item in raw:
        try:
            if isinstance(item, CellRange):
                min_col, min_row, max_col, max_row = item.bounds
            else:
                text = str(item).strip()
                if not text:
                    continue
                parsed = CellRange(text)
                min_col, min_row, max_col, max_row = parsed.bounds
        except Exception:
            continue
        result.append((min_col, min_row, max_col, max_row))
    return result


class OoxmlConditionalFormatEngine:
    """OOXML 规则触发检测后端（UOS 主路径；Windows 上用于对比验证）。"""

    source = SOURCE_OOXML
    name = "OOXML规则模式"

    def extract(self, workbook_path: Path, ranges: list[CopyRange],
                structure_ranges: list[CopyRange]) -> ConditionalFormatExtraction:
        """对报送文件执行规则求值，返回触发结果与不支持清单。

        ``ranges`` 为空时与 Windows COM 主路一致：回退为扫描每张工作表
        条件格式自身的 AppliesTo 范围。
        """
        workbook_path = Path(workbook_path)
        extraction = ConditionalFormatExtraction()
        book = load_workbook(workbook_path, data_only=True, keep_links=False)
        try:
            dxfs = getattr(book, "_differential_styles", None)
            live_ranges = [area for area in ranges if area.sheet_name in book.sheetnames]
            if not live_ranges:
                live_ranges = self._applies_to_ranges(book)
            extraction.scanned_ranges = live_ranges
            for area in live_ranges:
                ws = book[area.sheet_name]
                rule_specs = self._sheet_specs(ws, dxfs, extraction)
                extraction.results.extend(self._scan_area(ws, area.address, rule_specs, extraction))
            extraction.evaluated_cells = sum(
                self._area_size(area.address) for area in live_ranges)
            return extraction
        finally:
            book.close()

    def extract_issues(self, workbook_path: Path, ranges: list[CopyRange],
                       structure_ranges: list[CopyRange], *, period: str, batch_id: str,
                       audit_time: str,
                       source_file: Path) -> tuple[list[Issue], ConditionalFormatExtraction]:
        """规则求值 + 统一构造业务 Issue（字段口径与既有实现一致）。"""
        workbook_path = Path(workbook_path)
        extraction = self.extract(workbook_path, ranges, structure_ranges)
        static = load_workbook(workbook_path, data_only=True, keep_links=False)
        try:
            issues = results_to_issues(
                extraction.results,
                workbook_path=workbook_path, ranges=extraction.scanned_ranges or ranges,
                structure_ranges=structure_ranges, period=period,
                batch_id=batch_id, audit_time=audit_time,
                source_file=source_file, static_book=static,
            )
            return issues, extraction
        finally:
            static.close()

    # ---- 内部 ----

    def _applies_to_ranges(self, book) -> list[CopyRange]:
        """无名区域回退：把每张表条件格式的 AppliesTo 展开为扫描区域。"""
        ranges: list[CopyRange] = []
        seen: set[tuple[str, str]] = set()
        for ws in book.worksheets:
            for _rule, sqref, _order in _covering_rules(ws):
                for lo_col, lo_row, hi_col, hi_row in _sqref_bounds(sqref):
                    address = "{}{}:{}{}".format(
                        get_column_letter(lo_col), lo_row,
                        get_column_letter(hi_col), hi_row)
                    identity = (ws.title, address.upper())
                    if identity not in seen:
                        seen.add(identity)
                        ranges.append(CopyRange(ws.title, address))
        return ranges

    def _area_size(self, address: str) -> int:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(address)
        except ValueError:
            return 0
        return (max_col - min_col + 1) * (max_row - min_row + 1)

    def _sheet_specs(self, ws, dxfs, extraction):
        """预解析一张表的规则：颜色、锚点、优先级、stopIfTrue、支持性。"""
        specs = []
        for rule, sqref, order in _covering_rules(ws):
            extraction.rules_total += 1
            bounds = _sqref_bounds(sqref)
            if not bounds:
                continue
            rule_type = (rule.type or "").casefold()
            if rule_type not in ("cellis", "expression"):
                unsupported = ConditionalFormatUnsupported(
                    sheet=ws.title, rule_type=str(rule.type or "未知"),
                    detail=rule_label(rule, address=""),
                    reason="暂不支持的条件格式类型：{}".format(rule.type or "未知"),
                    priority=rule.priority,
                )
                extraction.unsupported.append(unsupported)
                # 不支持的规则也要进入覆盖清单：它可能本应命中并 stopIfTrue，
                # 离线无法判定，因此命中低优先级规则时须降级为 PARTIAL。
                anchor_col = min(b[0] for b in bounds)
                anchor_row = min(b[1] for b in bounds)
                specs.append({
                    "rule": rule, "bounds": bounds, "anchor_row": anchor_row,
                    "anchor_col": anchor_col, "colour": None,
                    "priority": priority_of(rule), "order": order,
                    "stop_if": bool(rule.stopIfTrue), "sqref": str(sqref),
                    "unsupported": unsupported,
                })
                continue
            extraction.rules_supported += 1
            rgb = _dxf_fill_rgb(dxfs, rule.dxfId)
            colour = None
            if rgb:
                try:
                    value = int(rgb[-6:], 16)
                    colour = None if value in NO_FILL_COLOURS else value
                except ValueError:
                    colour = None
            anchor_col = min(b[0] for b in bounds)
            anchor_row = min(b[1] for b in bounds)
            specs.append({
                "rule": rule, "bounds": bounds, "anchor_row": anchor_row,
                "anchor_col": anchor_col, "colour": colour,
                "priority": priority_of(rule), "order": order,
                "stop_if": bool(rule.stopIfTrue),
                "sqref": str(sqref), "unsupported": None,
            })
        return specs

    def _scan_area(self, ws, area_address: str, rule_specs: list[dict],
                   extraction: ConditionalFormatExtraction) -> list[ConditionalFormatResult]:
        """区域内逐格按 Excel 优先级顺序求值；命中的最高优先级规则定案。"""
        region_min_col, region_min_row, region_max_col, region_max_row = (
            range_boundaries(area_address))
        results: list[ConditionalFormatResult] = []
        for row in range(region_min_row, region_max_row + 1):
            for col in range(region_min_col, region_max_col + 1):
                value = ws.cell(row=row, column=col).value
                partial_reason = ""
                for spec in rule_specs:
                    inside = any(
                        lo_col <= col <= hi_col and lo_row <= row <= hi_row
                        for lo_col, lo_row, hi_col, hi_row in spec["bounds"]
                    )
                    if not inside:
                        continue
                    unsupported_entry = spec.get("unsupported")
                    if unsupported_entry is not None:
                        # 无法离线求值的规则（colorScale 等）：不是「未触发」。
                        # 记录原因并降级置信度；其 stopIfTrue 效果未知，继续
                        # 求值更低优先级规则，由 confidence 提示人工核实。
                        if not partial_reason:
                            partial_reason = unsupported_entry.reason
                        extraction.unsupported_cells[(ws.title, row, col)] = unsupported_entry.reason
                        continue
                    rule = spec["rule"]
                    evaluation = evaluate_ooxml_rule(
                        rule, value, ws,
                        anchor_row=spec["anchor_row"], anchor_col=spec["anchor_col"],
                        cur_row=row, cur_col=col,
                    )
                    if evaluation.status == "TRUE":
                        address = "{}{}".format(get_column_letter(col), row)
                        comment = ws.cell(row=row, column=col).comment
                        result = ConditionalFormatResult(
                            sheet=ws.title, cell=address,
                            range=spec["sqref"],
                            rule_type=(rule.type or "").casefold(),
                            formula=str((rule.formula or [""])[0] or ""),
                            triggered=True, source=self.source,
                            color=spec["colour"],
                            message=(comment.text.strip() if comment and comment.text and comment.text.strip()
                                     else rule_label(
                                         rule, address=address,
                                         dr=row - spec["anchor_row"],
                                         dc=col - spec["anchor_col"])),
                            confidence=CONFIDENCE_PARTIAL if partial_reason else CONFIDENCE_FULL,
                            unsupported_reason=partial_reason,
                        )
                        results.append(result)
                        break   # 最高优先级命中规则定案
                    if evaluation.status == "UNSUPPORTED" and not partial_reason:
                        partial_reason = evaluation.reason
                        extraction.unsupported_cells[(ws.title, row, col)] = evaluation.reason
                    elif evaluation.status == "ERROR":
                        extraction.unsupported.append(ConditionalFormatUnsupported(
                            sheet=ws.title, rule_type=(rule.type or "").casefold(),
                            detail=rule_label(
                                rule, address="{}{}".format(get_column_letter(col), row),
                                dr=row - spec["anchor_row"],
                                dc=col - spec["anchor_col"]),
                            reason="求值出错：{}".format(evaluation.reason),
                            priority=spec["priority"],
                        ))
                        continue
                # 全部规则求值完仍无命中：未触发（partial 只影响置信度，不产生结果）。
        return results


# ---------------------------------------------------------------------------
# 统一业务输出：ConditionalFormatResult → Issue（AuditResult）
# ---------------------------------------------------------------------------


def results_to_issues(results: list[ConditionalFormatResult], *,
                      workbook_path: Path, ranges: list[CopyRange],
                      structure_ranges: list[CopyRange], period: str, batch_id: str,
                      audit_time: str, source_file: Path, static_book,
                      resolver=None) -> list[Issue]:
    """把触发结果转换为业务 Issue（字段与既有条件格式提取完全一致）。

    统一经 :mod:`conditional_context` 的 Context/MessageResolver/Builder：
    指标定位走 :class:`IndicatorResolver`（COM ``_conditional_check_field``
    语义的静态复现），message 走「批注优先、公式兜底」。
    """
    from .conditional_context import ConditionalFormatContext, build_conditional_issues

    context = ConditionalFormatContext.build(
        workbook_path=workbook_path, static_book=static_book,
        structure_ranges=structure_ranges)
    issues, _stats = build_conditional_issues(
        results, context=context, ranges=ranges,
        workbook_path=workbook_path, source_file=source_file,
        period=period, batch_id=batch_id, audit_time=audit_time)
    return issues


def _cell_row(cell: str) -> int:
    return int("".join(ch for ch in cell if ch.isdigit()))


def _cell_col(cell: str) -> int:
    letters = "".join(ch for ch in cell if ch.isalpha())
    number = 0
    for ch in letters.upper():
        number = number * 26 + (ord(ch) - 64)
    return number


# ---------------------------------------------------------------------------
# Native 后端（Windows Excel/WPS COM 真实渲染；不含 LibreOffice UNO）
# ---------------------------------------------------------------------------


class NativeConditionalFormatEngine:
    """Excel/WPS COM 的真实渲染判定，用于验证 OOXML 结果或覆盖复杂规则。

    只做包装：判定逻辑仍在 ``ExcelSession.extract_conditional_format_issues``
    （DisplayFormat 与基础 Interior.Color 对比）；本类补齐与 OOXML 后端相同的
    :class:`ConditionalFormatResult` 输出。UOS 不提供此后端。
    """

    name = "Native模式"

    def __init__(self, session):
        self.session = session
        engine_name = str(getattr(session, "engine_name", "") or "")
        self.source = SOURCE_WPS_NATIVE if "WPS" in engine_name.upper() else SOURCE_EXCEL_NATIVE

    def extract_issues(self, workbook, *, mapping, structure_ranges, period, batch_id,
                       audit_time, org_code, org_name, source_file, workbook_path):
        """在已打开的 COM 工作簿上提取；返回 (issues, results, unsupported_count)。"""
        issues = self.session.extract_conditional_format_issues(
            workbook, mapping=mapping, structure_ranges=structure_ranges,
            period=period, batch_id=batch_id, audit_time=audit_time,
            org_code=org_code, org_name=org_name, source_file=source_file,
            workbook_path=workbook_path,
        )
        results = list(getattr(self.session, "last_conditional_results", []) or [])
        for item in results:
            item.source = self.source
        unsupported = int(getattr(self.session, "last_conditional_unsupported_count", 0) or 0)
        return issues, results, unsupported


# ---------------------------------------------------------------------------
# 工厂：DAG 节点唯一入口
# ---------------------------------------------------------------------------

CONDITIONAL_FORMAT_MODE_OOXML = "OOXML"
CONDITIONAL_FORMAT_MODE_NATIVE = "NATIVE"
CONDITIONAL_FORMAT_ENGINE_MODES = (CONDITIONAL_FORMAT_MODE_OOXML, CONDITIONAL_FORMAT_MODE_NATIVE)
DEFAULT_CONDITIONAL_FORMAT_ENGINE_MODE = CONDITIONAL_FORMAT_MODE_NATIVE


def create_conditional_format_engine(mode: str, *, pipeline: str, session=None):
    """按设置与平台选择条件格式引擎（DAG 节点唯一入口）。

    - ``pipeline="native"``（UOS）：固定返回 OOXML 引擎——UNO 读不到触发色，
      Native 后端在 UOS 上不存在；请求 NATIVE 时记一条原因。
    - ``pipeline="com"``（Windows）：NATIVE（默认，保持现有稳定行为）返回
      Native 引擎（需已启动的 ExcelSession）；OOXML 返回规则求值引擎。
    - 未知模式按默认处理，不抛错（保持流程可运行），但由调用方写日志。
    """
    requested = str(mode or "").strip().upper()
    if requested not in CONDITIONAL_FORMAT_ENGINE_MODES:
        requested = DEFAULT_CONDITIONAL_FORMAT_ENGINE_MODE
    if pipeline != "com":
        return OoxmlConditionalFormatEngine(), requested != CONDITIONAL_FORMAT_MODE_NATIVE
    if requested == CONDITIONAL_FORMAT_MODE_OOXML:
        return OoxmlConditionalFormatEngine(), True
    if session is None:
        raise RuntimeError("Native 条件格式模式需要已启动的 Excel/WPS 表格引擎")
    return NativeConditionalFormatEngine(session), requested == DEFAULT_CONDITIONAL_FORMAT_ENGINE_MODE
