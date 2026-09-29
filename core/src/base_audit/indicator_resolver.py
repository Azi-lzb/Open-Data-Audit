"""共享指标定位器（IndicatorResolver）：从静态工作表结构生成 check_field。

把旧 COM 路径 ``ExcelSession._conditional_check_field`` 已在真实数据上验证
的指标定位语义原样抽取为纯 Python 静态读取（不依赖办公套件、不读
DisplayFormat）：

- 表结构区域确定扫描起点（所有区域取 min_row / min_col）；
- 行指标：触发格所在行，从 min_col 向右扫到触发格前一列，收集去重文本；
- 列指标：触发格所在列，从 min_row 向下扫到触发行上一行，收集去重文本；
- 合并单元格：任意格取其合并区域左上角的值（COM ``MergeArea.Cells(1,1)``
  语义），使父级表头对其覆盖的子行/子列可见；
- 数值格不算指标（None、数值、含千分位的数字串都跳过）；
- 拼接：行指标间 ``_`` 相连、列指标间 ``_`` 相连，两段用 ``｜`` 相连。

RuleReader / Evaluator 不负责指标定位：条件格式触发格确定后，由本模块在
Issue 构建步统一补全 check_field。层级与分隔符以 COM 输出为金标准，
不在本模块内另行设计。
"""

from __future__ import annotations

import re
import time
from datetime import date, datetime, time as datetime_time
from typing import Any, Iterable

from .models import CopyRange

_NUMERIC_TEXT_RE = re.compile(r"[-+]?\d+(?:\.\d+)?")


class IndicatorResolver:
    """按 COM ``_conditional_check_field`` 语义解析触发格的校验指标。

    ``static_book`` 为已按 ``data_only=True`` 载入的 openpyxl 工作簿（触发
    判定用的同一份静态簿）；``structure_ranges`` 为模板“表结构区域”命名
    区域解析出的 :class:`CopyRange` 列表。
    """

    def __init__(self, static_book: Any, structure_ranges: Iterable[CopyRange]) -> None:
        self._book = static_book
        self._areas: dict[str, list[tuple[int, int]]] = {}
        for item in structure_ranges or ():
            sheet_name = str(getattr(item, "sheet_name", "") or "")
            address = str(getattr(item, "address", "") or "")
            if not sheet_name or not address:
                continue
            try:
                from openpyxl.utils.cell import range_boundaries

                min_col, min_row, _max_col, _max_row = range_boundaries(address)
            except Exception:
                continue
            self._areas.setdefault(sheet_name, []).append((min_row, min_col))
        self._anchor_cache: dict[str, dict[tuple[int, int], tuple[int, int]]] = {}
        self.resolve_seconds: float = 0.0

    def resolve(self, sheet_name: str, row: int, column: int) -> str:
        """生成 ``行指标_…｜列指标_…``；无表结构区域或缺层级时允许空串。"""
        started = time.perf_counter()
        try:
            areas = self._areas.get(sheet_name)
            if not areas or sheet_name not in self._book.sheetnames:
                return ""
            ws = self._book[sheet_name]
            anchors = self._anchors_for(sheet_name, ws)
            min_row = min(item[0] for item in areas)
            min_col = min(item[1] for item in areas)
            left: list[str] = []
            for col in range(min_col, column):
                value = self._display(ws, anchors, row, col)
                if value and value not in left:
                    left.append(value)
            top: list[str] = []
            for row_index in range(min_row, row):
                value = self._display(ws, anchors, row_index, column)
                if value and value not in top:
                    top.append(value)
            parts: list[str] = []
            if left:
                parts.append("_".join(left))
            if top:
                parts.append("_".join(top))
            return "｜".join(parts)
        finally:
            self.resolve_seconds += time.perf_counter() - started

    # ---- 内部 ----

    def _anchors_for(self, sheet_name: str, ws: Any) -> dict[tuple[int, int], tuple[int, int]]:
        """cell → 合并区域左上角 的映射（每表惰性构建并缓存）。"""
        cached = self._anchor_cache.get(sheet_name)
        if cached is not None:
            return cached
        anchors: dict[tuple[int, int], tuple[int, int]] = {}
        try:
            merged_ranges = list(ws.merged_cells.ranges)
        except Exception:
            merged_ranges = []
        for merged in merged_ranges:
            anchor = (int(merged.min_row), int(merged.min_col))
            for row in range(int(merged.min_row), int(merged.max_row) + 1):
                for col in range(int(merged.min_col), int(merged.max_col) + 1):
                    anchors[(row, col)] = anchor
        self._anchor_cache[sheet_name] = anchors
        return anchors

    @staticmethod
    def _display(ws: Any, anchors: dict[tuple[int, int], tuple[int, int]],
                 row: int, col: int) -> str:
        """读一格的指标文本；合并格取锚点，数值与数字串不算指标。"""
        anchor_row, anchor_col = anchors.get((row, col), (row, col))
        try:
            value = ws.cell(row=anchor_row, column=anchor_col).value
        except Exception:
            return ""
        if value is None or (
            isinstance(value, (int, float, date, datetime, datetime_time))
            and not isinstance(value, bool)
        ):
            # COM Value2 把日期读成数值序列，同样被数值过滤排除。
            return ""
        text = str(value).strip()
        if _NUMERIC_TEXT_RE.fullmatch(text.replace(",", "")):
            return ""
        return text
