"""跨平台审核结果工作簿写入器。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink

from ..history import HISTORY_HEADERS
from .s1_summary_writer import write_audit_summary_payload


def _plain(value):
    return "" if value is None else str(value)


def _issue_row(item):
    return (
        Path(item.source_file).name if item.source_file else "", item.sheet_name,
        item.target_cell, item.severity, item.check_type, item.check_field,
        item.detail or item.message,
        _plain(item.target_value), _plain(item.comparison_value), _plain(item.difference_value),
        item.issue_id, item.institution_feedback, item.auditor_opinion,
    )


def write_audit_summary(
    path: Path,
    issues,
    *,
    sheet_name: str = "本期审核结果",
    mode: str = "AUTO",
    on_engine=None,
) -> Path:
    """Write the S1-F01 result table with clickable links to audited workbooks."""
    path = Path(path)
    safe_sheet_name = "".join(
        "_" if char in r"\[]:*?/" else char
        for char in (sheet_name or "本期审核结果")
    )[:31] or "本期审核结果"
    issues = list(issues)
    rows = [_issue_row(item) for item in issues]
    hyperlinks = []
    for index, item in enumerate(issues):
        if not item.audit_file or not item.sheet_name or not item.target_cell:
            continue
        audit_path = Path(item.audit_file)
        if not audit_path.is_file():
            continue
        hyperlinks.append({
            "row": index + 1,
            "column": 2,
            "target": str(audit_path.resolve()),
            "location": "'{}'!{}".format(str(item.sheet_name).replace("'", "''"), item.target_cell),
            "display": str(item.target_cell),
        })

    payload = {
        "sheet_name": safe_sheet_name,
        "headers": list(HISTORY_HEADERS),
        "rows": rows,
        "hyperlinks": hyperlinks,
    }

    def write_openpyxl() -> Path:
        book = Workbook()
        try:
            sheet = book.active
            sheet.title = safe_sheet_name
            sheet.append(HISTORY_HEADERS)
            header_fill = PatternFill("solid", fgColor="D9EAD3")
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            for item, row_values in zip(issues, rows):
                sheet.append(row_values)
                row = sheet.max_row
                for cell in sheet[row]:
                    cell.number_format = "@"
                if not item.audit_file or not item.sheet_name or not item.target_cell:
                    continue
                audit_path = Path(item.audit_file)
                if not audit_path.is_file():
                    continue
                cell = sheet.cell(row, 3)
                cell.hyperlink = Hyperlink(
                    ref=cell.coordinate, target=str(audit_path.resolve()),
                    location="'{}'!{}".format(
                        str(item.sheet_name).replace("'", "''"), item.target_cell,
                    ),
                    display=str(item.target_cell),
                )
                cell.style = "Hyperlink"
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column, width in zip("ABCDEFGHIJKLM", (38, 24, 14, 16, 20, 28, 42, 18, 18, 18, 38, 34, 24)):
                sheet.column_dimensions[column].width = width
            path.parent.mkdir(parents=True, exist_ok=True)
            book.save(path)
        finally:
            book.close()
        return path

    return write_audit_summary_payload(
        payload, path, mode=mode, openpyxl_writer=write_openpyxl,
        on_engine=on_engine,
    )
