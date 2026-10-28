"""原生汇总模块（UOS/麒麟）：任意行/固定行汇总、汇总表合并、组合工作表。

任意行/固定行汇总对齐 Windows COM 版输出语义：来源文件列写“去日期”的
机构名（``normalize_template_name``），含全部身份字段的表末尾追加本地校验
历史三列（历史触发条数/历史分类说明/审核结果，复合键多行用 ``|`` 拼接）。
"""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
import re
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries

from ..discovery import source_workbooks
from ..history import history_append_rule, join_local_validation_history
from ..models import CopyRange
from ..name_config import (
    FIXED_ROW_SUMMARY_FUNCTION,
    USED_RANGE_SUMMARY_FUNCTION,
    FeatureMapping,
)
from ..template import normalize_template_name
from ..summary_fallback import split_summary_fallback_ranges
from .openpyxl_workbook import named_ranges, named_single_cells_with_prefix
from .s1_summary_writer import write_summary_payload
from .summary_source_reader import open_summary_source


def _bounds(address: str) -> tuple[int, int, int, int]:
    min_col, min_row, max_col, max_row = range_boundaries(address)
    return min_row, min_col, max_row, max_col


def _header_mapping(feature: FeatureMapping) -> FeatureMapping:
    return FeatureMapping(feature.name, feature.feature_type, ("表头区域",), False, "", "")


def _field_key(value: object) -> str:
    return str(value or "").replace(" ", "").replace("\n", "").casefold()


def _sheet_name(base: str, used: set[str]) -> str:
    text = re.sub(r"[\\/*?:\[\]]", "_", str(base))[:31] or "汇总"
    result = text
    number = 2
    while result.casefold() in used:
        result = "{}_{}".format(text[:28], number)
        number += 1
    used.add(result.casefold())
    return result


def _column_widths(headers, rows) -> list[float]:
    columns = max([len(headers), *(len(row) for row in rows)], default=0)
    widths = []
    for column in range(columns):
        values = [headers[column] if column < len(headers) else ""]
        values.extend(row[column] if column < len(row) else "" for row in rows)
        max_length = max((len(str(value or "")) for value in values), default=0)
        widths.append(float(min(40, max(12, max_length + 2))))
    return widths


def _history_context(
    config_path: Path | None, sheet_name: str, base_headers: list[str]
) -> tuple[dict[tuple[str, ...], list[str]], tuple[str, ...], tuple[str, ...]] | None:
    """按内置追加规则解析输出表的历史追加上下文；不追加返回 None。

    历史列未配置时自动取历史表中“非去重列且汇总表头中不存在”的全部列。
    """
    from ..history import history_append_rule, join_local_validation_history

    if config_path is None or not Path(config_path).is_file():
        return None
    key_fields, history_fields = history_append_rule(sheet_name)
    if not key_fields:
        return None
    key_keys = {_field_key(name) for name in key_fields}
    base_keys = {_field_key(name) for name in base_headers}
    if not key_keys or not key_keys <= base_keys:
        return None
    book = load_workbook(config_path, read_only=True, data_only=True, keep_links=False)
    try:
        if sheet_name not in book.sheetnames:
            return None
        rows = book[sheet_name].iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        if not history_fields:
            history_fields = tuple(
                name for name in headers
                if name and _field_key(name).casefold() != "none"
                and _field_key(name) not in key_keys and _field_key(name) not in base_keys
            )
        if not history_fields:
            return None
        lookup = join_local_validation_history(
            headers, [list(row) for row in rows],
            key_fields=key_fields, history_fields=history_fields,
        )
        return lookup, history_fields, key_fields
    finally:
        book.close()


def enrich_history_columns(
    summary_path: Path, config_path: Path, *, on_step: Callable[[str], None] | None = None
) -> Path:
    """把历史说明列追加到汇总工作簿的各张结果表（openpyxl 原地写回）。

    与 COM 版同一口径：对每张结果表，用表头中的去重列匹配历史工作簿的
    ``history_append_rule``（表名驱动），按复合键填入缺失的人工列。
    """
    say = on_step or (lambda _text: None)
    summary_path = Path(summary_path).resolve()
    config_path = Path(config_path).resolve()
    book = load_workbook(summary_path)
    try:
        for sheet in book.worksheets:
            headers = [str(cell.value or "").strip() for cell in sheet[1]]
            if not headers or headers == [""]:
                continue
            context = _history_context(config_path, sheet.title, headers)
            if not context:
                continue
            lookup, appended, keys = context
            missing = [name for name in appended if name not in headers]
            if not missing:
                continue  # 已有历史列，重复执行不叠加
            start = len(headers) + 1
            for offset, name in enumerate(missing):
                sheet.cell(1, start + offset, name)
            from ..history import _norm_key_value

            if any(name not in headers for name in keys):
                continue
            positions = {name: headers.index(name) + 1 for name in keys}
            for row_index in range(2, sheet.max_row + 1):
                identity = tuple(
                    _norm_key_value(sheet.cell(row_index, positions[name]).value)
                    for name in keys
                )
                joined = lookup.get(identity)
                if joined is None:
                    continue
                for offset, value in enumerate(joined):
                    sheet.cell(row_index, start + offset, value)
            say("历史说明富化：{}（新增列 {}）".format(sheet.title, "、".join(missing)))
        book.save(summary_path)
    finally:
        book.close()
    return summary_path


_INVALID = re.compile(r"[\[\]:*?/\\]")


def _copy_sheet(source, target, name, style_cache=None):
    """跨工作簿复制工作表；样式经缓存去重注册（与 Windows 实现同思路）。"""
    from copy import copy as _copy

    target_sheet = target.create_sheet(name)
    target_sheet.sheet_state = source.sheet_state
    target_sheet.freeze_panes = source.freeze_panes

    def cached_style(cell):
        if style_cache is None:
            return _copy(cell._style)
        key = cell.style_id if hasattr(cell, "style_id") else str(cell._style)
        if key not in style_cache:
            style_cache[key] = _copy(cell._style)
        return style_cache[key]

    for row in source.iter_rows():
        for cell in row:
            destination = target_sheet[cell.coordinate]
            destination.value = cell.value
            if cell.has_style:
                destination._style = cached_style(cell)
            if cell.comment:
                destination.comment = _copy(cell.comment)
            if cell.hyperlink:
                destination._hyperlink = _copy(cell.hyperlink)
    for key, dimension in source.row_dimensions.items():
        target_sheet.row_dimensions[key].height = dimension.height
        target_sheet.row_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.column_dimensions.items():
        target_sheet.column_dimensions[key].width = dimension.width
        target_sheet.column_dimensions[key].hidden = dimension.hidden
    for merged in source.merged_cells.ranges:
        target_sheet.merge_cells(str(merged))


def _unique_sheet_name(value: str, used: set[str]) -> str:
    base = _INVALID.sub("_", value)[:31] or "工作表"
    result = base
    number = 2
    while result.casefold() in used:
        tail = "_" + str(number)
        result = base[:31 - len(tail)] + tail
        number += 1
    used.add(result.casefold())
    return result


def _read_header(
    template_book, data_range, headers, *, header_template_sheet: str | None = None
) -> tuple[list[object], int, int]:
    """选与汇总区域同表且列范围覆盖最合适的表头区域，返回表头值与列范围。"""
    header_sheet = header_template_sheet or data_range.sheet_name
    matching = [area for area in headers if area.sheet_name == header_sheet]
    if not matching:
        raise ValueError("工作表“{}”缺少命名区域“表头区域”".format(header_sheet))
    top, left, bottom, right = _bounds(data_range.address)
    best = None
    for item in matching:
        h_top, h_left, h_bottom, h_right = _bounds(item.address)
        if h_bottom < h_top:
            continue
        lo, hi = max(left, h_left), min(right, h_right)
        if lo > hi:
            continue
        if best is None or (hi - lo) > (best[2] - best[1]):
            best = (h_top, lo, hi, h_bottom)
    if best is None:
        raise ValueError("汇总区域与表头区域没有共同列")
    h_top, lo, hi, h_bottom = best
    sheet = template_book[header_sheet]
    # 多行表头自上而下用下划线拼接（与 Windows 版 _combine_header_rows 一致）。
    header_values: list[object] = []
    columns = range(lo, hi + 1)
    merged: dict[int, list[str]] = {}
    for row in sheet.iter_rows(min_row=h_top, max_row=h_bottom, min_col=lo, max_col=hi, values_only=True):
        for offset, value in enumerate(row):
            text = str(value or "").strip()
            if text:
                merged.setdefault(offset, []).append(text)
    for column in columns:
        offset = column - lo
        header_values.append("_".join(merged.get(offset, [])) or "")
    return header_values, lo, hi


def run_region_summaries(
    *, template_path: Path, input_dir: Path, output_dir: Path,
    features: Iterable[FeatureMapping], recursive: bool = True,
    selected_files: Iterable[Path] | None = None,
    on_step: Callable[[str], None] | None = None,
    on_result: Callable[[list[tuple[str, int]], dict[str, int]], None] | None = None,
    output_name: str = "区域汇总",
    history_config_path: Path | None = None,
    source_reader: str = "AUTO",
    result_writer: str = "AUTO",
) -> Path:
    """运行配置的原生行汇总并返回输出工作簿路径。"""
    say = on_step or (lambda _text: None)
    features = [item for item in features if item.feature_type in {USED_RANGE_SUMMARY_FUNCTION, FIXED_ROW_SUMMARY_FUNCTION}]
    if not features:
        raise ValueError("没有启用“汇总_任意行汇总”或“汇总_固定行汇总”模块")
    files = source_workbooks(input_dir, recursive=recursive)
    if history_config_path is not None:
        history_path = Path(history_config_path).resolve()
        files = [path for path in files if path.resolve() != history_path]
    if selected_files:
        allowed = {path.resolve() for path in files}
        files = [Path(path).resolve() for path in selected_files if Path(path).resolve() in allowed]
    if not files:
        raise FileNotFoundError("没有可汇总的源工作簿")
    template = load_workbook(template_path, read_only=True, data_only=False, keep_links=False)
    try:
        source_engines: set[str] = set()

        def record_source_engine(value: str) -> None:
            source_engines.add(value)

        def report_source_fallback(reason: str) -> None:
            say("汇总源数据读取回退：" + reason)

        global_fields = named_single_cells_with_prefix(template, "全局单元格区域.")
        local_fields = named_single_cells_with_prefix(template, "单元格区域.")
        groups: dict[str, dict[str, object]] = {}
        feature_counts: dict[str, int] = {}
        for feature in features:
            feature_counts.setdefault(feature.name, 0)
            say("正在汇总：{}".format(feature.name))
            all_data_ranges = named_ranges(template, (feature,))
            all_headers = named_ranges(template, (_header_mapping(feature),))
            data_ranges, headers, fallback = split_summary_fallback_ranges(
                template.sheetnames, all_data_ranges, all_headers,
            )
            if not data_ranges and fallback is None:
                raise ValueError("汇总功能“{}”未找到命名区域：{}".format(feature.name, "、".join(feature.range_names)))
            if not headers and fallback is None:
                raise ValueError("汇总功能“{}”缺少命名区域“表头区域”".format(feature.name))
            for data_range in data_ranges:
                header_values, left, right = _read_header(template, data_range, headers)
                local = [(name, area) for name, area in local_fields if area.sheet_name == data_range.sheet_name]
                extras = [*global_fields, *local]
                fields = [name for name, _ in extras]
                key = data_range.sheet_name.casefold()
                base_headers = ["来源文件", "来源工作表", *fields, *header_values]
                history_context = _history_context(
                    history_config_path, data_range.sheet_name, base_headers
                )
                out_headers = (
                    [*base_headers, *history_context[1]] if history_context else base_headers
                )
                group = groups.setdefault(key, {
                    "title": data_range.sheet_name, "headers": out_headers, "rows": [],
                })
                top, _left, bottom, _right = _bounds(data_range.address)
                for file_path in files:
                    if file_path.suffix.casefold() == ".xls":
                        continue
                    with open_summary_source(
                        file_path, engine=source_reader,
                        on_engine=record_source_engine,
                        on_fallback=report_source_fallback,
                    ) as book:
                        if data_range.sheet_name not in book.sheetnames:
                            say(
                                "跳过源文件 {} 的工作表 {}：源文件缺少该工作表。".format(
                                    file_path.name, data_range.sheet_name
                                )
                            )
                            continue
                        sheet = book[data_range.sheet_name]
                        last_row = bottom if feature.feature_type == FIXED_ROW_SUMMARY_FUNCTION else max(top, sheet.max_row)
                        values = list(sheet.iter_rows(min_row=top, max_row=last_row, min_col=left, max_col=right, values_only=True))
                        metadata = []
                        for field, area in extras:
                            if area.sheet_name not in book.sheetnames:
                                say(
                                    "汇总提示：源文件 {} 缺少字段 {} 所在工作表 {}，该字段留空。".format(
                                        file_path.name, field, area.sheet_name
                                    )
                                )
                                metadata.append(None)
                                continue
                            row, col = _bounds(area.address)[:2]
                            metadata.append(book[area.sheet_name].cell(row, col).value)
                        for row in values:
                            if any(value not in (None, "") for value in row):
                                base_row = [normalize_template_name(file_path.stem), data_range.sheet_name, *metadata, *row]
                                if history_context:
                                    from ..history import _composite_key
                                    lookup, appended, keys = history_context
                                    identity = _composite_key(base_headers, base_row, keys)
                                    joined = lookup.get(identity)
                                    extra = joined if joined is not None else [""] * len(appended)
                                    group["rows"].append([*base_row, *extra])
                                else:
                                    group["rows"].append(base_row)
                                feature_counts[feature.name] += 1
            if fallback is not None:
                fallback_headers, fallback_left, fallback_right = _read_header(
                    template, fallback.data_range, [fallback.header_range],
                    header_template_sheet=fallback.sheet_name,
                )
                configured_names = {item.sheet_name.casefold() for item in data_ranges}
                fallback_top, _left, fallback_bottom, _right = _bounds(
                    fallback.data_range.address
                )
                fallback_local = [
                    (name, area) for name, area in local_fields
                    if area.sheet_name.casefold() == fallback.sheet_name.casefold()
                ]
                for file_path in files:
                    if file_path.suffix.casefold() == ".xls":
                        continue
                    with open_summary_source(
                        file_path, engine=source_reader,
                        on_engine=record_source_engine,
                        on_fallback=report_source_fallback,
                    ) as book:
                        source_only = [
                            name for name in book.sheetnames
                            if name.casefold() not in configured_names
                        ]
                        if source_only:
                            say(
                                "兜底汇总：{} 的未匹配工作表：{}".format(
                                    file_path.name, "、".join(source_only)
                                )
                            )
                        for source_sheet_name in source_only:
                            source_sheet = book[source_sheet_name]
                            last_row = (
                                fallback_bottom
                                if feature.feature_type == FIXED_ROW_SUMMARY_FUNCTION
                                else max(fallback_top, source_sheet.max_row)
                            )
                            source_rows = list(source_sheet.iter_rows(
                                min_row=fallback_top, max_row=last_row,
                                min_col=fallback_left, max_col=fallback_right,
                                values_only=True,
                            ))
                            nonempty_rows = [
                                row for row in source_rows
                                if any(value not in (None, "") for value in row)
                            ]
                            if not nonempty_rows:
                                say(
                                    "兜底模板：工作表“{}”没有有效数据行，不生成汇总表。".format(
                                        source_sheet_name
                                    )
                                )
                                continue

                            fallback_fields = [
                                (name, CopyRange(source_sheet_name, area.address))
                                for name, area in fallback_local
                            ]
                            extras = [*global_fields, *fallback_fields]
                            metadata = []
                            for field, area in extras:
                                if area.sheet_name not in book.sheetnames:
                                    say(
                                        "汇总提示：源文件 {} 缺少字段 {} 所在工作表 {}，该字段留空。".format(
                                            file_path.name, field, area.sheet_name
                                        )
                                    )
                                    metadata.append(None)
                                    continue
                                row, col = _bounds(area.address)[:2]
                                metadata.append(book[area.sheet_name].cell(row, col).value)

                            base_headers = [
                                "来源文件", "来源工作表",
                                *[name for name, _ in extras], *fallback_headers,
                            ]
                            history_context = _history_context(
                                history_config_path, source_sheet_name, base_headers
                            )
                            out_headers = (
                                [*base_headers, *history_context[1]]
                                if history_context else base_headers
                            )
                            key = source_sheet_name.casefold()
                            group = groups.setdefault(key, {
                                "title": source_sheet_name,
                                "headers": out_headers,
                                "rows": [],
                            })
                            for value_row in nonempty_rows:
                                base_row = [
                                    normalize_template_name(file_path.stem),
                                    source_sheet_name, *metadata, *value_row,
                                ]
                                if history_context:
                                    from ..history import _composite_key
                                    lookup, appended, keys = history_context
                                    identity = _composite_key(base_headers, base_row, keys)
                                    extra = lookup.get(identity)
                                    if extra is None:
                                        extra = [""] * len(appended)
                                    group["rows"].append([*base_row, *extra])
                                else:
                                    group["rows"].append(base_row)
                                feature_counts[feature.name] += 1
            say("完成：{}".format(feature.name))
    finally:
        template.close()
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "{}_{}.xlsx".format(output_name, datetime.now().strftime("%Y%m%d_%H%M%S"))
    output_rows: list[tuple[str, int]] = []
    tables = []
    used_names: set[str] = set()
    for group in groups.values():
        name = _sheet_name(str(group["title"]), used_names)
        headers, rows = group["headers"], group["rows"]
        output_rows.append((name, len(rows)))
        tables.append({
            "name": name, "headers": headers, "rows": rows,
            "widths": _column_widths(headers, rows),
            "freeze_top_rows": 1, "auto_filter": bool(headers),
        })
    if not tables:
        output_rows.append(("汇总", 0))
        tables.append({
            "name": "汇总", "headers": [], "rows": [], "widths": [],
            "freeze_top_rows": 0, "auto_filter": False,
        })

    def write_openpyxl() -> Path:
        output, used_names = Workbook(), set()
        try:
            for index, table in enumerate(tables):
                sheet = output.active if index == 0 else output.create_sheet()
                sheet.title = table["name"]
                headers, rows = table["headers"], table["rows"]
                if headers:
                    sheet.append(headers)
                    for row in rows:
                        sheet.append(row)
                    sheet.freeze_panes = "A2"
                    sheet.auto_filter.ref = sheet.dimensions
                    for column, width in enumerate(table["widths"], start=1):
                        sheet.column_dimensions[
                            sheet.cell(1, column).column_letter
                        ].width = width
            output.save(path)
        finally:
            output.close()
        return path

    path = write_summary_payload(
        {"sheets": tables}, path, mode=result_writer,
        openpyxl_writer=write_openpyxl,
        on_engine=lambda message: say("汇总结果写入：" + message),
    )
    for sheet_name, row_count in output_rows:
        say(f"输出工作表：{sheet_name}，数据 {row_count} 行")
    if source_engines:
        say("汇总源数据读取实际引擎=" + "、".join(sorted(source_engines)))
    if on_result is not None:
        on_result(output_rows, feature_counts)
    return path


def merge_by_first_filename_part(
    *, input_dir: Path, output_dir: Path, recursive: bool = True, selected_files=None,
    output_name: str = "组合联合核查表", on_step=None,
):
    """`修改_组合联合核查表`：按文件名第一段合并同一机构的工作簿。"""
    say = on_step or (lambda _text: None)
    files = source_workbooks(input_dir, recursive=recursive)
    if selected_files:
        allowed = {path.resolve() for path in files}
        files = [Path(p) for p in selected_files if Path(p).resolve() in allowed]
    groups: dict[str, list[Path]] = defaultdict(list)
    for path in files:
        groups[path.stem.split("_", 1)[0] or path.stem].append(path)
    target_dir = output_dir / (str(output_name or "组合联合核查表") + "_" + datetime.now().strftime("%Y%m%d%H%M%S"))
    target_dir.mkdir(parents=True, exist_ok=True)
    records = []
    for group, paths in sorted(groups.items()):
        out = target_dir / (group + "_合并_" + (datetime.now().strftime("%Y%m") or "本期") + ".xlsx")
        book = Workbook()
        del book[book.sheetnames[0]]
        used: set[str] = set()
        copied: list[str] = []
        try:
            for path in paths:
                source = load_workbook(path, data_only=False, keep_links=False)
                try:
                    style_cache: dict[str, object] = {}
                    for sheet in source.worksheets:
                        name = _unique_sheet_name(sheet.title, used)
                        _copy_sheet(sheet, book, name, style_cache)
                        copied.append(path.name + "｜" + sheet.title)
                finally:
                    source.close()
            if not book.sheetnames:
                book.create_sheet("空工作表")
            book.save(out)
            records.append((group, out.name, "；".join(copied), "成功"))
            say("已组合：{}（合并源文件 {} 个）".format(group, len(paths)))
        except Exception as exc:
            records.append((group, "", "", "失败：" + str(exc)))
        finally:
            book.close()
    return target_dir, records


def merge_workbook_tables(
    *, input_dir: Path, output_dir: Path, recursive: bool = True,
    output_name: str = "汇总表合并", selected_files: Iterable[Path] | None = None,
    on_step=None, source_reader: str = "AUTO", result_writer: str = "AUTO",
) -> Path:
    """`汇总_汇总表合并`：仅首行表头完全一致的工作表合并为一张表。"""
    say = on_step or (lambda _text: None)
    files = source_workbooks(input_dir, recursive=recursive)
    if selected_files is not None:
        allowed = {path.resolve() for path in files}
        files = [Path(path).resolve() for path in selected_files if Path(path).resolve() in allowed]
    if not files:
        raise FileNotFoundError("没有可合并的源工作簿")
    groups: dict[tuple[str, ...], dict[str, object]] = {}
    source_engines: set[str] = set()
    for path in files:
        if path.suffix.casefold() == ".xls":
            continue  # 旧版 .xls 由 Windows 版处理；原生管线聚焦 .xlsx
        with open_summary_source(
            path, engine=source_reader, on_engine=source_engines.add,
            on_fallback=lambda reason: say("汇总源数据读取回退：" + reason),
        ) as book:
            for sheet in book.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header = tuple(next(rows, ()))
                if not header or not any(value not in (None, "") for value in header):
                    continue
                signature = tuple(str(value or "").strip() for value in header)
                group = groups.setdefault(signature, {"headers": ["来源文件", "来源工作表", *header], "rows": [], "sheets": []})
                if sheet.title not in group["sheets"]:
                    group["sheets"].append(sheet.title)
                for row in rows:
                    if any(value not in (None, "") for value in row):
                        group["rows"].append([path.name, sheet.title, *row])
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "{}_{}.xlsx".format(output_name, datetime.now().strftime("%Y%m%d_%H%M%S"))
    tables = []
    used: set[str] = set()
    for group in groups.values():
        tables.append({
            "name": _sheet_name("_".join(group["sheets"][:2]) or "汇总", used),
            "headers": group["headers"], "rows": group["rows"], "widths": [],
            "freeze_top_rows": 1, "auto_filter": True,
        })
    if not tables:
        tables.append({
            "name": "汇总", "headers": [], "rows": [], "widths": [],
            "freeze_top_rows": 0, "auto_filter": False,
        })

    def write_openpyxl() -> Path:
        output = Workbook()
        try:
            for index, table in enumerate(tables):
                sheet = output.active if index == 0 else output.create_sheet()
                sheet.title = table["name"]
                if table["headers"]:
                    sheet.append(table["headers"])
                    for row in table["rows"]:
                        sheet.append(row)
                    sheet.freeze_panes = "A2"
                    sheet.auto_filter.ref = sheet.dimensions
            output.save(path)
        finally:
            output.close()
        return path

    path = write_summary_payload(
        {"sheets": tables}, path, mode=result_writer,
        openpyxl_writer=write_openpyxl,
        on_engine=lambda message: say("汇总结果写入：" + message),
    )
    say("汇总表合并完成：按表头分为 {} 张表".format(len(groups)))
    if source_engines:
        say("汇总源数据读取实际引擎=" + "、".join(sorted(source_engines)))
    return path
