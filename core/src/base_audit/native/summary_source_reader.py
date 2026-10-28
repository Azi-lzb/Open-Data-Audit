"""Value-only XLSX readers for S1-F02 summary source workbooks.

The summary template and configured defined names remain on OpenPyXL. This
module only reads source values, where Calamine can avoid OpenPyXL's Python
cell-object construction while retaining the existing row/column positions.
"""

from __future__ import annotations

from contextlib import contextmanager
from datetime import date, datetime
from importlib.metadata import PackageNotFoundError, version
from types import SimpleNamespace
from typing import Callable, Iterator

try:
    from python_calamine import load_workbook as _calamine_load
except Exception:  # pragma: no cover - missing/broken native extension
    _calamine_load = None


READERS = ("AUTO", "CALAMINE", "OPENPYXL")


def _normalize_value(value):
    """Match OpenPyXL's value-only XLSX representation used by S1-F02."""
    if isinstance(value, float) and value.is_integer() and abs(value) <= 2**53:
        return int(value)
    if value == "":
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime(value.year, value.month, value.day)
    return value


def _calamine_version() -> str:
    try:
        return version("python-calamine")
    except PackageNotFoundError:
        return "未知版本"


class _CalamineSheet:
    def __init__(self, book: "_CalamineBook", name: str) -> None:
        self._book = book
        self._sheet = book._book.get_sheet_by_name(name)
        self.title = name
        self._matrix = None
        self._start = (0, 0)

    def _values(self):
        if self._matrix is None:
            try:
                self._matrix = self._sheet.to_python()
                start = self._sheet.start
                self._start = start if isinstance(start, tuple) else (0, 0)
            except Exception as exc:
                if self._book.mode != "AUTO":
                    raise RuntimeError(
                        f"Calamine 读取工作表“{self.title}”失败：{exc}"
                    ) from exc
                self._book._switch_to_openpyxl(
                    f"工作表“{self.title}”读取失败：{type(exc).__name__}: {exc}"
                )
                raise _RetryOpenPyXL from exc
        return self._matrix

    @property
    def max_row(self) -> int:
        try:
            matrix = self._values()
        except _RetryOpenPyXL:
            return self._book._openpyxl()[self.title].max_row
        top = self._start[0]
        total = int(getattr(self._sheet, "total_height", 0) or 0)
        return max(total, top + len(matrix))

    @property
    def max_column(self) -> int:
        try:
            matrix = self._values()
        except _RetryOpenPyXL:
            return self._book._openpyxl()[self.title].max_column
        left = self._start[1]
        total = int(getattr(self._sheet, "total_width", 0) or 0)
        return max(total, left + max((len(row) for row in matrix), default=0))

    def cell(self, row: int, column: int):
        if row < 1 or column < 1:
            raise ValueError("Row and column values must be at least 1")
        try:
            matrix = self._values()
        except _RetryOpenPyXL:
            return self._book._openpyxl()[self.title].cell(row, column)
        local_row, local_col = row - 1 - self._start[0], column - 1 - self._start[1]
        value = None
        if 0 <= local_row < len(matrix) and 0 <= local_col < len(matrix[local_row]):
            value = matrix[local_row][local_col]
        return SimpleNamespace(value=_normalize_value(value), row=row, column=column)

    def iter_rows(
        self,
        min_row: int = 1,
        max_row: int | None = None,
        min_col: int = 1,
        max_col: int | None = None,
        values_only: bool = False,
    ) -> Iterator[tuple]:
        try:
            matrix = self._values()
        except _RetryOpenPyXL:
            yield from self._book._openpyxl()[self.title].iter_rows(
                min_row=min_row, max_row=max_row, min_col=min_col,
                max_col=max_col, values_only=values_only,
            )
            return
        final_row = self.max_row if max_row is None else max_row
        final_col = self.max_column if max_col is None else max_col
        top, left = self._start
        for row_number in range(min_row, final_row + 1):
            local_row = row_number - 1 - top
            source_row = matrix[local_row] if 0 <= local_row < len(matrix) else ()
            values = tuple(
                _normalize_value(
                    source_row[column - 1 - left]
                    if 0 <= column - 1 - left < len(source_row)
                    else None
                )
                for column in range(min_col, final_col + 1)
            )
            if values_only:
                yield values
            else:
                yield tuple(
                    SimpleNamespace(value=value, row=row_number, column=column)
                    for column, value in zip(range(min_col, final_col + 1), values)
                )


class _RetryOpenPyXL(Exception):
    """Internal marker: one Calamine sheet failed before it yielded any values."""


class _CalamineBook:
    def __init__(self, workbook, mode: str, on_engine, on_fallback) -> None:
        self._book = workbook
        self.mode = mode
        self.sheetnames = list(workbook.sheet_names)
        self._sheets = {name: _CalamineSheet(self, name) for name in self.sheetnames}
        self._openpyxl_book = None
        self._on_engine = on_engine
        self._on_fallback = on_fallback

    @property
    def worksheets(self):
        return [self[name] for name in self.sheetnames]

    def __getitem__(self, name: str) -> _CalamineSheet:
        return self._sheets[name]

    def __iter__(self):
        return iter(self.worksheets)

    def _openpyxl(self):
        if self._openpyxl_book is None:
            from openpyxl import load_workbook

            self._openpyxl_book = load_workbook(
                self._path, read_only=True, data_only=True, keep_links=False
            )
        return self._openpyxl_book

    def _switch_to_openpyxl(self, reason: str) -> None:
        if self._on_engine:
            self._on_engine("OpenPyXL")
        if self._on_fallback:
            self._on_fallback(reason)

    def close(self) -> None:
        close = getattr(self._book, "close", None)
        if close:
            close()
        if self._openpyxl_book is not None:
            self._openpyxl_book.close()


@contextmanager
def open_summary_source(
    path,
    *,
    engine: str = "AUTO",
    on_engine: Callable[[str], None] | None = None,
    on_fallback: Callable[[str], None] | None = None,
):
    """Open one source workbook in AUTO, CALAMINE or OPENPYXL mode.

    AUTO prefers Calamine and can fall back to OpenPyXL for an unsupported file
    or sheet. Explicit modes never silently route to another reader.
    """
    mode = str(engine or "AUTO").strip().upper()
    if mode not in READERS:
        raise ValueError("S1-F02 汇总源数据读取方式只能选择 AUTO、CALAMINE 或 OPENPYXL")
    if mode in {"AUTO", "CALAMINE"}:
        reason = ""
        try:
            if _calamine_load is None:
                raise ImportError("python-calamine 不可用")
            workbook = _calamine_load(path)
            wrapped = _CalamineBook(workbook, mode, on_engine, on_fallback)
            wrapped._path = path
        except Exception as exc:
            if mode == "CALAMINE":
                raise RuntimeError(f"Calamine 无法读取源工作簿“{path}”：{exc}") from exc
            reason = f"{type(exc).__name__}: {exc}"
            if on_engine:
                on_engine("OpenPyXL")
            if on_fallback:
                on_fallback(f"源工作簿打开失败：{reason}")
        else:
            if on_engine:
                on_engine(f"Calamine {_calamine_version()}")
            try:
                yield wrapped
            finally:
                wrapped.close()
            return
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if mode == "OPENPYXL" and on_engine:
            on_engine("OpenPyXL")
        yield workbook
    finally:
        workbook.close()
