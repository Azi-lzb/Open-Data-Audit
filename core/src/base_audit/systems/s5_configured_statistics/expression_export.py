"""Create local reference copies containing literal V_CUR / V_EXT expressions."""
from __future__ import annotations

import argparse
import json
import re
import uuid
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName

from s1audit_configured.dsl import parse_expression
from s1audit_configured.errors import ConfigurationError, RetrievalError
from s1audit_configured.external_layout import discover_external_layout, normalized
from s1audit_configured.layout import inferred_numeric_regions, named_regions, numeric_header_regions, numeric_regions, row_path_candidates, scan_rule_paths
from s1audit_configured.model import SourceBinding
from s1audit_configured.numeric_scope import project_numeric_regions
from s1audit_configured.sources import _Workbook, file_sha256
from s1audit_configured.workbook_config import _metric_profile


ROOT = Path(__file__).resolve().parents[5]
TEMPLATE_DIR = ROOT / "config" / "S5_模板文件"
REPORT_HEADERS = ("源文件", "布局模板", "工作表", "数字区域", "模板单元格", "源单元格", "行路径", "列路径或外部指标", "表达式", "状态", "说明")
GOOD_FILL = PatternFill("solid", fgColor="E8F3FC")
ERROR_FILL = PatternFill("solid", fgColor="FFF0CC")
SKIPPED_STATUSES = {"MERGED_COVERED", "HEADER_PLACEHOLDER"}


def _literal(value: str) -> str:
    return json.dumps(value, ensure_ascii=False)


def _path_literal(path: tuple[str, ...]) -> str:
    # A tuple preserves a literal @ inside a label; ordinary paths stay compact.
    if any("@" in part for part in path):
        return "(" + ", ".join(_literal(part) for part in path) + ("," if len(path) == 1 else "") + ")"
    return _literal("@".join(path))


def template_expression(sheet: str, row_path: tuple[str, ...], col_path: tuple[str, ...]) -> str:
    expression = f"V_CUR({_literal(sheet)}, {_path_literal(row_path)}, {_path_literal(col_path)})"
    if len(expression) > 32767:
        raise ValueError("表达式超过 Excel 单元格文本长度限制")
    parse_expression(expression)
    return expression


def _entry(source: Path, schema: Path | None = None, **values) -> dict[str, Any]:
    return {"file": str(source), "template": str(schema or ""), "sheet": "", "region": "", "template_cell": "",
            "cell": "", "row_path": "", "column_path": "", "expression": "", "status": "", "detail": "", **values}


def _put_expression(sheet, row: int, col: int, expression: str) -> None:
    cell = sheet.cell(row, col)
    cell.value = expression  # Text, deliberately without '=': Excel must not evaluate this DSL.
    cell.number_format = "@"
    cell.fill = GOOD_FILL
    cell.font = Font(name=cell.font.name or "Calibri", size=11, color="1565A6")
    cell.alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)
    letter = get_column_letter(col)
    sheet.column_dimensions[letter].width = max(sheet.column_dimensions[letter].width or 0, 38)
    sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 0, 65)


def _schema_for(source: _Workbook, template: Path | None, template_dir: Path, template_cache: dict) -> tuple[_Workbook, bool]:
    if template:
        schema = _Workbook(SourceBinding("TEMPLATE", template))
        if not numeric_regions(schema.book)[0]:
            schema.close()
            raise ValueError("所选模板没有数字区域、数值区域或条件格式区域命名")
        return schema, True
    candidates = []
    for path in sorted(template_dir.glob("*.xlsx")):
        if path not in template_cache:
            candidate = _Workbook(SourceBinding("TEMPLATE", path))
            try:
                template_cache[path] = {region.sheet for region in numeric_regions(candidate.book)[0]}
            finally:
                candidate.close()
        data_sheets = template_cache[path]
        if data_sheets and data_sheets.issubset(source.book.sheetnames):
            candidates.append((len(data_sheets), path))
    if not candidates:
        # A template itself can also be selected for expression generation.
        if numeric_regions(source.book)[0]:
            return source, False
        raise ValueError("源文件没有数字区域命名，且未自动匹配到模板；请选择对应的命名区域模板")
    score = max(item[0] for item in candidates)
    best = [path for value, path in candidates if value == score]
    if len(best) != 1:
        raise ValueError("匹配到多个布局模板，请明确选择模板：" + "、".join(path.name for path in best))
    return _Workbook(SourceBinding("TEMPLATE", best[0])), True


def _numeric_entries(schema: _Workbook, source_path: Path) -> list[dict]:
    regions, origin = numeric_regions(schema.book)
    entries = []
    explicit_sheets = {r.sheet for r in named_regions(schema.book, ("数字区域", "数值区域", "条件格式区域"))}
    if "表结构区域交叉推导" in origin:
        for region in inferred_numeric_regions(schema.book)[1]:
            if region.sheet in explicit_sheets:
                continue
            entries.append(_entry(source_path, schema.path, sheet=region.sheet, region=region.name,
                                  status="STRUCTURE_REGION_UNPAIRED",
                                  detail=f"表结构区域没有可交叉配对的行/列表头：{region.reference}；请修正区域归属或增加独立数字区域"))
    seen = set()
    total_cells = sum((r.max_row - r.min_row + 1) * (r.max_col - r.min_col + 1) for r in regions)
    if total_cells > 500_000:
        raise ValueError("数字区域超过 50 万个单元格，请缩小命名区域")
    for data in regions:
        try:
            row_header, col_header = numeric_header_regions(schema, data)
            header_error = ""
        except RetrievalError as exc:
            row_header = col_header = None
            header_error = f"{exc.code}: {exc.message}"
        row_path_index = (row_path_candidates(
            schema, data.sheet, row_header,
            candidate_rows=range(data.min_row, data.max_row + 1),
        ) if row_header else {})
        for row in range(data.min_row, data.max_row + 1):
            for col in range(data.min_col, data.max_col + 1):
                if (data.sheet, row, col) in seen:
                    continue
                seen.add((data.sheet, row, col))
                cell = schema.book[data.sheet].cell(row, col).coordinate
                entry = _entry(source_path, schema.path, sheet=data.sheet, region=data.name, template_cell=cell,
                               detail=f"布局来自{origin}")
                entries.append(entry)
                if header_error:
                    entry.update(status="HEADER_REGION_INVALID", detail=header_error)
                    continue
                if schema._anchor(data.sheet, row, col) != (row, col):
                    entry.update(status="MERGED_COVERED", detail="合并数据单元格的非左上角位置；只生成锚点表达式")
                    continue
                row_info = row_path_index.get(row)
                row_path = (row_info["primary"] if row_info else
                            schema.header_path(data.sheet, [(row, c) for c in range(row_header.min_col, row_header.max_col + 1)]))
                col_path = schema.header_path(data.sheet, [(r, col) for r in range(col_header.min_row, col_header.max_row + 1)])
                entry.update(row_path="@".join(row_path), column_path="@".join(col_path), _row_path=row_path, _col_path=col_path)
                if not row_path or not col_path:
                    entry.update(status="HEADER_PATH_EMPTY", detail="行路径或列路径为空；请检查表结构区域")
    return entries


def _export_template(source: _Workbook, output: Path, template: Path | None, template_dir: Path, template_cache: dict) -> list[dict]:
    schema, owns_schema = _schema_for(source, template, template_dir, template_cache)
    try:
        if source is not schema:
            projected = project_numeric_regions(source, schema)
        else:
            projected = numeric_regions(schema.book)[0]
        source._numeric_region_overrides = {
            sheet: [region for region in projected if region.sheet == sheet]
            for sheet in {region.sheet for region in projected}
        }
        entries = _numeric_entries(schema, source.path)
        key_counts = Counter((e["sheet"], e.get("_row_path"), e.get("_col_path")) for e in entries if not e["status"])
        indexes = {}
        for sheet in {e["sheet"] for e in entries if not e["status"]}:
            if sheet not in source.book.sheetnames:
                continue
            pairs = {(e["_row_path"], e["_col_path"]) for e in entries if e["sheet"] == sheet and not e["status"]}
            indexes[sheet] = scan_rule_paths(source, sheet, pairs)[0]
        named_cells = defaultdict(set)
        for entry in entries:
            if entry["status"]:
                if source is schema and entry["template_cell"] and entry["status"] != "MERGED_COVERED":
                    source.book[entry["sheet"]][entry["template_cell"]].fill = ERROR_FILL
                continue
            sheet = entry["sheet"]
            key = (entry["_row_path"], entry["_col_path"])
            matches = indexes.get(sheet, {}).get(key, [])
            entry["cell"] = ", ".join(f"{get_column_letter(col)}{row}" for row, col in matches)
            if key_counts[(sheet, *key)] > 1 or len(matches) > 1:
                entry.update(status="PATH_AMBIGUOUS", detail="完整行列路径重复，三参数表达式无法唯一定位")
            elif not matches:
                entry.update(status="PATH_NOT_FOUND", detail="源文件缺少模板中的完整行列路径")
            elif source is schema and entry["cell"] != entry["template_cell"]:
                entry.update(status="ROUNDTRIP_MISMATCH", detail="路径回读坐标与数字区域单元格不同")
            else:
                row, col = matches[0]
                try:
                    expression = template_expression(sheet, *key)
                except (ConfigurationError, ValueError) as exc:
                    entry.update(status="EXPRESSION_INVALID", detail=str(exc))
                    source.book[sheet].cell(row, col).fill = ERROR_FILL
                    continue
                if source._anchor(sheet, row, col) != (row, col):
                    entry.update(status="MERGED_COVERED", detail="源文件目标被合并；只生成合并锚点")
                elif source.book[sheet].cell(row, col).data_type in {"f", "e"}:
                    _put_expression(source.book[sheet], row, col, expression)
                    entry.update(status="SOURCE_VALUE_UNVERIFIED", expression=expression,
                                 detail="定位文本已生成，但原格含 Excel 公式或错误值；审核此输入前须提供可信数值源，不能视为取值已通过")
                else:
                    _put_expression(source.book[sheet], row, col, expression)
                    named_cells[(sheet, entry["region"])].add((row, col))
                    entry.update(status="READY", expression=expression)
            if entry["status"] not in {"READY", "MERGED_COVERED"}:
                for row, col in matches:
                    source.book[sheet].cell(*source._anchor(sheet, row, col)).fill = ERROR_FILL
        for index, ((sheet, _), cells) in enumerate(sorted(named_cells.items()), 1):
            rows, cols = {r for r, _ in cells}, {c for _, c in cells}
            if len(cells) == (max(rows) - min(rows) + 1) * (max(cols) - min(cols) + 1):
                area = f"{quote_sheetname(sheet)}!${get_column_letter(min(cols))}${min(rows)}:${get_column_letter(max(cols))}${max(rows)}"
                name = f"数字区域_{index:03d}"
                if name not in source.book.defined_names and not any(name in s.defined_names for s in source.book):
                    source.book.defined_names.add(DefinedName(name, attr_text=area))
        if file_sha256(schema.path) != schema.sha256:
            raise ValueError("布局模板在生成期间发生变化，请重新运行")
        source.book.save(output)
        return entries
    finally:
        if owns_schema:
            schema.close()


def _export_external(source: _Workbook, output: Path, unit: str, source_id: str) -> list[dict]:
    layout = discover_external_layout(source, support_time_series=True)
    entries = []
    seen = set()
    header_counts = Counter(normalized(c.header) for c in layout.columns)
    for column in layout.columns:
        keywords = []
        if header_counts[normalized(column.header)] > 1:
            keywords.append(f"sheet={_literal(column.sheet)}")
        if source_id != "EXT":
            keywords.append(f"source={_literal(source_id)}")
        if unit != "万元" and not re.search(r"[（(]\s*(元|万元|亿元|百分比|百分点|户|笔)\s*[）)]", column.header):
            # Unit overrides apply only to amount metrics, never rates or counts.
            if _metric_profile((column.header,))[0] == "amount":
                keywords.append(f"unit={_literal(unit)}")
        expression = f"V_EXT({_literal(column.header)}" + (", " + ", ".join(keywords) if keywords else "") + ")"
        parse_expression(expression)
        for report_item, rows in column.rows.items():
            for row in rows:
                key = (column.sheet, row, column.column)
                if key in seen:
                    continue
                seen.add(key)
                address = f"{get_column_letter(column.column)}{row}"
                entry = _entry(source.path, sheet=column.sheet, cell=address, column_path=column.header,
                               template_cell=f"{get_column_letter(column.column)}{column.header_row}")
                entries.append(entry)
                try:
                    resolved, resolved_row = layout.resolve(
                        column.header, report_item,
                        column.sheet if header_counts[normalized(column.header)] > 1 else "",
                        row_hint=row,
                    )
                    if (resolved.sheet, resolved_row, resolved.column) != key:
                        raise RetrievalError("ROUNDTRIP_MISMATCH", "外部表达式回读位置不一致")
                    cell = source.book[column.sheet].cell(row, column.column)
                    if cell.data_type in {"f", "e"}:
                        _put_expression(source.book[column.sheet], row, column.column, expression)
                        entry["expression"] = expression
                        raise RetrievalError("SOURCE_VALUE_UNVERIFIED", "源数据含公式或 Excel 错误值")
                    _put_expression(source.book[column.sheet], row, column.column, expression)
                    entry.update(status="READY", expression=expression, detail="机构由审核时选择的机构代码映射；表达式不固定机构行")
                except RetrievalError as exc:
                    entry.update(status=exc.code, detail=exc.message)
                    source.book[column.sheet].cell(row, column.column).fill = ERROR_FILL
    for column in layout.ignored_columns:
        for rows in column.rows.values():
            for row in rows:
                entries.append(_entry(source.path, sheet=column.sheet,
                                      cell=f"{get_column_letter(column.column)}{row}", column_path=column.header,
                                      status="HEADER_PLACEHOLDER", detail="null / none 等占位表头不是可唯一取数的业务指标，跳过表达式生成"))
    source.book.save(output)
    return entries


def _write_report(path: Path, entries: list[dict], counts: dict[str, int]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "表达式生成报告"
    sheet.append(["生成成功", "待确认或未生成", "合并或占位表头跳过", "处理单元格/文件项"])
    sheet.append([counts["ready"], counts["failed"], counts["skipped"], len(entries)])
    sheet.append([])
    sheet.append(REPORT_HEADERS)
    for row, entry in enumerate(entries, start=5):
        sheet.append([entry[key] for key in ("file", "template", "sheet", "region", "template_cell", "cell", "row_path", "column_path", "expression", "status", "detail")])
        if entry["status"] not in {"READY", *SKIPPED_STATUSES}:
            for cell in sheet[row]:
                cell.fill = ERROR_FILL
    for row in (1, 4):
        for cell in sheet[row]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(bold=True, color="FFFFFF")
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in enumerate((36, 36, 24, 24, 16, 18, 50, 45, 110, 28, 65), 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = "I5"
    sheet.auto_filter.ref = f"A4:K{max(sheet.max_row, 4)}"
    workbook.save(path)
    workbook.close()


def export_expressions(paths: list[str | Path], *, mode: str, output_root: str | Path | None = None,
                       template: str | Path | None = None, template_dir: str | Path | None = None,
                       external_unit: str = "万元", source_id: str = "EXT") -> dict:
    if mode not in {"vcur", "vext"}:
        raise ValueError("表达式生成类型必须是 vcur 或 vext")
    sources = list(dict.fromkeys(Path(p).expanduser().resolve() for p in paths if str(p).strip()))
    if not sources:
        raise ValueError("请至少选择一个源文件")
    if len(sources) > 200:
        raise ValueError("单次最多选择 200 个源文件")
    if external_unit not in {"万元", "亿元", "元"}:
        raise ValueError("外部金额单位必须是万元、亿元或元")
    output_root = Path(output_root or ROOT / "data" / "s5" / "维护输出" / "表达式参考").expanduser().resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    output_dir = output_root / f"{mode}_{uuid.uuid4().hex[:12]}"
    output_dir.mkdir(parents=True)
    template_path = Path(template).expanduser().resolve() if template else None
    template_dir = Path(template_dir or TEMPLATE_DIR).resolve()
    all_entries, files, template_cache = [], [], {}
    for index, path in enumerate(sources, 1):
        book = None
        try:
            if path.suffix.lower() != ".xlsx":
                raise ValueError("表达式参考副本仅支持 XLSX；请先将源文件另存为 XLSX")
            book = _Workbook(SourceBinding("EXPORT", path), editable=True)
            output = output_dir / f"{index:03d}_{path.stem}_{mode.upper()}.xlsx"
            entries = (_export_template(book, output, template_path, template_dir, template_cache) if mode == "vcur"
                       else _export_external(book, output, external_unit, source_id))
            if file_sha256(path) != book.sha256:
                raise ValueError("源文件在生成期间发生变化，请重新运行")
            all_entries.extend(entries)
            files.append({"source": str(path), "path": str(output), "ready": sum(e["status"] == "READY" for e in entries),
                          "input_unchanged": True})
        except Exception as exc:
            all_entries.append(_entry(path, status="FILE_FAILED", detail=f"{type(exc).__name__}: {exc}"))
        finally:
            if book:
                book.close()
    counts = {"ready": sum(e["status"] == "READY" for e in all_entries),
              "skipped": sum(e["status"] in SKIPPED_STATUSES for e in all_entries)}
    counts["failed"] = len(all_entries) - counts["ready"] - counts["skipped"]
    report = output_dir / "表达式生成报告.xlsx"
    _write_report(report, all_entries, counts)
    status = "COMPLETED" if not counts["failed"] else "COMPLETED_WITH_ISSUES" if counts["ready"] else "FAILED"
    return {"ok": status != "FAILED", "status": status, "mode": mode, "counts": counts, "files": files,
            "report_path": str(report), "output_dir": str(output_dir)}


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="批量生成 V_CUR 或 V_EXT 表达式参考副本")
    parser.add_argument("mode", choices=("vcur", "vext"))
    parser.add_argument("files", nargs="+")
    parser.add_argument("--template", help="源文件没有命名区域时使用的布局模板")
    parser.add_argument("--template-dir", type=Path, default=TEMPLATE_DIR)
    parser.add_argument("--unit", default="万元", choices=("万元", "亿元", "元"))
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    result = export_expressions(args.files, mode=args.mode, template=args.template, template_dir=args.template_dir,
                                external_unit=args.unit, output_root=args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "COMPLETED" else 2


if __name__ == "__main__":
    raise SystemExit(main())
