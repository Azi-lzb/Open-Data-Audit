"""Deterministic Excel validation-formula conversion into review-only XLSX rules."""
from __future__ import annotations

import hashlib
import json
import re
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

from s1audit_configured.dsl import Literal as DslLiteral, TupleLiteral as DslTupleLiteral, Variable as DslVariable
from s1audit_configured.dsl import parse_dictionary, parse_expression, calls_in, validate_dictionary_dependencies, validate_static_types, variables_in
from s1audit_configured.dsl import Binary as DslBinary, Unary as DslUnary, Call as DslCall, Literal as DslLiteral, TupleLiteral as DslTupleLiteral
from .expression_export import _numeric_entries, template_expression
from s1audit_configured.formula_error_scan import formula_error_tokens
from s1audit_configured.external_layout import discover_external_layout, normalized
from s1audit_configured.layout import named_regions, numeric_header_regions, numeric_regions
from s1audit_configured.model import SourceBinding
from s1audit_configured.sources import _Workbook, file_sha256
from s1audit_configured.workbook_config import _description_variables, _metric_profile


ROOT = Path(__file__).resolve().parents[5]
CANDIDATE_HEADERS = ("规则编号", "是否启用", "业务名称", "工作表名称", "校验公式单元格", "规则名称", "审核级别", "校验类型", "校验指标",
                     "审核表达式", "表达式布尔值取反", "规则描述", "当前值表达式", "对比值表达式", "差值表达式", "取值字典",
                     "数值运算模式")
SEVERITIES = {"错": "错误", "错误": "错误", "硬性": "错误", "硬性错误": "错误", "核实": "核实", "软性": "核实", "软性错误": "核实", "警告": "核实"}

# Formula reviewed by the AI triage where the output message does not use the
# normal ``severity|metric|description`` protocol. Require an exact source
# formula match so a future template change cannot silently inherit this rule.
AI_SUPPLEMENTAL_RULES = {
    ("！70.金融基础数据采集（债券业务）202608.xlsx", "存量债券发行", "AF88"): {
        "formula": '=IF(ABS(AB88*10000)>1,Y88&"|"&TEXT(AB88,"0.00")&"|"&AE88,"")',
        "name": "债券发行合计大集中核对",
        "metric": "债券发行合计",
        "external": 'V_EXT("（十四）债券负债（人民币合计）", sheet="集中系统数据", unit="亿元", standard_unit="万元")',
        "note": "模板本地侧最终引用 P88；该单元格为空且在数字区域外，按原公式空白数值计算为0。该候选忠实保留现有模板逻辑，但业务引用疑似不完整，需人工确认后再启用。",
    },
    ("！70.金融基础数据采集（债券业务）202608.xlsx", "存量债券投资", "AB155"): {
        "formula": '=IF(ABS(X155*10000)>1,U155&"|"&TEXT(X155,"0.00")&"|"&AA155,"")',
        "name": "债券投资合计大集中核对",
        "metric": "债券投资合计",
        "external": 'V_EXT("（二十五）债券资产（2017准则）（人民币合计）", sheet="集中系统数据", unit="亿元", standard_unit="万元") + V_EXT("（二十）债券资产（人民币合计）", sheet="集中系统数据", unit="亿元", standard_unit="万元")',
        "note": "模板本地侧最终引用 L155；该单元格为空且在数字区域外，按原公式空白数值计算为0。D155 存在可定位的本期余额路径，但替换为 D155 会改变原公式含义，因此本候选仍忠实按空白0生成并需人工确认模板是否错引。",
    },
}


def _ai_supplemental_candidate(template_name, business, sheet, address, formula, rule_sequences):
    spec = AI_SUPPLEMENTAL_RULES.get((template_name, sheet, address))
    if not spec:
        return None
    if str(formula or "").strip() != spec["formula"]:
        raise Pending("AI补充候选的模板公式指纹已变化，需重新分析后再生成")
    rule_sequences[(business, sheet)] += 1
    rule_id = f"{business}_{sheet}_{rule_sequences[(business, sheet)]}"
    dictionary = (
        '{"本期值": 0,\n'
        ' "对比值": ' + spec["external"] + ',\n'
        ' "差值": ({本期值} - {对比值})}'
    )
    return {
        "规则编号": rule_id, "是否启用": "否", "业务名称": business,
        "工作表名称": sheet, "校验公式单元格": address, "规则名称": spec["name"], "审核级别": "错误",
        "校验类型": "大集中不匹配", "校验指标": spec["metric"],
        "审核表达式": "ABS({差值}) > 1", "表达式布尔值取反": "是",
        "规则描述": spec["note"], "当前值表达式": "{本期值}",
        "对比值表达式": "{对比值}", "差值表达式": "{差值}",
        "取值字典": dictionary, "数值运算模式": "Excel双精度",
    }


class Pending(ValueError):
    pass


class _StaticFormulaError(Pending):
    """A deterministic Excel error which an enclosing IFERROR may catch."""


def _current_value_path(current, dictionary, rule_sheet):
    """Show every V_CUR path contributing to the displayed current value."""
    if not current:
        return ""
    try:
        variables = variables_in(parse_expression(current))
        if len(variables) != 1:
            return ""
        node = parse_expression(dictionary[variables.pop()])
    except Exception:
        return ""

    def path_text(argument):
        if isinstance(argument, DslLiteral) and argument.kind == "string":
            return argument.value
        if isinstance(argument, DslTupleLiteral):
            return "@".join(argument.values)
        return ""

    paths = []

    def collect(item):
        if isinstance(item, DslCall):
            if item.name in {"V_PRE", "V_EXT"}:
                return
            if item.name == "V_CUR":
                if len(item.arguments) != 3:
                    raise ValueError("本期取值路径参数不完整")
                sheet_arg = item.arguments[0]
                if not isinstance(sheet_arg, DslLiteral) or sheet_arg.kind != "string":
                    raise ValueError("本期取值工作表无法确定")
                row_path, column_path = (path_text(argument) for argument in item.arguments[1:])
                if not row_path or not column_path:
                    raise ValueError("本期取值路径为空")
                paths.append((sheet_arg.value, row_path, column_path))
                return
            for argument in item.arguments:
                collect(argument)
        elif isinstance(item, DslBinary):
            collect(item.left)
            collect(item.right)
        elif isinstance(item, DslUnary):
            collect(item.operand)

    try:
        collect(node)
    except ValueError:
        return ""
    paths = list(dict.fromkeys(paths))
    if not paths:
        return ""
    return "；".join(
        (f"{sheet}!" if sheet != rule_sheet else "") + f"{row_path} → {column_path}"
        for sheet, row_path, column_path in paths
    )


def _path_parts(argument):
    if isinstance(argument, DslLiteral) and argument.kind == "string":
        return tuple(argument.value.split("@"))
    if isinstance(argument, DslTupleLiteral):
        return argument.values
    return ()


def _template_path_checks(candidate, rule_id, business, formula_sheet, formula_cell, converter):
    try:
        dictionary = parse_dictionary(candidate["取值字典"])
    except Exception:
        return []
    checks = []
    for variable, expression in dictionary.items():
        for call in calls_in(expression):
            if call.name != "V_CUR" or len(call.arguments) != 3:
                continue
            sheet_arg, row_arg, column_arg = call.arguments
            if not isinstance(sheet_arg, DslLiteral) or sheet_arg.kind != "string":
                continue
            target_sheet = sheet_arg.value
            row_path, column_path = _path_parts(row_arg), _path_parts(column_arg)
            if not row_path or not column_path:
                continue
            expected = template_expression(target_sheet, row_path, column_path)
            locations = [f"{sheet}!{address}" for (sheet, address), (expression, _) in converter.cells.items()
                         if expression == expected]
            status = "数字区域唯一匹配" if len(locations) == 1 else "数字区域无匹配" if not locations else "数字区域路径重复"
            checks.append({
                "template": converter.book.path.name, "business": business, "rule_id": rule_id,
                "formula_sheet": formula_sheet, "formula_cell": formula_cell,
                "variable": variable, "source_sheet": target_sheet,
                "row_path": "@".join(row_path), "column_path": "@".join(column_path),
                "template_cells": "；".join(locations), "status": status,
            })
    return checks


def _serialize_dsl(node):
    if isinstance(node, DslLiteral):
        if node.kind == "string":
            return json.dumps(node.value, ensure_ascii=False)
        if node.kind == "boolean":
            return "TRUE" if node.value else "FALSE"
        return str(node.value)
    if isinstance(node, DslVariable):
        return "{" + node.name + "}"
    if isinstance(node, DslTupleLiteral):
        return "(" + ", ".join(json.dumps(value, ensure_ascii=False) for value in node.values) + ")"
    if isinstance(node, DslCall):
        arguments = [_serialize_dsl(argument) for argument in node.arguments]
        arguments.extend(f"{key}={json.dumps(value, ensure_ascii=False)}" for key, value in node.keywords)
        return f"{node.name}({', '.join(arguments)})"
    if isinstance(node, DslBinary):
        return f"({_serialize_dsl(node.left)} {node.operator} {_serialize_dsl(node.right)})"
    if isinstance(node, DslUnary):
        if node.operator == "NOT":
            return f"NOT({_serialize_dsl(node.operand)})"
        return f"{node.operator}({_serialize_dsl(node.operand)})"
    raise TypeError(f"不支持序列化的表达式节点：{type(node).__name__}")


def _hoist_dictionary_wrappers(dictionary, expressions):
    """Move display/coercion wrappers off dictionary definitions while preserving use-site behavior."""
    wrappers = {}
    parsed = {}
    for name, expression in dictionary.items():
        try:
            node = parse_expression(expression)
        except Exception:
            continue
        parsed[name] = node
        chain = []
        inner = node
        while isinstance(inner, DslCall):
            if inner.name == "SUM" and len(inner.arguments) == 1:
                chain.append(("SUM", None))
                inner = inner.arguments[0]
            elif (inner.name == "ROUND" and len(inner.arguments) == 2
                  and isinstance(inner.arguments[1], DslLiteral) and inner.arguments[1].kind == "number"):
                chain.append(("ROUND", _serialize_dsl(inner.arguments[1])))
                inner = inner.arguments[0]
            else:
                break
        if chain:
            # SUM of a scalar SUM result is identical to that single SUM; keep
            # one coercion for Excel blank semantics without noisy nesting.
            compact_chain = []
            for wrapper in chain:
                if wrapper[0] == "SUM" and compact_chain and compact_chain[-1][0] == "SUM":
                    continue
                compact_chain.append(wrapper)
            wrappers[name] = (tuple(compact_chain), inner)
    changed = True
    while changed:
        changed = False
        for name, node in parsed.items():
            if name not in wrappers and isinstance(node, DslVariable) and node.name in wrappers:
                chain, _ = wrappers[node.name]
                wrappers[name] = (chain, node)
                changed = True
    if not wrappers:
        return dictionary, expressions

    def wrap_references(text, *, include_wrappers=True):
        if not text:
            return text
        for name, (chain, _) in wrappers.items():
            if not include_wrappers:
                continue
            argument = "{" + name + "}"
            replacement = argument
            for function, digits in reversed(chain):
                replacement = f"{function}({replacement})" if digits is None else f"ROUND({replacement}, {digits})"
            text = text.replace(argument, replacement)
        return text

    normalized_dictionary = {}
    for name, expression in dictionary.items():
        if name in wrappers:
            normalized_dictionary[name] = _serialize_dsl(wrappers[name][1])
        else:
            normalized_dictionary[name] = wrap_references(expression)
    normalized_expressions = dict(expressions)
    for name, expression in normalized_expressions.items():
        # Descriptions interpolate values as text and cannot host executable DSL calls.
        normalized_expressions[name] = wrap_references(expression, include_wrappers=name != "规则描述")
    return normalized_dictionary, normalized_expressions


@dataclass(frozen=True)
class Expr:
    kind: str
    value: str = ""
    args: tuple["Expr", ...] = ()


class ExcelParser:
    def __init__(self, formula):
        try:
            self.tokens = [token for token in Tokenizer(formula).items if token.type != "WHITE-SPACE"]
        except Exception as exc:
            raise Pending(f"Excel公式分词失败：{exc}") from exc
        self.index = 0

    def parse(self):
        value = self.expression()
        if self.index != len(self.tokens):
            raise Pending(f"暂不支持的公式片段：{self.tokens[self.index].value}")
        return value

    def expression(self, minimum=0):
        if self.index >= len(self.tokens):
            raise Pending("公式缺少表达式")
        token = self.tokens[self.index]
        self.index += 1
        if token.type == "OPERATOR-PREFIX":
            left = Expr("unary", token.value, (self.expression(50),))
        elif token.type == "FUNC" and token.subtype == "OPEN":
            name, args = token.value[:-1].upper(), []
            while self.index < len(self.tokens) and self.tokens[self.index].subtype != "CLOSE":
                args.append(self.expression())
                if self.index < len(self.tokens) and self.tokens[self.index].type == "SEP":
                    self.index += 1
                else:
                    break
            if self.index >= len(self.tokens) or self.tokens[self.index].subtype != "CLOSE":
                raise Pending(f"函数{name}缺少右括号")
            self.index += 1
            left = Expr("call", name, tuple(args))
        elif token.type == "PAREN" and token.subtype == "OPEN":
            left = self.expression()
            if self.index >= len(self.tokens) or self.tokens[self.index].type != "PAREN" or self.tokens[self.index].subtype != "CLOSE":
                raise Pending("缺少右括号")
            self.index += 1
        elif token.type == "OPERAND":
            kind = {"RANGE": "ref", "TEXT": "text", "NUMBER": "number", "LOGICAL": "boolean"}.get(token.subtype)
            if not kind:
                raise Pending(f"不支持的常量：{token.value}")
            value = token.value[1:-1].replace('""', '"') if kind == "text" else token.value
            left = Expr(kind, value)
        else:
            raise Pending(f"无法解析：{token.value}")
        precedence = {"=": 10, "<>": 10, ">": 10, "<": 10, ">=": 10, "<=": 10, "&": 15, "+": 20, "-": 20, "*": 30, "/": 30}
        while self.index < len(self.tokens):
            token = self.tokens[self.index]
            if token.type == "OPERATOR-POSTFIX" and token.value == "%":
                self.index += 1
                left = Expr("binary", "/", (left, Expr("number", "100")))
                continue
            if token.type != "OPERATOR-INFIX" or token.value not in precedence or precedence[token.value] < minimum:
                break
            self.index += 1
            left = Expr("binary", token.value, (left, self.expression(precedence[token.value] + 1)))
        return left


def parse_excel(formula):
    return ExcelParser(formula).parse()


def join_parts(parts):
    if not parts:
        return Expr("text", "")
    result = parts[0]
    for item in parts[1:]:
        result = Expr("binary", "&", (result, item))
    return result


def concat_parts(node):
    if node.kind == "binary" and node.value == "&":
        return concat_parts(node.args[0]) + concat_parts(node.args[1])
    return [node]


def protocol_parts(node):
    segments = [[]]
    for part in concat_parts(node):
        if part.kind == "text":
            pieces = part.value.split("|")
            for index, piece in enumerate(pieces):
                if index:
                    segments.append([])
                if piece:
                    segments[-1].append(Expr("text", piece))
        else:
            segments[-1].append(part)
    result = [join_parts(parts) for parts in segments]
    prefix = "".join(part.value for part in segments[0] if part.kind == "text")
    # Some templates use a compact two-part message such as
    # ``错误|差值为"&<numeric expression>``.  The condition and displayed
    # value still carry the complete check; rule() derives the value/path from
    # the numeric message expression instead of requiring a redundant metric
    # label in the formula text.
    if len(result) < 2 or prefix not in SEVERITIES or any(part.kind != "text" for part in segments[0]):
        return None
    return SEVERITIES[prefix], result


def error_branches(node, guards=()):
    if node.kind == "call" and node.value == "IF" and len(node.args) == 3:
        yield from error_branches(node.args[1], guards + (node.args[0],))
        yield from error_branches(node.args[2], guards + (Expr("call", "NOT", (node.args[0],)),))
    else:
        parts = protocol_parts(node)
        if parts:
            yield guards, parts


def reference_parts(text, default_sheet):
    if "!" in text:
        prefix, address = text.rsplit("!", 1)
        prefix = prefix.strip("'").replace("''", "'")
        external = "[" in prefix
        sheet = prefix.split("]", 1)[-1] if external else prefix
    else:
        address, sheet, external = text, default_sheet, False
    return sheet, address.replace("$", ""), external


def _has_division(node):
    if isinstance(node, DslBinary):
        return node.operator == "/" or _has_division(node.left) or _has_division(node.right)
    if isinstance(node, DslUnary):
        return _has_division(node.operand)
    if isinstance(node, DslCall):
        return any(_has_division(arg) for arg in node.arguments)
    return False


class Converter:
    def __init__(self, book, external=None, external_unit="万元"):
        self.book, self.external, self.external_unit = book, external, external_unit
        self.cells = {}
        self.bad_cells = {}
        self.cell_paths = {}
        entries = _numeric_entries(book, book.path)
        counts = Counter((entry["sheet"], entry.get("_row_path"), entry.get("_col_path")) for entry in entries if not entry["status"])
        for entry in entries:
            key = (entry["sheet"], entry["template_cell"])
            self.cell_paths[key] = (tuple(entry.get("_row_path") or ()), tuple(entry.get("_col_path") or ()))
            if entry["status"] or counts[(entry["sheet"], entry.get("_row_path"), entry.get("_col_path"))] != 1:
                self.bad_cells[key] = entry["detail"] or "数字区域中的行列路径重复"
            else:
                self.cells[key] = (template_expression(entry["sheet"], entry["_row_path"], entry["_col_path"]), entry["row_path"])
        self.regions = named_regions(book.book, ("校验区域",))
        self.parsed = {}
        self.layout = None
        self.reset()

    def reset(self):
        self.dictionary, self.aliases, self.expression_aliases = {}, {}, {}
        self.notes = []

    def kind(self, node, sheet, stack=()):
        if len(stack) > 60:
            raise Pending("公式引用层级过深")
        if node.kind == "ref":
            name, address, cell = self.ref(node.value, sheet)
            key = (name, address)
            if key in stack:
                raise Pending(f"公式循环引用：{name}!{address}")
            if key in self.cells or key in self.bad_cells:
                return "number"
            expanded, name = self.expand(node, sheet)
            if expanded != node:
                return self.kind(expanded, name, stack + (key,))
            return "text" if isinstance(cell.value, str) or cell.value is None else "number"
        if node.kind == "binary":
            return "text" if node.value == "&" else "number"
        if node.kind == "unary":
            return self.kind(node.args[0], sheet, stack)
        if node.kind == "call":
            if node.value == "IF" and len(node.args) == 3:
                left, right = (self.kind(item, sheet, stack) for item in node.args[1:])
                return left if left == right else "mixed"
            return "number"
        return node.kind

    def in_validation(self, sheet, row, col):
        return any(region.sheet == sheet and region.min_row <= row <= region.max_row and region.min_col <= col <= region.max_col for region in self.regions)

    def ref(self, text, sheet):
        name, address, external = reference_parts(text, sheet)
        if external or name not in self.book.book.sheetnames:
            raise Pending(f"外部引用{name}!{address}需要可确认的外部查找公式，不能固定机构行")
        if not re.fullmatch(r"[A-Za-z]+[1-9]\d*", address):
            raise Pending(f"引用不是单个单元格：{text}")
        cell = self.book.book[name][address]
        return name, cell.coordinate, cell

    def expand(self, node, sheet):
        if node.kind == "ref":
            name, address, cell = self.ref(node.value, sheet)
            if cell.data_type == "f" and (name, address) not in self.cells:
                # Only scan output formulas inside 校验区域. A referenced helper
                # formula outside that region is a dependency, not a new rule;
                # follow just that cell instead of scanning the surrounding sheet.
                key = (name, address)
                if key not in self.parsed:
                    self.parsed[key] = parse_excel(cell.value)
                return self.parsed[key], name
        return node, sheet

    def static_text(self, node, sheet, stack=()):
        if len(stack) > 60:
            raise Pending("文本引用层级过深")
        if node.kind == "text":
            return node.value
        if node.kind == "number":
            return node.value
        if node.kind == "binary" and node.value == "&":
            return "".join(self.static_text(item, sheet, stack) for item in node.args)
        if node.kind == "unary" and node.value == "+":
            return self.static_text(node.args[0], sheet, stack)
        if node.kind == "ref":
            name, address, cell = self.ref(node.value, sheet)
            key = (name, address)
            if key in self.cells or key in self.bad_cells:
                raise Pending("文本包含动态数值")
            if key in stack:
                raise Pending(f"公式循环引用：{name}!{address}")
            if cell.data_type == "f":
                value = self.static_value(node, sheet, stack)
                if isinstance(value, Decimal) and value == value.to_integral_value():
                    return str(int(value))
                return str(value if value is not None else "")
            if cell.data_type == "e":
                raise Pending(f"模板辅助公式引用到 Excel 错误值：{name}!{address}={cell.value}")
            return str(cell.value if cell.value is not None else "")
        if node.kind == "call" and node.value == "ADDRESS":
            if len(node.args) < 2:
                raise Pending("ADDRESS参数不足")
            row = self.integer(node.args[0], sheet)
            col = self.integer(node.args[1], sheet)
            if len(node.args) > 3 and node.args[3].value.upper() not in {"1", "TRUE"}:
                raise Pending("暂不支持R1C1形式的INDIRECT")
            return f"${get_column_letter(col)}${row}"
        if node.kind == "call" and node.value == "CELL" and len(node.args) == 2 and node.args[0].kind == "text" and node.args[0].value.lower() == "address" and node.args[1].kind == "ref":
            _, _, cell = self.ref(node.args[1].value, sheet)
            return f"${cell.column_letter}${cell.row}"
        try:
            value = self.static_value(node, sheet, stack)
        except Pending as exc:
            raise Pending("文本包含不能静态确定的计算") from exc
        if isinstance(value, Decimal) and value == value.to_integral_value():
            return str(int(value))
        return str(value if value is not None else "")

    def _static_range(self, node, default_sheet):
        if node.kind != "ref":
            raise Pending("静态查找范围必须是单行或单列引用")
        name, address, external = reference_parts(node.value, default_sheet)
        if external or name not in self.book.book.sheetnames:
            raise Pending("静态查找不能跨工作簿")
        full_column = re.fullmatch(r"\$?([A-Za-z]{1,3}):\$?([A-Za-z]{1,3})", address)
        full_row = re.fullmatch(r"\$?([1-9]\d*):\$?([1-9]\d*)", address)
        if full_column and full_column.group(1).upper() == full_column.group(2).upper():
            col = column_index_from_string(full_column.group(1))
            return name, 1, col, max(1, self.book.book[name].max_row), col
        if full_row and full_row.group(1) == full_row.group(2):
            row = int(full_row.group(1))
            return name, row, 1, row, max(1, self.book.book[name].max_column)
        try:
            min_col, min_row, max_col, max_row = range_boundaries(address)
        except ValueError as exc:
            raise Pending(f"静态查找范围无法解析：{node.value}") from exc
        if not all((min_col, min_row, max_col, max_row)) or (min_col != max_col and min_row != max_row):
            raise Pending("静态查找只支持单行或单列范围")
        if (max_col-min_col+1)*(max_row-min_row+1) > 10000:
            raise Pending("静态查找范围超过10000格")
        return name, min_row, min_col, max_row, max_col

    def _static_equal(self, left, right):
        if left is None or left == "":
            return right is None or right == ""
        if right is None or right == "":
            return False
        if isinstance(left, bool) or isinstance(right, bool):
            return isinstance(left, bool) and isinstance(right, bool) and left == right
        if isinstance(left, (int, float, Decimal)) and isinstance(right, (int, float, Decimal)):
            try:
                return Decimal(str(left)) == Decimal(str(right))
            except Exception:
                return False
        return str(left).strip().casefold() == str(right).strip().casefold()

    def _static_match(self, node, sheet, stack):
        if len(node.args) not in {2, 3}:
            raise Pending("MATCH 静态解析只支持两至三个参数")
        mode = self.static_value(node.args[2], sheet, stack) if len(node.args) == 3 else 1
        if mode not in (0, Decimal(0)):
            raise Pending("MATCH 只支持精确匹配")
        expected = self.static_value(node.args[0], sheet, stack)
        name, min_row, min_col, max_row, max_col = self._static_range(node.args[1], sheet)
        ws = self.book.book[name]
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                actual = self.book.merged_value(name, row, col)
                if actual is not None and getattr(ws.cell(row, col), "data_type", "") == "f":
                    raise Pending(f"MATCH 查找列包含公式，无法静态确认：{name}!{ws.cell(row, col).coordinate}")
                if self._static_equal(actual, expected):
                    return row - min_row + 1 if min_col == max_col else col - min_col + 1
        raise _StaticFormulaError(f"MATCH 在模板辅助列 {name}!{get_column_letter(min_col)}:{get_column_letter(max_col)} 未找到精确键")

    def static_value(self, node, sheet, stack=()):
        """Evaluate only template-constant helper formulas used to build INDIRECT addresses."""
        if len(stack) > 80:
            raise Pending("静态地址引用层级过深")
        if node.kind == "number":
            return Decimal(node.value)
        if node.kind == "text":
            return node.value
        if node.kind == "boolean":
            return node.value.upper() == "TRUE"
        if node.kind == "ref":
            name, address, cell = self.ref(node.value, sheet)
            key = (name, address)
            if key in self.cells or key in self.bad_cells:
                raise Pending(f"动态数字格不能参与静态地址推导：{name}!{address}")
            if key in stack:
                raise Pending(f"静态地址公式循环引用：{name}!{address}")
            if cell.data_type == "e":
                raise _StaticFormulaError(f"静态地址依赖 Excel 错误值：{name}!{address}={cell.value}")
            if cell.data_type == "f":
                parsed = self.parsed.setdefault(key, parse_excel(cell.value))
                return self.static_value(parsed, name, stack + (key,))
            return cell.value
        if node.kind == "unary":
            value = self.static_value(node.args[0], sheet, stack)
            try:
                if node.value == "+":
                    return Decimal(str(value))
                if node.value == "-":
                    return -Decimal(str(value))
            except (InvalidOperation, ValueError, TypeError) as exc:
                raise Pending("静态地址公式的一元运算不是数值") from exc
            raise Pending(f"静态地址不支持一元运算：{node.value}")
        if node.kind == "binary":
            left = self.static_value(node.args[0], sheet, stack)
            right = self.static_value(node.args[1], sheet, stack)
            if node.value == "&":
                return ("" if left is None else str(left)) + ("" if right is None else str(right))
            if node.value in {"+", "-", "*", "/"}:
                try:
                    a, b = Decimal(str(left or 0)), Decimal(str(right or 0))
                except (InvalidOperation, ValueError, TypeError) as exc:
                    raise Pending("静态地址公式的算术运算包含非数值文本") from exc
                if node.value == "+": return a + b
                if node.value == "-": return a - b
                if node.value == "*": return a * b
                if b == 0: raise _StaticFormulaError("静态地址计算出现除以零")
                return a / b
            if node.value in {"=", "<>"}:
                equal = self._static_equal(left, right)
                return equal if node.value == "=" else not equal
            if node.value in {">", "<", ">=", "<="}:
                try:
                    a, b = Decimal(str(left)), Decimal(str(right))
                except Exception:
                    a, b = str(left).casefold(), str(right).casefold()
                return {">": a > b, "<": a < b, ">=": a >= b, "<=": a <= b}[node.value]
            raise Pending(f"静态地址不支持二元运算：{node.value}")
        if node.kind == "call":
            name = node.value.upper()
            if name == "IF" and len(node.args) == 3:
                condition = self.static_value(node.args[0], sheet, stack)
                return self.static_value(node.args[1] if condition else node.args[2], sheet, stack)
            if name == "IFERROR" and len(node.args) == 2:
                try:
                    return self.static_value(node.args[0], sheet, stack)
                except _StaticFormulaError:
                    return self.static_value(node.args[1], sheet, stack)
            if name == "MATCH":
                return self._static_match(node, sheet, stack)
            if name in {"SUM", "MIN", "MAX"}:
                values = []
                for arg in node.args:
                    if arg.kind == "ref" and (":" in arg.value or re.search(r"!\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}$", arg.value)):
                        ref_sheet, min_row, min_col, max_row, max_col = self._static_range(arg, sheet)
                        for row in range(min_row, max_row + 1):
                            for col in range(min_col, max_col + 1):
                                value = self.book.merged_value(ref_sheet, row, col)
                                if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
                                    values.append(Decimal(str(value)))
                    else:
                        value = self.static_value(arg, sheet, stack)
                        if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
                            values.append(Decimal(str(value)))
                if name == "SUM": return sum(values, Decimal(0))
                if not values: return Decimal(0)
                return min(values) if name == "MIN" else max(values)
            if name == "ROW" and len(node.args) == 1 and node.args[0].kind == "ref":
                return Decimal(self.ref(node.args[0].value, sheet)[2].row)
            if name == "COLUMN" and len(node.args) == 1 and node.args[0].kind == "ref":
                return Decimal(self.ref(node.args[0].value, sheet)[2].column)
        raise Pending("无法静态求值 INDIRECT/MATCH 地址公式")

    def integer(self, node, sheet):
        if node.kind == "number" and float(node.value).is_integer():
            return int(float(node.value))
        try:
            value = self.static_value(node, sheet)
            if isinstance(value, (int, float, Decimal)) and Decimal(str(value)) == Decimal(str(value)).to_integral_value():
                return int(Decimal(str(value)))
        except Pending:
            pass
        if node.kind == "call" and node.value in {"ROW", "COLUMN"} and len(node.args) == 1 and node.args[0].kind == "ref":
            _, _, cell = self.ref(node.args[0].value, sheet)
            return cell.row if node.value == "ROW" else cell.column
        raise Pending("地址/索引参数不能静态确定")

    def range_nodes(self, node, sheet):
        name, address, external = reference_parts(node.value, sheet)
        if external or name not in self.book.book.sheetnames:
            raise Pending("不支持直接展开外部范围")
        try:
            c1, r1, c2, r2 = range_boundaries(address)
        except ValueError as exc:
            raise Pending(f"无效引用范围：{node.value}") from exc
        if not all((c1, r1, c2, r2)) or (c2-c1+1)*(r2-r1+1) > 1000:
            raise Pending("范围引用必须为不超过1000格的有限区域")
        return [(Expr("ref", f"'{name.replace(chr(39), chr(39)*2)}'!{get_column_letter(col)}{row}"), name)
                for row in range(r1, r2+1) for col in range(c1, c2+1)]

    def _full_column(self, node, default_sheet):
        if node.kind != "ref":
            raise Pending("SUMIFS 区域必须是完整列引用")
        name, address, external = reference_parts(node.value, default_sheet)
        match = re.fullmatch(r"\$?([A-Za-z]{1,3}):\$?([A-Za-z]{1,3})", address)
        if external or not match or match.group(1).upper() != match.group(2).upper():
            raise Pending("SUMIFS 仅支持同工作簿完整单列引用")
        return name, column_index_from_string(match.group(1))

    def _sumifs(self, node, sheet):
        """Lower a structurally bounded Excel SUMIFS to a sum of V_CUR paths.

        This deliberately supports only exact-text criteria over columns inside
        the template's row-header region. That is enough to faithfully compile
        the tabular summary lookups in the current personal-loan template while
        keeping dynamic/wildcard Excel criteria explicit as pending.
        """
        if len(node.args) < 3 or len(node.args[1:]) % 2:
            raise Pending("SUMIFS 参数必须由求和列及成对筛选列/条件组成")
        data_sheet, sum_col = self._full_column(node.args[0], sheet)
        criteria = []
        for column_node, value_node in zip(node.args[1::2], node.args[2::2]):
            criteria_sheet, criteria_col = self._full_column(column_node, sheet)
            if criteria_sheet != data_sheet:
                raise Pending("SUMIFS 求和列与筛选列必须位于同一工作表")
            if value_node.kind in {"text", "number", "boolean"}:
                expected = value_node.value
            elif value_node.kind == "ref":
                expected = self.static_text(value_node, sheet)
            else:
                raise Pending("SUMIFS 条件必须是固定值或可静态读取的模板单元格")
            if isinstance(expected, str) and expected[:1] in {"<", ">", "="}:
                raise Pending("SUMIFS 运算符条件尚未支持；仅能确认精确匹配条件")
            criteria.append((criteria_col, expected))

        regions, _ = numeric_regions(self.book.book)
        layouts = []
        for region in regions:
            if region.sheet != data_sheet or not (region.min_col <= sum_col <= region.max_col):
                continue
            try:
                row_header, _ = numeric_header_regions(self.book, region)
            except Exception:
                continue
            if all(row_header.min_col <= col <= row_header.max_col for col, _ in criteria):
                layouts.append((region, row_header))
        if len(layouts) != 1:
            raise Pending("SUMIFS 求和列或筛选列无法唯一匹配模板数字区域与行表头")
        region, row_header = layouts[0]

        def matches(actual, expected):
            if actual is None:
                return expected is None or str(expected).strip() == ""
            if isinstance(actual, (int, float, Decimal)) and not isinstance(actual, bool):
                try:
                    return Decimal(str(actual)) == Decimal(str(expected))
                except Exception:
                    return False
            return str(actual).strip().casefold() == str(expected).strip().casefold()

        expressions = []
        for row in range(region.min_row, region.max_row + 1):
            if self.book._anchor(data_sheet, row, sum_col) != (row, sum_col):
                continue
            if not all(matches(self.book.merged_value(data_sheet, row, col), expected)
                       for col, expected in criteria):
                continue
            address = f"{get_column_letter(sum_col)}{row}"
            key = (data_sheet, address)
            if key in self.bad_cells:
                raise Pending(f"SUMIFS 命中 {data_sheet}!{address}，但其数字区域路径不可用：{self.bad_cells[key]}")
            if key not in self.cells:
                raise Pending(f"SUMIFS 命中 {data_sheet}!{address}，但该格未由完整行列路径唯一定位")
            expressions.append(self.cells[key][0])
        if not expressions:
            # Excel SUMIFS returns zero when no row meets the exact criteria.
            return "0"
        return "SUM(" + ", ".join(expressions) + ")"

    def _linear_terms(self, node, sign=1):
        """Flatten only top-level addition/subtraction, retaining each term's sign."""
        if node.kind == "binary" and node.value in {"+", "-"}:
            yield from self._linear_terms(node.args[0], sign)
            yield from self._linear_terms(node.args[1], sign if node.value == "+" else -sign)
        elif node.kind == "unary" and node.value in {"+", "-"}:
            yield from self._linear_terms(node.args[0], sign if node.value == "+" else -sign)
        else:
            yield node, sign

    def _numeric_cells(self, node, sheet, stack=()):
        """Find V_CUR-backed numeric cells behind a formula term, including helpers."""
        if len(stack) > 80:
            raise Pending("公式依赖过深")
        if node.kind == "ref":
            name, address, cell = self.ref(node.value, sheet)
            key = (name, address)
            if key in self.bad_cells:
                raise Pending(f"{name}!{address}：{self.bad_cells[key]}")
            if key in self.cells:
                return [key]
            if key in stack:
                raise Pending(f"公式循环引用：{name}!{address}")
            if cell.data_type == "f":
                expanded, name = self.expand(node, sheet)
                return self._numeric_cells(expanded, name, stack + (key,))
            return []
        if node.kind in {"binary", "unary"}:
            found = []
            for arg in node.args:
                found.extend(self._numeric_cells(arg, sheet, stack))
            return found
        if node.kind == "call":
            if node.value == "INDIRECT" and len(node.args) == 1:
                address = self.static_text(node.args[0], sheet)
                return self._numeric_cells(Expr("ref", address), sheet, stack)
            found = []
            for arg in node.args:
                if arg.kind == "ref" and ":" in arg.value and node.value in {"SUM", "MIN", "MAX"}:
                    for ref, name in self.range_nodes(arg, sheet):
                        found.extend(self._numeric_cells(ref, name, stack))
                else:
                    found.extend(self._numeric_cells(arg, sheet, stack))
            return found
        return []

    def _formula_signature(self, node, sheet, stack=()):
        """Normalize helper references and singleton SUM wrappers for parity checks."""
        if len(stack) > 80:
            raise Pending("公式签名展开过深")
        if node.kind == "ref":
            name, address, cell = self.ref(node.value, sheet)
            key = (name, address)
            if key in self.cells:
                return ("value", name, address)
            if key in self.bad_cells:
                raise Pending(f"{name}!{address}：{self.bad_cells[key]}")
            if key in stack:
                raise Pending(f"公式循环引用：{name}!{address}")
            if cell.data_type == "f":
                expanded, name = self.expand(node, sheet)
                return self._formula_signature(expanded, name, stack + (key,))
            return ("literal", cell.data_type, cell.value)
        if node.kind == "call" and node.value == "SUM" and len(node.args) == 1 and node.args[0].kind == "ref" and ":" in node.args[0].value:
            parts = []
            for ref, name in self.range_nodes(node.args[0], sheet):
                parts.append(self._formula_signature(ref, name, stack))
            return ("sum", tuple(parts))
        if node.kind == "call" and node.value == "SUM" and len(node.args) == 1 and node.args[0].kind != "ref":
            return self._formula_signature(node.args[0], sheet, stack)
        if node.kind == "call" and node.value == "SUM" and len(node.args) == 1 and node.args[0].kind == "ref":
            if ":" not in node.args[0].value:
                return self._formula_signature(node.args[0], sheet, stack)
        if node.kind == "call" and node.value == "INDIRECT" and len(node.args) == 1:
            address = self.static_text(node.args[0], sheet)
            return self._formula_signature(Expr("ref", address), sheet, stack)
        if node.kind in {"number", "text", "boolean"}:
            return (node.kind, node.value)
        if node.kind in {"binary", "unary", "call"}:
            return (node.kind, node.value, tuple(self._formula_signature(arg, sheet, stack) for arg in node.args))
        return (node.kind, node.value)

    def _linear_signature(self, node, sheet):
        terms = []
        for atom, sign in self._linear_terms(node):
            terms.append((sign, self._formula_signature(atom, sheet)))
        return tuple(terms)

    def _cross_period_shape(self, node, sheet):
        """Recognize a balance equation from explicit current/prior column paths."""
        terms = list(self._linear_terms(node))
        if len(terms) < 3 or terms[0][1] != 1:
            return None
        cells_by_term = []
        paths_by_term = []
        for atom, _sign in terms:
            cells = list(dict.fromkeys(self._numeric_cells(atom, sheet)))
            cells_by_term.append(cells)
            paths_by_term.append([self.cell_paths.get(key, ((), ())) for key in cells])
        current_paths = [column for _row, column in paths_by_term[0]]
        if not current_paths or not any(any(token in segment for segment in column for token in ("本期", "当前"))
                                        for column in current_paths):
            return None
        prior_indexes = [index for index, ((_, sign), paths) in enumerate(zip(terms, paths_by_term))
                         if sign == -1 and any("上期" in segment for _row, column in paths for segment in column)]
        if prior_indexes != [1]:
            return None
        occurrence_terms, settlement_terms = [], []
        for (atom, sign), cells in zip(terms[2:], cells_by_term[2:]):
            if not cells:
                return None
            (occurrence_terms if sign == -1 else settlement_terms).append(atom)
        if not occurrence_terms:
            return None
        return terms[0][0], terms[1][0], occurrence_terms, settlement_terms

    @staticmethod
    def _sum_terms(nodes):
        if not nodes:
            return None
        result = nodes[0]
        for node in nodes[1:]:
            result = Expr("binary", "+", (result, node))
        return result

    def _sum_nonzero_reference(self, node, sheet):
        """Return the single cell tested by SUM(ref) <> 0, if structurally exact."""
        if node.kind != "binary" or node.value != "<>":
            return None
        left, right = node.args
        if right.kind != "number" or Decimal(right.value) != 0:
            return None
        if left.kind != "call" or left.value != "SUM" or len(left.args) != 1:
            return None
        reference = left.args[0]
        if reference.kind != "ref" or ":" in reference.value:
            return None
        name, address, _ = self.ref(reference.value, sheet)
        return Expr("ref", f"'{name.replace(chr(39), chr(39) * 2)}'!{address}"), name

    def _summary_reconciliation(self, guards, sheet):
        """Recognize a total = sum(parts) guard wrapped in two SUM(ref) checks.

        Some templates put the actual difference in a helper formula and use a
        blank helper cell plus that difference in an OR(... <> 0) output. The
        blank cell is Excel's zero-valued placeholder; the meaningful rule is
        the helper's total-minus-parts equation. Only accept the exact shape
        when every numeric operand has a unique template path.
        """
        if len(guards) != 1 or guards[0].kind != "call" or guards[0].value != "OR":
            return None
        references = [self._sum_nonzero_reference(item, sheet) for item in guards[0].args]
        if len(references) != 2 or any(item is None for item in references):
            return None
        blank, helper = [], []
        for node, ref_sheet in references:
            name, address, cell = self.ref(node.value, ref_sheet)
            if cell.data_type == "f":
                helper.append((node, ref_sheet))
            elif cell.value is None:
                blank.append((node, ref_sheet))
        if len(blank) != 1 or len(helper) != 1:
            return None
        helper_node, helper_sheet = helper[0]
        expanded, expanded_sheet = self.expand(helper_node, helper_sheet)
        terms = list(self._linear_terms(expanded))
        if len(terms) != 3 or [sign for _, sign in terms] != [1, -1, -1]:
            return None
        if any(node.kind != "ref" for node, _ in terms):
            return None
        keys = []
        for node, _ in terms:
            cells = list(dict.fromkeys(self._numeric_cells(node, expanded_sheet)))
            if len(cells) != 1:
                return None
            keys.append(cells[0])
        return tuple(terms), expanded_sheet

    @staticmethod
    def _node_for_key(key):
        name, address = key
        escaped = name.replace("'", "''")
        return Expr("ref", f"'{escaped}'!{address}")

    def bind(self, name, node, sheet):
        expression = self.emit(node, sheet, bind=False)
        self.dictionary[name] = expression
        self.expression_aliases.setdefault(expression, name)
        if node.kind == "ref":
            n, address, _ = self.ref(node.value, sheet)
            self.aliases[(n, address)] = name
        return "{" + name + "}"

    def emit(self, node, sheet, *, bind=True, stack=()):
        if len(stack) > 80:
            raise Pending("公式依赖过深")
        if node.kind in {"number", "boolean"}:
            return node.value.upper()
        if node.kind == "text":
            return json.dumps(node.value, ensure_ascii=False)
        if node.kind == "unary":
            return node.value + "(" + self.emit(node.args[0], sheet, bind=bind, stack=stack) + ")"
        if node.kind == "ref":
            name, address, cell = self.ref(node.value, sheet)
            key = (name, address)
            if bind and key in self.aliases:
                return "{" + self.aliases[key] + "}"
            if key in stack:
                raise Pending(f"公式循环引用：{name}!{address}")
            if key in self.bad_cells:
                raise Pending(f"{name}!{address}：{self.bad_cells[key]}")
            if key in self.cells:
                expression = self.cells[key][0]
                if bind:
                    alias = self.expression_aliases.get(expression)
                    if not alias:
                        alias = "条件值" + str(1 + sum(key.startswith("条件值") for key in self.dictionary))
                        self.dictionary[alias] = expression
                        self.expression_aliases[expression] = alias
                    return "{" + alias + "}"
                return expression
            if cell.data_type == "f":
                expanded, name = self.expand(node, sheet)
                return self.emit(expanded, name, bind=bind, stack=stack + (key,))
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                raise Pending(f"{name}!{address}是模板中的固定数值，未由数字区域定位；为避免把模板示例值写死进配置，请补齐表结构/数字区域后重新生成")
            raise Pending(f"{name}!{address}不在数字区域，不能作为数值取数")
        if node.kind == "binary":
            left, right = node.args
            if node.value == "&":
                raise Pending("数值表达式中包含文本拼接")
            if node.value in {"=", "<>"} and right.kind == "text" and right.value == "":
                expression = "ISBLANK(" + self.emit(left, sheet, bind=bind, stack=stack) + ")"
                return "NOT(" + expression + ")" if node.value == "<>" else expression
            if node.value in {"*", "/"} and left.kind == "call" and left.value in {"INDEX", "VLOOKUP"}:
                return self.external_call(left, sheet, node.value, right)
            if node.value == "*" and right.kind == "call" and right.value in {"INDEX", "VLOOKUP"}:
                return self.external_call(right, sheet, node.value, left)
            return "(" + self.emit(left, sheet, bind=bind, stack=stack) + " " + node.value + " " + self.emit(right, sheet, bind=bind, stack=stack) + ")"
        if node.kind == "call":
            if node.value in {"AND", "OR", "NOT"}:
                return self.condition(node, sheet, bind=bind, stack=stack)
            if node.value == "SUMIFS":
                return self._sumifs(node, sheet)
            if node.value == "INDIRECT" and len(node.args) == 1:
                reference = self.static_text(node.args[0], sheet)
                return self.emit(Expr("ref", reference), sheet, bind=bind, stack=stack)
            if node.value in {"INDEX", "VLOOKUP"}:
                return self.external_call(node, sheet)
            if node.value == "TEXT" and len(node.args) == 2 and node.args[1].kind == "text":
                fmt = node.args[1].value
                if not re.fullmatch(r"0(?:\.0+)?", fmt):
                    raise Pending(f"TEXT格式尚未支持：{fmt}")
                digits = len(fmt.split(".")[1]) if "." in fmt else 0
                return f"ROUND({self.emit(node.args[0], sheet, bind=bind, stack=stack)}, {digits})"
            if node.value not in {"SUM", "ABS", "ROUND", "MIN", "MAX", "IF", "AND", "OR", "NOT", "ISBLANK", "ISNUMBER"}:
                raise Pending(f"暂不支持函数：{node.value}")
            arguments = []
            for arg in node.args:
                if arg.kind == "ref" and ":" in arg.value:
                    if node.value not in {"SUM", "MIN", "MAX"}:
                        raise Pending(f"{node.value}中的范围参数需人工处理")
                    arguments.extend(self.emit(item, name, bind=bind, stack=stack) for item, name in self.range_nodes(arg, sheet))
                else:
                    arguments.append(self.emit(arg, sheet, bind=bind, stack=stack))
            if node.value in {"AND", "OR"}:
                return "(" + (" " + node.value + " ").join(arguments) + ")"
            return node.value + "(" + ", ".join(arguments) + ")"
        raise Pending("暂不支持该表达式")

    def condition(self, node, sheet, *, bind=True, stack=()):
        if node.kind == "call" and node.value in {"AND", "OR"}:
            return "(" + (" " + node.value + " ").join(self.condition(arg, sheet, bind=bind, stack=stack) for arg in node.args) + ")"
        if node.kind == "call" and node.value == "NOT" and len(node.args) == 1:
            return "NOT(" + self.condition(node.args[0], sheet, bind=bind, stack=stack) + ")"
        if (node.kind == "call" and node.value == "SUM" and len(node.args) == 1
                and node.args[0].kind == "binary" and node.args[0].value in {"=", "<>", ">", "<", ">=", "<="}):
            # Excel converts SUM(single boolean) to 0/1; in IF this is the same predicate.
            return self.emit(node.args[0], sheet, bind=bind, stack=stack)
        return self.emit(node, sheet, bind=bind, stack=stack)

    def identity_lookup(self, node, sheet):
        if node.kind != "call" or node.value != "VLOOKUP" or len(node.args) != 4 or node.args[3].value.upper() not in {"0", "FALSE"}:
            raise Pending("外部机构匹配不是可确认的精确VLOOKUP")
        if node.args[0].kind != "ref" or node.args[1].kind != "ref":
            raise Pending("外部机构匹配字段不是代码到报表项目的已知结构")
        name, _, cell = self.ref(node.args[0].value, sheet)
        labels = [str(self.book.merged_value(name, cell.row, col) or "") for col in range(max(1, cell.column-3), cell.column)]
        if not any("信用代码" in label for label in labels):
            raise Pending("未确认外部查找使用当前机构信用代码")
        mapping_sheet, area, _ = reference_parts(node.args[1].value, sheet)
        index = self.integer(node.args[2], sheet)
        legacy = mapping_sheet == "参照表" and area.upper() == "A:B" and index == 2
        columns = range_boundaries(area)
        current = (mapping_sheet == self.layout.mapping_sheet and columns[0] == self.layout.identity_column
                   and columns[0] + index - 1 == self.layout.target_column and columns[2] >= self.layout.target_column)
        if not legacy and not current:
            raise Pending("外部机构映射范围不是已知的代码/报表项目两列结构")

    def external_call(self, node, sheet, operator="", factor=None):
        if self.external is None:
            raise Pending("包含外部查找，请选择外部文件并确认金额单位")
        if self.layout is None:
            self.layout = discover_external_layout(self.external, support_time_series=True)
        if node.value == "INDEX" and len(node.args) == 3:
            array, rows, cols = node.args
            if array.kind != "ref" or rows.kind != "call" or rows.value != "MATCH" or cols.kind != "call" or cols.value != "MATCH":
                raise Pending("INDEX不是已知的机构行/指标列精确匹配结构")
            if len(rows.args) != 3 or len(cols.args) != 3 or rows.args[2].value != "0" or cols.args[2].value != "0":
                raise Pending("INDEX/MATCH必须使用精确匹配")
            self.identity_lookup(rows.args[0], sheet)
            metric = self.static_text(cols.args[0], sheet)
            data_sheet, area, _ = reference_parts(array.value, sheet)
            row_sheet, row_area, _ = reference_parts(rows.args[1].value, sheet)
            col_sheet, col_area, _ = reference_parts(cols.args[1].value, sheet)
            columns = [column for column in self.layout.columns if normalized(column.header) == normalized(metric) and column.sheet == data_sheet]
            if len(columns) != 1:
                raise Pending("外部指标名无法唯一匹配：" + metric)
            column = columns[0]
            if (row_sheet != data_sheet or col_sheet != data_sheet or area != "1:1048576"
                    or row_area != f"{get_column_letter(column.key_column)}:{get_column_letter(column.key_column)}"
                    or col_area != f"{column.header_row}:{column.header_row}"):
                raise Pending("INDEX/MATCH的机构列或指标表头行与外部文件不一致")
        elif node.value == "VLOOKUP" and len(node.args) == 4 and node.args[3].value.upper() in {"0", "FALSE"}:
            self.identity_lookup(node.args[0], sheet)
            if node.args[1].kind != "ref":
                raise Pending("外部VLOOKUP缺少数据范围")
            data_sheet, area, _ = reference_parts(node.args[1].value, sheet)
            c1, _, c2, _ = range_boundaries(area)
            target = c1 + self.integer(node.args[2], sheet) - 1
            columns = [column for column in self.layout.columns if column.sheet == data_sheet and column.column == target and column.key_column == c1]
            if len(columns) != 1 or target > c2:
                raise Pending("外部VLOOKUP列号无法唯一匹配指标")
            column = columns[0]
        else:
            raise Pending("不支持该外部查找结构")
        metric = column.header
        unit_match = re.search(r"[（(]\s*(元|万元|亿元|户|笔|百分比|%)\s*[）)]", metric)
        indicated_unit = unit_match.group(1) if unit_match else ""
        amount = _metric_profile((metric,), source_unit=indicated_unit)[0] == "amount"
        unit = (indicated_unit or self.external_unit) if amount else ""
        scale = {"元": Decimal("0.0001"), "万元": Decimal(1), "亿元": Decimal(10000)}.get(unit, Decimal(1))
        if factor is not None and factor.kind != "number":
            raise Pending("外部换算倍数不是固定数字")
        applied = Decimal(factor.value) if factor is not None else Decimal(1)
        applied = Decimal(1) / applied if operator == "/" else applied
        if applied != scale:
            raise Pending("原外部公式的换算倍数与所选外部金额单位不一致")
        arguments = [json.dumps(metric, ensure_ascii=False), "sheet=" + json.dumps(column.sheet, ensure_ascii=False)]
        if unit:
            arguments.append("unit=" + json.dumps(unit, ensure_ascii=False))
        return "V_EXT(" + ", ".join(arguments) + ")"

    def description(self, node, sheet, stack=()):
        if len(stack) > 60:
            raise Pending("描述公式引用层级过深")
        if node.kind == "call" and node.value == "IF" and self.kind(node, sheet, stack) == "mixed":
            # Excel coerces mixed IF branch values to text when used under '&'.
            # The audit trigger remains fully expressible; a display-only branch
            # is safer to flag for review than to emit an invalid typed DSL IF.
            raise Pending("描述文字含有IF文字/数值混合分支，需人工核对展示内容")
        key = None
        if node.kind == "ref":
            name, address, _ = self.ref(node.value, sheet)
            key = (name, address)
            if key in stack:
                raise Pending(f"描述公式循环引用：{name}!{address}")
        try:
            return self.static_text(node, sheet).replace("{", "{{").replace("}", "}}")
        except Pending:
            pass
        expanded, new_sheet = self.expand(node, sheet)
        if expanded != node:
            return self.description(expanded, new_sheet, stack + (key,) if key else stack)
        if node.kind == "binary" and node.value == "&":
            return "".join(self.description(item, sheet, stack) for item in node.args)
        expression = self.emit(node, sheet)
        if re.fullmatch(r"\{[\w\u3400-\u9fff]+\}", expression):
            return expression
        alias = "描述值" + str(1 + sum(key.startswith("描述值") for key in self.dictionary))
        self.dictionary[alias] = expression
        return "{" + alias + "}"

    def rule(self, sheet, address, guards, payload):
        self.reset()
        severity, parts = payload
        if not guards:
            raise Pending("校验输出没有可确认的IF判断条件")
        text_only_value = None
        if len(parts) == 2:
            # Compact ticket formulas omit a separate metric segment and put
            # the difference in the message itself. Accept only one numeric
            # expression, so text-only or ambiguous messages remain pending.
            dynamic_parts = [part for part in concat_parts(parts[1]) if part.kind != "text"]
            if len(dynamic_parts) == 1 and self.kind(dynamic_parts[0], sheet) == "number":
                description_node = parts[1]
                display_nodes = dynamic_parts
            elif not dynamic_parts:
                referenced_cells = list(dict.fromkeys(self._numeric_cells(guards[0], sheet)))
                if len(referenced_cells) != 1:
                    raise Pending("两段式文本校验输出未引用唯一、可定位的数值指标")
                text_only_value = referenced_cells[0]
                description_node = parts[1]
                display_nodes = []
            else:
                raise Pending("两段式校验输出未包含唯一、可确定的数值表达式")
        else:
            description_node = parts[2]
            display_nodes = parts[3:]
        if len(parts) == 3 and self.kind(parts[2], sheet) == "number":
            # Compact severity|metric|<numeric value> output with an empty
            # description segment. Treat the numeric tail as the current value.
            description_node, display_nodes = Expr("text", ""), [parts[2]]
        elif len(parts) == 4 and self.kind(parts[2], sheet) == "number" and self.kind(parts[3], sheet) == "text":
            description_node, display_nodes = parts[3], [parts[2]]
        description_hint = ""
        try:
            description_hint = self.static_text(description_node, sheet)
        except Pending:
            pass
        header = ""
        for row in range(self.book.book[sheet][address].row-1, 0, -1):
            value = self.book.merged_value(sheet, row, self.book.book[sheet][address].column)
            if isinstance(value, str) and not value.startswith("=") and value.strip():
                header = value.strip()
                break
        classification = "不应有数" if any(token in header + description_hint for token in ("不应有数", "不应出现", "不应有")) else "条件核对"
        current = comparison = difference = ""
        reconciliation = self._summary_reconciliation(guards, sheet)
        if text_only_value:
            current = self.bind("本期值", self._node_for_key(text_only_value), text_only_value[0])
            comparison = "0"
            difference = "{本期值} - 0"
        if reconciliation:
            terms, reconciliation_sheet = reconciliation
            current = self.bind("本期合计", terms[0][0], reconciliation_sheet)
            component1 = self.bind("分项值1", terms[1][0], reconciliation_sheet)
            component2 = self.bind("分项值2", terms[2][0], reconciliation_sheet)
            self.dictionary["对比值"] = f"({component1} + {component2})"
            self.dictionary["差值"] = "{本期合计} - {对比值}"
            current, comparison, difference = "{本期合计}", "{对比值}", "{差值}"
            display_nodes = []
            description_node = parts[1]
        if len(parts) >= 6 and "跨期" in header:
            expanded, diff_sheet = self.expand(parts[5], sheet)
            if expanded.kind == "binary" and expanded.value == "-":
                left, right = expanded.args
                full = right.kind == "binary" and right.value == "-" and right.args[0].kind == "binary" and right.args[0].value == "+"
                partial = right.kind == "binary" and right.value == "+"
                if full or partial:
                    classification = "完整跨期" if full else "缺少本期结清的跨期"
                    current = self.bind("本期余额", left, diff_sheet)
                    prior, occurred = right.args[0].args if full else right.args
                    self.bind("上期余额", prior, diff_sheet)
                    self.bind("本期发生", occurred, diff_sheet)
                    comparison = "({上期余额} + {本期发生})"
                    if full:
                        self.bind("本期结清", right.args[1], diff_sheet)
                        comparison = "({上期余额} + {本期发生} - {本期结清})"
                    self.dictionary["跨期差额"] = "{本期余额} - " + comparison
                    difference = "{跨期差额}"
                    if parts[5].kind == "ref":
                        name, cell, _ = self.ref(parts[5].value, sheet)
                        self.aliases[(name, cell)] = "跨期差额"
        extra_descriptions = []
        compact_cross_period = False
        if not current and len(parts) == 2 and display_nodes:
            message_delta = display_nodes[0]
            for guard in guards:
                if (guard.kind != "binary" or guard.value != "<>" or len(guard.args) != 2
                        or guard.args[1].kind != "number" or Decimal(guard.args[1].value) != 0):
                    continue
                try:
                    if self._linear_signature(guard.args[0], sheet) != self._linear_signature(message_delta, sheet):
                        continue
                    shape = self._cross_period_shape(message_delta, sheet)
                except Pending:
                    raise
                if not shape:
                    continue
                current_node, prior_node, occurrence_nodes, settlement_nodes = shape
                current = self.bind("本期余额", current_node, sheet)
                self.bind("上期余额", prior_node, sheet)
                self.bind("本期发生", self._sum_terms(occurrence_nodes), sheet)
                comparison_expression = "({上期余额} + {本期发生})"
                if settlement_nodes:
                    self.bind("本期结清", self._sum_terms(settlement_nodes), sheet)
                    comparison_expression = "({上期余额} + {本期发生} - {本期结清})"
                    classification = "完整跨期"
                else:
                    classification = "缺少本期结清的跨期"
                self.dictionary["对比值"] = comparison_expression
                self.bind("跨期差额", message_delta, sheet)
                comparison = "{对比值}"
                difference = "{跨期差额}"
                compact_cross_period = True
                break
        if not current and display_nodes:
            if self.kind(display_nodes[0], sheet) == "text":
                extra_descriptions.append(display_nodes[0])
            else:
                current = self.bind("本期值", display_nodes[0], sheet)
        if classification not in {"不应有数", "完整跨期", "缺少本期结清的跨期"}:
            if len(display_nodes) > 1:
                if self.kind(display_nodes[1], sheet) == "text":
                    extra_descriptions.append(display_nodes[1])
                else:
                    comparison = self.bind("对比值", display_nodes[1], sheet)
            if len(display_nodes) > 2:
                if self.kind(display_nodes[2], sheet) == "text":
                    extra_descriptions.append(display_nodes[2])
                else:
                    difference = self.bind("差值", display_nodes[2], sheet)
        # Preserve the original trigger, including nested IF short-circuit guards.
        trigger = ("ABS({差值}) > 0" if reconciliation else
                   self.condition(guards[0], sheet) if len(guards) == 1 else
                   "(" + " AND ".join(self.condition(guard, sheet) for guard in guards) + ")")
        audit = trigger
        if text_only_value:
            audit = self.condition(guards[0], sheet)
        if compact_cross_period:
            description = "差值为{跨期差额}"
        else:
            try:
                description = self.description(description_node, sheet)
                for extra in extra_descriptions:
                    text = self.description(extra, sheet)
                    if text.strip():
                        description += "；" + text
            except Pending as exc:
                description = (header or classification) + "，请核实。"
                for caption, expression in (("本期值", current), ("对比值", comparison), ("差值", difference)):
                    if expression:
                        alias = "展示" + caption
                        self.dictionary[alias] = expression
                        description += caption + "={" + alias + "}；"
                self.notes.append("原描述未完整转换，已采用校验标题和取值占位；原描述逻辑请人工核对：" + str(exc))
        if reconciliation:
            message = self.static_text(parts[1], sheet)
            description = message + "；本期值={本期合计}；对比值={对比值}；差值={差值}"
        if not description.strip():
            description = header or "模板校验，请核实。"
        if len(guards) > 1:
            safety_guard = "(" + " AND ".join(self.condition(guard, sheet) for guard in guards[:-1]) + ")"
            guard_variables = variables_in(parse_expression(safety_guard))
            for alias, expression in list(self.dictionary.items()):
                if _has_division(parse_expression(expression)):
                    if alias in guard_variables:
                        raise Pending("保护条件本身依赖除法输出值，需要人工确认")
                    self.dictionary[alias] = f"IF({safety_guard}, {expression}, 0)"
        self.dictionary, normalized = _hoist_dictionary_wrappers(
            self.dictionary,
            {"审核表达式": audit, "当前值表达式": current, "对比值表达式": comparison,
             "差值表达式": difference, "规则描述": description},
        )
        audit = normalized["审核表达式"]
        current = normalized["当前值表达式"]
        comparison = normalized["对比值表达式"]
        difference = normalized["差值表达式"]
        description = normalized["规则描述"]
        try:
            metric = self.static_text(parts[1], sheet)
        except Pending:
            metric = " / ".join(dict.fromkeys(self.cells[key][1] for key in self.aliases if key in self.cells)) or address
        metric = metric.strip() or address
        metric = _current_value_path(current, self.dictionary, sheet) or metric
        dictionary = "{" + ",\n ".join(json.dumps(key, ensure_ascii=False) + ": " + value for key, value in self.dictionary.items()) + "}"
        nodes = parse_dictionary(dictionary)
        validate_dictionary_dependencies(nodes)
        for node in nodes.values():
            validate_static_types(node)
        for expression in (audit, current, comparison, difference):
            if expression:
                node = parse_expression(expression)
                validate_static_types(node)
                if variables_in(node) - nodes.keys():
                    raise Pending("生成表达式中存在未定义变量")
        if _description_variables(description) - nodes.keys():
            raise Pending("生成描述中存在未定义变量")
        if any(len(text) > 32767 for text in (audit, current, comparison, difference, dictionary, description, metric)):
            raise Pending("生成字段超过Excel单元格32767字符限制，需要人工拆分")
        return {"规则名称": header or classification, "审核级别": severity, "校验类型": classification, "校验指标": metric,
                "审核表达式": audit, "表达式布尔值取反": "是", "规则描述": description, "当前值表达式": current, "对比值表达式": comparison,
                "差值表达式": difference, "取值字典": dictionary, "数值运算模式": "Excel双精度"}


def _family(formula, origin):
    row = int(re.search(r"\d+", origin).group())
    col = column_index_from_string(re.match(r"[A-Za-z]+", origin).group())
    def replace(match):
        absolute_col, letters, absolute_row, digits = match.groups()
        c, r = column_index_from_string(letters), int(digits)
        return (f"C{c}" if absolute_col else f"C[{c-col}]") + (f"R{r}" if absolute_row else f"R[{r-row}]")
    try:
        tokens = Tokenizer(formula).items
        pieces = []
        for token in tokens:
            value = token.value
            if token.type == "OPERAND" and token.subtype == "RANGE":
                prefix, separator, address = value.rpartition("!")
                value = (prefix + separator if separator else "") + re.sub(r"(\$?)([A-Za-z]{1,3})(\$?)([1-9]\d*)", replace, address if separator else value)
            pieces.append(value)
        pattern = "".join(pieces)
    except Exception:
        pattern = formula
    return "F" + hashlib.sha256(pattern.encode("utf-8")).hexdigest()[:10], pattern


def _business_name(path):
    name = re.sub(r"^[！!\d.\s]+", "", path.stem)
    name = re.sub(r"\d{6,8}$", "", name)
    match = re.match(r"(?:金融)?基础数据采集[（(](.+)[）)]$", name)
    if match:
        return match.group(1)
    return re.sub(r"^(?:金融|融)?基础数据[-－]?", "", name) or path.stem


def _format(sheet, widths):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = "s"


def export_formula_candidates(files, *, output_root=None, business="", external=None, external_unit="万元"):
    paths = list(dict.fromkeys(Path(path).expanduser().resolve() for path in files))
    if not paths or len(paths) > 200:
        raise ValueError("请选择1至200个模板文件")
    output_root = Path(output_root or ROOT / "data" / "s5" / "维护输出" / "配置候选").expanduser().resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    if external_unit not in {"元", "万元", "亿元"}:
        raise ValueError("外部金额单位应为元、万元或亿元")
    directory = output_root / ("candidates_" + uuid.uuid4().hex)
    directory.mkdir(parents=True, exist_ok=False)
    external_book = _Workbook(SourceBinding("EXT", Path(external).resolve())) if external else None
    rows, details, summaries, copies, path_checks, formula_inventory = [], [], [], [], [], []
    rule_sequences = Counter()
    watched = {external_book.path: external_book.sha256} if external_book else {}
    try:
        for path in paths:
            business_name = business or _business_name(path)
            counts = Counter(formulas=0, outputs=0, candidates=0, pending=0, intermediate=0)
            book = None
            try:
                book = _Workbook(SourceBinding("TEMPLATE", path), editable=True)
                watched[path] = book.sha256
                converter = Converter(book, external_book, external_unit)
                if not converter.regions:
                    raise Pending("模板没有校验区域命名，未扫描整张工作表")
                seen, pending_cells = set(), {}
                for region in converter.regions:
                    for line in book.book[region.sheet].iter_rows(min_row=region.min_row, max_row=region.max_row, min_col=region.min_col, max_col=region.max_col):
                        for cell in line:
                            key = (region.sheet, cell.coordinate)
                            if key in seen or cell.data_type != "f":
                                continue
                            seen.add(key)
                            counts["formulas"] += 1
                            family, pattern = _family(cell.value, cell.coordinate)
                            is_output = "|" in cell.value
                            formula_inventory.append({
                                "template": path.name,
                                "business": business_name,
                                "sheet": region.sheet,
                                "cell": cell.coordinate,
                                "role": "校验输出公式" if is_output else "中间计算公式",
                                "family": family,
                                "formula": cell.value,
                                "rule_ids": "",
                                "classification": "",
                                "status": "待处理" if is_output else "中间公式（不生成独立审核规则）",
                                "reason": "" if is_output else "公式未包含审核结果分隔符“|”，作为模板中间计算公式保留，不单独生成审核规则。",
                            })
                            # Validation outputs use the established error|metric|description protocol.
                            static_errors = formula_error_tokens(cell.value)
                            if static_errors:
                                counts["pending"] += 1
                                error_reason = ("校验区域公式包含 Excel 错误引用 "
                                                + "、".join(static_errors)
                                                + "；该单元格不会生成可启用配置，请先修复模板公式")
                                inventory = formula_inventory[-1]
                                inventory.update(
                                    family="模板公式错误",
                                    status="模板公式错误（待修复）",
                                    reason=error_reason,
                                )
                                if is_output:
                                    counts["outputs"] += 1
                                base = {"template": str(path), "sheet": region.sheet, "cell": cell.coordinate,
                                        "formula": cell.value, "family": "模板公式错误", "pattern": pattern,
                                        "rule_id": "", "classification": "", "status": "模板公式错误（待修复）",
                                        "reason": error_reason}
                                details.append(base)
                                pending_cells[key] = error_reason
                                continue
                            if not is_output:
                                counts["intermediate"] += 1
                                continue
                            counts["outputs"] += 1
                            base = {"template": str(path), "sheet": region.sheet, "cell": cell.coordinate, "formula": cell.value, "family": family,
                                    "pattern": pattern, "rule_id": "", "classification": "", "status": "待处理", "reason": ""}
                            try:
                                branches = list(error_branches(parse_excel(cell.value)))
                                if not branches:
                                    supplemental = _ai_supplemental_candidate(
                                        path.name, business_name, region.sheet, cell.coordinate,
                                        cell.value, rule_sequences,
                                    )
                                    if supplemental is None:
                                        raise Pending("未识别出可确认的IF校验输出分支")
                                    dictionary_nodes = parse_dictionary(supplemental["取值字典"])
                                    validate_dictionary_dependencies(dictionary_nodes)
                                    for node in dictionary_nodes.values():
                                        validate_static_types(node)
                                    for field in ("审核表达式", "当前值表达式", "对比值表达式", "差值表达式"):
                                        expression = supplemental.get(field)
                                        if expression:
                                            node = parse_expression(expression)
                                            validate_static_types(node)
                                            if variables_in(node) - dictionary_nodes.keys():
                                                raise Pending(f"AI补充候选的{field}引用未定义变量")
                                    rows.append(supplemental)
                                    counts["candidates"] += 1
                                    base.update(
                                        rule_id=supplemental["规则编号"],
                                        classification=supplemental["校验类型"],
                                        status="AI补充候选（未启用，待人工确认）",
                                        reason=supplemental["规则描述"],
                                    )
                                    details.append(base)
                                    pending_cells[key] = supplemental["规则描述"]
                                    continue
                                for guards, payload in branches:
                                    entry = dict(base)
                                    try:
                                        candidate = converter.rule(region.sheet, cell.coordinate, guards, payload)
                                        sequence_key = (business_name, region.sheet)
                                        rule_sequences[sequence_key] += 1
                                        rule_id = f"{business_name}_{region.sheet}_{rule_sequences[sequence_key]}"
                                        row = {"规则编号": rule_id, "是否启用": "否", "业务名称": business_name,
                                               "工作表名称": region.sheet, "校验公式单元格": cell.coordinate, **candidate}
                                        rows.append(row)
                                        path_checks.extend(_template_path_checks(
                                            candidate, rule_id, business_name, region.sheet, cell.coordinate, converter
                                        ))
                                        counts["candidates"] += 1
                                        entry.update(rule_id=rule_id, classification=candidate["校验类型"], status="已生成候选（未启用）", reason="按原公式转换；请审核口径及单位后复制到正式配置并启用" + ("；" + "；".join(converter.notes) if converter.notes else ""))
                                    except Exception as exc:
                                        counts["pending"] += 1
                                        entry["reason"] = str(exc)
                                        pending_cells[key] = str(exc)
                                    details.append(entry)
                            except Exception as exc:
                                counts["pending"] += 1
                                base["reason"] = str(exc)
                                details.append(base)
                                pending_cells[key] = str(exc)
                copy_dir = directory / "模板标注副本"
                copy_dir.mkdir(exist_ok=True)
                target = copy_dir / f"{len(copies)+1:03d}_{path.name}"
                for (sheet, address), reason in pending_cells.items():
                    cell = book.book[sheet][address]
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
                    old_comment = cell.comment.text + "\n\n" if cell.comment else ""
                    cell.comment = Comment(old_comment + "配置候选转换待处理：" + reason, "审核工具-S5")
                book.book.save(target)
                copies.append({"path": str(target), "name": target.name})
            except Exception as exc:
                counts["pending"] += 1
                details.append({"template": str(path), "sheet": "", "cell": "", "formula": "", "family": "文件级", "pattern": "", "rule_id": "", "classification": "", "status": "待处理", "reason": str(exc)})
            finally:
                if book:
                    book.close()
            summaries.append({"template": str(path), **dict(counts)})
        if any(file_sha256(path) != digest for path, digest in watched.items()):
            raise ValueError("生成期间模板或外部文件发生变化，候选不可用，请重新生成")
        detail_by_formula = defaultdict(list)
        for entry in details:
            detail_by_formula[(Path(entry["template"]).name, entry.get("sheet", ""), entry.get("cell", ""))].append(entry)
        candidate_statuses = {"已生成候选（未启用）", "AI补充候选（未启用，待人工确认）"}
        for inventory in formula_inventory:
            if inventory["role"] != "校验输出公式" or inventory["status"] == "模板公式错误（待修复）":
                continue
            entries = detail_by_formula.get((inventory["template"], inventory["sheet"], inventory["cell"]), [])
            if not entries:
                inventory["status"] = "待处理"
                inventory["reason"] = "校验输出公式没有对应的转换明细记录。"
                continue
            statuses = {entry.get("status", "") for entry in entries}
            inventory["rule_ids"] = "；".join(dict.fromkeys(entry.get("rule_id", "") for entry in entries if entry.get("rule_id")))
            inventory["classification"] = "；".join(dict.fromkeys(entry.get("classification", "") for entry in entries if entry.get("classification")))
            reasons = list(dict.fromkeys(entry.get("reason", "") for entry in entries if entry.get("reason")))
            inventory["reason"] = "；".join(reasons)
            if statuses <= candidate_statuses:
                inventory["status"] = ("AI补充候选（未启用，待人工确认）"
                                        if "AI补充候选（未启用，待人工确认）" in statuses
                                        else "已生成候选（未启用）")
            elif statuses & candidate_statuses:
                inventory["status"] = "部分候选已生成，另有分支待处理"
            elif "模板公式错误（待修复）" in statuses:
                inventory["status"] = "模板公式错误（待修复）"
            else:
                inventory["status"] = "待处理"
        workbook_path = directory / "公式转配置审核工作簿.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "审核规则"
        sheet.append(CANDIDATE_HEADERS)
        for row in rows:
            sheet.append([row.get(header, "") for header in CANDIDATE_HEADERS])
        _format(sheet, (28, 12, 34, 32, 18, 34, 14, 28, 45, 90, 14, 90, 55, 55, 55, 110, 18))
        sheet = workbook.create_sheet("转换统计")
        sheet.title = "转换统计"
        sheet.append(["模板文件", "校验区公式总数", "校验输出公式数", "生成候选规则数", "待处理项数", "中间公式数"])
        for summary in summaries:
            sheet.append([summary.get(key, 0) for key in ("template", "formulas", "outputs", "candidates", "pending", "intermediate")])
        _format(sheet, (65, 22, 22, 24, 22, 22))
        sheet = workbook.create_sheet("公式族")
        sheet.append(["公式族", "涉及项数", "生成候选数", "待处理数", "示例工作表", "示例单元格", "示例公式", "主要待处理原因"])
        families = defaultdict(list)
        for entry in details:
            families[entry["family"]].append(entry)
        for family, items in families.items():
            ready = sum(item["status"] in {
                "已生成候选（未启用）", "AI补充候选（未启用，待人工确认）",
            } for item in items)
            sample = items[0]
            sheet.append([family, len(items), ready, len(items)-ready, sample["sheet"], sample["cell"], sample["formula"], "；".join(dict.fromkeys(item["reason"] for item in items if item["status"] == "待处理"))[:30000]])
        _format(sheet, (24, 18, 20, 18, 34, 20, 105, 95))
        sheet = workbook.create_sheet("转换明细")
        sheet.append(["模板文件", "工作表", "公式单元格", "公式族", "原公式", "生成规则编号", "校验类型", "转换状态", "说明"])
        for entry in details:
            sheet.append([entry[key] for key in ("template", "sheet", "cell", "family", "formula", "rule_id", "classification", "status", "reason")])
        _format(sheet, (60, 34, 20, 24, 110, 30, 28, 28, 100))
        sheet = workbook.create_sheet("模板取数定位")
        sheet.append(["模板文件", "业务名称", "规则编号", "公式工作表", "公式单元格", "取值变量", "取数工作表", "行路径", "列路径", "模板数字单元格", "定位状态"])
        for check in path_checks:
            sheet.append([check[key] for key in ("template", "business", "rule_id", "formula_sheet", "formula_cell", "variable", "source_sheet", "row_path", "column_path", "template_cells", "status")])
        _format(sheet, (58, 30, 38, 32, 18, 20, 32, 72, 72, 48, 24))
        sheet = workbook.create_sheet("公式清单")
        sheet.append(["模板文件", "业务名称", "工作表名称", "公式单元格", "公式类别", "公式族", "原公式", "生成规则编号", "校验类型", "转换状态", "说明"])
        for entry in formula_inventory:
            sheet.append([entry[key] for key in ("template", "business", "sheet", "cell", "role", "family", "formula", "rule_ids", "classification", "status", "reason")])
        _format(sheet, (62, 32, 34, 20, 22, 24, 110, 42, 30, 40, 105))
        workbook.save(workbook_path)
        workbook.close()
        counts = {key: sum(summary.get(key, 0) for summary in summaries) for key in ("formulas", "outputs", "candidates", "pending", "intermediate")}
        path_statuses = Counter(check["status"] for check in path_checks)
        receipt = {"counts": counts, "inputs_unchanged": True, "input_hashes": {str(path): digest for path, digest in watched.items()}, "all_candidates_disabled": True}
        (directory / "生成记录.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8")
        return {"ok": True, "status": "COMPLETED_WITH_PENDING" if counts["pending"] else "COMPLETED", "counts": counts,
                "template_path_checks": {"total": len(path_checks), **dict(path_statuses)},
                "workbook_path": str(workbook_path), "candidate_path": str(workbook_path),
                "report_path": str(workbook_path), "files": copies, "output_dir": str(directory)}
    finally:
        if external_book:
            external_book.close()
