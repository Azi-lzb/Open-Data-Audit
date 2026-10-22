"""S5 host adapter for its independent configured audit engine."""
from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from base_audit.discovery import TemplateCatalog, read_xlsx_sheet_names
from s1audit_configured.batch import inspect_institutions, run_batch
from s1audit_configured.dsl import Call, Literal, TupleLiteral, calls_in, parse_dictionary, parse_expression
from s1audit_configured.errors import ConfigurationError, RetrievalError
from s1audit_configured.external_layout import discover_external_layout, normalized
from s1audit_configured.institutions import InstitutionCatalog
from s1audit_configured.layout import scan_rule_paths, split_path
from s1audit_configured.model import SourceBinding
from s1audit_configured.runtime import RuntimePaths, check_cancel
from s1audit_configured.sources import _Workbook
from s1audit_configured.workbook_config import load_configuration

from .contracts import S5Request


def _path_argument(node):
    if isinstance(node, Literal) and node.kind == "string":
        return split_path(node.value)
    if isinstance(node, TupleLiteral):
        return tuple(str(value).strip() for value in node.values)
    return ()


def _rule_formula_nodes(rule):
    nodes = []
    dictionary = parse_dictionary(rule.value_dictionary)
    nodes.extend(("取值字典", node) for node in dictionary.values())
    for field, text in (
        ("审核表达式", rule.audit_expression),
        ("当前值表达式", rule.current_expression),
        ("对比值表达式", rule.comparison_expression),
        ("差值表达式", rule.difference_expression),
        ("适用表达式", rule.applicable_expression),
    ):
        if text:
            nodes.append((field, parse_expression(text)))
    return nodes


def _configuration_rule_locations(config_path, progress=None):
    """Map stable rule identity to its physical Excel row for diagnostics."""
    workbook = load_workbook(config_path, read_only=True, data_only=False)
    try:
        sheet = workbook["审核规则"]
        header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [str(value or "").strip() for value in header_values]
        id_header = next((name for name in ("规则ID", "规则编号", "规则序号") if name in headers), "")
        if not id_header:
            return {}
        id_col = headers.index(id_header) + 1
        revision_col = headers.index("规则修订") + 1 if "规则修订" in headers else 0
        columns = {header: index for index, header in enumerate(headers, 1) if header}
        columns.setdefault("规则ID", id_col)
        locations = defaultdict(list)
        total = max(0, sheet.max_row - 1)
        first_col = min(id_col, revision_col) if revision_col else id_col
        last_col = max(id_col, revision_col) if revision_col else id_col
        for position, values in enumerate(
                sheet.iter_rows(min_row=2, min_col=first_col, max_col=last_col, values_only=True), 1):
            row = position + 1
            rule_id = str(values[id_col - first_col] or "").strip()
            if not rule_id:
                if progress and (position % 100 == 0 or position == total):
                    progress(position, total, f"建立规则错误定位索引 {position}/{total}")
                continue
            revision = values[revision_col - first_col] if revision_col else 1
            try:
                revision = int(revision or 1)
            except (TypeError, ValueError):
                revision = 1
            locations[(rule_id, revision)].append({"row": row, "columns": columns})
            if progress and (position % 100 == 0 or position == total):
                progress(position, total, f"建立规则错误定位索引 {position}/{total}")
        return dict(locations)
    finally:
        workbook.close()


def _rule_references(configuration, business="", progress=None):
    template_refs = defaultdict(list)
    external_refs = defaultdict(list)
    dynamic_refs = []
    rules = [rule for rule in configuration.rules if rule.enabled and (not business or rule.business == business)]
    total = len(rules)
    for position, rule in enumerate(rules, 1):
        for field, node in _rule_formula_nodes(rule):
            for call in calls_in(node):
                if call.name in {"V_CUR", "V_PRE"}:
                    if len(call.arguments) == 3:
                        sheet_node, row_node, column_node = call.arguments
                        sheet = str(sheet_node.value).strip() if isinstance(sheet_node, Literal) else ""
                    else:
                        row_node, column_node = call.arguments
                        sheet = dict(call.keywords).get("sheet", "")
                    sheet = sheet or rule.default_sheet
                    if isinstance(row_node, Literal) and row_node.kind == "string" and row_node.value.strip() == "@当前行":
                        dynamic_refs.append((rule, field, sheet, _path_argument(column_node), call.name))
                        continue
                    row_path, column_path = _path_argument(row_node), _path_argument(column_node)
                    if not sheet or not row_path or not column_path:
                        continue  # The shared configuration loader reports malformed paths.
                    key = (rule.business, sheet, row_path, column_path)
                    reference = {"rule_id": rule.rule_id, "revision": rule.revision, "field": field,
                                 "source": call.name}
                    if reference not in template_refs[key]:
                        template_refs[key].append(reference)
                elif call.name == "V_EXT":
                    metric = str(call.arguments[0].value).strip() if call.arguments and isinstance(call.arguments[0], Literal) else ""
                    keywords = dict(call.keywords)
                    key = (keywords.get("source", "EXT") or "EXT", metric, keywords.get("sheet", ""))
                    reference = {"rule_id": rule.rule_id, "revision": rule.revision, "field": field}
                    if metric and reference not in external_refs[key]:
                        external_refs[key].append(reference)
        if progress and (position % 20 == 0 or position == total):
            progress(position, total, f"提取启用规则地址引用 {position}/{total}")
    return template_refs, external_refs, dynamic_refs


def _config_locations(rule_refs, row_map):
    locations = []
    for reference in rule_refs:
        for item in row_map.get((reference["rule_id"], reference["revision"]), ()):
            column = item["columns"].get(reference["field"])
            locations.append({"sheet": "审核规则", "row": item["row"], "column": column,
                              "cell": f"{get_column_letter(column)}{item['row']}" if column else "",
                              "field": reference["field"], "rule_id": reference["rule_id"]})
    return locations


def _configured_business_sheets(config_path):
    """Read lightweight business-to-sheet hints without parsing rule expressions."""
    workbook = load_workbook(config_path, read_only=True, data_only=True)
    result = defaultdict(set)
    try:
        if "审核规则" not in workbook.sheetnames:
            return result
        sheet = workbook["审核规则"]
        headers = [str(value or "").strip() for value in next(
            sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()
        )]
        business_header = next((name for name in ("业务名称", "业务类型") if name in headers), "")
        sheet_header = next((name for name in ("工作表名称", "默认Sheet") if name in headers), "")
        enabled_header = "是否启用" if "是否启用" in headers else ""
        if business_header and sheet_header:
            business_col, sheet_col = headers.index(business_header), headers.index(sheet_header)
            enabled_col = headers.index(enabled_header) if enabled_header else -1
            for values in sheet.iter_rows(min_row=2, values_only=True):
                business = str(values[business_col] or "").strip()
                sheet_name = str(values[sheet_col] or "").strip()
                if not business or not sheet_name:
                    continue
                enabled = str(values[enabled_col] or "").strip().lower() if enabled_col >= 0 else "是"
                if enabled in {"是", "是的", "yes", "y", "true", "1", "启用", "已确认"}:
                    result[business].add(sheet_name)
        if "模板指标目录" in workbook.sheetnames:
            sheet = workbook["模板指标目录"]
            headers = [str(value or "").strip() for value in next(
                sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()
            )]
            if "业务类型" in headers and "工作表名称" in headers:
                business_col, sheet_col = headers.index("业务类型"), headers.index("工作表名称")
                for values in sheet.iter_rows(min_row=2, values_only=True):
                    business = str(values[business_col] or "").strip()
                    sheet_name = str(values[sheet_col] or "").strip()
                    if business and sheet_name and not result.get(business):
                        result[business].add(sheet_name)
        return result
    finally:
        workbook.close()


def _infer_business_from_template(config_path, template_path, configured_sheets=None):
    """Infer the one business most strongly represented by a matched template."""
    template_path = Path(template_path).expanduser().resolve()
    if not template_path.is_file():
        return "", f"匹配模板不存在：{template_path}"
    sheets = set(read_xlsx_sheet_names(template_path))
    if not sheets:
        return "", f"无法读取匹配模板的工作表清单：{template_path.name}"
    configured_sheets = configured_sheets or _configured_business_sheets(config_path)
    scored = []
    for business, required in configured_sheets.items():
        if not required:
            continue
        overlap = len(required & sheets)
        if overlap:
            scored.append((overlap / len(required), business))
    if not scored:
        return "", f"模板“{template_path.name}”的工作表无法对应到配置业务，请核实模板和业务工作表配置"
    best = max(score for score, _ in scored)
    candidates = [business for score, business in scored if score == best]
    if len(candidates) > 1:
        name_matches = [business for business in candidates if business in template_path.stem]
        if len(name_matches) == 1:
            candidates = name_matches
    if len(candidates) != 1:
        return "", f"模板“{template_path.name}”同时匹配多个业务（{'、'.join(sorted(candidates))}），无法唯一确定当前业务"
    return candidates[0], ""


def _select_business_template(business, required_sheets, profiles):
    """Find a unique template containing all sheets referenced by an enabled business."""
    if not required_sheets:
        return None, ""
    exact = [profile for profile in profiles if required_sheets.issubset(set(profile.sheet_names))]
    if len(exact) == 1:
        return exact[0], ""
    candidates = exact
    if not candidates:
        candidates = [profile for profile in profiles
                      if business in Path(profile.path).stem or business in profile.report_name]
        if len(candidates) == 1:
            return candidates[0], ""
    if len(candidates) > 1:
        named = [profile for profile in candidates if business in Path(profile.path).stem]
        if len(named) == 1:
            return named[0], ""
        names = "、".join(Path(profile.path).name for profile in candidates[:5])
        return None, f"业务“{business}”对应多个模板，无法安全选择：{names}"
    best = max((len(required_sheets & set(profile.sheet_names)) for profile in profiles), default=0)
    if best:
        names = "、".join(Path(profile.path).name for profile in profiles
                          if len(required_sheets & set(profile.sheet_names)) == best)[:400]
        return None, f"模板目录没有工作表覆盖业务“{business}”所需的路径；最接近的模板：{names}"
    return None, f"模板目录中没有找到业务“{business}”所需工作表：{'、'.join(sorted(required_sheets))}"


def _check_template_addresses(configuration, template_path, business, row_map, progress=None, references=None):
    if references is None:
        references, _, dynamic = _rule_references(configuration, business)
    else:
        references, _, dynamic = references
    dynamic = [item for item in dynamic if item[0].business == business]
    if not references and not dynamic:
        return {"status": "no_references", "business": business, "checked": 0,
                "resolved": 0, "errors": [], "dynamic_references": len(dynamic)}
    if not template_path:
        return {"status": "not_checked", "business": business, "checked": 0,
                "resolved": 0, "errors": [], "dynamic_references": len(dynamic),
                "message": "尚未选择或识别匹配模板，无法检查模板取数地址"}
    path = Path(template_path).expanduser().resolve()
    if not path.is_file():
        error = {"message": f"匹配模板不存在：{path}", "locations": []}
        return {"status": "failed", "business": business, "checked": 0, "resolved": 0, "errors": [error],
                "dynamic_references": len(dynamic)}
    if not business:
        error = {"message": "已选择匹配模板，但尚未识别审核业务；不能确定应使用哪组字典地址", "locations": []}
        return {"status": "not_checked", "business": "", "checked": 0, "resolved": 0, "errors": [],
                "dynamic_references": len(dynamic), "message": error["message"]}

    grouped = defaultdict(set)
    references_by_key = defaultdict(list)
    for key, owners in references.items():
        if key[0] == business:
            grouped[key[1]].add((key[2], key[3]))
            references_by_key[key].extend(owners)
    errors, checked, resolved = [], 0, 0
    book = None
    try:
        book = _Workbook(SourceBinding("S5_CONFIG_CHECK", path, business=business))
        total_sheets = len(grouped)
        for sheet_index, (sheet_name, key_pairs) in enumerate(grouped.items(), 1):
            if progress:
                progress("template", sheet_index - 1, total_sheets,
                         f"扫描模板工作表 {sheet_index}/{total_sheets}：{sheet_name}")
            if sheet_name not in book.book.sheetnames:
                locations = []
                for row_path, column_path in key_pairs:
                    owners = references_by_key[(business, sheet_name, row_path, column_path)]
                    locations.extend(_config_locations(owners, row_map))
                    checked += 1
                errors.append({"message": f"匹配模板缺少工作表：{sheet_name}", "locations": locations})
                if progress:
                    progress("template", sheet_index, total_sheets,
                             f"模板工作表检查完成 {sheet_index}/{total_sheets}：{sheet_name}")
                continue
            try:
                found, row_locations, column_locations, _ = scan_rule_paths(book, sheet_name, key_pairs)
            except (KeyError, RetrievalError, ValueError) as exc:
                for row_path, column_path in key_pairs:
                    checked += 1
                    owners = references_by_key[(business, sheet_name, row_path, column_path)]
                    errors.append({"message": f"无法扫描模板地址 {sheet_name} / {'@'.join(row_path)} / {'@'.join(column_path)}：{exc}",
                                   "locations": _config_locations(owners, row_map)})
                if progress:
                    progress("template", sheet_index, total_sheets,
                             f"模板工作表检查完成 {sheet_index}/{total_sheets}：{sheet_name}")
                continue
            for row_path, column_path in key_pairs:
                checked += 1
                matches = found.get((row_path, column_path), ())
                if len(matches) == 1:
                    resolved += 1
                    continue
                owners = references_by_key[(business, sheet_name, row_path, column_path)]
                if len(matches) > 1:
                    reason = "匹配到多个单元格：" + ", ".join(
                        f"{get_column_letter(col)}{row}" for row, col in matches[:10]
                    )
                elif not row_locations.get(row_path):
                    reason = "未找到完整行路径"
                elif not column_locations.get(column_path):
                    reason = "未找到完整列路径"
                else:
                    reason = "行路径和列路径存在，但没有唯一有效交点"
                errors.append({
                    "message": f"模板取值地址无法唯一定位：{sheet_name} / {'@'.join(row_path)} / {'@'.join(column_path)}；{reason}",
                    "locations": _config_locations(owners, row_map),
                })
            if progress:
                progress("template", sheet_index, total_sheets,
                         f"模板工作表检查完成 {sheet_index}/{total_sheets}：{sheet_name}")
    except (OSError, RetrievalError, ValueError) as exc:
        errors.append({"message": f"匹配模板无法读取或扫描：{exc}", "locations": []})
    finally:
        if book is not None:
            book.close()
    status = "failed" if errors else "partial" if dynamic else "passed"
    message = f"业务“{business}”检查 {checked} 个静态模板取数地址，成功定位 {resolved} 个"
    if dynamic:
        message += f"；另有 {len(dynamic)} 个 @当前行 地址尚未展开检查，需结合适用范围核对"
    return {"status": status, "business": business, "checked": checked, "resolved": resolved,
            "errors": errors, "dynamic_references": len(dynamic), "message": message}


def _check_external_addresses(configuration, external_path, business, row_map, progress=None, references=None):
    if references is None:
        _, references, _ = _rule_references(configuration, business)
    else:
        _, references, _ = references
    if not references:
        return {"status": "no_references", "checked": 0, "resolved": 0, "errors": []}
    if not external_path:
        return {"status": "not_checked", "checked": 0, "resolved": 0, "errors": [],
                "message": "配置引用 V_EXT，但尚未选择外部文件，无法检查指标表头地址"}
    path = Path(external_path).expanduser().resolve()
    if not path.is_file():
        return {"status": "failed", "checked": 0, "resolved": 0,
                "errors": [{"message": f"外部文件不存在：{path}", "locations": []}]}
    book = None
    try:
        if progress:
            progress("external", 0, 1, "读取大集中外部文件并识别布局")
        book = _Workbook(SourceBinding("S5_CONFIG_CHECK_EXT", path))
        layout = discover_external_layout(book, support_time_series=True)
        errors, resolved = [], 0
        for (_source_id, metric, sheet_filter), owners in references.items():
            matches = [column for column in layout.columns
                       if normalized(column.header) == normalized(metric)
                       and (not sheet_filter or column.sheet == sheet_filter)]
            if len(matches) == 1:
                resolved += 1
                continue
            reason = "未找到指标表头" if not matches else "指标表头重复，不能唯一定位"
            locations = _config_locations(owners, row_map)
            errors.append({"message": f"外部文件 V_EXT 指标地址无法唯一定位：{metric}"
                                     + (f"（Sheet={sheet_filter}）" if sheet_filter else "")
                                     + f"；{reason}", "locations": locations})
    except (OSError, RetrievalError, ValueError) as exc:
        errors = [{"message": f"外部文件无法读取或识别映射/指标布局：{exc}", "locations": []}]
        resolved = 0
    finally:
        if book is not None:
            book.close()
    if progress:
        progress("external", 1, 1, "大集中外部文件指标表头检查完成")
    return {"status": "failed" if errors else "passed", "checked": len(references),
            "resolved": resolved, "errors": errors,
            "message": f"外部文件检查 {len(references)} 个 V_EXT 指标表头，成功定位 {resolved} 个"}


def validate_s5_config(config_path, *, template_path="", template_dir="", template_index_path="",
                       external_path="", business="", scope="current", on_progress=None):
    full_scope = scope in {"all", "all_enabled"}
    strict_current = scope in {"current_business", "current"} and scope == "current_business"
    if scope not in {"current", "current_business", "all", "all_enabled"}:
        raise ValueError(f"未知的 S5 配置检查范围：{scope}")
    scope_label = "全量启用配置" if full_scope else "当前业务启用配置"

    def progress(phase, current=0, total=0, message=""):
        if not on_progress:
            return
        bounds = (
            {"rules": (3, 40), "locations": (40, 47), "references": (47, 55),
             "template": (55, 90), "external": (90, 97), "complete": (100, 100)}
            if full_scope else
            {"rules": (3, 53), "locations": (53, 63), "references": (63, 73),
             "template": (73, 86), "external": (86, 96), "complete": (100, 100)}
        )
        start, end = bounds.get(phase, (1, 4))
        fraction = max(0.0, min(1.0, current / total)) if total else (1.0 if phase == "complete" else 0.0)
        on_progress({"phase": phase, "current": current, "total": total,
                     "percent": round(start + (end - start) * fraction, 1), "message": message})

    business = str(business or "").strip()
    if not full_scope and not business and template_path:
        business, inference_error = _infer_business_from_template(config_path, template_path)
        if inference_error and strict_current:
            return {
                "valid": False, "complete": False, "scope": scope, "scope_label": scope_label,
                "business": "", "message": inference_error, "path": str(Path(config_path).resolve()),
                "rules": 0, "enabled_rules": 0, "disabled_rules": 0,
                "checks": {name: {"status": "not_checked"} for name in
                           ("expressions", "dictionary_keys", "template_addresses", "external_addresses")},
                "warnings": [inference_error],
                "configuration_errors": [{"message": inference_error, "locations": []}],
            }
    if strict_current and (not business or not template_path):
        message = "当前业务检查需要先识别或选择匹配模板，再根据模板确定业务。"
        return {
            "valid": False, "complete": False, "scope": scope, "scope_label": scope_label,
            "business": business, "message": message, "path": str(Path(config_path).resolve()),
            "rules": 0, "enabled_rules": 0, "disabled_rules": 0,
            "checks": {name: {"status": "not_checked"} for name in
                       ("expressions", "dictionary_keys", "template_addresses", "external_addresses")},
            "warnings": [message], "configuration_errors": [{"message": message, "locations": []}],
        }

    progress("rules", 0, 1, "读取审核配置并检查启用规则")

    loader_snapshot = {}

    def loader_progress(event):
        loader_snapshot.update(event)
        progress("rules", event.get("current", 0), event.get("total", 0), event.get("message", ""))

    try:
        configuration = load_configuration(
            config_path, on_progress=loader_progress,
            business=business if not full_scope and business else "",
        )
    except ConfigurationError as exc:
        return {
            "valid": False, "complete": False, "scope": scope, "scope_label": scope_label,
            "business": business, "message": "S5 配置语法或字典引用检查未通过",
            "path": str(Path(config_path).resolve()),
            "rules": loader_snapshot.get("total", 0),
            "enabled_rules": loader_snapshot.get("enabled_rules", 0),
            "disabled_rules": max(0, loader_snapshot.get("total", 0) - loader_snapshot.get("enabled_rules", 0)),
            "checks": {
                "expressions": {"status": "failed"},
                "dictionary_keys": {"status": "failed"},
                "template_addresses": {"status": "not_checked"},
                "external_addresses": {"status": "not_checked"},
            },
            "warnings": [], "configuration_errors": [exc.as_dict()],
        }
    rules_by_business = Counter(rule.business for rule in configuration.rules)
    enabled_by_business = Counter(rule.business for rule in configuration.rules if rule.enabled)
    enabled_rules = sum(rule.enabled for rule in configuration.rules)
    all_businesses = sorted(
        name for name, count in enabled_by_business.items() if count
    ) if full_scope else sorted(rules_by_business)
    if not enabled_rules:
        if business and not full_scope:
            warning = (f"业务“{business}”的规则全部禁用；本次未执行表达式、字典键引用以及模板/外部文件地址检查。"
                       if configuration.rules else
                       f"配置中没有业务“{business}”的规则；本次未执行表达式、字典键引用以及模板/外部文件地址检查。")
        else:
            warning = "当前没有启用规则；表达式、字典键引用以及模板/外部文件地址检查均未执行。启用需要检查的规则后再运行检查。"
        progress("complete", 1, 1, "配置结构读取完成；没有启用规则，因此规则检查未执行")
        empty_message = (
            f"配置文件结构可读取，共 {len(configuration.rules)} 条规则，全部禁用。{warning}"
            if configuration.rules else
            f"配置文件结构可读取，{warning}"
        )
        return {
            "valid": True,
            "complete": False,
            "scope": scope, "scope_label": scope_label, "business": business,
            "message": empty_message,
            "path": str(Path(config_path).resolve()),
            "version": configuration.version,
            "version_status": configuration.version_status,
            "rules": len(configuration.rules),
            "enabled_rules": 0,
            "disabled_rules": len(configuration.rules),
            "metrics": len(configuration.metrics),
            "external_metrics": len(configuration.external_metrics),
            "source_contracts": len(configuration.source_contracts),
            "applicability_scopes": len(configuration.applicability_scopes),
            "businesses": [
                {"name": name, "rules": rules_by_business[name], "enabled_rules": 0}
                for name in all_businesses
            ],
            "checks": {
                "expressions": {"status": "not_checked", "rules": 0},
                "dictionary_keys": {"status": "not_checked", "rules": 0},
                "template_addresses": {"status": "not_checked", "checked": 0, "resolved": 0,
                                        "errors": [], "message": "没有启用规则，未扫描模板地址"},
                "external_addresses": {"status": "not_checked", "checked": 0, "resolved": 0,
                                       "errors": [], "message": "没有启用规则，未扫描外部文件地址"},
            },
            "warnings": [warning],
            "configuration_errors": [],
        }
    row_map = _configuration_rule_locations(
        config_path,
        progress=lambda current, total, message: progress("locations", current, total, message),
    )
    progress("references", 0, 1, "提取启用规则中的模板与外部文件地址引用")
    references = _rule_references(
        configuration, business if not full_scope else "",
        progress=lambda current, total, message: progress("references", current, total, message),
    )

    business_checks = []
    errors = []
    warnings = []
    if full_scope:
        template_refs, _, dynamic_refs = references
        required_sheets = defaultdict(set)
        for business_name, sheet_name, _, _ in template_refs:
            required_sheets[business_name].add(sheet_name)
        for rule, _, sheet_name, _, _ in dynamic_refs:
            required_sheets[rule.business].add(sheet_name)
        try:
            profiles = TemplateCatalog(
                Path(template_dir).expanduser().resolve(),
                Path(template_index_path).expanduser().resolve() if template_index_path else
                Path(template_dir).expanduser().resolve().parent / "模板索引.json",
            ).profiles() if template_dir and Path(template_dir).is_dir() else []
        except (OSError, ValueError) as exc:
            profiles = []
            catalog_error = f"无法读取模板目录：{exc}"
        else:
            catalog_error = "" if profiles else "模板目录中没有可读取的 XLSX 模板"

        template_checks = []
        template_checked = template_resolved = 0
        for index, business_name in enumerate(all_businesses, 1):
            required = required_sheets.get(business_name, set())
            if not required:
                check = {"status": "no_references", "business": business_name,
                         "checked": 0, "resolved": 0, "errors": [], "dynamic_references": 0,
                         "message": "该业务启用规则没有模板取数地址引用"}
                template_name = ""
            elif not profiles:
                message = catalog_error or "请先选择有效的模板文件目录"
                check = {"status": "not_checked", "business": business_name,
                         "checked": 0, "resolved": 0, "errors": [],
                         "dynamic_references": sum(rule.business == business_name for rule, *_ in dynamic_refs),
                         "message": f"业务“{business_name}”未检查模板地址：{message}"}
                template_name = ""
            else:
                profile, match_error = _select_business_template(business_name, required, profiles)
                if profile is None:
                    check = {"status": "not_checked", "business": business_name,
                             "checked": 0, "resolved": 0, "errors": [],
                             "dynamic_references": sum(rule.business == business_name for rule, *_ in dynamic_refs),
                             "message": match_error}
                    template_name = ""
                else:
                    template_name = Path(profile.path).name

                    def business_progress(phase, current=0, total=0, message="", *, offset=index - 1):
                        local_fraction = current / total if total else 1
                        fraction = (offset + max(0.0, min(1.0, local_fraction))) / max(1, len(all_businesses))
                        progress("template", fraction, 1, f"业务 {business_name}（模板 {template_name}）：{message}")

                    check = _check_template_addresses(
                        configuration, profile.path, business_name, row_map, business_progress, references,
                    )
                    check["template_path"] = profile.path
                    check["template_name"] = template_name
            template_checked += int(check.get("checked", 0) or 0)
            template_resolved += int(check.get("resolved", 0) or 0)
            if check.get("status") in {"not_checked", "partial"} and check.get("message"):
                warnings.append(check["message"])
            for issue in check.get("errors", []):
                issue = dict(issue)
                issue["business"] = business_name
                errors.append(issue)
            template_checks.append(check)
            business_checks.append({
                "name": business_name, "rules": rules_by_business[business_name],
                "enabled_rules": enabled_by_business[business_name],
                "template": template_name, "template_addresses": {
                    key: value for key, value in check.items() if key != "errors"
                },
            })
            progress("template", index, len(all_businesses),
                     f"全量模板检查完成 {index}/{len(all_businesses)}：{business_name}")
        template_statuses = {item["status"] for item in template_checks}
        template_status = (
            "failed" if "failed" in template_statuses else
            "not_checked" if "not_checked" in template_statuses else
            "partial" if "partial" in template_statuses else
            "no_references" if template_statuses == {"no_references"} else "passed"
        )
        template_check = {
            "status": template_status, "checked": template_checked, "resolved": template_resolved,
            "businesses": template_checks,
            "message": f"全量检查 {len(all_businesses)} 个启用业务的模板地址，成功定位 {template_resolved}/{template_checked} 个",
        }
        if catalog_error and any(item["status"] == "not_checked" for item in template_checks):
            warnings.append(catalog_error)
    else:
        progress("template", 0, 1, "检查当前业务启用规则引用的模板地址")
        template_check = _check_template_addresses(configuration, template_path, business, row_map, progress, references)
        business_checks = [{
            "name": business, "rules": rules_by_business.get(business, len(configuration.rules)),
            "enabled_rules": enabled_by_business.get(business, enabled_rules),
            "template": Path(template_path).name if template_path else "",
            "template_addresses": {key: value for key, value in template_check.items() if key != "errors"},
        }]
        if template_check.get("status") in {"not_checked", "partial"} and template_check.get("message"):
            warnings.append(template_check["message"])
        errors.extend(template_check.get("errors", []))

    progress("external", 0, 1, "检查启用规则引用的大集中外部指标地址")
    external_check = _check_external_addresses(
        configuration, external_path, "" if full_scope else business, row_map, progress, references,
    )
    errors.extend(external_check["errors"])
    if external_check.get("status") in {"not_checked", "partial"} and external_check.get("message"):
        warnings.append(external_check["message"])
    checks_complete = all(
        check["status"] not in {"not_checked", "partial"}
        for check in (template_check, external_check)
    )
    valid = not errors
    progress("complete", 1, 1, "配置检查完成")
    parser_summary = f"已用正式审核表达式解析器检查{scope_label}中的 {enabled_rules} 条启用规则，并校验表达式变量均能在对应取值字典键中找到"
    message_parts = ["配置语法与字典引用检查通过" if valid else "配置检查发现问题", parser_summary]
    if business and not full_scope:
        message_parts.append(f"当前业务：{business}")
    for check in (template_check, external_check):
        if check.get("message"):
            message_parts.append(check["message"])
    if warnings:
        message_parts.extend(warnings)
    return {
        "valid": valid,
        "complete": checks_complete,
        "scope": scope,
        "scope_label": scope_label,
        "business": business if not full_scope else "",
        "message": "\n".join(message_parts),
        "path": str(Path(config_path).resolve()),
        "version": configuration.version,
        "version_status": configuration.version_status,
        "rules": len(configuration.rules),
        "enabled_rules": enabled_rules,
        "disabled_rules": sum(not rule.enabled for rule in configuration.rules),
        "metrics": len(configuration.metrics),
        "external_metrics": len(configuration.external_metrics),
        "source_contracts": len(configuration.source_contracts),
        "applicability_scopes": len(configuration.applicability_scopes),
        "businesses": [
            {"name": name, "rules": rules_by_business[name], "enabled_rules": enabled_by_business[name]}
            for name in all_businesses
        ],
        "business_checks": business_checks,
        "checks": {
            "expressions": {"status": "passed", "rules": enabled_rules},
            "dictionary_keys": {"status": "passed", "rules": enabled_rules},
            "template_addresses": template_check,
            "external_addresses": external_check,
        },
        "warnings": warnings,
        "configuration_errors": errors,
    }


def inspect_s5_institutions(request: S5Request):
    catalog = InstitutionCatalog([request.external_path] if request.external_path else [])
    return inspect_institutions(list(request.files), catalog,
                                identity_sheet=request.identity_sheet, identity_cell=request.identity_cell)


def run_s5_f1(request: S5Request, *, data_dir: Path, on_progress=None, cancel_event=None):
    check_cancel(cancel_event)
    configuration = load_configuration(request.config_path)
    businesses = {rule.business for rule in configuration.rules}
    if not request.business or request.business not in businesses:
        raise ValueError("请选择审核配置中登记的一项业务；每次仅执行一个业务")
    if not any(rule.enabled and rule.business == request.business for rule in configuration.rules):
        raise ValueError(f"业务“{request.business}”没有启用规则，请先在审核配置中核实启停设置")
    result = run_batch(
        configuration, list(request.files), business=request.business, period=request.period,
        output_root=request.output_dir, catalog=InstitutionCatalog([request.external_path] if request.external_path else []),
        previous_files=list(request.previous_files), previous_period=request.previous_period,
        external=request.external_path, external_period=request.external_period,
        identity_sheet=request.identity_sheet, identity_cell=request.identity_cell,
        runtime_paths=RuntimePaths(template_path=request.template_path,
                                   schema_snapshot_dir=Path(data_dir) / "s5" / ".schema_snapshots"),
        on_progress=on_progress, cancel_event=cancel_event,
        allow_external_time_series=True,
    )
    errors = sum(result["counts"].get(key, 0) for key in
                 ("RETRIEVAL_FAILED", "STRUCTURE_ERROR", "CALCULATION_FAILED", "CONFIG_ERROR"))
    failed = result["status"] in {"FAILED", "INPUT_CHANGED_DURING_RUN"}
    partial = result["status"] != "COMPLETED" or errors or result["conditional_unresolved"]
    data_quality_warnings = result.get("data_quality_warnings", [])
    warnings = list(result.get("reader_warnings", []))
    if data_quality_warnings:
        warnings.append(
            f"外部数据质量提示 {len(data_quality_warnings)} 条，详细位置见结果工作簿“数据质量提示”表。"
        )
        for warning in data_quality_warnings:
            identity = warning.get("institution_name", "")
            location = "/".join(value for value in (
                warning.get("source_file", ""), warning.get("sheet", ""), warning.get("location", ""),
            ) if value)
            prefix = f"{identity}：" if identity else ""
            warnings.append(f"{prefix}{location}：{warning.get('message', '')}")
    result.update(
        system="S5", function="S5-F01", engine_status=result["status"],
        task_state="failed" if failed else "partial_failed" if partial else "succeeded",
        finding_counts={"rule_triggered": result["counts"].get("TRIGGERED", 0),
                        "conditional_triggered": result["conditional_triggered"],
                        "conditional_unresolved": result["conditional_unresolved"],
                        "rule_execution_errors": errors,
                        "data_quality_warnings": len(data_quality_warnings)},
        warnings=warnings,
    )
    return result
