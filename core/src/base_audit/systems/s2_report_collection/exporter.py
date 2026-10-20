"""S2-F01 Excel 输出器。

业务引擎负责生成审核结论和取数语义；本模块只负责把三类 finding
按各自的字段契约写入三个独立工作表，不从空值猜测业务含义。
"""

from __future__ import annotations

import math
import os
import tempfile
import zipfile
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Callable, Iterable
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.compat.strings import safe_string
from openpyxl.utils.exceptions import IllegalCharacterError
from openpyxl.styles import Alignment, Font, PatternFill

from .models import AuditFinding, AuditSummary, IndicatorRecord, STATUS_ABNORMAL, STATUS_NORMAL, V2RunPaths, SEVERITIES


def build_summary(records: Iterable[IndicatorRecord], findings: Iterable[AuditFinding]) -> AuditSummary:
    record_list, finding_list = list(records), list(findings)
    abnormal = [item for item in finding_list if item.status == STATUS_ABNORMAL]
    return AuditSummary(
        institutions=len({item.institution_id for item in record_list}), forms=len({item.form_code for item in record_list if item.form_code}),
        indicators=len(record_list), total_findings=len(finding_list), abnormal_findings=len(abnormal),
        by_type=dict(Counter(item.audit_type for item in abnormal)), by_severity=dict(Counter(item.severity for item in abnormal if item.severity)),
    )


PERIOD_HEADERS = [
    "地区", "承接行", "机构类别", "机构名称", "社会信用代码", "表单", "指标代码", "数据属性", "指标名称",
    "审核级别", "审核说明", "本期值", "上期值", "差异值", "变动幅度", "计算过程",
]

EXTERNAL_HEADERS = [
    "地区", "承接行", "机构类别", "机构名称", "社会信用代码", "报表项目", "表单", "指标代码", "指标名称",
    "外部指标名称", "审核级别", "审核说明", "报表值", "外部值", "差异值", "计算过程",
]

VALIDATION_HEADERS = [
    "地区", "承接行", "机构类别", "机构名称", "社会信用代码", "涉及表单", "规则编号", "规则类型",
    "审核级别", "审核说明", "取值明细", "计算过程",
]

SHEET_ORDER = ("环比规则", "外部核对规则", "校验规则")

_AUDIT_TYPE_TO_SHEET = {
    "环比": "环比规则",
    "环比规则": "环比规则",
    "外部核对": "外部核对规则",
    "外部核对规则": "外部核对规则",
    "规则": "校验规则",
    "校验规则": "校验规则",
}

_LONG_TEXT_HEADERS = {"审核说明", "取值明细", "计算过程", "报表项目"}


def _validate_finding(item: AuditFinding) -> None:
    """Fail loudly on status/severity conflicts instead of correcting in export."""
    if item.status == STATUS_NORMAL:
        if item.severity not in (None, ""):
            raise ValueError(
                f"S2-F01 审核级别冲突：正常 finding 不应有审核级别；"
                f"finding_id={item.finding_id} 机构={item.institution_name} 表单={item.form_code}"
                f" 规则={item.rule_id} 指标={item.indicator_code} 状态={item.status} 级别={item.severity}"
            )
        return
    if item.status == STATUS_ABNORMAL and item.severity in SEVERITIES:
        return
    raise ValueError(
        f"S2-F01 审核级别冲突：异常 finding 必须是硬性/软性；"
        f"finding_id={item.finding_id} 机构={item.institution_name} 表单={item.form_code}"
        f" 规则={item.rule_id} 指标={item.indicator_code} 状态={item.status} 级别={item.severity}"
    )


def _value_or_blank(value: object) -> object:
    """Preserve real zero and numeric values; only missing values become blank cells."""
    return "" if value is None else value


def _period_description(item: AuditFinding) -> str:
    return item.rule_description or item.description or ""


def _period_values(item: AuditFinding) -> tuple[object, ...]:
    return (
        item.region, item.handling_branch, item.institution_type, item.institution_name, item.social_credit_code,
        "、".join(item.form_codes) or item.form_code, item.indicator_code, item.data_type or "", item.indicator_name,
        item.severity or "", _period_description(item),
        _value_or_blank(item.current_value), _value_or_blank(item.previous_value), _value_or_blank(item.difference_value),
        _value_or_blank(item.change_rate), item.calculation_trace or "",
    )


def _external_values(item: AuditFinding) -> tuple[object, ...]:
    return (
        item.region, item.handling_branch, item.institution_type, item.institution_name, item.social_credit_code,
        item.report_item,
        "、".join(item.form_codes) or item.form_code, item.indicator_code, item.indicator_name,
        item.external_indicator_name, item.severity or "", item.rule_description or item.description,
        _value_or_blank(item.current_value), _value_or_blank(item.external_value), _value_or_blank(item.difference_value),
        item.calculation_trace or "",
    )


def _validation_form_scope(item: AuditFinding) -> str:
    related_forms = (record.form_code for record in item.related_records)
    forms = tuple(sorted({value for value in (*item.form_codes, item.form_code, *related_forms) if value}))
    if forms:
        return "、".join(forms)
    return "不适用" if item.rule_type in {"导入校验", "配置检查"} else "未识别"


def _validation_values(item: AuditFinding) -> tuple[object, ...]:
    rule_type = item.rule_type or "未识别"
    value_details = item.value_details or ("不适用" if rule_type in {"导入校验", "配置检查"} else "未识别")
    return (
        item.region or ("不适用" if not item.institution_name else ""),
        item.handling_branch or ("不适用" if not item.institution_name else ""),
        item.institution_type or ("不适用" if not item.institution_name else ""),
        item.institution_name or "全局", item.social_credit_code or ("不适用" if not item.institution_name else ""),
        _validation_form_scope(item), item.rule_id or "未配置", rule_type, item.severity or "",
        item.rule_description or item.description, value_details, item.calculation_trace or "",
    )


def _write_table(
    sheet, headers: list[str], rows: Iterable[Iterable[object]], *,
    percentage_rows: set[int] | None = None, point_rows: set[int] | None = None,
    on_progress: Callable[[int], None] | None = None,
) -> None:
    sheet.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    for row_number, row in enumerate(rows, start=1):
        sheet.append(list(row))
        if on_progress and row_number % 250 == 0:
            on_progress(row_number)

    percentage_rows = percentage_rows or set()
    point_rows = point_rows or set()
    if "变动幅度" in headers:
        column = headers.index("变动幅度") + 1
        for row_index in percentage_rows:
            sheet.cell(row=row_index, column=column).number_format = "0.00%"
        for row_index in point_rows:
            sheet.cell(row=row_index, column=column).number_format = "0.00"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_index = {value: index + 1 for index, value in enumerate(headers)}
    if "报表项目" in header_index:
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row_index, header_index["报表项目"])
            if cell.data_type == "f":
                # 项目标签始终是文本，不能因以等号开头而变成 Excel 公式。
                cell.data_type = "s"
    text_alignment = Alignment(vertical="top", wrap_text=True)
    for title in _LONG_TEXT_HEADERS.intersection(headers):
        column = header_index[title]
        for row in range(1, sheet.max_row + 1):
            sheet.cell(row=row, column=column).alignment = text_alignment
    for column in sheet.columns:
        letter = column[0].column_letter
        values = [str(cell.value or "") for cell in column]
        sheet.column_dimensions[letter].width = min(max(max((len(value) for value in values), default=0) + 2, 12), 42)
    if on_progress:
        on_progress(sheet.max_row - 1)


def _group_findings(findings: Iterable[AuditFinding]) -> dict[str, list[AuditFinding]]:
    grouped = {name: [] for name in SHEET_ORDER}
    for item in findings:
        _validate_finding(item)
        sheet_name = _AUDIT_TYPE_TO_SHEET.get(item.audit_type)
        if sheet_name is None:
            raise ValueError(f"S2-F01 遇到未知审核类型，未生成结果：{item.audit_type!r} finding_id={item.finding_id}")
        grouped[sheet_name].append(item)
    return grouped


def _hide_rows_with_description(sheet, description: str) -> None:
    """Keep findings in the workbook while hiding rows with an exact displayed description."""
    description_column = next(
        cell.column for cell in sheet[1] if cell.value == "审核说明"
    )
    for row_index in range(2, sheet.max_row + 1):
        if sheet.cell(row_index, description_column).value == description:
            sheet.row_dimensions[row_index].hidden = True


def _column_letter(number: int) -> str:
    """Convert a one-based column number to its Excel column letters."""
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _checked_text(value: str) -> str:
    """Match openpyxl's illegal-character rejection and cell text limit."""
    if ILLEGAL_CHARACTERS_RE.search(value):
        raise IllegalCharacterError(
            "S2-F01 导出值包含 Excel 不支持的控制字符，无法生成审核清单。"
        )
    return value[:32767]


def _direct_cell_xml(ref: str, value: object, style_id: int | None, *, literal: bool = False) -> str:
    attributes = f' r="{ref}"'
    if style_id:
        attributes += f' s="{style_id}"'
    if value is None:
        return f"<c{attributes}/>" if style_id else ""
    if value == "":
        return f'<c{attributes} t="inlineStr"/>'
    if isinstance(value, bool):
        return f'<c{attributes} t="b"><v>{1 if value else 0}</v></c>'
    if isinstance(value, (int, float, Decimal)):
        return f"<c{attributes}><v>{safe_string(value)}</v></c>"

    text = _checked_text(str(value))
    if not literal and text.startswith("="):
        formula = escape(text[1:])
        return f"<c{attributes}><f>{formula}</f><v></v></c>"
    escaped_text = escape(text)
    preserve = ' xml:space="preserve"' if text[:1].isspace() or text[-1:].isspace() else ""
    return f'<c{attributes} t="inlineStr"><is><t{preserve}>{escaped_text}</t></is></c>'


def _supports_direct_ooxml(tables: list[tuple[str, list[str], list[tuple[object, ...]], set[int], str | None]]) -> bool:
    """Keep unusual value types on openpyxl's general writer path."""
    for _sheet_name, _headers, rows, _point_rows, _hidden_description in tables:
        for row in rows:
            for value in row:
                if value is None or isinstance(value, (str, bool, int)):
                    continue
                if isinstance(value, float):
                    if math.isfinite(value):
                        continue
                    return False
                if isinstance(value, Decimal):
                    if value.is_finite():
                        continue
                    return False
                return False
    return True


def _direct_styles_xml(has_point_format: bool) -> str:
    point_style = (
        '<xf numFmtId="2" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
        if has_point_format else ""
    )
    style_count = 6 if has_point_format else 5
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<numFmts count="0"/>
<fonts count="2"><font><name val="Calibri"/><family val="2"/><color theme="1"/><sz val="11"/><scheme val="minor"/></font><font><b/><color rgb="00FFFFFF"/></font></fonts>
<fills count="3"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="001F4E78"/><bgColor rgb="00000000"/></patternFill></fill></fills>
<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
<cellXfs count="{style_count}">
<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1" applyAlignment="1"><alignment horizontal="center" vertical="center" wrapText="1"/></xf>
<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1" applyAlignment="1"><alignment vertical="top" wrapText="1"/></xf>
<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyAlignment="1"><alignment vertical="top" wrapText="1"/></xf>
<xf numFmtId="10" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>
{point_style}
</cellXfs>
<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles><dxfs count="0"/>
<tableStyles count="0" defaultTableStyle="TableStyleMedium2" defaultPivotStyle="PivotStyleLight16"/>
</styleSheet>'''


def _direct_sheet_xml(
    headers: list[str], rows: list[tuple[object, ...]], *, point_rows: set[int],
    hidden_description: str | None, on_progress: Callable[[int], None] | None,
) -> str:
    long_columns = {headers.index(name) + 1 for name in _LONG_TEXT_HEADERS.intersection(headers)}
    percentage_column = headers.index("变动幅度") + 1 if "变动幅度" in headers else 0
    description_column = headers.index("审核说明") if hidden_description is not None and "审核说明" in headers else -1
    max_lengths = [0] * len(headers)
    for row in [headers, *rows]:
        for index, value in enumerate(row):
            text = "" if not value else str(value)
            if isinstance(value, str):
                text = _checked_text(value)
            max_lengths[index] = max(max_lengths[index], len(text))
    columns_xml = "".join(
        f'<col min="{index}" max="{index}" width="{min(max(max_lengths[index - 1] + 2, 12), 42)}" customWidth="1"/>'
        for index in range(1, len(headers) + 1)
    )
    last_cell = f"{_column_letter(len(headers))}{len(rows) + 1}"
    output = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">',
        f'<dimension ref="A1:{last_cell}"/>',
        '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/><selection pane="bottomLeft" activeCell="A2" sqref="A2"/></sheetView></sheetViews>',
        '<sheetFormatPr defaultRowHeight="15"/>', f"<cols>{columns_xml}</cols>", "<sheetData>",
    ]
    for row_number, row in enumerate([headers, *rows], start=1):
        hidden = row_number > 1 and description_column >= 0 and row[description_column] == hidden_description
        output.append(f'<row r="{row_number}" hidden="1">' if hidden else f'<row r="{row_number}">')
        for column_number, value in enumerate(row, start=1):
            style_id = None
            if row_number == 1:
                style_id = 2 if column_number in long_columns else 1
            elif column_number in long_columns:
                style_id = 3
            if row_number > 1 and column_number == percentage_column:
                style_id = 5 if row_number in point_rows else 4
            is_literal_report_item = (
                "报表项目" in headers and column_number == headers.index("报表项目") + 1
            )
            output.append(_direct_cell_xml(
                f"{_column_letter(column_number)}{row_number}", value, style_id,
                literal=is_literal_report_item,
            ))
        output.append("</row>")
        if on_progress and row_number > 1 and (row_number - 1) % 250 == 0:
            on_progress(row_number - 1)
    if on_progress:
        on_progress(len(rows))
    output.extend(("</sheetData>", f'<autoFilter ref="A1:{last_cell}"/>',
                   '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>',
                   "</worksheet>"))
    return "".join(output)


def _write_direct_ooxml(
    path: Path,
    tables: list[tuple[str, list[str], list[tuple[object, ...]], set[int], str | None]],
    *, on_progress: Callable[[str, int, int], None] | None,
) -> None:
    total_rows = sum(len(rows) for _name, _headers, rows, _points, _hidden in tables)
    completed_rows = 0
    worksheets: list[str] = []
    for sheet_name, headers, rows, point_rows, hidden_description in tables:
        if on_progress:
            on_progress(sheet_name, completed_rows, total_rows)
        xml = _direct_sheet_xml(
            headers, rows, point_rows=point_rows, hidden_description=hidden_description,
            on_progress=(lambda count, base=completed_rows: on_progress(sheet_name, base + count, total_rows)) if on_progress else None,
        )
        worksheets.append(xml)
        completed_rows += len(rows)

    if on_progress:
        on_progress("save", completed_rows, total_rows)

    content_types = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/worksheets/sheet3.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
</Types>'''
    root_relationships = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>'''
    workbook = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<workbookPr/><bookViews><workbookView activeTab="0"/></bookViews>
<sheets><sheet name="环比规则" sheetId="1" r:id="rId1"/><sheet name="外部核对规则" sheetId="2" r:id="rId2"/><sheet name="校验规则" sheetId="3" r:id="rId3"/></sheets><calcPr calcId="124519"/></workbook>'''
    workbook_relationships = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet3.xml"/>
<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>'''
    app_properties = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><Application>审核工具</Application></Properties>'''
    timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
    core_properties = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"><dc:creator>审核工具</dc:creator><dcterms:created xsi:type="dcterms:W3CDTF">{timestamp}</dcterms:created><dcterms:modified xsi:type="dcterms:W3CDTF">{timestamp}</dcterms:modified></cp:coreProperties>'''
    package_items = [
        ("[Content_Types].xml", content_types), ("_rels/.rels", root_relationships),
        ("docProps/app.xml", app_properties), ("docProps/core.xml", core_properties),
        ("xl/workbook.xml", workbook), ("xl/_rels/workbook.xml.rels", workbook_relationships),
        ("xl/styles.xml", _direct_styles_xml(any(points for _n, _h, _r, points, _d in tables))),
        *((f"xl/worksheets/sheet{index}.xml", xml) for index, xml in enumerate(worksheets, start=1)),
    ]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as package:
        for name, content in package_items:
            package.writestr(name, content.encode("utf-8"))


class ExcelExporter:
    def export(
        self, *, output_dir: Path, current_records: list[IndicatorRecord], previous_records: list[IndicatorRecord],
        findings: list[AuditFinding], run_info: dict[str, str], hide_unchanged_period_rows: bool = False,
        hide_matching_external_rows: bool = False, on_progress: Callable[[str, int, int], None] | None = None,
    ) -> V2RunPaths:
        output_dir.mkdir(parents=True, exist_ok=True)
        result_path = output_dir / f"报表审核清单_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        grouped = _group_findings(findings)
        period_items = grouped["环比规则"]
        period_rows = [_period_values(item) for item in period_items]
        external_rows = [_external_values(item) for item in grouped["外部核对规则"]]
        validation_rows = [_validation_values(item) for item in grouped["校验规则"]]
        percentage_rows = {index for index, item in enumerate(period_items, start=2) if item.value_type != "百分数"}
        point_rows = {index for index, item in enumerate(period_items, start=2) if item.value_type == "百分数"}
        tables = [
            ("环比规则", PERIOD_HEADERS, period_rows, point_rows, "无变动" if hide_unchanged_period_rows else None),
            ("外部核对规则", EXTERNAL_HEADERS, external_rows, set(), "与大集中系统数据一致" if hide_matching_external_rows else None),
            ("校验规则", VALIDATION_HEADERS, validation_rows, set(), None),
        ]
        file_descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{result_path.stem}_", suffix=".tmp", dir=output_dir,
        )
        os.close(file_descriptor)
        temporary_path = Path(temporary_name)
        try:
            if _supports_direct_ooxml(tables):
                _write_direct_ooxml(temporary_path, tables, on_progress=on_progress)
            else:
                # 非常见日期/自定义数值类型继续交给 openpyxl 的通用写入器处理。
                book = Workbook()
                first = book.active
                first.title = SHEET_ORDER[0]
                for sheet_name in SHEET_ORDER[1:]:
                    book.create_sheet(sheet_name)
                completed_rows = 0
                for sheet_name, headers, rows, table_point_rows, _hidden_description in tables:
                    if on_progress:
                        on_progress(sheet_name, completed_rows, len(findings))
                    _write_table(
                        book[sheet_name], headers, rows,
                        percentage_rows=percentage_rows if sheet_name == "环比规则" else None,
                        point_rows=table_point_rows if sheet_name == "环比规则" else None,
                        on_progress=(lambda count, base=completed_rows, name=sheet_name: on_progress(name, base + count, len(findings))) if on_progress else None,
                    )
                    completed_rows += len(rows)
                if hide_unchanged_period_rows:
                    _hide_rows_with_description(book["环比规则"], "无变动")
                if hide_matching_external_rows:
                    _hide_rows_with_description(book["外部核对规则"], "与大集中系统数据一致")
                if on_progress:
                    on_progress("save", completed_rows, len(findings))
                book.save(temporary_path)
            os.replace(temporary_path, result_path)
        finally:
            if temporary_path.exists():
                temporary_path.unlink()
        return V2RunPaths(result_path)
