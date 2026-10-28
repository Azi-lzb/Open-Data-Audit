"""V2 源数据与大集中数据导入器；所有 Excel 细节停留在本层。"""

from __future__ import annotations

import re
import time
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from importlib.metadata import PackageNotFoundError, version as package_version
from pathlib import Path
from typing import Any, Callable, Tuple, Union

from openpyxl import load_workbook

from ...report_source import report_source_identity
from .models import IndicatorRecord, ReportAuditConfig, SourceRef
from .unit_conversion import AMOUNT_DATA_TYPES, convert_amount


class ImportErrorV2(ValueError):
    pass


# 快速读取引擎（python-calamine 0.3.1，Rust 原生扩展）。发行包冻结自带；
# 环境缺少该模块（如部分兼容构建）时自动回落 xlrd/openpyxl，不影响结果。
try:  # pragma: no cover - 取决于构建环境
    from python_calamine import load_workbook as _calamine_load
except ImportError:  # pragma: no cover
    _calamine_load = None


NAME_CODE = "20201001"
ID_CODE = "20201002"


def _code(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value or "").strip()
    # 兼容字符串形态的浮点指标代码（"20201028.0"），与 float 读出形态归一，
    # 避免与指标参照（"20201028"）lookup miss 后误按普通数值处理。
    if re.fullmatch(r"\d+\.0+", text):
        return str(int(float(text)))
    return text


def _is_yes(value: Any) -> bool:
    return str(value or "").strip().casefold() in {"是", "1", "true", "y", "yes"}


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(",", ""))
    except ValueError:
        return None


def _decimal_number(value: Any) -> Decimal | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    try:
        number = value if isinstance(value, Decimal) else Decimal(str(value).strip().replace(",", ""))
    except (InvalidOperation, TypeError, ValueError):
        return None
    return number if number.is_finite() else None


def _period(date: str) -> str:
    text = str(date or "")
    digits = "".join(ch for ch in text if ch.isdigit())
    return f"{digits[:4]}-{digits[4:6]}" if len(digits) >= 6 else text


def _normalize_fast_value(value: Any, *, xls: bool) -> Any:
    """把 calamine 的取值形态对齐传统引擎，保证两套读取逐值一致。

    .xls 对齐 xlrd：整型数值与布尔都还原成 xlrd 的 float/int 形态
    （xlrd 布尔读出 1/0，布尔经 _number 语义与数值不同，不能漂移）。
    .xlsx 对齐 openpyxl：整型 float 还原为 int（Excel 缓存为最简整数
    文本时 openpyxl 给 int）、空串对齐 None、date 补成零点 datetime。
    已知有意的边界差异：.xls 值列的日期单元格 xlrd 读作序列数 float
    参与数值计算，快速引擎读作 date 走非数值路径，后者更符合业务语义。
    """
    if xls:
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, int):
            return float(value)
        return value
    if isinstance(value, float) and value.is_integer() and abs(value) <= 2 ** 53:
        return int(value)
    if value == "":
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime(value.year, value.month, value.day)
    return value


def _open_fast(path: Path, warnings: list[str] | None):
    """Calamine 快速读取两期报表文件；不可用或失败时返回 None 走传统引擎。

    引擎缺失（兼容构建未打包）时静默回退；引擎在具体文件上失败时逐文件
    提示并回退，结果与直接使用传统引擎一致。
    """
    if _calamine_load is None:
        return None
    try:
        book = _calamine_load(path)
        names = list(book.sheet_names)
        xls = path.suffix.casefold() == ".xls"
        matrices = {}
        for name in names:
            sheet = book.get_sheet_by_name(name)
            matrix = sheet.to_python()
            # to_python 会裁掉前导空行；用 start（首个非空行的 0 基行号）
            # 补回等量空行，保证行号与 xlrd/openpyxl 逐行对齐——
            # rows[3:] 的表头跳过和 SourceRef 行号都依赖绝对行位。
            leading = sheet.start[0] if isinstance(sheet.start, tuple) else 0
            matrices[name] = (
                [[None, None, None] for _ in range(leading)]
                + [[_normalize_fast_value(cell, xls=xls) for cell in row[:3]] for row in matrix]
            )
    except Exception as exc:  # noqa: BLE001 - 单文件回退，任何异常都不中断导入
        if warnings is not None:
            warnings.append(f"快速读取引擎处理失败，已用兼容引擎重读：{path.name}（{exc}）")
        return None

    def rows(name: str) -> list[list[Any]]:
        matrix = matrices[name]
        return [row + [None] * (3 - len(row)) if len(row) < 3 else row for row in matrix]

    return names, rows, lambda: None


def _open(path: Path, warnings: list[str] | None = None):
    fast = _open_fast(path, warnings)
    if fast is not None:
        return fast
    if path.suffix.casefold() == ".xls":
        import xlrd
        book = xlrd.open_workbook(str(path))
        def rows(name: str):
            sheet = book.sheet_by_name(name)
            return [[sheet.cell_value(r, c) if c < sheet.ncols else None for c in range(3)] for r in range(sheet.nrows)]
        return list(book.sheet_names()), rows, lambda: None
    book = load_workbook(path, read_only=True, data_only=True)
    def rows(name: str):
        return [list(row) for row in book[name].iter_rows(min_col=1, max_col=3, values_only=True)]
    return list(book.sheetnames), rows, book.close


def _data_rows(rows: list[list[Any]]) -> list[tuple[int, str, str, Any]]:
    result = []
    for position, row in enumerate(rows[3:], start=4):
        if not row:
            continue
        code = _code(row[0] if len(row) > 0 else None)
        if not code:
            continue
        result.append((position, code, str(row[1] if len(row) > 1 and row[1] is not None else "").strip(), row[2] if len(row) > 2 else None))
    return result


@dataclass
class ImportResult:
    records: list[IndicatorRecord] = field(default_factory=list)
    duplicates: list[tuple[IndicatorRecord, IndicatorRecord]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def by_logical_key(self) -> dict[tuple[str, str, str], IndicatorRecord]:
        """重复项由 validation 处理；这里保留首条，绝不由后项静默覆盖。"""
        result: dict[tuple[str, str, str], IndicatorRecord] = {}
        for record in self.records:
            result.setdefault(record.logical_key, record)
        return result


def import_period_directory(
    directory: Path, *, config: ReportAuditConfig, label: str,
    on_progress: Callable[[int, int, str], None] | None = None,
) -> ImportResult:
    if not directory.is_dir():
        raise ImportErrorV2(f"{label}目录不存在：{directory}")
    files = [p for p in sorted(directory.iterdir()) if p.is_file() and p.suffix.casefold() in {".xls", ".xlsx"} and not p.name.startswith("~$")]
    if not files:
        raise ImportErrorV2(f"{label}目录没有可读取的 .xls/.xlsx 文件：{directory}")
    result = ImportResult()
    seen: dict[tuple[str, str, str], IndicatorRecord] = {}
    pending: list[tuple[Path, str, int, str, str, Any, str, str, str]] = []
    source_paths: dict[Path, str] = {}
    for file_index, path in enumerate(files):
        if on_progress:
            on_progress(file_index, len(files), path.name)
        if not report_source_identity(path, "").recognized:
            result.warnings.append(f"无法识别文件名，已跳过：{path.name}")
            continue
        # 同一工作簿的全部指标共用已解析路径，避免每条指标重复访问文件系统。
        source_paths[path] = str(path.resolve())
        sheets, read_rows, close = _open(path, result.warnings)
        try:
            for sheet_name in sheets:
                identity = report_source_identity(path, sheet_name)
                data = _data_rows(read_rows(sheet_name))
                for source_row, code, name, value in data:
                    pending.append((path, sheet_name, source_row, identity.institution_id, identity.institution_name, value, code, name, identity.form_code + "\x00" + identity.data_date))
        finally:
            close()
    if on_progress:
        on_progress(len(files), len(files), f"正在整理{label}指标")
    for path, sheet, row, raw_id, raw_name, value, code, name, form_date in pending:
        form, date = form_date.split("\x00", 1)
        profile = config.institutions_by_id.get(raw_id)
        if profile is None:
            candidate = config.institutions_by_name.get(raw_name)
            if candidate and (not raw_id or not candidate.social_credit_code or candidate.social_credit_code == raw_id):
                profile = candidate
        # 文件名/表名给出的机构代码和名称稳定优先；配置只补充缺失名称及分类。
        # 20201001/20201002 仍作为普通指标读取，正文空值/占位符不改变机构身份。
        institution_id = raw_id or (profile.institution_id if profile else raw_name)
        institution_name = raw_name or (profile.institution_name if profile else raw_id)
        indicator = config.indicators.get(code)
        expected_form = indicator.form_code if indicator else ""
        if expected_form and form and expected_form != form:
            result.warnings.append(f"表单归属不一致：{path.name}/{sheet} 第{row}行，声明表单 {form}，指标 {code} 配置归属 {expected_form}")
        is_text = bool(indicator and indicator.data_type == "文字")
        is_amount = bool(indicator and indicator.data_type in AMOUNT_DATA_TYPES)
        numeric = _decimal_number(value) if is_amount else _number(value)
        if is_text or numeric is None:
            normalized = str(value or "").strip() if is_text else value
            normalized_unit = ""
        elif is_amount:
            try:
                normalized = convert_amount(
                    numeric, config.unit_settings.source_unit,
                    config.unit_settings.output_file_unit,
                )
            except ValueError as exc:
                raise ImportErrorV2(
                    f"{path.name}/{sheet} 第{row}行指标 {code} 无法转换金额到输出单位：{exc}"
                ) from exc
            normalized_unit = config.unit_settings.output_file_unit
        else:
            # 百分数、个数及文字属性原值参与业务计算，不经过金额换算。
            normalized = numeric
            normalized_unit = ""
        record = IndicatorRecord(
            institution_id=institution_id, institution_name=institution_name,
            social_credit_code=raw_id or (profile.social_credit_code if profile else ""),
            institution_type=profile.institution_type if profile else "",
            handling_branch=profile.handling_branch if profile else "", region=profile.region if profile else "",
            form_code=form or expected_form, indicator_code=code, indicator_name=name or (indicator.indicator_name if indicator else ""),
            period=_period(date), source_date=date, value=value,
            # 保留配置的数据属性；百分数需在环比引擎按百分点差异处理，不能被统一降为“数值”。
            value_type=indicator.data_type if indicator and indicator.data_type else ("文字" if is_text else "数值"),
            source_unit=config.unit_settings.source_unit if is_amount else "原始单位",
            normalized_value=normalized, normalized_unit=normalized_unit,
            source=SourceRef(source_paths[path], sheet, row),
        )
        previous = seen.get(record.logical_key)
        if previous is not None:
            result.duplicates.append((previous, record))
            continue
        seen[record.logical_key] = record
        result.records.append(record)
    return result


@dataclass(frozen=True)
class ExternalValue:
    row_label: str
    indicator_name: str
    value: float | Decimal
    source: SourceRef
    fallback_reason: str = ""
    data_date: str = ""


ExternalValueKey = Union[Tuple[str, str], Tuple[str, str, str]]


class ExternalValues(dict):
    """保持原字典接口，并记录格式、行标签、指标列与可继续处理的数据提示。"""

    def __init__(
        self, *, time_series: bool = False, reader_engine: str = "",
        reader_version: str = "", reader_fallback_reason: str = "",
        reader_elapsed_seconds: float = 0.0,
    ) -> None:
        super().__init__()
        self.time_series = time_series
        self.reader_engine = reader_engine
        self.reader_version = reader_version
        self.reader_fallback_reason = reader_fallback_reason
        self.reader_elapsed_seconds = reader_elapsed_seconds
        self.available_dates: set[str] = set()
        self.date_labels: set[tuple[str, str]] = set()
        self.row_labels: set[str] = set()
        self.indicator_names: set[str] = set()
        self.blank_values: set[ExternalValueKey] = set()
        self.warnings: list[str] = []

    def add_warning(self, message: str) -> None:
        if message not in self.warnings:
            self.warnings.append(message)


def normalize_external_date(value: Any) -> str:
    """将 Excel 日期或完整日期文本统一为 YYYY-MM-DD；不猜测月份对应哪一天。"""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool) or value in (None, ""):
        return ""
    if isinstance(value, (int, float, Decimal)):
        try:
            if value != int(value):
                return ""
            text = str(int(value))
        except (ValueError, OverflowError):
            return ""
    else:
        text = str(value).strip()
    # strptime 的 %m/%d 允许单个数字，必须先拦住仅年月等短数字，
    # 避免将 202611（2026年11月）误读成 2026-01-01。
    if text.isdecimal() and len(text) != 8:
        return ""
    for pattern in ("%Y%m%d", "%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError:
        return ""


# Internal-only key for the explicit zero fallback used when an external
# workbook has no detectable institution row-label column.
MISSING_INSTITUTION_LABEL_FALLBACK_KEY = (
    "__S2_MISSING_INSTITUTION_ROW_LABEL__",
    "__S2_ZERO_FALLBACK__",
)
MISSING_INSTITUTION_LABEL_FALLBACK_REASON = "机构行标签列未识别，外部数按0参与比较"

_EXTERNAL_METADATA_HEADERS = {
    "数据日期", "报表名称", "报表代码", "版本号", "锁定信息", "报表组", "行顺序号", "行名称",
}


def _external_headers(row: list[Any]) -> list[str]:
    headers = [
        str(value if value is not None else "").strip().lstrip("\ufeff").strip()
        for value in row
    ]
    # Excel 导出中的 null/NULL 是空表头占位，不是可供外部规则匹配的指标名。
    return ["" if title.casefold() == "null" else title for title in headers]


def _has_external_value(value: Any) -> bool:
    return value is not None and not (isinstance(value, str) and not value.strip())


def _last_nonempty_row(matrix: list[list[Any]], start: int) -> int:
    return max(
        (index for index, row in enumerate(matrix) if index >= start and any(_has_external_value(value) for value in row)),
        default=start - 1,
    )


def _external_value_preview(value: Any, *, limit: int = 80) -> str:
    text = str(value).replace("\r", " ").replace("\n", " ")
    return text if len(text) <= limit else text[:limit - 3] + "..."


def _read_time_series_external(
    path: Path, matrix: list[list[Any]], header_row: int,
) -> ExternalValues:
    headers = _external_headers(matrix[header_row])
    for field_name in ("数据日期", "行名称"):
        if headers.count(field_name) != 1:
            raise ImportErrorV2(f"集中系统数据第{header_row + 1}行：时序格式必须有且只有一列“{field_name}”")
    date_col, label_col = headers.index("数据日期"), headers.index("行名称")
    indicators: dict[str, int] = {}
    for col, title in enumerate(headers):
        if not title or title in _EXTERNAL_METADATA_HEADERS:
            continue
        if title in indicators:
            raise ImportErrorV2(f"集中系统数据第{header_row + 1}行：时序格式的指标列“{title}”重复，无法确定取数列")
        indicators[title] = col
    if not indicators:
        raise ImportErrorV2(f"集中系统数据第{header_row + 1}行：时序格式缺少指标列")

    values = ExternalValues(time_series=True)
    values.indicator_names.update(indicators)
    seen_rows: dict[tuple[str, str], int] = {}
    source_file = str(path.resolve())
    last_nonempty_row = _last_nonempty_row(matrix, header_row + 1)
    for row_number, row in enumerate(matrix[header_row + 1:], start=header_row + 2):
        label = str(row[label_col] or "").strip() if label_col < len(row) else ""
        populated_indicators = [
            title for title, col in indicators.items()
            if col < len(row) and _has_external_value(row[col])
        ]
        if not label:
            if populated_indicators:
                titles = "、".join(f"“{title}”" for title in populated_indicators)
                values.add_warning(
                    f"集中系统数据第{row_number}行：行名称为空，但指标列{titles}有数据；"
                    "本行无法匹配报表项目，未参与外部核对。请检查行名称是否被误删。"
                )
            elif not any(_has_external_value(value) for value in row) and row_number - 1 < last_nonempty_row:
                values.add_warning(
                    f"集中系统数据第{row_number}行：行名称及指标数据均为空，已忽略该行；"
                    "如属导出空行，可忽略。"
                )
            continue
        raw_date = row[date_col] if date_col < len(row) else None
        data_date = normalize_external_date(raw_date)
        if not data_date:
            date_display = "空白" if raw_date is None or (isinstance(raw_date, str) and not raw_date.strip()) else f"“{_external_value_preview(raw_date)}”"
            values.add_warning(
                f"集中系统数据第{row_number}行、行名称“{label}”：数据日期{date_display}为空或无效，"
                "本行未参与时序匹配；请检查数据日期是否误删或格式错误。"
            )
            continue
        row_key = (data_date, label)
        if row_key in seen_rows:
            raise ImportErrorV2(
                f"集中系统数据第{seen_rows[row_key]}、{row_number}行：日期 {data_date}、报表项目“{label}”重复，"
                "请保留本轮使用的唯一记录，不自动合计或覆盖不同报表版本"
            )
        seen_rows[row_key] = row_number
        values.available_dates.add(data_date)
        values.date_labels.add(row_key)
        values.row_labels.add(label)
        source = SourceRef(source_file, "集中系统数据", row_number)
        blank_titles: list[str] = []
        non_numeric_titles: list[str] = []
        for title, col in indicators.items():
            raw_value = row[col] if col < len(row) else None
            if not _has_external_value(raw_value):
                values.blank_values.add((data_date, label, title))
                blank_titles.append(f"“{title}”为空白")
                number = Decimal(0)
            else:
                number = _decimal_number(raw_value)
                if number is None:
                    non_numeric_titles.append(
                        f"“{title}”内容“{_external_value_preview(raw_value)}”不是数值"
                    )
                    number = Decimal(0)
            values[(data_date, label, title)] = ExternalValue(
                label, title, number, source, data_date=data_date,
            )
        if blank_titles or non_numeric_titles:
            details = "；".join([*blank_titles, *non_numeric_titles])
            values.add_warning(
                f"集中系统数据第{row_number}行、报表项目“{label}”：{details}，已按 0 参与外部核对；"
                "请核查是否误删或改动了数据。"
            )
    return values


def _open_external(path: Path):
    """打开大集中数据文件，返回 (表名列表, rows(name)→矩阵, close)。

    仅支持 .xlsx：机构映射优先读取「机构参照表」，兼容旧「参照表」；
    没有映射表时由外部核对引擎回退到 S2 配置。当前 .xls/.csv 不满足此工作表契约。
    """
    if path.suffix.casefold() != ".xlsx":
        raise ImportErrorV2(
            f"大集中数据仅支持 .xlsx 工作簿：{path.name}"
            "（.xls/.csv 会丢失「机构参照表/参照表」机构映射，不支持）")
    book = load_workbook(path, read_only=True, data_only=True)
    def rows(name: str):
        return [list(row) for row in book[name].iter_rows(values_only=True)]
    return list(book.sheetnames), rows, book.close


def _open_external_calamine(path: Path):
    """Calamine 的 Rust 原生读取路径；保留工作表名与原始行号。"""
    if _calamine_load is None:
        raise ImportErrorV2("未安装 python-calamine")
    book = _calamine_load(path)
    names = list(book.sheet_names)
    matrices: dict[str, list[list[Any]]] = {}
    for name in names:
        sheet = book.get_sheet_by_name(name)
        matrix = sheet.to_python()
        # Calamine 会裁去首个非空单元格之前的空行；补回行位，保证数据行号、
        # 表头定位和 SourceRef 与 OpenPyXL 读取一致。列位置由 to_python 保留。
        start = sheet.start
        leading_rows = start[0] if isinstance(start, tuple) else 0
        matrices[name] = (
            [[None] for _ in range(leading_rows)]
            + [[_normalize_fast_value(value, xls=False) for value in row] for row in matrix]
        )

    def rows(name: str):
        return [list(row) for row in matrices[name]]

    # 当前 python-calamine 的工作簿对象没有需要用户关闭的文件句柄；保留对象
    # 到本次导入结束，避免不同后端版本在惰性读取时提前释放底层资源。
    def close():
        _ = book

    return names, rows, close


def _external_reader_version(engine: str) -> str:
    package = "python-calamine" if engine == "Calamine" else "openpyxl"
    try:
        return package_version(package)
    except PackageNotFoundError:
        return "未知"


def _read_external_org_map(sheet_name: str, rows: list[list[Any]]) -> dict[str, str]:
    if not rows:
        return {}
    headers = [str(value or "").strip() for value in rows[0]]
    code_col = next(
        (headers.index(name) for name in ("社会信用代码", "统一社会信用代码") if name in headers),
        None,
    )
    report_item_col = headers.index("报表项目") if "报表项目" in headers else None
    if code_col is None or report_item_col is None:
        raise ImportErrorV2(
            f"大集中文件“{sheet_name}”缺少必要表头：社会信用代码（或统一社会信用代码）、报表项目"
        )
    disabled_col = headers.index("禁用") if "禁用" in headers else None
    org_map: dict[str, str] = {}
    for row in rows[1:]:
        if disabled_col is not None and disabled_col < len(row) and _is_yes(row[disabled_col]):
            continue
        code = _code(row[code_col] if code_col < len(row) else None)
        report_item = str(row[report_item_col] or "").strip() if report_item_col < len(row) else ""
        if code and report_item:
            # 同一代码有多条有效映射时，后行覆盖前行，与旧“参照表”保持一致。
            org_map[code] = report_item
    return org_map


def import_external_workbook(
    path: Path, *, reader_engine: str = "AUTO",
) -> tuple[ExternalValues, dict[str, str]]:
    """读取 .xlsx 中的普通报表或时序布局；时序数值按完整日期匹配。

    AUTO 优先用 Calamine；引擎缺失或读文件失败时明确回退 OpenPyXL。手动选定
    的引擎严格执行，不在失败时静默切换。
    """
    if not path.is_file():
        raise ImportErrorV2(f"大集中文件不存在：{path}")
    if path.suffix.casefold() != ".xlsx":
        raise ImportErrorV2(
            f"大集中数据仅支持 .xlsx 工作簿：{path.name}"
            "（.xls/.csv 会丢失「机构参照表/参照表」机构映射，不支持）"
        )
    mode = str(reader_engine or "AUTO").strip().upper()
    if mode not in {"AUTO", "CALAMINE", "OPENPYXL"}:
        raise ImportErrorV2(f"未知的大集中数据读取方式：{reader_engine}")
    started = time.perf_counter()
    fallback_reason = ""
    if mode == "OPENPYXL":
        sheets, read_rows, close = _open_external(path)
        actual_engine = "OpenPyXL"
    else:
        try:
            sheets, read_rows, close = _open_external_calamine(path)
            actual_engine = "Calamine"
        except Exception as exc:  # noqa: BLE001 - AUTO 可按兼容路径重试
            if mode == "CALAMINE":
                raise ImportErrorV2(
                    f"已指定 Calamine 读取大集中数据，但读取失败：{exc}"
                ) from exc
            fallback_reason = f"Calamine 读取失败，自动回退 OpenPyXL：{exc}"
            if _calamine_load is None:
                fallback_reason = "未安装 python-calamine，自动回退 OpenPyXL"
            sheets, read_rows, close = _open_external(path)
            actual_engine = "OpenPyXL"
    try:
        if "集中系统数据" not in sheets:
            raise ImportErrorV2("大集中文件缺少“集中系统数据”工作表")
        matrix = read_rows("集中系统数据")
        header_row = next((i for i, row in enumerate(matrix) if any(isinstance(v, str) and v.strip() for v in row)), None)
        if header_row is None:
            raise ImportErrorV2("集中系统数据未找到表头")
        # 依据字段识别，允许表头前留空行/说明行。不能用“第一行非空”
        # 判断时序格式：普通格式也可能直接从第一行开始。
        time_header_row = next(
            (i for i, row in enumerate(matrix[:20]) if {"数据日期", "行名称"}.issubset(_external_headers(row))),
            None,
        )
        plain_header_row = next(
            (
                i for i, row in enumerate(matrix[:20])
                if {"行名称", "报表项目"}.intersection(_external_headers(row)[:4])
            ),
            None,
        )
        if time_header_row is not None:
            values = _read_time_series_external(path, matrix, time_header_row)
        else:
            partial_time_header_row = next(
                (i for i, row in enumerate(matrix[:20]) if "数据日期" in _external_headers(row)),
                None,
            )
            if partial_time_header_row is not None:
                raise ImportErrorV2(
                    f"集中系统数据第{partial_time_header_row + 1}行：时序表头不完整，必须同时包含“数据日期”和“行名称”"
                )
            values = _read_plain_external(path, matrix, plain_header_row if plain_header_row is not None else header_row)
        reference_sheet = next(
            (name for name in ("机构参照表", "参照表") if name in sheets),
            None,
        )
        org_map = (
            _read_external_org_map(reference_sheet, read_rows(reference_sheet))
            if reference_sheet is not None
            else {}
        )
        values.reader_engine = actual_engine
        values.reader_version = _external_reader_version(actual_engine)
        values.reader_fallback_reason = fallback_reason
        values.reader_elapsed_seconds = time.perf_counter() - started
        return values, org_map
    finally:
        close()


def _read_plain_external(path: Path, matrix: list[list[Any]], header_row: int) -> ExternalValues:
    """保留原普通格式的取数和缺少机构行标签列时的零值兜底。"""
    header = matrix[header_row]
    headers = _external_headers(header)
    label_col = next((
        c for c, title in enumerate(headers[:4])
        if title in {"行名称", "报表项目"}
    ), None)
    if label_col is None:
        label_col = next((
            c for c in range(min(4, len(header)))
            if not headers[c]
            and any(
                isinstance(r[c] if c < len(r) else None, str) and str(r[c]).strip()
                for r in matrix[header_row + 1:header_row + 5]
            )
        ), None)
    values = ExternalValues()
    if label_col is None:
        # 没有机构行标签时，任何行都无法可靠匹配到报表机构；按用户要求
        # 继续执行外部核对，并让规则明确使用 0，而不是错误退出。
        values[MISSING_INSTITUTION_LABEL_FALLBACK_KEY] = ExternalValue(
            "",
            "",
            Decimal(0),
            SourceRef(str(path.resolve()), "集中系统数据", header_row + 1),
            MISSING_INSTITUTION_LABEL_FALLBACK_REASON,
        )
    else:
        indicators = [
            (col, title)
            for col, title in enumerate(headers)
            if title
            and col != label_col
            and title not in _EXTERNAL_METADATA_HEADERS
        ]
        for _col, title in indicators:
            values.indicator_names.add(title)
        last_nonempty_row = _last_nonempty_row(matrix, header_row + 1)
        for row_number, row in enumerate(matrix[header_row + 1:], start=header_row + 2):
            label = str(row[label_col] or "").strip() if label_col < len(row) else ""
            populated_indicators = [
                title for col, title in indicators
                if col < len(row) and _has_external_value(row[col])
            ]
            if not label:
                if populated_indicators:
                    titles = "、".join(f"“{title}”" for title in populated_indicators)
                    values.add_warning(
                        f"集中系统数据第{row_number}行：行名称为空，但指标列{titles}有数据；"
                        "本行无法匹配报表项目，未参与外部核对。请检查行名称是否被误删。"
                    )
                elif not any(_has_external_value(value) for value in row) and row_number - 1 < last_nonempty_row:
                    values.add_warning(
                        f"集中系统数据第{row_number}行：行名称及指标数据均为空，已忽略该行；"
                        "如属导出空行，可忽略。"
                    )
                continue
            values.row_labels.add(label)
            blank_titles: list[str] = []
            non_numeric_titles: list[str] = []
            for col, title in indicators:
                raw_value = row[col] if col < len(row) else None
                if not _has_external_value(raw_value):
                    values.blank_values.add((label, title))
                    blank_titles.append(f"“{title}”为空白")
                    number = Decimal(0)
                else:
                    number = _decimal_number(raw_value)
                    if number is None:
                        non_numeric_titles.append(
                            f"“{title}”内容“{_external_value_preview(raw_value)}”不是数值"
                        )
                        number = Decimal(0)
                values[(label, title)] = ExternalValue(
                    label, title, number, SourceRef(str(path.resolve()), "集中系统数据", row_number)
                )
            if blank_titles or non_numeric_titles:
                details = "；".join([*blank_titles, *non_numeric_titles])
                values.add_warning(
                    f"集中系统数据第{row_number}行、报表项目“{label}”：{details}，已按 0 参与外部核对；"
                    "请核查是否误删或改动了数据。"
                )
    return values
