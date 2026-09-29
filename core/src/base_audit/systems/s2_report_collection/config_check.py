"""S2 报表采集配置的只读结构与规则检查。

检查器与运行时共用 :func:`load_config` 的字段解析，不会改写工作簿；
它返回可供 Flask、pywebview 和测试复用的简单字典，避免两个外壳各自实现
一套配置判断。
"""

from __future__ import annotations

import math
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import _code, _number, _severity, _text, _yes, load_config


REQUIRED_SHEETS = ("指标参照", "机构参照", "环比策略", "校验规则", "外部核对规则", "通用设置")
DATA_TYPES = {"余额", "累发", "百分数", "个数", "文字"}
SEVERITIES = {"", "提示", "关注", "严重"}
CHECK_DEFINITIONS = (
    ("file_read", "配置文件可读取"),
    ("required_sheets", "必需工作表齐全"),
    ("indicators", "指标参照字段与记录"),
    ("period_strategies", "环比策略区间与指标引用"),
    ("runtime_parse", "全部配置可被正式运行时解析"),
)
CHECK_LIMITATIONS = (
    "此项检查只读取 S2 配置工作簿，不读取本期、上期报表或大集中外部数据文件。",
    "此项检查不实际执行审核规则，也不判断报表结果是否符合业务预期。",
    "检查通过表示配置结构和内容可由当前运行时读取，不代表真实数据审核已经通过。",
)


def _issue(report: dict[str, Any], level: str, message: str, *, sheet: str = "", row: int | None = None) -> None:
    item = {"level": level, "message": message}
    if sheet:
        item["sheet"] = sheet
    if row is not None:
        item["row"] = row
    report["issues"].append(item)
    if level == "blocking":
        report["blocking_errors"] += 1
    else:
        report["warnings"] += 1


def _set_check(report: dict[str, Any], check_id: str, status: str, detail: str) -> None:
    for check in report["checks"]:
        if check["id"] == check_id:
            check["status"] = status
            check["detail"] = detail
            return
    raise KeyError(f"unknown S2 config check: {check_id}")


def _finish_check(report: dict[str, Any], check_id: str, issue_start: int, success_detail: str) -> None:
    issues = report["issues"][issue_start:]
    blocking = sum(issue["level"] == "blocking" for issue in issues)
    warnings = sum(issue["level"] == "warning" for issue in issues)
    if blocking:
        _set_check(report, check_id, "failed", f"发现 {blocking} 个阻断问题，详见下方问题明细。")
    elif warnings:
        _set_check(report, check_id, "warning", f"{success_detail}；另有 {warnings} 条警告，详见下方问题明细。")
    else:
        _set_check(report, check_id, "passed", success_detail)


def _header(sheet) -> dict[str, int]:
    return {_text(value): idx for idx, value in enumerate(next(sheet.iter_rows(values_only=True), ())) if _text(value)}


def _cell(row: tuple[Any, ...], headers: dict[str, int], name: str) -> Any:
    idx = headers.get(name)
    return row[idx] if idx is not None and idx < len(row) else None


def _check_ranges(report: dict[str, Any], rows: list[dict[str, Any]], group: str) -> None:
    enabled = [row for row in rows if not row["disabled"]]
    if not enabled:
        _issue(report, "blocking", f"环比策略组 {group} 没有启用的策略行", sheet="环比策略")
        return
    enabled.sort(key=lambda row: float("-inf") if row["lower"] is None else row["lower"])
    if enabled[0]["lower"] is not None:
        _issue(report, "blocking", f"环比策略组 {group} 左侧没有覆盖负无穷（首行下限应为空）", sheet="环比策略", row=enabled[0]["row"])
    if enabled[-1]["upper"] is not None:
        _issue(report, "blocking", f"环比策略组 {group} 右侧没有覆盖正无穷（末行上限应为空）", sheet="环比策略", row=enabled[-1]["row"])
    for left, right in zip(enabled, enabled[1:]):
        if left["upper"] is None or right["lower"] is None:
            _issue(report, "blocking", f"环比策略组 {group} 存在重叠：{left['id']} 与 {right['id']}", sheet="环比策略", row=right["row"])
            continue
        if left["upper"] < right["lower"]:
            _issue(report, "blocking", f"环比策略组 {group} 存在空档：{left['id']} 与 {right['id']}", sheet="环比策略", row=right["row"])
        elif left["upper"] > right["lower"]:
            _issue(report, "blocking", f"环比策略组 {group} 存在重叠：{left['id']} 与 {right['id']}", sheet="环比策略", row=right["row"])


def check_s2_config(path: Path) -> dict[str, Any]:
    """只读检查 S2 配置，返回稳定的 UI/测试报告结构。"""
    target = Path(path)
    report: dict[str, Any] = {
        "path": str(target), "passed": False, "blocking_errors": 0, "warnings": 0,
        "issues": [], "sheet_stats": [], "indicator_count": 0, "strategy_count": 0,
        "message": "",
        "checks": [
            {"id": check_id, "name": name, "status": "not_run", "detail": "未执行（前置检查未完成）。"}
            for check_id, name in CHECK_DEFINITIONS
        ],
        "limitations": list(CHECK_LIMITATIONS),
    }
    if not target.is_file():
        _issue(report, "blocking", f"配置文件不存在：{target}")
        _set_check(report, "file_read", "failed", "配置文件不存在，无法读取。")
        report["message"] = f"审核配置存在 {report['blocking_errors']} 个阻断错误，不能执行。"
        return report
    try:
        book = load_workbook(target, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - UI 需要可理解的失败原因
        _issue(report, "blocking", f"配置文件无法读取：{exc}")
        _set_check(report, "file_read", "failed", "配置文件无法打开，详见下方问题明细。")
        report["message"] = f"审核配置存在 {report['blocking_errors']} 个阻断错误，不能执行。"
        return report
    try:
        _set_check(report, "file_read", "passed", "配置工作簿已成功打开，并以只读方式检查。")
        names = set(book.sheetnames)
        sheet_check_start = len(report["issues"])
        for name in REQUIRED_SHEETS:
            if name not in names:
                _issue(report, "blocking", f"缺少工作表：{name}")
        for sheet in book.worksheets:
            max_row, max_column = sheet.max_row, sheet.max_column
            if max_row is None or max_column is None:
                # 一些兼容型 XLSX 没有 worksheet dimension 元素；流式解析仍
                # 能读到数据，但 openpyxl 会返回 None，故从实际行流计算统计。
                rows = list(sheet.iter_rows(values_only=True))
                max_row = len(rows)
                max_column = max((len(row) for row in rows), default=0)
            report["sheet_stats"].append({"name": sheet.title, "rows": max(max_row - 1, 0), "columns": max_column})
        _finish_check(
            report, "required_sheets", sheet_check_start,
            f"已确认 {len(REQUIRED_SHEETS)} 张必需工作表存在",
        )
        if report["blocking_errors"]:
            report["message"] = f"审核配置存在 {report['blocking_errors']} 个阻断错误，不能执行。"
            return report

        indicator_check_start = len(report["issues"])
        indicator_sheet = book["指标参照"]
        indicators = _header(indicator_sheet)
        inline_units = {"源数据单位", "大集中数据单位"}.intersection(indicators)
        if inline_units:
            _issue(
                report, "blocking",
                f"单位统一配置应放在本工作簿的“通用设置”表；请从“指标参照”删除旧字段：{'、'.join(sorted(inline_units))}",
                sheet="指标参照",
            )
        required_indicator_headers = ("指标代码", "指标名称", "表单代码", "数据属性", "环比启用", "环比策略组", "最小变动值", "禁用")
        for name in required_indicator_headers:
            if name not in indicators:
                _issue(report, "blocking", f"“指标参照”缺少字段：{name}", sheet="指标参照")
        seen_codes: set[str] = set()
        referenced_groups: set[str] = set()
        for row_number, row in enumerate(indicator_sheet.iter_rows(min_row=2, values_only=True), start=2):
            code = _code(_cell(row, indicators, "指标代码"))
            if not code:
                continue
            report["indicator_count"] += 1
            if code in seen_codes:
                _issue(report, "blocking", f"指标代码重复：{code}", sheet="指标参照", row=row_number)
            seen_codes.add(code)
            data_type = _text(_cell(row, indicators, "数据属性"))
            if data_type not in DATA_TYPES:
                _issue(report, "blocking", f"指标 {code} 数据属性无效：{data_type or '空白'}", sheet="指标参照", row=row_number)
            enabled_raw = _cell(row, indicators, "环比启用")
            if _text(enabled_raw) not in {"是", "否"}:
                _issue(report, "blocking", f"指标 {code} 环比启用必须为“是”或“否”：{enabled_raw!r}", sheet="指标参照", row=row_number)
            disabled_raw = _cell(row, indicators, "禁用")
            if _text(disabled_raw) not in {"是", "否", ""}:
                _issue(report, "blocking", f"指标 {code} 禁用必须为“是”或“否”：{disabled_raw!r}", sheet="指标参照", row=row_number)
            if not _yes(enabled_raw):
                continue
            group = _text(_cell(row, indicators, "环比策略组"))
            min_raw = _cell(row, indicators, "最小变动值")
            if data_type == "文字":
                if group != "不适用" or _text(min_raw) != "不适用":
                    _issue(report, "blocking", f"文字指标 {code} 的环比策略组和最小变动值必须均为“不适用”", sheet="指标参照", row=row_number)
            else:
                if not group or group == "不适用":
                    _issue(report, "blocking", f"数值指标 {code} 未配置有效环比策略组", sheet="指标参照", row=row_number)
                else:
                    referenced_groups.add(group)
                min_value = _number(min_raw)
                if min_value is None or not math.isfinite(min_value) or min_value < 0:
                    _issue(report, "blocking", f"数值指标 {code} 最小变动值必须为非负数：{min_raw!r}", sheet="指标参照", row=row_number)

        _finish_check(
            report, "indicators", indicator_check_start,
            f"已检查 {report['indicator_count']} 条指标记录的必需字段、重复代码、属性、启用值及环比设置",
        )

        strategy_check_start = len(report["issues"])
        strategy_sheet = book["环比策略"]
        strategies = _header(strategy_sheet)
        for name in ("策略组", "规则编号", "数据属性", "下限", "上限", "审核级别", "分类说明", "禁用"):
            if name not in strategies:
                _issue(report, "blocking", f"“环比策略”缺少字段：{name}", sheet="环比策略")
        groups: dict[str, list[dict[str, Any]]] = {}
        seen_pair: set[tuple[str, str]] = set()
        for row_number, row in enumerate(strategy_sheet.iter_rows(min_row=2, values_only=True), start=2):
            group, rule_id = _text(_cell(row, strategies, "策略组")), _text(_cell(row, strategies, "规则编号"))
            if not group and not rule_id:
                continue
            report["strategy_count"] += 1
            if not group or not rule_id:
                _issue(report, "blocking", "策略组与规则编号不能为空", sheet="环比策略", row=row_number)
                continue
            pair = (group, rule_id)
            if pair in seen_pair:
                _issue(report, "blocking", f"策略组+规则编号重复：{group}/{rule_id}", sheet="环比策略", row=row_number)
            seen_pair.add(pair)
            data_type = _text(_cell(row, strategies, "数据属性"))
            if data_type not in DATA_TYPES - {"文字"}:
                _issue(report, "blocking", f"策略 {group}/{rule_id} 数据属性无效：{data_type or '空白'}", sheet="环比策略", row=row_number)
            lower, upper = _number(_cell(row, strategies, "下限")), _number(_cell(row, strategies, "上限"))
            if ((lower is not None and not math.isfinite(lower)) or
                    (upper is not None and not math.isfinite(upper))):
                _issue(report, "blocking", f"策略 {group}/{rule_id} 边界必须是有限数字或空白", sheet="环比策略", row=row_number)
            if lower is not None and upper is not None and upper <= lower:
                _issue(report, "blocking", f"策略 {group}/{rule_id} 上限必须大于下限", sheet="环比策略", row=row_number)
            severity = _severity(_cell(row, strategies, "审核级别"))
            if severity not in SEVERITIES:
                _issue(report, "blocking", f"策略 {group}/{rule_id} 审核级别无效：{severity}", sheet="环比策略", row=row_number)
            disabled_raw = _cell(row, strategies, "禁用")
            if _text(disabled_raw) not in {"是", "否", ""}:
                _issue(report, "blocking", f"策略 {group}/{rule_id} 禁用必须为“是”或“否”：{disabled_raw!r}", sheet="环比策略", row=row_number)
            groups.setdefault(group, []).append({"id": rule_id, "lower": lower, "upper": upper, "disabled": _yes(disabled_raw), "data_type": data_type, "row": row_number})
        for group, rows in groups.items():
            types = {row["data_type"] for row in rows}
            if len(types) != 1:
                _issue(report, "blocking", f"环比策略组 {group} 的数据属性必须一致：{'、'.join(sorted(types))}", sheet="环比策略")
            _check_ranges(report, rows, group)
        missing_groups = sorted(referenced_groups - set(groups))
        for group in missing_groups:
            _issue(report, "blocking", f"指标参照引用了不存在的环比策略组：{group}", sheet="指标参照")
        for group in sorted(set(groups) - referenced_groups):
            _issue(report, "warning", f"环比策略组 {group} 当前没有被任何启用指标引用", sheet="环比策略")

        _finish_check(
            report, "period_strategies", strategy_check_start,
            f"已检查 {report['strategy_count']} 条策略记录的编号、区间覆盖/空档/重叠及指标组引用",
        )

        runtime_check_start = len(report["issues"])
        try:
            loaded = load_config(target)
        except Exception as exc:  # noqa: BLE001
            _issue(report, "blocking", f"配置无法进入运行时解析：{exc}")
        else:
            for warning in loaded.warnings:
                # 解析器的兼容性提示仍展示给用户；任何跳过/无效都阻断，避免启用规则静默丢弃。
                level = "blocking" if any(token in warning for token in ("已跳过", "不存在", "无效", "暂不支持")) else "warning"
                _issue(report, level, warning)
        _finish_check(
            report, "runtime_parse", runtime_check_start,
            "指标参照、机构参照、校验规则、外部核对规则和通用设置均已交给正式配置解析器检查",
        )
    finally:
        book.close()
    report["passed"] = report["blocking_errors"] == 0
    report["message"] = "审核配置检查通过，可执行报表规则。" if report["passed"] else f"审核配置存在 {report['blocking_errors']} 个阻断错误，不能执行。"
    return report


def format_config_report(report: dict[str, Any]) -> str:
    """将结构化报告转换为 UI 可读文本。"""
    lines = [report.get("message", "")]
    lines.append(f"配置文件：{report.get('path', '')}")
    lines.append(f"指标参照：{report.get('indicator_count', 0)} 条；环比策略：{report.get('strategy_count', 0)} 条")
    lines.append("实际检查项目：")
    status_labels = {"passed": "通过", "warning": "有警告", "failed": "未通过", "not_run": "未检查"}
    for check in report.get("checks", []):
        label = status_labels.get(check.get("status"), "未检查")
        lines.append(f"[{label}] {check.get('name', '')}：{check.get('detail', '')}")
    if report.get("limitations"):
        lines.append("检查范围说明：")
        lines.extend(f"- {item}" for item in report["limitations"])
    for stat in report.get("sheet_stats", []):
        lines.append(f"工作表 {stat['name']}：{stat['rows']} 行，{stat['columns']} 列")
    for issue in report.get("issues", []):
        location = issue.get("sheet", "")
        if issue.get("row") is not None:
            location += f" 第{issue['row']}行"
        prefix = "阻断" if issue["level"] == "blocking" else "警告"
        lines.append(f"[{prefix}] {location + '：' if location else ''}{issue['message']}")
    return "\n".join(lines)
