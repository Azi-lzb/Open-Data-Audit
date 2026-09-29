"""大集中统计系统：对外稳定接口（UI/调度层只允许调用这里）。

``run_comparison`` / ``run_cross_period_check`` / ``render_financial_forms``
三个固定入口；内部流水线的模块划分对调用方不可见。
"""

from __future__ import annotations

import time
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

from .config import load_central_config
from .comparison_engine import _append_explanation, build_comparison, sort_comparison_rows
from .comparison_exporter import write_comparison_workbook, write_log_workbook
from .cross_period_engine import DEFAULT_TITLE as _CHECK_TITLE
from .csv_importer import read_central_csv
from .models import CentralDataset, CheckResult, ComparisonResult


def _load_dataset(paths: Iterable[str | Path], *, exempt: set[str], factor: float) -> CentralDataset:
    """读取一期（可多个 CSV，追加合并；VBA 多文件语义）。"""
    combined = CentralDataset()
    for item in paths:
        part = str(item).strip()
        if not part:
            continue
        dataset = read_central_csv(Path(part), unit_factor=factor, exempt_indicators=exempt)
        combined.records.extend(dataset.records)
        combined.issues.extend(dataset.issues)
    dates = {record.record_date for record in combined.records if record.record_date}
    if len(dates) > 1:
        raise ValueError(f"同一期文件数据日期不一致：{sorted(dates)}")
    combined.record_date = next(iter(dates), "")
    combined.build_index()
    return combined


def run_comparison(
    *,
    current_csv: str | Path | list,
    previous_csv: str | Path | list,
    output_path: str | Path,
    config_path: str | Path,
    source_unit: str = "元",
    target_unit: str = "亿元",
    on_step: Callable[[str], None] | None = None,
    write_flow_logs: bool = True,
    rule_engine: str = "v3",
    office_evaluator=None,
    expression_mode: str | None = None,
    expression_backend: str | None = None,
    expression_schema: str | None = None,
) -> ComparisonResult:
    """执行比较：两期导数（CSV/XLSX）→ 比较结果.xlsx。

    ``write_flow_logs`` 与全局“全量运行日志”设置对应；关闭时仍保留结果内的
    业务字段，但不额外落盘运行日志工作簿。

    ``rule_engine``：``v3``（默认，「规则动作」表 + 启动过滤 + 指标索引）或
    当前正式规则引擎固定为 ``v3``（规则动作表）。

    ``expression_mode``/``expression_backend``：表达式 2×2 覆盖（SBE/LAE ×
    PYTHON/OFFICE）；缺省时读配置「运行参数」，再缺省 SBE+PYTHON。

    ``office_evaluator``：表达式求值方式=OFFICE 时的任务级
    ``office_eval.OfficeEvaluationAdapter``（provider 须由调用方一次解析并固定）；
    未提供而配置要求 OFFICE 时明确报错，不静默回退 Python。
    """
    started = time.monotonic()
    say = on_step or (lambda _text: None)
    current_list = [current_csv] if not isinstance(current_csv, list) else current_csv
    previous_list = [previous_csv] if not isinstance(previous_csv, list) else previous_csv
    if not current_list or not str(current_list[0]).strip():
        raise ValueError("执行比较：请先选择本期导数文件")
    if not previous_list or not str(previous_list[0]).strip():
        raise ValueError("执行比较：请先选择上期导数文件")

    config = load_central_config(config_path)
    target_unit = config.param("默认目标单位", target_unit)
    source_unit = config.param("默认源数据单位", source_unit)
    factor = 1.0
    from .comparison_engine import get_unit_factor

    factor = get_unit_factor(source_unit) / get_unit_factor(target_unit)
    exempt = config.exempt_indicators()

    say("读取本期数据")
    current = _load_dataset(current_list, exempt=exempt, factor=factor)
    say(f"本期 {len(current.records)} 行")
    say("读取上期数据")
    previous = _load_dataset(previous_list, exempt=exempt, factor=factor)
    say(f"上期 {len(previous.records)} 行")

    say("两期配对与环比计算")
    rows = build_comparison(current, previous, config, target_unit=target_unit)
    # 排序须在规则富化之前：VBA 的软性规则去重依赖排序后的行序。
    rows = sort_comparison_rows(rows)

    # ---- 规则富化：复杂校验 / 累计与特殊指标 ----
    from .complex_rule_engine import enrich_complex_rules
    from .indicator_rule_engine import enrich_indicator_rules

    if rule_engine != "v3":
        raise ValueError(f"未知规则引擎：{rule_engine}（当前仅支持 v3）")
    from . import expression_backend as _expression_backend_module
    from . import expression_parser as _expression_parser_module
    from . import expression_lae as _expression_lae_module

    _expression_parser_module.reset_cache_stats()
    _expression_lae_module.reset_cache_stats()
    resolved_mode, resolved_backend = _expression_backend_module.resolve_expression_settings(config.run_params)
    expr_mode = (expression_mode or resolved_mode).strip().upper()
    expr_backend = (expression_backend or resolved_backend).strip().upper()
    if (expr_mode, expr_backend) != (resolved_mode, resolved_backend):
        # 显式参数与配置口径不一致时以显式参数为准并校验合法性。
        _expression_backend_module.resolve_expression_settings({
            _expression_backend_module.MODE_PARAM: expr_mode,
            _expression_backend_module.BACKEND_PARAM: expr_backend,
        })
    from .complex_rule_engine import (
        SCHEMA_FIVE_SEGMENT_V1, SCHEMA_LEGACY_8,
        audit_five_segment_keys,
    )
    expression_schema = (expression_schema or "").strip().upper() or SCHEMA_LEGACY_8
    if expression_schema not in (SCHEMA_LEGACY_8, SCHEMA_FIVE_SEGMENT_V1):
        raise ValueError(
            f"未知 3.1 表达式规则语法：{expression_schema!r}"
            "（应为 LEGACY_8 / FIVE_SEGMENT_V1）")
    say(f"表达式语法：{expression_schema}")
    if expression_schema == SCHEMA_FIVE_SEGMENT_V1:
        # 五段取数键唯一性审计（安全门）：同机构+地区桶内 5 段键出现多条
        # 记录时无法确定取哪条——阻止 5 段正式执行，8 段不受影响。
        for label, dataset in (("本期", current), ("上期", previous)):
            if dataset is None:
                continue
            conflicts = audit_five_segment_keys(dataset)
            if conflicts:
                detail = "；".join(
                    f"机构 {c['org']} 地区 {c['region']} 指标 {c['indicator']}"
                    f"（{c['attr']}/{c['currency']}/{c['frequency']}/{c['batch']}）"
                    f" 共 {c['count']} 条（业务类 {'、'.join(c['biz_classes'])}）"
                    for c in conflicts[:10])
                raise ValueError(
                    f"五段式取数键不唯一（{label}共 {len(conflicts)} 组），"
                    "无法确定应使用哪条数据，已阻止五段式执行："
                    + detail
                    + "。请继续使用兼容8段式，或检查数据结构。")
        say("五段式取数键唯一性审计通过")
    exempt = config.exempt_indicators()
    previous_date = previous.record_date if previous is not None else ""
    soft_color = int(config.param("复杂校验软性颜色", "46") or 46)
    fres = {""}
    for dataset in (current, previous):
        if dataset is not None:
            for record in dataset.records:
                fres.add(record.frequency)
                fres.add(record.frequency + record.batch)
    expression_stats: dict = {}
    enrich_complex_rules(
        rows, config,
        current_index=current.key_index,
        previous_index=previous.key_index if previous else {},
        target_unit=target_unit, exempt=exempt, soft_color=soft_color, fres=fres,
        expr_mode=expr_mode, expr_backend=expr_backend,
        expression_dates=(current.record_date, previous_date),
        office_evaluator=office_evaluator,
        stats_out=expression_stats,
        expression_schema=expression_schema,
    )
    from .rule_action_engine import prepare_action_rules

    prepared = prepare_action_rules(config, current_date=current.record_date, fres=fres)
    if not config.action_rules:
        raise ValueError("V3 配置缺少「规则动作」工作表")
    enrich_indicator_rules(
        rows, config,
        current_index=current.key_index,
        previous_index=previous.key_index if previous else {},
        current_date=current.record_date, previous_date=previous_date,
        target_unit=target_unit, exempt=exempt, soft_color=soft_color,
        prepared=prepared,
    )
    # 环比警戒、表达式、累计和特殊指标都可写“是否说明”。统一拆分、去重、
    # 再以同一分隔符合并，避免后执行的规则覆盖或重复前面已经命中的说明。
    for row in rows:
        row.need_explain = _append_explanation("", row.need_explain)
    rows = sort_comparison_rows(rows)

    output_path = Path(output_path)
    say("写出比较结果")
    write_comparison_workbook(rows, output_path, target_unit=target_unit, sheet_date=current.record_date)

    # ---- 运行日志：配置载入统计 + 数据读取 + 配对 + 规则命中 + 数据异常 ----
    def _is_disabled(rule_row) -> bool:
        return str(rule_row.get("禁用") or "").strip() in {"是", "1", "true", "True", "Y", "y"}

    def _rule_count(rule_rows, source: str | None = None) -> tuple[int, int, int]:
        pool = [r for r in rule_rows if not source or r.get("来源") == source]
        total = len(pool)
        disabled = sum(1 for r in pool if _is_disabled(r))
        return total, disabled, total - disabled

    log_rows: list[tuple[str, str]] = []
    for source in ("自定义", "单频", "跨期", "年报", "结转", "跨期核对"):
        total, disabled, enabled = _rule_count(config.complex_rules, source)
        if total:
            log_rows.append(("配置载入", f"复杂校验-{source}：读取 {total} 条，实际载入 {enabled} 条，禁用 {disabled} 条"))
    action_stats = prepared.stats
    log_rows.append(("配置载入", (
        f"规则动作表（V3）：读取 {action_stats.get('loaded', 0)} 条，"
        f"累计 {action_stats.get('accu', 0)} 条；特殊规则启动过滤后 "
        f"{action_stats.get('special', 0)} 条（剔除 {action_stats.get('filtered_out', 0)} 条）")))
    cross_total, cross_disabled, cross_enabled = _rule_count(config.cross_rules)
    log_rows.append(("配置载入", f"本期数值核对规则：读取 {cross_total} 条，实际载入 {cross_enabled} 条，禁用 {cross_disabled} 条"))
    log_rows.append(("配置载入", f"环比警戒档：{len(config.alerts)} 档；单位不转换指标：{len(exempt)} 个"))

    for label, dataset in (("本期", current), ("上期", previous)):
        names = "、".join(Path(item).name for item in
                          ((current_list if label == "本期" else previous_list)))
        log_rows.append(("数据读取", f"{label} {names}：{len(dataset.records)} 行，数据日期 {dataset.record_date or '未知'}"))

    matched = sum(1 for row in rows if row.prev_value is not None)
    cur_only = sum(1 for row in rows if row.remark == "本期有，上期无")
    pre_only = sum(1 for row in rows if row.remark == "本期无，上期有")
    log_rows.append(("两期配对", f"配对 {matched} 行；仅本期有 {cur_only} 行；仅上期有 {pre_only} 行"))

    remark_hits = sum(1 for row in rows if row.remark and row.remark not in ("本期有，上期无", "本期无，上期有"))
    explain_hits = sum(1 for row in rows if row.need_explain)
    process_hits = sum(1 for row in rows if row.process)
    log_rows.append(("规则命中", f"警戒区间 {remark_hits} 行；复杂校验/特殊指标写入是否说明 {explain_hits} 行；计算过程 {process_hits} 行"))

    if expression_stats:
        provider = expression_stats.get("provider", "PYTHON")
        observation = (
            f"表达式求值：mode={expression_stats.get('mode')} backend={expression_stats.get('backend')}"
            f" provider={provider} Office求值 {expression_stats.get('office_evaluations', 0)} 次"
            f"（单元格/重算回退 {expression_stats.get('office_cell_path', 0)}）"
        )
        if expression_stats.get("mode") == "LAE":
            observation += (
                f"；AST 缓存命中 {expression_stats.get('ast_cache_hit_count', 0)}"
                f"/未命中 {expression_stats.get('ast_cache_miss_count', 0)}；"
                f"引用解析 {expression_stats.get('reference_resolve_count', 0)} 次、"
                f"函数求值 {expression_stats.get('function_eval_count', 0)} 次、"
                f"AND 短路 {expression_stats.get('and_short_circuit_count', 0)}、"
                f"OR 短路 {expression_stats.get('or_short_circuit_count', 0)}、"
                f"IF/IFS 跳过分支 {expression_stats.get('if_skipped_branch_count', 0)}"
                f"/{expression_stats.get('ifs_skipped_branch_count', 0)}、"
                f"少求值节点 {expression_stats.get('skipped_ast_node_count', 0)}"
            )
        else:
            observation += (
                f"；表达式解析 {expression_stats.get('expression_parse_count', 0)} 次"
                f"（缓存命中 {expression_stats.get('expression_parse_cache_hits', 0)}）"
            )
        log_rows.append(("规则命中", observation))

    for issue in current.issues:
        log_rows.append(("数据异常", f"本期第 {issue.row} 行：{issue.message}"))
    for issue in previous.issues:
        log_rows.append(("数据异常", f"上期第 {issue.row} 行：{issue.message}"))

    seconds = time.monotonic() - started
    log_rows.append(("结果输出", f"输出 {output_path.name}，共 {len(rows)} 行；耗时 {seconds:.1f} 秒"))

    result = ComparisonResult(
        output_path=output_path,
        rows=rows,
        log_rows=log_rows,
        seconds=seconds,
    )
    if write_flow_logs:
        log_path = output_path.with_name(
            f"执行比较_运行日志_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        )
        write_log_workbook(result.log_rows, log_path, ("阶段", "描述"))
        say(f"完成：{output_path.name}（运行日志 {log_path.name}）")
    else:
        say(f"完成：{output_path.name}")
    return result


def run_cross_period_check(
    *,
    current_csv: str | Path | list,
    previous_csv: str | Path | list = "",
    output_path: str | Path,
    config_path: str | Path,
    source_unit: str = "元",
    target_unit: str = "亿元",
    on_step: Callable[[str], None] | None = None,
    write_flow_logs: bool = True,
) -> CheckResult:
    """跨期/数值核对：两期导数（CSV/XLSX）→ 结果工作簿（上期可空，含 {} 的规则自动跳过）。"""
    from .comparison_engine import get_unit_factor
    from .cross_period_engine import DEFAULT_TITLE, run_cross_period

    started = time.monotonic()
    say = on_step or (lambda _text: None)
    current_list = [current_csv] if not isinstance(current_csv, list) else current_csv
    previous_list = [previous_csv] if not isinstance(previous_csv, list) else previous_csv
    if not current_list or not str(current_list[0]).strip():
        raise ValueError("本期数值核对：请先选择数据文件")

    config = load_central_config(config_path)
    target_unit = config.param("默认目标单位", target_unit)
    source_unit = config.param("默认源数据单位", source_unit)
    factor = get_unit_factor(source_unit) / get_unit_factor(target_unit)
    exempt = config.exempt_indicators()

    current = _load_dataset(current_list, exempt=exempt, factor=factor)
    say(f"本期 {len(current.records)} 行")
    previous = None
    if previous_list and str(previous_list[0]).strip():
        previous = _load_dataset(previous_list, exempt=exempt, factor=factor)
        say(f"上期 {len(previous.records)} 行")
    else:
        say("未选择上期数据文件：含 {...} 引用的规则将跳过")

    say("求值本期数值核对规则")
    rows, log_rows = run_cross_period(
        current=current, previous=previous, config=config, target_unit=target_unit,
    )
    # 与执行比较同口径的装载统计（插在规则错误行之前，便于阅读）。
    cross_total = len(config.cross_rules)
    cross_disabled = sum(
        1 for row in config.cross_rules
        if str(row.get("禁用") or "").strip() in {"是", "1", "true", "True", "Y", "y"})
    summary_rows = [
        ("配置载入", f"本期数值核对规则：读取 {cross_total} 条，启用 {cross_total - cross_disabled} 条，禁用 {cross_disabled} 条"),
        ("配置载入", f"数据日期：本期 {current.record_date or '未知'}"
         + (f"；上期 {previous.record_date}" if previous is not None else "（未选上期，{...} 规则跳过）")),
        ("规则命中", f"输出 {len(rows)} 行"),
    ]
    log_rows[:0] = summary_rows
    output_path = Path(output_path)
    _write_check_workbook(rows, output_path, target_unit=target_unit)
    result = CheckResult(
        output_path=output_path, rows=rows, log_rows=log_rows,
        seconds=time.monotonic() - started,
    )
    if write_flow_logs:
        log_path = output_path.with_name(f"本期数值核对_运行日志_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        write_log_workbook(log_rows, log_path, ("机构类代码", "地区代码", "错误类型", "规则", "详情"))
    say(f"完成：{output_path.name}")
    return result


def _write_check_workbook(rows: list[dict], output_path, *, target_unit: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = list(_CHECK_TITLE)
    book = Workbook()
    sheet = book.active
    sheet.title = "对比结果"
    sheet.append(headers)
    for row in rows:
        values = []
        for name in headers:
            value = row.get(name)
            base_name = name.replace("(%)" , "").replace("(单位)", "")
            if base_name in {"左值", "右值", "差异", "差异绝对值"} and name != "差异幅度(%)":
                label = f"{base_name}({target_unit})" if base_name in {"差异"} else base_name
                values.append(row.get(label, row.get(base_name)))
            else:
                values.append(value)
        sheet.append(values)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    for index, width in enumerate((12, 10, 22, 10, 10, 10, 22, 10, 14, 12, 12, 12, 12, 12, 24, 24, 44), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    book.save(output_path)
    book.close()


def render_financial_forms(
    *,
    comparison_file: str | Path,
    template_file: str | Path | None = None,
    output_dir: str | Path,
    config_path: str | Path,
    hide_empty_rows: bool = True,
    delete_empty_sheets: bool = True,
    render_mode: str = "OPENPYXL",
    on_step: Callable[[str], None] | None = None,
) -> "object":
    """系统导数转 Excel：比较结果 + 内嵌金融表单 → 每机构一个 xlsx。

    ``template_file`` 仅保留给旧调用方及测试使用；正式 3.3 配置已经内嵌
    金融表单时，省略它即可直接运行。
    """
    from io import BytesIO

    from openpyxl import load_workbook

    from .comparison_exporter import RESULT_HEADERS
    from .renderers import RendererFactory, output_file_name
    from .models import FileSetResult

    started = time.monotonic()
    say = on_step or (lambda _text: None)
    config = load_central_config(config_path)
    target_unit = config.param("默认目标单位", "亿元")
    hide_empty_rows = config.toform_params.get("隐藏空行", "是") != "否"
    delete_empty_sheets = config.toform_params.get("空表删除", "是") != "否"
    if template_file:
        template_path = Path(template_file)
    else:
        from .form_template import has_embedded_form_template

        template_path = next(
            (candidate for candidate in reversed(config.paths)
             if has_embedded_form_template(candidate)),
            None,
        )
        if template_path is None:
            raise ValueError("指标比较拆分配置未内嵌金融表单；请重置 3.3 配置文件")

    raw = Path(comparison_file).read_bytes()
    if raw[:2] != b"PK":
        raise ValueError(f"比较结果须为 xlsx（不支持旧 .xls）：{Path(comparison_file).name}")
    book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    sheet = None
    for candidate in book.sheetnames:
        if "对比结果" in candidate:
            sheet = book[candidate]
            break
    if sheet is None:
        raise ValueError("比较结果文件缺少「对比结果」工作表")
    rows = [row for row in sheet.iter_rows(values_only=True) if any(v not in (None, "") for v in row)]
    book.close()
    header, data_rows = list(rows[0]), rows[1:]
    if [str(h) for h in header[:14]] != [h.format(unit=target_unit) if "增减额" in str(h) else str(h) for h in RESULT_HEADERS[:14]]:
        say("提示：比较结果表头与标准 21 列不完全一致，按固定列位置读取")

    orgs: dict[str, dict] = {}
    for row in data_rows:
        org_code = str(row[2] or "").replace("'", "").strip()
        region_code = str(row[4] or "").replace("'", "").strip()
        org_name = str(row[3] or "").strip()
        region_name = str(row[5] or "").strip()
        record_date = str(row[1] or "")[:10]
        key = f"{row[0]}{row[7]}{row[9]}{row[10]}"   # 业务类+指标代码+数据属性+币种
        freq_batch = f"{row[11]}{row[12]}"
        org = orgs.setdefault(org_code + region_code, {
            "org_name": org_name, "region_name": region_name,
            "record_date": record_date, "data": {}, "freq_batch": freq_batch,
        })
        # 比较结果数值已按目标单位（表单口径）；同键首条胜出（VBA 语义）。
        org["data"].setdefault(key, (row[13], row[14]))

    alerts = config.alerts

    say(f"Renderer: {render_mode}")
    say("编译金融表单模板")
    factory = RendererFactory(template_path, alerts=alerts,
                              hide_empty_rows=hide_empty_rows,
                              delete_empty_sheets=delete_empty_sheets)

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    result = FileSetResult(directory=output_dir)
    result.stats["renderer"] = render_mode

    tasks = []
    for org in orgs.values():
        name = output_file_name(org["record_date"], org["org_name"], org["region_name"])
        tasks.append({
            "org_name": org["org_name"], "region_name": org["region_name"],
            "record_date": org["record_date"], "data": org["data"],
            "output_path": str(output_dir / name), "name": name,
        })

    # 多机构并行：每个工作进程经 RendererFactory 只构建一次渲染器
    # （标准=解析一次模板元数据；高速=编译一次 OOXML 缓存）。
    import os

    workers = min(len(tasks), os.cpu_count() or 1, 12)
    done = False
    pool_started = False
    if workers >= 2:
        try:
            from concurrent.futures import ProcessPoolExecutor, as_completed

            say(f"并行生成 {len(tasks)} 个机构文件（{workers} 进程，Renderer: {render_mode}）……")
            spawn = __import__("multiprocessing").get_context("spawn")
            # 进程池上下文参数名为 mp_context（Python 3.8+ 均为该名；context 不是
            # 合法关键字参数，写成 context= 会在 3.13 直接抛 TypeError）。
            with ProcessPoolExecutor(
                max_workers=workers, mp_context=spawn,
                initializer=factory.worker_init,
                initargs=(render_mode, str(template_path), alerts,
                          hide_empty_rows, delete_empty_sheets),
            ) as pool:
                pool_started = True
                futures = {pool.submit(factory.worker_task, task): task for task in tasks}
                for future in as_completed(futures):
                    summary = future.result()
                    result.files.append(Path(summary["path"]))
                    result.stats[f"org[{summary['name']}]"] = summary["stats"]
                    for message in summary["messages"]:
                        say(message)
            done = True
        except Exception as exc:
            # 两种情况都退回顺序执行（保持既有韧性），但如实区分真实原因：
            # 把渲染错误误报成“并行不可用”会掩盖真正的问题。
            reason = f"{type(exc).__name__}: {exc}" if str(exc) else type(exc).__name__
            if pool_started:
                say(f"并行渲染中断（{reason}），改为顺序重试")
            else:
                say(f"并行渲染不可用（{reason}），改为顺序执行")
            result.files = []
            result.stats.clear()

    if not done:
        renderer = factory.get_renderer(render_mode)
        for index, task in enumerate(tasks, start=1):
            rendered = renderer.render_org(
                org_name=task["org_name"], region_name=task["region_name"],
                record_date=task["record_date"], data=task["data"],
                output_path=Path(task["output_path"]),
            )
            result.files.append(rendered.output_path)
            result.stats[f"org[{task['name']}]"] = rendered.stats
            for message in rendered.messages:
                say(message)

    result.stats["total_time"] = time.monotonic() - started
    result.seconds = result.stats["total_time"]
    return result
