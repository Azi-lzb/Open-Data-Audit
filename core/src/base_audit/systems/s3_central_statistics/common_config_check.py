"""Read-only preflight checks for the S3 common (3.0) workbook."""

from __future__ import annotations

import math
from pathlib import Path
from typing import Any

from .config import (
    ALERT_SHEET,
    COMMON_SHEETS,
    DEFAULT_RUN_PARAMS,
    ORG_REFERENCE_SHEET,
    RUN_SHEET,
    UNIT_EXCEPTION_SHEET,
    CentralConfigError,
    _as_number,
    _read_rows,
    _text,
)


_ALLOWED_UNITS = frozenset({"元", "万元", "亿元"})
_UNIT_PARAMETERS = ("默认源数据单位", "默认目标单位")


def _check(name: str, status: str, detail: str) -> dict[str, str]:
    return {"name": name, "status": status, "detail": detail}


def _missing_workbook_report(path: Path | None, detail: str) -> dict[str, Any]:
    return {
        "path": str(path.resolve()) if path is not None else "",
        "passed": False,
        "errors": [detail],
        "warnings": [],
        "checks": [_check("工作簿读取", "error", detail)],
        "sheet_stats": [],
    }


def _alert_interval_issues(rows: list[dict[str, str]]) -> tuple[list[str], list[str]]:
    """Return blocking malformed-row errors and non-blocking continuity warnings.

    Runtime alert matching uses workbook row order and half-open intervals
    ``[下限, 上限)``.  The preflight therefore compares adjacent valid rows in
    that same order and never fills or normalizes a boundary.
    """
    errors: list[str] = []
    warnings: list[str] = []
    intervals: list[tuple[int, float | None, float | None]] = []

    for row in rows:
        excel_row = row.get("__行号__", "?")
        lower_text = _text(row.get("下限"))
        upper_text = _text(row.get("上限"))
        lower = _as_number(lower_text) if lower_text else None
        upper = _as_number(upper_text) if upper_text else None

        if lower_text and lower is None:
            errors.append(f"「{ALERT_SHEET}」第{excel_row}行下限不是有效数字：{lower_text!r}。")
        if upper_text and upper is None:
            errors.append(f"「{ALERT_SHEET}」第{excel_row}行上限不是有效数字：{upper_text!r}。")
        if (lower_text and lower is None) or (upper_text and upper is None):
            continue
        if lower is None and upper is None:
            errors.append(f"「{ALERT_SHEET}」第{excel_row}行上下限同时为空，无法形成警戒区间。")
            continue
        if lower is not None and upper is not None and lower >= upper:
            errors.append(
                f"「{ALERT_SHEET}」第{excel_row}行区间无效：下限 {lower:g} 必须小于上限 {upper:g}。"
            )
            continue
        intervals.append((int(excel_row) if excel_row.isdigit() else -1, lower, upper))

    if not intervals:
        if not errors:
            errors.append(f"「{ALERT_SHEET}」没有可用的警戒区间。")
        return errors, warnings

    for previous, current in zip(intervals, intervals[1:]):
        prev_row, _prev_lower, prev_upper = previous
        curr_row, curr_lower, _curr_upper = current
        prev_location = str(prev_row) if prev_row >= 0 else "?"
        curr_location = str(curr_row) if curr_row >= 0 else "?"
        if prev_upper is None:
            warnings.append(
                f"「{ALERT_SHEET}」第{prev_location}行上限为空，后续第{curr_location}行与其重叠；"
                "程序未自动调整区间，建议人工复核。"
            )
        elif curr_lower is None:
            warnings.append(
                f"「{ALERT_SHEET}」第{curr_location}行下限为空，位于第{prev_location}行之后会与前段重叠；"
                "程序未自动调整区间，建议人工复核。"
            )
        elif math.isclose(prev_upper, curr_lower, rel_tol=0.0, abs_tol=1e-10):
            continue
        elif curr_lower > prev_upper:
            warnings.append(
                f"「{ALERT_SHEET}」第{prev_location}行上限 {prev_upper:g} 与第{curr_location}行下限 "
                f"{curr_lower:g} 不连续，存在缺口；程序未自动补值，建议人工复核。"
            )
        else:
            warnings.append(
                f"「{ALERT_SHEET}」第{prev_location}行上限 {prev_upper:g} 与第{curr_location}行下限 "
                f"{curr_lower:g} 不连续，区间有重叠；程序未自动调整，建议人工复核。"
            )
    return errors, warnings


def check_common_config(path: str | Path | None) -> dict[str, Any]:
    """Inspect the 3.0 common workbook without changing it.

    The flexible ``单位换算例外`` and ``机构地区参照`` sheets are presence-only:
    their headers, cells, and business rows are deliberately not validated.
    """
    from openpyxl import load_workbook

    target = Path(path) if path is not None else None
    if target is None or not target.is_file():
        label = str(target) if target is not None else "（未选择文件）"
        return _missing_workbook_report(target, f"大集中通用配置文件不存在：{label}。")

    report: dict[str, Any] = {
        "path": str(target.resolve()),
        "passed": False,
        "errors": [],
        "warnings": [],
        "checks": [],
        "sheet_stats": [],
    }

    try:
        workbook = load_workbook(target, read_only=True, data_only=True)
    except Exception as exc:
        detail = f"大集中通用配置文件无法读取：{target.name}（{exc}）。"
        report["errors"].append(detail)
        report["checks"].append(_check("工作簿读取", "error", detail))
        return report

    try:
        missing = [sheet for sheet in COMMON_SHEETS if sheet not in workbook.sheetnames]
        for sheet in COMMON_SHEETS:
            present = sheet in workbook.sheetnames
            scope = (
                "仅检查工作表是否存在"
                if sheet in (UNIT_EXCEPTION_SHEET, ORG_REFERENCE_SHEET)
                else "检查表头及可前置校验内容"
            )
            report["sheet_stats"].append({
                "sheet": sheet,
                "present": present,
                "scope": scope,
            })
        if missing:
            detail = f"3.0大集中通用配置缺少工作表：{'、'.join(missing)}。"
            report["errors"].append(detail)
            report["checks"].append(_check("必需工作表", "error", detail))
        else:
            report["checks"].append(_check(
                "必需工作表", "passed", f"已找到全部必需工作表：{'、'.join(COMMON_SHEETS)}。"
            ))

        if RUN_SHEET in workbook.sheetnames:
            try:
                run_rows = _read_rows(workbook, RUN_SHEET)
                report["sheet_stats"][COMMON_SHEETS.index(RUN_SHEET)]["data_rows"] = len(run_rows)
            except CentralConfigError as exc:
                detail = f"「{RUN_SHEET}」表结构不完整：{exc}。"
                report["errors"].append(detail)
                report["checks"].append(_check("运行参数表结构", "error", detail))
                run_rows = []

            configured_values: dict[str, list[tuple[str, str]]] = {
                parameter: [] for parameter in _UNIT_PARAMETERS
            }
            for row in run_rows:
                parameter = _text(row.get("参数"))
                if parameter in configured_values:
                    configured_values[parameter].append((
                        row.get("__行号__", "?"), _text(row.get("值"))
                    ))

            for parameter in _UNIT_PARAMETERS:
                entries = configured_values[parameter]
                if not entries:
                    default_value = DEFAULT_RUN_PARAMS.get(parameter, "")
                    detail = (
                        f"「{RUN_SHEET}」缺少参数“{parameter}”；当前程序默认值为“{default_value}”，"
                        "建议在配置中明确填写并人工复核。"
                    )
                    report["warnings"].append(detail)
                    report["checks"].append(_check(parameter, "warning", detail))
                    continue
                if len(entries) > 1:
                    row_numbers = "、".join(row_number for row_number, _value in entries)
                    detail = f"「{RUN_SHEET}」参数“{parameter}”重复出现在第{row_numbers}行；程序未自动选择，建议人工复核。"
                    report["errors"].append(detail)
                    report["checks"].append(_check(parameter, "error", detail))
                    continue
                row_number, value = entries[0]
                if value not in _ALLOWED_UNITS:
                    detail = (
                        f"「{RUN_SHEET}」第{row_number}行参数“{parameter}”的值为“{value or '空白'}”，"
                        "只允许“元”“万元”“亿元”；请人工修正，程序未自动修改配置。"
                    )
                    report["errors"].append(detail)
                    report["checks"].append(_check(parameter, "error", detail))
                else:
                    report["checks"].append(_check(
                        parameter, "passed", f"「{RUN_SHEET}」第{row_number}行单位“{value}”有效。"
                    ))

            soft_color_entries = [
                (row.get("__行号__", "?"), _text(row.get("值")))
                for row in run_rows
                if _text(row.get("参数")) == "复杂校验软性颜色"
            ]
            if len(soft_color_entries) > 1:
                detail = f"「{RUN_SHEET}」参数“复杂校验软性颜色”重复，请保留一行。"
                report["errors"].append(detail)
                report["checks"].append(_check("复杂校验软性颜色", "error", detail))
            elif soft_color_entries:
                row_number, raw_color = soft_color_entries[0]
                try:
                    int(raw_color or DEFAULT_RUN_PARAMS["复杂校验软性颜色"])
                except ValueError:
                    detail = (
                        f"「{RUN_SHEET}」第{row_number}行“复杂校验软性颜色”必须是整数；"
                        f"当前值为“{raw_color}”，执行时无法读取。"
                    )
                    report["errors"].append(detail)
                    report["checks"].append(_check("复杂校验软性颜色", "error", detail))
                else:
                    report["checks"].append(_check(
                        "复杂校验软性颜色", "passed", f"「{RUN_SHEET}」第{row_number}行颜色编号可解析。"
                    ))
            else:
                detail = "「运行参数」未配置“复杂校验软性颜色”，执行时使用程序默认值 46。"
                report["warnings"].append(detail)
                report["checks"].append(_check("复杂校验软性颜色", "warning", detail))

        if ALERT_SHEET in workbook.sheetnames:
            try:
                alert_rows = _read_rows(workbook, ALERT_SHEET)
                report["sheet_stats"][COMMON_SHEETS.index(ALERT_SHEET)]["data_rows"] = len(alert_rows)
            except CentralConfigError as exc:
                detail = f"「{ALERT_SHEET}」表结构不完整：{exc}。"
                report["errors"].append(detail)
                report["checks"].append(_check("环比警戒表结构", "error", detail))
            else:
                alert_errors, alert_warnings = _alert_interval_issues(alert_rows)
                report["errors"].extend(alert_errors)
                report["warnings"].extend(alert_warnings)
                status = "error" if alert_errors else ("warning" if alert_warnings else "passed")
                if alert_errors:
                    detail = f"「{ALERT_SHEET}」发现 {len(alert_errors)} 项无效区间。"
                elif alert_warnings:
                    detail = f"「{ALERT_SHEET}」发现 {len(alert_warnings)} 项区间不连续提示。"
                else:
                    detail = f"「{ALERT_SHEET}」{len(alert_rows)} 条警戒区间格式有效且相邻区间连续。"
                report["checks"].append(_check("环比警戒区间", status, detail))

        report["passed"] = not report["errors"]
        return report
    finally:
        workbook.close()


__all__ = ["check_common_config"]
