"""从大集中比较结果导出人工选定的指标说明文件。

这是比较完成后的人工复核工具：用户在比较结果的「输出说明」列标记“是”，
本模块仅整理这些行，不重算、不修改比较结果，也不依赖 Office/COM。
"""

from __future__ import annotations

from datetime import datetime
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from shutil import copy2

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_REQUIRED_HEADERS = (
    "机构类代码", "机构类名称", "地区代码", "地区名称", "指标代码", "指标名称", "数据值", "上期值",
    "输出说明", "说明内容",
)
_SELECTED_VALUES = {"是", "y", "yes", "true", "1", "√", "✓"}
_ORG_REFERENCE_SHEET = "机构地区参照"


@dataclass(frozen=True)
class ExplanationImportResult:
    imported: int
    unmatched: int
    conflicts: int
    backup_path: Path
    conflict_report_path: Path | None = None


def _header_index(sheet) -> dict[str, int]:
    headers = {
        str(cell.value).strip(): cell.column
        for cell in sheet[1]
        if cell.value not in (None, "")
    }
    missing = [header for header in _REQUIRED_HEADERS if header not in headers]
    if missing:
        raise ValueError("比较结果缺少说明导出所需列：" + "、".join(missing))
    return headers


def _selected(value: object) -> bool:
    return str(value or "").strip().lower() in _SELECTED_VALUES


def _load_org_reference(config_path: Path | None) -> dict[tuple[str, str], str]:
    if config_path is None or not Path(config_path).is_file():
        return {}
    book = load_workbook(config_path, read_only=True, data_only=True)
    try:
        if _ORG_REFERENCE_SHEET not in book.sheetnames:
            return {}
        sheet = book[_ORG_REFERENCE_SHEET]
        headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value not in (None, "")}
        required = ("机构类代码", "地区代码", "机构名称")
        if any(name not in headers for name in required):
            return {}
        mapping: dict[tuple[str, str], str] = {}
        names_by_org: dict[str, set[str]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) < max(headers[name] for name in required):
                continue
            org_code = str(row[headers["机构类代码"] - 1] or "").replace("'", "").strip()
            region_code = str(row[headers["地区代码"] - 1] or "").replace("'", "").strip()
            name = str(row[headers["机构名称"] - 1] or "").strip()
            if not org_code or not name:
                continue
            mapping[(org_code, region_code)] = name
            names_by_org.setdefault(org_code, set()).add(name)
        # 用户可在配置中显式填写地区代码“*”；只有唯一名称的机构也自动视为全地区适用。
        for org_code, names in names_by_org.items():
            if len(names) == 1:
                mapping.setdefault((org_code, "*"), next(iter(names)))
        return mapping
    finally:
        book.close()


def export_selected_explanations(comparison_file: Path, output_path: Path, *, config_path: Path | None = None) -> int:
    """导出「输出说明」标记为“是”的行，返回导出条数。

    输出采用旧说明文件兼容的八个业务列。
    """
    comparison_file, output_path = Path(comparison_file), Path(output_path)
    source = load_workbook(comparison_file, read_only=True, data_only=True)
    try:
        sheet = next(
            (item for item in source.worksheets if "指标代码" in {
                str(cell.value).strip() for cell in item[1] if cell.value not in (None, "")
            }),
            None,
        )
        if sheet is None:
            raise ValueError("比较结果中未找到包含“指标代码”的结果工作表")
        index = _header_index(sheet)
        unit_header = str(sheet.cell(1, index["数据值"]).value or "本期金额")
        change_header = next((name for name in index if name.startswith("增减额")), "增减额")
        org_reference = _load_org_reference(config_path)
        # 先扫全表统计各机构出现的地区代码：一家机构对应多个地区时（如同一
        # 农商行既报省级又报市级），说明文件的机构代码拼上地区代码才能区分，
        # 机构反馈回填的长码也能被四项标识唯一匹配。单地区机构保持原码。
        regions_by_org: dict[str, set[str]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            org = str(row[index["机构类代码"] - 1] or "").replace("'", "").strip() if len(row) >= index["机构类代码"] else ""
            region = str(row[index["地区代码"] - 1] or "").replace("'", "").strip() if len(row) >= index["地区代码"] else ""
            if org and region:
                regions_by_org.setdefault(org, set()).add(region)
        selected_rows = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not _selected(row[index["输出说明"] - 1] if len(row) >= index["输出说明"] else None):
                continue
            org_code = str(row[index["机构类代码"] - 1] or "").replace("'", "").strip()
            region_code = str(row[index["地区代码"] - 1] or "").replace("'", "").strip()
            org_name = str(row[index["机构类名称"] - 1] or "").strip()
            region_name = str(row[index["地区名称"] - 1] or "").strip()
            base_name = (org_reference.get((org_code, region_code))
                         or org_reference.get((org_code, "*"))
                         or org_name)
            # 地区始终取比较文件的原始值：映射只负责纠正机构名称，不能抹掉区县维度。
            display_name = f"{base_name}（{region_name}）" if region_name else base_name
            # 长码是两条独立的规则，判断口径不同：
            # 1) 村镇银行：机构类代码全国重号（7020 系），对齐 VBA 惯例——无论该批
            #    比较结果里出现几个地区，一律拼地区代码长码；
            # 2) 非村镇银行：机构类代码本身唯一，仅当该机构在比较结果中覆盖多个
            #    地区时才拼地区代码区分；单地区保持原码，机构按原习惯填短码即可。
            is_village_bank = org_code.startswith("7020") or "村镇银行" in org_name
            multi_region = len(regions_by_org.get(org_code, ())) > 1
            if region_code and (is_village_bank or multi_region):
                display_code = f"{org_code}{region_code}"
            else:
                display_code = org_code
            selected_rows.append((
                display_code, display_name,
                row[index["指标代码"] - 1], row[index["指标名称"] - 1],
                row[index["数据值"] - 1], row[index["上期值"] - 1],
                row[index[change_header] - 1] if len(row) >= index[change_header] else None,
                row[index["说明内容"] - 1],
            ))
    finally:
        source.close()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    explanation = book.active
    explanation.title = "说明文件"
    explanation.merge_cells("A1:H1")
    explanation["A1"] = "xx说明"
    explanation["A1"].font = Font(bold=True, size=14)
    explanation["A1"].alignment = Alignment(horizontal="center", vertical="center")
    headers = (
        "机构代码", "机构名称", "指标代码", "指标名称", unit_header,
        "上期值", change_header, "变动原因",
    )
    explanation.append(headers)
    for row in selected_rows:
        explanation.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in explanation[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    explanation.freeze_panes = "A3"
    explanation.auto_filter.ref = f"A2:H{max(3, explanation.max_row)}"
    for column, width in enumerate((14, 26, 14, 36, 14, 14, 14, 48), start=1):
        explanation.column_dimensions[get_column_letter(column)].width = width
    # E/F/G（数据值/上期值/增减额）显示两位小数；只设显示格式，单元格数值不动
    # （openpyxl 不参与计算，四舍五入交给查看者与导入匹配的舍入规则）。
    for column in (5, 6, 7):
        for cell in explanation[get_column_letter(column)][2:]:
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = "0.00"

    info = book.create_sheet("生成信息")
    info.append(("项目", "内容"))
    info.append(("来源比较文件", str(comparison_file)))
    info.append(("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
    info.append(("选定指标数", len(selected_rows)))
    info.append(("选取规则", "比较结果“输出说明”列填写“是”"))
    for cell in info[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 90
    book.save(output_path)
    book.close()
    return len(selected_rows)


def _find_explanation_sheet(book):
    if "说明文件" in book.sheetnames:
        return book["说明文件"]
    for sheet in book.worksheets:
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10), values_only=True):
            if "机构代码" in {str(value).strip() for value in row if value not in (None, "")}:
                return sheet
    raise ValueError("未找到说明文件工作表或“机构代码”表头")


def _explanation_header_index(sheet) -> tuple[int, dict[str, int]]:
    for row_number in range(1, min(sheet.max_row, 10) + 1):
        headers = {
            str(cell.value).strip(): cell.column
            for cell in sheet[row_number]
            if cell.value not in (None, "")
        }
        if all(name in headers for name in ("机构代码", "机构名称", "指标代码", "指标名称")):
            value_header = next((name for name in headers if name == "数据值" or name.startswith("本期金额")), None)
            change_header = next((name for name in headers if name.startswith("增减额")), None)
            if value_header and change_header and "变动原因" in headers:
                return row_number, {
                    "机构代码": headers["机构代码"],
                    "机构名称": headers["机构名称"],
                    "指标代码": headers["指标代码"],
                    "指标名称": headers["指标名称"],
                    "数据值": headers[value_header],
                    "增减额": headers[change_header],
                    "变动原因": headers["变动原因"],
                }
    raise ValueError("说明文件缺少机构代码、机构名称、指标代码、指标名称、数据值、增减额或变动原因列")


def _code_value(value: object) -> str:
    """机构/指标代码的匹配值：文本化、去 CSV 撇号、数字去尾随 .0、统一大写。

    代码是标识不是数量，禁止数值化——Decimal 归一会把 7020 变成 7.02E+3，
    与机构按文本填写的代码永不相等（真实批次实测教训）。大小写归一对齐
    Excel/VBA 时代的宽松比较（龙门 6K0u vs 6k0u 实测案例）。
    """
    if value is None:
        return ""
    text = str(value).strip().replace("'", "")
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text.upper()


def _amount_value(value: object) -> str:
    """数据值/增减额的匹配值：按 4 位小数四舍五入（ROUND_HALF_UP）。

    机构誊抄说明时常保留 4 位小数（0.4185 vs 0.418462 实测案例），精确
    相等会误伤合法反馈；两侧同规则舍入后再比较。非数值文本原样返回，
    不与数字互相误匹配。
    """
    if value is None:
        return ""
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = Decimal(str(value))
    else:
        text = str(value).strip()
        try:
            number = Decimal(text)
        except InvalidOperation:
            return text
    try:
        quantized = number.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
    except InvalidOperation:
        return str(number)
    return f"#{quantized:f}"


def _identity(row: tuple[object, ...], columns: dict[str, int]) -> tuple[str, str, str, str]:
    """四项业务标识：机构代码+指标代码+数据值+增减额。

    机构/指标名称不参与匹配：机构常改写名称（加层级前缀、括号注记）或
    复制他行忘改机构名（博罗文件通篇写惠东的实测案例），代码加两值足以
    定位唯一行；同标识出现多行时按冲突处理，不悄悄写入。
    """
    return (
        _code_value(row[columns["机构代码"] - 1]) if len(row) >= columns["机构代码"] else "",
        _code_value(row[columns["指标代码"] - 1]) if len(row) >= columns["指标代码"] else "",
        _amount_value(row[columns["数据值"] - 1]) if len(row) >= columns["数据值"] else "",
        _amount_value(row[columns["增减额"] - 1]) if len(row) >= columns["增减额"] else "",
    )


def import_explanation_feedback(target_path: Path, feedback_paths: list[Path]) -> ExplanationImportResult:
    """按四项业务标识（机构代码+指标代码+数据值+增减额）回写机构反馈。

    代码统一大写、金额按 4 位小数舍入后比较，名称不参与（机构常改写或
    填错名称，实测详见 _identity）。同一标识出现不同反馈时不写入，避免
    批量导入悄悄覆盖人工内容。写入前在同目录创建带时间戳的完整备份。
    """
    target_path = Path(target_path)
    if not feedback_paths:
        raise ValueError("请至少选择一个机构反馈说明文件")
    target = load_workbook(target_path)
    try:
        target_sheet = _find_explanation_sheet(target)
        target_header_row, target_columns = _explanation_header_index(target_sheet)
        targets: dict[tuple[str, str, str, str], list[tuple[int, tuple[object, ...]]]] = {}
        for row_number in range(target_header_row + 1, target_sheet.max_row + 1):
            row = tuple(target_sheet.cell(row_number, column).value for column in range(1, target_sheet.max_column + 1))
            key = _identity(row, target_columns)
            if any(key):
                targets.setdefault(key, []).append((row_number, row))

        # key → {反馈原因 → 来源文件名集合}。保留来源才能在冲突时给用户可核查的明细。
        feedback: dict[tuple[str, str, str, str], dict[str, set[str]]] = {}
        for feedback_path in feedback_paths:
            source = load_workbook(feedback_path, read_only=True, data_only=True)
            try:
                sheet = _find_explanation_sheet(source)
                header_row, columns = _explanation_header_index(sheet)
                for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                    key = _identity(row, columns)
                    reason = str(row[columns["变动原因"] - 1] or "").strip() if len(row) >= columns["变动原因"] else ""
                    if any(key) and reason:
                        feedback.setdefault(key, {}).setdefault(reason, set()).add(feedback_path.name)
            finally:
                source.close()

        imported = unmatched = conflicts = 0
        conflict_rows: list[tuple[object, ...]] = []

        def target_field(row: tuple[object, ...], name: str) -> object:
            position = target_columns[name] - 1
            return row[position] if len(row) > position else ""

        for key, reason_sources in feedback.items():
            target_rows = targets.get(key)
            if not target_rows:
                unmatched += 1
                continue
            if len(reason_sources) != 1 or len(target_rows) != 1:
                conflicts += 1
                reasons_text = "\n".join(
                    f"{reason}（来源：{'、'.join(sorted(sources))}）"
                    for reason, sources in sorted(reason_sources.items())
                )
                causes = []
                if len(target_rows) != 1:
                    causes.append(f"目标说明文件存在 {len(target_rows)} 行相同标识")
                if len(reason_sources) != 1:
                    causes.append("多个反馈文件给出不同变动原因")
                sample_number, sample_row = target_rows[0]
                conflict_rows.append((
                    "冲突",
                    key[0], str(target_field(sample_row, "机构名称")),
                    key[1], str(target_field(sample_row, "指标名称")),
                    key[2].lstrip("#"), key[3].lstrip("#"),
                    "、".join(str(number) for number, _row in target_rows),
                    reasons_text, "；".join(causes),
                ))
                continue
            target_sheet.cell(target_rows[0][0], target_columns["变动原因"]).value = next(iter(reason_sources))
            imported += 1

        backup_path = target_path.with_name(
            f"{target_path.stem}_导入前备份_{datetime.now():%Y%m%d_%H%M%S}{target_path.suffix}"
        )
        copy2(target_path, backup_path)
        target.save(target_path)
        conflict_report_path = None
        if conflict_rows:
            conflict_report_path = target_path.with_name(
                f"{target_path.stem}_导入冲突清单_{datetime.now():%Y%m%d_%H%M%S}{target_path.suffix}"
            )
            report = Workbook()
            sheet = report.active
            sheet.title = "冲突明细"
            sheet.append((
                "状态", "机构代码", "机构名称", "指标代码", "指标名称", "数据值", "增减额",
                "目标说明文件行号", "反馈文件及原因", "冲突原因",
            ))
            for row in conflict_rows:
                sheet.append(row)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="C00000")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:J{sheet.max_row}"
            for column, width in enumerate((10, 14, 26, 14, 32, 14, 14, 20, 72, 38), start=1):
                sheet.column_dimensions[get_column_letter(column)].width = width
            report.save(conflict_report_path)
            report.close()
        return ExplanationImportResult(imported, unmatched, conflicts, backup_path, conflict_report_path)
    finally:
        target.close()
