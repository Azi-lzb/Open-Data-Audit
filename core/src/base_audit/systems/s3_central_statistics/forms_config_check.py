"""Read-only preflight checks for the S3 3.3 forms configuration workbook."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import (
    FORM_SHEET,
    SHEET_HEADERS,
    TOFORM_SHEET,
    CentralConfigError,
    _locate_header,
)
from .form_template import has_embedded_form_template


REQUIRED_TOFORM_SETTINGS = {
    "隐藏空行": "boolean",
    "空表删除": "boolean",
    "分析文件数据单位": "unit",
}
ALLOWED_UNITS = {"元", "万元", "亿元"}
ALLOWED_SWITCHES = {"是", "否"}
REQUIRED_SHEETS = (TOFORM_SHEET, FORM_SHEET)
CHECK_DEFINITIONS = (
    ("file_read", "配置文件可读取"),
    ("required_sheets", "必需工作表及表头齐全"),
    ("settings", "转表设置完整且取值有效"),
    ("embedded_template", "内嵌金融表单模板可识别"),
    ("form_references", "报表清单引用的表单工作表存在"),
)


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _add_issue(
    report: dict[str, Any],
    collection: str,
    message: str,
    *,
    sheet: str | None = None,
    row: int | None = None,
    field: str | None = None,
) -> None:
    issue: dict[str, Any] = {"message": message}
    if sheet:
        issue["sheet"] = sheet
    if row is not None:
        issue["row"] = row
    if field:
        issue["field"] = field
    report[collection].append(issue)


def _set_check(report: dict[str, Any], check_id: str, status: str, detail: str) -> None:
    for check in report["checks"]:
        if check["id"] == check_id:
            check["status"] = status
            check["detail"] = detail
            return
    raise KeyError(f"unknown S3 forms config check: {check_id}")


def _finish_check(
    report: dict[str, Any], check_id: str, error_start: int, success_detail: str
) -> None:
    if len(report["errors"]) > error_start:
        count = len(report["errors"]) - error_start
        _set_check(report, check_id, "failed", f"发现 {count} 个问题，详见问题明细。")
    else:
        _set_check(report, check_id, "passed", success_detail)


def _headers(sheet, sheet_name: str, report: dict[str, Any]) -> dict[str, int] | None:
    """Use the same header-row locator as the S3 runtime configuration reader."""
    try:
        rows = list(sheet.iter_rows(values_only=True))
        header_index, positions = _locate_header(rows, SHEET_HEADERS[sheet_name], sheet_name)
    except CentralConfigError as exc:
        _add_issue(report, "errors", str(exc), sheet=sheet_name)
        return None
    return {name: column for name, column in positions.items() if name}


def _sheet_stats(book) -> list[dict[str, Any]]:
    stats = []
    for sheet in book.worksheets:
        stats.append({
            "name": sheet.title,
            "rows": max(int(sheet.max_row or 0) - 1, 0),
            "columns": int(sheet.max_column or 0),
        })
    return stats


def check_forms_config(path: Path | str) -> dict[str, Any]:
    """Check the 3.3 workbook's runnable settings and embedded form references.

    The workbook is only opened in read-only mode. Embedded form sheet contents
    are deliberately left to the form renderer; this preflight checks only the
    report-list contract and whether each runtime-recognized form code resolves
    to a worksheet.
    """
    target = Path(path)
    report: dict[str, Any] = {
        "path": str(target),
        "passed": False,
        "errors": [],
        "warnings": [],
        "checks": [
            {"id": check_id, "name": name, "status": "not_run", "detail": "未执行。"}
            for check_id, name in CHECK_DEFINITIONS
        ],
        "sheet_stats": [],
    }

    if not target.is_file():
        _add_issue(report, "errors", f"配置文件不存在：{target}")
        _set_check(report, "file_read", "failed", "配置文件不存在，无法读取。")
        _set_check(report, "required_sheets", "not_run", "配置文件未能打开。")
        _set_check(report, "settings", "not_run", "配置文件未能打开。")
        _set_check(report, "embedded_template", "not_run", "配置文件未能打开。")
        _set_check(report, "form_references", "not_run", "配置文件未能打开。")
        return report

    try:
        book = load_workbook(target, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - report a user-facing workbook read error
        _add_issue(report, "errors", f"配置文件无法读取：{exc}")
        _set_check(report, "file_read", "failed", "配置文件无法打开，详见问题明细。")
        for check_id in ("required_sheets", "settings", "embedded_template", "form_references"):
            _set_check(report, check_id, "not_run", "配置文件未能打开。")
        return report

    try:
        _set_check(report, "file_read", "passed", "配置工作簿已成功打开，并以只读方式检查。")
        report["sheet_stats"] = _sheet_stats(book)
        names = set(book.sheetnames)

        required_start = len(report["errors"])
        for sheet_name in REQUIRED_SHEETS:
            if sheet_name not in names:
                _add_issue(report, "errors", f"缺少工作表：{sheet_name}", sheet=sheet_name)

        toform_headers = None
        if TOFORM_SHEET in names:
            toform_headers = _headers(book[TOFORM_SHEET], TOFORM_SHEET, report)

        form_headers = None
        if FORM_SHEET in names:
            # The renderer reads this sheet from row 1 and uses columns A:D by
            # position, so check that concrete runtime contract explicitly.
            first_row = next(book[FORM_SHEET].iter_rows(min_row=1, max_row=1, values_only=True), ())
            expected = SHEET_HEADERS[FORM_SHEET]
            actual = tuple(_text(value) for value in first_row[: len(expected)])
            if actual != expected:
                missing = [name for name in expected if name not in actual]
                if missing:
                    detail = f"缺少必需列：{'、'.join(missing)}"
                else:
                    detail = f"表头需按运行时顺序排列：{'、'.join(expected)}"
                _add_issue(report, "errors", f"工作表“{FORM_SHEET}”{detail}", sheet=FORM_SHEET, row=1)
            else:
                form_headers = {name: index for index, name in enumerate(expected)}

        _finish_check(
            report,
            "required_sheets",
            required_start,
            f"已确认必需工作表及其运行时表头：{'、'.join(REQUIRED_SHEETS)}",
        )

        settings_start = len(report["errors"])
        if toform_headers is not None:
            settings_sheet = book[TOFORM_SHEET]
            header_rows = list(settings_sheet.iter_rows(values_only=True))
            header_index, _ = _locate_header(header_rows, SHEET_HEADERS[TOFORM_SHEET], TOFORM_SHEET)
            columns = toform_headers
            values_by_key: dict[str, list[tuple[int, Any]]] = {}
            for row_number, row in enumerate(
                settings_sheet.iter_rows(min_row=header_index + 2, values_only=True),
                start=header_index + 2,
            ):
                key_index = columns.get("参数")
                value_index = columns.get("值")
                key = _text(row[key_index]) if key_index is not None and key_index < len(row) else ""
                value = row[value_index] if value_index is not None and value_index < len(row) else None
                if key:
                    values_by_key.setdefault(key, []).append((row_number, value))

            for key in REQUIRED_TOFORM_SETTINGS:
                occurrences = values_by_key.get(key, [])
                if not occurrences:
                    _add_issue(report, "errors", f"缺少转表设置项：{key}", sheet=TOFORM_SHEET, field=key)
                elif len(occurrences) > 1:
                    for row_number, _ in occurrences[1:]:
                        _add_issue(
                            report,
                            "errors",
                            f"转表设置项重复：{key}",
                            sheet=TOFORM_SHEET,
                            row=row_number,
                            field=key,
                        )

            for key, occurrences in values_by_key.items():
                if key in REQUIRED_TOFORM_SETTINGS or len(occurrences) < 2:
                    continue
                for row_number, _ in occurrences[1:]:
                    _add_issue(
                        report,
                        "errors",
                        f"转表设置项重复：{key}",
                        sheet=TOFORM_SHEET,
                        row=row_number,
                        field=key,
                    )

            for key, kind in REQUIRED_TOFORM_SETTINGS.items():
                occurrences = values_by_key.get(key, [])
                if len(occurrences) != 1:
                    continue
                row_number, raw_value = occurrences[0]
                value = _text(raw_value)
                if kind == "unit" and value not in ALLOWED_UNITS:
                    _add_issue(
                        report,
                        "errors",
                        f"“{key}”必须为元、万元或亿元，当前值为：{value or '空白'}",
                        sheet=TOFORM_SHEET,
                        row=row_number,
                        field=key,
                    )
                elif kind == "boolean" and value not in ALLOWED_SWITCHES:
                    _add_issue(
                        report,
                        "errors",
                        f"“{key}”必须为“是”或“否”，当前值为：{value or '空白'}",
                        sheet=TOFORM_SHEET,
                        row=row_number,
                        field=key,
                    )

            _finish_check(
                report,
                "settings",
                settings_start,
                f"已检查 {len(REQUIRED_TOFORM_SETTINGS)} 项必需设置、重复键及允许值",
            )
        else:
            _set_check(report, "settings", "not_run", "转表设置工作表或必需表头缺失。")

        embedded_start = len(report["errors"])
        try:
            has_template = has_embedded_form_template(target)
        except Exception as exc:  # noqa: BLE001 - keep a readable preflight report
            has_template = False
            _add_issue(
                report,
                "errors",
                f"内嵌金融表单模板无法检查：{exc}",
                sheet=FORM_SHEET,
            )
        if not has_template and not any(
            issue.get("sheet") == FORM_SHEET and "内嵌金融表单模板无法检查" in issue["message"]
            for issue in report["errors"][embedded_start:]
        ):
            _add_issue(
                report,
                "errors",
                f"配置工作簿未包含可识别的内嵌金融表单模板；请检查“{FORM_SHEET}”及其对应表单工作表。",
                sheet=FORM_SHEET,
            )
        _finish_check(
            report,
            "embedded_template",
            embedded_start,
            "已由正式表单模板识别逻辑确认工作簿内嵌模板。",
        )

        references_start = len(report["errors"])
        if FORM_SHEET in names and form_headers is not None:
            code_column = form_headers["报表代码"]
            for row_number, row in enumerate(
                book[FORM_SHEET].iter_rows(min_row=2, values_only=True), start=2
            ):
                code = _text(row[code_column]) if code_column < len(row) else ""
                # TemplateMeta and has_embedded_form_template recognize form
                # codes containing A, matching the established runtime rule.
                if not code or "A" not in code:
                    continue
                if code not in names:
                    _add_issue(
                        report,
                        "errors",
                        f"报表代码“{code}”未对应到内嵌表单工作表。",
                        sheet=FORM_SHEET,
                        row=row_number,
                        field="报表代码",
                    )
            _finish_check(
                report,
                "form_references",
                references_start,
                "已确认报表清单中运行时识别的表单代码均有对应工作表。",
            )
        else:
            _set_check(report, "form_references", "not_run", "报表清单工作表或运行时表头缺失。")
    finally:
        book.close()

    report["passed"] = not report["errors"]
    return report
