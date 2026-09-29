"""组合工作表的用户设置及旧 ``config.xlsx`` 一次性迁移。

组合方案不是 DAG 图的一部分：它只是“组合工作表（按配置）”的业务参数。
流程拓扑、节点顺序和命名区域均由代码固定；本模块只负责用户可维护的分组方案。
"""

from __future__ import annotations

import json
import re
from copy import deepcopy
from pathlib import Path
from typing import Any


DEFAULT_REGEX_PLAN = {
    "id": "default",
    "name": "按文件名正则提取组合",
    "mode": "regex",
    "pattern": r"(?P<组合>.+?)_金融基础数据-",
    "output_marker": "正则组合",
    "groups": [],
}

DEFAULT_KEYWORD_PLAN = {
    "id": "keyword",
    "name": "按关键字分组",
    "mode": "name",
    "pattern": "",
    "output_marker": "按关键字组合",
    "groups": [{"name": "请先配置", "keywords": ["请替换"]}],
}


def default_combine_plans() -> list[dict[str, object]]:
    """返回可安全修改的默认方案副本。"""
    return deepcopy([DEFAULT_REGEX_PLAN, DEFAULT_KEYWORD_PLAN])


def _text(value: object) -> str:
    return str(value or "").strip()


def validate_combine_plans(plans: object, active_plan_id: object) -> list[str]:
    """验证设置中心提交的组合方案，返回面向用户的错误列表。"""
    if not isinstance(plans, list) or not plans:
        return ["至少保留一个组合方案"]
    ids: set[str] = set()
    names: set[str] = set()
    errors: list[str] = []
    for index, raw in enumerate(plans, start=1):
        if not isinstance(raw, dict):
            errors.append(f"第 {index} 个组合方案格式无效")
            continue
        plan_id, name, mode = _text(raw.get("id")), _text(raw.get("name")), _text(raw.get("mode"))
        if not plan_id:
            errors.append(f"第 {index} 个组合方案缺少编号")
        elif plan_id in ids:
            errors.append(f"组合方案编号重复：{plan_id}")
        ids.add(plan_id)
        if not name:
            errors.append(f"第 {index} 个组合方案名称不能为空")
        elif name in names:
            errors.append(f"组合方案名称重复：{name}")
        names.add(name)
        if mode not in {"regex", "name"}:
            errors.append(f"组合方案“{name or index}”的分组方式无效")
            continue
        if mode == "regex":
            pattern = _text(raw.get("pattern"))
            if not pattern:
                errors.append(f"组合方案“{name or index}”未填写文件名正则")
                continue
            try:
                compiled = re.compile(pattern)
            except re.error as exc:
                errors.append(f"组合方案“{name or index}”正则无效：{exc}")
                continue
            if "组合" not in compiled.groupindex and "group" not in compiled.groupindex and not compiled.groups:
                errors.append(f"组合方案“{name or index}”的正则必须包含命名分组“组合”或至少一个括号分组")
        else:
            groups = raw.get("groups")
            if not isinstance(groups, list) or not groups:
                errors.append(f"组合方案“{name or index}”至少需要一条关键字分组")
                continue
            group_names: set[str] = set()
            for group_index, group in enumerate(groups, start=1):
                if not isinstance(group, dict):
                    errors.append(f"组合方案“{name or index}”第 {group_index} 条分组格式无效")
                    continue
                group_name = _text(group.get("name"))
                keywords = group.get("keywords")
                if not group_name:
                    errors.append(f"组合方案“{name or index}”第 {group_index} 条未填写组合名称")
                elif group_name in group_names:
                    errors.append(f"组合方案“{name or index}”存在重复组合名称：{group_name}")
                group_names.add(group_name)
                if not isinstance(keywords, list) or not any(_text(item) for item in keywords):
                    errors.append(f"组合方案“{name or index}”第 {group_index} 条未填写匹配关键字")
    if _text(active_plan_id) not in ids:
        errors.append("当前使用方案必须是已保存的组合方案")
    return errors


def normalize_combine_settings(
    plans: object, active_plan_id: object,
) -> tuple[list[dict[str, object]], str]:
    """规范化可 JSON 序列化的方案，缺失时回退到两个内置默认方案。"""
    if not isinstance(plans, list):
        plans = default_combine_plans()
    normalized: list[dict[str, object]] = []
    for raw in plans:
        if not isinstance(raw, dict):
            continue
        mode = _text(raw.get("mode")) or "regex"
        groups = raw.get("groups") if isinstance(raw.get("groups"), list) else []
        normalized.append({
            "id": _text(raw.get("id")),
            "name": _text(raw.get("name")),
            "mode": mode,
            "pattern": _text(raw.get("pattern")),
            "output_marker": _text(raw.get("output_marker")) or "组合",
            "groups": [
                {
                    "name": _text(item.get("name")),
                    "keywords": [_text(word) for word in item.get("keywords", []) if _text(word)],
                }
                for item in groups if isinstance(item, dict)
            ],
        })
    active = _text(active_plan_id)
    if validate_combine_plans(normalized, active):
        normalized = default_combine_plans()
        active = _text(normalized[0]["id"])
    return normalized, active


def get_combine_plan(
    plans: object, active_plan_id: object, plan_id: str | None = None,
) -> dict[str, object]:
    """返回指定或当前组合方案；错误在执行前明确暴露。"""
    normalized, active = normalize_combine_settings(plans, active_plan_id)
    selected = _text(plan_id) or active
    for plan in normalized:
        if _text(plan.get("id")) == selected:
            return deepcopy(plan)
    raise ValueError(f"组合工作表方案不存在：{selected}")


def load_legacy_combine_settings(project_root: Path) -> tuple[list[dict[str, object]], str] | None:
    """从退役配置读取一次方案；读取失败不会阻止应用启动。"""
    root = Path(project_root).resolve()
    # 开发布局的用户设置在 core/data，而运行配置在仓库根 config；
    # 打包布局则二者同级。两种布局都只读一次，不创建旧文件。
    roots = (root, root.parent) if root.parent != root else (root,)
    for candidate_root in roots:
        from_workbook = _plans_from_legacy_workbook(candidate_root / "config" / "config.xlsx")
        if from_workbook is not None:
            return from_workbook
    for candidate_root in roots:
        from_json = _plans_from_legacy_json(candidate_root / "data" / "流程配置.json")
        if from_json is not None:
            return from_json
    return None


def _plans_from_legacy_json(path: Path) -> tuple[list[dict[str, object]], str] | None:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return None
    if not isinstance(payload, dict):
        return None
    plans = payload.get("combineSheetsPlans")
    active = payload.get("activeCombineSheetsPlanId")
    normalized, normalized_active = normalize_combine_settings(plans, active)
    if not isinstance(plans, list):
        return None
    return normalized, normalized_active


def _plans_from_legacy_workbook(path: Path) -> tuple[list[dict[str, object]], str] | None:
    if not path.is_file():
        return None
    try:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        if "组合方案" not in book.sheetnames:
            return None
        plan_rows = _sheet_rows(book["组合方案"], {"方案名称", "分组方式", "方案状态"})
        group_rows = _sheet_rows(book["组合分组"], {"方案名称", "组合名称", "匹配关键字"}) if "组合分组" in book.sheetnames else []
    finally:
        book.close()
    plans: list[dict[str, object]] = []
    active = ""
    for row in plan_rows:
        name = _text(row.get("方案名称"))
        state = _text(row.get("方案状态"))
        mode_text = _text(row.get("分组方式"))
        # 固定核查表预置归代码所有，不作为可编辑方案迁移。
        if not name or state == "固定只读":
            continue
        mode = "regex" if mode_text in {"文件名正则", "正则", "regex"} else "name"
        plan = {
            "id": "legacy-" + str(len(plans) + 1),
            "name": name,
            "mode": mode,
            "pattern": _text(row.get("文件名正则")),
            "output_marker": _text(row.get("输出标识")) or "组合",
            "groups": [],
        }
        if mode == "name":
            plan["groups"] = [
                {
                    "name": _text(group.get("组合名称")),
                    "keywords": [
                        item.strip() for item in _text(group.get("匹配关键字"))
                        .replace("、", ",").replace("，", ",").split(",") if item.strip()
                    ],
                }
                for group in group_rows if _text(group.get("方案名称")) == name
                and _text(group.get("启用")) not in {"否", "不", "0", "false", "False"}
            ]
        plans.append(plan)
        if state == "当前使用":
            active = str(plan["id"])
    if not plans:
        return None
    active = active or str(plans[0]["id"])
    errors = validate_combine_plans(plans, active)
    if errors:
        return None
    return plans, active


def _sheet_rows(sheet, required: set[str]) -> list[dict[str, str]]:
    header: dict[str, int] | None = None
    start = 0
    for row_index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        possible = {_text(value): index for index, value in enumerate(values) if _text(value)}
        if required.issubset(possible):
            header, start = possible, row_index
            break
    if header is None:
        return []
    rows: list[dict[str, str]] = []
    for values in sheet.iter_rows(min_row=start + 1, values_only=True):
        row = {name: _text(values[index]) if index < len(values) else "" for name, index in header.items()}
        if any(row.values()):
            rows.append(row)
    return rows
