"""模板的只读静态检查（检查模板：公式明确错误 + 命名区域盘点）。

这不是公式计算器：它只识别公式文本/名称定义中的明确错误
（#REF!、#NAME?、#N/A、#VALUE! 等）。易失函数（INDIRECT/OFFSET 等）与
查询引用（INDEX/MATCH/VLOOKUP 等）不属于错误，不逐条列出——公式是否
可算一律以 Office 计算引擎真实重算为准。

命名区域盘点只做提示、不做阻断：按 V3 模板协议统计三族业务命名区域
（校验区域 / 表结构区域 / 条件格式区域，含旧 VBA 变体），逐条列出
名称与 range 供人工核对；Excel 内部名称（_FilterDatabase、Print_Area、
Print_Titles 等）单独忽略清单；其余自定义名称列入“其他”，供确认是否
遗留。缺失业务命名区域不阻断任何流程——正式的必需性校验仍由
模板体检（表结构比对）在执行流程中完成。
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook


_ERROR_TOKEN = re.compile(r"#(?:REF!|VALUE!|NAME\?|N/A|DIV/0!|NUM!|NULL!|CALC!|SPILL!)", re.I)
_STRING_LITERAL = re.compile(r'"[^"]*"')

# 业务命名区域族：完整名或以族名开头（如 校验区域#1、条件格式区域_附加）。
_RANGE_FAMILIES = ("校验区域", "表结构区域", "条件格式区域")
# Excel 自动生成的内部名称（含 _xlnm 前缀变体）——不是业务命名区域。
_INTERNAL_NAMES = ("_filterdatabase", "print_area", "print_titles")


def _error_token(expression: str) -> str | None:
    """公式/名称表达式中第一个错误值 token；字符串字面量不算（"#N/A 文本"是常量）。"""
    match = _ERROR_TOKEN.search(_STRING_LITERAL.sub('""', expression))
    return match.group(0).upper() if match else None


def _classify_name(name: str) -> str:
    """命名区域分类：三族业务区域 / Excel 内部名称 / 其他。"""
    lowered = name.casefold()
    if lowered.startswith("_xlnm") or any(
            lowered == internal or lowered.startswith(internal)
            for internal in _INTERNAL_NAMES):
        return "internal"
    for family in _RANGE_FAMILIES:
        if name == family or name.startswith(family):
            return family
    if "校验区域" in name:
        # 旧 VBA 协议变体：本地校验区域1 / 联表校验区域#1 等。
        return "校验区域"
    return "other"


def _defined_names(book):
    """全部命名区域：(名称, range文本, 作用域)。作用域=工作簿或工作表名。"""
    values = [(item, str(getattr(item, "attr_text", "") or ""), "工作簿")
              for item in book.defined_names.values()]
    for sheet in book.worksheets:
        values.extend((item, str(getattr(item, "attr_text", "") or ""), sheet.title)
                      for item in sheet.defined_names.values())
    return values


def inspect_template_formulas(template_path: Path) -> dict:
    """Return serialisable formula-error and named-range findings (read-only)."""
    path = Path(template_path)
    if not path.is_file():
        raise FileNotFoundError("模板文件不存在：{}".format(path))
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    errors: list[dict] = []
    family_items: dict[str, list[dict]] = {family: [] for family in _RANGE_FAMILIES}
    internal_items: list[dict] = []
    other_items: list[dict] = []
    formula_count = 0
    try:
        for defined, range_text, scope in _defined_names(book):
            name = str(getattr(defined, "name", "") or "")
            if not name:
                continue
            entry = {"name": name, "range": range_text, "scope": scope}
            family = _classify_name(name)
            if family == "internal":
                # Excel 自动生成的内部名称不是业务命名区域，且常残留过期
                # 引用（如自动筛选的 #REF!）——只进忽略清单，不报错误。
                internal_items.append(entry)
                continue
            if family == "other":
                other_items.append(entry)
            else:
                family_items[family].append(entry)
            token = _error_token(range_text)
            if token:
                errors.append({
                    "kind": "命名区域", "location": name, "token": token,
                    "message": "名称定义引用了错误值：{}".format(token),
                })
        for sheet in book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    formula = cell.value
                    if not isinstance(formula, str) or not formula.startswith("="):
                        continue
                    formula_count += 1
                    location = "{}!{}".format(sheet.title, cell.coordinate)
                    token = _error_token(formula)
                    if token:
                        errors.append({
                            "kind": "公式", "location": location, "token": token,
                            "message": "公式文本含错误引用：{}".format(token),
                        })
    finally:
        book.close()
    named_groups = [
        {"family": family, "count": len(items), "items": items}
        for family, items in family_items.items()
    ]
    return {
        "template": str(path), "formulaCount": formula_count, "errors": errors,
        "namedRangeGroups": named_groups,
        "internalNames": internal_items,
        "otherNames": other_items,
    }
