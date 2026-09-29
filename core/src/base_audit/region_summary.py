from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Optional

from .excel_com import ExcelSession
from .feature_log import FeatureLog
from .models import CopyRange
from .models import SourceMatch
from .name_config import (
    FIXED_ROW_SUMMARY_FUNCTION,
    FeatureMapping,
    USED_RANGE_SUMMARY_FUNCTION,
    WORKBOOK_TABLE_MERGE_FUNCTION,
    features_of_type,
    load_feature_mappings,
)
from .template import normalize_template_name
from .history import (
    _composite_key,
    history_append_rule,
)


def _bounds(address: str) -> tuple[int, int, int, int]:
    def cell(value: str) -> tuple[int, int]:
        value = value.replace("$", "").upper()
        letters = "".join(char for char in value if char.isalpha())
        digits = "".join(char for char in value if char.isdigit())
        column = 0
        for letter in letters:
            column = column * 26 + ord(letter) - 64
        return int(digits), column

    first, _, last = address.partition(":")
    top, left = cell(first)
    bottom, right = cell(last or first)
    return min(top, bottom), min(left, right), max(top, bottom), max(left, right)


def _value_matrix(value: object) -> list[list[object]]:
    if value is None:
        return [[""]]
    if not isinstance(value, tuple):
        return [[value]]
    if value and not isinstance(value[0], tuple):
        return [list(value)]
    return [list(row) for row in value]


@dataclass(frozen=True)
class RegionSummaryItem:
    feature_name: str
    row_count: int
    message: str = ""


@dataclass(frozen=True)
class RegionSummaryResult:
    output_path: Path
    items: list[RegionSummaryItem]
    preflight_path: Path | None = None
    skipped_files: tuple[tuple[Path, SourceMatch], ...] = ()
    log_path: Path | None = None

    def summary_text(self) -> str:
        details = "；".join(
            f"{item.feature_name}：{item.row_count} 行" + (f"（{item.message}）" if item.message else "")
            for item in self.items
        )
        lines = [f"汇总完成：{details}", f"输出文件：{self.output_path}"]
        if self.preflight_path:
            lines.append(f"审核前检查：{self.preflight_path}")
        if self.skipped_files:
            names = "、".join(path.name for path, _ in self.skipped_files)
            lines.append(f"已跳过 {len(self.skipped_files)} 个结构不匹配文件：{names}")
        if self.log_path:
            lines.append(f"运行日志：{self.log_path}")
        return "\n".join(lines)


def _output_filename_prefix(flow_name: str | None) -> str:
    """Give the explanation-report flow a business-facing output name."""
    if flow_name == "汇总校验结果说明":
        return "校验结果与报送说明汇总"
    return "区域汇总"


def _source_files(
    input_dir: Path,
    selected_files: list[Path] | None,
    *,
    recursive: bool = True,
) -> list[Path]:
    iterator = input_dir.rglob("*.xlsx") if recursive else input_dir.glob("*.xlsx")
    discovered = sorted(
        path for path in iterator
        if not path.name.startswith("~$") and "_审核版" not in path.stem
    )
    # 说明报送目录常在根目录保留上一版“汇总信息.xlsx”，而机构文件放在
    # 子目录中。存在子目录文件时优先使用它们，避免把历史汇总再次汇总。
    nested_files = [path for path in discovered if path.parent != input_dir]
    all_files = nested_files or discovered if recursive else discovered
    # 区域汇总的源文件常不在审核文件清单中。未选择任何文件时，视为不限制。
    if not selected_files:
        return all_files
    allowed = {path.resolve() for path in all_files}
    return [Path(path).resolve() for path in selected_files if Path(path).resolve() in allowed]


def _header_mapping(feature: FeatureMapping) -> FeatureMapping:
    return FeatureMapping(
        feature.name, feature.feature_type, ("表头区域",),
        feature.workbook_limited, feature.workbook_keyword, feature.remark,
    )


def _intersection_columns(header: CopyRange, data: CopyRange) -> tuple[int, int] | None:
    _, header_left, _, header_right = _bounds(header.address)
    _, data_left, _, data_right = _bounds(data.address)
    left, right = max(header_left, data_left), min(header_right, data_right)
    return (left, right) if left <= right else None


def _output_sheet_name(
    feature_name: str, source_sheets: list[str], used_names: set[str]
) -> str:
    """Use the source sheet name; add a suffix only when Excel requires it."""
    base = source_sheets[0] if source_sheets else feature_name
    base = "".join("_" if char in r"[]:*?/\\" else char for char in base).strip()[:31] or "汇总"
    candidate = base
    suffix = 2
    while candidate.casefold() in used_names:
        marker = f"_{suffix}"
        candidate = base[: 31 - len(marker)] + marker
        suffix += 1
    used_names.add(candidate.casefold())
    return candidate


def _header_signature(values: list[object]) -> tuple[str, ...]:
    return tuple("" if value is None else str(value).strip() for value in values)


def _combine_header_rows(rows: list[list[object]]) -> list[str]:
    """Join each header column from top to bottom, e.g. ``贷款_余额``.

    ``表头区域`` may cover more than one row.  The output table is always
    one row wide, so each source column needs a stable, readable field name.
    Empty cells are ignored here; merged-cell expansion is handled by the COM
    reader before this function is called.
    """
    width = max((len(row) for row in rows), default=0)
    headers: list[str] = []
    for column in range(width):
        parts = []
        for row in rows:
            value = row[column] if column < len(row) else None
            text = "" if value is None else str(value).strip()
            if text:
                parts.append(text)
        headers.append("_".join(parts))
    return headers


def _read_header_values(
    sheet: object,
    header: CopyRange,
    left: int,
    right: int,
) -> list[str]:
    """Read all rows in a named header range and expand merged titles."""
    top, _, bottom, _ = _bounds(header.address)
    matrix = _value_matrix(
        sheet.Range(sheet.Cells(top, left), sheet.Cells(bottom, right)).Value2
    )
    for row_offset, row in enumerate(matrix):
        for column_offset, value in enumerate(row):
            if value not in (None, ""):
                continue
            try:
                cell = sheet.Cells(top + row_offset, left + column_offset)
                if bool(cell.MergeCells):
                    value = cell.MergeArea.Cells(1, 1).Value2
                    if value not in (None, ""):
                        row[column_offset] = value
            except Exception:
                # A non-merged blank is a valid header layout.  Do not make a
                # summary fail merely because an older spreadsheet engine does
                # not expose MergeArea for a particular cell.
                continue
    return _combine_header_rows(matrix)


def _sort_extra_fields(
    fields: list[tuple[str, CopyRange]]
) -> list[tuple[str, CopyRange]]:
    """Keep appended metadata columns stable and visible in worksheet order."""
    return sorted(
        fields,
        key=lambda item: (
            _bounds(item[1].address)[0],  # top to bottom
            _bounds(item[1].address)[1],  # then left to right
            item[0],
        ),
    )


def _field_key(value: object) -> str:
    return "" if value is None else str(value).strip().casefold()


def _metadata_plan(
    global_fields: list[tuple[str, CopyRange]],
    local_fields: list[tuple[str, CopyRange]],
    header_values: list[object],
    sheet_name: str,
) -> tuple[list[tuple[str, CopyRange]], list[str]]:
    """Return metadata fields and names that repeat an existing table header."""
    all_fields = [*global_fields, *local_fields]
    seen: set[str] = set()
    for field_name, _ in all_fields:
        key = _field_key(field_name)
        if key in seen:
            raise ValueError(f"工作表“{sheet_name}”存在重复的单元格区域字段名：{field_name}")
        seen.add(key)
    header_keys = {_field_key(value) for value in header_values}
    return all_fields, [name for name, _ in all_fields if _field_key(name) in header_keys]


def _history_context(
    config_path: Path | None,
    sheet_name: str,
    base_headers: list[str],
):
    """按内置追加规则解析一张输出表的历史追加上下文。

    返回 (查找表, 追加列, 去重列) 或 None（该表不追加）。追加列未在规则中
    指定时，自动取历史表中“非去重列且汇总表头中不存在”的全部列。
    """
    from openpyxl import load_workbook
    from .history import history_append_rule, join_local_validation_history

    if config_path is None or not Path(config_path).is_file():
        return None
    key_fields, history_fields = history_append_rule(sheet_name)
    if not key_fields:
        return None
    key_keys = {_field_key(name) for name in key_fields}
    base_keys = {_field_key(name) for name in base_headers}
    if not key_keys or not key_keys <= base_keys:
        return None
    workbook = load_workbook(config_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return None
        rows = workbook[sheet_name].iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        if not history_fields:
            history_fields = tuple(
                name for name in headers
                if name and _field_key(name).casefold() != "none"
                and _field_key(name) not in key_keys and _field_key(name) not in base_keys
            )
        if not history_fields:
            return None
        lookup = join_local_validation_history(
            headers, [list(row) for row in rows],
            key_fields=key_fields, history_fields=history_fields,
        )
        return lookup, history_fields, key_fields
    finally:
        workbook.close()



def run_region_summaries(
    *,
    template_path: Path,
    input_dir: Path,
    output_dir: Path,
    config_path: Path,
    selected_files: list[Path] | None = None,
    feature_names: tuple[str, ...] | None = None,
    flow_name: str | None = None,
    recursive: bool = True,
    on_step: Optional[Callable[[str], None]] = None,
    named_range_features: tuple[FeatureMapping, ...] = (),
    output_name: str | None = None,
    feature_log: Optional[FeatureLog] = None,
    copies_dir: Path | None = None,
    engine_preference: str = "自动",
    history_config_path: Path | None = None,
    summary_features: tuple[FeatureMapping, ...] = (),
) -> RegionSummaryResult:
    """按命名区域汇总多个源工作簿。

    ``config_path`` 只用于读取模块化功能映射；``summary_features`` 非空时直接
    使用调用方（DAG 汇总节点按 config「区域组」解析）给出的映射，不再读内置
    默认区域，保证检查与汇总使用同一份区域定义。``history_config_path`` 非空
    时在汇总输出上追加历史说明列；DAG 已把历史富化拆为独立节点，因此该参数
    默认为空：汇总节点只产基础列，历史列由 ``summary.history_enrich`` 追加。
    """
    template_path, input_dir, output_dir = template_path.resolve(), input_dir.resolve(), output_dir.resolve()
    if summary_features:
        summaries = [
            feature for feature in summary_features
            if feature.feature_type in {
                USED_RANGE_SUMMARY_FUNCTION, FIXED_ROW_SUMMARY_FUNCTION, WORKBOOK_TABLE_MERGE_FUNCTION,
            }
        ]
    else:
        features = load_feature_mappings(config_path, template_path)
        summaries = [
            *features_of_type(features, USED_RANGE_SUMMARY_FUNCTION),
            *features_of_type(features, FIXED_ROW_SUMMARY_FUNCTION),
            *features_of_type(features, WORKBOOK_TABLE_MERGE_FUNCTION),
        ]
        if feature_names is not None:
            summaries = [feature for feature in summaries if feature.name in feature_names]
    if not summaries:
        raise ValueError("模块化功能中没有“任意行汇总”“固定行汇总”或“汇总表合并”功能")
    if copies_dir is not None:
        # “处理对象”指向某副本目录：直接汇总该目录下的审核副本。
        sources = sorted(copies_dir.glob("*.xlsx"))
    else:
        sources = _source_files(input_dir, selected_files, recursive=recursive)
    if not sources:
        raise ValueError("源数据目录中没有可汇总的 .xlsx 文件")
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{output_name or _output_filename_prefix(flow_name)}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    items: list[RegionSummaryItem] = []

    with ExcelSession(engine_preference) as excel:
        if on_step is not None:
            on_step(f"已连接表格引擎：{excel.engine_name}")
        template = excel.open_workbook(template_path, read_only=True)
        output = excel.excel.Workbooks.Add()
        try:
            for mapping in named_range_features:
                if on_step is not None:
                    on_step(f"正在执行：{mapping.name}")
                if feature_log is not None:
                    survey = excel.survey_named_ranges(template, mapping)
                    feature_log.add_sheet(
                        mapping.name, ("工作表", "命名区域名", "覆盖区域", "结果"), survey
                    )
                    if on_step is not None:
                        found = [row for row in survey if row[3] == "通过"]
                        if found:
                            per_sheet: dict[str, int] = {}
                            for row in found:
                                per_sheet[row[0]] = per_sheet.get(row[0], 0) + 1
                            detail = "、".join(f"{sheet} {count} 处" for sheet, count in per_sheet.items())
                            on_step(f"{mapping.name}：{detail}")
                        else:
                            on_step(f"{mapping.name}：未找到 {mapping.range_names}")
                excel.require_named_ranges(template, mapping)
                if on_step is not None:
                    on_step(f"完成：{mapping.name}")
            # Excel may create 1 or 3 blank sheets according to a user's
            # local Excel setting. Keep one neutral sheet, then create every
            # result sheet explicitly so blank Sheet2/Sheet3 cannot leak into
            # the output or displace a named result sheet.
            while output.Worksheets.Count > 1:
                output.Worksheets(output.Worksheets.Count).Delete()
            first_output_sheet = output.Worksheets(1)
            output_sheet_count = 0
            output_names: set[str] = set()
            history_cache: dict[str, tuple | None] = {}

            def history_context_for(sheet_name: str, base_headers: list[str]):
                if history_config_path is None:
                    return None
                cache_key = sheet_name.casefold()
                if cache_key not in history_cache:
                    history_cache[cache_key] = _history_context(
                        history_config_path, sheet_name, base_headers
                    )
                return history_cache[cache_key]
            for feature in summaries:
                if on_step is not None:
                    on_step(f"正在汇总：{feature.name}")
                if feature.feature_type == WORKBOOK_TABLE_MERGE_FUNCTION:
                    # Deliberately separate from named-range summary: source
                    # sheets are merged only when their actual header rows match.
                    groups: dict[tuple[str, ...], dict[str, object]] = {}
                    for source_path in sources:
                        source = excel.open_workbook(source_path, read_only=True)
                        try:
                            for index in range(1, source.Worksheets.Count + 1):
                                source_sheet = source.Worksheets(index)
                                values = _value_matrix(source_sheet.UsedRange.Value2)
                                values = [row for row in values if any(value not in (None, "") for value in row)]
                                if not values or not any(value not in (None, "") for value in values[0]):
                                    continue
                                header_values = values[0]
                                group = groups.setdefault(_header_signature(header_values), {
                                    "headers": ["来源文件", "来源工作表", *header_values],
                                    "rows": [],
                                    "source_sheets": [],
                                })
                                source_sheets = group["source_sheets"]
                                if source_sheet.Name not in source_sheets:
                                    source_sheets.append(source_sheet.Name)
                                for value_row in values[1:]:
                                    group["rows"].append([source_path.name, source_sheet.Name, *value_row])
                        finally:
                            excel.close_workbook(source)
                    total_rows = 0
                    output_rows: list[tuple[str, int]] = []
                    for group in groups.values():
                        if output_sheet_count == 0:
                            target = first_output_sheet
                        else:
                            target = output.Worksheets.Add(After=output.Worksheets(output.Worksheets.Count))
                        output_sheet_count += 1
                        target.Name = _output_sheet_name(feature.name, group["source_sheets"], output_names)
                        excel._write_table(target, [group["headers"], *group["rows"]], freeze=True)
                        total_rows += len(group["rows"])
                        output_rows.append((target.Name, len(group["rows"])))
                    items.append(RegionSummaryItem(feature.name, total_rows, f"按表头分为 {len(groups)} 张表"))
                    if feature_log is not None:
                        feature_log.add_sheet(
                            feature.name, ("输出工作表", "行数"),
                            output_rows or [("（无汇总行）", 0)],
                        )
                    if on_step is not None:
                        on_step(f"完成：{feature.name}")
                    continue
                data_ranges = excel._named_ranges_for_features(template, [feature])
                headers = excel._named_ranges_for_features(template, [_header_mapping(feature)])
                global_extras = _sort_extra_fields(
                    excel._named_ranges_with_prefix(template, "全局单元格区域.")
                )
                extras_by_sheet: dict[str, list[tuple[str, CopyRange]]] = {}
                for field_name, cell_range in excel._named_ranges_with_prefix(template, "单元格区域."):
                    extras_by_sheet.setdefault(cell_range.sheet_name, []).append((field_name, cell_range))
                for sheet_name, fields in extras_by_sheet.items():
                    extras_by_sheet[sheet_name] = _sort_extra_fields(fields)
                if not data_ranges:
                    raise ValueError(
                        f"汇总功能“{feature.name}”未找到命名区域：" + "、".join(feature.range_names)
                    )
                if not headers:
                    raise ValueError(
                        f"汇总功能“{feature.name}”缺少命名区域“表头区域”，"
                        "请在模板名称管理器中框选表头后补充该名称"
                    )
                # 普通命名区域汇总按来源工作表名称归并。表头是否相同不再
                # 影响合并策略；需要跨表头合并时使用“汇总表合并”。
                groups: dict[str, dict[str, object]] = {}
                repeated_global_fields: list[str] = []
                for data_range in data_ranges:
                    matching_headers = [item for item in headers if item.sheet_name == data_range.sheet_name]
                    if not matching_headers:
                        raise ValueError(
                            f"汇总功能“{feature.name}”在工作表“{data_range.sheet_name}”缺少命名区域“表头区域”，"
                            "请在该工作表中补充表头区域"
                        )
                    candidates = [
                        (item, _intersection_columns(item, data_range))
                        for item in matching_headers
                    ]
                    candidates = [item for item in candidates if item[1] is not None]
                    if not candidates:
                        raise ValueError(
                            f"汇总功能“{feature.name}”中，工作表“{data_range.sheet_name}”的“表头区域”"
                            "与汇总区域没有共同列，请重新框选表头区域"
                        )
                    header, columns = max(
                        candidates, key=lambda item: item[1][1] - item[1][0]
                    )
                    if not columns:
                        continue
                    left, right = columns
                    template_sheet = template.Worksheets(data_range.sheet_name)
                    header_values = _read_header_values(
                        template_sheet, header, left, right
                    )
                    # 全局字段对所有汇总表都有效，包括字段所在的工作表。
                    # 即使名称已在表头中，也单独输出一列，避免覆盖原表数据。
                    extra_fields, _ = _metadata_plan(
                        global_extras,
                        extras_by_sheet.get(data_range.sheet_name, []),
                        header_values,
                        data_range.sheet_name,
                    )
                    for name, _ in global_extras:
                        if _field_key(name) in {_field_key(value) for value in header_values}:
                            label = f"{data_range.sheet_name}!{name}"
                            if label not in repeated_global_fields:
                                repeated_global_fields.append(label)
                    extra_names = [item[0] for item in extra_fields]
                    base_headers = ["来源文件", "来源工作表", *extra_names, *header_values]
                    history_context = history_context_for(data_range.sheet_name, base_headers)
                    expected_headers = (
                        [*base_headers, *history_context[1]]
                        if history_context
                        else base_headers
                    )
                    group = groups.setdefault(data_range.sheet_name.casefold(), {
                        "headers": expected_headers,
                        "rows": [],
                        "source_sheets": [],
                    })
                    source_sheets = group["source_sheets"]
                    if data_range.sheet_name not in source_sheets:
                        source_sheets.append(data_range.sheet_name)
                    data_top, _, data_bottom, _ = _bounds(data_range.address)
                    for source_path in sources:
                        source = excel.open_workbook(source_path, read_only=True)
                        try:
                            try:
                                source_sheet = source.Worksheets(data_range.sheet_name)
                            except Exception:
                                continue
                            if feature.feature_type == USED_RANGE_SUMMARY_FUNCTION:
                                used = source_sheet.UsedRange
                                first_row = max(data_top, int(used.Row))
                                last_row = int(used.Row) + int(used.Rows.Count) - 1
                            else:
                                first_row, last_row = data_top, data_bottom
                            if last_row < first_row:
                                continue
                            values = _value_matrix(source_sheet.Range(source_sheet.Cells(first_row, left), source_sheet.Cells(last_row, right)).Value2)
                            metadata_values: dict[str, object] = {}
                            for field_name, cell_range in extra_fields:
                                try:
                                    metadata_values[_field_key(field_name)] = source.Worksheets(
                                        cell_range.sheet_name
                                    ).Range(cell_range.address).Value2
                                except Exception as exc:
                                    raise ValueError(
                                        f"源文件“{source_path.name}”缺少单元格区域字段“{field_name}”"
                                        f"所在工作表“{cell_range.sheet_name}”"
                                    ) from exc
                            extra_values = [metadata_values[_field_key(name)] for name in extra_names]
                            for value_row in values:
                                if any(value not in (None, "") for value in value_row):
                                    row = [normalize_template_name(source_path.stem), data_range.sheet_name, *extra_values, *value_row]
                                    if history_context:
                                        lookup, _fields, keys = history_context
                                        key = _composite_key(base_headers, row, keys)
                                        history_values = lookup.get(key, [""] * len(_fields))
                                        group["rows"].append([*row, *history_values])
                                    else:
                                        group["rows"].append(row)
                        finally:
                            excel.close_workbook(source)
                total_rows = 0
                output_rows: list[tuple[str, int]] = []
                for group in groups.values():
                    if output_sheet_count == 0:
                        target = first_output_sheet
                    else:
                        target = output.Worksheets.Add(After=output.Worksheets(output.Worksheets.Count))
                    output_sheet_count += 1
                    target.Name = _output_sheet_name(
                        feature.name, group["source_sheets"], output_names
                    )
                    excel._write_table(
                        target, [group["headers"], *group["rows"]], freeze=True
                    )
                    total_rows += len(group["rows"])
                    output_rows.append((target.Name, len(group["rows"])))
                repeat_note = ""
                if repeated_global_fields:
                    repeat_note = "同名全局字段已另列输出：" + "、".join(repeated_global_fields)
                items.append(RegionSummaryItem(feature.name, total_rows, repeat_note))
                if feature_log is not None:
                    feature_log.add_sheet(
                        feature.name, ("输出工作表", "行数"),
                        output_rows or [("（无汇总行）", 0)],
                    )
                if on_step is not None:
                    on_step(f"完成：{feature.name}")
            if output_sheet_count == 0:
                first_output_sheet.Name = "汇总"
            output.SaveAs(str(output_path), FileFormat=51)
        finally:
            excel.close_workbook(output)
            excel.close_workbook(template)
    return RegionSummaryResult(output_path, items)


def enrich_history_columns(
    summary_path: Path, config_path: Path, *, on_step: Optional[Callable[[str], None]] = None
) -> Path:
    """按内置规则把历史说明列追加到汇总工作簿的各张结果表（原地写回）。

    与既有汇总内的历史富化同一口径：对每张结果表，用其表头中的去重列匹配
    历史工作簿的人工作表，取出待追加的人工列并按复合键填入。
    """
    from openpyxl import load_workbook

    say = on_step or (lambda _text: None)
    summary_path = Path(summary_path).resolve()
    config_path = Path(config_path).resolve()
    book = load_workbook(summary_path)
    try:
        history_book = load_workbook(config_path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                headers = [str(cell.value or "").strip() for cell in sheet[1]]
                if not headers or headers == [""]:
                    continue
                context = _history_context(config_path, sheet.title, headers)
                if not context:
                    continue
                lookup, appended, keys = context
                if not appended:
                    continue
                missing = [name for name in appended if name not in headers]
                if not missing:
                    continue  # 该表已有历史列，重复执行不叠加
                key_fields = [str(key) for key in keys]
                if any(name not in headers for name in key_fields):
                    continue  # 汇总表缺少去重列，无法配对，不补空表头
                start = len(headers) + 1
                for offset, name in enumerate(missing):
                    sheet.cell(1, start + offset, name)
                positions = {name: headers.index(name) + 1 for name in key_fields}
                for row_index in range(2, sheet.max_row + 1):
                    identity = tuple(
                        _norm_join(sheet.cell(row_index, positions[name]).value)
                        for name in key_fields
                    )
                    joined = lookup.get(identity)
                    if joined is None:
                        continue
                    for offset, value in enumerate(joined):
                        sheet.cell(row_index, start + offset, value)
                say(f"历史说明富化：{sheet.title}（新增列 {'、'.join(missing)}）")
        finally:
            history_book.close()
        book.save(summary_path)
    finally:
        book.close()
    return summary_path


def _norm_join(value: object) -> str:
    from .history import _norm_key_value

    return _norm_key_value(value)
