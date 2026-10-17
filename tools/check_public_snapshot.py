"""Check the staged public snapshot before committing or pushing it."""

from __future__ import annotations

import subprocess
import sys
import json
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
FORBIDDEN_PARTS = {
    "config-real", "2026-07-31", "test-output", "reference", "dist",
    "runtime", "wheels", "build", "data", "__pycache__",
}
# 各配置册的「使用说明」表由 core/src/base_audit/config_guide.py 按配置版本整体
# 重建，只含静态指引文字，属公开内容；其余表仍必须表头以下为空。若 core 侧表名
# 变更而此处未同步，检查会重新拦截并暴露差异，失败方向安全。
GUIDE_SHEET = "使用说明"


def workbook_privacy_errors(book, *, external_reference: bool = False) -> list[str]:
    """仅登记的外部模板允许第2行表头；首行、保留列和数据区仍须为空。"""
    errors = []
    for sheet in book:
        if sheet.title == GUIDE_SHEET:
            continue
        first_data_row = 2
        if external_reference and sheet.title == "集中系统数据":
            if any(c.value not in (None, "") for row in sheet.iter_rows(max_row=1) for c in row):
                errors.append(f"外部模板首行必须为空：{sheet.title}")
            if any(str(c.value or "").strip() for row in sheet.iter_rows(min_row=2, max_row=2, max_col=2) for c in row):
                errors.append(f"外部模板A2、B2须为空：{sheet.title}")
            first_data_row = 3
        for row in sheet.iter_rows(min_row=first_data_row):
            if any(cell.value not in (None, "") for cell in row):
                errors.append(f"发现表头以下内容：{sheet.title}")
                break
    return errors


def main() -> int:
    raw = subprocess.check_output(["git", "diff", "--cached", "--name-only", "-z"], cwd=ROOT)
    paths = [Path(p.decode("utf-8")) for p in raw.split(b"\0") if p]
    if not paths:
        print("没有暂存文件。")
        return 1
    errors: list[str] = []
    books = 0
    # 清单与工作簿都检查暂存版本，避免工作区改名扩大豁免范围。
    manifest_raw = subprocess.check_output(
        ["git", "show", ":版本管理/VERSION_MANIFEST.json"], cwd=ROOT
    )
    external_file = json.loads(manifest_raw).get("configs", {}).get("S2-CFG-02", {}).get("file")
    for path in paths:
        if any(part.lower() in FORBIDDEN_PARTS for part in path.parts):
            errors.append(f"禁止公开的目录：{path}")
        if path.suffix.lower() in {".csv", ".xls", ".xlsm", ".zip", ".deb", ".exe", ".dll"}:
            errors.append(f"禁止公开的文件类型：{path}")
        if path.suffix.lower() != ".xlsx":
            continue
        if not path.parts or path.parts[0] != "config":
            errors.append(f"工作簿只能位于 config/：{path}")
            continue
        books += 1
        staged = subprocess.check_output(["git", "show", f":{path.as_posix()}"], cwd=ROOT)
        book = load_workbook(BytesIO(staged), read_only=True, data_only=False)
        try:
            errors.extend(
                f"{path} / {message}" for message in workbook_privacy_errors(
                    book, external_reference=path.as_posix() in {
                        external_file,
                        f"config/默认配置/{Path(external_file).name}" if external_file else "",
                    },
                )
            )
        finally:
            book.close()
    if errors:
        print("\n".join(errors))
        return 1
    print(f"公开文件检查通过：{len(paths)} 个暂存文件，{books} 个空表头工作簿。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
